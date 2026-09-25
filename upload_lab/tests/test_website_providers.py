"""Tests for the website provider registry (MIN-69 task 2).

Contract upload.workflow.v1 §4/§6.1: the production catalog exposes exactly
`nam_dinh`; unknown ids are rejected; fake websites are only injectable from
tests through a fresh registry — never into the production list.
"""
from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from batch_scan import connect_registry, get_row_by_file_key, upsert_registry_record
from playwright_uploader import NamDinhUploaderSession
import providers
from providers import (
    ProviderDataDirError,
    UnknownWebsiteError,
    WebsiteProvider,
    WebsiteRegistry,
)

NAM_DINH_URL = "https://congchungnamdinh.ninhbinh.gov.vn"
EXPECTED_CAPABILITIES = [
    "login", "download_export", "audit_excel", "scan", "prepare",
    "staff_options", "reconcile",
]


class _FakeProvider(WebsiteProvider):
    """Provider gia — chi ton tai trong test, khong dang ky production."""

    def __init__(self, website_id: str):
        super().__init__(
            website_id=website_id,
            label=f"Fake {website_id}",
            display_url=f"https://{website_id}.example.test",
            capabilities=("scan",),
            status="available",
        )


def _seed_registry_row(db_dir: Path, *, file_key: str, status: str) -> int:
    source_file = db_dir / f"{file_key}.docx"
    source_file.write_text("dummy", encoding="utf-8")
    conn = connect_registry(db_dir / "registry.sqlite3")
    try:
        upsert_registry_record(
            conn,
            file_key=file_key,
            file_path=source_file,
            stat_result=source_file.stat(),
            customer_folder="KhachA",
            contract_no="1/2026",
            status=status,
            run_id="run_test",
        )
        row = get_row_by_file_key(conn, file_key)
        return int(row["id"])
    finally:
        conn.close()


class ProductionCatalogTests(unittest.TestCase):
    def test_catalog_contains_only_nam_dinh_with_contract_fields(self):
        websites = providers.list_websites()

        self.assertEqual([w["website_id"] for w in websites], ["nam_dinh"])
        entry = websites[0]
        self.assertEqual(entry["label"], "Nam Định")
        self.assertEqual(entry["display_url"], NAM_DINH_URL)
        self.assertEqual(entry["capabilities"], EXPECTED_CAPABILITIES)
        self.assertEqual(entry["status"], "available")

    def test_get_provider_returns_nam_dinh_provider(self):
        provider = providers.get_provider("nam_dinh")
        self.assertEqual(provider.website_id, "nam_dinh")
        self.assertEqual(provider.display_url, NAM_DINH_URL)

    def test_unknown_website_rejected(self):
        for bad in ("khong_co", "NAM_DINH", "namdinh", "../nam_dinh",
                    "nam_dinh/x", "", None, 123):
            with self.assertRaises(UnknownWebsiteError, msg=f"id={bad!r}"):
                providers.get_provider(bad)
        err = None
        try:
            providers.get_provider("khong_co")
        except UnknownWebsiteError as exc:
            err = exc
        self.assertIsNotNone(err)
        self.assertEqual(err.code, "unknown_website")

    def test_catalog_result_is_defensive_copy(self):
        first = providers.list_websites()
        first[0]["website_id"] = "bi_sua"
        first[0]["capabilities"].append("hack")
        again = providers.list_websites()
        self.assertEqual(again[0]["website_id"], "nam_dinh")
        self.assertNotIn("hack", again[0]["capabilities"])

    def test_production_list_has_no_fake_websites(self):
        ids = [w["website_id"] for w in providers.list_websites()]
        self.assertNotIn("fake_a", ids)
        self.assertNotIn("fake_b", ids)
        self.assertEqual(ids, ["nam_dinh"])


class RegistryInjectionTests(unittest.TestCase):
    """Website gia chi inject qua registry rieng cua test."""

    def test_fake_websites_only_via_test_registry(self):
        registry = WebsiteRegistry()
        registry.register(_FakeProvider("fake_a"))
        registry.register(_FakeProvider("fake_b"))

        self.assertEqual(
            [w["website_id"] for w in registry.list_websites()],
            ["fake_a", "fake_b"])
        # production list khong bi anh huong
        self.assertEqual(
            [w["website_id"] for w in providers.list_websites()],
            ["nam_dinh"])
        with self.assertRaises(UnknownWebsiteError):
            providers.get_provider("fake_a")

    def test_same_record_id_two_websites_do_not_share_state(self):
        """Hai website gia cung record_id nhung registry.sqlite3 tach biet:
        trang thai uploaded_success cua web A khong lam web B bo nham."""
        registry = WebsiteRegistry()
        provider_a = _FakeProvider("fake_a")
        provider_b = _FakeProvider("fake_b")
        registry.register(provider_a)
        registry.register(provider_b)

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dir_a = provider_a.ensure_data_layout(
                provider_a.website_data_dir(root))
            dir_b = provider_b.ensure_data_layout(
                provider_b.website_data_dir(root))
            self.assertNotEqual(dir_a, dir_b)

            # Cung mot file_key → moi web mot registry rieng
            id_a = _seed_registry_row(
                dir_a, file_key="same-key", status="uploaded_success")
            id_b = _seed_registry_row(
                dir_b, file_key="same-key", status="extracted")
            self.assertEqual(id_a, id_b)  # cung record_id nhu nhau

            conn_a = connect_registry(dir_a / "registry.sqlite3")
            conn_b = connect_registry(dir_b / "registry.sqlite3")
            try:
                self.assertEqual(
                    get_row_by_file_key(conn_a, "same-key")["status"],
                    "uploaded_success")
                self.assertEqual(
                    get_row_by_file_key(conn_b, "same-key")["status"],
                    "extracted")
            finally:
                conn_a.close()
                conn_b.close()


class NamDinhDataDirTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.data_root = Path(self.tempdir.name) / "upload_lab"
        self.provider = providers.get_provider("nam_dinh")

    def test_website_data_dir_uses_websites_subdir(self):
        data_dir = self.provider.website_data_dir(self.data_root)
        self.assertEqual(data_dir.parent.name, "websites")
        self.assertEqual(data_dir.name, "nam_dinh")
        self.assertEqual(
            data_dir, self.data_root / "websites" / "nam_dinh")

    def test_ensure_data_layout_creates_contract_subdirs(self):
        data_dir = self.provider.ensure_data_layout(
            self.provider.website_data_dir(self.data_root))
        for sub in ("output", "runs", "downloads", "upload_runs", "logs"):
            self.assertTrue((data_dir / sub).is_dir(), f"thieu {sub}")
        # idempotent
        again = self.provider.ensure_data_layout(data_dir)
        self.assertEqual(again, data_dir)

    def test_assert_data_dir_rejects_other_website_dir(self):
        wrong = self.data_root / "websites" / "khac"
        with self.assertRaises(ProviderDataDirError):
            self.provider.assert_data_dir(wrong)
        flat = self.data_root / "nam_dinh"  # thieu tang websites/
        with self.assertRaises(ProviderDataDirError):
            self.provider.assert_data_dir(flat)

    def test_create_browser_uses_website_data_dir(self):
        data_dir = self.provider.website_data_dir(self.data_root)
        session = self.provider.create_browser(data_dir)

        self.assertIsInstance(session, NamDinhUploaderSession)
        self.assertEqual(session.working_dir, data_dir.resolve())
        self.assertTrue(
            str(session.settings.storage_state_path).startswith(
                str(data_dir.resolve())))
        self.assertEqual(
            session.staff_options_cache_path.parent, data_dir.resolve())

    def test_create_browser_reads_per_website_env(self):
        data_dir = self.provider.ensure_data_layout(
            self.provider.website_data_dir(self.data_root))
        (data_dir / ".env").write_text(
            "ND_BASE_URL=https://congchungnamdinh.ninhbinh.gov.vn\n"
            "ND_MAX_PREPARED_TABS=7\n",
            encoding="utf-8")
        session = self.provider.create_browser(data_dir)
        self.assertEqual(session.settings.max_prepared_tabs, 7)
        self.assertEqual(session.settings.base_url, NAM_DINH_URL)

    def test_create_browser_rejects_wrong_dir(self):
        with self.assertRaises(ProviderDataDirError):
            self.provider.create_browser(self.data_root / "websites" / "khac")

    def test_public_config_returns_only_schema_values(self):
        data_dir = self.provider.ensure_data_layout(
            self.provider.website_data_dir(self.data_root))
        (data_dir / ".env").write_text(
            "ND_BASE_URL=https://congchungnamdinh.ninhbinh.gov.vn\n"
            "ND_MAX_PREPARED_TABS=15\n"
            "MAT_KHAU_BI_MAT=khong-duoc-lo\n"
            "ND_CUSTOM_EXTRA=abc\n",
            encoding="utf-8")
        config = self.provider.public_config(data_dir)
        self.assertLessEqual(
            set(config),
            {"website_id", "base_url", "browser_channel",
             "max_prepared_tabs", "storage_state_exists", "env_exists"})
        self.assertEqual(config["base_url"], NAM_DINH_URL)
        self.assertEqual(config["max_prepared_tabs"], 15)
        self.assertFalse(config["storage_state_exists"])
        self.assertTrue(config["env_exists"])
        # khong ro tri gia tri nhay cam / key ngoai schema
        self.assertNotIn("MAT_KHAU_BI_MAT", json.dumps(config))
        self.assertNotIn("ND_CUSTOM_EXTRA", json.dumps(config))

    def test_staff_options_cache_is_per_website(self):
        data_dir = self.provider.ensure_data_layout(
            self.provider.website_data_dir(self.data_root))
        (data_dir / "uploader_staff_options.json").write_text(
            json.dumps({"cong_chung_vien": ["CCV A"], "thu_ky": ["TK B"]},
                       ensure_ascii=False),
            encoding="utf-8")
        cached = self.provider.load_staff_options_cache(data_dir)
        self.assertEqual(cached["cong_chung_vien"], ["CCV A"])
        self.assertEqual(cached["thu_ky"], ["TK B"])


class NamDinhAuditTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.root = Path(self.tempdir.name)
        self.provider = providers.get_provider("nam_dinh")

    def _make_book(self, rows) -> Path:
        path = self.root / "book.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "SO CONG CHUNG"
        sheet["B1"] = "NGAY, THANG, NAM CONG CHUNG"
        for index, (contract_no, date_value) in enumerate(rows, start=2):
            sheet.cell(row=index, column=1).value = contract_no
            sheet.cell(row=index, column=2).value = date_value
        workbook.save(path)
        workbook.close()
        return path

    def test_audit_excel_returns_common_analysis_structure(self):
        book = self._make_book([
            ("1/2026", "01/01/2026"),
            ("3/2026", "02/01/2026"),
            ("abc", "03/01/2026"),
        ])
        analysis = self.provider.audit_excel(
            book, from_date="01/01/2026", to_date="31/12/2026")

        self.assertEqual(analysis.summary.excel_total, 3)
        self.assertEqual(analysis.summary.valid_count, 2)
        self.assertEqual(
            [m.contract_no for m in analysis.missing_numbers], ["2/2026"])
        self.assertEqual(len(analysis.issue_rows), 1)
        # tap so hop le phuc vu doi chieu sau nay
        self.assertTrue(hasattr(analysis, "valid_rows"))


if __name__ == "__main__":
    unittest.main()
