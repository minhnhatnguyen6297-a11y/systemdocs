"""Fictitious, native payload regressions for TASK-10.3 source shapes."""
import tempfile
import unittest
from pathlib import Path
from docx import Document
from extraction_adapter import extract
from regex_workbench.engine import contract_extract as engine


class BoundedContractTests(unittest.TestCase):
    def native(self, text, tables=()):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'fixture.docx'
            doc = Document()
            for line in text.splitlines():
                doc.add_paragraph(line)
            for rows in tables:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for row, cells in zip(table.rows, rows):
                    for cell, value in zip(row.cells, cells):
                        cell.text = value
            doc.save(path)
            return extract(path)['raw']

    def test_representative_shapes_keep_sides_and_addresses(self):
        for label in ['Đại diện: ông', 'Đại diện ông', ' | | Đại diện: ông']:
            text = f'''HỢP ĐỒNG THẾ CHẤP
BÊN NHẬN THẾ CHẤP:
Ngân hàng Mẫu
{label} Lê Văn Mẫu – Phó Giám đốc CN; CCCD số 012345678901 do Cơ quan Mẫu cấp ngày 01/02/2024.
 | | BÊN THẾ CHẤP:
Ông Trần Văn Ví Dụ; Sinh năm 1980; CCCD số 012345678902 do Cơ quan Hai cấp ngày 02/03/2024.
Bà Nguyễn Thị Ví Dụ; Sinh năm 1981; CCCD số 012345678903 do Cơ quan Hai cấp ngày 02/03/2024.
Cả hai cùng thường trú tại: phố Mẫu, tỉnh Mẫu.
ĐIỀU 1: Nội dung
Theo uỷ quyền của ông Người Trong Câu Kể; CCCD số 012345678904.
LỜI CHỨNG'''
            with self.subTest(label=label):
                raw = self.native(text)
                self.assertEqual([p['ho_ten'] for p in raw['ben_a']['nguoi']], ['Trần Văn Ví Dụ', 'Nguyễn Thị Ví Dụ'])
                representative, = raw['ben_b']['nguoi']
                self.assertEqual((representative['ho_ten'], representative['cccd'], representative['noi_cap'], representative['dia_chi']),
                                 ('Lê Văn Mẫu', '012345678901', 'Cơ quan Mẫu', ''))
                self.assertTrue(all(p['dia_chi'] == 'phố Mẫu, tỉnh Mẫu' for p in raw['ben_a']['nguoi']))

    def test_certificate_cannot_supply_person_identity(self):
        for token in ['012345678901do Bộ Coongan cấp ngày 01/02/2024.', '012345678901.']:
            raw = self.native(f'''GIẤY UỶ QUYỀN
Tôi là
Ông Lê Văn Mẫu; Sinh năm 1980;
CCCD số {token}
Nơi cư trú: Nhà Mẫu.
Tôi là chủ sử dụng đất:
1. Thửa đất có địa chỉ tại: Nhà đất Mẫu theo Giấy chứng nhận quyền sử dụng đất số AA 123456; Số vào sổ cấp GCN: VP123 do Cơ quan Đất cấp ngày 03/04/2025.
Nay bằng giấy này tôi uỷ quyền cho:
Ông Trần Văn Mẫu; Sinh năm 1990; CCCD số 012345678902.
LỜI CHỨNG''')
            person, = raw['ben_a']['nguoi']
            self.assertEqual(person['cccd'], '012345678901')
            self.assertNotEqual(person['noi_cap'], 'Cơ quan Đất')
            self.assertNotEqual(person['ngay_cap_cccd'], '03/04/2025')

    def test_bulleted_asset_issuance_cannot_supply_person_identity(self):
        for bullet in ['-', '*', '•', '–']:
            with self.subTest(bullet=bullet):
                raw = self.native(f'''GIẤY UỶ QUYỀN
Ông Lê Văn Mẫu; Sinh năm 1980; CCCD số 012345678901.
{bullet} Thửa đất số: 11; do Cơ quan Đất cấp ngày 03/04/2025.
LỜI CHỨNG''')
                person, = raw['ben_a']['nguoi']
                self.assertEqual((person['noi_cap'], person['ngay_cap_cccd']), ('', ''))
                self.assertIn('do Cơ quan Đất cấp ngày 03/04/2025', person['raw_text'])

    def test_mixed_case_items_retain_shared_certificate_without_overriding_own(self):
        for own in ['', ' theo Giấy chứng nhận quyền sử dụng đất số BB 654321; Số vào sổ cấp GCN: OWN2; do Cơ quan Riêng cấp ngày 04/05/2025.']:
            raw = self.native(f'''HỢP ĐỒNG CHUYỂN NHƯỢNG
BÊN A:
Ông Người Mẫu; CCCD số 012345678901.
ĐIỀU 1: Đối tượng
Các thửa dưới đây theo Giấy chứng nhận quyền sử dụng đất số AA 123456; Số vào sổ cấp GCN: SHARED1; do Cơ quan Chung cấp ngày 01/02/2024.
1. thửa đất có địa chỉ tại: Địa chỉ Một
2. THỬA ĐẤT có địa chỉ tại: Địa chỉ Hai{own}
ĐIỀU 2: Nội dung
LỜI CHỨNG''')
            rows = raw['danh_sach_thua_dat']
            self.assertEqual(len(rows), 2, 'shared prefix must not create a phantom parcel')
            self.assertEqual((rows[0]['dia_chi'], rows[0]['so_gcn'], rows[0]['so_vao_so'], rows[0]['co_quan_cap']),
                             ('Địa chỉ Một', 'AA123456', 'SHARED1', 'Cơ quan Chung'))
            self.assertEqual((rows[1]['dia_chi'], rows[1]['so_gcn'], rows[1]['so_vao_so'], rows[1]['co_quan_cap']),
                             ('Địa chỉ Hai', 'BB654321' if own else 'AA123456', 'OWN2' if own else 'SHARED1', 'Cơ quan Riêng' if own else 'Cơ quan Chung'))

    def test_entry_book_field_delimiters_are_read_natively(self):
        for suffix in [';', ',', '; do Cơ quan Mẫu cấp ngày 01/02/2024.']:
            raw = self.native(f'''GIẤY UỶ QUYỀN
1. Thửa đất có địa chỉ tại: Nhà đất Mẫu theo Giấy chứng nhận quyền sử dụng đất số AA 123456; Số vào sổ cấp GCN: VP123{suffix}
Nay bằng giấy này tôi uỷ quyền cho:
Ông Người Nhận Mẫu; CCCD số 012345678901.
LỜI CHỨNG''')
            row, = raw['danh_sach_thua_dat']
            self.assertEqual(row['so_vao_so'], 'VP123')

    def test_two_and_three_certificate_items_are_separate(self):
        for count in [2, 3]:
            items = '\n'.join(f'{n}. Thửa đất có địa chỉ tại: Địa chỉ {n} theo Giấy chứng nhận quyền sử dụng đất số AA {123450+n}; Số vào sổ cấp giấy chứng nhận số: VP{n} do Cơ quan {n} cấp ngày 0{n}/04/2025.' for n in range(1, count + 1))
            raw = self.native(f'GIẤY UỶ QUYỀN\n{items}\nNay bằng giấy này tôi uỷ quyền cho:\nÔng Người Nhận Mẫu; CCCD số 012345678901.\nĐịa chỉ: Nhà người nhận.\nLỜI CHỨNG')
            rows = raw['danh_sach_thua_dat']
            self.assertEqual(len(rows), count)
            for n, row in enumerate(rows, 1):
                self.assertEqual((row['so_thua_dat'], row['so_to_ban_do'], row['dia_chi'], row['so_vao_so'], row['co_quan_cap']), ('', '', f'Địa chỉ {n}', f'VP{n}', f'Cơ quan {n}'))

    def test_commitment_title_and_address_before_parcel(self):
        body = '''Bà Lê Thị Mẫu là chủ sử dụng 01 thửa đất, cụ thể như sau:
* Thửa đất: Theo Giấy chứng nhận quyền sử dụng đất số AA 123456; Vào sổ cấp giấy chứng nhận: PV123 do Cơ quan Đất cấp ngày 03/04/2025.
- Địa chỉ thửa đất: Phố Mẫu, tỉnh Mẫu
- Thửa đất số: 11 | - Tờ bản đồ số: 205; | Diện tích: 360,0 m2;
- Loại đất: Đất ở; Thời hạn sử dụng: Lâu dài;
1. Tôi – Người Mẫu cam kết thửa đất trên là tài sản riêng.
LỜI CHỨNG
Tại Văn phòng Mẫu, địa chỉ: Địa chỉ Văn phòng.'''
        raw = self.native('VĂN BẢN CAM KẾT\nVỀ TÀI SẢN RIÊNG CỦA VỢ CHỒNG\n' + body)
        self.assertEqual(raw['document_kind'], engine.DOC_KIND_ASSET_COMMITMENT)
        row, = raw['danh_sach_thua_dat']
        self.assertEqual((row['so_thua_dat'], row['so_to_ban_do'], row['dia_chi'], row['ngay_cap_gcn']), ('11', '205', 'Phố Mẫu, tỉnh Mẫu', '03/04/2025'))
        moved = body.replace('- Địa chỉ thửa đất: Phố Mẫu, tỉnh Mẫu\n', '').replace('- Loại đất:', '- Địa chỉ thửa đất: Phố Mẫu, tỉnh Mẫu\n- Loại đất:')
        self.assertEqual(engine.tach_danh_sach_thua_dat(body), engine.tach_danh_sach_thua_dat(moved))
        negative = self.native('VĂN BẢN CAM KẾT\nÔng Người Mẫu; CCCD số 012345678901 do Cơ quan Người cấp ngày 01/02/2024.\nLỜI CHỨNG\nĐịa chỉ: Văn phòng Mẫu.')
        self.assertEqual(negative['danh_sach_thua_dat'], [])

    def test_land_empty_table_cells_keep_their_position(self):
        rows = engine.tach_danh_sach_thua_dat(engine._normalize_plain_text_for_extract('Số thửa\tSố tờ\tDiện tích\tMục đích\tThời hạn\n11\t\t360\tĐất ở\tLâu dài'))
        self.assertEqual((rows[0]['so_thua_dat'], rows[0]['so_to_ban_do'], rows[0]['dien_tich']), ('11', '', '360'))

    def test_native_docx_table_inline_representative_and_side_labels(self):
        rows = [['BÊN NHẬN THẾ CHẤP:', 'Ngân hàng Mẫu', 'Đại diện: ông Lê Văn Mẫu – Phó Giám đốc CN; CCCD số 012345678901 do Cơ quan Mẫu cấp ngày 01/02/2024.'],
                ['', 'BÊN THẾ CHẤP:', 'Ông Trần Văn Ví Dụ; Sinh năm 1980; CCCD số 012345678902. Nơi thường trú tại: Nhà Mẫu.'],
                ['ĐIỀU 1: Nội dung', '', '']]
        raw = self.native('HỢP ĐỒNG THẾ CHẤP', [rows])
        person, = raw['ben_a']['nguoi']
        representative, = raw['ben_b']['nguoi']
        self.assertEqual((person['ho_ten'], person['dia_chi']), ('Trần Văn Ví Dụ', 'Nhà Mẫu'))
        self.assertEqual((representative['ho_ten'], representative['cccd'], representative['dia_chi']), ('Lê Văn Mẫu', '012345678901', ''))

    def test_native_docx_compact_details_preserve_blank_term_and_aligned_cells(self):
        for term in ('Lâu dài', ''):
            with self.subTest(term=term):
                table = [['Số tờ bản đồ', 'Số thửa', 'Diện tích', 'Mục đích', 'Thời hạn'],
                         ['4', '169', '427', '', ''], ['Trong đó', '', '', '', ''],
                         ['200', 'Đất ở', term, '', ''], ['', '', '227', 'Đất vườn', ''],
                         ['5', '170', '100', 'Đất ở', 'Lâu dài']]
                raw = self.native('HỢP ĐỒNG CHUYỂN NHƯỢNG\nĐIỀU 1: Đối tượng', [table])
                first, second = raw['danh_sach_thua_dat']
                self.assertEqual((first['so_thua_dat'], second['so_thua_dat']), ('169', '170'))
                details = first['chi_tiet_dien_tich']
                self.assertEqual(len(details), 2)
                self.assertEqual((details[0]['dien_tich'], details[0]['muc_dich_su_dung'], details[0]['thoi_han_su_dung']), ('200', 'Đất ở', term))
                self.assertEqual((details[1]['dien_tich'], details[1]['muc_dich_su_dung']), ('227', 'Đất vườn'))
                self.assertFalse(any(d.get('so_thua_dat') or d.get('so_to_ban_do') for d in details))

    def test_native_docx_land_table_empty_cells_keep_their_position(self):
        rows = [['Số thửa', 'Số tờ', 'Diện tích', 'Mục đích', 'Thời hạn'],
                ['11', '', '360', 'Đất ở', 'Lâu dài'],
                ['12', '5', '100', '', '2040']]
        raw = self.native('HỢP ĐỒNG CHUYỂN NHƯỢNG\nĐIỀU 1: Đối tượng của hợp đồng', [rows])
        first, second = raw['danh_sach_thua_dat']
        self.assertEqual((first['so_thua_dat'], first['so_to_ban_do'], first['dien_tich']), ('11', '', '360'))
        self.assertEqual((second['so_thua_dat'], second['so_to_ban_do'], second['muc_dich_su_dung'], second['thoi_han_su_dung']), ('12', '5', '', '2040'))
