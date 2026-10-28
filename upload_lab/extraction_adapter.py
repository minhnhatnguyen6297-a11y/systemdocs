"""Ranh consumer của lõi luật Word: đọc file (.docx/.doc) rồi gọi lõi.

Luật trích xuất nằm ở `tools/regex_workbench/engine/contract_extract.py` —
một đường luật duy nhất cho Document Intake (Notary), Upload Lab và tool
nhân viên. Adapter này chỉ giữ phần biên:

- đọc file -> text (python-docx cho `.docx`, Windows IFilter cho `.doc`);
- chuẩn hóa text `.doc` (TAB -> ` | `) trước khi vào luật;
- quét số công chứng theo file (scan/lọc trước khi bóc);
- giữ payload `{case, web_form, raw}`; Case là metadata canonical,
  web_form chỉ presentation/portal projection, không nguồn Case thứ hai;
- CLI `python -m ...` để bóc 1 file ra JSON.

Không chép luật, không giữ bản engine thứ hai; sửa luật thì sửa ở lõi.
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path
from typing import Iterable

_HERE = Path(__file__).resolve()
# Lõi nằm cạnh adapter theo từng layout: repo (`upload_lab/../tools`), bản
# đóng gói cạnh sidecar/tool (`<root>/tools` hoặc `<root>/upload_lab/tools`).
for _candidate in (_HERE.parent / "tools", _HERE.parents[1] / "tools"):
    if _candidate.is_dir() and str(_candidate) not in sys.path:
        sys.path.insert(0, str(_candidate))

from regex_workbench.engine import contract_extract as core  # noqa: E402
from regex_workbench.engine.contract_extract import (  # noqa: E402,F401
    ENGINE_ID,
    ENGINE_VERSION,
    DEFAULT_CCV,
    get_missing_web_form_fields,
    case_from_payload,
    case_portal_projection,
    validate_case_fields,
)


_ASSET_RENDER_FIELDS = (
    ("Số serial", "so_gcn_raw", "so_gcn_raw_values"),
    ("Số vào sổ", "so_vao_so", "so_vao_so_values"),
    ("Loại giấy", "loai_gcn", "loai_gcn_raw_values"),
    ("Thửa đất số", "so_thua_dat", "so_thua_dat_values"),
    ("Tờ bản đồ số", "so_to_ban_do", "so_to_ban_do_values"),
    ("Địa chỉ", "dia_chi", "dia_chi_values"),
    ("Ngày cấp", "ngay_cap_gcn", "ngay_cap_values"),
    ("Cơ quan cấp", "co_quan_cap", "co_quan_cap_values"),
    ("Nguồn gốc", "nguon_goc", "nguon_goc_values"),
    ("Hình thức sử dụng", "hinh_thuc_su_dung", "hinh_thuc_su_dung_values"),
)


def _uniq(values) -> list[str]:
    return list(dict.fromkeys(str(value).strip() for value in values or [] if str(value).strip()))


def _asset_value(parcel: dict, key: str, values_key: str) -> str:
    values = _uniq(parcel.get(values_key))
    if len(values) > 1:
        return ""
    return str(parcel.get(key) or (values[0] if values else "")).strip()


def _source_total(parcel: dict) -> str:
    totals = _uniq(parcel.get("dien_tich_tong_values"))
    return totals[0] if len(totals) == 1 else ""


def _detail_value(detail: dict, key: str, values_key: str) -> str:
    values = _uniq(detail.get(values_key))
    if len(values) > 1:
        return ""
    return str(detail.get(key) or (values[0] if values else "") or "").strip()


def _detail_term(detail: dict) -> str:
    return _detail_value(detail, "thoi_han_su_dung", "thoi_han_su_dung_values")


def _render_land_row(land_type: str, area: str, term: str) -> str:
    parts = [
        f"Loại đất: {land_type or '—'}",
        f"Diện tích: {area or '—'}",
        f"Thời hạn: {term or '—'}",
    ]
    return "; ".join(parts)


def _render_parcel(parcel: dict) -> list[str]:
    lines = []
    for label, key, values_key in _ASSET_RENDER_FIELDS:
        value = _asset_value(parcel, key, values_key)
        if value:
            lines.append(f"{label}: {value}")
    total = _source_total(parcel)
    if total:
        lines.append(f"Diện tích: {total}")
    details = [detail for detail in parcel.get("chi_tiet_dien_tich") or []
               if isinstance(detail, dict) and not detail.get("o_tho")]
    # Rows are already typed/flat-classified by the shared parser; do not
    # dedupe a known detail row by equality with a total value.

    if details:
        for detail in details:
            lines.append(_render_land_row(
                _detail_value(detail, "muc_dich_su_dung", "muc_dich_su_dung_values"),
                _detail_value(detail, "dien_tich", "dien_tich_values"),
                _detail_term(detail)))
    return lines


def _render_group(group: dict) -> list[str]:
    lines = []
    fields = (
        ("Số serial", "so_serial", "so_serial_values"),
        ("Số vào sổ", "so_vao_so", "so_vao_so_values"),
        ("Loại giấy", "loai_so", "loai_so_values"),
        ("Thửa đất số", "so_thua_dat", "so_thua_dat_values"),
        ("Tờ bản đồ số", "so_to_ban_do", "so_to_ban_do_values"),
        ("Địa chỉ", "dia_chi", "dia_chi_values"),
        ("Ngày cấp", "ngay_cap", "ngay_cap_values"),
        ("Cơ quan cấp", "co_quan_cap", "co_quan_cap_values"),
        ("Nguồn gốc", "nguon_goc", "nguon_goc_values"),
        ("Hình thức sử dụng", "hinh_thuc_su_dung", "hinh_thuc_su_dung_values"),
    )
    for label, key, values_key in fields:
        value = _asset_value(group, key, values_key)
        if value:
            lines.append(f"{label}: {value}")
    totals = _uniq(group.get("dien_tich_tong_values"))
    if len(totals) == 1:
        lines.append(f"Diện tích: {totals[0]}")
    rows = [row for row in group.get("land_rows") or [] if isinstance(row, dict)]
    for row in rows:
        lines.append(_render_land_row(
            str(row.get("loai_dat") or "").strip(),
            str(row.get("dien_tich") or "").strip(),
            str(row.get("thoi_han") or "").strip()))
    return lines


def render_asset_text(payload: dict) -> str:
    """Render asset editor text from approved canonical fields only.

    A09-A: no residual/notes/raw block. Ambiguous values (including total)
    are omitted; grouped asset shapes are not dropped; blank land-row positions
    stay represented by placeholders. Unsupported/no structured asset returns
    empty rather than raw fallback.
    """
    source = str((payload.get("web_form") or {}).get("tai_san") or "")
    groups = core.tach_nhom_tai_san(source, land_rows_as_details=True)
    blocks: list[str] = []
    if groups:
        # Grouped source owns its windows; do not mix/duplicate raw parcels.
        for group in groups:
            lines = _render_group(group)
            if lines:
                blocks.append("\n".join(lines))
    else:
        # Reparse the preserved source in detailed mode so historical payloads
        # replay the same typed rows/totals as fresh extraction; never trust
        # legacy raw parcel scalar fields as canonical detail.
        for parcel in core.tach_danh_sach_thua_dat(source, land_rows_as_details=True):
            lines = _render_parcel(parcel)
            if lines:
                blocks.append("\n".join(lines))
    return "\n\n".join(blocks).strip()


# ============================================================
# DOCX / DOC HELPERS
# ============================================================
def _append_text_lines(lines: list[str], values: Iterable[str]) -> None:
    for raw in values:
        text = str(raw or "").strip()
        if text:
            lines.append(text)


def _append_table_lines(lines: list[str], tables) -> None:
    for table in tables:
        for row in table.rows:
            row_values = []
            seen_cells = set()
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue  # A merged physical cell is exposed more than once by python-docx.
                seen_cells.add(cell._tc)
                cell_text = " ".join(
                    p.text.strip()
                    for p in cell.paragraphs
                    if p.text and p.text.strip()
                ).strip()
                row_values.append(cell_text)
            if any(row_values):
                lines.append(" | ".join(row_values))


def _append_docx_body_lines(lines: list[str], doc) -> None:
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            _append_text_lines(lines, [Paragraph(child, doc).text])
            continue
        if isinstance(child, CT_Tbl):
            _append_table_lines(lines, [Table(child, doc)])


def read_doc_via_ifilter(filepath) -> str:
    """Trich xuat text tu .doc bang Windows IFilter (cung co che voi Agent Ransack).
    WHY: Doc .doc tren Windows ma khong can Word/Office.
         Dung cho scan (tim so hop dong) — chi can plain text, khong can table structure.
    RISK: IFilter tra plain text thuan — mat cau truc bang va header/footer.
          Flow extract() cho .doc se chay theo plain text nay.
    REQUIRE: query.dll (co san tren Windows 7+, khong can cai them).
    """
    import ctypes

    abs_path = str(Path(filepath).resolve())
    ole32 = ctypes.windll.ole32
    ole32.CoInitialize(None)
    try:
        query = ctypes.windll.query
        # c_long (signed) de compare HRESULT am — WINFUNCTYPE(HRESULT) tu raise exception
        query.LoadIFilter.restype = ctypes.c_long
        query.LoadIFilter.argtypes = [
            ctypes.c_wchar_p,
            ctypes.c_void_p,
            ctypes.POINTER(ctypes.c_void_p),
        ]
        p_filter = ctypes.c_void_p()
        hr = query.LoadIFilter(abs_path, None, ctypes.byref(p_filter))
        if hr != 0 or not p_filter.value:
            raise OSError(f"LoadIFilter that bai (hr=0x{hr & 0xFFFFFFFF:08X})")

        # IFilter VTBL layout (IUnknown 0-2, IFilter 3-7):
        #   3=Init, 4=GetChunk, 5=GetText, 6=GetValue, 7=BindRegion
        VTBL_INIT = 3
        VTBL_GET_CHUNK = 4
        VTBL_GET_TEXT = 5
        VTBL_RELEASE = 2

        vtbl_ptr = ctypes.cast(p_filter, ctypes.POINTER(ctypes.c_void_p))
        vtbl = ctypes.cast(vtbl_ptr[0], ctypes.POINTER(ctypes.c_void_p))

        # Init(grfFlags, cAttributes, aAttributes, pFlags) — c_long tranh auto-raise
        InitFunc = ctypes.WINFUNCTYPE(
            ctypes.c_long, ctypes.c_void_p,
            ctypes.c_uint, ctypes.c_ulong, ctypes.c_void_p,
            ctypes.POINTER(ctypes.c_uint),
        )
        out_flags = ctypes.c_uint(0)
        InitFunc(vtbl[VTBL_INIT])(p_filter, 0, 0, None, ctypes.byref(out_flags))

        # STAT_CHUNK tren x64 Windows:
        #   offset  0: idChunk   (ULONG 4B)
        #   offset  4: breakType (enum  4B)
        #   offset  8: flags     (CHUNKSTATE 4B) — CHUNK_TEXT=0x1, CHUNK_VALUE=0x2
        #   offset 12: locale    (LCID  4B)
        #   offset 16: attribute (FULLPROPSPEC = GUID(16B) + PROPSPEC(4+4pad+8B) = 32B)
        #   offset 48: idChunkSource (4B), cwcStartSource (4B), cwcLenSource (4B)
        #   total: 60B — dung 64B buffer cho an toan tren moi alignment
        CHUNK_BUF_SIZE = 64
        CHUNK_FLAGS_OFFSET = 8
        CHUNK_TEXT = 0x1  # CHUNKSTATE: CHUNK_TEXT=1 (khong phai 2)

        GetChunkFunc = ctypes.WINFUNCTYPE(
            ctypes.c_long, ctypes.c_void_p, ctypes.c_char_p,
        )
        get_chunk_fn = GetChunkFunc(vtbl[VTBL_GET_CHUNK])

        BUF_SIZE = 4096
        GetTextFunc = ctypes.WINFUNCTYPE(
            ctypes.c_long, ctypes.c_void_p,
            ctypes.POINTER(ctypes.c_ulong), ctypes.c_wchar_p,
        )
        get_text_fn = GetTextFunc(vtbl[VTBL_GET_TEXT])

        ReleaseFunc = ctypes.WINFUNCTYPE(ctypes.c_ulong, ctypes.c_void_p)
        release_fn = ReleaseFunc(vtbl[VTBL_RELEASE])

        # HRESULT codes (signed 32-bit)
        FILTER_E_END_OF_CHUNKS = ctypes.c_long(0x80041700).value  # negative
        FILTER_E_NO_MORE_TEXT  = ctypes.c_long(0x80041701).value  # negative
        # FILTER_S_LAST_TEXT = 0x00041709 — SUCCESS: text duoc ghi VA day la phan cuoi cung
        # Phai capture text TRUOC khi check status nay, roi break
        FILTER_S_LAST_TEXT = 0x00041709  # positive (success with info)

        parts: list[str] = []
        chunk_buf = ctypes.create_string_buffer(CHUNK_BUF_SIZE)

        while True:
            hr_chunk = get_chunk_fn(p_filter, chunk_buf)
            if hr_chunk == FILTER_E_END_OF_CHUNKS:
                break
            if hr_chunk != 0:
                break

            flags = ctypes.c_uint.from_buffer_copy(chunk_buf, CHUNK_FLAGS_OFFSET).value
            if not (flags & CHUNK_TEXT):
                continue

            while True:
                buf_len = ctypes.c_ulong(BUF_SIZE)
                text_buf = ctypes.create_unicode_buffer(BUF_SIZE)
                hr_text = get_text_fn(p_filter, ctypes.byref(buf_len), text_buf)
                # Capture text TRUOC khi check status — FILTER_S_LAST_TEXT tra text hop le
                if buf_len.value > 0:
                    parts.append(text_buf.value[: buf_len.value])
                if hr_text == FILTER_E_NO_MORE_TEXT:
                    break
                if hr_text == FILTER_S_LAST_TEXT:
                    break  # text da duoc lay, day la phan cuoi cua chunk nay
                if hr_text != 0:
                    break

        release_fn(p_filter)
        return "".join(parts)
    finally:
        ole32.CoUninitialize()


def read_docx(filepath, *, use_ifilter_for_doc: bool = False):
    """Doc .docx hoac .doc -> text thuan.
    - .docx: dung python-docx truc tiep (nhanh).
    - .doc: dung IFilter/plain text tren Windows, khong can Word.
    """
    path = Path(filepath)
    if path.suffix.lower() == ".doc":
        return read_doc_via_ifilter(filepath)

    from docx import Document
    doc = Document(str(filepath))
    lines: list[str] = []
    _append_docx_body_lines(lines, doc)
    for section in doc.sections:
        _append_text_lines(lines, (p.text for p in section.header.paragraphs))
        _append_text_lines(lines, (p.text for p in section.footer.paragraphs))
    return "\n".join(lines)


def scan_docx_for_contract_no(filepath, *, include_text: bool = False):
    """Tim so cong chung trong noi dung file .docx hoac .doc.
    .doc dung IFilter (nhanh) — chi can plain text de tim regex so cong chung.
    """
    path = Path(filepath)
    use_ifilter_for_doc = path.suffix.lower() == ".doc"
    try:
        text = read_docx(path, use_ifilter_for_doc=use_ifilter_for_doc)
        if use_ifilter_for_doc:
            text = core._normalize_plain_text_for_extract(text)
    except Exception as exc:
        result = {
            "is_contract": False,
            "contract_no": "",
            "reason": f"Khong doc duoc file: {exc}",
        }
        if include_text:
            result["text"] = ""
        return result

    result = core._scan_contract_text(text, file_name=path.name)
    if include_text:
        result["text"] = text
    return result


# ============================================================
# HAM CHINH
# ============================================================
def extract(filepath, *, scan_result: dict | None = None, preloaded_text: str | None = None):
    path = Path(filepath)
    use_ifilter_for_doc = path.suffix.lower() == ".doc"

    text = preloaded_text
    if text is None:
        text = read_docx(path, use_ifilter_for_doc=use_ifilter_for_doc)

    return core.extract_text(
        text,
        file_goc=str(path),
        file_name=path.name,
        scan_result=scan_result,
        use_ifilter_for_doc=use_ifilter_for_doc,
    )


def main():
    if len(sys.argv) < 2:
        print("Cách dùng:")
        print('  python -m upload_lab.extraction_adapter "hop_dong.docx"')
        print('  python -m upload_lab.extraction_adapter "hop_dong.doc"')
        sys.exit(1)

    fp = sys.argv[1]
    if not os.path.exists(fp):
        print(f"Không tìm thấy file: {fp}")
        sys.exit(1)

    result = extract(fp)
    output = json.dumps(result, ensure_ascii=False, indent=2)
    print(output)

    json_path = fp.rsplit(".", 1)[0] + "_extracted.json"
    with open(json_path, "w", encoding="utf-8") as file_obj:
        file_obj.write(output)
    print(f"\n→ Đã lưu: {json_path}", file=sys.stderr)


if __name__ == "__main__":
    main()
