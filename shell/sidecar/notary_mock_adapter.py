"""Mock backend cho 7 command notary.* — notary.case-drafting.v1 (MIN-106).

Chi duoc goi qua notary_gateway khi G1_DEV_NOTARY_MOCK=1 + khong packaged.
SOT wire: contracts/notary-case-drafting.md — mock tra dung shape/error
code cua contract; khong chinh contract, khong phat minh code moi.

Nguon du lieu: scenario fixtures JSON o
`shell/test/fixtures/notary-case-drafting/*.json` (merge tat ca case, key =
case_id). Fixture la database gia cua mock — "mock tra fixture bat bien".
Thieu thu muc fixture (packaged build, khong bao gio vi gateway chan) ->
fallback seed toi thieu trong code.

Scenario mac dinh: 42 ready | 43 empty | 44 locked | 45 gift (unsupported)
| 46 ready + pool chua gan. Khac -> case_not_found.

Marker deterministic (khong PII, khong parse that):
  - text/file name chua "mock_fail"       -> errors[] intake.parse_failed
  - chua "mock_unsupported"               -> errors[] intake.unsupported_target
  - diagram.render_model="auto" trong seed -> tinh render_model luc load

word_export_batch tao DOCX that bang python-docx (dep co san) theo naming
`<stem>_HS-<case_id>[_n].docx`, reservation trong batch, KHONG BAO GIO ghi
de — ghi vao destination file_ref, khong viet ra folder nguoi dung khac.

Gioi han platform da biet (khong sua trong task nay — jobstore):
  - job failed/canceled -> result=null (jobstore huy result); mock van tra
    error code + details dung contract; thong tin per-file nam trong
    error.details.documents / tren dia.
  - status "partial": handler tra result kem key "partial" lam marker —
    JobStore doc roi POP khoi result truoc khi len wire (jobstore.py).
"""
import copy
import hashlib
import json
import re
import threading
import time
from pathlib import Path

from errors import CommandError
from fileref import validate_file_ref

SCHEMA_VERSION = "notary.case-drafting.v1"
SUPPORTED_CASE_TYPE = "inheritance"

# Gioi han contract §5.2 — mock giu nguyen, khong noi.
MAX_SOURCES = 8
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 100_000

# Delay mo phong moi file word — cho cancel co cua so; test monkeypatch.
WORD_DOC_DELAY = 0.05

UUID4_RX = re.compile(
    r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-4[0-9a-fA-F]{3}"
    r"-[89abAB][0-9a-fA-F]{3}-[0-9a-fA-F]{12}$")
DATE_OR_YEAR_RX = re.compile(r"^\d{4}(-\d{2}-\d{2})?$")
DATE_FULL_RX = re.compile(r"^\d{4}-\d{2}-\d{2}$")
SERIAL_RX = re.compile(r"^[A-Z]{2}\d{6,8}$")
DOC_KEY_RX = re.compile(r"^[a-z][a-z0-9_]*$")
GENDERS = ("Nam", "Nữ", None)
INTAKE_KINDS = ("image", "pdf", "docx", "xlsx", "text")
NODE_BOOL_FIELDS = ("isLandOwner", "willReceive", "hidden", "deleted")
NODE_FIELDS = {"id", "personId", "parentSlotIds", "spouseSlotId",
               "isLandOwner", "willReceive", "hidden", "deleted"}
PERSON_FIELDS = ("row_id", "entity_id", "ho_ten", "gioi_tinh", "ngay_sinh",
                 "ngay_chet", "so_giay_to", "ngay_cap", "noi_cap",
                 "dia_chi", "place_of_origin")
ASSET_FIELDS = ("row_id", "entity_id", "is_primary", "so_serial",
                "so_vao_so", "so_thua_dat", "so_to_ban_do", "dia_chi",
                "loai_so", "hinh_thuc_su_dung", "thoi_han", "nguon_goc",
                "ngay_cap", "co_quan_cap")

# Catalog van ban v1 (contract §8.1) + filename_stem ASCII do backend so huu.
DOC_CATALOG = {
    "khai_nhan_di_san": {
        "display_name": "Văn bản khai nhận di sản",
        "filename_stem": "Van_ban_khai_nhan_di_san",
    },
    "thoa_thuan_phan_chia": {
        "display_name": "Thỏa thuận phân chia di sản",
        "filename_stem": "Thoa_thuan_phan_chia_di_san",
    },
    "niem_yet": {
        "display_name": "Thông báo niêm yết",
        "filename_stem": "Thong_bao_niem_yet",
    },
}
# niem_yet chua co template -> luon bi block (contract §8.1 catalog note).
NO_TEMPLATE_KEYS = {"niem_yet"}

_FIXTURE_DIR = (Path(__file__).resolve().parent.parent
                / "test" / "fixtures" / "notary-case-drafting")

_FAIL_MARK = "mock_fail"
_UNSUPPORTED_MARK = "mock_unsupported"


# ---------- seed / state ----------

def _resolve_case_ref(case, fixture_dir, cid):
    """Case seed co the la {'$ref_scenario': '<name>'} — nap tu file khac."""
    if isinstance(case, dict) and "$ref_scenario" in case:
        ref = fixture_dir / (case["$ref_scenario"] + ".json")
        try:
            doc = json.loads(ref.read_text(encoding="utf-8"))
        except OSError:
            return None
        cases = doc.get("cases") or {}
        return copy.deepcopy(cases.get(cid) or
                             cases.get(next(iter(cases), None)))
    return case


def _builtin_cases():
    """Fallback toi thieu khi fixture dir khong ton tai (chi defensive —
    mock khong chay packaged)."""
    return {
        "43": {
            "case_type": "inheritance", "document_type": "khai_nhan",
            "status": "draft", "locked": False, "revision": 1,
            "stage": {"people": [], "assets": []},
            "diagram": {"state": {"version": 2, "nodes": []},
                        "render_model": None, "warnings": []},
        },
    }


# Thu tu merge fixture — case 42 duoc khai o nhieu file (ready goc,
# conflict/intake-partial la ban giam); seed mac dinh lay phien ban
# "primary" truoc. Seed rieng trong test dung reset_backend(seed).
_FIXTURE_PRIORITY = (
    "ready", "empty", "locked", "unsupported", "diagram-warning",
    "intake-partial", "conflict", "word-collision", "word-partial",
    "word-all-failed", "word-canceled",
)


def _load_default_cases():
    """Merge moi fixture *.json thanh case map; fallback builtin."""
    cases = {}
    if _FIXTURE_DIR.is_dir():
        files = sorted(
            _FIXTURE_DIR.glob("*.json"),
            key=lambda f: (
                _FIXTURE_PRIORITY.index(f.stem)
                if f.stem in _FIXTURE_PRIORITY else 99, f.stem))
        for f in files:
            try:
                doc = json.loads(f.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            for cid, c in (doc.get("cases") or {}).items():
                c = _resolve_case_ref(c, _FIXTURE_DIR, cid)
                if isinstance(c, dict) and cid not in cases:
                    cases[cid] = c
    return cases or _builtin_cases()


class _MockState:
    """Server gia trong bo nho — case map + counter entity_id."""

    def __init__(self, seed=None):
        self._lock = threading.Lock()
        self._next_entity_id = 500
        raw = (seed or {}).get("cases") or _load_default_cases()
        self.cases = {}
        for cid, c in raw.items():
            prepared = self._prepare(c)
            prepared["id"] = int(cid)
            self.cases[int(cid)] = prepared

    def _prepare(self, c):
        case = copy.deepcopy(c)
        case.setdefault("case_type", "inheritance")
        case.setdefault("document_type", "khai_nhan")
        case.setdefault("status", "draft")
        case.setdefault("locked", False)
        case.setdefault("revision", 1)
        case.setdefault("stage", {"people": [], "assets": []})
        dg = case.setdefault("diagram", {})
        dg.setdefault("domain", "inheritance")
        dg.setdefault("state", {"version": 2, "nodes": []})
        dg.setdefault("warnings", [])
        if dg.get("render_model") == "auto":
            # render_model luon khop state hien tai (contract §4)
            dg["render_model"] = _render(dg["state"], case["stage"])
        elif "render_model" not in dg:
            dg["render_model"] = None
        return case

    def assign_entity_id(self):
        eid = self._next_entity_id
        self._next_entity_id += 1
        return eid


_STATE = None
_STATE_LOCK = threading.Lock()


def _state():
    global _STATE
    with _STATE_LOCK:
        if _STATE is None:
            _STATE = _MockState()
        return _STATE


def reset_backend(seed=None):
    """Test hook: reset state; seed={'cases': {...}} hoac None = default."""
    global _STATE
    with _STATE_LOCK:
        _STATE = _MockState(seed)


# ---------- helpers ----------

def _result(kind, data, warnings=None, source_files=None,
            evidence=None, partial=False):
    r = {"kind": kind, "data": data,
         "evidence": evidence or [], "warnings": warnings or [],
         "source_files": source_files or []}
    if partial:
        r["partial"] = True        # marker jobstore -> status 'partial'
    return r


def _case(payload):
    cid = (payload or {}).get("case_id")
    if not isinstance(cid, int) or isinstance(cid, bool) or cid < 1:
        raise CommandError("validation_error",
                           f"case_id={cid!r} phai la int >= 1")
    case = _state().cases.get(cid)
    if case is None:
        raise CommandError("case_not_found",
                           f"Không tìm thấy hồ sơ #{cid}",
                           details={"case_id": cid})
    return case


def _check_supported(case):
    ct = case.get("case_type")
    if ct != SUPPORTED_CASE_TYPE:
        raise CommandError("case_type_unsupported",
                           f"loại việc {ct!r} chưa hỗ trợ trong v1",
                           details={"case_type": ct})


def _check_writable(case):
    _check_supported(case)
    if case.get("locked"):
        raise CommandError("workspace_locked",
                           f"hồ sơ #{case['id']} đã khóa")


def _check_base_revision(case, payload):
    br = (payload or {}).get("base_revision")
    if not isinstance(br, int) or isinstance(br, bool) or br < 1:
        raise CommandError("validation_error",
                           f"base_revision={br!r} phai la int >= 1")
    srv = case["revision"]
    if br != srv:
        raise CommandError("workspace_conflict",
                           f"base_revision={br} khác server_revision={srv}",
                           details={"server_revision": srv})


def _det_uuid4(seed):
    """UUID4-shaped deterministic id cho suggestion (suggestion_id do
    backend sinh — fixture bat bien giua hai lan goi)."""
    h = hashlib.md5(seed.encode("utf-8")).hexdigest()
    return (f"{h[0:8]}-{h[8:12]}-4{h[13:16]}-"
            f"{('8', '9', 'a', 'b')[int(h[16], 16) % 4]}{h[17:20]}-{h[20:32]}")


def _person_map(case):
    return {p["row_id"]: p for p in case["stage"]["people"]
            if isinstance(p, dict)}


# ---------- diagram validation + render (mock engine v2) ----------

def _validate_diagram_state(state, case):
    """Tra (errors, outside_pid). errors dung engine code contract §7.3."""
    errors = []
    if not isinstance(state, dict):
        return [{"code": "invalid_input", "message": "state không phải object"}], None
    if state.get("version") != 2:
        errors.append({"code": "invalid_version",
                       "message": f"version={state.get('version')!r}, cần 2"})
    nodes = state.get("nodes")
    if not isinstance(nodes, list):
        errors.append({"code": "invalid_nodes", "message": "nodes không phải list"})
        return errors, None
    ids = set()
    persons = {}
    for i, n in enumerate(nodes):
        w = f"nodes[{i}]"
        if not isinstance(n, dict):
            errors.append({"code": "invalid_node", "node": w,
                           "message": f"{w} không phải object"})
            continue
        nid = n.get("id")
        if not (isinstance(nid, str) and nid.strip()):
            errors.append({"code": "missing_node_id", "node": w,
                           "message": f"{w} thiếu id"})
            nid = None
        elif nid in ids:
            errors.append({"code": "duplicate_node_id", "node": nid,
                           "message": f"trùng node id {nid!r}"})
        else:
            ids.add(nid)
        for f in NODE_BOOL_FIELDS:
            if not isinstance(n.get(f), bool):
                errors.append({"code": "invalid_boolean", "node": nid or w,
                               "message": f"{w}.{f} phải là boolean"})
        pid = n.get("personId")
        if pid is not None:
            if not (isinstance(pid, str) and UUID4_RX.match(pid)):
                errors.append({"code": "invalid_node", "node": nid or w,
                               "message": f"{w}.personId không phải uuid4"})
            else:
                persons.setdefault(pid, []).append(nid or w)
        for k in set(n) - NODE_FIELDS:
            errors.append({"code": "invalid_node", "node": nid or w,
                           "message": f"{w} field lạ {k!r}"})
        ps = n.get("parentSlotIds")
        if not (isinstance(ps, list)
                and all(isinstance(x, str) and x.strip() for x in ps)):
            errors.append({"code": "invalid_parent_slots", "node": nid or w,
                           "message": f"{w}.parentSlotIds phải là "
                                      "list[string non-empty]"})
            ps = []
        if len(ps) > 2:
            errors.append({"code": "too_many_parents", "node": nid or w,
                           "message": f"{w} quá 2 parent slots"})
        if nid and nid in ps:
            errors.append({"code": "self_parent", "node": nid,
                           "message": f"{nid} tự làm cha"})
        ss = n.get("spouseSlotId")
        if ss is not None and not (isinstance(ss, str) and ss.strip()):
            errors.append({"code": "invalid_node", "node": nid or w,
                           "message": f"{w}.spouseSlotId phải là "
                                      "string non-empty/null"})
        if nid and ss == nid:
            errors.append({"code": "self_spouse", "node": nid,
                           "message": f"{nid} tự làm vợ/chồng"})
    # dangling refs + spouse conflict
    node_by_id = {n.get("id"): n for n in nodes
                  if isinstance(n, dict) and n.get("id")}
    for pid, slots in persons.items():
        active = [s for s in slots
                  if not (node_by_id.get(s) or {}).get("deleted")]
        if len(active) > 1:
            errors.append({"code": "duplicate_person",
                           "message": f"personId {pid} gán trên nhiều node"})
    for i, n in enumerate(nodes):
        if not isinstance(n, dict):
            continue
        nid = n.get("id")
        for p in n.get("parentSlotIds") or []:
            if isinstance(p, str) and p not in ids:
                errors.append({"code": "dangling_parent", "node": nid,
                               "message": f"{nid} tham chiếu parent {p!r} "
                                          "không tồn tại"})
        ss = n.get("spouseSlotId")
        if isinstance(ss, str):
            if ss not in ids:
                errors.append({"code": "dangling_spouse", "node": nid,
                               "message": f"{nid} tham chiếu spouse {ss!r} "
                                          "không tồn tại"})
            else:
                other = node_by_id.get(ss) or {}
                if other.get("spouseSlotId") != nid:
                    errors.append({"code": "spouse_conflict", "node": nid,
                                   "message": f"{nid}↔{ss} spouse không đối xứng"})
    # ancestry cycle qua parentSlotIds
    for nid in ids:
        # DFS don gian: di theo parent edges, quay lai chinh minh = cycle
        seen = set()
        frontier = list((node_by_id.get(nid) or {}).get("parentSlotIds") or [])
        while frontier:
            cur = frontier.pop()
            if cur == nid:
                errors.append({"code": "ancestry_cycle", "node": nid,
                               "message": f"chu kỳ tổ tiên qua {nid}"})
                frontier = []
                break
            if cur in seen:
                continue
            seen.add(cur)
            frontier.extend(
                (node_by_id.get(cur) or {}).get("parentSlotIds") or [])
    # personId ngoai stage da commit -> job error rieng (contract §7.1)
    stage_ids = set(_person_map(case))
    for n in nodes:
        if not isinstance(n, dict):
            continue
        pid = n.get("personId")
        if isinstance(pid, str) and UUID4_RX.match(pid) \
                and pid not in stage_ids:
            return errors, pid
    return errors, None


def _render(state, stage):
    """Render_model deterministic cua mock engine (v2 shape, contract §7.2).

    Rule don gian mot estate: landowner da chet -> chia deu cho cac node
    willReceive co personId. Khong co landowner -> invalid
    missing_land_owner. Pool person chua gan -> incomplete + warning.
    """
    stage_people = _person_map({"stage": stage})
    nodes = [n for n in (state.get("nodes") or [])
             if isinstance(n, dict) and not n.get("deleted")]
    empty = {"engineVersion": 2, "status": "invalid", "allocations": {},
             "breakdowns": [], "requiredSlots": [], "warnings": [],
             "errors": [{"code": "missing_land_owner",
                         "message": "Sơ đồ chưa có chủ đất"}],
             "unresolvedEstates": [],
             "conservation": {"allocated": "0", "unresolved": "0",
                              "total": "0"}}
    landowners = [n for n in nodes
                  if n.get("isLandOwner") and n.get("personId")]
    if not landowners:
        if not nodes:
            return empty
        rm = copy.deepcopy(empty)
        rm["requiredSlots"] = _required_slots(nodes)
        return rm
    deceased = [n for n in landowners
                if (stage_people.get(n["personId"]) or {}).get("ngay_chet")]
    assigned = {n["personId"] for n in nodes if n.get("personId")}
    unassigned_pool = [rid for rid in stage_people if rid not in assigned]
    warnings = []
    if unassigned_pool:
        warnings.append({
            "code": "diagram.unassigned_pool_person",
            "message": "Còn người trong Pool chưa được gán trên sơ đồ"})
    receivers = [n for n in nodes
                 if n.get("willReceive") and n.get("personId")]
    rm = {"engineVersion": 2, "status": "complete", "allocations": {},
          "breakdowns": [], "requiredSlots": _required_slots(nodes),
          "warnings": warnings, "errors": [], "unresolvedEstates": [],
          "conservation": {"allocated": "0", "unresolved": "0",
                           "total": "0"}}
    if not deceased:
        rm["status"] = "incomplete"
        rm["warnings"] = warnings + [{
            "code": "diagram.no_active_estate",
            "message": "Chưa có chủ đất đã chết — chưa phát sinh di sản"}]
        for n in nodes:
            if n.get("personId"):
                rm["allocations"][n["personId"]] = _alloc("0", "0", "0", "0")
        return rm
    if not receivers:
        src = deceased[0]["personId"]
        ev = (stage_people.get(src) or {}).get("ngay_chet")
        # unresolved_estate.eventDate la date_full; ngay_chet nam-le -> null
        if not (isinstance(ev, str) and DATE_FULL_RX.match(ev)):
            ev = None
        rm["status"] = "incomplete"
        rm["unresolvedEstates"] = [{
            "sourcePersonId": src, "eventDate": ev,
            "fraction": "1", "reason": "no_valid_heir"}]
        rm["conservation"] = {"allocated": "0", "unresolved": "1",
                              "total": "1"}
        for n in nodes:
            if n.get("personId"):
                rm["allocations"][n["personId"]] = _alloc("0", "0", "0", "0")
        return rm
    n_recv = len(receivers)
    share = f"1/{n_recv}" if n_recv > 1 else "1"
    pct = f"{100.0 / n_recv:.2f}"
    src_id = deceased[0]["personId"]
    for n in nodes:
        pid = n.get("personId")
        if not pid:
            continue
        if n in landowners and pid == src_id:
            rm["allocations"][pid] = _alloc("1", "0", "1", "0")
        elif n.get("willReceive"):
            rm["allocations"][pid] = _alloc("0", share, "0", share)
            rm["allocations"][pid]["displayPercent"] = pct
            rm["breakdowns"].append({
                "personId": pid, "total": share,
                "terms": [{"kind": "inheritance", "fraction": share,
                           "sourcePersonId": src_id,
                           "viaBranchPersonIds": []}]})
        else:
            rm["allocations"][pid] = _alloc("0", "0", "0", "0")
    rm["conservation"] = {"allocated": "1", "unresolved": "0", "total": "1"}
    if unassigned_pool:
        rm["status"] = "incomplete"
    return rm


def _alloc(base, inherited, distributed, final):
    return {"baseShare": base, "inheritedShare": inherited,
            "distributedShare": distributed, "finalShare": final,
            "displayPercent": "0.00"}


def _required_slots(nodes):
    return [{"anchorSlotId": n["id"], "reason": "active_estate",
             "slotTypes": ["father", "mother", "spouse", "child"],
             "minimumEmptyChildSlots": 0}
            for n in nodes
            if isinstance(n, dict) and n.get("id") and not n.get("personId")
            and not n.get("deleted")]


def _payload_diagram_state(payload):
    """Lay diagram.state tu payload — thieu field -> validation_error
    (missing required field); state co mat nhung sai shape -> de
    _evaluate_state_or_raise lo (diagram_invalid_state)."""
    dg = (payload or {}).get("diagram")
    if not isinstance(dg, dict) or "state" not in dg:
        raise CommandError("validation_error",
                           "payload.diagram.state bắt buộc")
    return dg["state"]


def _evaluate_state_or_raise(case, state):
    """Validate + render. Raise CommandError theo contract §7."""
    errors, outside_pid = _validate_diagram_state(state, case)
    if errors:
        raise CommandError("diagram_invalid_state",
                           "diagram state không hợp lệ",
                           details={"errors": errors})
    if outside_pid is not None:
        raise CommandError("diagram_reference_outside_stage",
                           f"personId {outside_pid} không thuộc Stage "
                           "đã commit", details={"personId": outside_pid})
    return _render(state, case["stage"])


def _prune_diagram(case):
    """Sau commit stage: bo personId/slot tham chieu phan tu da xoa (§6.1)."""
    stage_ids = set(_person_map(case))
    nodes = (case["diagram"].get("state") or {}).get("nodes") or []
    node_ids = {n.get("id") for n in nodes if isinstance(n, dict)}
    for n in nodes:
        if not isinstance(n, dict):
            continue
        if n.get("personId") and n["personId"] not in stage_ids:
            n["personId"] = None
        if isinstance(n.get("parentSlotIds"), list):
            n["parentSlotIds"] = [p for p in n["parentSlotIds"]
                                  if p in node_ids]
        ss = n.get("spouseSlotId")
        if ss is not None and ss not in node_ids:
            n["spouseSlotId"] = None


# ---------- stage validation ----------

def _err(row_id, field, code, message):
    return {"row_id": row_id, "field": field, "code": code,
            "message": message}


def _err_row(rid, label):
    """field_error.row_id schema bat uuid4 — row_id loi thi dung uuid4
    deterministic tu label (van truy vet duoc qua message)."""
    if isinstance(rid, str) and UUID4_RX.match(rid):
        return rid
    return _det_uuid4(f"field_error:{label}:{rid!r}")


def _validate_stage(stage):
    """Tra (field_errors, people_norm, assets_norm). Norm = strip field la
    + dien null cho nullable thieu (producer strip — contract §4.1).

    Code parity voi validator oracle: loi CONTAINER shape (stage khong
    dict / people|assets thieu hoac khong list) -> validation_error;
    stage_validation_error chi cho loi ROW-level trong field_errors."""
    if not isinstance(stage, dict):
        raise CommandError("validation_error",
                           "stage phải là object")
    missing = [k for k in ("people", "assets")
               if not isinstance(stage.get(k), list)]
    if missing:
        raise CommandError(
            "validation_error",
            f"stage.{', stage.'.join(missing)} phải là list")
    errs = []
    people = stage["people"]
    assets = stage["assets"]
    seen = set()
    people_norm = []
    for i, r in enumerate(people):
        rid = r.get("row_id") if isinstance(r, dict) else None
        w = _err_row(rid, f"people[{i}]")
        if not isinstance(r, dict):
            errs.append(_err(w, "row", "invalid_type",
                             f"people[{i}] không phải object"))
            continue
        if not (isinstance(rid, str) and UUID4_RX.match(rid)):
            errs.append(_err(w, "row_id", "invalid_format",
                             f"row_id phải là uuid4 (nhận {rid!r})"))
        elif rid in seen:
            errs.append(_err(rid, "row_id", "duplicate_row_id",
                             "row_id trùng trong payload"))
        seen.add(rid)
        eid = r.get("entity_id")
        if eid is not None and (not isinstance(eid, int)
                                or isinstance(eid, bool)):
            errs.append(_err(w, "entity_id", "invalid_type",
                             "entity_id phải là int/null"))
        ht = r.get("ho_ten")
        if not (isinstance(ht, str) and ht.strip()):
            errs.append(_err(w, "ho_ten", "required", "ho_ten bắt buộc"))
        if r.get("gioi_tinh") not in GENDERS:
            errs.append(_err(w, "gioi_tinh", "invalid_enum",
                             "gioi_tinh ∈ {Nam, Nữ, null}"))
        for f in ("ngay_sinh", "ngay_chet", "ngay_cap"):
            v = r.get(f)
            if v is not None and not (isinstance(v, str)
                                      and DATE_OR_YEAR_RX.match(v)):
                errs.append(_err(w, f, "invalid_date",
                                 f"{f} phải YYYY-MM-DD/YYYY/null"))
        people_norm.append({f: r.get(f) for f in PERSON_FIELDS})
    assets_norm = []
    primaries = 0
    for i, r in enumerate(assets):
        rid = r.get("row_id") if isinstance(r, dict) else None
        w = _err_row(rid, f"assets[{i}]")
        if not isinstance(r, dict):
            errs.append(_err(w, "row", "invalid_type",
                             f"assets[{i}] không phải object"))
            continue
        if not (isinstance(rid, str) and UUID4_RX.match(rid)):
            errs.append(_err(w, "row_id", "invalid_format",
                             f"row_id phải là uuid4 (nhận {rid!r})"))
        elif rid in seen:
            errs.append(_err(rid, "row_id", "duplicate_row_id",
                             "row_id trùng trong payload"))
        seen.add(rid)
        eid = r.get("entity_id")
        if eid is not None and (not isinstance(eid, int)
                                or isinstance(eid, bool)):
            errs.append(_err(w, "entity_id", "invalid_type",
                             "entity_id phải là int/null"))
        if not isinstance(r.get("is_primary"), bool):
            errs.append(_err(w, "is_primary", "invalid_type",
                             "is_primary phải là boolean"))
        elif r["is_primary"]:
            primaries += 1
        serial = r.get("so_serial")
        if not (isinstance(serial, str) and serial.strip()):
            errs.append(_err(w, "so_serial", "required",
                             "so_serial bắt buộc"))
        elif not SERIAL_RX.match(serial):
            errs.append(_err(w, "so_serial", "invalid_format",
                             "so_serial phải canonical [A-Z]{2}\\d{6,8}"))
        dc = r.get("dia_chi")
        if not (isinstance(dc, str) and dc.strip()):
            errs.append(_err(w, "dia_chi", "required", "dia_chi bắt buộc"))
        nc = r.get("ngay_cap")
        if nc is not None and not (isinstance(nc, str)
                                   and DATE_FULL_RX.match(nc)):
            errs.append(_err(w, "ngay_cap", "invalid_date",
                             "ngay_cap tài sản chỉ nhận YYYY-MM-DD"))
        lr = r.get("land_rows")
        if lr is not None and not isinstance(lr, list):
            errs.append(_err(w, "land_rows", "invalid_type",
                             "land_rows phải là list/null"))
        elif isinstance(lr, list):
            for j, x in enumerate(lr):
                if not isinstance(x, dict):
                    errs.append(_err(w, f"land_rows[{j}]", "invalid_type",
                                     "land_row phải là object"))
                    continue
                dt = x.get("dien_tich")
                if dt is not None and (not isinstance(dt, (int, float))
                                       or isinstance(dt, bool)):
                    errs.append(_err(w, f"land_rows[{j}].dien_tich",
                                     "invalid_type",
                                     "dien_tich phải là number/null"))
                for lf in ("loai_dat", "thoi_han"):
                    lv = x.get(lf)
                    if lv is not None and not (
                            isinstance(lv, str) and lv.strip()):
                        errs.append(_err(w, f"land_rows[{j}].{lf}",
                                         "invalid_type",
                                         f"{lf} phải là string "
                                         "non-empty/null"))
        row = {f: r.get(f) for f in ASSET_FIELDS}
        row["land_rows"] = [
            {"loai_dat": x.get("loai_dat"),
             "dien_tich": x.get("dien_tich"),
             "thoi_han": x.get("thoi_han")}
            for x in lr if isinstance(x, dict)] \
            if isinstance(lr, list) else None
        assets_norm.append(row)
    if assets and primaries != 1:
        errs.append(_err(
            _err_row(assets_norm[0].get("row_id"), "assets[0]")
            if assets_norm else _det_uuid4("field_error:primary_count"),
            "is_primary", "primary_count",
            f"cần đúng 1 tài sản chính, có {primaries}"))
    return errs, people_norm, assets_norm


# ---------- commands ----------

def workspace_get(job, payload):
    case = _case(payload)
    caps = case.get("case_type") == SUPPORTED_CASE_TYPE
    data = {
        "schema_version": SCHEMA_VERSION,
        "backend_mode": "mock",          # contract §4 — nhan 'Dữ liệu mô phỏng'
        "case": {
            "id": int(payload["case_id"]),
            "case_type": case.get("case_type"),
            "document_type": case.get("document_type"),
            "status": case.get("status"),
            "locked": bool(case.get("locked")),
            "revision": case.get("revision"),
        },
        "stage": copy.deepcopy(case["stage"]),
        "diagram": {
            "domain": case["diagram"].get("domain", "inheritance"),
            "state": copy.deepcopy(case["diagram"]["state"]),
            "render_model": copy.deepcopy(case["diagram"]["render_model"]),
            "warnings": copy.deepcopy(case["diagram"].get("warnings") or []),
        },
        "capabilities": {
            "intake": list(INTAKE_KINDS) if caps else [],
            "diagram": caps,
            "word_export": caps,
        },
    }
    job.check_cancel()
    return _result("workspace_get", data)


def intake_analyze(job, payload):
    case = _case(payload)
    _check_writable(case)          # §5.3: locked/type -> loi job-level
    sources = (payload or {}).get("sources")
    if not isinstance(sources, list) or not sources:
        raise CommandError("validation_error", "sources phải là list 1..8")
    if len(sources) > MAX_SOURCES:
        raise CommandError("intake_too_many_sources",
                           f"{len(sources)} nguồn > {MAX_SOURCES}",
                           details={"count": len(sources),
                                    "limit": MAX_SOURCES})
    seen = set()
    for s in sources:
        if not isinstance(s, dict):
            raise CommandError("validation_error", "source không phải object")
        sid = s.get("source_id")
        if not (isinstance(sid, str) and UUID4_RX.match(sid)):
            raise CommandError("validation_error",
                               f"source_id={sid!r} phải là uuid4")
        if sid in seen:
            raise CommandError("validation_error",
                               f"source_id trùng: {sid}")
        seen.add(sid)
        kind = s.get("kind")
        if kind not in INTAKE_KINDS:
            raise CommandError("intake_unsupported_source",
                               f"kind={kind!r} ngoài enum",
                               details={"source_id": sid, "kind": kind})
        if kind == "text":
            if "file_ref" in s:
                raise CommandError("validation_error",
                                   "kind=text cấm file_ref")
            t = s.get("text")
            if not isinstance(t, str) or not t:
                raise CommandError("validation_error",
                                   "kind=text bắt buộc text non-empty")
            if len(t) > MAX_TEXT_CHARS:
                raise CommandError("intake_text_too_long",
                                   f"text {len(t)} > {MAX_TEXT_CHARS}",
                                   details={"source_id": sid,
                                            "length": len(t),
                                            "limit": MAX_TEXT_CHARS})
        else:
            if "text" in s:
                raise CommandError("validation_error",
                                   f"kind={kind} cấm text")
            ref = s.get("file_ref")
            if not isinstance(ref, dict):
                raise CommandError("validation_error",
                                   f"kind={kind} bắt buộc file_ref")
            if ref.get("is_dir") is True:
                raise CommandError("validation_error",
                                   "intake file_ref is_dir=true")
            sb = ref.get("size_bytes")
            if not isinstance(sb, int) or isinstance(sb, bool) or sb < 0:
                raise CommandError("validation_error",
                                   "intake file_ref bắt buộc size_bytes "
                                   "int >= 0")
            if sb > MAX_FILE_BYTES:
                raise CommandError("intake_source_too_large",
                                   f"size_bytes={sb} > {MAX_FILE_BYTES}",
                                   details={"source_id": sid,
                                            "limit": MAX_FILE_BYTES})
    suggestions = []
    errors = []
    ok_ids = []
    bad_ids = []
    for i, s in enumerate(sources):
        job.check_cancel()
        job.report_progress(i, len(sources),
                            f"nguồn {i + 1}/{len(sources)}")
        sid = s["source_id"]
        blob = s.get("text") if s["kind"] == "text" else (
            (s.get("file_ref") or {}).get("path") or "")
        low = str(blob).lower()
        if _FAIL_MARK in low:
            errors.append({"source_id": sid, "code": "intake.parse_failed",
                           "message": "Không trích được thực thể nào "
                                      "(mô phỏng)"})
            bad_ids.append(sid)
            continue
        if _UNSUPPORTED_MARK in low:
            errors.append({"source_id": sid,
                           "code": "intake.unsupported_target",
                           "message": "Loại nguồn không map được "
                                      "(mô phỏng)"})
            bad_ids.append(sid)
            continue
        suggestions.append(_canned_suggestion(s))
        ok_ids.append(sid)
    job.report_progress(len(sources), len(sources), "Hoàn tất")
    data = {"schema_version": SCHEMA_VERSION,
            "suggestions": suggestions, "errors": errors}
    partial = bool(errors)
    if partial:
        data["breakdown"] = {"succeeded": ok_ids, "failed": bad_ids}
    src_files = [{"path": s["file_ref"]["path"], "scope": "machine_local"}
                 for s in sources
                 if s["kind"] != "text" and isinstance(s.get("file_ref"), dict)
                 and isinstance(s["file_ref"].get("path"), str)]
    return _result("intake_analyze", data, source_files=src_files,
                   partial=partial)


def _canned_suggestion(src):
    """Suggestion fixture bat bien theo kind — observation_state hop le,
    khong bao gio 'confirmed'."""
    sid = src["source_id"]
    if src["kind"] in ("text", "image"):
        ref = {"span": [0, 10]} if src["kind"] == "text" else {"page": 1}
        fields = {
            "ho_ten": {"raw_value": "NGƯỜI MẪU I",
                       "normalized_value": "Người Mẫu I",
                       "observation_state": "normalized",
                       "confidence": None, "source_refs": [ref]},
            "ngay_sinh": {"raw_value": "1950",
                          "normalized_value": "1950",
                          "observation_state": "normalized",
                          "confidence": None, "source_refs": [ref]},
            "gioi_tinh": {"raw_value": "Nam",
                          "normalized_value": "Nam",
                          "observation_state": "observed",
                          "confidence": None, "source_refs": [ref]},
        }
        target = "person"
    else:
        ref = {"page": 1}
        fields = {
            "so_serial": {"raw_value": "MM 000010",
                          "normalized_value": "MM000010",
                          "observation_state": "inferred",
                          "confidence": 0.62, "source_refs": [ref]},
            "dia_chi": {"raw_value": "Địa chỉ mẫu intake",
                        "normalized_value": None,
                        "observation_state": "observed",
                        "confidence": None, "source_refs": [ref]},
            "so_thua_dat": {"raw_value": "123",
                            "normalized_value": "123",
                            "observation_state": "normalized",
                            "confidence": None,
                            "source_refs": [{"page": 2}]},
        }
        target = "asset"
    return {"suggestion_id": _det_uuid4(sid + ":" + src["kind"]),
            "source_id": sid, "target": target, "fields": fields,
            "warnings": []}


def workspace_commit_stage(job, payload):
    case = _case(payload)
    _check_writable(case)
    _check_base_revision(case, payload)
    stage = (payload or {}).get("stage")
    errs, people_norm, assets_norm = _validate_stage(stage)
    if errs:
        raise CommandError("stage_validation_error",
                           f"{len(errs)} lỗi field trong Stage",
                           details={"field_errors": errs})
    # Atomic: assign entity_id + swap stage + prune + re-evaluate + rev+1
    for row in people_norm + assets_norm:
        if row.get("entity_id") is None:
            row["entity_id"] = _state().assign_entity_id()
    case["stage"] = {"people": people_norm, "assets": assets_norm}
    _prune_diagram(case)
    case["diagram"]["render_model"] = _render(
        case["diagram"]["state"], case["stage"])
    case["revision"] += 1
    job.check_cancel()
    return _result("workspace_commit_stage", {
        "schema_version": SCHEMA_VERSION,
        "revision": case["revision"],
        "stage": copy.deepcopy(case["stage"]),
        "diagram": {
            "state": copy.deepcopy(case["diagram"]["state"]),
            "render_model": copy.deepcopy(
                case["diagram"]["render_model"]),
        },
    })


def diagram_evaluate(job, payload):
    """Read-only — duoc phep tren case locked (§7.4); khong persist."""
    case = _case(payload)
    _check_supported(case)
    state = _payload_diagram_state(payload)
    rm = _evaluate_state_or_raise(case, state)
    job.check_cancel()
    return _result("diagram_evaluate", {
        "schema_version": SCHEMA_VERSION,
        "evaluated_revision": case["revision"],
        "render_model": rm,
    })


def diagram_save(job, payload):
    case = _case(payload)
    _check_writable(case)
    _check_base_revision(case, payload)
    state = _payload_diagram_state(payload)
    rm = _evaluate_state_or_raise(case, state)
    case["diagram"]["state"] = copy.deepcopy(state)
    case["diagram"]["render_model"] = rm
    case["revision"] += 1
    job.check_cancel()
    return _result("diagram_save", {
        "schema_version": SCHEMA_VERSION,
        "revision": case["revision"],
        "diagram": {
            "state": copy.deepcopy(case["diagram"]["state"]),
            "render_model": copy.deepcopy(rm),
        },
    })


def _word_block_reason(case, document_key):
    """Readiness theo validation that (contract §8.1 mapping)."""
    if document_key in NO_TEMPLATE_KEYS:
        return "word.template_missing"
    assets = case["stage"]["assets"]
    nodes = [n for n in (case["diagram"].get("state") or {}).get("nodes")
             or [] if isinstance(n, dict) and not n.get("deleted")]
    if not assets:
        return "word.no_assets"
    if len(assets) > 5:
        return "word.too_many_assets"
    stage_people = _person_map(case)
    landowners = [n for n in nodes
                  if n.get("isLandOwner") and n.get("personId")]
    if not landowners:
        return "word.no_landowner"
    deceased = [n for n in landowners
                if (stage_people.get(n["personId"]) or {}).get("ngay_chet")]
    if not deceased:
        return "word.no_deceased_landowner"
    receivers = [n for n in nodes
                 if n.get("willReceive") and n.get("personId")]
    if not receivers:
        return "word.no_receiver"
    assigned = [n for n in nodes if n.get("personId")]
    if len(assigned) > 20:
        return "word.too_many_people"
    if len(receivers) > 20:
        return "word.too_many_signers"
    return None


def word_export_options(job, payload):
    """Read-only — duoc phep tren locked/unsupported (chi workspace_get
    quyet dinh capability); block_reason theo validation that."""
    case = _case(payload)
    docs = []
    for key, meta in DOC_CATALOG.items():
        reason = _word_block_reason(case, key)
        docs.append({"document_key": key,
                     "display_name": meta["display_name"],
                     "ready": reason is None,
                     "block_reason": reason})
    job.check_cancel()
    return _result("word_export_options",
                   {"schema_version": SCHEMA_VERSION, "documents": docs})


def _write_docx(out, title, case):
    """Tao DOCX that bang python-docx; fallback zip toi thieu hop le.
    `out` la file object mo san mode wb (exclusive create)."""
    try:
        import docx
        d = docx.Document()
        d.add_heading(title, level=1)
        d.add_paragraph(
            f"Dữ liệu mô phỏng — hồ sơ #{case['id']} "
            f"(notary.case-drafting.v1 mock).")
        d.add_paragraph("Nội dung mẫu, không phải văn bản pháp lý.")
        d.save(out)
        return
    except ImportError:
        pass
    # Fallback: docx toi thieu hop le (zip + 2 part bat buoc)
    import zipfile
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/'
        'content-types"><Default Extension="rels" ContentType="application/'
        'vnd.openxmlformats-package.relationships+xml"/><Default '
        'Extension="xml" ContentType="application/xml"/><Override '
        'PartName="/word/document.xml" ContentType="application/'
        'vnd.openxmlformats-officedocument.wordprocessingml.document.main'
        '+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/'
            'package/2006/relationships"><Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/officeDocument" Target="word/document.xml"/>'
            '</Relationships>')
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Dữ liệu mô '
        'phỏng — mock notary.case-drafting.v1</w:t></w:r></w:p></w:body>'
        '</w:document>')
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)


def _reserve_and_write(dest_dir, stem, case_id, taken, title, case):
    """Naming <stem>_HS-<id>[_n].docx (n>=2) + ghi file bang
    open('xb') exclusive-create: ten trong `taken` + ten da ton tai deu
    bo qua — KHONG BAO GIO ghi de (§8.3) va khong TOCTOU giua
    exists-check va write. Tra actual_filename."""
    base = f"{stem}_HS-{case_id}"
    n = 1
    while True:
        suffix = "" if n == 1 else f"_{n}"
        name = f"{base}{suffix}.docx"
        n += 1
        if name in taken:
            continue
        path = dest_dir / name
        try:
            fh = open(path, "xb")          # exclusive create — atomic
        except FileExistsError:
            continue
        taken.add(name)
        try:
            _write_docx(fh, title, case)
            fh.close()
        except Exception:
            try:
                fh.close()
            finally:
                try:
                    path.unlink()          # khong de file hong lai
                except OSError:
                    pass
            raise
        return name


def word_export_batch(job, payload):
    case = _case(payload)
    _check_writable(case)
    keys = (payload or {}).get("document_keys")
    # Oracle parity: thieu/null/khong list -> validation_error;
    # word_no_documents_selected CHI cho list rong [].
    if not isinstance(keys, list):
        raise CommandError("validation_error",
                           "document_keys phải là list")
    if not keys:
        raise CommandError("word_no_documents_selected",
                           "document_keys rỗng")
    if len(keys) != len(set(keys)):
        dup = next(k for k in keys if keys.count(k) > 1)
        raise CommandError("word_duplicate_document_key",
                           f"document_key lặp: {dup}",
                           details={"document_key": dup})
    for k in keys:
        if not (isinstance(k, str) and DOC_KEY_RX.match(k)) \
                or k not in DOC_CATALOG:
            raise CommandError("word_unknown_document_key",
                               f"document_key ngoài catalog: {k!r}",
                               details={"document_key": k})
    dest = (payload or {}).get("destination")
    if not isinstance(dest, dict) or dest.get("is_dir") is not True:
        raise CommandError("validation_error",
                           "destination phải là file_ref is_dir:true")
    dest_dir = validate_file_ref(dest)      # scope/UNC/abs — envelope §6
    if not dest_dir.is_dir():
        raise CommandError("file_not_found",
                           "destination không phải thư mục đang tồn tại")
    docs = []
    taken = set()
    failed = []
    saved = []
    for i, key in enumerate(keys):
        job.check_cancel()                   # cancel giua batch
        meta = DOC_CATALOG[key]
        reason = _word_block_reason(case, key)
        if reason is not None:
            docs.append({"document_key": key,
                         "display_name": meta["display_name"],
                         "status": "failed", "actual_filename": None,
                         "output_file": None,
                         "error": {"code": reason,
                                   "message": _block_message(reason)}})
            failed.append(key)
        else:
            time.sleep(WORD_DOC_DELAY)       # mo phong render docx
            job.check_cancel()               # truoc khi ghi file
            name = _reserve_and_write(
                dest_dir, meta["filename_stem"],
                int(payload["case_id"]), taken,
                meta["display_name"], {"id": payload["case_id"]})
            out_path = dest_dir / name
            docs.append({"document_key": key,
                         "display_name": meta["display_name"],
                         "status": "saved", "actual_filename": name,
                         "output_file": {"path": str(out_path),
                                         "scope": "machine_local"},
                         "error": None})
            saved.append(key)
        job.report_progress(i + 1, len(keys),
                            f"văn bản {i + 1}/{len(keys)}")
    breakdown = {"succeeded": saved, "failed": failed, "skipped": []}
    data = {"schema_version": SCHEMA_VERSION,
            "destination": {"path": str(dest_dir), "scope": "machine_local",
                            "is_dir": True},
            "documents": docs, "breakdown": breakdown}
    if len(failed) == len(keys):
        raise CommandError(                  # all failed -> job failed §8.4
            "word_batch_failed",
            "Toàn bộ văn bản trong lượt xuất đều lỗi",
            retryable=True, next_action="retry",
            details={"documents": [
                {"document_key": d["document_key"],
                 "code": d["error"]["code"],
                 "message": d["error"]["message"]} for d in docs]})
    return _result("word_export_batch", data, partial=bool(failed))


def _block_message(code):
    return {
        "word.no_assets": "Hồ sơ chưa có tài sản",
        "word.no_landowner": "Chưa có chủ đất trên Diagram",
        "word.no_deceased_landowner": "Không có chủ đất đã chết",
        "word.no_receiver": "Chưa có người nhận",
        "word.too_many_assets": "Quá 5 tài sản",
        "word.too_many_people": "Quá 20 người trên Diagram",
        "word.too_many_signers": "Quá 20 người ký",
        "word.template_missing": "Văn bản chưa có template",
    }.get(code, code)
