"""MIN-69 — Task 4: browser session, download, per-batch prepare.

Contract upload.workflow.v1 §6.4-§6.8/§6.12/§6.14 + §7 against a FAKE
localhost portal (`test/fixtures/upload_portal.py`): real HTTP pages/endpoints
on 127.0.0.1, a ``PortalBrowserSession`` that mimics
``NamDinhUploaderSession`` and records the thread id of every engine call.

Covered:
- one browser thread owns login/download/staff/prepare/poll/close
- confirm/finish_review bound to website+browser+job — wrong_job never
  releases another job's wait
- status queries read snapshots only (never enqueue ops / spawn browser)
- idle browser loop keeps polling login + prepared pages while jobs wait
- prepare: run binding, backend exclusions, chunk_size, dry-run, chi tien
  do prepared (khong ket luan Luu), tab resource facts, partial breakdown
- download_export: completed file only, ISO->dd/mm/yyyy normalization
- browser lost / session expiry / early+late confirm / cancel-stop
- no focus stealing (engine keeps window minimized)

Run from the ``shell/`` directory:
    python test/test_upload_browser_workflow.py
"""

import gc
import json
import os
import sqlite3
import sys
import tempfile
import threading
import time
import unittest
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

_HERE = Path(__file__).resolve().parent
_SHELL = _HERE.parent
sys.path.insert(0, str(_SHELL / "sidecar"))
sys.path.insert(0, str(_HERE / "fixtures"))

from errors import CommandError                      # noqa: E402
from jobstore import JobStore                         # noqa: E402
import command_registry as reg                        # noqa: E402
import engine_roots                                   # noqa: E402
import upload_adapter                                 # noqa: E402
import upload_session                                 # noqa: E402
import upload_workspace                               # noqa: E402
import upload_portal as fixture                       # noqa: E402


V1 = "upload.workflow.v1"
WEBSITE = "fake_portal"
TERMINAL = {"succeeded", "failed", "canceled", "partial"}

CONTRACT_NOS = ["101/2026/CCGD", "102/2026/CCGD", "103/2026/CCGD",
                "104/2026/CCGD", "105/2026/CCGD"]


def _engines_available():
    try:
        engine_roots.engine_root("upload_lab")
        import openpyxl  # noqa: F401
        return True
    except Exception:
        return False


def _write_excel(path, contract_nos, header=True):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if header:
        ws.append(["Số công chứng", "Ngày công chứng"])
    for cn in contract_nos:
        ws.append([cn, "15/04/2026"])
    wb.save(str(path))


@unittest.skipUnless(_engines_available(),
                     "upload_lab engine/openpyxl missing")
class BrowserWorkflowCase(unittest.TestCase):
    """Real JobStore + per-test _BrowserWorker over FakePortal."""

    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.addCleanup(gc.collect)
        self.data_root = Path(self.tempdir.name) / "upload-data"
        self._old_env = os.environ.get("G1_UPLOAD_DATA_DIR")
        os.environ["G1_UPLOAD_DATA_DIR"] = str(self.data_root)
        upload_workspace.reset_store_for_tests()
        self.addCleanup(upload_workspace.reset_store_for_tests)
        self.addCleanup(self._restore_env)

        self.portal = fixture.FakePortal()
        self.addCleanup(self.portal.stop)
        self.provider = fixture.register_fake_portal_website(
            self.portal, website_id=WEBSITE)
        self.addCleanup(self._unregister_fake)

        # Safety net: NEVER let a real NamDinhUploaderSession be created in
        # these tests (TDD-RED can still hit the legacy path). working_dir is
        # redirected to temp so fake registry ops can't touch real data.
        self.fake_wd = Path(self.tempdir.name) / "fake-wd"
        self.fake_wd.mkdir()
        uploader_mod = engine_roots.import_engine_module(
            "upload_lab", "playwright_uploader")
        patcher = mock.patch.object(
            uploader_mod, "NamDinhUploaderSession",
            lambda settings, working_dir: fixture.PortalBrowserSession(
                self.portal, self.fake_wd))
        patcher.start()
        self.addCleanup(patcher.stop)

        self.jobs = JobStore(max_workers=6)
        self.addCleanup(self.jobs.drain)

        self.browser = upload_session._BrowserWorker()
        self.browser._poll_interval = 0.05
        self.addCleanup(self._shutdown_browser)
        worker_patch = mock.patch.object(
            upload_adapter, "worker", lambda: self.browser)
        worker_patch.start()
        self.addCleanup(worker_patch.stop)

        self.store = upload_workspace.open_store()
        self.folder = Path(self.tempdir.name) / "ho so"
        self.folder.mkdir()

    # ------------------------------------------------------------ plumbing
    def _restore_env(self):
        if self._old_env is None:
            os.environ.pop("G1_UPLOAD_DATA_DIR", None)
        else:
            os.environ["G1_UPLOAD_DATA_DIR"] = self._old_env

    def _unregister_fake(self):
        providers = engine_roots.import_engine_module(
            "upload_lab", "providers")
        providers.DEFAULT_REGISTRY._providers.pop(WEBSITE, None)

    def _shutdown_browser(self):
        if getattr(self.browser, "shutdown", None):
            self.browser.shutdown()

    @property
    def session(self):
        """The PortalBrowserSession the worker is currently driving."""
        return self.provider.last_session

    # ------------------------------------------------------------ helpers
    def submit(self, command, payload):
        return self.jobs.submit(
            str(uuid.uuid4()), command, reg.COMMANDS[command], payload)

    def wait_status(self, job, status, timeout=15.0):
        deadline = time.time() + timeout
        while time.time() < deadline:
            snap = job.snapshot()
            if snap["status"] == status:
                return snap
            time.sleep(0.05)
        self.fail(f"job {job.job_id} khong vao {status} "
                  f"(hien tai {job.snapshot()['status']})")

    def wait_terminal(self, job, timeout=15.0):
        deadline = time.time() + timeout
        while time.time() < deadline:
            snap = job.snapshot()
            if snap["status"] in TERMINAL:
                return snap
            time.sleep(0.05)
        self.fail(f"job {job.job_id} khong ket thuc "
                  f"(hien tai {job.snapshot()['status']})")

    def wait_until(self, predicate, timeout=15.0, what="condition"):
        deadline = time.time() + timeout
        while time.time() < deadline:
            if predicate():
                return True
            time.sleep(0.05)
        self.fail(f"timeout cho {what}")

    def run_scan(self, contract_nos=None):
        self.provider.scan_contract_nos = list(
            contract_nos or CONTRACT_NOS[:3])
        job = self.submit("upload.scan", {
            "workflow_version": V1,
            "website_id": WEBSITE,
            "folder": {"path": str(self.folder), "scope": "machine_local"},
            "expected_revision": self.store.revision(),
            "full_rescan": False,
            "modified_since": None,
        })
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        data = snap["result"]["data"]
        return data["run_id"], [r["record_id"] for r in data["records"]]

    def run_audit(self, contract_nos, name="so.xlsx", website_id=WEBSITE):
        excel = Path(self.tempdir.name) / name
        _write_excel(excel, contract_nos)
        job = self.submit("upload.audit_excel", {
            "workflow_version": V1,
            "website_id": website_id,
            "file_ref": {"path": str(excel), "scope": "machine_local"},
            "from_date": "2026-01-01",
            "to_date": "2026-12-31",
        })
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        return snap["result"]["data"]["audit_id"]

    def queue_get(self, run_id, audit_id=None):
        job = self.submit("upload.queue_get", {
            "workflow_version": V1,
            "website_id": WEBSITE,
            "run_id": run_id,
            "audit_id": audit_id,
        })
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        return snap["result"]["data"]

    def v1(self, **kw):
        return {"workflow_version": V1, "website_id": WEBSITE, **kw}

    def session_start(self):
        """Submit upload.session_start; returns the waiting job."""
        job = self.submit("upload.session_start", self.v1(
            expected_revision=self.store.revision()))
        self.wait_status(job, "waiting_user")
        return job

    def confirm(self, browser_id, target_job_id):
        job = self.submit("upload.confirm_login", self.v1(
            browser_id=browser_id, target_job_id=target_job_id))
        return job

    def finish_review(self, browser_id, target_job_id):
        job = self.submit("upload.finish_review", self.v1(
            browser_id=browser_id, target_job_id=target_job_id))
        return job

    def session_status(self, browser_id):
        job = self.submit("upload.session_status", self.v1(
            browser_id=browser_id))
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        return snap["result"]["data"]

    def login(self):
        """Full login flow -> browser_id (job succeeded)."""
        job = self.session_start()
        bid = self.browser.browser_id
        self.assertIsNotNone(bid)
        self.portal.user_login()
        cjob = self.confirm(bid, job.job_id)
        snap = self.wait_terminal(cjob)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        data = snap["result"]["data"]
        self.assertEqual(data["browser_id"], bid)
        self.assertEqual(data["login"]["status"], "authenticated")
        return bid

    def prepare(self, browser_id, run_id, record_ids, *,
                audit_id=None, queue_revision=None, chunk_size=10,
                cong_chung_vien=None, thu_ky=None):
        if queue_revision is None:
            queue_revision = self.store.queue_revision(run_id)
        return self.submit("upload.prepare", self.v1(
            browser_id=browser_id, run_id=run_id, audit_id=audit_id,
            queue_revision=queue_revision, record_ids=list(record_ids),
            chunk_size=chunk_size, cong_chung_vien=cong_chung_vien,
            thu_ky=thu_ky))

    def registry_status(self, record_id):
        site_dir = upload_workspace.website_data_dir(WEBSITE)
        import sqlite3
        conn = sqlite3.connect(str(site_dir / "registry.sqlite3"))
        try:
            row = conn.execute(
                "SELECT status FROM file_registry WHERE id=?",
                (record_id,)).fetchone()
            return row[0] if row else None
        finally:
            conn.close()


# =====================================================================
# Browser thread invariant + session lifecycle
# =====================================================================

def _batch_scan():
    """Module batch_scan that (engine that) — dung identity file that khi
    quet lai, khac fixture scan gia."""
    return engine_roots.import_engine_module("upload_lab", "batch_scan")


@unittest.skipUnless(_engines_available(),
                     "upload_lab engine/openpyxl missing")
class SingleBrowserThreadTest(BrowserWorkflowCase):
    def test_all_engine_calls_run_on_one_browser_thread(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        session = self.session
        self.assertIsNotNone(session)

        # staff refresh (portal) + prepare + poll + download + close
        so = self.submit("upload.staff_options", self.v1(
            browser_id=bid, refresh=True))
        snap = self.wait_terminal(so)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        self.assertEqual(snap["result"]["data"]["source"], "portal")
        self.assertEqual(
            snap["result"]["data"]["cong_chung_vien"],
            self.portal.staff["cong_chung_vien"])

        pjob = self.prepare(bid, run_id, ids)
        self.wait_status(pjob, "waiting_user")
        self.assertEqual(
            sorted(self.session_status(bid)["tabs"]["open_record_ids"]),
            sorted(ids))
        self.session.user_save(ids[0])
        self.session.user_close_tab(ids[1])
        self.wait_until(
            lambda: ids[1] in self.session_status(bid)
            ["tabs"]["closed_record_ids"],
            what="poll phat hien tab dong (tai nguyen)")
        fj = self.finish_review(bid, pjob.job_id)
        self.wait_terminal(fj)
        self.assertEqual(self.wait_terminal(pjob)["status"], "succeeded")

        dl = self.submit("upload.download_export", self.v1(
            browser_id=bid, from_date="2026-01-01", to_date="2026-09-30"))
        snap = self.wait_terminal(dl)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        self.assertTrue(
            Path(snap["result"]["data"]["file_ref"]["path"]).is_file())

        cj = self.submit("upload.session_close", self.v1(browser_id=bid))
        snap = self.wait_terminal(cj)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        self.assertTrue(snap["result"]["data"]["closed"])

        # THE invariant: every engine-facing call ran on one thread — the
        # worker's browser thread (login/download/staff/prepare/poll/close).
        idents = {t for _name, t in session.thread_log}
        self.assertEqual(len(idents), 1,
                         f"calls spread over threads: {session.thread_log}")
        self.assertEqual(idents.pop(), self.browser._thread.ident)
        # and all the expected operations really happened
        for op in ("open_manual_login", "fetch_staff_options",
                   "prepare_manifest", "poll_prepared_pages",
                   "download_contract_book_export", "close"):
            self.assertIn(op, session.calls)

    def test_idle_poll_keeps_running_while_job_waits(self):
        """Poll tiếp trong lúc chờ người review — status cập nhật mà không
        cần finish_review trước."""
        run_id, ids = self.run_scan()
        bid = self.login()
        pjob = self.prepare(bid, run_id, ids[:2])
        self.wait_status(pjob, "waiting_user")
        self.session.user_save(ids[0])
        self.session.user_close_tab(ids[1])
        # Idle poll must report the closed tab as a resource fact while the
        # job keeps waiting — khong suy ket qua Luu tu thao tac nguoi dung.
        self.wait_until(
            lambda: ids[1] in self.session_status(bid)
            ["tabs"]["closed_record_ids"],
            what="closed_record_ids cap nhat qua idle poll")
        self.assertEqual(pjob.snapshot()["status"], "waiting_user")

    def test_login_latched_authenticated_across_idle_polls(self):
        """Engine tra 'authenticated' MOT LAN (one-shot) roi 'idle' mai —
        snapshot phai giu authenticated sau nhieu nhip idle poll, va
        session_start khong the tieu thu nham one-shot do."""
        bid = self.login()
        # Nhieu nhip idle poll sau khi one-shot da bi tieu thu.
        time.sleep(0.3)  # ~6 poll ticks o 0.05s
        data = self.session_status(bid)
        self.assertEqual(data["login"]["status"], "authenticated")
        # Chung minh raw engine da la 'idle' — latch la thu giu status,
        # khong phai poll lap lai authenticated.
        self.assertEqual(
            self.session.poll_manual_login()["status"], "idle")
        time.sleep(0.15)
        data = self.session_status(bid)
        self.assertEqual(data["login"]["status"], "authenticated")
        # Op can login van chay duoc sau hang chuc nhip poll 'idle'.
        dl = self.submit("upload.download_export", self.v1(
            browser_id=bid, from_date="2026-01-01", to_date="2026-09-30"))
        self.assertEqual(self.wait_terminal(dl)["status"], "succeeded")

    def test_session_status_reads_snapshot_never_spawns_browser(self):
        # 1) Unknown browser before any session: no spawn, scope_violation.
        job = self.submit("upload.session_status", self.v1(
            browser_id="brw_nope"))
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "scope_violation")
        self.assertIsNone(self.browser._thread,
                          "status query must not spawn browser thread")

        # 2) After login: status returns snapshot without ANY browser
        # work. Idle poll van ghi session.calls moi nhip → dong bang no
        # bang _poll_interval lon roi snapshot tap call TRUOC/SAU: status
        # impl nao enqueue op (ke ca op poll gian tiep) deu bi lo.
        run_id, ids = self.run_scan()
        bid = self.login()
        session = self.session
        self.prepare(bid, run_id, ids[:1])
        self.wait_until(lambda: "prepare_manifest" in session.calls,
                        what="prepare chay xong")
        self.browser._poll_interval = 600
        time.sleep(0.25)  # cho nhip poll dang cho ket thuc
        session.calls.clear()
        before = list(session.calls)
        data = self.session_status(bid)
        self.assertEqual(data["login"]["status"], "authenticated")
        self.assertEqual(
            list(session.calls), before,
            f"session_status cham vao browser thread: {session.calls}")
        self.browser._poll_interval = 0.05

    def test_one_mutating_browser_op_at_a_time(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        session = self.session
        session.prepare_gate = threading.Event()  # closed gate

        pjob = self.prepare(bid, run_id, ids)
        self.wait_until(lambda: "prepare_manifest" in session.calls,
                        what="prepare da vao browser thread")
        # While prepare_manifest holds the browser, another mutating op
        # must be rejected immediately — not queued behind it.
        dl = self.submit("upload.download_export", self.v1(
            browser_id=bid, from_date="2026-01-01", to_date="2026-09-30"))
        snap = self.wait_terminal(dl)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "browser_busy")
        self.assertTrue(snap["error"]["retryable"])
        session.prepare_gate.set()
        self.wait_status(pjob, "waiting_user")


# =====================================================================
# Per-job confirmation identity
# =====================================================================

@unittest.skipUnless(_engines_available(),
                     "upload_lab engine/openpyxl missing")
class WaitIdentityTest(BrowserWorkflowCase):
    def test_confirm_releases_only_its_own_job(self):
        run_id, ids = self.run_scan()
        bid = self.login()

        # Job A: a second session_start -> waiting login on same browser.
        job_a = self.session_start()
        # Job B: prepare -> waiting review on same browser.
        job_b = self.prepare(bid, run_id, ids)
        self.wait_status(job_b, "waiting_user")

        self.portal.user_login()  # keep portal authed for A's forced poll
        ca = self.confirm(bid, job_a.job_id)
        self.assertEqual(self.wait_terminal(ca)["status"], "succeeded")
        self.assertEqual(self.wait_terminal(job_a)["status"], "succeeded")
        # Job B was not woken by A's confirmation.
        self.assertEqual(job_b.snapshot()["status"], "waiting_user")

        # Wrong-kind confirm for B -> wrong_job, still waiting.
        cb = self.confirm(bid, job_b.job_id)
        snap = self.wait_terminal(cb)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "wrong_job")
        self.assertEqual(job_b.snapshot()["status"], "waiting_user")

        # finish_review targeting the (already terminal) job A -> wrong_job.
        fb = self.finish_review(bid, job_a.job_id)
        snap = self.wait_terminal(fb)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "wrong_job")
        self.assertEqual(job_b.snapshot()["status"], "waiting_user")

        # Correct finish_review for B releases exactly B.
        fj = self.finish_review(bid, job_b.job_id)
        self.assertEqual(self.wait_terminal(fj)["status"], "succeeded")
        self.assertEqual(self.wait_terminal(job_b)["status"], "succeeded")

    def test_legacy_confirm_never_releases_versioned_wait(self):
        """confirm_login/finish_review LEGACY (payload khong
        workflow_version → release_any_wait) chi giai phong wait khong
        scope — job versioned dang cho KHONG bi danh thuc nham."""
        run_id, ids = self.run_scan()
        bid = self.login()
        pjob = self.prepare(bid, run_id, ids[:1])
        self.wait_status(pjob, "waiting_user")

        # Legacy finish_review → release_any_wait("review") khong scope.
        legacy = self.submit("upload.finish_review", {})
        self.assertEqual(self.wait_terminal(legacy)["status"], "succeeded")
        self.assertEqual(pjob.snapshot()["status"], "waiting_user")

        # Legacy confirm_login → release_any_wait("login") khong scope.
        legacy = self.submit("upload.confirm_login", {})
        self.assertEqual(self.wait_terminal(legacy)["status"], "succeeded")
        self.assertEqual(pjob.snapshot()["status"], "waiting_user")

        # Confirm versioned dung scope van giai phong binh thuong.
        fj = self.finish_review(bid, pjob.job_id)
        self.assertEqual(self.wait_terminal(fj)["status"], "succeeded")
        self.assertEqual(self.wait_terminal(pjob)["status"], "succeeded")

    def test_release_any_wait_only_touches_legacy_scope(self):
        """Unit-level: release_any_wait bo qua wait co scope day du."""
        w = self.browser
        ev_versioned = w.register_wait(
            "j-versioned", waiting_on="login",
            website_id=WEBSITE, browser_id="brw_x")
        ev_legacy = w.register_wait(
            "j-legacy", waiting_on="login",
            website_id=None, browser_id=None)
        w.release_any_wait("login")
        self.assertFalse(ev_versioned.is_set())
        self.assertTrue(ev_legacy.is_set())

    def test_confirm_early_and_late(self):
        """Bấm xác nhận sớm (job chưa vào waiting) -> wrong_job và KHÔNG
        giải phóng; bấm đúng lúc sau đó vẫn giải phóng bình thường."""
        # Arm the gate BEFORE the session exists: open_manual_login blocks
        # inside the browser op so the job stays 'running' (not waiting).
        gate = threading.Event()
        self.provider.next_session_open_gate = gate
        job = self.submit("upload.session_start", self.v1(
            expected_revision=self.store.revision()))
        self.wait_until(lambda: self.provider.last_session is not None
                        and "open_manual_login" in
                        self.provider.last_session.calls,
                        what="open_manual_login duoc goi")
        self.assertEqual(job.snapshot()["status"], "running")

        bid = self.browser.browser_id
        early = self.confirm(bid, job.job_id)
        snap = self.wait_terminal(early)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "wrong_job")
        self.assertEqual(job.snapshot()["status"], "running")

        gate.set()  # release the fake's open_manual_login
        self.wait_status(job, "waiting_user")
        self.portal.user_login()
        late = self.confirm(bid, job.job_id)
        self.assertEqual(self.wait_terminal(late)["status"], "succeeded")
        self.assertEqual(self.wait_terminal(job)["status"], "succeeded")

        # Stale confirm for the now-terminal job -> wrong_job.
        stale = self.confirm(bid, job.job_id)
        snap = self.wait_terminal(stale)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "wrong_job")

    def test_confirm_wrong_website_and_browser_rejected(self):
        self.login()
        bid = self.browser.browser_id
        job = self.session_start()
        # wrong website: browser belongs to fake_portal
        bad_site = self.submit("upload.confirm_login", {
            "workflow_version": V1, "website_id": "nam_dinh",
            "browser_id": bid, "target_job_id": job.job_id})
        snap = self.wait_terminal(bad_site)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "website_mismatch")
        # unknown browser id -> scope_violation
        bad_browser = self.confirm("brw_khongco", job.job_id)
        snap = self.wait_terminal(bad_browser)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "scope_violation")
        self.assertEqual(job.snapshot()["status"], "waiting_user")
        # cleanup: release the real wait
        self.portal.user_login()
        self.confirm(bid, job.job_id)
        self.wait_terminal(job)


# =====================================================================
# Prepare: run binding, exclusions, chunking, dry-run, saves
# =====================================================================

@unittest.skipUnless(_engines_available(),
                     "upload_lab engine/openpyxl missing")
class PrepareTest(BrowserWorkflowCase):
    def test_empty_record_ids_validation_error(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        job = self.prepare(bid, run_id, [])
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "validation_error")

    def test_record_outside_run_scope_violation(self):
        """N+1: id khong thuoc run -> scope_violation truoc khi mo tab."""
        run_id, ids = self.run_scan()
        bid = self.login()
        job = self.prepare(bid, run_id, ids + [999999])
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "scope_violation")
        self.assertEqual(self.session.tabs, {})

    def test_same_file_rescan_moves_rows_and_stale_run_ids_rejected(self):
        """Quet lai CUNG file (identity that cua batch_scan) chuyen row sang
        run moi: 'Tiep tuc' tren run cu voi ids goc bi scope_violation
        truoc khi engine mo tab nao."""
        data_dir = upload_workspace.website_data_dir(WEBSITE)
        folder = Path(self.tempdir.name) / "hs-real"
        from docx import Document

        for name, no in (("HD-701", "701/2026/CCGD"),
                         ("HD-702", "702/2026/CCGD")):
            folder.mkdir(parents=True, exist_ok=True)
            doc = Document()
            doc.add_paragraph(f"Số công chứng {no}")
            doc.save(str(folder / f"{name}.docx"))

        manifest_a = _batch_scan().run_batch_scan(
            folder, working_dir=data_dir)
        manifest_a_path = _batch_scan().finalize_manifest(
            manifest_a, data_dir / "runs")
        run_a = upload_workspace.register_run(
            WEBSITE, manifest_a_path)["run_id"]
        ids_a = self._run_record_ids(data_dir, run_a)
        self.assertEqual(len(ids_a), 2, ids_a)

        bid = self.login()
        job = self.prepare(bid, run_a, ids_a)
        self.wait_status(job, "waiting_user")
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        self.wait_terminal(job)
        tabs_before = sorted(self.session.tabs)
        prepare_calls_before = self.session.prepare_calls

        # Quet lai cung folder: file_key = path|mtime|size (identity that)
        # nen upsert chuyen chinh cac row do sang run B.
        time.sleep(1.1)  # ten manifest co do phan giai 1 giay
        manifest_b = _batch_scan().run_batch_scan(
            folder, working_dir=data_dir)
        manifest_b_path = _batch_scan().finalize_manifest(
            manifest_b, data_dir / "runs")
        run_b = upload_workspace.register_run(
            WEBSITE, manifest_b_path)["run_id"]
        self.assertNotEqual(run_a, run_b)
        self.assertEqual(self._run_record_ids(data_dir, run_a), [])
        self.assertEqual(
            sorted(self._run_contract_nos(data_dir, run_b)),
            ["701/2026/CCGD", "702/2026/CCGD"])

        stale = self.prepare(bid, run_a, ids_a, chunk_size=2)
        snap = self.wait_terminal(stale)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "scope_violation")
        self.assertEqual(
            self.session.prepare_calls, prepare_calls_before,
            "scope_violation van goi engine mo tab")
        self.assertEqual(sorted(self.session.tabs), tabs_before)

    def _run_record_ids(self, data_dir, run_id):
        conn = sqlite3.connect(str(data_dir / "registry.sqlite3"))
        try:
            conn.row_factory = sqlite3.Row
            rows = _batch_scan().fetch_registry_records_for_run(
                conn, run_id)
            return sorted(int(r["id"]) for r in rows)
        finally:
            conn.close()

    def _run_contract_nos(self, data_dir, run_id):
        conn = sqlite3.connect(str(data_dir / "registry.sqlite3"))
        try:
            conn.row_factory = sqlite3.Row
            rows = _batch_scan().fetch_registry_records_for_run(
                conn, run_id)
            return [str(r["contract_no"]) for r in rows]
        finally:
            conn.close()

    def test_queue_revision_stale(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        job = self.prepare(bid, run_id, ids, queue_revision=999)
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "stale_revision")
        self.assertTrue(snap["error"]["retryable"])

    def test_one_record_full_cycle(self):
        run_id, ids = self.run_scan(CONTRACT_NOS[:1])
        bid = self.login()
        job = self.prepare(bid, run_id, ids)
        self.wait_status(job, "waiting_user")
        status = self.session_status(bid)
        self.assertEqual(status["tabs"]["open_record_ids"], ids)
        self.session.user_save(ids[0])
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded")
        data = snap["result"]["data"]
        # Ket qua chi mo ta tien do chuan bi + tai nguyen: khong co ket qua
        # Luu, tab nguoi dung da Luu van duoc bao la dang mo.
        self.assertNotIn("saved_record_ids", data)
        self.assertNotIn("needs_reconcile_record_ids", data)
        self.assertEqual(data["summary"]["prepared_count"], 1)
        self.assertEqual(data["summary"]["open_record_ids"], ids)
        self.assertEqual(
            [b["record_id"] for b in data["breakdown"]["succeeded"]], ids)
        self.assertEqual(
            [b["stage"] for b in data["breakdown"]["succeeded"]], ["prepared"])
        # Khong ghi ket qua Luu vao registry.
        self.assertEqual(self.registry_status(ids[0]), "prepared_dry_run")

    def test_n_records_and_chunk_cap(self):
        """N ho so, chunk_size nho hon N -> toi da chunk tab mo mot dot."""
        run_id, ids = self.run_scan(CONTRACT_NOS)
        bid = self.login()
        job = self.prepare(bid, run_id, ids, chunk_size=2)
        self.wait_status(job, "waiting_user")
        kwargs = self.session.last_prepare_kwargs
        self.assertEqual(kwargs["chunk_size"], 2)
        self.assertEqual(sorted(self.session.tabs), sorted(ids[:2]))
        status = self.session_status(bid)
        self.assertEqual(status["tabs"]["open_record_ids"], sorted(ids[:2]))
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        snap = self.wait_terminal(job)
        data = snap["result"]["data"]
        self.assertEqual(data["summary"]["prepared_count"], 2)
        self.assertEqual(data["summary"]["remaining"], 3)

    def test_audit_excludes_existing_contract_nos(self):
        """Backend tinh exclusion tu audit that — so da co tren web bi loai
        khoi batch ke ca khi renderer van gui record_id do."""
        run_id, ids = self.run_scan()
        audit_id = self.run_audit([CONTRACT_NOS[0], CONTRACT_NOS[1]])
        queue = self.queue_get(run_id, audit_id)
        bid = self.login()
        job = self.prepare(bid, run_id, ids, audit_id=audit_id,
                           queue_revision=queue["queue_revision"])
        self.wait_status(job, "waiting_user")
        kwargs = self.session.last_prepare_kwargs
        self.assertIn("101/2026", kwargs["exclude_contract_nos"])
        self.assertIn("102/2026", kwargs["exclude_contract_nos"])
        self.assertEqual(sorted(self.session.tabs), [ids[2]])
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        snap = self.wait_terminal(job)
        data = snap["result"]["data"]
        self.assertEqual(data["summary"]["excluded_duplicates"], 2)

    def test_all_records_excluded_returns_without_wait(self):
        """0 ho so con lai sau exclusion -> ket qua ngay, khong waiting."""
        run_id, ids = self.run_scan(CONTRACT_NOS[:2])
        audit_id = self.run_audit(list(CONTRACT_NOS[:2]))
        queue = self.queue_get(run_id, audit_id)
        bid = self.login()
        job = self.prepare(bid, run_id, ids, audit_id=audit_id,
                           queue_revision=queue["queue_revision"])
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded")
        data = snap["result"]["data"]
        self.assertEqual(data["summary"]["prepared_count"], 0)
        self.assertEqual(data["summary"]["open_record_ids"], [])
        self.assertEqual(self.session.tabs, {})

    def test_zero_padded_contract_excluded_consistently(self):
        """Record '0101/2026/CCGD' + Excel '101/2026': queue danh 'da co
        trong Excel' bang canonical (cat so 0) thi prepare PHAI loai no
        khoi batch — normalize tho giu so 0 o mot phia se bo sot va
        upload trung len web."""
        run_id, ids = self.run_scan(["0101/2026/CCGD", "0102/2026/CCGD"])
        audit_id = self.run_audit(["101/2026"])
        queue = self.queue_get(run_id, audit_id)
        rows = {r["contract_no"]: r for r in queue["folder_rows"]}
        self.assertEqual(
            rows["0101/2026/CCGD"]["ghi_chu"], "da co trong Excel")
        self.assertFalse(rows["0101/2026/CCGD"]["selected"])
        self.assertTrue(rows["0102/2026/CCGD"]["selected"])

        bid = self.login()
        job = self.prepare(bid, run_id, ids, audit_id=audit_id,
                           queue_revision=queue["queue_revision"])
        self.wait_status(job, "waiting_user")
        # Record dem so 0 KHONG mo tab — chi 0102 duoc chuan bi; exclude
        # set gui cho engine mang dang zero-keeping no se tu tinh.
        self.assertEqual(sorted(self.session.tabs), [ids[1]])
        kwargs = self.session.last_prepare_kwargs
        self.assertEqual(kwargs["selected_record_ids"], {ids[1]})
        self.assertIn("0101/2026", kwargs["exclude_contract_nos"])
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        data = snap["result"]["data"]
        self.assertEqual(data["summary"]["excluded_duplicates"], 1)
        self.assertEqual(data["summary"]["prepared_count"], 1)

    def test_one_item_fails_partial_breakdown(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        self.session.fail_record_ids = {ids[1]}
        job = self.prepare(bid, run_id, ids)
        self.wait_status(job, "waiting_user")
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "partial")
        data = snap["result"]["data"]
        failed = data["breakdown"]["failed"]
        self.assertEqual(len(failed), 1)
        self.assertEqual(failed[0]["record_id"], ids[1])
        self.assertEqual(failed[0]["stage"], "prepared")
        succeeded_ids = [b["record_id"] for b in
                         data["breakdown"]["succeeded"]]
        self.assertEqual(sorted(succeeded_ids), [ids[0], ids[2]])

    def test_tab_closed_without_save_is_resource_fact_only(self):
        """Tab dong tay chi la su that tai nguyen: khong ket luan Luu, khong
        ghi store, khong ha trang thai da chuan bi cua ho so."""
        run_id, ids = self.run_scan()
        bid = self.login()
        job = self.prepare(bid, run_id, ids[:2])
        self.wait_status(job, "waiting_user")
        self.session.user_close_tab(ids[1])
        self.wait_until(
            lambda: ids[1] in self.session_status(bid)
            ["tabs"]["closed_record_ids"],
            what="poll phat hien tab dong")
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded")
        data = snap["result"]["data"]
        self.assertNotIn("needs_reconcile_record_ids", data)
        self.assertNotIn("saved_record_ids", data)
        # Ho so van duoc coi la DA CHUAN BI trong luot (khong bi phat lai).
        succeeded = [b["record_id"] for b in data["breakdown"]["succeeded"]]
        self.assertEqual(sorted(succeeded), sorted(ids[:2]))
        self.assertEqual(self.registry_status(ids[1]), "prepared_dry_run")
        # Chi tab da dong mat quyen so huu; tab con mo van duoc theo doi.
        self.assertEqual(self.store.open_tab_record_ids(WEBSITE), [ids[0]])

    def test_finish_review_keeps_tabs_and_writes_nothing(self):
        """'Xong kiem tra' chi ket thuc buoc cho review: tab van mo, khong
        ghi ket qua Luu vao registry."""
        run_id, ids = self.run_scan()
        bid = self.login()
        job = self.prepare(bid, run_id, ids[:1])
        self.wait_status(job, "waiting_user")
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        snap = self.wait_terminal(job)
        data = snap["result"]["data"]
        self.assertNotIn("saved_record_ids", data)
        self.assertEqual(data["summary"]["open_record_ids"], ids[:1])
        self.assertEqual(self.registry_status(ids[0]), "prepared_dry_run")
        self.assertEqual(self.store.open_tab_record_ids(WEBSITE), ids[:1])

    def test_continue_batch_never_reopens_prepared_records(self):
        """AC2: 'Tiep tuc' chi mo ho so da chon CHUA chuan bi; 5 ho so da
        chuan bi nhung chua Luu khong tu quay lai lo sau."""
        run_id, ids = self.run_scan([
            "101/2026/CCGD", "102/2026/CCGD", "103/2026/CCGD",
            "104/2026/CCGD", "105/2026/CCGD", "106/2026/CCGD",
            "107/2026/CCGD", "108/2026/CCGD"])
        bid = self.login()
        job = self.prepare(bid, run_id, ids, chunk_size=4)
        self.wait_status(job, "waiting_user")
        # Nguoi dung Luu 2 ho so, dong 1 tab chua Luu, de 1 tab mo.
        self.session.user_save(ids[0])
        self.session.user_save(ids[1])
        self.session.user_close_tab(ids[2])
        self.wait_until(
            lambda: ids[2] in self.session_status(bid)
            ["tabs"]["closed_record_ids"],
            what="tab chua Luu da dong")
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        self.wait_terminal(job)

        queue = self.queue_get(run_id)
        queued_ids = [r["record_id"] for r in queue["folder_rows"]]
        self.assertIn(ids[0], queued_ids)

        open_before = set(
            self.session_status(bid)["tabs"]["open_record_ids"])
        job2 = self.prepare(bid, run_id, ids, chunk_size=4)
        self.wait_status(job2, "waiting_user")
        fj2 = self.finish_review(bid, job2.job_id)
        self.wait_terminal(fj2)
        snap2 = self.wait_terminal(job2)
        data2 = snap2["result"]["data"]
        # Lo sau chi mo ho so chua chuan bi: khong co tab moi nao cho cac
        # ID da chuan bi (ke ca ho so da Luu va ho so tab da dong).
        open_after = set(
            self.session_status(bid)["tabs"]["open_record_ids"])
        self.assertEqual(sorted(open_after - open_before), ids[4:])
        self.assertEqual(data2["summary"]["prepared_count"], 4)
        self.assertEqual(data2["summary"]["remaining"], 0)
        for rid in ids:
            self.assertEqual(self.registry_status(rid), "prepared_dry_run")

    def test_cancel_during_prepare_releases_browser(self):
        """Dung dat stop event; engine dung sau don vi dang xu ly va nha
        browser cho dot moi."""
        run_id, ids = self.run_scan(CONTRACT_NOS)
        bid = self.login()
        self.session.record_delay_s = 0.25
        job = self.prepare(bid, run_id, ids)
        self.wait_until(
            lambda: self.session.calls.count("prepare_manifest") >= 1,
            what="prepare bat dau")
        time.sleep(0.35)
        self.jobs.cancel(job.job_id)
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "canceled")
        # browser released -> next mutating op works, not browser_busy
        dl = self.submit("upload.download_export", self.v1(
            browser_id=bid, from_date="2026-01-01", to_date="2026-09-30"))
        snap = self.wait_terminal(dl)
        self.assertEqual(snap["status"], "succeeded", snap["error"])



# =====================================================================
# Final review F1–F4 — gates chi tren duong v1
# =====================================================================

@unittest.skipUnless(_engines_available(),
                     "upload_lab engine/openpyxl missing")
class FinalReviewV1Test(BrowserWorkflowCase):
    def test_scan_cancel_stops_engine_mid_scan(self):
        """F4: cancel giua scan phai cat engine ngay — CancelledByUser la
        Exception bi emit_progress cua engine nuot; adapter doi sang
        _ScanAbortRequested (BaseException) de thoat ra ngoai; job ket
        thuc 'canceled' (khong 'failed'), engine khong chay het folder."""
        emitted = []

        def looping_scan(folder, data_dir, **kw):
            cb = kw.get("progress_callback")
            for i in range(50):
                emitted.append(i)
                if cb:
                    cb({"total_files": 50,
                        "processed_files": i + 1,
                        "current_file": f"HD-{i}.docx"})
                time.sleep(0.02)
            raise AssertionError(
                "scan khong bi cat — cancel khong thoat emit_progress")

        self.provider.run_scan = looping_scan
        job = self.submit("upload.scan", self.v1(
            folder={"path": str(self.folder), "scope": "machine_local"},
            expected_revision=self.store.revision(),
            full_rescan=False, modified_since=None))
        self.wait_until(
            lambda: job.snapshot()["progress"] is not None,
            what="scan da bat dau emit progress")
        self.jobs.cancel(job.job_id)
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "canceled", snap["error"])
        self.assertLess(len(emitted), 50)
        self.assertEqual(
            self.store.job_for(job.job_id)["status"], "canceled")

    def test_reconcile_command_is_retired(self):
        """Command/capability reconcile da bo: registry khong con command
        (consumer nhan command_unknown) va catalog het capability."""
        self.assertNotIn("upload.reconcile", reg.COMMANDS)
        wj = self.submit("upload.websites", self.v1())
        snap = self.wait_terminal(wj)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        for site in snap["result"]["data"]["websites"]:
            self.assertNotIn("reconcile", site["capabilities"])


# =====================================================================
# Session problems: expiry, browser lost, close semantics
# =====================================================================

@unittest.skipUnless(_engines_available(),
                     "upload_lab engine/openpyxl missing")
class SessionProblemTest(BrowserWorkflowCase):
    def test_expired_session_prepare_fails_structured(self):
        """Session portal het han (engine RuntimeError 'Session da het
        han') → upload.login_not_confirmed + login_required — UI dua
        dang nhap lai thay vi retry mu."""
        run_id, ids = self.run_scan()
        bid = self.login()
        self.portal.expire()
        job = self.prepare(bid, run_id, ids)
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "upload.login_not_confirmed")
        self.assertTrue(snap["error"]["retryable"])
        self.assertEqual(snap["error"]["next_action"], "login_required")
        self.assertEqual(self.session.tabs, {})
        # Latch da bi bo: status khong con bao authenticated.
        self.assertEqual(self.browser.snapshot()["login"]["status"],
                         "unknown")

    def test_download_requires_login(self):
        run_id, ids = self.run_scan()
        # session open but NOT logged in yet
        job = self.session_start()
        bid = self.browser.browser_id
        dl = self.submit("upload.download_export", self.v1(
            browser_id=bid, from_date="2026-01-01", to_date="2026-09-30"))
        snap = self.wait_terminal(dl)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "upload.login_not_confirmed")
        self.portal.user_login()
        self.confirm(bid, job.job_id)
        self.wait_terminal(job)

    def test_browser_lost_mid_review_fails_job_and_frees_resource(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        job = self.prepare(bid, run_id, ids[:2])
        self.wait_status(job, "waiting_user")
        self.session.crash()
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "engine_unavailable")
        # Loi ha tang thuong: retryable, khong khoa vi ket qua Luu chua ro.
        self.assertTrue(snap["error"]["retryable"])
        self.assertEqual(snap["error"]["next_action"], "retry")
        details = snap["error"].get("details") or {}
        self.assertNotIn("needs_reconcile_record_ids", details)
        # Tab mat chu so huu: store khong giu quyen so huu nua, registry
        # giu nguyen tien do da chuan bi (khong suy ket qua Luu).
        self.assertEqual(self.store.open_tab_record_ids(WEBSITE), [])
        for rid in ids[:2]:
            self.assertEqual(self.registry_status(rid), "prepared_dry_run")

    def test_session_close_reports_resource_ids_and_never_saves(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        job = self.prepare(bid, run_id, ids[:2])
        self.wait_status(job, "waiting_user")
        self.session.user_save(ids[0])
        self.session.user_close_tab(ids[1])
        self.wait_until(
            lambda: ids[1] in self.session_status(bid)
            ["tabs"]["closed_record_ids"],
            what="tab dong truoc khi close phien")
        cj = self.submit("upload.session_close", self.v1(browser_id=bid))
        snap = self.wait_terminal(cj)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        data = snap["result"]["data"]
        self.assertTrue(data["closed"])
        self.assertIn(ids[1], data["closed_record_ids"])
        self.assertNotIn("verified_record_ids", data)
        self.assertNotIn("needs_reconcile_record_ids", data)
        # Khong ghi ket qua Luu nao; quyen so huu tab da duoc xa.
        for rid in ids[:2]:
            self.assertEqual(self.registry_status(rid), "prepared_dry_run")
        self.assertEqual(self.store.open_tab_record_ids(WEBSITE), [])
        # the waiting prepare job dies with the browser
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")

    def test_website_mismatch_other_website_ops(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        # session bound to fake_portal; ops for nam_dinh -> website_mismatch
        job = self.submit("upload.session_status", {
            "workflow_version": V1, "website_id": "nam_dinh",
            "browser_id": bid})
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "website_mismatch")
        # legacy op against versioned session -> workflow_conflict
        with self.assertRaises(CommandError) as ctx:
            self.browser.call("poll_manual_login")
        self.assertEqual(ctx.exception.code, "workflow_conflict")


# =====================================================================
# Download + staff options
# =====================================================================

@unittest.skipUnless(_engines_available(),
                     "upload_lab engine/openpyxl missing")
class DownloadStaffTest(BrowserWorkflowCase):
    def test_download_export_completed_file_and_date_normalization(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        job = self.submit("upload.download_export", self.v1(
            browser_id=bid, from_date="2026-01-01", to_date="2026-09-30"))
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        data = snap["result"]["data"]
        ref = data["file_ref"]
        path = Path(ref["path"])
        self.assertTrue(path.is_file())
        self.assertEqual(ref["scope"], "machine_local")
        self.assertEqual(ref["size_bytes"], path.stat().st_size)
        self.assertEqual(
            ref["sha256"], upload_workspace.sha256_file(path))
        self.assertEqual(data["from_date"], "2026-01-01")
        self.assertEqual(data["to_date"], "2026-09-30")
        # ISO wire dates normalized to dd/mm/yyyy at the adapter boundary
        self.assertEqual(self.session.last_download_dates,
                         ("01/01/2026", "30/09/2026"))
        self.assertEqual(self.portal.export_requests,
                         [("01/01/2026", "30/09/2026")])

    def test_interrupted_download_not_returned(self):
        run_id, ids = self.run_scan()
        bid = self.login()
        self.portal.fail_next_export = True
        dl_dir = upload_workspace.website_data_dir(WEBSITE) / "downloads"
        before = set(dl_dir.glob("*.xlsx")) if dl_dir.is_dir() else set()
        job = self.submit("upload.download_export", self.v1(
            browser_id=bid, from_date="2026-01-01", to_date="2026-09-30"))
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        after = set(dl_dir.glob("*.xlsx")) if dl_dir.is_dir() else set()
        self.assertEqual(before, after, "partial download left a file")
        self.assertTrue(self.portal.export_requests)

    def test_staff_options_refresh_and_cache(self):
        run_id, ids = self.run_scan()
        # cache read without browser -> empty cache, source cache
        job = self.submit("upload.staff_options", self.v1(
            browser_id=None, refresh=False))
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        self.assertEqual(snap["result"]["data"]["source"], "cache")

        # refresh needs browser_id
        job = self.submit("upload.staff_options", self.v1(
            browser_id=None, refresh=True))
        snap = self.wait_terminal(job)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "validation_error")

        # refresh before login -> login_not_confirmed
        job = self.session_start()
        bid = self.browser.browser_id
        so = self.submit("upload.staff_options", self.v1(
            browser_id=bid, refresh=True))
        snap = self.wait_terminal(so)
        self.assertEqual(snap["status"], "failed")
        self.assertEqual(snap["error"]["code"], "upload.login_not_confirmed")
        self.portal.user_login()
        self.confirm(bid, job.job_id)
        self.wait_terminal(job)

        so = self.submit("upload.staff_options", self.v1(
            browser_id=bid, refresh=True))
        snap = self.wait_terminal(so)
        self.assertEqual(snap["status"], "succeeded", snap["error"])
        data = snap["result"]["data"]
        self.assertEqual(data["source"], "portal")
        self.assertEqual(data["cong_chung_vien"],
                         self.portal.staff["cong_chung_vien"])
        self.assertTrue(data["fetched_at"])

        # cache read now returns the fetched list for this website
        job = self.submit("upload.staff_options", self.v1(
            browser_id=None, refresh=False))
        snap = self.wait_terminal(job)
        data = snap["result"]["data"]
        self.assertEqual(data["source"], "cache")
        self.assertEqual(data["cong_chung_vien"],
                         self.portal.staff["cong_chung_vien"])
        self.assertTrue(data["fetched_at"])


# =====================================================================
# No focus steal
# =====================================================================

@unittest.skipUnless(_engines_available(),
                     "upload_lab engine/openpyxl missing")
class FocusTest(BrowserWorkflowCase):
    def test_prepare_never_steals_focus(self):
        run_id, ids = self.run_scan(CONTRACT_NOS[:2])
        bid = self.login()
        job = self.prepare(bid, run_id, ids)
        self.wait_status(job, "waiting_user")
        self.assertTrue(self.session.focus_events,
                        "fake session should record window handling")
        self.assertTrue(
            all(ev == "keep_minimized" for ev in self.session.focus_events),
            f"focus steal events: {self.session.focus_events}")
        fj = self.finish_review(bid, job.job_id)
        self.wait_terminal(fj)
        self.wait_terminal(job)


# =====================================================================
# Engine-level F1/F2 — real NamDinhUploaderSession + fake page
# =====================================================================

class _EmptyLoc:
    """Locator count()==0 — get_by_role/get_by_text/locator miss."""

    @property
    def first(self):
        return self

    def count(self):
        return 0

    def nth(self, _index):
        return self

    def locator(self, *_a, **_k):
        return self

    def is_visible(self):
        return False


class _ComboLoc:
    """Input combobox cua ten_hop_dong — gia tri doc qua input_value."""

    def __init__(self):
        self.value = ""
        self.clicks = 0
        self.presses = []
        self.type_values = []

    @property
    def first(self):
        return self

    def count(self):
        return 1

    def evaluate(self, script):
        if "tagName" in script:
            return "input"
        return None

    def input_value(self):
        return self.value

    def click(self, **_kw):
        self.clicks += 1

    def fill(self, value):
        self.value = value

    def type(self, value, delay=0):
        self.value += value
        self.type_values.append(value)

    def press(self, key):
        self.presses.append(key)
        if key == "Backspace" and "Control+A" in self.presses:
            self.value = ""


class _SaveLoc:
    """Save button counter for the non-submitting preparation boundary."""

    def __init__(self):
        self.clicks = []

    @property
    def first(self):
        return self

    def count(self):
        return 1

    def is_visible(self):
        return True

    def click(self, **kw):
        self.clicks.append(kw)

    def locator(self, *_a, **_k):
        return _EmptyLoc()


class _FormPage:
    """Page Playwright toi thieu cho _fill_dropdown/poll_prepared_pages."""

    def __init__(self, url="https://portal.test/ho-so-cong-chung/tao-moi-nhanh"):
        self.url = url
        self.keyboard = SimpleNamespace(press=lambda _k: None)
        self.closed = False
        self.handlers = {}

    def wait_for_timeout(self, _ms):
        return None

    def locator(self, *_a, **_k):
        return _EmptyLoc()

    def get_by_text(self, *_a, **_k):
        return _EmptyLoc()

    def get_by_role(self, *_a, **_k):
        return _EmptyLoc()

    def evaluate(self, _script):
        return "complete"

    def on(self, event, callback):
        self.handlers[event] = callback

    def is_closed(self):
        return self.closed

    def close(self):
        self.closed = True


@unittest.skipUnless(_engines_available(), "upload_lab engine missing")
class EngineTabResourceTest(unittest.TestCase):
    """Real uploader with fake locators: tab resources and non-submit commit."""

    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.workdir = Path(self.tempdir.name) / "wd"
        self.workdir.mkdir()
        self.uploader = engine_roots.import_engine_module(
            "upload_lab", "playwright_uploader")
        self.batch_scan = engine_roots.import_engine_module(
            "upload_lab", "batch_scan")
        self.settings = self.uploader.UploaderSettings(
            base_url="https://portal.test",
            login_url="https://portal.test/dang-nhap",
            create_url="https://portal.test/ho-so-cong-chung/tao-moi-nhanh",
            storage_state_path=self.workdir / "state.json")
        self.logs = []
        self.session = self.uploader.NamDinhUploaderSession(
            self.settings, working_dir=self.workdir,
            log_callback=self.logs.append)

    def _seed_record(self, contract_no, *, file_key="k1",
                     status="prepared_dry_run"):
        src = self.workdir / f"{file_key}.docx"
        src.write_text("dummy", encoding="utf-8")
        conn = self.batch_scan.connect_registry(
            self.workdir / "registry.sqlite3")
        try:
            self.batch_scan.upsert_registry_record(
                conn, file_key=file_key, file_path=src,
                stat_result=src.stat(), customer_folder="Khach",
                contract_no=contract_no, status=status, run_id="r1")
            row = conn.execute(
                "SELECT id FROM file_registry WHERE file_key=?",
                (file_key,)).fetchone()
            return int(row[0])
        finally:
            conn.close()

    def _registry_status(self, record_id):
        conn = self.batch_scan.connect_registry(
            self.workdir / "registry.sqlite3")
        try:
            row = conn.execute(
                "SELECT status FROM file_registry WHERE id=?",
                (record_id,)).fetchone()
            return row[0] if row else None
        finally:
            conn.close()

    def _register_page(self, record_id, url):
        page = _FormPage(url)
        self.session._register_prepared_page(
            SimpleNamespace(record_id=record_id,
                            contract_no="101/2026/CCGD"),
            page)
        return page

    # ------------------------------------------------ tai nguyen tab (cutover)
    def test_navigated_away_tab_is_still_reported_open(self):
        """Tab roi trang tao moi (ke ca da bam Luu) van la tab dang mo:
        poll khong doc POST/dieu huong, khong dong tab, khong ghi registry."""
        rid = self._seed_record("101/2026/CCGD")
        page = self._register_page(
            rid, "https://portal.test/ho-so-cong-chung")

        result = self.session.poll_prepared_pages()
        self.assertEqual(result["closed_record_ids"], [])
        self.assertEqual(result["open_record_ids"], [rid])
        self.assertFalse(page.closed)
        self.assertIn(rid, self.session.prepared_pages)
        self.assertEqual(self._registry_status(rid), "prepared_dry_run")

        # POST /api/hoso 2xx khong con duoc nghe: khong co listener save.
        self.assertNotIn("response", page.handlers)
        again = self.session.poll_prepared_pages()
        self.assertEqual(again["open_record_ids"], [rid])
        self.assertEqual(again["closed_record_ids"], [])
        self.assertFalse(page.closed)
        self.assertEqual(self._registry_status(rid), "prepared_dry_run")

    def test_closed_tab_reports_closed_and_untracks(self):
        """Tab dong that: bao closed + bo tracking; khong ghi ket qua Luu."""
        rid = self._seed_record("103/2026/CCGD")
        page = self._register_page(
            rid, "https://portal.test/ho-so-cong-chung")
        page.closed = True
        result = self.session.poll_prepared_pages()
        self.assertEqual(result["closed_record_ids"], [rid])
        self.assertEqual(result["open_record_ids"], [])
        self.assertNotIn(rid, self.session.prepared_pages)
        self.assertEqual(self._registry_status(rid), "prepared_dry_run")

    def test_free_text_commit_never_clicks_save_or_presses_enter(self):
        combo = _ComboLoc()
        save = _SaveLoc()
        self.session._resolve_control_locator = (
            lambda page, field, **kw:
            (combo, "test") if field == "ten_hop_dong"
            else (_EmptyLoc(), "none"))
        self.session._locator_from_strategy = (
            lambda page, strategy, **kw: save)
        page = _FormPage()
        self.assertTrue(self.session._fill_dropdown(page, "ten_hop_dong", "HD A"))
        self.assertEqual(combo.value, "HD A")
        self.assertEqual(save.clicks, [])
        self.assertNotIn("Enter", combo.presses)
        self.assertIn("Tab", combo.presses)


if __name__ == "__main__":
    unittest.main()
