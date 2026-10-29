"""MIN-128 — mock backend notary case-drafting v2 (contract-faithful).

Chay:  python -m pytest shell/test/test_notary_mock_adapter.py -q

Gateway (`notary_gateway`) chon mock khi G1_DEV_NOTARY_MOCK=1 VA sidecar
khong packaged (sys.frozen); mac dinh real. Scenario fixtures o
`fixtures/notary-case-drafting/*.json` — ten/dia chi deu la mau gia.

Contract SOT: contracts/notary-case-drafting.md §13 (notary.case-drafting.v2)
— test reuse chinh validator cua contract (`validate_examples.violations`)
de kiem chung wire shape.
"""
import copy
import importlib.util
import json
import os
import subprocess
import sys
import threading
import time
import uuid
from pathlib import Path

import pytest

_HERE = Path(__file__).resolve().parent
_SHELL = _HERE.parent
_REPO = _SHELL.parent
sys.path.insert(0, str(_SHELL / "sidecar"))

from errors import CommandError            # noqa: E402
from jobstore import Job, JobStore          # noqa: E402
import command_registry as reg              # noqa: E402
import notary_gateway as gw                 # noqa: E402
import notary_mock_adapter as mock          # noqa: E402

FIXTURES = _HERE / "fixtures" / "notary-case-drafting"
VALIDATOR = (_REPO / "contracts" / "notary-case-drafting"
             / "validate_examples.py")

_spec = importlib.util.spec_from_file_location("ncd_validate", VALIDATOR)
ncd = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(ncd)

SCHEMA = "notary.case-drafting.v2"

DRAFTING_COMMANDS = [
    "notary.workspace_get",
    "notary.workspace_create",
    "notary.intake_analyze",
    "notary.workspace_commit_stage",
    "notary.workspace_update",
    "notary.workspace_convert_type",
    "notary.diagram_evaluate",
    "notary.diagram_save",
    "notary.word_export_options",
    "notary.word_export_batch",
    "notary.address_lookup",
]

P_OWNER = "11111111-1111-4111-8111-111111111111"
P_SPOUSE = "22222222-2222-4222-8222-222222222222"
P_CHILD1 = "33333333-3333-4333-8333-333333333333"
P_CHILD2 = "44444444-4444-4444-8444-444444444444"
P_POOL = "77777777-7777-4777-8777-777777777777"
P_TP1 = "aaaaaaaa-0000-4000-8000-000000000101"
P_TP2 = "bbbbbbbb-0000-4000-8000-000000000102"


# ---------- helpers ----------

def _load_fixture(name):
    """Doc scenario fixture; resolve '$ref_scenario' bang file cung ten."""
    doc = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    cases = {}
    for cid, c in (doc.get("cases") or {}).items():
        if isinstance(c, dict) and "$ref_scenario" in c:
            ref = _load_fixture(c["$ref_scenario"] + ".json")
            cases[cid] = copy.deepcopy(ref["cases"][cid])
        else:
            cases[cid] = c
    doc["cases"] = cases
    return doc


def _job(command="notary.workspace_get"):
    j = Job(str(uuid.uuid4()), command)
    j.status = "running"
    return j


def _dir_link(link, target):
    """Directory link `link` -> `target`: os.symlink truoc, fallback
    junction `mklink /J` tren Windows (khong can quyen admin)."""
    try:
        os.symlink(target, link, target_is_directory=True)
        return
    except OSError:
        pass
    if os.name == "nt":
        r = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(target)],
            capture_output=True)
        if r.returncode == 0:
            return
    pytest.skip("không tạo được directory link (thiếu quyền)")


def _file_link(link, target):
    """File symlink — Windows can SeCreateSymbolicLinkPrivilege
    (WinError 1314 tren may khong co quyen) -> skip."""
    try:
        os.symlink(target, link)
    except OSError as exc:
        pytest.skip(f"không tạo được file symlink: {exc}")


WINDOWS_ONLY = pytest.mark.skipif(
    os.name != "nt", reason="file_ref yêu cầu path tuyệt đối Windows")


def _call(fn_name, payload):
    """Di qua gateway dispatch (mock da bat qua autouse fixture)."""
    return gw.dispatch(fn_name, _job(f"notary.{fn_name}"), payload)


def _node(nid, person_id=None, parents=(), spouse=None,
          own=(), receive=(), hidden=False, deleted=False):
    """diagram node v3 (§13.4) — ownPositions/receivePositions ⊆ {1..5}."""
    return {
        "id": nid, "personId": person_id,
        "parentSlotIds": list(parents),
        "spouseSlotIds": [spouse] if spouse else [],
        "spouseSlotId": spouse, "spouseSuggestionId": None,
        "ownPositions": list(own), "receivePositions": list(receive),
        "hidden": hidden, "deleted": deleted,
    }


def _ready_state():
    """State v3 domain inheritance khop fixture ready.json (2 assets)."""
    return {
        "version": 3, "domain": "inheritance",
        "nodes": [
            _node("owner", P_OWNER, spouse="spouse", own=[1, 2]),
            _node("spouse", P_SPOUSE, spouse="owner", receive=[1, 2]),
            _node("child_1", P_CHILD1, parents=("owner", "spouse"),
                  receive=[1, 2]),
            _node("child_2", P_CHILD2, parents=("owner", "spouse"),
                  receive=[1, 2]),
        ],
    }


def _tp_state():
    """State v3 domain two_party canonical 30 slot (p1/p2 da gan)."""
    nodes = [{"id": f"p{i}", "personId": None,
              "hidden": False, "deleted": False}
             for i in range(1, 31)]
    nodes[0]["personId"] = P_TP1
    nodes[1]["personId"] = P_TP2
    return {"version": 3, "domain": "two_party", "nodes": nodes}


def _job_doc(command, payload, result=None, error=None,
             status="succeeded", fixture_context=None):
    doc = {
        "contract_version": "desktopcommand.v1",
        "job_id": "j_mock_test",
        "command_id": "bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb",
        "command": command,
        "payload": payload,
        "status": status,
        "waiting_on": None,
        "progress": None,
        "result": result,
        "error": error,
        "updated_at": "2026-09-24T00:00:00Z",
    }
    ctx = {"draft_v2": True}                     # §13 — flag ngoai wire
    if fixture_context:
        ctx.update(fixture_context)
    doc["fixture_context"] = ctx
    return doc


def _assert_contract(doc):
    """Chay chinh validator cua contract len job snapshot cua mock."""
    viols = ncd.violations(doc)
    assert viols == [], f"contract violations: {viols}"


def _no_confirmed_key(obj):
    if isinstance(obj, dict):
        for k, v in obj.items():
            assert str(k).lower() != "confirmed", "co key 'confirmed' bi cam"
            _no_confirmed_key(v)
    elif isinstance(obj, list):
        for v in obj:
            _no_confirmed_key(v)


def _no_legacy_keys(obj):
    """v2 wire: is_primary/isLandOwner/willReceive bi cam o moi noi."""
    if isinstance(obj, dict):
        for k, v in obj.items():
            assert k not in ("is_primary", "isLandOwner", "willReceive"), k
            _no_legacy_keys(v)
    elif isinstance(obj, list):
        for v in obj:
            _no_legacy_keys(v)


def _wait_terminal(store, job_id, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = store.get(job_id)
        if job.status in ("succeeded", "failed", "canceled", "partial"):
            return job
        time.sleep(0.02)
    raise AssertionError(f"job {job_id} khong terminal trong {timeout}s")


@pytest.fixture(autouse=True)
def _mock_env(monkeypatch):
    """Moi test chay voi mock ON va state seed mac dinh sach."""
    monkeypatch.setenv("G1_DEV_NOTARY_MOCK", "1")
    monkeypatch.delattr(sys, "frozen", raising=False)
    gw._warned_packaged = False          # doc lap thu tu test
    mock.reset_backend()
    yield
    mock.reset_backend()


# ---------- fixtures / registration ----------

class TestScaffolding:
    def test_commands_registered(self):
        for cmd in DRAFTING_COMMANDS:
            assert cmd in reg.COMMANDS, f"thieu {cmd} trong COMMANDS"

    def test_address_lookup_dispatch_envelope(self):
        """registry → gateway → adapter → service → _result envelope.

        Mock delegate sang adapter that → doc ledger local that
        (`notary_v2/data/administrative_units`), khong can fixture.
        """
        res = _call("address_lookup", {
            "query": "Thị trấn Lâm, huyện Ý Yên, tỉnh Nam Định",
            "presentation": "drafting",
        })
        assert res["kind"] == "address_lookup"
        assert res["data"]["schema_version"] == "notary.address-lookup.v2"
        assert res["data"]["status"] in (
            "found", "not_found", "too_many", "invalid_input",
            "reload_failed")
        assert res["data"]["presentation"] in ("drafting", "full", None)
        assert res["data"]["total"] == len(res["data"]["results"]) or \
            res["data"]["status"] == "too_many"
        bad = _call("address_lookup", {
            "query": "xã A", "presentation": "drafting",
            "source_date": "2024-06-01"})
        assert bad["data"]["status"] == "invalid_input"

    def test_fixtures_loadable_and_no_real_pii(self):
        names = ["empty", "ready", "done", "conflict", "intake-partial",
                 "diagram-warning", "word-collision", "word-partial",
                 "word-all-failed", "word-canceled", "unsupported",
                 "ready-two-party"]
        for n in names:
            doc = _load_fixture(f"{n}.json")
            assert doc["scenario"] == n
            assert doc.get("cases"), n
            raw = json.dumps(doc, ensure_ascii=False)
            assert "Người Mẫu" in raw or n in (
                "empty", "unsupported", "intake-partial",
                "word-all-failed"), n   # word-all-failed ref -> empty
            _no_legacy_keys(doc)

    def test_fixture_invariant_v2(self):
        """Moi fixture: diagram state version 3 + domain; asset khong
        is_primary; inheritance stage co owner_row_id; two_party stage
        KHONG co owner_row_id + diagram 30 slot canonical."""
        for f in FIXTURES.glob("*.json"):
            doc = _load_fixture(f.name)
            for cid, c in doc["cases"].items():
                st = (c.get("diagram") or {}).get("state") or {}
                assert st.get("version") == 3, (f.name, cid)
                assert st.get("domain") in ("inheritance", "two_party"), \
                    (f.name, cid)
                stage = c.get("stage") or {}
                for a in stage.get("assets") or []:
                    assert "is_primary" not in a, (f.name, cid)
                ct = c.get("case_type")
                if ct == "inheritance":
                    assert "owner_row_id" in stage, (f.name, cid)
                elif ct == "two_party":
                    assert "owner_row_id" not in stage, (f.name, cid)
                    ids = [n.get("id") for n in st.get("nodes") or []]
                    assert ids == [f"p{i}" for i in range(1, 31)], \
                        (f.name, cid)

    def test_default_backend_covers_scenarios(self):
        # Default seed (fixture files) phai co du case cho moi scenario.
        for cid in (42, 43, 44, 45, 46, 47, 48):
            res = _call("workspace_get", {"case_id": cid})
            assert res["data"]["case"]["id"] == cid


# ---------- gateway ----------

class TestGateway:
    def test_mock_selected_when_env_and_not_packaged(self, monkeypatch):
        monkeypatch.setenv("G1_DEV_NOTARY_MOCK", "1")
        monkeypatch.delattr(sys, "frozen", raising=False)
        assert gw.use_mock() is True

    def test_real_by_default(self, monkeypatch):
        monkeypatch.delenv("G1_DEV_NOTARY_MOCK", raising=False)
        assert gw.use_mock() is False

    def test_flag_value_other_than_1_ignored(self, monkeypatch):
        monkeypatch.setenv("G1_DEV_NOTARY_MOCK", "yes")
        assert gw.use_mock() is False

    def test_packaged_ignores_flag_with_warning(self, monkeypatch, capsys):
        monkeypatch.setenv("G1_DEV_NOTARY_MOCK", "1")
        # PyInstaller packaged — sys.frozen chi ton tai tren ban dong goi
        monkeypatch.setattr(sys, "frozen", True, raising=False)
        assert gw.use_mock() is False
        err = capsys.readouterr().err
        assert "G1_DEV_NOTARY_MOCK" in err

    def test_packaged_dispatch_goes_real(self, monkeypatch):
        monkeypatch.setenv("G1_DEV_NOTARY_MOCK", "1")
        monkeypatch.setattr(sys, "frozen", True, raising=False)
        # Packaged bo qua flag mock -> dispatch phai toi real adapter
        # (MIN-107+), khong phai mock va khong phai engine_not_installed.
        import notary_adapter
        sentinel = object()
        monkeypatch.setattr(
            notary_adapter, "workspace_get",
            lambda job, payload: sentinel)
        assert gw.dispatch("workspace_get", _job(), {"case_id": 42}) is sentinel

    def test_dispatch_mock(self):
        res = gw.dispatch("workspace_get", _job(), {"case_id": 42})
        assert res["data"]["backend_mode"] == "mock"


# ---------- workspace_get ----------

class TestWorkspaceGet:
    def test_empty(self):
        mock.reset_backend(_load_fixture("empty.json"))
        payload = {"case_id": 43}
        res = _call("workspace_get", payload)
        assert res["kind"] == "workspace_get"
        data = res["data"]
        assert data["schema_version"] == SCHEMA
        assert data["backend_mode"] == "mock"
        assert data["case"]["revision"] == 1
        assert "locked" not in data["case"]
        assert data["stage"] == {"owner_row_id": None,
                                 "people": [], "assets": []}
        assert data["diagram"]["render_model"] is None
        assert data["diagram"]["domain"] == "inheritance"
        assert data["capabilities"]["diagram"] is True
        _assert_contract(_job_doc("notary.workspace_get", payload, result=res))

    def test_ready_stage_full(self):
        mock.reset_backend(_load_fixture("ready.json"))
        res = _call("workspace_get", {"case_id": 42})
        data = res["data"]
        assert len(data["stage"]["people"]) == 4
        assert len(data["stage"]["assets"]) == 2
        assert data["stage"]["owner_row_id"] == P_OWNER
        rm = data["diagram"]["render_model"]
        assert rm is not None and rm["status"] == "complete"
        assert data["case"]["revision"] == 7
        # personId trong diagram deu thuoc stage da commit
        stage_ids = {p["row_id"] for p in data["stage"]["people"]}
        for n in data["diagram"]["state"]["nodes"]:
            assert set(n) <= {"id", "personId", "parentSlotIds",
                              "spouseSlotIds", "spouseSlotId",
                              "spouseSuggestionId",
                              "ownPositions", "receivePositions",
                              "hidden", "deleted"}
            if n["personId"] is not None:
                assert n["personId"] in stage_ids
        _assert_contract(_job_doc(
            "notary.workspace_get", {"case_id": 42}, result=res,
            fixture_context={"stage_row_ids": sorted(stage_ids)}))

    def test_unsupported_case_type(self):
        mock.reset_backend(_load_fixture("unsupported.json"))
        res = _call("workspace_get", {"case_id": 45})
        data = res["data"]
        assert data["case"]["case_type"] == "gift"
        assert data["capabilities"]["word_export"] is False
        assert data["capabilities"]["diagram"] is False
        assert data["capabilities"]["intake"] == []

    def test_two_party_case(self):
        mock.reset_backend(_load_fixture("ready-two-party.json"))
        res = _call("workspace_get", {"case_id": 48})
        data = res["data"]
        assert data["case"]["case_type"] == "two_party"
        assert data["case"]["document_type"] == "chuyen_nhuong"
        assert "owner_row_id" not in data["stage"]
        st = data["diagram"]["state"]
        assert st["version"] == 3 and st["domain"] == "two_party"
        assert [n["id"] for n in st["nodes"]] == \
            [f"p{i}" for i in range(1, 31)]
        rm = data["diagram"]["render_model"]
        assert rm["status"] == "unsupported"
        assert rm["allocations"] == {}
        assert data["capabilities"]["word_export"] is True   # NAIA-8
        assert data["capabilities"]["diagram"] is True

    def test_case_not_found(self):
        with pytest.raises(CommandError) as exc:
            _call("workspace_get", {"case_id": 9999})
        assert exc.value.code == "case_not_found"


# ---------- workspace_create (§13.6) ----------

_NEW_UNSET = object()


def _mk_person(**kw):
    row = {
        "row_id": str(uuid.uuid4()), "entity_id": None,
        "ho_ten": "Người Mẫu A", "gioi_tinh": "Nam",
        "ngay_sinh": "1950", "ngay_chet": "2020-01-15",
        "so_giay_to": None, "ngay_cap": None, "noi_cap": None,
        "dia_chi": None, "place_of_origin": None,
        "loai_giay_to": None, "loai_dia_chi": None,
    }
    row.update(kw)
    return row


def _mk_asset(**kw):
    row = {
        "row_id": str(uuid.uuid4()), "entity_id": None,
        "so_serial": "MM000001", "so_vao_so": None, "so_thua_dat": "123",
        "so_to_ban_do": "45", "dia_chi": "Địa chỉ mẫu tài sản 1",
        "loai_so": None, "hinh_thuc_su_dung": None, "thoi_han": None,
        "nguon_goc": None, "ngay_cap": None, "co_quan_cap": None,
        "land_rows": None,
    }
    row.update(kw)
    return row


def _create_payload(key=_NEW_UNSET, case_meta=_NEW_UNSET,
                    with_diagram=True):
    owner = _mk_person()
    spouse = _mk_person(ho_ten="Người Mẫu B", gioi_tinh="Nữ",
                        ngay_chet=None)
    p = {
        "idempotency_key": (key if key is not _NEW_UNSET
                            else str(uuid.uuid4())),
        "case": (case_meta if case_meta is not _NEW_UNSET
                 else {"document_type": "khai_nhan"}),
        "stage": {"owner_row_id": owner["row_id"],
                  "people": [owner, spouse], "assets": [_mk_asset()]},
    }
    if with_diagram:
        p["diagram"] = {"state": {"version": 3, "domain": "inheritance",
                                  "nodes": [
            _node("owner", owner["row_id"], spouse="spouse", own=[1]),
            _node("spouse", spouse["row_id"], spouse="owner", receive=[1]),
        ]}}
    return p


class TestWorkspaceCreate:
    def test_create_success_shape(self):
        res = _call("workspace_create", _create_payload(
            case_meta={"document_type": "khai_nhan",
                       "ngay_lap_ho_so": "2026-09-26",
                       "noi_niem_yet": "xã Mẫu", "ghi_chu": "Nháp"}))
        assert res["kind"] == "workspace_create"
        data = res["data"]
        assert data["schema_version"] == SCHEMA
        assert data["backend_mode"] == "mock"
        assert data["created"] is True
        c = data["case"]
        assert c["case_type"] == "inheritance"
        assert c["status"] == "draft"
        assert "locked" not in c
        assert c["revision"] == 1
        assert c["ngay_lap_ho_so"] == "2026-09-26"
        assert c["noi_niem_yet"] == "xã Mẫu"
        assert c["ghi_chu"] == "Nháp"
        assert all(p["entity_id"] for p in data["stage"]["people"])
        assert data["stage"]["owner_row_id"] == \
            data["stage"]["people"][0]["row_id"]
        assert data["diagram"]["state"]["version"] == 3
        assert data["diagram"]["state"]["domain"] == "inheritance"
        assert data["diagram"]["render_model"]["engineVersion"] == 2
        assert data["capabilities"]["word_export"] is True
        _no_legacy_keys(data)
        # case ton tai — get lai duoc
        got = _call("workspace_get", {"case_id": c["id"]})
        assert got["data"]["case"]["revision"] == 1

    def test_create_without_diagram_seeds_owner(self):
        """diagram absent -> server seed node 'owner' = owner_row_id,
        render_model null (§13.6)."""
        p = _create_payload(with_diagram=False)
        res = _call("workspace_create", p)
        data = res["data"]
        st = data["diagram"]["state"]
        assert st["version"] == 3 and st["domain"] == "inheritance"
        owner = next((n for n in st["nodes"] if n["id"] == "owner"), None)
        assert owner is not None
        assert owner["personId"] == data["stage"]["owner_row_id"]
        assert data["diagram"]["render_model"] is None

    def test_create_two_party_seeds_30_slots(self):
        tp1 = _mk_person(ho_ten="Bên A 1")
        tp2 = _mk_person(ho_ten="Bên B 1", gioi_tinh="Nữ")
        p = {
            "idempotency_key": str(uuid.uuid4()),
            "case": {"case_type": "two_party",
                     "document_type": "chuyen_nhuong"},
            "stage": {"people": [tp1, tp2], "assets": [_mk_asset()]},
        }
        res = _call("workspace_create", p)
        data = res["data"]
        assert data["case"]["case_type"] == "two_party"
        assert "owner_row_id" not in data["stage"]
        st = data["diagram"]["state"]
        assert st["domain"] == "two_party"
        assert [n["id"] for n in st["nodes"]] == \
            [f"p{i}" for i in range(1, 31)]
        assert data["diagram"]["render_model"] is None
        assert data["capabilities"]["word_export"] is True   # NAIA-8

    def test_two_party_stage_owner_row_id_rejected(self):
        p = {
            "idempotency_key": str(uuid.uuid4()),
            "case": {"case_type": "two_party",
                     "document_type": "chuyen_nhuong"},
            "stage": {"owner_row_id": None,
                      "people": [_mk_person()], "assets": [_mk_asset()]},
        }
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "validation_error"

    def test_two_party_wrong_document_type(self):
        p = {
            "idempotency_key": str(uuid.uuid4()),
            "case": {"case_type": "two_party",
                     "document_type": "khai_nhan"},   # enum inheritance
            "stage": {"people": [_mk_person()], "assets": [_mk_asset()]},
        }
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "validation_error"

    def test_idempotent_replay(self):
        key = str(uuid.uuid4())
        first = _call("workspace_create", _create_payload(key=key))
        second = _call("workspace_create", _create_payload(key=key))
        assert first["data"]["created"] is True
        assert second["data"]["created"] is False
        assert second["data"]["case"]["id"] == first["data"]["case"]["id"]
        third = _call("workspace_create", _create_payload())
        assert third["data"]["created"] is True
        assert third["data"]["case"]["id"] != first["data"]["case"]["id"]

    def test_bad_idempotency_key(self):
        for bad in (None, "nope", 42, str(uuid.uuid1())):
            with pytest.raises(CommandError) as exc:
                _call("workspace_create", _create_payload(key=bad))
            assert exc.value.code == "validation_error", bad

    def test_missing_owner_pointer(self):
        """inheritance: owner_row_id null/missing/khong thuoc people ->
        workspace_owner_required (§13.6)."""
        for bad in (None, "khong-uuid", str(uuid.uuid4())):
            p = _create_payload()
            p["stage"]["owner_row_id"] = bad
            with pytest.raises(CommandError) as exc:
                _call("workspace_create", p)
            assert exc.value.code == "workspace_owner_required", bad
        p = _create_payload()
        del p["stage"]["owner_row_id"]
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "workspace_owner_required"

    def test_owner_node_mismatch_rejected(self):
        """Node owner personId != stage.owner_row_id ->
        diagram_owner_mismatch — server KHONG tu sua (§13.6)."""
        p = _create_payload()
        # owner.personId -> spouse row; node spouse giai phong de khong
        # bi duplicate_person (wire check chay truoc mirror check).
        p["diagram"]["state"]["nodes"][0]["personId"] = \
            p["stage"]["people"][1]["row_id"]
        p["diagram"]["state"]["nodes"][1]["personId"] = None
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "diagram_owner_mismatch"
        assert exc.value.details["owner_row_id"] == \
            p["stage"]["owner_row_id"]

    def test_entity_hint_accepted_on_create(self):
        """`entity_id` trên dòng mới là hint nội bộ — không chặn
        (parity server thật, owner 23/10/2026)."""
        p = _create_payload()
        p["stage"]["people"][0]["entity_id"] = 5
        res = _call("workspace_create", p)
        assert res["data"]["stage"]["people"][0]["entity_id"] == 5

    def test_empty_stage(self):
        """create stage rong: assets la required (stage_validation_error);
        people rong + owner_row_id treo -> workspace_owner_required
        (pointer check chay truoc field errors — real parity)."""
        p = _create_payload()
        p["stage"]["assets"] = []
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "stage_validation_error"
        fe = exc.value.details["field_errors"]
        assert any(e["code"] == "required" for e in fe)
        # people rong -> owner pointer khong tro vao row nao
        p = _create_payload()
        p["stage"]["people"] = []
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "workspace_owner_required"

    def test_is_primary_rejected_as_unknown_field(self):
        """v2: is_primary bi cam tren wire — truong la ->
        validation_error (§13.3)."""
        p = _create_payload()
        p["stage"]["assets"][0]["is_primary"] = True
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "validation_error"

    def test_asset_limit_create(self):
        p = _create_payload()
        p["stage"]["assets"] += [
            _mk_asset(so_serial=f"MM{i + 1:06d}")
            for i in range(mock.MAX_ASSETS)]
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "stage_validation_error"
        fe = exc.value.details["field_errors"]
        assert any(e["code"] == "asset_limit" for e in fe)

    def test_outside_stage_person(self):
        p = _create_payload()
        p["diagram"]["state"]["nodes"].append(
            _node("child_1", str(uuid.uuid4()), parents=("owner",),
                  receive=[1]))
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "diagram_reference_outside_stage"

    def test_domain_mismatch(self):
        """state.domain hop le nhung khac case.case_type ->
        diagram_domain_mismatch (§13.5)."""
        p = _create_payload()
        p["diagram"]["state"]["domain"] = "two_party"
        p["diagram"]["state"]["nodes"] = [
            {"id": f"p{i}", "personId": None,
             "hidden": False, "deleted": False} for i in range(1, 31)]
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", p)
        assert exc.value.code == "diagram_domain_mismatch"
        assert exc.value.details["expected"] == "inheritance"

    def test_case_meta_invalid(self):
        for meta in ({"document_type": "khac"},
                     {"case_type": "gift"},
                     {"document_type": "khai_nhan",
                      "ngay_lap_ho_so": "26/09/2026"},
                     {"document_type": "khai_nhan", "noi_niem_yet": "  "}):
            with pytest.raises(CommandError) as exc:
                _call("workspace_create",
                      _create_payload(case_meta=meta))
            assert exc.value.code == "validation_error", meta

    def test_contract_shape_via_validator(self):
        payload = _create_payload()
        res = _call("workspace_create", payload)
        _assert_contract(_job_doc(
            "notary.workspace_create", payload, result=res))


# ---------- workspace_commit_stage ----------

# ---------- workspace_save_case (§13.19) ----------


class TestWorkspaceSaveCase:
    def test_new_done_replay_and_idempotency_conflict(self):
        payload = _create_payload()
        first = _call("workspace_save_case", payload)
        assert first["data"]["created"] is True
        assert first["data"]["case"]["status"] == "done"
        assert first["data"]["case"]["revision"] == 1
        assert first["data"]["diagram"]["render_model"] is not None

        replay = _call("workspace_save_case", payload)
        assert replay["data"]["created"] is False
        assert replay["data"]["case"]["id"] == first["data"]["case"]["id"]
        assert replay["data"]["case"]["revision"] == 1

        other = copy.deepcopy(payload)
        other["case"] = {"document_type": "khai_nhan", "ghi_chu": "khác"}
        with pytest.raises(CommandError) as exc:
            _call("workspace_save_case", other)
        assert exc.value.code == "idempotency_conflict"

    def test_existing_done_and_revision_guard(self):
        case_id = _call("workspace_get", {"case_id": 42})[
            "data"]["case"]["id"]
        payload = _create_payload(case_meta={"document_type": "khai_nhan"})
        payload.pop("idempotency_key")
        payload["case_id"] = case_id
        payload["base_revision"] = 7
        saved = _call("workspace_save_case", payload)
        assert saved["data"]["created"] is False
        assert saved["data"]["case"]["status"] == "done"
        assert saved["data"]["case"]["revision"] == 8

        with pytest.raises(CommandError) as exc:
            _call("workspace_save_case", payload)
        assert exc.value.code == "workspace_conflict"
        assert exc.value.details["server_revision"] == 8

    def test_new_save_keeps_entity_hints_from_source(self):
        """Dòng mới mang entity_id nguồn (source_latest/nhân bản) KHÔNG bị
        chặn — parity server thật: entity_id là hint nội bộ (§4.1 mục 5)."""
        payload = _create_payload()
        for row, hint in zip(payload["stage"]["people"], (101, 102)):
            row["entity_id"] = hint
        payload["stage"]["assets"][0]["entity_id"] = 201

        saved = _call("workspace_save_case", payload)

        assert saved["data"]["created"] is True
        assert [r["entity_id"] for r in saved["data"]["stage"]["people"]] == \
            [101, 102]
        assert saved["data"]["stage"]["assets"][0]["entity_id"] == 201

    def test_new_save_drops_garbage_entity_hint(self):
        """Hint rác (không phải int ≥ 1) bị bỏ và backend cấp ID mới."""
        payload = _create_payload()
        payload["stage"]["people"][0]["entity_id"] = "hint-rác"
        payload["stage"]["assets"][0]["entity_id"] = 0

        saved = _call("workspace_save_case", payload)

        assert isinstance(
            saved["data"]["stage"]["people"][0]["entity_id"], int)
        assert saved["data"]["stage"]["assets"][0]["entity_id"] >= 1


class TestCommitStage:
    def _payload(self, base_revision=1, people=None, assets=None,
                 owner_row_id=None, case_id=43):
        return {"case_id": case_id, "base_revision": base_revision,
                "stage": {"owner_row_id": owner_row_id,
                          "people": people or [], "assets": assets or []}}

    def _person(self, row_id, name="Người Mẫu X", entity_id=None):
        return {
            "row_id": row_id, "entity_id": entity_id,
            "ho_ten": name, "gioi_tinh": "Nam",
            "ngay_sinh": "1990-01-01", "ngay_chet": None,
            "so_giay_to": None, "ngay_cap": None, "noi_cap": None,
            "dia_chi": "Địa chỉ mẫu", "place_of_origin": None,
        }

    def _asset(self, row_id, serial="MM000010"):
        return {
            "row_id": row_id, "entity_id": None,
            "so_serial": serial, "so_vao_so": None, "so_thua_dat": "9",
            "so_to_ban_do": "1", "dia_chi": "Địa chỉ mẫu tài sản",
            "loai_so": None, "hinh_thuc_su_dung": None, "thoi_han": None,
            "nguon_goc": None, "ngay_cap": None, "co_quan_cap": None,
            "land_rows": [],
        }

    def test_commit_success_assigns_entity_and_bumps_revision(self):
        mock.reset_backend(_load_fixture("empty.json"))
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        payload = self._payload(people=[self._person(rid)],
                                owner_row_id=rid)
        res = _call("workspace_commit_stage", payload)
        data = res["data"]
        assert data["revision"] == 2                     # 1 -> 2
        row = data["stage"]["people"][0]
        assert row["row_id"] == rid
        assert isinstance(row["entity_id"], int)         # backend gan
        assert data["stage"]["owner_row_id"] == rid
        assert data["diagram"]["render_model"] is not None  # re-evaluate §6.1
        _no_legacy_keys(data)
        _assert_contract(_job_doc(
            "notary.workspace_commit_stage", payload, result=res))
        # get lai thay stage da persist
        res2 = _call("workspace_get", {"case_id": 43})
        assert res2["data"]["case"]["revision"] == 2
        assert len(res2["data"]["stage"]["people"]) == 1

    def test_commit_syncs_owner_node(self):
        """§13.6: commit sync node 'owner' (neu co, chua deleted) :=
        owner_row_id — SOT la stage pointer. Node khong ton tai thi
        khong tu tao."""
        mock.reset_backend(_load_fixture("ready.json"))
        stage = copy.deepcopy(
            _load_fixture("ready.json")["cases"]["42"]["stage"])
        stage["owner_row_id"] = P_SPOUSE       # doi chu so huu
        res = _call("workspace_commit_stage", {
            "case_id": 42, "base_revision": 7, "stage": stage})
        nodes = res["data"]["diagram"]["state"]["nodes"]
        owner = next((n for n in nodes if n["id"] == "owner"), None)
        assert owner is not None and owner["personId"] == P_SPOUSE
        assert res["data"]["stage"]["owner_row_id"] == P_SPOUSE

    def test_commit_empty_diagram_no_owner_created(self):
        mock.reset_backend(_load_fixture("empty.json"))
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        res = _call("workspace_commit_stage",
                    self._payload(people=[self._person(rid)],
                                  owner_row_id=rid))
        nodes = res["data"]["diagram"]["state"]["nodes"]
        assert not [n for n in nodes if n.get("id") == "owner"]

    def test_revision_monotonic_two_commits(self):
        mock.reset_backend(_load_fixture("empty.json"))
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        p = self._payload(people=[self._person(rid)], owner_row_id=rid)
        assert _call("workspace_commit_stage", p)["data"]["revision"] == 2
        p["base_revision"] = 2
        assert _call("workspace_commit_stage", p)["data"]["revision"] == 3

    def test_stale_revision_conflict(self):
        mock.reset_backend(_load_fixture("conflict.json"))
        payload = _load_fixture("conflict.json")["request"][
            "workspace_commit_stage"]
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", payload)
        assert exc.value.code == "workspace_conflict"
        assert exc.value.details["server_revision"] == 7

    def test_future_revision_also_conflict(self):
        mock.reset_backend(_load_fixture("empty.json"))
        p = self._payload(base_revision=99)
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", p)
        assert exc.value.code == "workspace_conflict"

    def test_atomic_bad_row_rolls_back(self):
        mock.reset_backend(_load_fixture("empty.json"))
        bad = self._person("ffffffff-ffff-4fff-8fff-ffffffffffff")
        bad["ho_ten"] = ""                                # vi pham required
        good = self._person("eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee")
        payload = self._payload(people=[good, bad], owner_row_id=good[
            "row_id"])
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", payload)
        assert exc.value.code == "stage_validation_error"
        fe = exc.value.details["field_errors"]
        assert any(e["row_id"] == bad["row_id"] and e["field"] == "ho_ten"
                   and e["code"] == "required" for e in fe)
        # Stage khong doi (atomic) — dong good cung khong duoc ghi
        res = _call("workspace_get", {"case_id": 43})
        assert res["data"]["stage"]["people"] == []
        assert res["data"]["case"]["revision"] == 1

    def test_owner_pointer_required(self):
        mock.reset_backend(_load_fixture("empty.json"))
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        for bad_owner in (None, str(uuid.uuid4())):
            p = self._payload(people=[self._person(rid)],
                              owner_row_id=bad_owner)
            with pytest.raises(CommandError) as exc:
                _call("workspace_commit_stage", p)
            assert exc.value.code == "workspace_owner_required", bad_owner
        # stage khong owner_row_id key -> cung workspace_owner_required
        p = self._payload(people=[self._person(rid)])
        del p["stage"]["owner_row_id"]
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", p)
        assert exc.value.code == "workspace_owner_required"

    def test_duplicate_row_id(self):
        mock.reset_backend(_load_fixture("empty.json"))
        row = self._person("eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee")
        payload = self._payload(people=[row, dict(row)],
                                owner_row_id=row["row_id"])
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", payload)
        assert exc.value.code == "stage_validation_error"
        assert any(e["code"] == "duplicate_row_id"
                   for e in exc.value.details["field_errors"])

    def test_duplicate_entity_id_allowed(self):
        """Hai dòng cùng identity là hợp lệ (§4.2 step 6) — mock parity
        server thật sau 23/10/2026."""
        mock.reset_backend(_load_fixture("empty.json"))
        a = self._person("eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee",
                         entity_id=9)
        b = self._person("ffffffff-ffff-4fff-8fff-ffffffffffff",
                         entity_id=9)
        payload = self._payload(people=[a, b], owner_row_id=a["row_id"])
        res = _call("workspace_commit_stage", payload)
        assert [r["entity_id"] for r in res["data"]["stage"]["people"]] \
            == [9, 9]

    def test_duplicate_so_giay_to(self):
        mock.reset_backend(_load_fixture("empty.json"))
        a = self._person("eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee")
        b = self._person("ffffffff-ffff-4fff-8fff-ffffffffffff")
        a["so_giay_to"] = b["so_giay_to"] = "001122334455"
        payload = self._payload(people=[a, b], owner_row_id=a["row_id"])
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", payload)
        assert any(e["field"] == "so_giay_to"
                   for e in exc.value.details["field_errors"])

    def test_duplicate_serial_allowed_on_distinct_assets(self):
        mock.reset_backend(_load_fixture("empty.json"))
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        a1 = self._asset("aaaaaaa1-1111-4111-8111-111111111111",
                         serial="S-001")
        a2 = self._asset("aaaaaaa2-2222-4222-8222-222222222222",
                         serial="S-001")
        a2["so_thua_dat"] = "10"
        payload = self._payload(people=[self._person(rid)],
                                owner_row_id=rid, assets=[a1, a2])
        res = _call("workspace_commit_stage", payload)
        assets = res["data"]["stage"]["assets"]
        assert [a["so_serial"] for a in assets] == ["S-001", "S-001"]
        assert assets[0]["entity_id"] != assets[1]["entity_id"]

    def test_asset_limit(self):
        """§13.3: toi da 5 asset — dong thu 6 -> asset_limit field error
        gan row_id dong thua."""
        mock.reset_backend(_load_fixture("empty.json"))
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        assets = [self._asset(
            f"aaaaaaa{i}-1111-4111-8111-11111111111{i}",
            serial=f"MM{i:06d}") for i in range(1, 7)]
        payload = self._payload(people=[self._person(rid)],
                                owner_row_id=rid, assets=assets[:5])
        saved = _call("workspace_commit_stage", payload)
        assert len(saved["data"]["stage"]["assets"]) == 5
        payload["base_revision"] = saved["data"]["revision"]
        payload["stage"]["assets"] = assets
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", payload)
        fe = exc.value.details["field_errors"]
        over = [e for e in fe if e["code"] == "asset_limit"]
        assert over and over[0]["row_id"] == assets[5]["row_id"]

    def test_is_primary_key_rejected(self):
        mock.reset_backend(_load_fixture("empty.json"))
        a = self._asset("aaaaaaa1-1111-4111-8111-111111111111")
        a["is_primary"] = True                       # truong la tren v2
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        payload = self._payload(people=[self._person(rid)],
                                owner_row_id=rid, assets=[a])
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", payload)
        assert exc.value.code == "validation_error"

    def test_serial_is_verbatim_and_may_be_empty(self):
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        for serial in ("khong-canonical", ""):
            mock.reset_backend(_load_fixture("empty.json"))
            a = self._asset("aaaaaaa1-1111-4111-8111-111111111111",
                            serial=serial)
            payload = self._payload(people=[self._person(rid)],
                                    owner_row_id=rid, assets=[a])
            res = _call("workspace_commit_stage", payload)
            assert res["data"]["stage"]["assets"][0]["so_serial"] == serial

    def test_commit_prunes_diagram_refs(self):
        """Xoa nguoi khoi stage -> node tham chieu bi prune trong cung
        transaction, render_model khop state moi (§6.1/§13.6)."""
        mock.reset_backend(_load_fixture("ready.json"))
        keep = _load_fixture("ready.json")["cases"]["42"]["stage"][
            "people"][0]
        payload = {"case_id": 42, "base_revision": 7,
                   "stage": {"owner_row_id": keep["row_id"],
                             "people": [keep],
                             "assets": _load_fixture("ready.json")
                             ["cases"]["42"]["stage"]["assets"]}}
        res = _call("workspace_commit_stage", payload)
        nodes = res["data"]["diagram"]["state"]["nodes"]
        for n in nodes:
            if n["personId"] is not None:
                assert n["personId"] == keep["row_id"]

    def test_commit_prunes_positions_above_asset_count(self):
        """Giam assets 2 -> 1: dau chon vi tri 2 bi prune trong cung
        transaction + warning diagram.selection_pruned (§13.4)."""
        mock.reset_backend(_load_fixture("ready.json"))
        stage = copy.deepcopy(
            _load_fixture("ready.json")["cases"]["42"]["stage"])
        stage["assets"] = stage["assets"][:1]           # con 1 vi tri
        payload = {"case_id": 42, "base_revision": 7, "stage": stage}
        res = _call("workspace_commit_stage", payload)
        dg = res["data"]["diagram"]
        codes = [w["code"] for w in dg.get("warnings") or []]
        assert "diagram.selection_pruned" in codes
        for n in dg["state"]["nodes"]:
            assert all(p == 1 for p in n.get("ownPositions") or [])
            assert all(p == 1 for p in n.get("receivePositions") or [])

    def test_container_shape_is_validation_error(self):
        """Oracle parity: stage khong dict / people|assets thieu hoac
        khong list -> validation_error (KHONG stage_validation_error;
        code do chi cho loi row-level)."""
        mock.reset_backend(_load_fixture("empty.json"))
        for bad_stage in (
                "khong-phai-object",
                {"people": "x", "assets": []},
                {"people": []},                       # thieu assets
                {"assets": []},                       # thieu people
                {"people": [], "assets": "y"},
                {"people": None, "assets": []}):
            p = {"case_id": 43, "base_revision": 1, "stage": bad_stage}
            with pytest.raises(CommandError) as exc:
                _call("workspace_commit_stage", p)
            assert exc.value.code == "validation_error", bad_stage

    def test_unsupported_case_type(self):
        mock.reset_backend(_load_fixture("unsupported.json"))
        p = {"case_id": 45, "base_revision": 1,
             "stage": {"people": [], "assets": []}}
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", p)
        assert exc.value.code == "case_type_unsupported"
        assert exc.value.details["case_type"] == "gift"

    def test_two_party_commit(self):
        """two_party: commit stage khong owner_row_id; prune personId
        ngoai stage nhung giu 30 slot; render_model unsupported."""
        mock.reset_backend(_load_fixture("ready-two-party.json"))
        people = copy.deepcopy(_load_fixture("ready-two-party.json")
                               ["cases"]["48"]["stage"]["people"])
        payload = {"case_id": 48, "base_revision": 2,
                   "stage": {"people": people[:1],
                             "assets": []}}
        res = _call("workspace_commit_stage", payload)
        data = res["data"]
        assert data["revision"] == 3
        assert "owner_row_id" not in data["stage"]
        st = data["diagram"]["state"]
        assert [n["id"] for n in st["nodes"]] == \
            [f"p{i}" for i in range(1, 31)]
        # p2 bi xoa khoi stage -> personId null nhung slot con
        assert st["nodes"][1]["personId"] is None
        assert st["nodes"][0]["personId"] == people[0]["row_id"]
        assert data["diagram"]["render_model"]["status"] == "unsupported"

    def test_two_party_owner_row_id_rejected(self):
        mock.reset_backend(_load_fixture("ready-two-party.json"))
        p = {"case_id": 48, "base_revision": 2,
             "stage": {"owner_row_id": None, "people": [], "assets": []}}
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", p)
        assert exc.value.code == "validation_error"

    def test_two_party_people_limit(self):
        """two_party: toi da 30 nguoi — dong 31 -> people_limit."""
        mock.reset_backend(_load_fixture("ready-two-party.json"))
        people = [self._person(
            f"eeeeeeee-{i:04d}-4eee-8eee-eeeeeeeeeeee".replace("x", "0"))
            for i in range(31)]
        payload = {"case_id": 48, "base_revision": 2,
                   "stage": {"people": people, "assets": []}}
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", payload)
        fe = exc.value.details["field_errors"]
        assert any(e["code"] == "people_limit" for e in fe)


# ---------- workspace_update (§6.2 — NAIA-8) ----------

class TestWorkspaceUpdate:
    def test_update_meta_persists_and_reloads(self):
        mock.reset_backend(_load_fixture("ready.json"))
        rev = _call("workspace_get", {"case_id": 42})[
            "data"]["case"]["revision"]
        res = _call("workspace_update", {
            "case_id": 42, "base_revision": rev,
            "case": {"noi_niem_yet": "xã Yên Sở"}})
        assert res["kind"] == "workspace_update"
        assert res["data"]["schema_version"] == SCHEMA
        assert res["data"]["revision"] == rev + 1
        assert res["data"]["case"]["noi_niem_yet"] == "xã Yên Sở"
        got = _call("workspace_get", {"case_id": 42})
        assert got["data"]["case"]["noi_niem_yet"] == "xã Yên Sở"
        assert got["data"]["case"]["revision"] == rev + 1
        # Stage/Diagram khong bi cham.
        assert len(got["data"]["stage"]["people"]) == 4

    def test_conflict(self):
        mock.reset_backend(_load_fixture("ready.json"))
        rev = _call("workspace_get", {"case_id": 42})[
            "data"]["case"]["revision"]
        with pytest.raises(CommandError) as exc:
            _call("workspace_update", {
                "case_id": 42, "base_revision": rev + 9,
                "case": {"ghi_chu": "x"}})
        assert exc.value.code == "workspace_conflict"
        assert exc.value.details["server_revision"] == rev

    def test_validation_errors(self):
        mock.reset_backend(_load_fixture("ready.json"))
        rev = _call("workspace_get", {"case_id": 42})[
            "data"]["case"]["revision"]
        for bad_case in ({"case_type": "gift"}, {}, "chuoi", None):
            with pytest.raises(CommandError) as exc:
                _call("workspace_update", {
                    "case_id": 42, "base_revision": rev, "case": bad_case})
            assert exc.value.code == "validation_error", bad_case
        for bad_rev in (0, "1", None, True):
            with pytest.raises(CommandError) as exc:
                _call("workspace_update", {
                    "case_id": 42, "base_revision": bad_rev,
                    "case": {"ghi_chu": "x"}})
            assert exc.value.code == "validation_error", bad_rev
        with pytest.raises(CommandError) as exc:
            _call("workspace_update", {
                "case_id": 999, "base_revision": 1,
                "case": {"ghi_chu": "x"}})
        assert exc.value.code == "case_not_found"

    def test_invalid_patch_keeps_meta_and_revision(self):
        # NAIA-8 review: patch gom field dung + sai phai loi nguyen khoi —
        # khong luu mot phan (parity real update_meta, §6.2 atomic).
        mock.reset_backend(_load_fixture("ready.json"))
        before = _call("workspace_get", {"case_id": 42})["data"]["case"]
        rev = before["revision"]
        with pytest.raises(CommandError) as exc:
            _call("workspace_update", {
                "case_id": 42, "base_revision": rev,
                "case": {"noi_niem_yet": "xã Đã Đổi", "ghi_chu": 123}})
        assert exc.value.code == "validation_error"
        after = _call("workspace_get", {"case_id": 42})["data"]["case"]
        assert after["noi_niem_yet"] == before["noi_niem_yet"]
        assert after["ghi_chu"] == before["ghi_chu"]
        assert after["revision"] == rev


# ---------- intake_analyze ----------

class TestIntakeAnalyze:
    def _text_src(self, sid, text):
        return {"source_id": sid, "kind": "text", "text": text}

    def test_succeeded(self):
        mock.reset_backend(_load_fixture("empty.json"))
        sid = "aaaaaaaa-0000-4000-8000-000000000009"
        payload = {"case_id": 43,
                   "sources": [self._text_src(sid, "Người Mẫu I, sinh 1950")]}
        res = _call("intake_analyze", payload)
        assert res["kind"] == "intake_analyze"
        assert res["data"]["schema_version"] == SCHEMA
        sug = res["data"]["suggestions"][0]
        assert sug["source_id"] == sid
        assert sug["target"] in ("person", "asset")
        for fv in sug["fields"].values():
            assert fv["observation_state"] in (
                "observed", "normalized", "inferred")
        _no_confirmed_key(res)
        _assert_contract(_job_doc("notary.intake_analyze", payload, result=res))

    def test_partial_with_breakdown(self):
        doc = _load_fixture("intake-partial.json")
        mock.reset_backend(doc)
        payload = doc["request"]["intake_analyze"]
        res = _call("intake_analyze", payload)
        assert res.get("partial") is True            # marker cho jobstore
        bd = res["data"]["breakdown"]
        assert len(bd["succeeded"]) == 2
        assert bd["failed"] == ["aaaaaaaa-0000-4000-8000-000000000003"]
        assert res["data"]["errors"][0]["code"] == "intake.parse_failed"
        _assert_contract(_job_doc(
            "notary.intake_analyze", payload, result=res, status="partial"))

    def test_progress_reported(self):
        mock.reset_backend(_load_fixture("empty.json"))
        job = _job("notary.intake_analyze")
        payload = {"case_id": 43, "sources": [
            self._text_src("aaaaaaaa-0000-4000-8000-000000000001", "x")]}
        gw.dispatch("intake_analyze", job, payload)
        assert job.progress == {"done": 1, "total": 1,
                                "current_label": job.progress["current_label"]}
        assert job.progress["done"] == 1

    def test_too_many_sources(self):
        mock.reset_backend(_load_fixture("empty.json"))
        sources = [self._text_src(
            f"aaaaaaaa-0000-4000-8000-00000000000{i}", "x")
            for i in range(9)]
        with pytest.raises(CommandError) as exc:
            _call("intake_analyze", {"case_id": 43, "sources": sources})
        assert exc.value.code == "intake_too_many_sources"
        assert exc.value.details["limit"] == 8

    def test_unsupported_kind(self):
        mock.reset_backend(_load_fixture("empty.json"))
        s = {"source_id": "aaaaaaaa-0000-4000-8000-000000000001",
             "kind": "zalo_media", "text": "x"}
        with pytest.raises(CommandError) as exc:
            _call("intake_analyze", {"case_id": 43, "sources": [s]})
        assert exc.value.code == "intake_unsupported_source"
        assert exc.value.details["kind"] == "zalo_media"

    def test_text_too_long(self):
        mock.reset_backend(_load_fixture("empty.json"))
        s = self._text_src("aaaaaaaa-0000-4000-8000-000000000001",
                           "x" * 100_001)
        with pytest.raises(CommandError) as exc:
            _call("intake_analyze", {"case_id": 43, "sources": [s]})
        assert exc.value.code == "intake_text_too_long"

    def test_source_too_large(self):
        mock.reset_backend(_load_fixture("empty.json"))
        s = {"source_id": "aaaaaaaa-0000-4000-8000-000000000001",
             "kind": "image",
             "file_ref": {"path": "D:/mock-intake/big.jpg",
                          "scope": "machine_local",
                          "size_bytes": 21 * 1024 * 1024}}
        with pytest.raises(CommandError) as exc:
            _call("intake_analyze", {"case_id": 43, "sources": [s]})
        assert exc.value.code == "intake_source_too_large"

    def test_duplicate_source_id(self):
        mock.reset_backend(_load_fixture("empty.json"))
        s = self._text_src("aaaaaaaa-0000-4000-8000-000000000001", "x")
        with pytest.raises(CommandError) as exc:
            _call("intake_analyze", {"case_id": 43, "sources": [s, dict(s)]})
        assert exc.value.code == "validation_error"

    def test_file_ref_required_for_nontext(self):
        mock.reset_backend(_load_fixture("empty.json"))
        s = {"source_id": "aaaaaaaa-0000-4000-8000-000000000001",
             "kind": "pdf"}
        with pytest.raises(CommandError) as exc:
            _call("intake_analyze", {"case_id": 43, "sources": [s]})
        assert exc.value.code == "validation_error"

    def test_draft_mode_no_case_id(self):
        """§2.1a — case_id absent = nhap moi; bo qua kiem tra case."""
        sid = "aaaaaaaa-0000-4000-8000-000000000009"
        payload = {"sources": [self._text_src(sid, "Người Mẫu I, sinh 1950")]}
        res = _call("intake_analyze", payload)
        assert res["data"]["suggestions"][0]["source_id"] == sid
        _assert_contract(_job_doc("notary.intake_analyze", payload,
                                  result=res))

    def test_null_case_id_is_validation_error(self):
        s = self._text_src("aaaaaaaa-0000-4000-8000-000000000001", "x")
        with pytest.raises(CommandError) as exc:
            _call("intake_analyze", {"case_id": None, "sources": [s]})
        assert exc.value.code == "validation_error"


# ---------- diagram_evaluate / diagram_save ----------

class TestDiagram:
    def test_evaluate_complete(self):
        mock.reset_backend(_load_fixture("ready.json"))
        payload = {"case_id": 42, "diagram": {"state": _ready_state()}}
        res = _call("diagram_evaluate", payload)
        assert res["kind"] == "diagram_evaluate"
        data = res["data"]
        assert data["evaluated_revision"] == 7
        rm = data["render_model"]
        assert rm["engineVersion"] == 2
        assert rm["status"] == "complete"
        # 3 nguoi nhan, moi nguoi 1/3
        assert rm["allocations"][P_SPOUSE]["finalShare"] == "1/3"
        _assert_contract(_job_doc(
            "notary.diagram_evaluate", payload, result=res,
            fixture_context={"stage_row_ids": [
                P_OWNER, P_SPOUSE, P_CHILD1, P_CHILD2],
                "stage_owner_row_id": P_OWNER}))

    def test_evaluate_unassigned_pool_warning(self):
        doc = _load_fixture("diagram-warning.json")
        mock.reset_backend(doc)
        payload = doc["request"]["diagram_evaluate"]
        res = _call("diagram_evaluate", payload)
        rm = res["data"]["render_model"]
        codes = [w["code"] for w in rm["warnings"]]
        assert "diagram.unassigned_pool_person" in codes
        assert rm["status"] == "incomplete"

    def test_evaluate_personid_outside_stage(self):
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        state["nodes"][0]["personId"] = (
            "99999999-9999-4999-8999-999999999999")
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate",
                  {"case_id": 42, "diagram": {"state": state}})
        assert exc.value.code == "diagram_reference_outside_stage"

    def test_evaluate_invalid_state(self):
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        state["version"] = 2                              # v3 bat buoc
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate",
                  {"case_id": 42, "diagram": {"state": state}})
        assert exc.value.code == "diagram_invalid_state"
        assert exc.value.details["errors"]

    def test_evaluate_legacy_node_flags_invalid(self):
        """isLandOwner/willReceive tren node -> invalid_node (truong la)."""
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        state["nodes"][0]["isLandOwner"] = True
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate",
                  {"case_id": 42, "diagram": {"state": state}})
        assert exc.value.code == "diagram_invalid_state"
        codes = {e["code"] for e in exc.value.details["errors"]}
        assert "invalid_node" in codes

    def test_evaluate_position_out_of_range(self):
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        state["nodes"][1]["receivePositions"] = [6]
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate",
                  {"case_id": 42, "diagram": {"state": state}})
        assert exc.value.code == "diagram_invalid_state"
        codes = {e["code"] for e in exc.value.details["errors"]}
        assert "invalid_position" in codes

    def test_evaluate_domain_mismatch(self):
        """state.domain hop le nhung khac case_type ->
        diagram_domain_mismatch TRUOC wire check khac (§13.5)."""
        mock.reset_backend(_load_fixture("ready.json"))
        state = _tp_state()                     # domain two_party tren
        with pytest.raises(CommandError) as exc:    # case inheritance
            _call("diagram_evaluate",
                  {"case_id": 42, "diagram": {"state": state}})
        assert exc.value.code == "diagram_domain_mismatch"
        assert exc.value.details["got"] == "two_party"

    def test_evaluate_draft_mode(self):
        """§2.1a/§7.4 — case_id absent + stage payload; khong tao case."""
        owner = _mk_person()
        spouse = _mk_person(ho_ten="Người Mẫu B", gioi_tinh="Nữ",
                            ngay_chet=None)
        stage = {"owner_row_id": owner["row_id"],
                 "people": [owner, spouse], "assets": [_mk_asset()]}
        state = {"version": 3, "domain": "inheritance", "nodes": [
            _node("owner", owner["row_id"], spouse="spouse", own=[1]),
            _node("spouse", spouse["row_id"], spouse="owner", receive=[1]),
        ]}
        res = _call("diagram_evaluate",
                    {"stage": stage, "diagram": {"state": state}})
        data = res["data"]
        assert data["schema_version"] == SCHEMA
        assert data["evaluated_revision"] is None
        assert data["render_model"]["engineVersion"] == 2
        # khong ghi gi vao state — case list khong thay doi
        before = _call("case_list", {})["data"]["total"]
        _call("diagram_evaluate",
              {"stage": stage, "diagram": {"state": state}})
        assert _call("case_list", {})["data"]["total"] == before

    def test_evaluate_draft_two_party(self):
        """Draft two_party: case.case_type hint + canonical 30 slot ->
        unsupported render_model (§13.5)."""
        stage = {"people": [_mk_person(), _mk_person()],
                 "assets": [_mk_asset()]}
        res = _call("diagram_evaluate", {
            "case": {"case_type": "two_party"},
            "stage": stage,
            "diagram": {"state": {
                "version": 3, "domain": "two_party",
                "nodes": [{"id": f"p{i}", "personId": None,
                           "hidden": False, "deleted": False}
                          for i in range(1, 31)]}}})
        rm = res["data"]["render_model"]
        assert rm["status"] == "unsupported"
        assert res["data"]["evaluated_revision"] is None

    def test_evaluate_draft_owner_required(self):
        """Draft inheritance thieu owner_row_id -> workspace_owner_required."""
        stage = {"people": [_mk_person()], "assets": [_mk_asset()]}
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate", {
                "stage": stage,
                "diagram": {"state": {"version": 3,
                                      "domain": "inheritance",
                                      "nodes": []}}})
        assert exc.value.code == "workspace_owner_required"

    def test_evaluate_draft_outside_stage(self):
        owner = _mk_person()
        stage = {"owner_row_id": owner["row_id"],
                 "people": [owner], "assets": [_mk_asset()]}
        state = {"version": 3, "domain": "inheritance", "nodes": [
            _node("owner", owner["row_id"], own=[1]),
            _node("child", str(uuid.uuid4()), parents=("owner",),
                  receive=[1]),
        ]}
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate",
                  {"stage": stage, "diagram": {"state": state}})
        assert exc.value.code == "diagram_reference_outside_stage"

    def test_evaluate_draft_stage_errors(self):
        owner = _mk_person()
        stage = {"owner_row_id": owner["row_id"],
                 "people": [_mk_person(ho_ten="   "),
                            owner],
                 "assets": [_mk_asset()]}
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate", {
                "stage": stage,
                "diagram": {"state": {"version": 3,
                                      "domain": "inheritance",
                                      "nodes": []}}})
        assert exc.value.code == "stage_validation_error"
        assert exc.value.details["field_errors"]

    def test_evaluate_draft_requires_stage(self):
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate", {
                "diagram": {"state": {"version": 3,
                                      "domain": "inheritance",
                                      "nodes": []}}})
        assert exc.value.code == "validation_error"

    def test_evaluate_missing_diagram_is_validation_error(self):
        """diagram/diagram.state thieu -> validation_error; state co mat
        ma sai -> diagram_invalid_state (oracle parity)."""
        mock.reset_backend(_load_fixture("ready.json"))
        for bad in (
                {},                                    # thieu diagram
                {"diagram": None},
                {"diagram": "x"},                      # khong phai object
                {"diagram": {}},                       # thieu state
                {"diagram": {"stae": {}}},             # key sai chinh ta
        ):
            p = {"case_id": 42}
            p.update(bad)
            with pytest.raises(CommandError) as exc:
                _call("diagram_evaluate", p)
            assert exc.value.code == "validation_error", bad
        # state co mat nhung khong phai object -> diagram_invalid_state
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate",
                  {"case_id": 42, "diagram": {"state": None}})
        assert exc.value.code == "diagram_invalid_state"

    def test_save_missing_diagram_is_validation_error(self):
        mock.reset_backend(_load_fixture("ready.json"))
        with pytest.raises(CommandError) as exc:
            _call("diagram_save",
                  {"case_id": 42, "base_revision": 7})
        assert exc.value.code == "validation_error"

    def test_evaluate_string_bool_is_invalid(self):
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        state["nodes"][0]["hidden"] = "false"     # chuoi, khong phai bool
        with pytest.raises(CommandError) as exc:
            _call("diagram_evaluate",
                  {"case_id": 42, "diagram": {"state": state}})
        assert exc.value.code == "diagram_invalid_state"

    def test_evaluate_read_only_on_done_case(self):
        mock.reset_backend(_load_fixture("ready.json"))
        res = _call("diagram_evaluate", {
            "case_id": 42,
            "diagram": {"state": _load_fixture("ready.json")
                        ["cases"]["42"]["diagram"]["state"]}})
        assert res["data"]["evaluated_revision"] == 7

    def test_evaluate_two_party_unsupported(self):
        """two_party committed: evaluate tra unsupported render, khong
        chay engine thua ke (§13.5)."""
        mock.reset_backend(_load_fixture("ready-two-party.json"))
        res = _call("diagram_evaluate", {
            "case_id": 48,
            "diagram": {"state": _tp_state()}})
        rm = res["data"]["render_model"]
        assert rm["status"] == "unsupported"
        assert rm["engineVersion"] == 2
        assert rm["allocations"] == {}
        assert rm["conservation"]["total"] == "0"
        codes = [w["code"] for w in rm["warnings"]]
        assert "diagram.two_party_unsupported" in codes
        assert res["data"]["evaluated_revision"] == 2

    def test_evaluate_does_not_persist(self):
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        state["nodes"].append(_node("extra_slot"))   # slot trong, hop le
        _call("diagram_evaluate", {
            "case_id": 42, "diagram": {"state": state}})
        res = _call("workspace_get", {"case_id": 42})
        assert res["data"]["case"]["revision"] == 7      # khong doi
        assert len(res["data"]["diagram"]["state"]["nodes"]) == 4

    def test_save_persists_and_bumps_revision(self):
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        payload = {"case_id": 42, "base_revision": 7,
                   "diagram": {"state": state}}
        res = _call("diagram_save", payload)
        data = res["data"]
        assert data["revision"] == 8
        assert data["diagram"]["render_model"]["status"] == "complete"
        assert data["diagram"]["state"]["domain"] == "inheritance"
        _no_legacy_keys(data)
        _assert_contract(_job_doc(
            "notary.diagram_save", payload, result=res,
            fixture_context={"stage_row_ids": [
                P_OWNER, P_SPOUSE, P_CHILD1, P_CHILD2],
                "stage_owner_row_id": P_OWNER}))
        got = _call("workspace_get", {"case_id": 42})
        assert got["data"]["case"]["revision"] == 8

    @pytest.mark.parametrize("error_code", [
        "dangling_parent", "dangling_spouse", "self_parent", "self_spouse",
        "ancestry_cycle", "too_many_parents", "invalid_node",
    ])
    def test_save_invalid_inheritance_reopens_exact_state_and_errors(
            self, monkeypatch, error_code):
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        owner = state["nodes"][0]
        if error_code == "dangling_parent":
            owner["parentSlotIds"] = ["absent"]
        elif error_code == "dangling_spouse":
            owner.update(spouseSlotIds=["absent"], spouseSlotId="absent")
        elif error_code == "self_parent":
            owner["parentSlotIds"] = ["owner"]
        elif error_code == "self_spouse":
            owner.update(spouseSlotIds=["owner"], spouseSlotId="owner")
        elif error_code == "ancestry_cycle":
            owner["parentSlotIds"] = ["child_1"]
        elif error_code == "invalid_node":
            owner["spouseUndetermined"] = True
        else:
            invalid = copy.deepcopy(_call(
                "workspace_get", {"case_id": 42})["data"]["diagram"]["render_model"])
            invalid.update(status="invalid", errors=[{
                "code": error_code, "message": "Một người chỉ có tối đa hai cha/mẹ.",
                "nodeId": "child_1"}])
            monkeypatch.setattr(mock, "_render", lambda *_args: invalid)
        saved = _call("diagram_save", {
            "case_id": 42, "base_revision": 7,
            "diagram": {"state": state}})["data"]
        assert saved["revision"] == 8
        assert saved["diagram"]["state"] == state
        rm = saved["diagram"]["render_model"]
        assert rm["status"] == "invalid"
        assert error_code in {e["code"] for e in rm["errors"]}
        reopened = _call("workspace_get", {"case_id": 42})["data"]
        assert reopened["case"]["revision"] == 8
        assert reopened["diagram"]["state"] == state
        assert reopened["diagram"]["render_model"] == rm
        monkeypatch.undo()

    @pytest.mark.parametrize("field", ["so_thua_dat", "so_to_ban_do"])
    def test_save_blank_parcel_or_sheet_preserves_workspace(self, field):
        doc = _load_fixture("ready.json")
        doc["cases"]["42"]["stage"]["assets"][0][field] = "  "
        mock.reset_backend(doc)
        before = _call("workspace_get", {"case_id": 42})["data"]
        with pytest.raises(CommandError) as exc:
            _call("diagram_save", {
                "case_id": 42, "base_revision": 7,
                "diagram": {"state": _ready_state()}})
        assert exc.value.code == "stage_validation_error"
        assert any(e["field"] == field and e["code"] == "required"
                   for e in exc.value.details["field_errors"])
        assert _call("workspace_get", {"case_id": 42})["data"] == before

    def test_save_owner_mismatch(self):
        """Node owner personId != owner_row_id da commit ->
        diagram_owner_mismatch (§13.4 mirror rule)."""
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        # owner.personId -> spouse row; giai phong node spouse de khong
        # bi duplicate_person (wire check chay truoc mirror check).
        state["nodes"][0]["personId"] = P_SPOUSE
        state["nodes"][1]["personId"] = None
        with pytest.raises(CommandError) as exc:
            _call("diagram_save", {
                "case_id": 42, "base_revision": 7,
                "diagram": {"state": state}})
        assert exc.value.code == "diagram_owner_mismatch"
        assert exc.value.details["owner_row_id"] == P_OWNER

    def test_save_prunes_positions(self):
        """Dau chon vi tri > len(assets) bi prune + warning
        diagram.selection_pruned (§13.4)."""
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        # 2 assets -> position 3 luon vuot
        state["nodes"][1]["receivePositions"] = [1, 2, 3]
        res = _call("diagram_save", {
            "case_id": 42, "base_revision": 7,
            "diagram": {"state": state}})
        dg = res["data"]["diagram"]
        codes = [w["code"] for w in dg.get("warnings") or []]
        assert "diagram.selection_pruned" in codes
        node = next(n for n in dg["state"]["nodes"] if n["id"] == "spouse")
        assert node["receivePositions"] == [1, 2]

    def test_save_unassigned_pool_warning_parity(self):
        """Save cũng giữ warning Pool + incomplete như evaluate (§6)."""
        doc = _load_fixture("diagram-warning.json")
        mock.reset_backend(doc)
        state = doc["request"]["diagram_evaluate"]["diagram"]["state"]
        res = _call("diagram_save", {
            "case_id": 46, "base_revision": 5,
            "diagram": {"state": state}})
        rm = res["data"]["diagram"]["render_model"]
        codes = [w["code"] for w in rm["warnings"]]
        assert "diagram.unassigned_pool_person" in codes
        assert rm["status"] == "incomplete"
        assert res["data"]["revision"] == 6

    def test_save_missing_owner_saves_incomplete_with_warning(self):
        mock.reset_backend(_load_fixture("ready.json"))
        state = _ready_state()
        for node in state["nodes"]:
            node["ownPositions"] = []
        saved = _call("diagram_save", {
            "case_id": 42, "base_revision": 7,
            "diagram": {"state": state}})["data"]
        assert saved["revision"] == 8
        assert saved["diagram"]["state"] == state
        rm = saved["diagram"]["render_model"]
        assert rm["status"] == "incomplete"
        assert any(w["code"] == "diagram.missing_land_owner"
                   for w in rm["warnings"])
        reopened = _call("workspace_get", {"case_id": 42})["data"]
        assert reopened["case"]["revision"] == 8
        assert reopened["diagram"]["state"] == state
        assert reopened["diagram"]["render_model"] == rm
        evaluated = _call("diagram_evaluate", {
            "case_id": 42, "diagram": {"state": state}})["data"]["render_model"]
        assert evaluated["status"] == "invalid"
        assert any(e["code"] == "missing_land_owner" for e in evaluated["errors"])
        assert any(w["code"] == "diagram.missing_land_owner"
                   for w in evaluated["warnings"])

    def test_save_two_party(self):
        mock.reset_backend(_load_fixture("ready-two-party.json"))
        res = _call("diagram_save", {
            "case_id": 48, "base_revision": 2,
            "diagram": {"state": _tp_state()}})
        data = res["data"]
        assert data["revision"] == 3
        assert data["diagram"]["state"]["domain"] == "two_party"
        assert data["diagram"]["render_model"]["status"] == "unsupported"

    def test_save_two_party_missing_slot_invalid(self):
        """29 slot -> diagram_invalid_state (missing_position) —
        canonical 30 bat buoc (§13.5)."""
        mock.reset_backend(_load_fixture("ready-two-party.json"))
        state = _tp_state()
        state["nodes"] = state["nodes"][:29]
        with pytest.raises(CommandError) as exc:
            _call("diagram_save", {
                "case_id": 48, "base_revision": 2,
                "diagram": {"state": state}})
        assert exc.value.code == "diagram_invalid_state"

    def test_save_conflict(self):
        mock.reset_backend(_load_fixture("ready.json"))
        with pytest.raises(CommandError) as exc:
            _call("diagram_save", {
                "case_id": 42, "base_revision": 3,
                "diagram": {"state": _ready_state()}})
        assert exc.value.code == "workspace_conflict"
        assert exc.value.details["server_revision"] == 7

    def test_save_done_case_allowed(self):
        mock.reset_backend(_load_fixture("ready.json"))
        res = _call("diagram_save", {
            "case_id": 42, "base_revision": 7,
            "diagram": {"state": _ready_state()}})
        assert res["data"]["revision"] == 8

    @pytest.mark.parametrize("state", [
        None, [], {"version": 2, "domain": "inheritance", "nodes": []},
        {"version": 3, "domain": "inheritance", "nodes": "bad"},
        {"version": 3, "domain": "inheritance", "nodes": [None]},
        {"version": 3, "domain": "inheritance", "nodes": [
            dict(_node("owner"), parentSlotIds=[{}])]},
        {"version": 3, "domain": "inheritance", "nodes": [
            dict(_node("owner"), hidden="false")]},
    ])
    def test_save_invalid_state_not_persisted(self, state):
        mock.reset_backend(_load_fixture("ready.json"))
        before = _call("workspace_get", {"case_id": 42})["data"]
        with pytest.raises(CommandError) as exc:
            _call("diagram_save", {
                "case_id": 42, "base_revision": 7,
                "diagram": {"state": state}})
        assert exc.value.code == "diagram_invalid_state"
        assert _call("workspace_get", {"case_id": 42})["data"] == before


# ---------- word_export_options / word_export_batch ----------

class TestWordExport:
    def _dest(self, tmp_path):
        return {"path": str(tmp_path), "scope": "machine_local",
                "is_dir": True}

    @pytest.mark.parametrize("case_type,people,assets,has_diagram,expected", [
        # Phần W (owner 23/10/2026): thiếu dữ kiện không còn là block.
        ("inheritance", 1, 1, False, None),
        ("inheritance", 0, 0, False, None),
        ("inheritance", 1, 0, False, None),
        ("inheritance", 1, 6, False, "word.too_many_assets"),
        ("inheritance", 21, 1, False, "word.too_many_people"),
        ("two_party", 1, 1, False, None),
        ("two_party", 1, 1, True, None),
        ("two_party", 31, 1, True, "word.too_many_people"),
        ("two_party", 1, 0, False, None),
    ])
    def test_folder_readiness_matches_real_technical_checks(
            self, case_type, people, assets, has_diagram, expected, monkeypatch):
        from types import SimpleNamespace
        monkeypatch.syspath_prepend(str(_REPO / "notary_v2"))
        from services.word_engine import word_block_reason

        case = {"case_type": case_type,
                "stage": {"people": [{"row_id": str(i)} for i in range(people)],
                          "assets": [{} for _ in range(assets)]},
                "diagram": {"state": {"nodes": [{}] if has_diagram else []}}}
        context = SimpleNamespace(case_type=case_type,
            persons=case["stage"]["people"], assets=case["stage"]["assets"],
            has_diagram_state=has_diagram)
        assert mock._word_block_reason(case, static=True) == expected
        assert word_block_reason(context, static=True) == expected

    @WINDOWS_ONLY
    def test_folder_options_and_export_allow_no_inheritance_owner_or_receiver(self, tmp_path):
        import docx

        seed = _load_fixture("ready.json")
        seed["cases"]["42"]["diagram"]["state"]["nodes"] = []
        mock.reset_backend(seed)
        templates, dest = tmp_path / "templates", tmp_path / "out"
        templates.mkdir()
        dest.mkdir()
        docx.Document().save(templates / "static.docx")
        payload = {"case_id": 42, "template_dir": self._dest(templates)}
        options = _call("word_export_options", payload)
        assert options["data"]["documents"][0]["ready"] is True
        batch = _call("word_export_batch", {**payload,
            "document_keys": ["static.docx"], "destination": self._dest(dest)})
        assert batch["data"]["documents"][0]["status"] == "saved"
        # Catalog callers dùng cùng kiểm tra kỹ thuật: thiếu chủ đất không block.
        legacy = _call("word_export_options", {"case_id": 42})
        docs = {d["document_key"]: d for d in legacy["data"]["documents"]}
        assert docs["khai_nhan_di_san"]["block_reason"] is None

    def test_options_ready_case(self):
        mock.reset_backend(_load_fixture("ready.json"))
        res = _call("word_export_options", {"case_id": 42})
        docs = {d["document_key"]: d for d in res["data"]["documents"]}
        assert docs["khai_nhan_di_san"]["ready"] is True
        assert docs["khai_nhan_di_san"]["block_reason"] is None
        assert docs["niem_yet"]["ready"] is False
        assert docs["niem_yet"]["block_reason"] == "word.template_missing"
        _assert_contract(_job_doc(
            "notary.word_export_options", {"case_id": 42}, result=res))

    def test_options_empty_case_ready(self):
        # W6: thiếu dữ kiện không chặn — readiness chỉ còn mức kỹ thuật/mẫu.
        mock.reset_backend(_load_fixture("empty.json"))
        res = _call("word_export_options", {"case_id": 43})
        docs = {d["document_key"]: d for d in res["data"]["documents"]}
        assert docs["khai_nhan_di_san"]["ready"] is True
        assert docs["khai_nhan_di_san"]["block_reason"] is None
        assert docs["niem_yet"]["block_reason"] == "word.template_missing"

    @WINDOWS_ONLY
    def test_two_party_word_allowed(self, tmp_path):
        """NAIA-8: word_export_* tren two_party khong con
        case_type_unsupported — fixture ready-two-party du dieu kien:
        options ready, batch ghi DOCX that."""
        mock.reset_backend(_load_fixture("ready-two-party.json"))
        res = _call("word_export_options", {"case_id": 48})
        docs = {d["document_key"]: d for d in res["data"]["documents"]}
        assert docs["khai_nhan_di_san"]["ready"] is True
        assert docs["niem_yet"]["block_reason"] == "word.template_missing"
        res = _call("word_export_batch", {
            "case_id": 48,
            "document_keys": ["khai_nhan_di_san"],
            "destination": self._dest(tmp_path)})
        d = res["data"]["documents"][0]
        assert d["status"] == "saved"
        assert Path(d["output_file"]["path"]).is_file()

    @WINDOWS_ONLY
    def test_batch_writes_real_docx(self, tmp_path):
        mock.reset_backend(_load_fixture("ready.json"))
        payload = {"case_id": 42,
                   "document_keys": ["khai_nhan_di_san",
                                     "thoa_thuan_phan_chia"],
                   "destination": self._dest(tmp_path)}
        res = _call("word_export_batch", payload)
        assert res.get("partial") is not True
        data = res["data"]
        assert data["breakdown"] == {
            "succeeded": ["khai_nhan_di_san", "thoa_thuan_phan_chia"],
            "failed": [], "skipped": []}
        import docx
        for d in data["documents"]:
            assert d["status"] == "saved"
            out = Path(d["output_file"]["path"])
            assert out.parent == tmp_path         # trong destination
            assert out.is_file()
            docx.Document(str(out))               # DOCX hop le mo duoc
            assert ".." not in d["actual_filename"]
        _assert_contract(_job_doc(
            "notary.word_export_batch", payload, result=res))

    @WINDOWS_ONLY
    def test_batch_collision_suffix(self, tmp_path):
        doc = _load_fixture("word-collision.json")
        mock.reset_backend(doc)
        # pre-existing files: base + _2 -> file moi phai la _3
        for name in doc["preexisting_files"]:
            (tmp_path / name).write_bytes(b"cu")
        payload = dict(doc["request"]["word_export_batch"])
        payload["destination"] = self._dest(tmp_path)
        res = _call("word_export_batch", payload)
        d = res["data"]["documents"][0]
        assert d["actual_filename"] == (
            "Van_ban_khai_nhan_di_san_HS-42_3.docx")
        # file cu khong bi ghi de
        assert (tmp_path / doc["preexisting_files"][0]).read_bytes() == b"cu"

    @WINDOWS_ONLY
    def test_batch_partial(self, tmp_path):
        doc = _load_fixture("word-partial.json")
        mock.reset_backend(doc)
        payload = dict(doc["request"]["word_export_batch"])
        payload["destination"] = self._dest(tmp_path)
        res = _call("word_export_batch", payload)
        assert res.get("partial") is True
        docs = {d["document_key"]: d for d in res["data"]["documents"]}
        assert docs["khai_nhan_di_san"]["status"] == "saved"
        assert docs["niem_yet"]["status"] == "failed"
        assert docs["niem_yet"]["error"]["code"] == "word.template_missing"
        assert res["data"]["breakdown"]["failed"] == ["niem_yet"]

    @WINDOWS_ONLY
    def test_batch_all_failed(self, tmp_path):
        doc = _load_fixture("word-all-failed.json")
        mock.reset_backend(doc)
        payload = dict(doc["request"]["word_export_batch"])
        payload["destination"] = self._dest(tmp_path)
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", payload)
        assert exc.value.code == "word_batch_failed"
        assert len(exc.value.details["documents"]) == 2
        assert list(tmp_path.iterdir()) == []      # khong file nao duoc tao
        # MIN-115: failed job van mang result.data.breakdown.failed
        bd = exc.value.result["data"]["breakdown"]
        assert len(bd["failed"]) == 2
        assert all(d["status"] == "failed"
                   for d in exc.value.result["data"]["documents"])

    def test_batch_validates_before_writing(self, tmp_path):
        mock.reset_backend(_load_fixture("ready.json"))
        # document_keys thieu/null/khong list -> validation_error
        # (word_no_documents_selected CHI cho list rong)
        for bad_keys in (None, "khai_nhan_di_san", 3, {"a": 1}):
            with pytest.raises(CommandError) as exc:
                _call("word_export_batch", {
                    "case_id": 42, "document_keys": bad_keys,
                    "destination": self._dest(tmp_path)})
            assert exc.value.code == "validation_error", bad_keys
        p_no_keys = {"case_id": 42, "destination": self._dest(tmp_path)}
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", p_no_keys)
        assert exc.value.code == "validation_error"
        # document_keys rong [] -> word_no_documents_selected
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": [],
                "destination": self._dest(tmp_path)})
        assert exc.value.code == "word_no_documents_selected"
        # key trung
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42,
                "document_keys": ["khai_nhan_di_san", "khai_nhan_di_san"],
                "destination": self._dest(tmp_path)})
        assert exc.value.code == "word_duplicate_document_key"
        # key ngoai catalog
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": ["khong_co_trong_catalog"],
                "destination": self._dest(tmp_path)})
        assert exc.value.code == "word_unknown_document_key"
        # key sai pattern
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": ["BAD-KEY"],
                "destination": self._dest(tmp_path)})
        assert exc.value.code == "word_unknown_document_key"
        assert list(tmp_path.iterdir()) == []

    @WINDOWS_ONLY
    def test_batch_destination_rules(self, tmp_path):
        mock.reset_backend(_load_fixture("ready.json"))
        # is_dir false -> validation_error
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": ["khai_nhan_di_san"],
                "destination": {"path": str(tmp_path),
                                "scope": "machine_local", "is_dir": False}})
        assert exc.value.code == "validation_error"
        # thieu is_dir -> validation_error
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": ["khai_nhan_di_san"],
                "destination": {"path": str(tmp_path),
                                "scope": "machine_local"}})
        assert exc.value.code == "validation_error"
        # folder khong ton tai -> file_not_found
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": ["khai_nhan_di_san"],
                "destination": {"path": str(tmp_path / "khong-co"),
                                "scope": "machine_local", "is_dir": True}})
        assert exc.value.code == "file_not_found"

    @WINDOWS_ONLY
    def test_batch_dest_write_denied_file_locked(
            self, tmp_path, monkeypatch):
        """MIN-116 parity engine that: dest ACL deny-write -> probe fail
        -> MOI doc per-file file_locked (ke ca niem_yet co
        template_missing), all-failed -> word_batch_failed kem
        result.data.breakdown — khong hang, khong exception kho."""
        real_open = open
        root = str(tmp_path)

        def _deny(file, mode="r", *a, **kw):
            if str(file).startswith(root) and \
                    any(c in mode for c in "wax+"):
                raise PermissionError(13, "Access is denied", str(file))
            return real_open(file, mode, *a, **kw)

        monkeypatch.setattr("builtins.open", _deny)
        mock.reset_backend(_load_fixture("ready.json"))
        keys = ["khai_nhan_di_san", "thoa_thuan_phan_chia", "niem_yet"]
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": keys,
                "destination": self._dest(tmp_path)})
        assert exc.value.code == "word_batch_failed"
        data = exc.value.result["data"]
        assert data["breakdown"] == {
            "succeeded": [], "failed": keys, "skipped": []}
        for d in data["documents"]:
            assert d["status"] == "failed"
            assert d["error"]["code"] == "file_locked"
        assert list(tmp_path.iterdir()) == []     # khong file nao sot lai

    def test_batch_unsupported_case_type(self, tmp_path):
        mock.reset_backend(_load_fixture("unsupported.json"))
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 45, "document_keys": ["khai_nhan_di_san"],
                "destination": self._dest(tmp_path)})
        assert exc.value.code == "case_type_unsupported"

    @WINDOWS_ONLY
    def test_batch_progress(self, tmp_path):
        mock.reset_backend(_load_fixture("ready.json"))
        job = _job("notary.word_export_batch")
        gw.dispatch("word_export_batch", job, {
            "case_id": 42, "document_keys": ["khai_nhan_di_san"],
            "destination": self._dest(tmp_path)})
        assert job.progress["done"] == 1
        assert job.progress["total"] == 1

    @WINDOWS_ONLY
    def test_batch_subdir_link_escape_rejected(self, tmp_path):
        """Parity engine that (review NAIA-8): destination/HS-<id> la
        junction/symlink tro ra ngoai -> mkdir(exist_ok) van pass; phai
        word_path_traversal va khong ghi file nao ben ngoai."""
        mock.reset_backend(_load_fixture("ready.json"))
        dest = tmp_path / "selected"
        dest.mkdir()
        outside = tmp_path / "outside"
        outside.mkdir()
        _dir_link(dest / "HS-42", outside)
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": ["khai_nhan_di_san"],
                "destination": {"path": str(dest),
                                "scope": "machine_local", "is_dir": True},
                "create_case_subdir": True})
        assert exc.value.code == "word_path_traversal"
        # khong file nao duoc tao — ngoai outside lan trong selected
        assert list(outside.iterdir()) == []
        assert [p.name for p in dest.iterdir()] == ["HS-42"]

    @WINDOWS_ONLY
    def test_batch_subdir_result_destination_is_granted_dir(
            self, tmp_path):
        """Contract §8.2: create_case_subdir=true -> ghi vao
        <dest>/HS-<id>/ nhung data.destination van la thu muc nguoi
        dung chon."""
        mock.reset_backend(_load_fixture("ready.json"))
        res = _call("word_export_batch", {
            "case_id": 42, "document_keys": ["khai_nhan_di_san"],
            "destination": self._dest(tmp_path),
            "create_case_subdir": True})
        data = res["data"]
        assert data["destination"]["path"] == str(tmp_path)
        assert (tmp_path / "HS-42"
                / "Van_ban_khai_nhan_di_san_HS-42.docx").is_file()

    @WINDOWS_ONLY
    def test_folder_mode_skips_link_to_outside(self, tmp_path):
        """Mau symlink tro ra ngoai thu muc mau: bi loai khoi scan ->
        options khong list, batch tu choi word_unknown_document_key,
        khong doc file ngoai thu muc duoc cap quyen."""
        mock.reset_backend(_load_fixture("ready.json"))
        tdir = tmp_path / "templates"
        tdir.mkdir()
        (tdir / "mau.docx").write_bytes(b"real")
        outside = tmp_path / "outside.docx"
        outside.write_bytes(b"outside-bytes")
        _file_link(tdir / "linked.docx", outside)
        tref = {"path": str(tdir), "scope": "machine_local",
                "is_dir": True}
        res = _call("word_export_options", {"case_id": 42,
                                            "template_dir": tref})
        keys = [d["document_key"] for d in res["data"]["documents"]]
        assert keys == ["mau.docx"]
        dest = tmp_path / "out"
        dest.mkdir()
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": ["linked.docx"],
                "template_dir": tref,
                "destination": {"path": str(dest),
                                "scope": "machine_local",
                                "is_dir": True}})
        assert exc.value.code == "word_unknown_document_key"
        assert list(dest.iterdir()) == []

    @WINDOWS_ONLY
    def test_folder_template_swapped_for_link(self, tmp_path,
                                              monkeypatch):
        """Parity real: mau quet duoc bi doi thanh symlink ra ngoai
        truoc luc ghi -> per-file word.template_missing (kiem tra lai
        containment tai thoi diem mo)."""
        mock.reset_backend(_load_fixture("ready.json"))
        tdir = tmp_path / "templates"
        tdir.mkdir()
        target = tdir / "mau.docx"
        target.write_bytes(b"real")
        outside = tmp_path / "outside.docx"
        outside.write_bytes(b"outside-bytes")
        # snapshot scan luc file con that — bang cach gia lap ket qua
        # scan (fixture ready chi co document_keys catalog; folder mode
        # quet lai moi lan nen phai patch de mo phong scan cu)
        monkeypatch.setattr(mock, "_scan_template_dir",
                            lambda _d: [{"document_key": "mau.docx",
                                         "display_name": "mau"}])
        target.unlink()
        _file_link(target, outside)
        dest = tmp_path / "out"
        dest.mkdir()
        with pytest.raises(CommandError) as exc:
            _call("word_export_batch", {
                "case_id": 42, "document_keys": ["mau.docx"],
                "template_dir": {"path": str(tdir),
                                 "scope": "machine_local",
                                 "is_dir": True},
                "destination": {"path": str(dest),
                                "scope": "machine_local",
                                "is_dir": True}})
        err = exc.value
        assert err.code == "word_batch_failed"
        d = err.result["data"]["documents"][0]
        assert d["status"] == "failed"
        assert d["error"]["code"] == "word.template_missing"
        assert list(dest.iterdir()) == []


# ---------- job lifecycle qua JobStore that ----------

class TestJobLifecycle:
    def _store(self):
        return JobStore(max_workers=2)

    def test_succeeded_status(self):
        store = self._store()
        try:
            job = store.submit(str(uuid.uuid4()), "notary.workspace_get",
                               reg.COMMANDS["notary.workspace_get"],
                               {"case_id": 42})
            final = _wait_terminal(store, job.job_id)
            assert final.status == "succeeded"
            assert final.result["data"]["backend_mode"] == "mock"
        finally:
            store.drain(2)

    def test_partial_status_and_breakdown(self):
        mock.reset_backend(_load_fixture("intake-partial.json"))
        store = self._store()
        try:
            payload = _load_fixture("intake-partial.json")[
                "request"]["intake_analyze"]
            job = store.submit(str(uuid.uuid4()), "notary.intake_analyze",
                               reg.COMMANDS["notary.intake_analyze"], payload)
            final = _wait_terminal(store, job.job_id)
            assert final.status == "partial"
            assert set(final.result["data"]["breakdown"]) >= {
                "succeeded", "failed"}
        finally:
            store.drain(2)

    def test_failed_error_shape(self):
        store = self._store()
        try:
            job = store.submit(str(uuid.uuid4()), "notary.workspace_get",
                               reg.COMMANDS["notary.workspace_get"],
                               {"case_id": 9999})
            final = _wait_terminal(store, job.job_id)
            assert final.status == "failed"
            err = final.error
            assert err["code"] == "case_not_found"
            for k in ("code", "message", "retryable", "next_action",
                      "job_id", "details"):
                assert k in err
        finally:
            store.drain(2)

    @WINDOWS_ONLY
    def test_idempotent_command_id_no_duplicate_files(self, tmp_path):
        """Retry cung command_id -> cung job -> khong tao them file."""
        mock.reset_backend(_load_fixture("ready.json"))
        store = self._store()
        try:
            cid = str(uuid.uuid4())
            payload = {"case_id": 42,
                       "document_keys": ["khai_nhan_di_san"],
                       "destination": {"path": str(tmp_path),
                                       "scope": "machine_local",
                                       "is_dir": True}}
            handler = reg.COMMANDS["notary.word_export_batch"]
            j1 = store.submit(cid, "notary.word_export_batch", handler,
                              payload)
            _wait_terminal(store, j1.job_id)
            j2 = store.submit(cid, "notary.word_export_batch", handler,
                              payload)
            assert j2.job_id == j1.job_id
            files = list(tmp_path.glob("*.docx"))
            assert len(files) == 1            # khong _2
        finally:
            store.drain(2)

    @WINDOWS_ONLY
    def test_word_export_canceled_mid_batch(self, tmp_path, monkeypatch):
        """Cancel giua batch: file da ghi giu nguyen, file chua bat dau
        khong duoc tao; job canceled{user_canceled} VAN mang
        result.data.breakdown.skipped len wire (MIN-115)."""
        monkeypatch.setattr(mock, "WORD_DOC_DELAY", 0.4)
        doc = _load_fixture("word-canceled.json")
        mock.reset_backend(doc)
        store = self._store()
        try:
            payload = dict(doc["request"]["word_export_batch"])
            payload["destination"] = {"path": str(tmp_path),
                                      "scope": "machine_local",
                                      "is_dir": True}
            job = store.submit(str(uuid.uuid4()),
                               "notary.word_export_batch",
                               reg.COMMANDS["notary.word_export_batch"],
                               payload)
            # cho van ban dau duoc ghi roi cancel
            deadline = time.time() + 8
            while time.time() < deadline:
                cur = store.get(job.job_id)
                if cur.progress and cur.progress["done"] >= 1:
                    break
                time.sleep(0.02)
            assert store.cancel(job.job_id) == "ok"
            final = _wait_terminal(store, job.job_id)
            assert final.status == "canceled"
            assert final.error["code"] == "user_canceled"
            files = sorted(p.name for p in tmp_path.glob("*.docx"))
            # file dau giu lai; hai file sau khong bao gio duoc tao
            assert files == ["Van_ban_khai_nhan_di_san_HS-42.docx"]
            # contract §8.4: canceled job van tra breakdown.skipped
            bd = final.result["data"]["breakdown"]
            assert bd["succeeded"] == ["khai_nhan_di_san"]
            assert bd["skipped"] == ["thoa_thuan_phan_chia", "niem_yet"]
        finally:
            store.drain(2)


# ---------- cross-cutting contract rules ----------

class TestContractInvariants:
    def test_schema_version_on_all_kinds(self, tmp_path):
        mock.reset_backend(_load_fixture("ready.json"))
        ready_stage = _load_fixture("ready.json")["cases"]["42"]["stage"]
        calls = {
            "workspace_get": {"case_id": 42},
            "intake_analyze": {"case_id": 42, "sources": [
                {"source_id": "aaaaaaaa-0000-4000-8000-000000000001",
                 "kind": "text", "text": "mau"}]},
            "workspace_commit_stage": {
                "case_id": 42, "base_revision": 7,
                "stage": copy.deepcopy(ready_stage)},   # giu stage de
                                                       # personId hop le
            "diagram_evaluate": {
                "case_id": 42, "diagram": {"state": _ready_state()}},
            "diagram_save": {
                "case_id": 42, "base_revision": 8,
                "diagram": {"state": _ready_state()}},
            "word_export_options": {"case_id": 42},
            "word_export_batch": {
                "case_id": 42, "document_keys": ["khai_nhan_di_san"],
                "destination": {"path": str(tmp_path),
                                "scope": "machine_local", "is_dir": True}},
        }
        if os.name != "nt":
            calls.pop("word_export_batch")   # destination FileRef Windows-only
        for name, payload in calls.items():
            res = _call(name, payload)
            assert res["kind"] == name
            assert res["data"]["schema_version"] == SCHEMA, name

    def test_no_v1_fields_anywhere(self):
        """workspace_get tra wire v2: khong is_primary/isLandOwner/
        willReceive o bat ky level nao."""
        for fixture in ("ready.json", "ready-two-party.json",
                        "done.json"):
            mock.reset_backend(_load_fixture(fixture))
            cid = 48 if fixture == "ready-two-party.json" else (
                44 if fixture == "done.json" else 42)
            res = _call("workspace_get", {"case_id": cid})
            _no_legacy_keys(res["data"])

    def test_no_confirmed_anywhere(self):
        mock.reset_backend(_load_fixture("intake-partial.json"))
        res = _call("intake_analyze", _load_fixture("intake-partial.json")
                    ["request"]["intake_analyze"])
        _no_confirmed_key(res)

    def test_deterministic_suggestion_ids(self):
        mock.reset_backend(_load_fixture("empty.json"))
        payload = {"case_id": 43, "sources": [
            {"source_id": "aaaaaaaa-0000-4000-8000-000000000001",
             "kind": "text", "text": "Người Mẫu I"}]}
        r1 = _call("intake_analyze", payload)
        r2 = _call("intake_analyze", payload)
        assert (r1["data"]["suggestions"][0]["suggestion_id"]
                == r2["data"]["suggestions"][0]["suggestion_id"])


# ---------- case_list (MIN-112 — mock parity cho overview) ----------

class TestCaseListMock:
    """notary.case_list route qua gateway: mock tra danh sach fixture case
    dung shape _case_row cua real adapter de overview chay duoc
    khong can engine that."""

    def test_registry_routes_case_list_via_gateway(self):
        job = _job("notary.case_list")
        res = reg.COMMANDS["notary.case_list"](job, {})
        assert res["kind"] == "case_list"
        ids = {c["id"] for c in res["data"]["cases"]}
        assert {42, 43, 44, 45, 46, 48} <= ids
        assert res["data"]["total"] == len(res["data"]["cases"])

    def test_case_list_row_shape_and_query(self):
        res = _call("case_list", {})
        c42 = next(c for c in res["data"]["cases"] if c["id"] == 42)
        # shape parity voi notary_adapter._case_row
        for k in ("id", "nguoi_chet", "tai_san", "ngay_lap_ho_so",
                  "loai_van_ban", "trang_thai"):
            assert k in c42, f"thieu key {k}"
        assert c42["nguoi_chet"]["ho_ten"] == "Người Mẫu A"
        assert c42["tai_san"]["so_serial"] == "MM000001"
        assert c42["trang_thai"] == "draft"
        # query filter giong real adapter: substring tren row json —
        # '"id": 45' chi khop case 45 (row luon co truong "id": <so>).
        res2 = _call("case_list", {"query": '"id": 45'})
        assert {c["id"] for c in res2["data"]["cases"]} == {45}
        res3 = _call("case_list", {"query": "khong-co-gi"})
        assert res3["data"]["cases"] == []


# ---------- MIN-141 đợt 3: case meta + danh bạ ủy quyền ----------

class TestCaseMetaMock:
    """Parity voi real backend (đợt 3): payload.case di qua create/commit,
    canonical alias fold + conflict, uq id resolve qua danh ba gia."""

    def _commit_payload(self, case_id=43, base_revision=1, meta=None):
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        p = {"case_id": case_id, "base_revision": base_revision,
             "stage": {"owner_row_id": rid,
                       "people": [{"row_id": rid, "entity_id": None,
                                   "ho_ten": "Người Mẫu X"}],
                       "assets": []}}
        if meta is not None:
            p["case"] = meta
        return p

    def test_create_meta_roundtrip(self):
        res = _call("workspace_create", _create_payload(
            case_meta={"document_type": "khai_nhan",
                       "noi_niem_yet": "xã Yên Sở",
                       "nguoi_nhan_uy_quyen": "Người UQ Tự Do",
                       "noi_dung_viec": "Khai nhận thừa kế",
                       "ghi_chu": "Ghi chú"}))
        c = res["data"]["case"]
        assert c["noi_niem_yet"] == "xã Yên Sở"
        assert c["nguoi_nhan_uy_quyen"] == "Người UQ Tự Do"
        assert c["nguoi_nhan_uy_quyen_id"] is None   # free text → id null
        assert c["noi_dung_viec"] == "Khai nhận thừa kế"
        assert c["ghi_chu"] == "Ghi chú"
        # workspace_get tra lai dung cac field moi.
        got = _call("workspace_get", {"case_id": c["id"]})
        for k in ("noi_niem_yet", "nguoi_nhan_uy_quyen",
                  "nguoi_nhan_uy_quyen_id", "noi_dung_viec"):
            assert k in got["data"]["case"], f"thieu {k} tren wire"

    def test_create_meta_canonical_alias(self):
        res = _call("workspace_create", _create_payload(
            case_meta={"document_type": "khai_nhan",
                       "noiniemyet": "xã Canonical",
                       "noidungviec": "Nội dung canonical"}))
        c = res["data"]["case"]
        assert c["noi_niem_yet"] == "xã Canonical"
        assert c["noi_dung_viec"] == "Nội dung canonical"

    def test_create_meta_alias_conflict_rejected(self):
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", _create_payload(
                case_meta={"document_type": "khai_nhan",
                           "noiniemyet": "xã A",
                           "noi_niem_yet": "xã B"}))
        assert exc.value.code == "validation_error"

    def test_create_uq_id_unknown_rejected(self):
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", _create_payload(
                case_meta={"document_type": "khai_nhan",
                           "nguoi_nhan_uy_quyen_id": 777}))
        assert exc.value.code == "validation_error"
        assert "không tồn tại" in exc.value.message

    def test_create_uq_id_resolves_name_from_catalog(self):
        mock.reset_backend({"customers": [
            {"id": 9, "ho_ten": "Người Danh Bạ", "so_giay_to": None}],
            "cases": _load_fixture("empty.json")["cases"]})
        res = _call("workspace_create", _create_payload(
            case_meta={"document_type": "khai_nhan",
                       "nguoi_nhan_uy_quyen_id": 9}))
        c = res["data"]["case"]
        assert c["nguoi_nhan_uy_quyen"] == "Người Danh Bạ"
        assert c["nguoi_nhan_uy_quyen_id"] == 9

    def test_create_uq_name_id_mismatch_rejected(self):
        mock.reset_backend({"customers": [
            {"id": 9, "ho_ten": "Người Danh Bạ", "so_giay_to": None}],
            "cases": _load_fixture("empty.json")["cases"]})
        with pytest.raises(CommandError) as exc:
            _call("workspace_create", _create_payload(
                case_meta={"document_type": "khai_nhan",
                           "nguoi_nhan_uy_quyen": "Tên Khác",
                           "nguoi_nhan_uy_quyen_id": 9}))
        assert exc.value.code == "validation_error"

    def test_commit_meta_atomic(self):
        mock.reset_backend(_load_fixture("empty.json"))
        meta = {"noi_niem_yet": "xã Commit", "noi_dung_viec": "Việc",
                "nguoi_nhan_uy_quyen": "Người UQ", "ghi_chu": None}
        res = _call("workspace_commit_stage",
                    self._commit_payload(meta=meta))
        assert res["data"]["revision"] == 2
        got = _call("workspace_get", {"case_id": 43})["data"]["case"]
        assert got["noi_niem_yet"] == "xã Commit"
        assert got["nguoi_nhan_uy_quyen"] == "Người UQ"
        assert got["nguoi_nhan_uy_quyen_id"] is None
        assert got["noi_dung_viec"] == "Việc"

    def test_commit_meta_canonical_alias(self):
        mock.reset_backend(_load_fixture("empty.json"))
        res = _call("workspace_commit_stage", self._commit_payload(
            meta={"noiniemyet": "xã Canonical", "noidungviec": "N"}))
        assert res["data"]["revision"] == 2
        got = _call("workspace_get", {"case_id": 43})["data"]["case"]
        assert got["noi_niem_yet"] == "xã Canonical"

    def test_commit_meta_invalid_rolls_back_stage(self):
        """Meta loi → ca commit roll back (atomic): stage khong doi."""
        mock.reset_backend(_load_fixture("empty.json"))
        p = self._commit_payload(meta={"case_type": "two_party"})
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", p)
        assert exc.value.code == "validation_error"
        got = _call("workspace_get", {"case_id": 43})["data"]
        assert got["case"]["revision"] == 1
        assert got["stage"]["people"] == []

    def test_commit_without_meta_keeps_case(self):
        mock.reset_backend(_load_fixture("empty.json"))
        res = _call("workspace_commit_stage", self._commit_payload())
        assert res["data"]["revision"] == 2
        got = _call("workspace_get", {"case_id": 43})["data"]["case"]
        assert got["noi_niem_yet"] is None   # default, khong bi xoa

    def test_person_canonical_alias_fold(self):
        """Canonical person keys (ten/diachi/…) fold vao legacy snake."""
        mock.reset_backend(_load_fixture("empty.json"))
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        p = {"case_id": 43, "base_revision": 1,
             "stage": {"owner_row_id": rid,
                       "people": [{"row_id": rid, "entity_id": None,
                                   "ten": "Người Canonical",
                                   "diachi": "Địa chỉ canonical",
                                   "loaigiayto": "CCCD",
                                   "loaidiachi": "Thường trú"}],
                       "assets": []}}
        res = _call("workspace_commit_stage", p)
        row = res["data"]["stage"]["people"][0]
        assert row["ho_ten"] == "Người Canonical"
        assert row["dia_chi"] == "Địa chỉ canonical"
        assert row["loai_giay_to"] == "CCCD"
        assert row["loai_dia_chi"] == "Thường trú"

    def test_person_alias_conflict_rejected(self):
        mock.reset_backend(_load_fixture("empty.json"))
        rid = "eeeeeeee-eeee-4eee-8eee-eeeeeeeeeeee"
        p = {"case_id": 43, "base_revision": 1,
             "stage": {"owner_row_id": rid,
                       "people": [{"row_id": rid, "entity_id": None,
                                   "ten": "A", "ho_ten": "B"}],
                       "assets": []}}
        with pytest.raises(CommandError) as exc:
            _call("workspace_commit_stage", p)
        assert exc.value.code == "validation_error"


class TestCustomerCatalogMock:
    """notary.customer_list/customer_create qua gateway (đợt 3)."""

    def test_customer_list_empty_default(self):
        res = _call("customer_list", {})
        assert res["kind"] == "customer_list"
        assert res["data"]["customers"] == []
        assert res["data"]["total"] == 0

    def test_customer_create_then_list(self):
        r = _call("customer_create", {"ho_ten": "Người Danh Bạ A"})
        assert r["kind"] == "customer_upsert"
        cust = r["data"]["customer"]
        assert cust["id"] >= 900
        assert cust["ho_ten"] == "Người Danh Bạ A"
        assert r["data"]["updated"] is False
        res = _call("customer_list", {})
        assert [c["id"] for c in res["data"]["customers"]] == [cust["id"]]

    def test_customer_create_requires_name(self):
        with pytest.raises(CommandError) as exc:
            _call("customer_create", {"ho_ten": "  "})
        assert exc.value.code == "validation_error"

    def test_customer_upsert_same_sogiayto(self):
        r1 = _call("customer_create",
                   {"ho_ten": "Tên Cũ", "so_giay_to": "S001"})
        r2 = _call("customer_create",
                   {"ho_ten": "Tên Mới", "so_giay_to": "S001"})
        assert r2["data"]["updated"] is True
        assert r2["data"]["customer"]["id"] == r1["data"]["customer"]["id"]
        assert r2["data"]["customer"]["ho_ten"] == "Tên Mới"

    def test_customer_list_query_filter(self):
        mock.reset_backend({"customers": [
            {"id": 1, "ho_ten": "Nguyễn An", "so_giay_to": "A1",
             "dia_chi": "Hà Nội"},
            {"id": 2, "ho_ten": "Trần Bình", "so_giay_to": "B2",
             "dia_chi": "Huế"}]})
        res = _call("customer_list", {"query": "nguyễn"})
        assert [c["id"] for c in res["data"]["customers"]] == [1]
        res = _call("customer_list", {"query": "huế"})
        assert [c["id"] for c in res["data"]["customers"]] == [2]

    def test_case_commit_with_catalog_uq(self):
        """UQ tu danh ba gia: commit id → get tra id + ten master."""
        mock.reset_backend({"customers": [
            {"id": 9, "ho_ten": "Người Danh Bạ", "so_giay_to": None}],
            "cases": _load_fixture("empty.json")["cases"]})
        p = TestCaseMetaMock()._commit_payload(
            meta={"nguoi_nhan_uy_quyen_id": 9})
        res = _call("workspace_commit_stage", p)
        assert res["data"]["revision"] == 2
        got = _call("workspace_get", {"case_id": 43})["data"]["case"]
        assert got["nguoi_nhan_uy_quyen"] == "Người Danh Bạ"
        assert got["nguoi_nhan_uy_quyen_id"] == 9


def test_mock_saved_conversion_has_real_boundary_shape_and_preserves_stage():
    mock.reset_backend(_load_fixture("ready.json"))
    before = _call("workspace_get", {"case_id": 42})["data"]
    result = _call("workspace_convert_type", {
        "case_id": 42, "base_revision": before["case"]["revision"],
        "case_type": "two_party", "document_type": "dat_coc",
        "confirm_diagram_reset": True})["data"]
    assert result["case"]["case_type"] == "two_party"
    assert result["stage"]["people"] == before["stage"]["people"]
    assert result["stage"]["assets"] == before["stage"]["assets"]
    assert "owner_row_id" not in result["stage"]
    assert len(result["diagram"]["state"]["nodes"]) == 30


# ---------- workspace_source_latest (§13.20, vòng 2) ----------

def test_workspace_source_latest_mock_parity():
    res = _call("workspace_source_latest", {})
    assert res["kind"] == "workspace_source_latest"
    data = res["data"]
    assert data["source_case_id"] == 44
    assert data["source_status"] == "done"
    stage = data["stage"]
    ws = _call("workspace_get", {"case_id": 44})["data"]
    assert stage["owner_row_id"] == stage["people"][0]["row_id"]
    original = {row["row_id"]
                for row in ws["stage"]["people"] + ws["stage"]["assets"]}
    returned = [row["row_id"]
                for row in stage["people"] + stage["assets"]]
    assert returned and all(rid not in original for rid in returned)
    assert [p["entity_id"] for p in stage["people"]] == [
        p["entity_id"] for p in ws["stage"]["people"]]
    assert [a["entity_id"] for a in stage["assets"]] == [
        a["entity_id"] for a in ws["stage"]["assets"]]
    assert set(data) == {"schema_version", "backend_mode",
                         "source_case_id", "source_status", "stage"}
    assert set(stage) == {"owner_row_id", "people", "assets"}


def test_workspace_source_latest_mock_empty_and_payload_errors():
    # Không có case Done (chỉ draft) → nguồn rỗng thành công.
    mock.reset_backend({"cases": {
        1: {"case_type": "inheritance", "status": "draft",
            "stage": {"people": [], "assets": []}},
    }})
    data = _call("workspace_source_latest", {})["data"]
    assert data["source_case_id"] is None
    assert data["source_status"] is None
    assert data["stage"] == {"people": [], "assets": []}
    for payload in (None, {"case_id": 44}, {"filter": "x"}):
        with pytest.raises(CommandError) as exc:
            _call("workspace_source_latest", payload)
        assert exc.value.code == "validation_error"


def test_workspace_source_latest_mock_unsupported_and_saved_at_order():
    stage = {"people": [], "assets": []}
    mock.reset_backend({"cases": {
        1: {"case_type": "gift", "status": "done", "stage": stage},
    }})
    with pytest.raises(CommandError) as exc:
        _call("workspace_source_latest", {})
    assert exc.value.code == "case_type_unsupported"
    assert exc.value.details["case_type"] == "gift"

    mock.reset_backend({"cases": {
        1: {"case_type": "inheritance", "status": "done",
            "updated_at": "2026-10-22T08:00:00.100001", "stage": stage},
        2: {"case_type": "inheritance", "status": "done",
            "updated_at": "2026-10-22T08:00:00.100002", "stage": stage},
    }})
    assert _call("workspace_source_latest", {})["data"][
        "source_case_id"] == 2
    # Tie tuyệt đối → id cao thắng.
    mock.reset_backend({"cases": {
        1: {"case_type": "inheritance", "status": "done",
            "updated_at": "2026-10-22T08:00:01.000000", "stage": stage},
        2: {"case_type": "inheritance", "status": "done",
            "updated_at": "2026-10-22T08:00:01.000000", "stage": stage},
    }})
    assert _call("workspace_source_latest", {})["data"][
        "source_case_id"] == 2


# ---------- workspace_duplicate (§13.21, vòng 2) ----------

def test_workspace_duplicate_mock_parity_and_replay():
    mock.reset_backend(_load_fixture("ready.json"))
    before = _call("workspace_get", {"case_id": 42})["data"]
    key = str(uuid.uuid4())
    data = _call("workspace_duplicate", {
        "source_case_id": 42, "idempotency_key": key})["data"]
    assert data["created"] is True
    assert data["source_case_id"] == 42
    assert data["backend_mode"] == "mock"
    assert data["case"]["status"] == "draft"
    assert data["case"]["revision"] == 1
    copy_id = data["case"]["id"]
    assert copy_id != 42
    assert data["case"]["document_type"] == before["case"]["document_type"]

    old_person = [row["row_id"] for row in before["stage"]["people"]]
    new_person = [row["row_id"] for row in data["stage"]["people"]]
    old_asset = [row["row_id"] for row in before["stage"]["assets"]]
    new_asset = [row["row_id"] for row in data["stage"]["assets"]]
    assert [uuid.UUID(rid).version for rid in new_person + new_asset] == \
        [4] * (len(new_person) + len(new_asset))
    assert not set(new_person) & set(old_person)
    assert not set(new_asset) & set(old_asset)
    assert [row["entity_id"] for row in data["stage"]["people"]] == \
        [row["entity_id"] for row in before["stage"]["people"]]
    person_map = dict(zip(old_person, new_person))
    assert data["stage"]["owner_row_id"] == person_map[
        before["stage"]["owner_row_id"]]
    for old_node, new_node in zip(before["diagram"]["state"]["nodes"],
                                  data["diagram"]["state"]["nodes"]):
        assert old_node["id"] == new_node["id"]
        assert new_node["personId"] == (
            person_map.get(old_node["personId"])
            if old_node["personId"] else None)
        for field in ("parentSlotIds", "spouseSlotId", "ownPositions",
                      "receivePositions", "hidden", "deleted"):
            assert new_node[field] == old_node[field]
    assert _call("workspace_get", {"case_id": 42})["data"] == before

    replay = _call("workspace_duplicate", {
        "source_case_id": 42, "idempotency_key": key})["data"]
    assert replay["created"] is False
    assert replay["case"]["id"] == copy_id
    assert replay["case"]["revision"] == 1

    with pytest.raises(CommandError) as exc:
        _call("workspace_duplicate", {
            "source_case_id": 43, "idempotency_key": key})
    assert exc.value.code == "idempotency_conflict"


def test_workspace_duplicate_mock_payload_errors_and_unsupported():
    mock.reset_backend(_load_fixture("ready.json"))
    key = str(uuid.uuid4())
    for payload in (
            None,
            {},
            {"source_case_id": 42},
            {"source_case_id": True, "idempotency_key": key},
            {"source_case_id": 42, "idempotency_key": "no"},
            {"source_case_id": 42, "idempotency_key": key, "case_id": 42},
    ):
        with pytest.raises(CommandError) as exc:
            _call("workspace_duplicate", payload)
        assert exc.value.code == "validation_error"
    with pytest.raises(CommandError) as exc:
        _call("workspace_duplicate", {
            "source_case_id": 987654, "idempotency_key": key})
    assert exc.value.code == "case_not_found"
    assert exc.value.details["source_case_id"] == 987654

    mock.reset_backend({"cases": {
        1: {"case_type": "gift", "status": "done",
            "stage": {"people": [], "assets": []}},
    }})
    with pytest.raises(CommandError) as exc:
        _call("workspace_duplicate", {
            "source_case_id": 1, "idempotency_key": key})
    assert exc.value.code == "case_type_unsupported"
