"""MIN-69 — Task 3: versioned upload workflow (`upload.workflow.v1`).

Tests the `upload.scan` / `upload.audit_excel` / `upload.queue_get` contract
commands end-to-end against the real engine services: fixtures are generated
DOCX/Excel files in a TemporaryDirectory, commands are dispatched through the
real command registry with real ``jobstore.Job`` objects, and every command
runs against a temporary ``G1_UPLOAD_DATA_DIR`` workspace.

Run from the ``shell/`` directory:
    python test/test_upload_workflow.py
"""

import gc
import json
import os
import sys
import tempfile
import time
import unittest
import uuid
from pathlib import Path
from unittest import mock

_HERE = Path(__file__).resolve().parent
_SHELL = _HERE.parent
sys.path.insert(0, str(_SHELL / "sidecar"))

from errors import CommandError                      # noqa: E402
from jobstore import Job                              # noqa: E402
import command_registry as reg                        # noqa: E402
import engine_roots                                   # noqa: E402
import upload_adapter                                 # noqa: E402
import upload_workspace                               # noqa: E402


V1 = "upload.workflow.v1"
WEBSITE = "nam_dinh"
FAKE_WEBSITE = "fake_b"

EXCEL_FROM = "2026-01-01"
EXCEL_TO = "2026-12-31"


def _engines_available():
    try:
        engine_roots.engine_root("upload_lab")
        import docx      # noqa: F401
        import openpyxl  # noqa: F401
        return True
    except Exception:
        return False


def _write_docx(path, lines):
    """Write a real .docx file (python-docx) so the engine scan extracts it."""
    import docx

    doc = docx.Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(str(path))


def _write_excel(path, rows, header=True):
    """Write a real .xlsx file: col A = so cong chung, col B = ngay."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if header:
        ws.append(["Số công chứng", "Ngày công chứng"])
    for row in rows:
        ws.append(list(row))
    wb.save(str(path))


def _contract_lines(so, *, full=True, date_text="15 tháng 04 năm 2026"):
    """DOCX body lines that the engine extractor reads as one contract.

    ``full=False`` omits the title/asset paragraphs so the record keeps
    missing fields (ghi_chu 'missing: ...' on the queue row).
    """
    lines = []
    if full:
        lines.append("HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT")
    lines.append(f"Số công chứng {so}")
    lines.append("Hôm nay, tại Văn phòng Công chứng Fixture, chúng tôi gồm:")
    lines.append("BÊN A: Ông Nguyễn Văn A, CCCD 001099000111.")
    lines.append("BÊN B: Bà Trần Thị B, CCCD 001088000222.")
    if full:
        lines.append(
            "Đối tượng của Hợp đồng này là quyền sử dụng đất "
            "có địa chỉ tại: thửa đất số 12, tờ bản đồ số 34, "
            "xã Fixture, tỉnh Nam Định.")
    lines.append("LỜI CHỨNG")
    lines.append(f"Văn phòng công chứng chứng nhận, ngày {date_text}")
    lines.append("CÔNG CHỨNG VIÊN")
    return lines


DEFAULT_DOCS = {
    "GD-138.docx": "138/2026/CCGD",
    "GD-139.docx": "139/2026/CCGD",
}


@unittest.skipUnless(_engines_available(), "upload_lab engine/docx/openpyxl missing")
class UploadWorkflowCase(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        # sqlite3.Connection giu statement-cache tao self-cycle → chi duoc
        # cyclic GC giai phong. Collect truoc khi xoa tempdir de nha handle
        # tren Windows (cleanup chay LIFO: lenh nay chay truoc tempdir).
        self.addCleanup(gc.collect)
        self.data_root = Path(self.tempdir.name) / "upload-data"
        self._old_env = os.environ.get("G1_UPLOAD_DATA_DIR")
        os.environ["G1_UPLOAD_DATA_DIR"] = str(self.data_root)
        upload_workspace.reset_store_for_tests()
        self.addCleanup(upload_workspace.reset_store_for_tests)
        self.addCleanup(self._restore_env)
        self.providers = engine_roots.import_engine_module(
            "upload_lab", "providers")
        self._register_fake()
        self.store = upload_workspace.open_store()
        self.folder = Path(self.tempdir.name) / "ho so"
        self.folder.mkdir()
        self.excel_path = Path(self.tempdir.name) / "so.xlsx"
        self.last_job = None

    # ---------------------------------------------------------- fixtures
    def _register_fake(self):
        class _Fake(self.providers.WebsiteProvider):
            def __init__(self):
                super().__init__(
                    website_id=FAKE_WEBSITE,
                    label="Cong chung fake B",
                    display_url="https://fake-b.example",
                    capabilities=("audit_excel",))

        self.providers.DEFAULT_REGISTRY.register(_Fake())
        self.addCleanup(
            lambda: self.providers.DEFAULT_REGISTRY._providers.pop(
                FAKE_WEBSITE, None))

    def _restore_env(self):
        if self._old_env is None:
            os.environ.pop("G1_UPLOAD_DATA_DIR", None)
        else:
            os.environ["G1_UPLOAD_DATA_DIR"] = self._old_env

    # ---------------------------------------------------------- helpers
    def _run(self, command, payload):
        """Dispatch through the real command registry with a real Job."""
        job = Job(str(uuid.uuid4()), command)
        job.status = "running"
        self.last_job = job
        return reg.COMMANDS[command](job, payload)

    def run_scan_fixture(self, docs=None, folder=None,
                         expected_revision=None):
        """Create DOCX fixtures in the temp folder, then run upload.scan.

        The scan runs the real engine batch scan against the temporary
        workspace; the returned result carries the real manifest FileRef.
        """
        docs = DEFAULT_DOCS if docs is None else docs
        target = Path(folder) if folder else self.folder
        for name, so in docs.items():
            _write_docx(target / name, _contract_lines(so))
        revision = (self.store.revision() if expected_revision is None
                    else expected_revision)
        return self._run("upload.scan", {
            "workflow_version": V1,
            "website_id": WEBSITE,
            "folder": {"path": str(target), "scope": "machine_local"},
            "expected_revision": revision,
            "full_rescan": False,
            "modified_since": None,
        })

    def run_audit_fixture(self, excel_path=None, from_date=EXCEL_FROM,
                          to_date=EXCEL_TO):
        excel = Path(excel_path) if excel_path else self.excel_path
        return self._run("upload.audit_excel", {
            "workflow_version": V1,
            "website_id": WEBSITE,
            "file_ref": {"path": str(excel), "scope": "machine_local"},
            "from_date": from_date,
            "to_date": to_date,
        })

    def queue_get(self, run_id, audit_id=None, website_id=WEBSITE):
        return self._run("upload.queue_get", {
            "workflow_version": V1,
            "website_id": website_id,
            "run_id": run_id,
            "audit_id": audit_id,
        })

    def _rows_by_contract(self, data):
        return {r["contract_no"]: r for r in data["folder_rows"]}


@unittest.skipUnless(_engines_available(), "upload_lab engine/docx/openpyxl missing")
class ScanAuditQueueTest(UploadWorkflowCase):

    # ---------------------------------------------------------- upload.scan
    def test_scan_v1_binds_manifest_to_run(self):
        result = self.run_scan_fixture()
        ref = result["data"]["manifest_ref"]
        self.assertTrue(Path(ref["path"]).is_file())
        manifest = json.loads(Path(ref["path"]).read_text(encoding="utf-8"))
        self.assertEqual(manifest["run_id"], result["data"]["run_id"])
        self.assertEqual(result["data"]["website_id"], "nam_dinh")

        data = result["data"]
        self.assertEqual(data["workflow_version"], V1)
        self.assertEqual(ref["scope"], "machine_local")
        self.assertEqual(ref["size_bytes"], Path(ref["path"]).stat().st_size)
        self.assertEqual(ref["sha256"],
                         upload_workspace.sha256_file(Path(ref["path"])))
        # Manifest lives under the website data dir, not the engine dir.
        site_dir = upload_workspace.website_data_dir(WEBSITE)
        self.assertIn(str(site_dir), str(Path(ref["path"]).parent))
        # The binding is stored in the workspace for later resolution.
        run = self.store.run_for(data["run_id"])
        self.assertIsNotNone(run)
        self.assertEqual(run["website_id"], WEBSITE)
        self.assertEqual(Path(run["manifest_path"]), Path(ref["path"]))
        # Records returned belong to this run and carry real IDs.
        self.assertEqual(len(data["records"]), 2)
        for rec in data["records"]:
            self.assertEqual(rec["status"], "extracted")
            self.assertIn("/CCGD", rec["contract_no"])
            self.assertGreaterEqual(int(rec["record_id"]), 1)
            self.assertTrue(rec["file_path"].endswith(".docx"))
        self.assertEqual(data["stats"]["processed_files"], 2)
        self.assertEqual(data["stats"]["candidates_found"], 2)
        self.assertEqual(data["revision"], self.store.revision())

    def test_scan_v1_result_fields_and_job_row(self):
        result = self.run_scan_fixture()
        data = result["data"]
        self.assertEqual(result["kind"], "scan_report")
        self.assertEqual(
            result["source_files"],
            [{"path": str(self.folder), "scope": "machine_local"}])
        job_row = self.store.job_for(self.last_job.job_id)
        self.assertIsNotNone(job_row)
        self.assertEqual(job_row["command"], "upload.scan")
        self.assertEqual(job_row["website_id"], WEBSITE)
        self.assertEqual(job_row["run_id"], data["run_id"])
        self.assertEqual(job_row["status"], "succeeded")

    def test_scan_v1_expected_revision_gate(self):
        self.run_scan_fixture()
        stale = self.store.revision() - 1
        with self.assertRaises(CommandError) as ctx:
            self.run_scan_fixture(expected_revision=stale)
        self.assertEqual(ctx.exception.code, "stale_revision")
        self.assertTrue(ctx.exception.retryable)
        self.assertEqual(ctx.exception.next_action, "retry")

    def test_scan_v1_folder_validation(self):
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.scan", {
                "workflow_version": V1,
                "website_id": WEBSITE,
                "folder": {"path": str(self.folder / "khong-co"),
                           "scope": "machine_local"},
                "expected_revision": self.store.revision(),
            })
        self.assertEqual(ctx.exception.code, "file_not_found")
        self.assertEqual(ctx.exception.next_action, "pick_files")

    def test_scan_v1_modified_since_and_key_validation(self):
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.scan", {
                "workflow_version": V1,
                "website_id": WEBSITE,
                "folder": {"path": str(self.folder), "scope": "machine_local"},
                "expected_revision": self.store.revision(),
                "modified_since": "",
            })
        self.assertEqual(ctx.exception.code, "validation_error")
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.scan", {
                "workflow_version": V1,
                "website_id": WEBSITE,
                "folder": {"path": str(self.folder), "scope": "machine_local"},
                "expected_revision": self.store.revision(),
                "modified_since": "15/04/2026",
            })
        self.assertEqual(ctx.exception.code, "validation_error")
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.scan", {
                "workflow_version": V1,
                "website_id": WEBSITE,
                "folder": {"path": str(self.folder), "scope": "machine_local"},
                "expected_revision": self.store.revision(),
                "portal_url": "https://x",
            })
        self.assertEqual(ctx.exception.code, "validation_error")

    # ---------------------------------------------------- upload.audit_excel
    def _excel_main(self):
        _write_excel(self.excel_path, [
            ("1/2026", "05/01/2026"),
            ("2/2026", "06/01/2026"),
            ("4/2026", "08/01/2026"),
            ("007/2026", "09/01/2026"),
        ])
        return self.excel_path

    def test_audit_excel_v1_summary_and_rows(self):
        excel = self._excel_main()
        result = self.run_audit_fixture(excel)
        data = result["data"]
        self.assertEqual(result["kind"], "audit_report")
        self.assertEqual(data["workflow_version"], V1)
        self.assertEqual(data["website_id"], WEBSITE)
        self.assertTrue(str(data["audit_id"]).startswith("aud_"))
        self.assertEqual(data["from_date"], EXCEL_FROM)
        self.assertEqual(data["to_date"], EXCEL_TO)
        summary = data["summary"]
        self.assertEqual(summary["excel_total"], 4)
        self.assertEqual(summary["valid_count"], 4)
        self.assertEqual(summary["missing_count"], 3)
        self.assertEqual(summary["issue_count"], 0)
        self.assertEqual(summary["duplicate_count"], 0)
        # Missing numbers: 3/2026, 5/2026, 6/2026 within observed 1..7.
        self.assertEqual([m["so_cong_chung"] for m in data["missing"]],
                         ["3/2026", "5/2026", "6/2026"])
        for i, row in enumerate(data["missing"], 1):
            self.assertEqual(row["stt"], i)
            self.assertIsNone(row["ngay"])
            self.assertEqual(row["ghi_chu"], "Thieu that")
        self.assertEqual(data["issues"], [])
        self.assertEqual(data["revision"], self.store.revision())
        # Audit binding persisted for queue classification.
        audit = self.store.audit_for(data["audit_id"])
        self.assertIsNotNone(audit)
        self.assertEqual(audit["website_id"], WEBSITE)
        self.assertEqual(Path(audit["file_path"]), excel)
        self.assertEqual(audit["file_sha256"],
                         upload_workspace.sha256_file(excel))
        self.assertEqual(audit["from_date"], EXCEL_FROM)
        self.assertEqual(audit["to_date"], EXCEL_TO)
        job_row = self.store.job_for(self.last_job.job_id)
        self.assertEqual(job_row["audit_id"], data["audit_id"])
        self.assertEqual(job_row["status"], "succeeded")

    def test_audit_excel_v1_issue_rows(self):
        _write_excel(self.excel_path, [
            ("abc", "x"),
            ("5/2025", "10/01/2025"),
            ("2/2026", "06/01/2026"),
            ("2/2026", "07/01/2026"),
        ])
        data = self.run_audit_fixture()["data"]
        summary = data["summary"]
        self.assertEqual(summary["valid_count"], 0)
        self.assertEqual(summary["issue_count"], 4)
        self.assertEqual(summary["duplicate_count"], 2)
        kinds = [row["ghi_chu"].split(":")[0] for row in data["issues"]]
        self.assertIn("sai_format", kinds)
        self.assertIn("sai_nam", kinds)
        self.assertEqual(kinds.count("trung_so"), 2)
        wrong_year = next(r for r in data["issues"]
                          if r["ghi_chu"].startswith("sai_nam"))
        self.assertEqual(wrong_year["so_cong_chung"], "5/2025")
        self.assertEqual(wrong_year["ngay"], "2025-01-10")
        bad = next(r for r in data["issues"]
                   if r["ghi_chu"].startswith("sai_format"))
        self.assertEqual(bad["so_cong_chung"], "abc")

    def test_audit_excel_v1_validation(self):
        # .xls has no verified read path -> validation_error.
        fake_xls = Path(self.tempdir.name) / "so.xls"
        _write_excel(fake_xls, [("1/2026", "05/01/2026")])
        with self.assertRaises(CommandError) as ctx:
            self.run_audit_fixture(fake_xls)
        self.assertEqual(ctx.exception.code, "validation_error")
        # Missing file -> file_not_found.
        with self.assertRaises(CommandError) as ctx:
            self.run_audit_fixture(Path(self.tempdir.name) / "ko-co.xlsx")
        self.assertEqual(ctx.exception.code, "file_not_found")
        # Bad dates -> validation_error (dd/mm/yyyy on the wire is rejected).
        self._excel_main()
        with self.assertRaises(CommandError) as ctx:
            self.run_audit_fixture(from_date="01/01/2026")
        self.assertEqual(ctx.exception.code, "validation_error")
        with self.assertRaises(CommandError) as ctx:
            self.run_audit_fixture(from_date="")
        self.assertEqual(ctx.exception.code, "validation_error")
        with self.assertRaises(CommandError) as ctx:
            self.run_audit_fixture(from_date=EXCEL_FROM,
                                   to_date="2025-01-01")
        self.assertEqual(ctx.exception.code, "validation_error")

    # ------------------------------------------------------- upload.queue_get
    def _scan_and_audit(self, docs, excel_rows):
        for name, so in docs.items():
            _write_docx(self.folder / name, _contract_lines(so))
        _write_excel(self.excel_path, excel_rows)
        scan = self.run_scan_fixture(docs={})   # files already written
        audit = self.run_audit_fixture()
        return scan["data"], audit["data"]

    def test_queue_get_without_excel(self):
        scan = self.run_scan_fixture()["data"]
        data = self.queue_get(scan["run_id"])["data"]
        self.assertEqual(data["workflow_version"], V1)
        self.assertEqual(data["website_id"], WEBSITE)
        self.assertEqual(data["run_id"], scan["run_id"])
        self.assertIsNone(data["audit_id"])
        self.assertEqual(data["queue_revision"], 0)
        self.assertFalse(data["has_excel"])
        self.assertEqual(len(data["folder_rows"]), 2)
        for row in data["folder_rows"]:
            self.assertFalse(row["selected"])
            self.assertEqual(row["ghi_chu"], "chua load Excel")
            self.assertEqual(row["status"], "extracted")
            self.assertEqual(row["missing_fields"], [])
            self.assertFalse(row["has_issue"])
            self.assertEqual(row["ngay"], "2026-04-15")
            self.assertTrue(row["file_path"].endswith(".docx"))
            self.assertRegex(row["normalized_contract_no"], r"^\d+/2026$")
        self.assertEqual(data["missing_in_excel_record_ids"], [])

    def test_queue_get_with_excel_classification(self):
        docs = {
            "A-002.docx": "2/2026/CCGD",
            "B-003.docx": "3/2026/CCGD",
            "C-007.docx": "7/2026/CCGD",
            "D-009.docx": "9/2025/CCGD",
            "E-008.docx": "8/2026/XYZ/CCGD",
        }
        scan, audit = self._scan_and_audit(docs, [
            ("1/2026", "05/01/2026"),
            ("2/2026", "06/01/2026"),
            ("4/2026", "08/01/2026"),
            ("007/2026", "09/01/2026"),
        ])
        data = self.queue_get(scan["run_id"], audit["audit_id"])["data"]
        self.assertTrue(data["has_excel"])
        self.assertEqual(data["audit_id"], audit["audit_id"])
        self.assertEqual(data["queue_revision"], 1)
        rows = self._rows_by_contract(data)

        r2 = rows["2/2026/CCGD"]
        self.assertEqual(r2["normalized_contract_no"], "2/2026")
        self.assertEqual(r2["ghi_chu"], "da co trong Excel")
        self.assertFalse(r2["selected"])
        self.assertFalse(r2["has_issue"])

        r3 = rows["3/2026/CCGD"]
        self.assertEqual(r3["ghi_chu"], "chua co trong Excel")
        self.assertTrue(r3["selected"])

        # Excel "007/2026" vs folder "7/2026" — canonical match.
        r7 = rows["7/2026/CCGD"]
        self.assertEqual(r7["normalized_contract_no"], "7/2026")
        self.assertEqual(r7["ghi_chu"], "da co trong Excel")
        self.assertFalse(r7["selected"])

        r9 = rows["9/2025/CCGD"]
        self.assertIn("sai nam", r9["ghi_chu"])
        self.assertTrue(r9["has_issue"])
        self.assertFalse(r9["selected"])

        r8 = rows["8/2026/XYZ/CCGD"]
        self.assertIn("sai format", r8["ghi_chu"])
        self.assertTrue(r8["has_issue"])
        self.assertFalse(r8["selected"])

        missing_ids = data["missing_in_excel_record_ids"]
        self.assertEqual(missing_ids, [r3["record_id"]])

        # Re-fetch with the same audit does not bump queue_revision again.
        again = self.queue_get(scan["run_id"], audit["audit_id"])["data"]
        self.assertEqual(again["queue_revision"], 1)

    def test_queue_get_local_duplicate_flagged(self):
        docs = {
            "X-020a.docx": "20/2026/CCGD",
            "X-020b.docx": "20/2026/CCGD",
        }
        scan, audit = self._scan_and_audit(docs, [("1/2026", "05/01/2026")])
        data = self.queue_get(scan["run_id"], audit["audit_id"])["data"]
        self.assertEqual(len(data["folder_rows"]), 2)
        ids = set()
        for row in data["folder_rows"]:
            self.assertIn("trung trong folder", row["ghi_chu"])
            self.assertIn("chua co trong Excel", row["ghi_chu"])
            ids.add(row["record_id"])
        self.assertEqual(sorted(ids),
                         data["missing_in_excel_record_ids"])

    def test_queue_get_missing_fields_note(self):
        for name, so in {"F-077.docx": "77/2026/CCGD"}.items():
            _write_docx(self.folder / name,
                        _contract_lines(so, full=False))
        scan = self.run_scan_fixture(docs={})["data"]
        data = self.queue_get(scan["run_id"])["data"]
        row = data["folder_rows"][0]
        self.assertIn("missing: ", row["ghi_chu"])
        self.assertIn("ten_hop_dong", row["missing_fields"])
        self.assertIn("tai_san", row["missing_fields"])

    def test_queue_get_excel_changed_after_audit(self):
        excel = self._excel_main()
        scan = self.run_scan_fixture()["data"]
        audit = self.run_audit_fixture(excel)["data"]
        data = self.queue_get(scan["run_id"], audit["audit_id"])["data"]
        self.assertTrue(data["has_excel"])
        # Rewrite the Excel file -> hash no longer matches the binding.
        _write_excel(excel, [("9/2026", "10/02/2026")])
        with self.assertRaises(CommandError) as ctx:
            self.queue_get(scan["run_id"], audit["audit_id"])
        self.assertEqual(ctx.exception.code, "stale_revision")
        self.assertTrue(ctx.exception.retryable)
        self.assertEqual(ctx.exception.next_action, "retry")
        # Deleting the bound file -> file_not_found.
        excel.unlink()
        with self.assertRaises(CommandError) as ctx:
            self.queue_get(scan["run_id"], audit["audit_id"])
        self.assertEqual(ctx.exception.code, "file_not_found")
        # Audit-free request still works even while the file is gone.
        data = self.queue_get(scan["run_id"])["data"]
        self.assertFalse(data["has_excel"])

    def test_queue_get_registry_read_error(self):
        scan = self.run_scan_fixture()["data"]
        site_dir = upload_workspace.website_data_dir(WEBSITE)
        # Corrupt the engine registry -> structured error, not empty list.
        (site_dir / "registry.sqlite3").write_bytes(b"khong phai sqlite")
        with self.assertRaises(CommandError) as ctx:
            self.queue_get(scan["run_id"])
        self.assertEqual(ctx.exception.code, "engine_unavailable")
        self.assertTrue(ctx.exception.retryable)

    def test_queue_get_scope_and_binding(self):
        scan = self.run_scan_fixture()["data"]
        with mock.patch.object(upload_adapter, "worker") as w:
            # Unknown run_id -> scope_violation.
            with self.assertRaises(CommandError) as ctx:
                self.queue_get("khong-co-run")
            self.assertEqual(ctx.exception.code, "scope_violation")
            self.assertFalse(ctx.exception.retryable)
            # Run of nam_dinh requested under another website.
            with self.assertRaises(CommandError) as ctx:
                self.queue_get(scan["run_id"], website_id=FAKE_WEBSITE)
            self.assertEqual(ctx.exception.code, "website_mismatch")
            # Unknown audit -> scope_violation.
            with self.assertRaises(CommandError) as ctx:
                self.queue_get(scan["run_id"], audit_id="aud_ko-co")
            self.assertEqual(ctx.exception.code, "scope_violation")
            # Audit of a different website -> website_mismatch.
            self.store.add_audit(
                FAKE_WEBSITE, "aud_b1", str(self.excel_path),
                file_sha256="x",
                from_date=EXCEL_FROM, to_date=EXCEL_TO)
            with self.assertRaises(CommandError) as ctx:
                self.queue_get(scan["run_id"], audit_id="aud_b1")
            self.assertEqual(ctx.exception.code, "website_mismatch")
            w.assert_not_called()

    def test_two_scans_have_independent_manifests(self):
        folder_b = Path(self.tempdir.name) / "ho so B"
        folder_b.mkdir()
        scan_a = self.run_scan_fixture(
            docs={"A-138.docx": "138/2026/CCGD"})["data"]
        time.sleep(1.1)  # manifest filename has 1-second resolution
        scan_b = self.run_scan_fixture(
            docs={"B-200.docx": "200/2026/CCGD"}, folder=folder_b)["data"]
        self.assertNotEqual(scan_a["run_id"], scan_b["run_id"])
        man_a = Path(scan_a["manifest_ref"]["path"])
        man_b = Path(scan_b["manifest_ref"]["path"])
        self.assertNotEqual(man_a, man_b)
        self.assertTrue(man_a.is_file() and man_b.is_file())

        rows_a = self.queue_get(scan_a["run_id"])["data"]["folder_rows"]
        self.assertEqual([r["contract_no"] for r in rows_a],
                         ["138/2026/CCGD"])
        rows_b = self.queue_get(scan_b["run_id"])["data"]["folder_rows"]
        self.assertEqual([r["contract_no"] for r in rows_b],
                         ["200/2026/CCGD"])

        # A request for run B whose manifest was replaced by run A's
        # manifest must fail before any browser work — no "latest file"
        # fallback, no silent re-binding.
        man_b.write_text(man_a.read_text(encoding="utf-8"),
                         encoding="utf-8")
        with mock.patch.object(upload_adapter, "worker") as w:
            with self.assertRaises(CommandError) as ctx:
                self.queue_get(scan_b["run_id"])
            self.assertEqual(ctx.exception.code, "manifest_mismatch")
            w.assert_not_called()

        # Deleting run A's manifest fails instead of falling back.
        man_a.unlink()
        with mock.patch.object(upload_adapter, "worker") as w:
            with self.assertRaises(CommandError) as ctx:
                self.queue_get(scan_a["run_id"])
            self.assertEqual(ctx.exception.code, "file_not_found")
            w.assert_not_called()

    # ------------------------------------------------------- dispatch gates
    def test_v1_gate_errors(self):
        # queue_get without workflow_version -> unsupported_workflow_version.
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.queue_get", {
                "website_id": WEBSITE,
                "run_id": "r1",
                "audit_id": None,
            })
        self.assertEqual(ctx.exception.code, "unsupported_workflow_version")
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.queue_get", {
                "workflow_version": "upload.workflow.v0",
                "website_id": WEBSITE,
                "run_id": "r1",
                "audit_id": None,
            })
        self.assertEqual(ctx.exception.code, "unsupported_workflow_version")
        # Unknown key -> validation_error.
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.queue_get", {
                "workflow_version": V1,
                "website_id": WEBSITE,
                "run_id": "r1",
                "audit_id": None,
                "record_ids": [1],
            })
        self.assertEqual(ctx.exception.code, "validation_error")
        # audit_id empty string -> validation_error (null only).
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.queue_get", {
                "workflow_version": V1,
                "website_id": WEBSITE,
                "run_id": "r1",
                "audit_id": "",
            })
        self.assertEqual(ctx.exception.code, "validation_error")
        # Unknown website -> unknown_website.
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.queue_get", {
                "workflow_version": V1,
                "website_id": "khong-co",
                "run_id": "r1",
                "audit_id": None,
            })
        self.assertEqual(ctx.exception.code, "unknown_website")

    def test_legacy_dispatch_unchanged(self):
        # Payloads without workflow_version keep legacy behavior: the legacy
        # handlers validate their own (different) payload shape first.
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.scan", {
                "folder": {"path": str(self.folder / "ko-co"),
                           "scope": "machine_local"},
            })
        self.assertEqual(ctx.exception.code, "file_not_found")
        with self.assertRaises(CommandError) as ctx:
            self._run("upload.audit_excel", {
                "file": {"path": str(self.excel_path / "ko-co.xlsx"),
                         "scope": "machine_local"},
                "from_date": "01/01/2026",
                "to_date": "31/12/2026",
            })
        self.assertEqual(ctx.exception.code, "file_not_found")


if __name__ == "__main__":
    unittest.main()
