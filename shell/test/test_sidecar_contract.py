"""Contract conformance test cho sidecar desktopcommand.v1.

Chay: python test/test_sidecar_contract.py   (can fastapi/uvicorn/httpx/docx)
Spawn app.py that tren port dong voi SIDECAR_TOKEN/PORT nhu Electron main lam.
"""
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import unittest
import uuid
from pathlib import Path

SIDECAR_DIR = Path(__file__).resolve().parent.parent / "sidecar"
TOKEN = "test-token-" + uuid.uuid4().hex


def _free_port():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def _new_cmd(command, payload=None, command_id=None):
    return {
        "contract_version": "desktopcommand.v1",
        "command_id": command_id or str(uuid.uuid4()),
        "command": command,
        "payload": payload,
        "client_meta": {"shell_version": "0.1.0", "module": "test"},
    }


class SidecarContractTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.port = _free_port()
        cls.base = f"http://127.0.0.1:{cls.port}"
        env = dict(os.environ, SIDECAR_PORT=str(cls.port),
                   SIDECAR_TOKEN=TOKEN,
                   # MIN-106: dev flag -> gateway chon mock cho notary.* v1
                   G1_DEV_NOTARY_MOCK="1")
        cls.proc = subprocess.Popen(
            [sys.executable, str(SIDECAR_DIR / "app.py")],
            cwd=str(SIDECAR_DIR), env=env,
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        import httpx
        cls.http = httpx.Client(base_url=cls.base,
                                headers={"authorization": f"Bearer {TOKEN}"},
                                timeout=10)
        deadline = time.time() + 30
        while time.time() < deadline:
            try:
                if cls.http.get("/healthz").status_code == 200:
                    break
            except Exception:
                pass
            if cls.proc.poll() is not None:
                raise RuntimeError("sidecar exit som")
            time.sleep(0.2)
        else:
            raise RuntimeError("sidecar khong healthy trong 30s")

    @classmethod
    def tearDownClass(cls):
        try:
            cls.http.post("/shutdown")
        except Exception:
            pass
        try:
            cls.proc.wait(timeout=8)
        except Exception:
            cls.proc.kill()

    def _wait_job(self, job_id, timeout=15):
        deadline = time.time() + timeout
        while time.time() < deadline:
            job = self.http.get(f"/v1/jobs/{job_id}").json()
            if job["status"] in ("succeeded", "failed", "canceled", "partial"):
                return job
            time.sleep(0.15)
        self.fail(f"job {job_id} khong terminal trong {timeout}s")

    def test_healthz_no_auth(self):
        import httpx
        r = httpx.get(f"{self.base}/healthz", timeout=5)
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertIn("desktopcommand.v1", body["supported_versions"])
        self.assertTrue(body["engine_instance_id"])

    def test_auth_required(self):
        import httpx
        r = httpx.post(f"{self.base}/v1/commands", json=_new_cmd("x.y"),
                       timeout=5)
        self.assertEqual(r.status_code, 401)
        self.assertEqual(r.json()["error"]["code"], "auth_unauthorized")

    def test_origin_rejected(self):
        r = self.http.post("/v1/commands", json=_new_cmd("x.y"),
                           headers={"origin": "http://evil.local"})
        self.assertEqual(r.status_code, 403)

    def test_bad_contract_version(self):
        cmd = _new_cmd("file.inspect")
        del cmd["contract_version"]
        r = self.http.post("/v1/commands", json=cmd)
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"],
                         "unsupported_contract_version")

    def test_sensitive_key_rejected_recursive(self):
        r = self.http.post("/v1/commands", json=_new_cmd(
            "file.inspect",
            {"file": {"path": "D:/x", "scope": "machine_local"},
             "opts": {"nd_password": "pw"}}))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"],
                         "payload_rejected_sensitive_key")

    def test_unknown_command(self):
        r = self.http.post("/v1/commands", json=_new_cmd("hack.run"))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"], "command_unknown")

    def test_file_inspect_real_docx(self):
        import docx
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "mau.docx"
            d = docx.Document()
            d.add_paragraph("Hợp đồng chuyển nhượng thửa 12 tờ 34")
            d.save(str(p))
            job = self.http.post("/v1/commands", json=_new_cmd(
                "file.inspect",
                {"file": {"path": str(p), "scope": "machine_local"}})).json()
            final = self._wait_job(job["job_id"])
            self.assertEqual(final["status"], "succeeded")
            data = final["result"]["data"]
            self.assertIn("chuyển nhượng", data["text_preview"])
            self.assertEqual(len(data["sha256"]), 64)
            self.assertEqual(final["result"]["kind"], "file_inspect")

    def test_unc_path_rejected(self):
        # contract §6: UNC → 400 file_scope_not_supported ngay luc submit
        r = self.http.post("/v1/commands", json=_new_cmd(
            "file.inspect",
            {"file": {"path": "\\\\maychu\\share\\a.docx",
                      "scope": "machine_local"}}))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"], "file_scope_not_supported")

    def test_lan_scope_rejected(self):
        r = self.http.post("/v1/commands", json=_new_cmd(
            "file.inspect",
            {"file": {"path": "D:/hs/a.docx", "scope": "lan_share"}}))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"], "file_scope_not_supported")

    def test_unc_extended_prefix_rejected(self):
        # \\?\UNC\server\share la dang extended cua UNC — van phai bi chan
        r = self.http.post("/v1/commands", json=_new_cmd(
            "file.inspect",
            {"file": {"path": "\\\\?\\UNC\\maychu\\share\\a.docx",
                      "scope": "machine_local"}}))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"], "file_scope_not_supported")

    def test_long_path_prefix_allowed_shape(self):
        # \\?\D:\... hop le ve scope (ton tai file moi that bai → file_not_found)
        r = self.http.post("/v1/commands", json=_new_cmd(
            "file.inspect",
            {"file": {"path": "\\\\?\\D:\\khong-ton-tai-xyz.docx",
                      "scope": "machine_local"}}))
        self.assertEqual(r.status_code, 200)
        final = self._wait_job(r.json()["job_id"])
        self.assertEqual(final["error"]["code"], "file_not_found")

    def test_bad_client_meta(self):
        cmd = _new_cmd("file.inspect")
        cmd["client_meta"] = {}
        r = self.http.post("/v1/commands", json=cmd)
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"], "validation_error")

    def test_bad_command_id(self):
        cmd = _new_cmd("file.inspect", command_id="khong-phai-uuid")
        r = self.http.post("/v1/commands", json=cmd)
        self.assertEqual(r.status_code, 400)

    def test_missing_file(self):
        job = self.http.post("/v1/commands", json=_new_cmd(
            "file.inspect",
            {"file": {"path": "D:/khong-ton-tai-xyz.docx",
                      "scope": "machine_local"}})).json()
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "failed")
        self.assertEqual(final["error"]["code"], "file_not_found")

    def test_idempotent_command_id(self):
        cid = str(uuid.uuid4())
        cmd = _new_cmd("diag.slow_task", {"steps": 2}, command_id=cid)
        j1 = self.http.post("/v1/commands", json=cmd).json()
        j2 = self.http.post("/v1/commands", json=cmd).json()
        self.assertEqual(j1["job_id"], j2["job_id"])
        self._wait_job(j1["job_id"])

    def test_cancel_running_and_terminal_409(self):
        job = self.http.post("/v1/commands", json=_new_cmd(
            "diag.slow_task", {"steps": 60})).json()
        r = self.http.post(f"/v1/jobs/{job['job_id']}/cancel")
        self.assertEqual(r.status_code, 200)
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "canceled")
        self.assertEqual(final["error"]["code"], "user_canceled")
        r2 = self.http.post(f"/v1/jobs/{job['job_id']}/cancel")
        self.assertEqual(r2.status_code, 409)
        self.assertEqual(r2.json()["error"]["code"], "job_already_terminal")

    def test_job_not_found(self):
        r = self.http.get("/v1/jobs/j_khong-co")
        self.assertEqual(r.status_code, 404)
        self.assertEqual(r.json()["error"]["code"], "job_not_found")

    def test_env_check_tra_checklist(self):
        job = self.http.post("/v1/commands", json=_new_cmd(
            "diag.env_check")).json()
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "succeeded")
        self.assertEqual(final["result"]["kind"], "env_check")
        data = final["result"]["data"]
        names = {c["name"] for c in data["checks"]}
        self.assertTrue({"python", "fastapi", "uvicorn"} <= names)
        self.assertTrue(data["required_ok"])
        for c in data["checks"]:
            self.assertIn("ok", c)
            self.assertIn("detail", c)

    def test_waiting_task_waiting_roi_resume(self):
        # waiting_user quan sat duoc tu ben ngoai, roi job tu resume
        job = self.http.post("/v1/commands", json=_new_cmd(
            "diag.waiting_task", {"wait_seconds": 2})).json()
        deadline = time.time() + 10
        seen_waiting = False
        while time.time() < deadline:
            cur = self.http.get(f"/v1/jobs/{job['job_id']}").json()
            if cur["status"] == "waiting_user":
                seen_waiting = True
                self.assertEqual(cur["waiting_on"], "review")
                break
            time.sleep(0.1)
        self.assertTrue(seen_waiting, "khong thay waiting_user")
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "succeeded")
        self.assertIsNone(final["waiting_on"])

    def test_waiting_task_cancel_ngay_khi_waiting(self):
        job = self.http.post("/v1/commands", json=_new_cmd(
            "diag.waiting_task", {"wait_seconds": 60})).json()
        deadline = time.time() + 10
        while time.time() < deadline:
            cur = self.http.get(f"/v1/jobs/{job['job_id']}").json()
            if cur["status"] == "waiting_user":
                break
            time.sleep(0.1)
        r = self.http.post(f"/v1/jobs/{job['job_id']}/cancel")
        self.assertEqual(r.status_code, 200)
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "canceled")
        self.assertEqual(final["error"]["code"], "user_canceled")

    # ----- notary.* case-drafting v1 qua mock backend (G1_DEV_NOTARY_MOCK=1)

    def test_notary_workspace_get_mock(self):
        job = self.http.post("/v1/commands", json=_new_cmd(
            "notary.workspace_get", {"case_id": 42})).json()
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "succeeded")
        res = final["result"]
        self.assertEqual(res["kind"], "workspace_get")
        data = res["data"]
        self.assertEqual(data["schema_version"], "notary.case-drafting.v1")
        self.assertEqual(data["backend_mode"], "mock")
        self.assertEqual(data["case"]["id"], 42)
        self.assertIn("capabilities", data)

    def test_notary_case_not_found(self):
        job = self.http.post("/v1/commands", json=_new_cmd(
            "notary.workspace_get", {"case_id": 9999})).json()
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "failed")
        self.assertEqual(final["error"]["code"], "case_not_found")

    def test_notary_intake_partial_breakdown(self):
        sources = [
            {"source_id": str(uuid.uuid4()), "kind": "text",
             "text": "Người Mẫu I, sinh 1950"},
            {"source_id": str(uuid.uuid4()), "kind": "text",
             "text": "mock_fail — khong trich duoc (mo phong)"},
        ]
        job = self.http.post("/v1/commands", json=_new_cmd(
            "notary.intake_analyze",
            {"case_id": 43, "sources": sources})).json()
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "partial")
        # marker "partial" la noi bo jobstore — khong len wire
        self.assertNotIn("partial", final["result"])
        bd = final["result"]["data"]["breakdown"]
        self.assertEqual(len(bd["succeeded"]), 1)
        self.assertEqual(bd["failed"], [sources[1]["source_id"]])

    def test_notary_locked_write_rejected(self):
        job = self.http.post("/v1/commands", json=_new_cmd(
            "notary.workspace_commit_stage",
            {"case_id": 44, "base_revision": 3,
             "stage": {"people": [], "assets": []}})).json()
        final = self._wait_job(job["job_id"])
        self.assertEqual(final["status"], "failed")
        self.assertEqual(final["error"]["code"], "workspace_locked")

    def test_notary_word_export_idempotent_no_duplicate(self):
        import docx
        with tempfile.TemporaryDirectory() as td:
            cid = str(uuid.uuid4())
            cmd = _new_cmd("notary.word_export_batch", {
                "case_id": 42, "document_keys": ["khai_nhan_di_san"],
                "destination": {"path": td, "scope": "machine_local",
                                "is_dir": True}}, command_id=cid)
            j1 = self.http.post("/v1/commands", json=cmd).json()
            final = self._wait_job(j1["job_id"])
            self.assertEqual(final["status"], "succeeded")
            doc = final["result"]["data"]["documents"][0]
            self.assertEqual(doc["status"], "saved")
            out = Path(doc["output_file"]["path"])
            self.assertTrue(out.is_file())
            docx.Document(str(out))        # DOCX hop le
            # retry cung command_id -> cung job -> khong tao file moi
            j2 = self.http.post("/v1/commands", json=cmd).json()
            self.assertEqual(j2["job_id"], j1["job_id"])
            self.assertEqual(len(list(Path(td).glob("*.docx"))), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
