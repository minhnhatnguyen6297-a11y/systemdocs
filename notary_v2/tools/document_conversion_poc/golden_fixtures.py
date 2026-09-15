"""Create deterministic synthetic files for the GD-01..07 manifest."""

from __future__ import annotations

from datetime import datetime, timezone
import io
from pathlib import Path
import re
import zipfile

import fitz
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage

# PNG bytes written as literals so fixture hashes do not depend on the
# installed Pillow version (encoder output differs across releases).
_WHITE_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000002000000020802000000fdd49a"
    "730000001649444154789c63fcffff3f030303130303030303030024060301fc"
    "35de9b0000000049454e44ae426082"
)
_BLACK_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000002000000020802000000fdd49a"
    "730000000b49444154789c6360400600000e0001a99173b10000000049454e44"
    "ae426082"
)


def _pdf(path: Path, *, text: str | None) -> None:
    document = fitz.open()
    document.set_metadata({"format": "PDF 1.7", "title": "Synthetic", "author": "POC", "creationDate": "D:20200101000000Z"})
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    data = document.tobytes()
    data = re.sub(
        rb"/ID\s*\[<[^>]+><[^>]+>\]",
        b"/ID [<00000000000000000000000000000000><11111111111111111111111111111111>]",
        data,
    )
    path.write_bytes(data)
    document.close()


def _normalize_office_zip(path: Path) -> None:
    """Normalize ZIP metadata so identical synthetic inputs have identical bytes."""
    source = zipfile.ZipFile(path)
    try:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                normalized = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                normalized.compress_type = zipfile.ZIP_DEFLATED
                normalized.external_attr = info.external_attr
                payload = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    payload = re.sub(
                        rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                        rb"\g<1>2020-01-01T00:00:00Z\g<2>",
                        payload,
                    )
                target.writestr(normalized, payload)
        path.write_bytes(buffer.getvalue())
    finally:
        source.close()


def materialize_golden_fixtures(directory: Path) -> list[Path]:
    """Write canonical synthetic GD-01..07 fixtures in manifest order."""
    directory.mkdir(parents=True, exist_ok=True)
    _pdf(directory / "gd-01-text.pdf", text="Synthetic PDF text")
    _pdf(directory / "gd-02-scanned.pdf", text=None)

    document = Document()
    document.add_paragraph("Synthetic DOCX text")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Party"
    table.cell(1, 1).text = "Synthetic A"
    image_buffer = io.BytesIO(_WHITE_PNG)
    document.add_picture(image_buffer, width=Inches(0.1))
    fixed = datetime(2020, 1, 1, tzinfo=timezone.utc)
    document.core_properties.created = fixed
    document.core_properties.modified = fixed
    document.core_properties.last_printed = fixed
    document.save(directory / "gd-03-contract.docx")
    _normalize_office_zip(directory / "gd-03-contract.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic"
    sheet.append(["Field", "Value"])
    sheet.append(["Party", "Synthetic A"])
    second_sheet = workbook.create_sheet("Second")
    second_sheet.append(["Marker", "Synthetic B"])
    image_buffer = io.BytesIO(_BLACK_PNG)
    sheet.add_image(ExcelImage(image_buffer), "D2")
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    workbook.save(directory / "gd-04-sheet.xlsx")
    _normalize_office_zip(directory / "gd-04-sheet.xlsx")

    (directory / "gd-05-legacy.doc").write_bytes(b"\xd0\xcf\x11\xe0" + b"synthetic")
    (directory / "gd-06-id.png").write_bytes(_WHITE_PNG)
    (directory / "gd-07-unsupported.bin").write_bytes(b"synthetic unsupported")
    return [directory / f"gd-0{index}-{name}" for index, name in (
        (1, "text.pdf"), (2, "scanned.pdf"), (3, "contract.docx"),
        (4, "sheet.xlsx"), (5, "legacy.doc"), (6, "id.png"),
        (7, "unsupported.bin"),
    )]
