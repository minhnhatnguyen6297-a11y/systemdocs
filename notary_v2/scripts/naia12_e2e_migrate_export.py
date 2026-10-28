"""E2E NAIA-12: migrate bộ mẫu thừa kế rồi export qua ĐÚNG luồng batch.

Chạy:

    python -m scripts.naia12_e2e_migrate_export SRC_DIR OUT_DIR

hoặc từ ``notary_v2``:

    python scripts/naia12_e2e_migrate_export.py SRC_DIR OUT_DIR

SRC_DIR    thư mục mẫu nguồn (.doc/.docx) — KHÔNG bị sửa, mọi đầu ra ghi vào
           OUT_DIR. Script từ chối khi nguồn/đích trùng hoặc chồng lấn.
OUT_DIR    nhận ``migrated/`` (mẫu đã migrate + ``migration_manifest.json``),
           ``exports/`` (DOCX xuất thử qua batch), ``exports_neg/`` (phép thử
           phản ví dụ thiếu người được ủy quyền), ``e2e_report.json``.

Luồng export đi qua API THẬT giống UI gọi: ``scan_template_dir`` quét thư mục
mẫu đã migrate -> ``template_path_resolver`` -> ``export_options`` ->
``export_batch`` -> kiểm ``status``/``actual_filename`` từng văn bản. Phần
"engine dry-run" trong report chỉ là chẩn đoán phụ — KHÔNG thay thế batch.

Exit code: 0 khi mọi kiểm chứng đạt; khác 0 khi migrate lỗi, batch lỗi, nội
dung mong đợi thiếu, nguồn bị đổi, hoặc còn placeholder sót.

Bằng chứng sinh ra (tách theo tầng):
  * ``migration``  — manifest, sha256 nguồn trước/sau, excluded files;
  * ``engine``     — dry-run replace/unresolved theo từng mẫu (diagnostic);
  * ``batch``      — export_options + từng document result + assertion nội
                     dung (người được ủy quyền, nơi niêm yết, diện tích từng
                     dòng đất) trích từ DOCX đầu ra;
  * ``negative``   — cùng batch trên hồ sơ THIẾU vai trò ủy quyền phải fail
                     bằng ``word.missing_authorized_person``;
  * ``format_compare`` — so khung cấu trúc/định dạng nguồn->đích (docx->docx:
                     section, bảng, đoạn, header/footer, style, font) và
                     ComputeStatistics qua Word COM khi có Word;
  * ``ui``         — trạng thái kiểm chứng Electron UI (chưa chạy -> báo rõ).

Giới hạn đã biết được ghi trong ``report['limitations']``: export tạo file
bằng python-docx; Word COM chỉ đo được thống kê trang/dòng nếu máy có Word —
KHÔNG phải so sánh trực quan định dạng; UI Electron chưa được chạy.
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path
from types import SimpleNamespace

# Cho phép chạy trực tiếp "python scripts/naia12_e2e_migrate_export.py" —
# thêm notary_v2 vào sys.path để import services.*.
_HERE = Path(__file__).resolve().parent
if str(_HERE.parent) not in sys.path:
    sys.path.insert(0, str(_HERE.parent))

from services import word_batch_export as wbe  # noqa: E402
from services import word_engine  # noqa: E402
from services import word_template_migration as wtm  # noqa: E402

PROXY_NAME = "Trần Thị Ủy"
NIEM_YET = "UBND xã Minh Tân (giá trị synthetic)"


def _person(cid, name, *, gender="Nam", born=None, dead=None, doc_no="",
            address=""):
    return SimpleNamespace(
        id=cid, ho_ten=name, gioi_tinh=gender, ngay_sinh=born,
        ngay_chet=dead, so_giay_to=doc_no, ngay_cap=None,
        loai_giay_to="Căn cước công dân",
        noi_cap="Cục cảnh sát quản lý hành chính về trận tự xã hội",
        dia_chi=address, loai_dia_chi="Thường trú tại",
    )


def _participant(person, role, *, receive=True, share=0):
    return SimpleNamespace(
        customer=person, customer_id=person.id, vai_tro=role,
        co_nhan_tai_san=receive, ty_le=share, hang_thua_ke=1,
        parent_customer_id=None,
    )


def _synthetic_case(*, with_proxy: bool = True):
    """Hồ sơ tổng hợp đủ owner/người nhận/3 tài sản có land_rows.

    Người được ủy quyền (``proxy``) CHỈ có vai trò lưu ở participants, không
    có node Diagram — đúng tình huống luồng nhập lưu vai trò đại diện đã xác
    nhận (đại diện không nằm trên sơ đồ thừa kế). ``with_proxy=False`` dựng
    phản ví dụ hồ sơ thiếu người được ủy quyền.
    """
    owner = _person(1, "Nguyễn Văn Chủ", born=date(1950, 1, 1),
                    dead=date(2024, 2, 3), doc_no="001050000001")
    receiver = _person(2, "Nguyễn Văn Con", born=date(1980, 8, 9),
                       doc_no="001080000003")
    proxy = _person(3, PROXY_NAME, gender="Nữ", born=date(1975, 5, 5),
                    doc_no="001075000005")

    def prop(serial, rows):
        return SimpleNamespace(
            so_serial=serial, so_vao_so="CS123", so_thua_dat="45",
            so_to_ban_do="12", dia_chi="Thửa đất tại xã Minh Tân",
            loai_dat="Đất ở tại nông thôn", dien_tich=120.5,
            loai_so="Giấy chứng nhận quyền sử dụng đất",
            land_rows_json=json.dumps(rows, ensure_ascii=False),
            hinh_thuc_su_dung="Sử dụng riêng", thoi_han="Lâu dài",
            nguon_goc="Nhà nước công nhận quyền sử dụng đất",
            ngay_cap=date(2020, 7, 8), co_quan_cap="UBND huyện Ý Yên",
        )

    properties = [
        prop("BM 1451111", [{"loai_dat": "ONT", "dien_tich": "80",
                             "thoi_han": "Lâu dài"},
                            {"loai_dat": "CLN", "dien_tich": "40.5"}]),
        prop("BM 1452222", [{"loai_dat": "NTS", "dien_tich": "55"}]),
        # Tài sản 3 — mẫu "DK18, thuế 3" cần token [ONT 3]/[CLN 3]...;
        # thiếu nó thì export đúng cách báo placeholder sót.
        prop("BM 1453333", [{"loai_dat": "LUC", "dien_tich": "66"},
                            {"loai_dat": "CLN", "dien_tich": "12"},
                            {"loai_dat": "ONT", "dien_tich": "7"},
                            {"loai_dat": "NTS", "dien_tich": "3"}]),
    ]
    people = (owner, receiver, proxy) if with_proxy else (owner, receiver)
    stage = [
        {"id": str(p.id), "ho_ten": p.ho_ten, "gioi_tinh": p.gioi_tinh,
         "ngay_sinh": p.ngay_sinh.isoformat() if p.ngay_sinh else None,
         "ngay_chet": p.ngay_chet.isoformat() if p.ngay_chet else None,
         "so_giay_to": p.so_giay_to, "ngay_cap": None, "noi_cap": p.noi_cap,
         "dia_chi": p.dia_chi, "loai_dia_chi": p.loai_dia_chi}
        for p in people
    ]
    case_state = json.dumps({
        "schemaVersion": 1, "stage": stage,
        "diagram": {"engineState": {"nodes": [
            {"id": "n1", "personId": "1", "role": "Owner",
             "isLandOwner": True, "hidden": False, "deleted": False},
            {"id": "n2", "personId": "2", "role": "Con",
             "willReceive": True, "isLandOwner": False,
             "hidden": False, "deleted": False},
        ]}},
    }, ensure_ascii=False)
    participants = [_participant(receiver, "Con")]
    if with_proxy:
        participants.append(
            _participant(proxy, "Người được ủy quyền", receive=False))
    return SimpleNamespace(
        id=99, nguoi_chet=owner, tai_san=properties[0],
        property_links=[
            SimpleNamespace(property=p, is_primary=(p is properties[0]), id=i)
            for i, p in enumerate(properties)],
        ngay_lap_ho_so=date(2026, 7, 3), loai_van_ban="thoa_thuan",
        noi_niem_yet=NIEM_YET,
        ghi_chu="Hồ sơ kiểm chứng NAIA-12",
        participants=participants,
        case_state_json=case_state,
    )


# ------------------------------------------------------------- docx helpers


def _qn():
    from docx.oxml.ns import qn
    return qn


def _doc_paragraph_texts(doc) -> list[str]:
    """Text từng đoạn trên MỌI part (body, bảng lồng, header/footer,
    textbox, footnotes/endnotes), gồm xuống dòng và tab Word thật.

    Đọc trực tiếp XML để oracle không phụ thuộc cách engine thay token;
    nội dung paragraph lồng thuộc riêng paragraph đó, không ghép vào cha."""
    qn = _qn()
    text_tag, paragraph_tag = qn("w:t"), qn("w:p")
    controls = {qn("w:br"): "\n", qn("w:cr"): "\n", qn("w:tab"): "\t"}
    texts: list[str] = []
    for part_root in word_engine._iter_word_part_roots(doc):
        for paragraph in word_engine._iter_part_paragraphs(
                part_root.element, qn):
            values = []
            for node in paragraph.iter():
                if node.tag != text_tag and node.tag not in controls:
                    continue
                owner = next((parent for parent in node.iterancestors()
                              if parent.tag == paragraph_tag), None)
                if owner is paragraph:
                    values.append(node.text or "" if node.tag == text_tag
                                  else controls[node.tag])
            texts.append("".join(values))
    return texts


def _doc_text(path: Path) -> str:
    import docx
    return "\n".join(_doc_paragraph_texts(docx.Document(str(path))))


def _norm_ws(text: str) -> str:
    """Collapse mọi khoảng trắng (kể cả \n trong block nhiều dòng)."""
    return " ".join(text.split())


# ------------------------------------------------------ expected assertions

# Nhóm token "đáng kiểm chứng" theo Leader: người được ủy quyền, nơi niêm yết,
# diện tích TỪNG dòng đất ([ONT]/[CLN]/[NTS]/[LUC]( N)? — token vị trí),
# loại đất từng dòng/tổng, diện tích tổng, và dạng canonical
# "[Tài sản N - <trường>]" (review NAIA-12: canonical trước đây bị bỏ qua
# nên giá trị sai vẫn qua kiểm nội dung).
_LAND_ROW_RE = re.compile(r"^(ont|cln|nts|luc)(\s+\d+)?$")
_INTERESTING_TOKEN_RE = re.compile(
    r"nguoi duoc uy quyen|niem yet|noi niem yet|^dien tich|^loai dat"
    r"|muc dich su dung|^tai san( \d+)? - ")


def _is_interesting(token: str) -> bool:
    normalized = word_engine._normalize_token(token[1:-1])
    return bool(
        _INTERESTING_TOKEN_RE.search(normalized)
        or _LAND_ROW_RE.match(normalized))


def _fixture_case_data(case) -> dict:
    """Suy expected TRỰC TIẾP từ fixture — không đi qua mapping engine.

    Nếu expected và output cùng dùng một mapping sai thì phép so content vẫn
    pass (review NAIA-12): mọi giá trị ở đây đọc từ dữ liệu hồ sơ giả lập.
    """
    props = [link.property for link in case.property_links]
    rows_per_asset = [
        json.loads(prop.land_rows_json or "[]") or [{}]
        for prop in props]

    def _rows_of(asset: int) -> list:
        return (rows_per_asset[asset - 1]
                if 0 < asset <= len(rows_per_asset) else [])

    def _row(asset: int, position: int) -> dict:
        rows = _rows_of(asset)
        return rows[position - 1] if position <= len(rows) else {}

    def _area_total(asset: int) -> str:
        if not 0 < asset <= len(props):
            return ""
        total = 0.0
        for row in _rows_of(asset):
            try:
                total += float(row.get("dien_tich") or 0)
            except (TypeError, ValueError):
                pass
        return f"{total:g}" if total else _safe_num(
            props[asset - 1].dien_tich)

    def _purpose(asset: int) -> str:
        return "; ".join(
            f"{row.get('loai_dat', '')}: {row.get('dien_tich', '')} m2"
            for row in _rows_of(asset)
            if row.get("loai_dat") or row.get("dien_tich"))

    def _area_total_num(asset: int) -> float:
        """Tổng diện tích SỐ của tài sản — cho oracle `[Diện tích chữ]`."""
        total = 0.0
        for row in _rows_of(asset):
            try:
                total += float(row.get("dien_tich") or 0)
            except (TypeError, ValueError):
                pass
        if total or asset > len(props):
            return total
        value = getattr(props[asset - 1], "dien_tich", None)
        return float(value) if isinstance(value, (int, float)) else 0.0

    def _safe_num(value) -> str:
        return f"{value:g}" if isinstance(value, (int, float)) else str(value or "")

    # Đại diện theo vai trò ĐÃ LƯU trên participants (độc lập engine): giữ
    # (person-view, vai_tro, tỷ_lệ) để oracle dựng block/field giống hợp đồng
    # _person_detail_block/_person_placeholder_values mà không đi qua
    # build_template_mapping. Dữ liệu người lấy từ STAGE ĐÃ LƯU
    # (case_state_json['stage']) — đúng nguồn engine merge vai trò; object
    # participant.customer có thể giàu trường hơn stage (vd loai_giay_to)
    # và sẽ làm oracle lệch thực tế xuất.
    def _stage_by_id() -> dict[str, dict]:
        try:
            state = json.loads(getattr(case, "case_state_json", "") or "{}")
        except Exception:  # noqa: BLE001 — fixture lỗi thì oracle rỗng
            return {}
        out = {}
        for person in state.get("stage") or []:
            key = person.get("entity_id", person.get("id"))
            if key is not None:
                out[str(key)] = person
        return out

    stage_people = _stage_by_id()

    def _person_view(source) -> dict:
        get = (source.get if isinstance(source, dict)
               else lambda key, default=None: getattr(source, key, default))
        return {
            "ho_ten": get("ho_ten") or "",
            "gioi_tinh": get("gioi_tinh") or "",
            "ngay_sinh": get("ngay_sinh"),
            "ngay_chet": get("ngay_chet"),
            "so_giay_to": get("so_giay_to") or "",
            "loai_giay_to": get("loai_giay_to") or "",
            "ngay_cap": get("ngay_cap"),
            "noi_cap": get("noi_cap") or "",
            "dia_chi": get("dia_chi") or "",
            "loai_dia_chi": get(
                "loai_dia_chi", get("place_of_origin")) or "",
        }

    authorized_people = []
    for participant in case.participants:
        role = str(getattr(participant, "vai_tro", "") or "")
        if not word_engine._normalize_token(role).startswith(
                "nguoi duoc uy quyen"):
            continue
        customer = participant.customer
        stage_row = stage_people.get(
            str(getattr(participant, "customer_id",
                        getattr(customer, "id", ""))))
        authorized_people.append((
            _person_view(stage_row if stage_row is not None else customer),
            role,
            getattr(participant, "ty_le", "")))

    def _title(person: dict) -> str:
        gender = word_engine._normalize_token(person["gioi_tinh"])
        if gender == "nam":
            return "ông"
        if gender in {"nu", "nữ"}:
            return "bà"
        return "ông/bà"

    def _relation(role: str) -> str:
        return word_engine.ROLE_LABELS.get(
            word_engine._role_key(role), role.strip())

    def _block(index: int) -> str:
        """Block `[Dòng người được ủy quyền N]` theo contract engine nhưng
        dữ liệu lấy thẳng từ fixture/stage."""
        if index > len(authorized_people):
            return ""
        person, role, _share = authorized_people[index - 1]
        first = f"{index}. {_title(person)} {person['ho_ten']}".strip()
        if person["ngay_sinh"]:
            first += ("; Sinh ngày: "
                      + word_engine._fmt_birth_or_year(person["ngay_sinh"]))
        first += "."
        lines = [first]
        parts = []
        if person["loai_giay_to"]:
            parts.append(person["loai_giay_to"])
        if person["so_giay_to"]:
            parts.append(f"số: {person['so_giay_to']}")
        doc_line = " ".join(parts)
        if person["noi_cap"]:
            doc_line += f" do {person['noi_cap']} cấp"
        if person["ngay_cap"]:
            doc_line += f" ngày {word_engine._fmt_date(person['ngay_cap'])}"
        if doc_line:
            lines.append(doc_line + ".")
        if person["dia_chi"]:
            lines.append(
                f"{person['loai_dia_chi'] or 'Địa chỉ'}: "
                f"{person['dia_chi']}.")
        relation = _relation(role)
        if relation:
            lines.append(f"Là {relation}.")
        return "\n".join(lines)

    def _inline(index: int) -> str:
        if index > len(authorized_people):
            return ""
        person, _role, _share = authorized_people[index - 1]
        name = f"{_title(person)} {person['ho_ten']}".strip()
        if index == 1:
            return name
        if index == len(authorized_people):
            return f" và {name}"
        return f", {name}"

    def _name_list(titled: bool) -> str:
        parts = []
        for person, _role, _share in authorized_people:
            if not person["ho_ten"]:
                continue
            parts.append(
                f"{_title(person)} {person['ho_ten']}".strip()
                if titled else person["ho_ten"])
        if not parts:
            return ""
        if len(parts) == 1:
            return parts[0]
        if len(parts) == 2:
            return f"{parts[0]} và {parts[1]}"
        return ", ".join(parts[:-1]) + f" và {parts[-1]}"

    # Trường lẻ `[<Trường> người được ủy quyền N]` / `[Người được ủy quyền N -
    # <Trường>]` — cùng tập trường _person_placeholder_values của engine.
    _PERSON_FIELDS = {
        "xung ho": _title,
        "ho ten": lambda person: person["ho_ten"],
        "gioi tinh": lambda person: person["gioi_tinh"],
        "ngay sinh": lambda person: word_engine._fmt_date(person["ngay_sinh"]),
        "nam sinh": lambda person: word_engine._fmt_year(person["ngay_sinh"]),
        "ngay chet": lambda person: word_engine._fmt_date(person["ngay_chet"]),
        "nam chet": lambda person: word_engine._fmt_year(person["ngay_chet"]),
        "loai giay to": lambda person: person["loai_giay_to"],
        "so giay to": lambda person: person["so_giay_to"],
        "ngay cap": lambda person: word_engine._fmt_date(person["ngay_cap"]),
        "noi cap": lambda person: person["noi_cap"],
        "nhan dia chi": lambda person: person["loai_dia_chi"] or "Địa chỉ",
        "dia chi": lambda person: person["dia_chi"],
    }

    def _field(index: int, field_norm: str) -> str:
        if index > len(authorized_people):
            return ""
        person, role, share = authorized_people[index - 1]
        if field_norm == "quan he":
            return _relation(role)
        if field_norm == "vai tro":
            return role
        if field_norm == "trang thai":
            # Fixture: đại diện không nằm trong receivers/landowners.
            return "Người không nhận"
        if field_norm == "ty le":
            if share in (None, ""):
                return ""
            return f"{share:g}" if isinstance(share, (int, float)) else str(share)
        getter = _PERSON_FIELDS.get(field_norm)
        return getter(person) if getter else ""

    # Slot người `[Tên N]`/`[CCCD N]`/... — oracle tự chọn theo đúng quy
    # tắc _pick_core_people (slot 1/2 = owner+vợ/chồng, slot 3 = người
    # nhận tỷ lệ cao nhất, slot 4+ = còn lại) để token neighbor được giải
    # từ fixture, không phải wildcard nuốt giá trị sai (review NAIA-12:
    # ` SAI` nối vào tên UQ trước `[Tên 1]` từng pass).
    def _person_slots() -> list:
        owner = getattr(case, "nguoi_chet", None)
        participants = getattr(case, "participants", []) or []
        spouse = None
        for participant in participants:
            if word_engine._role_key(
                    getattr(participant, "vai_tro", "")) == "vo_chong":
                spouse = getattr(participant, "customer", None)
                break
        pair = [c for c in (owner, spouse) if c is not None]
        nam = [c for c in pair if word_engine._normalize_token(
            getattr(c, "gioi_tinh", "")) == "nam"]
        nu = [c for c in pair if word_engine._normalize_token(
            getattr(c, "gioi_tinh", "")) in {"nu", "nữ"}]
        if len(nam) == 1 and len(nu) == 1:
            person1, person2 = nam[0], nu[0]
        elif len(pair) == 2:
            person1, person2 = pair
        elif len(pair) == 1:
            person1, person2 = pair[0], None
        else:
            person1 = person2 = None
        excluded = {getattr(c, "id", None)
                    for c in (person1, person2) if c is not None}
        receivers = sorted(
            (p for p in participants
             if getattr(p, "co_nhan_tai_san", False)
             and getattr(p, "customer_id", None) not in excluded),
            key=lambda p: (-(getattr(p, "ty_le", 0) or 0),
                           getattr(p, "customer_id", 0) or 0))
        non_receivers = sorted(
            (p for p in participants
             if not getattr(p, "co_nhan_tai_san", False)
             and getattr(p, "customer_id", None) not in excluded),
            key=lambda p: getattr(p, "customer_id", 0) or 0)
        person3 = (getattr(receivers[0], "customer", None)
                   if receivers else None)
        slots = [None] * 21
        slots[1], slots[2], slots[3] = person1, person2, person3
        rest = ([p.customer for p in receivers[1:]]
                + [p.customer for p in non_receivers])
        for idx, c in enumerate(rest[:17], start=4):
            slots[idx] = c
        return slots

    people_slots = _person_slots()

    def _person_slot(slot: int, field_norm: str) -> str:
        person = (people_slots[slot]
                  if 0 < slot < len(people_slots) else None)
        if person is None:
            return ""
        value = getattr(person, {
            "ho_ten": "ho_ten", "so_giay_to": "so_giay_to",
            "dia_chi": "dia_chi", "loai_giay_to": "loai_giay_to",
            "noi_cap": "noi_cap", "loai_dia_chi": "loai_dia_chi",
        }.get(field_norm, field_norm), None)
        if field_norm == "ngay_cap":
            return word_engine._fmt_date(value)
        if field_norm in {"ngay_sinh", "ngay_chet"}:
            return word_engine._fmt_year(value)
        return word_engine._safe_text(value)

    return {
        "props": props, "rows": rows_per_asset, "row": _row,
        "area_total": _area_total, "area_total_num": _area_total_num,
        "purpose": _purpose,
        "authorized": [
            person["ho_ten"] for person, _r, _s in authorized_people],
        "authorized_block": _block,
        "authorized_inline": _inline,
        "authorized_field": _field,
        "authorized_list": _name_list,
        "niem_yet": case.noi_niem_yet,
        "person_slot": _person_slot,
        "today": date.today(),
        "case": case,
    }


def _fixture_lookup(token: str, fixture: dict):
    """(found, value) từ fixture cho token đáng kiểm chứng — oracle độc lập
    với ``build_template_mapping``."""
    normalized = word_engine._normalize_token(token[1:-1])

    m = _LAND_ROW_RE.match(normalized)
    if m:
        asset = int(m.group(2) or 1)
        row = fixture["row"](asset, word_engine.LAND_ROW_SLOTS[
            m.group(1).upper()])
        return True, str(row.get("dien_tich", "") or "")

    m = re.match(r"^loai dat tai san (\d+)\.(\d+)$", normalized)
    if m:
        row = fixture["row"](int(m.group(1)), int(m.group(2)))
        return True, str(row.get("loai_dat", "") or "")

    m = re.match(r"^loai dat (\d+)\.(\d+) - (loai dat|dien tich|thoi han)$",
                 normalized)
    if m:
        row = fixture["row"](int(m.group(1)), int(m.group(2)))
        key = {"loai dat": "loai_dat", "dien tich": "dien_tich",
               "thoi han": "thoi_han"}[m.group(3)]
        return True, str(row.get(key, "") or "")

    m = re.match(
        r"^(dien tich|thoi han|tho han) loai dat tai san (\d+)\.(\d+)$",
        normalized)
    if m:
        row = fixture["row"](int(m.group(2)), int(m.group(3)))
        key = "dien_tich" if m.group(1) == "dien tich" else "thoi_han"
        return True, str(row.get(key, "") or "")

    m = re.match(r"^muc dich su dung (?:tai san )?(\d+)$", normalized)
    if m:
        return True, fixture["purpose"](int(m.group(1)))

    m = re.match(r"^dong muc dich su dung (?:tai san )?(\d+)$", normalized)
    if m:
        return True, f"Mục đích sử dụng: {fixture['purpose'](int(m.group(1)))}."

    m = re.match(
        r"^(dien tich|dia chi dat|serial|so vao so|so thua|so to|loai so"
        r"|hinh thuc su dung|thoi han|tho han|nguon goc|ngay cap so"
        r"|co quan cap so|loai dat) tai san (\d+)$", normalized)
    if m:
        asset = int(m.group(2))
        if asset > len(fixture["props"]):
            return True, ""
        return True, _fixture_prop_field(
            fixture["props"][asset - 1], m.group(1), fixture, asset)

    m = re.match(r"^tai san (\d+) - (.+)$", normalized)
    if m:
        asset = int(m.group(1))
        if asset > len(fixture["props"]):
            return True, ""
        return True, _fixture_prop_field(
            fixture["props"][asset - 1], m.group(2), fixture, asset)

    # Canonical mặc định `[Tài sản - <trường>]` -> tài sản 1 (engine map
    # tương tự; review NAIA-12: oracle_missing trên dạng hợp lệ).
    m = re.match(r"^tai san - (.+)$", normalized)
    if m:
        return True, _fixture_prop_field(
            fixture["props"][0], m.group(1), fixture, 1)

    if normalized in ("noi niem yet", "niem yet"):
        value = fixture["niem_yet"] or ""
        # [NIÊM YẾT] toàn HOA -> giá trị HOA (engine map riêng).
        inner = token[1:-1]
        if any(ch.isalpha() for ch in inner) and inner == inner.upper():
            value = value.upper()
        return True, value

    # ---- Người được ủy quyền: oracle dựng từ participant vai trò đã lưu —
    # KHÔNG bao giờ qua engine mapping (review NAIA-12).
    m = re.match(r"^dong nguoi duoc uy quyen (\d+)$", normalized)
    if m:
        return True, fixture["authorized_block"](int(m.group(1)))
    m = re.match(r"^nguoi duoc uy quyen inline (\d+)$", normalized)
    if m:
        return True, fixture["authorized_inline"](int(m.group(1)))
    m = re.match(r"^nguoi duoc uy quyen (\d+) - (.+)$", normalized)
    if m:
        return True, fixture["authorized_field"](
            int(m.group(1)), m.group(2))
    if normalized == "danh sach nguoi duoc uy quyen":
        return True, fixture["authorized_list"](titled=False)
    if normalized == "danh sach nguoi duoc uy quyen inline":
        return True, fixture["authorized_list"](titled=True)
    m = re.match(r"^(.+) nguoi duoc uy quyen (\d+)$", normalized)
    if m:
        field = m.group(1)
        if field.startswith("danh sach"):
            return True, fixture["authorized_list"](
                titled="inline" in field)
        return True, fixture["authorized_field"](
            int(m.group(2)), field)

    # ---- Alias property không hậu tố "tài sản": engine ghi đè bằng giá trị
    # TÀI SẢN N (xem _add_property_placeholders) nên oracle theo asset N.
    m = re.match(r"^dien tich chu (?:tai san )?(\d+)$", normalized)
    if m:
        return True, _area_in_words(fixture, int(m.group(1)))
    if normalized == "dien tich chu":
        return True, _area_in_words(fixture, 1)
    m = re.match(r"^dien tich (\d+)$", normalized)
    if m:
        return True, fixture["area_total"](int(m.group(1)))
    if normalized == "dien tich":
        return True, fixture["area_total"](1)
    m = re.match(
        r"^(dia chi dat|serial|so vao so|so thua|so to|loai so"
        r"|hinh thuc su dung|nguon goc|ngay cap so|co quan cap so) (\d+)$",
        normalized)
    if m:
        asset = int(m.group(2))
        if asset > len(fixture["props"]):
            return True, ""
        return True, _fixture_prop_field(
            fixture["props"][asset - 1], m.group(1), fixture, asset)
    if normalized in (
            "dia chi dat", "serial", "so vao so", "so thua", "so to",
            "loai so", "hinh thuc su dung", "nguon goc", "ngay cap so",
            "co quan cap so"):
        return True, _fixture_prop_field(
            fixture["props"][0], normalized, fixture, 1)
    # [Loại đất N]/[Thời hạn N] (không chữ "tài sản"): alias ROW N của
    # tài sản 1 — ghi đè sau cùng tại mirror "Thờ hạn"->"Thời hạn" trong
    # build_template_mapping (đã kiểm chứng trên fixture: [Thời hạn 2]="").
    m = re.match(r"^loai dat (\d+)$", normalized)
    if m:
        return True, str(
            fixture["row"](1, int(m.group(1))).get("loai_dat", "") or "")
    m = re.match(r"^thoi han (\d+)$|^tho han (\d+)$", normalized)
    if m:
        row_n = int(m.group(1) or m.group(2))
        row = fixture["row"](1, row_n)
        value = str(row.get("thoi_han", "") or "")
        if not value and row_n == 1:
            value = str(getattr(
                fixture["props"][0], "thoi_han", "") or "")
        return True, value
    if normalized == "loai dat":
        return True, str(getattr(
            fixture["props"][0], "loai_dat", "") or "")
    if normalized in ("thoi han", "tho han"):
        return True, str(getattr(
            fixture["props"][0], "thoi_han", "") or "")
    # "muc dich su dung" trần không có suffix — engine KHÔNG map (export sẽ
    # fail placeholder sót); oracle vẫn trả expected tài sản 1 để soi nội
    # dung khi engine đổi contract.
    if normalized == "muc dich su dung":
        return True, fixture["purpose"](1)

    # ---- Token NEIGHBOR giải độc lập (không qua mapping engine): slot
    # người `[Tên N]`/… theo quy tắc _pick_core_people và token ngày/tháng/
    # năm — để wildcard kế bên không nuốt giá trị sai (review NAIA-12:
    # " SAI" nối vào tên UQ trước `[Tên 1]`/`[Ngày]`, "1.80" thay 80 bị
    # wildcard che).
    m = re.match(
        r"^(ten|nam sinh|cccd|ngay cap|dia chi|loai cc|noi cap cc"
        r"|thuong tru|nam chet) (\d+)$", normalized)
    if m:
        attr = {
            "ten": "ho_ten", "nam sinh": "ngay_sinh",
            "cccd": "so_giay_to", "ngay cap": "ngay_cap",
            "dia chi": "dia_chi", "loai cc": "loai_giay_to",
            "noi cap cc": "noi_cap", "thuong tru": "loai_dia_chi",
            "nam chet": "ngay_chet",
        }[m.group(1)]
        return True, fixture["person_slot"](int(m.group(2)), attr)
    if normalized == "nam chet":
        return True, fixture["person_slot"](1, "ngay_chet")
    today = fixture["today"]
    if normalized == "ngay":
        return True, f"{today.day:02d}"
    if normalized == "thang":
        return True, f"{today.month:02d}"
    if normalized == "nam":
        return True, str(today.year)
    if normalized == "ngay chu":
        return True, word_engine._ngay_chu(today.day)
    if normalized == "thang chu":
        return True, word_engine._so_thanh_chu(today.month)
    if normalized == "ngay lap ho so":
        return True, word_engine._fmt_date(
            getattr(fixture["case"], "ngay_lap_ho_so", None))
    # `[Ghi chú]` — field case độc lập từ fixture (engine map
    # `case.ghi_chu`): neighbor cạnh dấu `,`/`.` phải có expected riêng
    # thì rào số của token số kế bên mới phân biệt được (review NAIA-12:
    # `[ONT 3].[Ghi chú]` wildcard từng nuốt mutant `66.5`).
    if normalized == "ghi chu":
        return True, word_engine._safe_text(
            getattr(fixture["case"], "ghi_chu", ""))

    return False, ""


def _area_in_words(fixture: dict, asset: int) -> str:
    """`[Diện tích chữ (N)]` — engine chỉ viết HOA chữ cái đầu, không đơn vị."""
    total = fixture["area_total_num"](asset)
    return (word_engine._so_thanh_chu(total).capitalize() if total else "")


_PROP_FIELD_ATTR = {
    "dien tich": None,            # tổng diện tích — qua fixture['area_total']
    "dia chi dat": "dia_chi",
    "dia chi": "dia_chi",
    "serial": "so_serial",
    "so vao so": "so_vao_so",
    "so thua": "so_thua_dat",
    "so to": "so_to_ban_do",
    "loai so": "loai_so",
    "hinh thuc su dung": "hinh_thuc_su_dung",
    "thoi han": "thoi_han",
    "tho han": "thoi_han",
    "nguon goc": "nguon_goc",
    "ngay cap so": "ngay_cap",
    "co quan cap so": "co_quan_cap",
    "loai dat": "loai_dat",
    "muc dich su dung": None,     # qua fixture['purpose']
}


def _fixture_prop_field(prop, field_norm: str, fixture: dict, asset: int) -> str:
    """Giá trị field tổng của tài sản lấy thẳng từ fixture."""
    if field_norm == "dien tich":
        return fixture["area_total"](asset)
    if field_norm == "dien tich chu":
        return _area_in_words(fixture, asset)
    if field_norm == "muc dich su dung":
        return fixture["purpose"](asset)
    attr = _PROP_FIELD_ATTR.get(field_norm)
    if attr is None:
        return ""
    value = getattr(prop, attr, "")
    if attr == "ngay_cap":
        return value.strftime("%d/%m/%Y") if value else ""
    if isinstance(value, float):
        return f"{value:g}"
    return str(value or "")


def _expected_patterns(
    template: Path, fixture: dict,
) -> tuple[list[tuple[int, str]], list[str]]:
    """((chỉ-số-đoạn, pattern), oracle_missing) cho đoạn chứa token đáng
    kiểm chứng.

    Mỗi đoạn thành regex khớp TRỌN ĐÚNG ĐOẠN ĐÓ của file xuất bằng
    ``re.fullmatch`` trên text đã collapse whitespace — một global
    ``re.search`` trên cả văn bản dùng đoạn đúng để che đoạn lặp bị sai,
    còn substring trong đoạn thì bỏ qua ký tự thừa nối vào giá trị (review
    NAIA-12: ``Nguyễn Văn Chủ`` được nối `` SAI`` vẫn pass; ``120.5`` có
    dấu chấm câu literal đi sau phải hợp lệ còn ``180`` thay ``80`` phải
    trượt — fullmatch trọn đoạn cho cả hai phía):

    - literal giữ nguyên (khoảng trắng khớp ``\\s*`` — đề phòng ``w:br``
      không sinh text node giữa hai phần đoạn; dấu câu đứng liền giá trị
      như ``[X].`` so đúng ``giá trị.``);
    - token ĐÁNG KIỂM CHỨNG -> giá trị oracle FIXTURE escape nguyên văn,
      bọc ranh giới số ``(?<!\\d)``/``(?!\\.?\\d)`` khi giá trị đầu/cuối là
      chữ số — wildcard đoạn kế bên ``[^\\[\\]]*?`` KHÔNG được nuốt thêm
      phần số: ``80`` không khớp ``180`` (số thừa trước) lẫn ``80.5``
      (đuôi thập phân), ``Nguyễn Văn Chủ`` không khớp đoạn có `` SAI``.
      Giá trị RỖNG vẫn phát pattern — đoạn chỉ-chứa-token expected rỗng
      cho pattern ``""`` khớp đúng đoạn rỗng, output ``999`` trượt
      (review NAIA-12: pattern rỗng từng bị bỏ qua nên đoạn expected trống
      không được kiểm).
      Không có fallback engine (review NAIA-12): token interesting mà
      oracle không phủ được gom vào ``oracle_missing`` -> E2E FAIL, buộc
      bổ sung rule thay vì âm thầm chấp nhận giá trị sai (diện tích 999);
    - mã biểu mẫu cố định ``[NN]/[NN.N]`` -> literal — engine giữ nguyên
      dạng token nên đoạn chứa nó VẪN được kiểm phần literal/data;
    - token khác (slot người, ngày/tháng/năm, ...) -> cũng được giải qua
      fixture khi oracle có rule — neighbor giá trị THẬT thì wildcard kế
      bên không còn nuốt được phần sai nối vào giá trị oracle (review
      NAIA-12: ` SAI` sau tên UQ trước `[Tên 1]`/`[Ngày]` từng được
      `[^\\[\\]]*?` chấp nhận). Token oracle không phủ -> wildcard
      ``[^\\[\\]]*?`` không vượt qua ranh giới token.
    """
    import docx

    patterns: list[tuple[int, str]] = []
    oracle_missing: list[str] = []
    doc = docx.Document(str(template))
    for para_idx, text in enumerate(_doc_paragraph_texts(doc)):
        if "[" not in text:
            continue
        tokens = word_engine._PLACEHOLDER_RE.findall(text)
        if not any(_is_interesting(t) for t in tokens):
            continue
        parts: list[str] = []
        pos = 0
        matches = list(word_engine._PLACEHOLDER_RE.finditer(text))
        # Phân loại token TRƯỚC khi phát pattern — rào số qua dấu `,`/`.`
        # cần biết slot kề phía separator có giải được expected độc lập
        # không (review NAIA-12).
        kinds: list[str] = []        # "fixed" | "value" | "wildcard"
        values: list[str] = []
        for match in matches:
            token = match.group(0)
            if word_engine.is_fixed_form_code(token):
                kinds.append("fixed")
                values.append(token)
                continue
            found, value = _fixture_lookup(token, fixture)
            if found:
                kinds.append("value")
                values.append(value or "")
            else:
                if _is_interesting(token):
                    oracle_missing.append(f"{template.name}: {token}")
                kinds.append("wildcard")
                values.append("")
        for index, match in enumerate(matches):
            before_raw = match.string[pos:match.start()]
            after_raw = match.string[match.end():(
                matches[index + 1].start() if index + 1 < len(matches)
                else len(match.string))]
            parts.append(_ws_flex(before_raw))
            if kinds[index] == "fixed":
                parts.append(re.escape(values[index]))
            elif kinds[index] == "wildcard":
                parts.append(r"[^\[\]]*?")
            else:
                value = values[index]
                # Separator `,`/`.` ĐỨNG MỘT MÌNH giữa giá trị số và token
                # wildcard: hai slot hai bên separator không còn expected
                # độc lập đủ phân biệt — `66`+`.5` hay `66.5` đều khớp
                # được. Báo oracle thiếu bằng chứng thay vì âm thầm chấp
                # nhận (review NAIA-12: `[ONT 3].[Ghi chú]` chấp `66.5`,
                # `[Ghi chú].[ONT 3]` chấp `1.66` thay expected `66`).
                # Separator kèm literal khác (`.x`, `, `...) vẫn phân
                # biệt được nên không đếm.
                if (value[:1].isdigit()
                        and index > 0 and kinds[index - 1] == "wildcard"
                        and before_raw.strip() in (",", ".")) or (
                        value[-1:].isdigit()
                        and index + 1 < len(kinds)
                        and kinds[index + 1] == "wildcard"
                        and after_raw.strip() in (",", ".")):
                    oracle_missing.append(
                        f"{template.name}: {match.group(0)} separator "
                        f"cạnh wildcard")
                parts.append(
                    _value_pattern(value, before_raw, after_raw))
            pos = match.end()
        parts.append(_ws_flex(text[pos:]))
        pattern = "".join(parts)
        # Luôn ghi nhận pattern — kể cả "" (đoạn chỉ-chứa token expected
        # rỗng): fullmatch "" chỉ khớp đoạn rỗng, nên output "999" trượt
        # (review NAIA-12: `if pattern` bỏ qua expected rỗng -> lỗ hổng).
        patterns.append((para_idx, pattern))
    return patterns, oracle_missing


def _value_pattern(value: str, before: str = "", after: str = "") -> str:
    """Pattern cho giá trị fixture — escape + ranh giới số hai phía.

    ``(?<![\\d.,])``/``(?![\\d.,]?\\d)`` chặn prefix/suffix số lẫn dấu
    phân tách (``80`` không khớp ``180``, ``1.80``/``1,80``, ``80.5``,
    ``805`` hay ``80,5``) nhưng vẫn cho dấu chấm câu literal ``120.5.``
    đi sau. ``before``/``after`` là literal NGUỒN liền kề token: khi
    literal kết thúc/mở đầu bằng ``,``/``.`` thì dấu đó đã được literal
    match riêng — nó là DẤU PHÂN TÁCH giữa hai giá trị, không phải phần
    của giá trị — nên ranh giới chỉ còn chặn CHỮ SỐ dính trực tiếp
    (review NAIA-12: nguồn ``[ONT 3],[CLN 3]`` xuất đúng ``66,12`` từng
    bị coi dấu ``,`` literal là phần số -> fail oan). Chỉ bọc khi mép
    giá trị là chữ số để không thu hẹp giá trị chữ bình thường."""
    body = _ws_flex(value)
    # ``rstrip()[-1:]`` trên literal CHỈ-khoảng-trắng cho ``''`` — và
    # ``'' in ",."`` luôn True trong Python (chuỗi rỗng là substring của
    # mọi chuỗi) nên so membership trên TUPLE ký tự: separator chỉ đúng
    # khi ký tự thực sự là ``,``/``.`` (review NAIA-12: ``[ONT 3]
    # [Ghi chú]`` có after=" " -> sep_after True -> wildcard neighbor
    # nuốt mutant 66.5/1.66 thay expected 66).
    sep_before = before.rstrip()[-1:] in (",", ".")
    sep_after = after.lstrip()[:1] in (",", ".")
    if body and value[:1].isdigit():
        body = (r"(?<!\d)" if sep_before else r"(?<![\d.,])") + body
    if body and value[-1:].isdigit():
        body = body + (r"(?!\d)" if sep_after else r"(?![\d.,]?\d)")
    return body


def _ws_flex(text: str) -> str:
    """Escape literal; mọi khoảng trắng (kể cả đầu/cuối đoạn) -> ``\\s*`` —
    khớp cả khi export viết xuống dòng bằng ``w:br`` không sinh text node
    (hai phần đoạn dính liền) lẫn khi whitespace thừa."""
    parts = [re.escape(chunk) for chunk in text.split() if chunk]
    if not parts:
        return r"\s*" if text and text.isspace() else ""
    pattern = r"\s*".join(parts)
    if text[:1].isspace():
        pattern = r"\s*" + pattern
    if text[-1:].isspace():
        pattern = pattern + r"\s*"
    return pattern


# ----------------------------------------------------- context label oracle

# Chỉ số tài sản đọc từ token — kiểm nhãn `tài sản N` có khớp ngữ cảnh đoạn
# NGUỒN tương ứng không. Anchor (không phải dòng đất) thiết lập ngữ cảnh;
# token dòng đất phải thuộc tập tài sản nguồn cho phép — nhãn sai nhưng tự
# nhất quán (mutant "dùng tài sản 2 trong block tài sản 3") vẫn bị bắt vì
# expected suy từ file NGUỒN, không phải file migrated (review NAIA-12).
_TAI_SAN_IDX_RE = re.compile(r"tai san\s*(\d+)(?:\s*\.\s*(\d+))?")
# `[ONT]` trần = dòng vị trí 1 của TÀI SẢN 1 (spec word-output: token vị trí
# không suffix mặc định tài sản 1) — hoán đổi `[ONT]`->`[CLN]` không được
# manifest phê duyệt phải fail như `[ONT 3]`->`[CLN 3]` (review NAIA-12).
_LAND_ROW_TOKEN_RE = re.compile(r"^(ont|cln|nts|luc)(?:\s+(\d+))?$")
_CANONICAL_ROW_TOKEN_RE = re.compile(
    r"^loai dat\s+(\d+)\.(\d+)\s*-\s*(.+)$")
# Alias dạng "X [loại đất] tài sản N.M" — cầu nối "loại đất" là tùy chọn
# ([Loại đất tài sản N.M] chính là field "loai dat"; [Diện tích loại đất
# tài sản N.M] có prefix field "dien tich").
_TAI_SAN_ROW_TOKEN_RE = re.compile(
    r"^(?:(.*?)\s+)?(loai dat\s+)?tai san\s*(\d+)\s*\.\s*(\d+)$")
_TAI_SAN_FIELD_TOKEN_RE = re.compile(r"^tai san\s*(\d+)\s*-\s*(.+)$")
_FIELD_TAI_SAN_RE = re.compile(r"^(.+?)\s+tai san\s*(\d+)$")
_LEGACY_ROW_TOKEN_RE = re.compile(r"^(thoi han|tho han|loai dat)\s+(\d+)$")
_ANCHOR_TOKEN_RE = re.compile(r"^(so thua|so to|dia chi dat)\s+(\d+)$")


def _canon_field(text: str) -> str:
    """Chuẩn hoá tên trường trong tuple (tài sản, dòng, trường) — alias
    spelling 'thờ hạn' về 'thoi han' để hai dạng viết cùng một trường."""
    normalized = word_engine._normalize_token(text)
    return {"tho han": "thoi han"}.get(normalized, normalized)


def _asset_ref(token: str):
    """-> (tài_sản, dòng, trường, is_row, dạng) | None — chỉ nhận token có
    chỉ số TÀI SẢN rõ ràng, kèm tuple dòng/trường để đối chiếu occurrence.

    Tuple (tài_sản, dòng, trường) là danh tính SEMANTIC của occurrence —
    mọi alias cùng tuple (canonical `[Loại đất N.M - diện tích]`, alias
    `[Diện tích loại đất tài sản N.M]`, vị trí `[ONT N]`…) là counterpart
    của nhau, còn sai tài sản/dòng/trường là occurrence khác (review
    NAIA-12: oracle chỉ ghép theo dạng syntax nên đổi canonical sang alias
    sai tài sản/trường hoặc ONT->CLN vẫn pass).

    `[Tên 3]`/`[CCCD 3]` là chỉ số NGƯỜI, không phải tài sản — loại khỏi
    oracle context để tránh nhầm hai miền chỉ số. Token DÒNG đất (phải khớp
    ngữ cảnh, không tự thiết lập): `[X tài sản N.M]`/`[X loại đất tài sản
    N.M]`/`[Loại đất tài sản N.M]` (dạng ``taisan``), `[ONT|CLN|NTS|LUC N]`
    (``landrow``: dòng đất chứa diện tích của tài sản N — suffix N là tài
    sản, VỊ TRÍ dòng lấy theo mã), canonical `[Loại đất N.M - <trường>]`
    (``canonical``, có đủ tuple để tự xác minh cả khi đoạn không có anchor
    riêng — review NAIA-12: mutant 3.1 -> 2.1/3.2 trên canonical row-only
    từng pass), và dạng legacy `[Thời hạn|Loại đất N]` (``legacy``) = DÒNG
    N CỦA TÀI SẢN 1 — KHÔNG phải tài sản N (engine map "Thờ/Thời hạn N" ->
    dòng N tài sản 1; review NAIA-12: `[Thời hạn 2]` hợp lệ từng bị báo
    "sai tài sản 2"). Anchor (được tin để thiết lập ngữ cảnh): `[Số
    thửa|Số tờ|Địa chỉ đất N]` (``anchor``), `[<trường> tài sản N]` /
    `[Tài sản N - <trường>]` (``taisan``, không dòng). ``dòng``/``trường``
    là None khi token không mang thông tin đó.
    """
    normalized = word_engine._normalize_token(token[1:-1])
    m = _CANONICAL_ROW_TOKEN_RE.match(normalized)
    if m:
        return (int(m.group(1)), int(m.group(2)),
                _canon_field(m.group(3).strip()), True, "canonical")
    m = _LAND_ROW_TOKEN_RE.match(normalized)
    if m:
        return (int(m.group(2) or 1),
                word_engine.LAND_ROW_SLOTS[m.group(1).upper()],
                "dien tich", True, "landrow")
    m = _LEGACY_ROW_TOKEN_RE.match(normalized)
    if m:
        return (1, int(m.group(2)), _canon_field(m.group(1)), True,
                "legacy")
    m = _TAI_SAN_ROW_TOKEN_RE.match(normalized)
    if m and (m.group(1) or m.group(2)):
        field = m.group(1) or "loai dat"
        return (int(m.group(3)), int(m.group(4)),
                _canon_field(field), True, "taisan")
    m = _TAI_SAN_FIELD_TOKEN_RE.match(normalized)
    if m:
        return (int(m.group(1)), None, _canon_field(m.group(2)),
                False, "taisan")
    m = _FIELD_TAI_SAN_RE.match(normalized)
    if m:
        return (int(m.group(2)), None, _canon_field(m.group(1)),
                False, "taisan")
    m = _ANCHOR_TOKEN_RE.match(normalized)
    if m:
        return (int(m.group(2)), None, m.group(1), False, "anchor")
    m = _TAI_SAN_IDX_RE.search(normalized)
    if m:
        return (int(m.group(1)),
                int(m.group(2)) if m.group(2) else None,
                None, bool(m.group(2)), "taisan")
    return None


def _para_signature(text: str) -> str:
    """Chữ ký đoạn để ánh xạ nguồn->migrated: placeholder và số là hai phần
    migration được phép đổi — thay bằng marker chung rồi collapse."""
    sig = word_engine._PLACEHOLDER_RE.sub("¤", text)
    sig = re.sub(r"\d+", "", sig)
    return " ".join(sig.split()).lower()


def _predict_removed_source_paras(src_texts: list[str]) -> set[int]:
    """Chỉ số đoạn nguồn sẽ bị drop ``remove_line`` xoá khi migrate — lặp
    đúng quy tắc của ``drop_tokens_in_doc`` nhưng chỉ đọc, không ghi."""
    removed: set[int] = set()
    for idx, text in enumerate(src_texts):
        remainder = text
        for match in word_engine._PLACEHOLDER_RE.finditer(text):
            action = wtm._DROP_LOOKUP.get(wtm._token_key(match.group(0)))
            if action is None:
                continue
            remainder = remainder.replace(match.group(0), action[1])
        if remainder != text and wtm._ORPHAN_LINE_RE.match(
                remainder.strip()):
            removed.add(idx)
    return removed


def _align_source_paragraphs(
        src_texts: list[str], mig_texts: list[str]) -> dict[int, int]:
    """Ánh xạ chỉ số đoạn migrated -> chỉ số đoạn nguồn.

    Migration chỉ XOÁ đoạn (drop ``remove_line``) và đổi nội dung chỗ —
    không thêm/xáo thứ tự — nên dãy migrated là dãy nguồn bỏ bớt đoạn. Ưu
    tiên đoán đoạn bị drop rồi khớp theo thứ tự; lệch số lượng thì LCS trên
    ``_para_signature`` ('equal' ghép đoạn giữ nguyên, 'replace' ghép cùng
    thứ tự trong cụm đổi)."""
    removed = _predict_removed_source_paras(src_texts)
    kept = [i for i in range(len(src_texts)) if i not in removed]
    if len(kept) == len(mig_texts):
        return {j: kept[j] for j in range(len(mig_texts))}
    import difflib

    matcher = difflib.SequenceMatcher(
        a=[_para_signature(t) for t in src_texts],
        b=[_para_signature(t) for t in mig_texts],
        autojunk=False)
    mapping: dict[int, int] = {}
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for k in range(j2 - j1):
                mapping[j1 + k] = i1 + k
        elif tag == "replace":
            for k in range(min(i2 - i1, j2 - j1)):
                mapping[j1 + k] = i1 + k
    return mapping


def _stem_asset_index(source_name: str) -> int | None:
    """Chỉ số tài sản N của stem per-asset đã xác minh (mẫu DK18-N) —
    ràng buộc oracle ĐỘC LẬP, không suy từ nội dung file."""
    stem = wtm._normalize(Path(source_name).stem)
    if stem in wtm._LAND_LABEL_STEMS or stem in wtm._CONTEXT_TOKEN_RENAMES:
        match = re.search(r"(\d+)\s*$", stem)
        if match:
            return int(match.group(1))
    return None


def _replay_source_paragraph(
        src_text: str,
        stem_key: str | None,
        token_actions: dict[str, list[dict]],
        literal_budget: Counter) -> str:
    """Dựng TEXT expected của đoạn nguồn bằng cách replay đúng thứ tự
    action producer áp lên tài liệu — renamed -> dropped -> span literal
    (``_context_spans``) -> context token rename — khớp
    ``migrate_templates`` (renames trước, ``drop_tokens_in_doc``, rồi
    ``_apply_context_fixes``{literal trước, token rename sau}).

    Mọi action chỉ tiêu hao budget ``applied`` của nó: ``remaining``
    của record và ``literal_budget`` bị trừ dần theo từng occurrence/
    span thực sự replay — record không được chứng nhận chỗ khác (review
    NAIA-12). Record ``context_fix`` thuộc lineage ``removed_tokens``
    của một literal fix (``removal_only``) KHÔNG được replay: token đó
    đã nằm trong span literal thay thế rồi — ``removed_tokens`` chỉ là
    metadata đếm, không phải budget gỡ thứ hai.
    """
    text = src_text

    def _sub(action: dict) -> None:
        nonlocal text
        n = action["remaining"]
        old = action["old_text"]
        if n <= 0 or old not in text:
            return
        k = min(n, text.count(old))
        text = text.replace(old, action["new_text"], k)
        action["remaining"] -= k

    for action in token_actions["renamed"]:
        _sub(action)
    for action in token_actions["dropped"]:
        _sub(action)
    if (stem_key in wtm._LAND_LABEL_STEMS
            or stem_key in wtm._PURPOSE_ENUM_STEMS):
        # Span literal dự đoán trên text SAU rename/drop — đúng trạng
        # thái producer thấy khi _apply_context_fixes chạy: drop trước
        # có thể làm literal rời thành enum hợp lệ ("... m2; ; CLN:...")
        # và merge vẫn được phép (review NAIA-12: drop->merge hợp lệ
        # từng bị oracle từ chối oan vì dự đoán span trên text gốc).
        spans, _warn = wtm._context_spans(text, stem_key)
        applied_spans = []
        for s0, e0, new_tok, old_desc in sorted(spans):
            key = (wtm._normalize(old_desc), wtm._normalize(new_tok))
            if literal_budget[key] <= 0:
                continue
            literal_budget[key] -= 1
            applied_spans.append((s0, e0, new_tok))
        for s0, e0, new_tok in sorted(applied_spans, reverse=True):
            text = text[:s0] + new_tok + text[e0:]
    for action in token_actions["context_fix"]:
        if not action["removal_only"]:
            _sub(action)
    return text


def _slot_segments(text: str) -> list[tuple[str, object]]:
    """Tách đoạn thành dãy slot đan xen (literal, token) — literal gộp
    whitespace; token ref tài sản lấy tuple semantic (tài_sản, dòng,
    trường), token thường lấy text normalize — để so VỊ TRÍ chứ không
    chỉ multiset/thứ tự ref."""
    segs: list[tuple[str, object]] = []
    pos = 0
    for match in word_engine._PLACEHOLDER_RE.finditer(text):
        segs.append(("lit", " ".join(text[pos:match.start()].split())))
        ref = _asset_ref(match.group(0))
        segs.append(("tok", ref[:3] if ref is not None else
                     ("txt", wtm._normalize(match.group(0)))))
        pos = match.end()
    segs.append(("lit", " ".join(text[pos:].split())))
    return segs


def _context_label_violations(
    template_path: Path,
    source_path: Path | None = None,
    stem_asset: int | None = None,
    manifest_entry: dict | None = None,
):
    """Nhãn `tài sản N` trong file migrated có đúng ngữ cảnh NGUỒN không.

    Expected suy từ SOURCE + manifest (không phải file migrated): đoạn
    migrated ánh xạ về đoạn nguồn tương ứng qua ``_align_source_paragraphs``,
    rồi ``_replay_source_paragraph`` replay các action manifest đã phê
    duyệt THEO ĐÚNG THỨ TỰ producer (rename -> drop -> span literal ->
    context token rename) để dựng TEXT expected của từng đoạn nguồn.
    Phép so cuối đối chiếu SEQUENCE slot đan xen (literal, token) giữa
    expected và migrated qua ``_slot_segments``: ref tài sản so theo
    tuple semantic ``(tài_sản, dòng, trường)`` — mọi dạng viết cùng tuple
    (canonical, alias, vị trí ONT/CLN/NTS/LUC, legacy) là counterpart —
    còn mốc literal giữa các token cũng phải nguyên vị trí, nên dời
    token qua nhãn khác trong cùng đoạn, hoán hai token, hay xoá token
    vừa được context-fix đều trật slot -> violation (review NAIA-12:
    dời ``[Loại đất tài sản 3.1]`` từ nhãn ``Loại đất 1:`` sang ``Ghi
    chú:``, hoán ``[ONT 3]``/``[Ngày]``, và xoá ``[Địa chỉ đất tài sản
    3]`` sau context-fix từng pass exit 0).

    Budget literal chỉ tiêu hao tại đúng span ``_context_spans`` dự đoán
    lại trên text NGUỒN SAU drop — drop ``[Người ủy quyền2]`` làm enum
    rời thành cụm liên tiếp hợp lệ thì merge ``[Mục đích sử dụng N]``
    vẫn được replay đúng chỗ (review NAIA-12: dự đoán span trên text
    gốc từng từ chối oan drop->gộp enum hợp lệ). ``removed_tokens`` là
    lineage — record context_fix trùng lineage được đánh dấu
    ``removal_only`` và KHÔNG replay lần hai (không cấp budget gỡ thêm
    cho occurrence ngoài span). Token action tiêu hao ``applied`` theo
    TỪNG occurrence nguồn theo thứ tự — budget không dùng lại cho vị
    trí khác (review NAIA-12).

    Đoạn migrated không map được đoạn nguồn (hoặc nguồn .doc không đọc
    được) -> đếm ``unverified_rows`` theo lối bảo thủ, KHÔNG kết luận
    pass (review NAIA-12); stem per-asset DK18-N vẫn ràng buộc ref phải
    đúng tài sản N.
    -> (violations, {"unverified_rows": n})."""
    import docx

    mig_texts = _doc_paragraph_texts(docx.Document(str(template_path)))
    violations: list[str] = []
    unverified = 0

    def _refs(text: str):
        refs = []
        for match in word_engine._PLACEHOLDER_RE.finditer(text):
            ref = _asset_ref(match.group(0))
            if ref is not None:
                refs.append((match.group(0), ref))
        return refs  # [(token, (asset, row, field, is_row, form))]

    src_texts: list[str] | None = None
    if source_path is not None:
        try:
            src_texts = _doc_paragraph_texts(
                docx.Document(str(source_path)))
        except Exception:  # noqa: BLE001 — .doc legacy không mở bằng python-docx
            src_texts = None

    if src_texts is None:
        # Fallback: ngữ cảnh tự file migrated (đoạn anchor đồng nhất ->
        # ngữ cảnh chạy; row phải khớp) — dùng khi source là .doc legacy.
        current = None
        for idx, text in enumerate(mig_texts):
            refs = _refs(text)
            anchors = {ref[0] for _t, ref in refs if not ref[3]}
            expected = (
                next(iter(anchors)) if len(anchors) == 1 else current)
            for token, ref in refs:
                if not ref[3]:
                    continue
                if expected is None:
                    unverified += 1
                elif ref[0] != expected:
                    violations.append(
                        f"đoạn {idx}: {token} thuộc tài sản {ref[0]} nhưng "
                        f"ngữ cảnh là tài sản {expected}")
            if len(anchors) == 1:
                current = next(iter(anchors))
        return violations, {"unverified_rows": unverified}

    # ---- Action producer đã duyệt (manifest schema v2) ------------------
    # Budget literal theo (mô tả literal nguồn, token mới) — chỉ tiêu hao
    # tại đúng span ``_context_spans`` dự đoán lại trên đoạn nguồn.
    literal_auth_add: Counter = Counter()   # (norm old_desc, norm new)
    removal_keys: set = set()   # (norm token, norm new) bị literal gỡ
    entry = manifest_entry or {}
    for fix in entry.get("literal_fixes") or []:
        applied = int(fix.get("applied") or 0)
        new_key = wtm._normalize(str(fix.get("new") or ""))
        if applied > 0:
            literal_auth_add[
                (wtm._normalize(str(fix.get("old") or "")),
                 new_key)] += applied
        # ``removed_tokens`` là lineage đếm — KHÔNG phải budget removal
        # thứ hai (review NAIA-12: đếm đôi -> xoá thêm occurrence pass).
        # Record token context_fix trùng lineage chỉ để replay biết nó là
        # hạch toán của việc gỡ TRONG span, không phải rename thường.
        for tok in (fix.get("removed_tokens") or {}):
            removal_keys.add((wtm._normalize(str(tok)), new_key))

    # Token action replay theo ĐÚNG THỨ TỰ producer áp lên tài liệu
    # (migrate_templates: renames -> drops -> _apply_context_fixes{span
    # literal -> token rename}). Mọi token đều tham gia replay — kể cả
    # token KHÔNG phải ref tài sản như [Người ủy quyền2] bị drop — vì
    # expected là TEXT đoạn nguồn sau toàn bộ action, không chỉ dãy ref
    # (review NAIA-12: oracle so dãy ref nên drop->gộp enum hợp lệ bị từ
    # chối oan và xoá token sau context-fix bị bỏ qua).
    token_actions: dict[str, list[dict]] = {
        "renamed": [], "dropped": [], "context_fix": []}
    for rec in entry.get("tokens") or []:
        status = rec.get("status")
        if status not in token_actions:
            continue
        applied = int(rec.get("applied") or 0)
        if applied <= 0:
            continue
        old_key = wtm._normalize(str(rec.get("old") or ""))
        new_text = str(rec.get("new") or "")
        if status == "dropped":
            # Tra theo đúng khóa producer match token (`_token_key` — bỏ
            # dấu): record `[So cong chung]` không dấu được producer gỡ
            # qua lookup chuẩn hoá thì replay cũng nhận (review NAIA-12:
            # `_normalize` giữ dấu từng bỏ record -> expected sai).
            lookup = wtm._DROP_LOOKUP.get(
                wtm._token_key(str(rec.get("old") or "")))
            if lookup is None:   # producer cũng skip token không lọt DROP
                continue
            new_text = lookup[1]
        token_actions[status].append({
            "old_text": str(rec.get("old") or ""),
            "new_text": new_text,
            "remaining": applied,
            "removal_only": status == "context_fix"
                and (old_key, wtm._normalize(new_text)) in removal_keys})

    # Tập tài sản được phép theo đoạn NGUỒN: anchor của chính đoạn + ngữ
    # cảnh chạy hiện tại + stem per-asset — đoạn nhiều anchor mâu thuẫn cho
    # phép cả hai (không suy đoán). Mọi ref migrated ngoài tập đó là
    # violation — lớp này bắt ref ĐÚNG tuple nhưng sai ngữ cảnh tài sản
    # (replay một mình không bắt được vì src và out cùng sai giống nhau).
    src_allowed: list[set[int]] = []
    current = None
    for text in src_texts:
        anchors = {ref[0] for _t, ref in _refs(text) if not ref[3]}
        if len(anchors) == 1:
            current = next(iter(anchors))
        allowed = set(anchors)
        if current is not None:
            allowed.add(current)
        if stem_asset is not None:
            allowed.add(stem_asset)
        src_allowed.append(allowed)

    stem_key = None
    src_name = str(entry.get("source_name") or "")
    if src_name:
        stem_key = wtm._normalize(Path(src_name).stem)
    elif source_path is not None:
        stem_key = wtm._normalize(Path(source_path).stem)

    mapping = _align_source_paragraphs(src_texts, mig_texts)
    mapped_src: set[int] = set()
    for idx, text in enumerate(mig_texts):
        mig_refs = _refs(text)
        src_idx = mapping.get(idx)
        if src_idx is None:
            # Đoạn migrated mới không có counterpart nguồn — bảo thủ đếm
            # unverified trừ khi stem per-asset chứng minh sai tài sản.
            for token, ref in mig_refs:
                if ref[3]:
                    if stem_asset is not None and ref[0] != stem_asset:
                        violations.append(
                            f"đoạn {idx}: {token} tài sản {ref[0]} trong "
                            f"mẫu per-asset tài sản {stem_asset}")
                    else:
                        unverified += 1
            continue
        mapped_src.add(src_idx)
        # Expected = TEXT đoạn nguồn sau khi replay đúng thứ tự action
        # manifest đã phê duyệt (rename -> drop -> span literal -> context
        # token rename) rồi so SEQUENCE slot (literal, token) đan xen —
        # vị trí token trong đoạn và mốc literal phải nguyên: dời token
        # qua literal khác, hoán thứ tự token, hay xoá token sau
        # context-fix đều trật slot -> violation (review NAIA-12).
        expected_text = _replay_source_paragraph(
            src_texts[src_idx], stem_key, token_actions, literal_auth_add)
        exp_segs = _slot_segments(expected_text)
        act_segs = _slot_segments(text)
        if exp_segs != act_segs:
            exp_toks = [v for k, v in exp_segs if k == "tok"]
            act_toks = [v for k, v in act_segs if k == "tok"]
            missing = Counter(exp_toks) - Counter(act_toks)
            extra = Counter(act_toks) - Counter(exp_toks)
            for key, n in missing.items():
                violations.append(
                    f"đoạn {idx}: thiếu {n} ref {key} của đoạn nguồn")
            for key, n in extra.items():
                violations.append(
                    f"đoạn {idx}: thừa {n} ref {key} không có ở nguồn")
            if not missing and not extra and exp_toks != act_toks:
                violations.append(
                    f"đoạn {idx}: thứ tự ref bị hoán đổi "
                    f"{exp_toks} -> {act_toks}")
            exp_lits = [v for k, v in exp_segs if k == "lit"]
            act_lits = [v for k, v in act_segs if k == "lit"]
            if exp_lits != act_lits:
                violations.append(
                    f"đoạn {idx}: literal lệch vị trí/ngữ cảnh nguồn "
                    f"(expected {exp_lits} -> migrated {act_lits})")
        # Ngữ cảnh: mọi ref migrated (row lẫn anchor) phải thuộc tập tài
        # sản mà đoạn NGUỒN cho phép — mutant đổi SANG tài sản khác tuy
        # giữ đúng dạng/tuple-counterpart trong src cũng sai nguồn vẫn bị
        # bắt (review NAIA-12).
        allowed = src_allowed[src_idx]
        for token, ref in mig_refs:
            if allowed and ref[0] not in allowed:
                violations.append(
                    f"đoạn {idx}: {token} thuộc tài sản {ref[0]} nhưng "
                    f"ngữ cảnh nguồn cho phép {sorted(allowed)}")
    # Đoạn nguồn còn ref mà không map được (ngoài đoạn drop dự đoán) ->
    # unverified thay vì kết luận pass.
    predicted_removed = _predict_removed_source_paras(src_texts)
    for src_idx, src_text in enumerate(src_texts):
        if src_idx in mapped_src or src_idx in predicted_removed:
            continue
        for _t, ref in _refs(src_text):
            if ref[3]:
                unverified += 1
    return violations, {"unverified_rows": unverified}


def _required_authorized_block_slots(doc) -> set[int]:
    """Slot N của token FULL BLOCK `[Dòng người được ủy quyền N]` — chỉ dạng
    này đòi cả block nhiều dòng; token UQ inline/trường lẻ/danh sách được
    oracle kiểm GIÁ TRỊ trong batch.content, không suy ra block bắt buộc
    (review NAIA-12: `required_authorized_slots` gom mọi dạng UQ nên mẫu
    chỉ có trường tên bị đòi block -> fail oan)."""
    slots: set[int] = set()
    for token in word_engine._collect_placeholder_occurrences(doc):
        m = re.fullmatch(
            r"dong nguoi duoc uy quyen (\d+)",
            word_engine._normalize_token(token[1:-1]))
        if m:
            slots.add(int(m.group(1)))
    return slots



def _structural_signature(path: Path) -> dict:
    """Dấu vân cấu trúc/định dạng so được bằng python-docx.

    So sánh nguồn->đích theo khung (KHÔNG phải render trực quan): bố cục
    trang/lề của từng section, số bảng & đoạn trên mọi part, số part
    header/footer, tập style đoạn và font run đang dùng."""
    import docx

    qn = _qn()
    doc = docx.Document(str(path))
    partnames = [
        str(getattr(p, "partname", "")) for p in doc.part.package.iter_parts()]
    header_footer = sum(
        1 for name in partnames
        if "/word/header" in name or "/word/footer" in name)
    styles: set[str] = set()
    fonts: set[str] = set()
    paragraphs = 0
    tables = 0
    for part_root in word_engine._iter_word_part_roots(doc):
        paragraphs += len(list(part_root.element.iter(qn("w:p"))))
        tables += len(list(part_root.element.iter(qn("w:tbl"))))
        for para in part_root.element.iter(qn("w:p")):
            ppr = para.find(qn("w:pPr"))
            if ppr is not None:
                pstyle = ppr.find(qn("w:pStyle"))
                if pstyle is not None and pstyle.get(qn("w:val")):
                    styles.add(pstyle.get(qn("w:val")))
            for run in para.iter(qn("w:r")):
                rpr = run.find(qn("w:rPr"))
                if rpr is None:
                    continue
                rfonts = rpr.find(qn("w:rFonts"))
                sz = rpr.find(qn("w:sz"))
                fonts.add(str((
                    rfonts.get(qn("w:ascii")) if rfonts is not None else None,
                    sz.get(qn("w:val")) if sz is not None else None,
                    rpr.find(qn("w:b")) is not None)))
    return {
        "sections": [
            {"page_width": str(s.page_width), "page_height": str(s.page_height),
             "orientation": str(s.orientation),
             "margins": [str(s.left_margin), str(s.right_margin),
                         str(s.top_margin), str(s.bottom_margin)]}
            for s in doc.sections
        ],
        "tables": tables,
        "paragraphs": paragraphs,
        "header_footer_parts": header_footer,
        "styles": sorted(styles),
        "fonts": sorted(fonts),
    }


_POWERSHELL_STATS = r"""
$ErrorActionPreference = 'Continue'
$files = [System.IO.File]::ReadAllText($env:NAIA_STAT_FILES, [System.Text.Encoding]::UTF8) | ConvertFrom-Json
$word = $null
try {
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $word.DisplayAlerts = 0
  $word.AutomationSecurity = 3
  try { $word.Options.UpdateLinksAtOpen = $false } catch {}
  $out = @{}
  foreach ($f in $files) {
    $doc = $null
    try {
      $doc = $word.Documents.Open($f, $false, $true, $false)
      $out[$f] = @{
        pages = $doc.ComputeStatistics(2)
        words = $doc.ComputeStatistics(0)
        lines = $doc.ComputeStatistics(1)
        paragraphs = $doc.ComputeStatistics(4)
      }
    } catch {
      $out[$f] = @{ error = $_.Exception.Message }
    } finally {
      if ($doc -ne $null) { try { $doc.Close(0) } catch {} }
    }
  }
  $json = $out | ConvertTo-Json -Depth 4
  [System.IO.File]::WriteAllText($env:NAIA_STAT_OUT, $json, [System.Text.Encoding]::UTF8)
} finally {
  if ($word -ne $null) {
    try { $word.Quit() } catch {}
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
  }
}
"""


def _word_com_stats(paths: list[Path], work_dir: Path) -> dict[str, dict]:
    """ComputeStatistics qua Word COM cho từng file — bằng chứng Word đọc/
    phân trang được. Trả {} khi máy không có Word (ghi nhận limitation).

    Trao đổi qua file JSON UTF-8 thay vì stdout: console PowerShell mặc định
    cp1252 làm mất ký tự tiếng Việt trong đường dẫn."""
    import os
    import subprocess
    import tempfile

    if not paths:
        return {}
    work_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
            mode="w", encoding="utf-8", suffix=".json",
            dir=work_dir, delete=False) as fh:
        json.dump([str(p) for p in paths], fh, ensure_ascii=False)
        list_path = Path(fh.name)
    out_path = work_dir / f"{list_path.stem}.out.json"
    env = dict(os.environ)
    env["NAIA_STAT_FILES"] = str(list_path)
    env["NAIA_STAT_OUT"] = str(out_path)
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command",
             _POWERSHELL_STATS],
            capture_output=True, timeout=600, env=env)
    except (OSError, subprocess.SubprocessError):
        return {}
    finally:
        try:
            list_path.unlink()
        except OSError:
            pass
    try:
        # utf-8-sig: PowerShell 5.1 [Text.Encoding]::UTF8 ghi BOM đầu file.
        data = json.loads(out_path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        data = {}
    try:
        out_path.unlink()
    except OSError:
        pass
    try:
        work_dir.rmdir()  # chỉ xoá được khi đã trống — dọn scratch
    except OSError:
        pass
    return {str(k): v for k, v in data.items()} if isinstance(data, dict) else {}


# ------------------------------------------------------------------- run


def _preflight_write_targets(
    src_dir: Path, out_dir: Path, targets: list[Path], report_path: Path,
) -> None:
    """Chặn mọi cấu hình ghi nguy hiểm TRƯỚC khi tạo thư mục/file nào.

    Tái dùng guard containment của ``word_template_migration``: nguồn ≡ đích
    hoặc chồng lấn (kể cả qua junction/symlink) -> từ chối; mọi write target
    phải resolve nằm trong ``out_dir``; ``e2e_report.json`` đã tồn tại hoặc là
    liên kết -> từ chối ghi đè/đuổi theo link. Raise trước mọi mkdir/ghi nên
    nguồn và đích giữ nguyên bytes (review NAIA-12: run(src, src) từng tạo
    exports/report ngay trong nguồn dù migration từ chối)."""
    src_key = wtm._norm_key(src_dir).rstrip("\\/")
    out_key = wtm._norm_key(out_dir).rstrip("\\/")
    if (src_key == out_key
            or wtm._contained_key(out_key, src_key)
            or wtm._contained_key(src_key, out_key)):
        raise ValueError(
            "thư mục đích chồng lấn thư mục nguồn: "
            f"{out_dir} vs {src_dir}")
    for target in targets:
        if not wtm._contained_key(wtm._norm_key(target), out_key):
            raise ValueError(
                f"write target resolve ra ngoài thư mục đích: {target}")
    if report_path.exists() or report_path.is_symlink():
        raise FileExistsError(
            "e2e_report.json đã tồn tại hoặc là liên kết, từ chối ghi đè: "
            f"{report_path}")


def _format_status(
    signature_equal: bool | None,
    pages_src: int | None, pages_dst: int | None,
) -> str:
    """Phán quyết format theo BẰNG CHỨNG đo được (review NAIA-12: pages
    lệch mà ok=true từng xảy ra):

    - bất kỳ bằng chứng đo được nào lệch (signature khác / số trang khác /
      lỗi đọc signature -> signature_equal=False) -> ``failed``;
    - cả signature lẫn số trang đo được và khớp -> ``passed``;
    - thiếu bằng chứng (DOC chưa so được signature -> ``None``, COM không
      đo được số trang) -> ``not_verified`` — không mặc định pass.
    """
    if signature_equal is False:
        return "failed"
    if pages_src is not None and pages_dst is not None:
        if pages_src != pages_dst:
            return "failed"
        return "passed" if signature_equal is True else "not_verified"
    return "not_verified"


def _snapshot_source_tree(src_dir: Path) -> dict[str, str]:
    """sha256 của MỌI file template dưới nguồn + tên toàn bộ file con —
    file mới xuất hiện (exports/report lọt vào nguồn) cũng là mutation."""
    snapshot: dict[str, str] = {}
    for path in sorted(src_dir.rglob("*")):
        if path.is_file():
            snapshot[path.relative_to(src_dir).as_posix()] = (
                wtm.sha256_file(path))
    return snapshot


def run(src_dir: Path, out_dir: Path) -> dict:
    import docx  # noqa: F401 — dependency bắt buộc của toàn script

    # resolve(): Word COM tự phân giải path tương đối theo cwd của WORD,
    # không phải cwd của script -> path tuyệt đối cho mọi thao tác.
    src_dir, out_dir = Path(src_dir).resolve(), Path(out_dir).resolve()
    migrated_dir = out_dir / "migrated"
    exports_dir = out_dir / "exports"
    exports_neg_dir = out_dir / "exports_neg"
    scratch_dir = out_dir / "_scratch"
    report_path = out_dir / "e2e_report.json"
    # Preflight TRƯỚC mọi mkdir/ghi: nguồn/đích trùng hoặc chồng lấn, write
    # target thoát scope qua junction, report đã tồn tại/là link -> từ chối
    # ngay khi nguồn còn nguyên bytes.
    _preflight_write_targets(
        src_dir, out_dir,
        [migrated_dir, exports_dir, exports_neg_dir, scratch_dir,
         report_path],
        report_path)
    out_dir.mkdir(parents=True, exist_ok=True)
    # Batch export yêu cầu thư mục đích đã tồn tại (giống UI: user chọn thư
    # mục có sẵn) — tạo trước, migrate_templates tự tạo migrated/.
    exports_dir.mkdir(parents=True, exist_ok=True)
    exports_neg_dir.mkdir(parents=True, exist_ok=True)

    failures: list[str] = []
    checks: list[dict] = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        checks.append({"name": name, "ok": bool(ok), "detail": detail})
        if not ok:
            failures.append(f"{name}: {detail}" if detail else name)

    # ---- nguồn: snapshot trước (mọi file — file mới lọt vào cũng là ghi) ---
    src_files, excluded = wtm.scan_source_templates(src_dir)
    snapshot_before = _snapshot_source_tree(src_dir)

    case = _synthetic_case()
    mapping = word_engine.build_template_mapping(case)
    context = word_engine.build_word_context(case)
    fixture = _fixture_case_data(case)
    check(
        "engine.authorized_from_saved_role",
        [p.ho_ten for p in context.authorized_people] == [PROXY_NAME],
        f"authorized_people={[p.ho_ten for p in context.authorized_people]}")
    check(
        "engine.mapping_uq_slot1",
        PROXY_NAME in mapping.get("[Dòng người được ủy quyền 1]", ""),
        "mapping['[Dòng người được ủy quyền 1]'] không chứa đại diện đã gán")

    # ---- stage 1: migration --------------------------------------------------
    try:
        manifest = wtm.migrate_templates(
            src_dir, migrated_dir,
            registry=mapping.keys(),       # token engine đã biết -> "kept"
            no_source=word_engine.NO_SOURCE_PLACEHOLDERS)
    except Exception as exc:  # noqa: BLE001 — report rồi exit!=0
        check("migration", False, f"{type(exc).__name__}: {exc}")
        manifest = None
    else:
        check("migration", manifest["file_count"] == len(src_files),
              f"{manifest['file_count']}/{len(src_files)} file")

    # ---- stage 2: engine dry-run (diagnostic — KHÔNG thay thế batch E2E) -----
    engine_results = []
    if manifest is not None:
        for entry in manifest["files"]:
            template = migrated_dir / entry["output_name"]
            doc = docx.Document(str(template))
            replaced = word_engine.replace_in_doc(doc, mapping)
            engine_results.append({
                "template": entry["output_name"],
                "replaced": replaced,
                "unresolved": word_engine.find_unresolved_placeholders(
                    doc, template_ref=template),
                "required_authorized_slots": sorted(
                    word_engine.required_authorized_slots(doc)),
            })

    # ---- stage 3: batch export — đường THẬT của UI ---------------------------
    specs: list = []
    batch_documents: list[dict] = []
    options_documents: list[dict] = []
    if manifest is not None:
        specs = wbe.scan_template_dir(migrated_dir)
        resolve = wbe.template_path_resolver(migrated_dir)
        options = wbe.export_options(
            case, resolve_template=resolve, doc_specs=specs)
        options_documents = options["documents"]
        batch_error = None
        try:
            batch = wbe.export_batch(
                case, case_id=case.id,
                document_keys=[spec.document_key for spec in specs],
                dest_dir=exports_dir,
                resolve_template=resolve,
                doc_specs=specs)
        except wbe.WordBatchError as exc:
            batch = exc.result_data or {"documents": []}
            batch_error = exc.code
        batch_documents = batch.get("documents", [])
        if batch_error:
            check("batch.export", False,
                  f"export_batch raise WordBatchError: {batch_error}")

        # ---- coverage: manifest == scan == batch — không mẫu nào bị bỏ sót --
        manifest_keys = {e["output_name"] for e in manifest["files"]}
        scanned_keys = {s.document_key for s in specs}
        check(
            "batch.coverage.scan",
            scanned_keys == manifest_keys,
            f"scan thấy {len(scanned_keys)}/{len(manifest_keys)} mẫu; "
            f"chưa quét={sorted(manifest_keys - scanned_keys)[:5]}")
        batch_keys = {d["document_key"] for d in batch_documents}
        check(
            "batch.coverage.export",
            batch_keys == scanned_keys,
            f"batch trả {len(batch_keys)}/{len(scanned_keys)} văn bản; "
            f"thiếu={sorted(scanned_keys - batch_keys)[:5]}")

        # ---- context nguồn->migrated: nhãn `tài sản N` khớp ngữ cảnh NGUỒN
        # Expected suy từ anchor trong file NGUỒN tương ứng + stem per-asset
        # (DK18-N) — KHÔNG suy từ file migrated: mutant đổi cả anchor lẫn
        # rows sang tài sản khác (tự nhất quán) vẫn fail (review NAIA-12).
        context_unverified = 0
        for entry in manifest["files"]:
            violations, stats = _context_label_violations(
                migrated_dir / entry["output_name"],
                src_dir / entry["source_name"],
                _stem_asset_index(entry["source_name"]),
                manifest_entry=entry)
            context_unverified += stats["unverified_rows"]
            check(
                f"migration.context_labels[{entry['output_name']}]",
                not violations,
                f"{len(violations)} nhãn sai ngữ cảnh: {violations[:3]}")
        check("migration.context_labels.coverage", True,
              f"{context_unverified} dòng đất chưa có anchor để đối chiếu")

        # ---- stage 3b: kiểm từng file + nội dung mong đợi ---------------------
        oracle_missing: list[str] = []
        for d in batch_documents:
            key = d["document_key"]
            if d["status"] != "saved":
                check(f"batch.saved[{key}]", False,
                      f"status={d['status']} error={d.get('error')}")
                continue
            out_path = exports_dir / d["actual_filename"]
            if not out_path.is_file():
                check(f"batch.output_exists[{key}]", False,
                      f"thiếu {d['actual_filename']}")
                continue
            template = migrated_dir / key
            patterns, missing_tokens = _expected_patterns(template, fixture)
            oracle_missing.extend(missing_tokens)
            text_norm = _norm_ws(_doc_text(out_path))
            # Pattern khớp TRỌN ĐÚNG ĐOẠN của file xuất (re.fullmatch) —
            # đoạn lặp sai không được đoạn đúng khác chứng nhận thay, còn
            # giá trị bị nối thêm ("Nguyễn Văn Chủ SAI") cũng trượt vì so
            # trọn đoạn (review NAIA-12).
            out_paras = [
                _norm_ws(t)
                for t in _doc_paragraph_texts(docx.Document(str(out_path)))]
            missing = [
                pattern for para_idx, pattern in patterns
                if para_idx >= len(out_paras)
                or not re.fullmatch(pattern, out_paras[para_idx])]
            check(
                f"batch.content[{key}]", not missing,
                f"thiếu {len(missing)} đoạn mong đợi: {missing[:3]}")
            # Người được ủy quyền: chỉ token FULL BLOCK `[Dòng người được ủy
            # quyền N]` bắt buộc cả block của đúng người — token inline/
            # trường lẻ/danh sách đã được kiểm giá trị ở batch.content
            # (review NAIA-12: mẫu chỉ có [Họ tên người được ủy quyền 1]
            # từng bị đòi full block -> fail oan).
            requires_uq = _required_authorized_block_slots(
                docx.Document(str(template)))
            uq_wrong = [
                slot for slot in sorted(requires_uq)
                if not fixture["authorized_block"](slot)
                or _norm_ws(fixture["authorized_block"](slot))
                not in text_norm]
            check(
                f"batch.authorized[{key}]", not uq_wrong,
                f"slot UQ sai/thiếu: {uq_wrong}")
            leftover = word_engine.find_unresolved_placeholders(
                docx.Document(str(out_path)), template_ref=template)
            check(f"batch.unresolved[{key}]", not leftover,
                  ", ".join(leftover[:10]))
        check(
            "batch.oracle_missing", not oracle_missing,
            f"{len(oracle_missing)} token đáng kiểm oracle chưa phủ: "
            f"{sorted(set(oracle_missing))[:8]}")

        # ---- stage 3c: phản ví dụ — hồ sơ THIẾU vai trò ủy quyền phải fail ----
        negative: dict = {"status": "skipped",
                          "reason": "không mẫu nào yêu cầu người được ủy quyền"}
        uq_keys = [
            s.document_key for s in specs
            if word_engine.required_authorized_slots(
                docx.Document(str(migrated_dir / s.document_key)))]
        if uq_keys:
            case_no_uq = _synthetic_case(with_proxy=False)
            try:
                neg_batch = wbe.export_batch(
                    case_no_uq, case_id=case_no_uq.id, document_keys=uq_keys,
                    dest_dir=exports_neg_dir,
                    resolve_template=resolve, doc_specs=specs)
                neg_docs = neg_batch.get("documents", [])
            except wbe.WordBatchError as exc:
                neg_docs = (exc.result_data or {}).get("documents", [])
            # Kết quả phải đủ từng mẫu UQ với đúng lỗi per-file — danh sách
            # rỗng (job lỗi không result_data) KHÔNG được tính là pass.
            returned_keys = {d["document_key"] for d in neg_docs}
            missing_results = sorted(set(uq_keys) - returned_keys)
            bad = [
                d["document_key"] for d in neg_docs
                if d["status"] != "failed"
                or (d.get("error") or {}).get("code")
                != "word.missing_authorized_person"]
            leaked = list(exports_neg_dir.glob("*_HS-*.docx"))
            ok_neg = (not bad and not leaked and not missing_results
                      and len(neg_docs) == len(uq_keys))
            negative = {
                "status": "passed" if ok_neg else "failed",
                "templates_checked": uq_keys,
                "documents": [
                    {"document_key": d["document_key"],
                     "status": d["status"],
                     "error_code": (d.get("error") or {}).get("code")}
                    for d in neg_docs],
                "missing_results": missing_results,
                "bad_documents": bad,
                "outputs_leaked": [p.name for p in leaked],
            }
            check(
                "negative.missing_authorized_fails", ok_neg,
                f"{len(neg_docs)}/{len(uq_keys)} văn bản UQ có kết quả; "
                f"bad={bad[:5]}; thiếu result={missing_results[:5]}; "
                f"leaked={len(leaked)}")

    # ---- stage 4: so sánh cấu trúc/định dạng nguồn -> đích -------------------
    format_compare: list[dict] = []
    if manifest is not None:
        com_paths: list[Path] = []
        for entry in manifest["files"]:
            src_path = src_dir / entry["source_name"]
            dst_path = migrated_dir / entry["output_name"]
            item: dict = {
                "source": entry["source_name"],
                "output": entry["output_name"],
                "converted_from_doc": entry["converted_from_doc"],
            }
            com_paths += [src_path, dst_path]
            if not entry["converted_from_doc"]:
                try:
                    sig_src = _structural_signature(src_path)
                    sig_dst = _structural_signature(dst_path)
                except Exception as exc:  # noqa: BLE001
                    item["signature_equal"] = False
                    item["error"] = f"{type(exc).__name__}: {exc}"
                    check(f"format.signature[{entry['output_name']}]", False,
                          item["error"])
                else:
                    expected_paras = sig_src["paragraphs"] - entry.get(
                        "removed_lines", 0)
                    diffs = {
                        k: {"src": sig_src[k], "dst": sig_dst[k]}
                        for k in sig_src
                        if k != "paragraphs" and sig_src[k] != sig_dst[k]}
                    if sig_dst["paragraphs"] != expected_paras:
                        diffs["paragraphs"] = {
                            "src": sig_src["paragraphs"],
                            "dst": sig_dst["paragraphs"],
                            "expected_dst": expected_paras,
                            "removed_lines": entry.get("removed_lines", 0)}
                    item["signature_equal"] = not diffs
                    item["paragraph_delta_expected"] = entry.get(
                        "removed_lines", 0)
                    if diffs:
                        item["diff"] = diffs
                    check(
                        f"format.signature[{entry['output_name']}]",
                        not diffs, str(diffs)[:400])
            format_compare.append(item)

        com_stats = _word_com_stats(com_paths, scratch_dir)
        for item in format_compare:
            src_abs = str(src_dir / item["source"])
            dst_abs = str(migrated_dir / item["output"])
            item["word_com"] = {
                "source": com_stats.get(src_abs),
                "migrated": com_stats.get(dst_abs),
            }
            src_pages = (com_stats.get(src_abs) or {}).get("pages")
            dst_pages = (com_stats.get(dst_abs) or {}).get("pages")
            com_ok = src_pages is not None and dst_pages is not None
            item["word_pages_match"] = (
                src_pages == dst_pages if com_ok else None)
            # Trạng thái format theo BẰNG CHỨNG: sai lệch đã đo được -> failed;
            # thiếu Word COM/chưa so được -> not_verified, KHÔNG mặc định pass.
            status = _format_status(
                item.get("signature_equal"), src_pages, dst_pages)
            item["status"] = status
            check(
                f"format.pages[{item['output']}]",
                status != "failed",
                (f"nguồn={src_pages} trang, đích={dst_pages} trang"
                 if com_ok else "không đo được qua Word COM"))
        word_com_available = bool(com_stats)
    else:
        word_com_available = False

    # ---- nguồn không đổi (snapshot trước report — lần ghi cuối check lại) ---
    snapshot_after = _snapshot_source_tree(src_dir)
    check("sources_unchanged", snapshot_before == snapshot_after,
          "file nguồn thêm/xoá/sửa sau khi chạy")

    limitations = [
        "Export tạo DOCX bằng python-docx (đúng pipeline render của engine); "
        "chưa mở/duyệt bằng Microsoft Word thật ngoài ComputeStatistics.",
        "UI Electron chưa được chạy — luồng chọn thư mục/mẫu/đích qua giao "
        "diện chưa kiểm chứng (báo cáo batch chỉ chứng minh service path).",
        "format_compare so khung cấu trúc (section/bảng/đoạn/header-footer/"
        "style/font) + thống kê Word — KHÔNG phải so sánh trực quan từng "
        "trang render; chưa đủ để kết luận định dạng hiển thị giống hệt.",
    ]
    if not word_com_available:
        limitations.append(
            "Máy không có/không gọi được Word COM -> không có thống kê "
            "ComputeStatistics; phần so sánh còn hạn chế hơn nữa.")

    report = {
        "schema": "naia12.e2e-report.v5",
        "source_dir": str(src_dir),
        "migrated_dir": str(migrated_dir),
        "exports_dir": str(exports_dir),
        "ok": not failures,
        "checks": checks,
        "failures": failures,
        "migration": None if manifest is None else {
            "file_count": manifest["file_count"],
            "excluded_files": excluded,
            "totals": manifest["totals"],
            "files": manifest["files"],
        },
        "engine": {
            "authorized_people": [
                p.ho_ten for p in context.authorized_people],
            "dry_run": engine_results,
        },
        "batch": {
            "documents_scanned": len(specs),
            "options": options_documents,
            "documents": batch_documents,
            "saved": sum(1 for d in batch_documents if d["status"] == "saved"),
            "failed": sum(1 for d in batch_documents if d["status"] == "failed"),
        },
        "negative": negative if manifest is not None else {"status": "skipped"},
        "format_compare": format_compare,
        "format_status": {
            status: sum(1 for i in format_compare if i.get("status") == status)
            for status in ("passed", "failed", "not_verified")},
        "word_com_available": word_com_available,
        "ui": {
            "status": "not_run",
            "reason": "chưa chạy UI Electron — luồng chọn thư mục mẫu/mẫu/"
                      "đích trên giao diện chưa kiểm chứng trong E2E này",
        },
        "sources_unchanged": snapshot_before == snapshot_after,
        "source_sha256": snapshot_after,
        "limitations": limitations,
    }
    # Report = LẦN GHI CUỐI: 'xb' + re-check containment ngay trước ghi (không
    # ghi đè, không đuổi theo junction mới tạo giữa chừng). Danh tính report
    # lấy từ CHÍNH fd 'xb' còn mở (os.fstat) — chứng minh ownership trước
    # cả khi handle đóng; không stat() path sau đóng handle (khe đó cho
    # file cạnh tranh vào thay trước khi đọc identity — review NAIA-12).
    import os

    if not wtm._contained_key(wtm._norm_key(report_path),
                              wtm._norm_key(out_dir).rstrip("\\/")):
        raise ValueError(
            f"e2e_report.json resolve ra ngoài thư mục đích: {report_path}")
    initial_payload = json.dumps(
        report, ensure_ascii=False, indent=2).encode("utf-8")
    with open(report_path, "xb") as handle:
        handle.write(initial_payload)
        report_identity = _file_key(os.fstat(handle.fileno()))
    # Kiểm nguồn SAU lần ghi cuối — bắt cả file mới (report/export lọt vào
    # nguồn) lẫn sửa nội dung; lệch thì sửa lại report thành kết quả fail.
    # Report phải còn đúng file run tạo Ở CẢ HAI NHÁNH: report bị thay sau
    # capture nhưng nguồn nguyên cũng là failure (review NAIA-12: chỉ nhánh
    # nguồn đổi kiểm ownership nên trả ok=true trên report của file ngoài).
    # Khi report còn đúng danh tính run (inode + đúng bytes run vừa ghi —
    # Linux có thể tái dùng inode ngay sau unlink nên inode một mình không
    # đủ): ghi lại qua fd r+b ràng buộc inode — KHÔNG stat->unlink (khe
    # thay file giữa hai thao tác). Khi report bị thay/xoá: GIỮ NGUYÊN file
    # ngoài (không xoá/ghi đè), bản fail ghi sang e2e_report.failed.json
    # 'xb' — hoặc vào đúng tên chuẩn nếu path trống.
    snapshot_final = _snapshot_source_tree(src_dir)
    source_changed = snapshot_final != snapshot_before
    if source_changed:
        report["sources_unchanged"] = False
        report["ok"] = False
        report["checks"].append({
            "name": "sources_unchanged_final", "ok": False,
            "detail": "nguồn đổi sau lần ghi report cuối cùng"})
        report["failures"].append(
            "sources_unchanged_final: nguồn đổi sau lần ghi report cuối cùng")
        report["source_sha256"] = snapshot_final
    try:
        report_now = report_path.stat()
    except OSError:
        report_now = None
    report_bytes = b""
    if report_now is not None:
        try:
            report_bytes = report_path.read_bytes()
        except OSError:
            report_bytes = b""
    report_missing = report_now is None
    report_foreign = (
        report_missing
        or _file_key(report_now) != report_identity
        or report_bytes != initial_payload)
    if report_foreign:
        report["ok"] = False
        report["checks"].append({
            "name": "e2e_report.ownership", "ok": False,
            "detail": "e2e_report.json không còn là file run tạo — giữ "
                      "nguyên file đó, report fail ghi ở "
                      "e2e_report.failed.json"})
        report["failures"].append(
            "e2e_report.ownership: report bị thay/xoá trước snapshot "
            "cuối, không ghi đè file ngoài")
    if source_changed or report_foreign:
        payload = json.dumps(
            report, ensure_ascii=False, indent=2).encode("utf-8")
        if not wtm._contained_key(
                wtm._norm_key(report_path),
                wtm._norm_key(out_dir).rstrip("\\/")):
            raise ValueError(
                "e2e_report.json resolve ra ngoài thư mục đích: "
                f"{report_path}")
        if report_missing:
            # File run tạo đã bị xoá — path report còn trống nên 'xb'
            # vẫn ghi được bản fail vào đúng tên chuẩn mà không đụng file
            # ngoài nào.
            with open(report_path, "xb") as handle:
                handle.write(payload)
        elif report_foreign:
            failed_path = report_path.with_name("e2e_report.failed.json")
            if not wtm._contained_key(
                    wtm._norm_key(failed_path),
                    wtm._norm_key(out_dir).rstrip("\\/")):
                raise ValueError(
                    "e2e_report.failed.json resolve ra ngoài thư mục đích: "
                    f"{failed_path}")
            with open(failed_path, "xb") as handle:
                handle.write(payload)
        else:
            # Report còn đúng inode run — ghi lại qua fd ràng buộc inode:
            # mở r+b rồi fstat đối chiếu danh tính; file bị thay ngay trong
            # khe kiểm->ghi thì KHÔNG ghi đè file ngoài mà ghi failed.json.
            rewritten = _rewrite_owned_report(
                report_path, report_identity, initial_payload, payload)
            if not rewritten:
                report["ok"] = False
                report["checks"].append({
                    "name": "e2e_report.ownership", "ok": False,
                    "detail": "e2e_report.json bị thay giữa kiểm và ghi "
                              "lại — giữ nguyên file ngoài, report fail "
                              "ghi ở e2e_report.failed.json"})
                report["failures"].append(
                    "e2e_report.ownership: report bị thay giữa kiểm và "
                    "ghi lại, không ghi đè file ngoài")
                payload = json.dumps(
                    report, ensure_ascii=False, indent=2).encode("utf-8")
                failed_path = report_path.with_name(
                    "e2e_report.failed.json")
                if not wtm._contained_key(
                        wtm._norm_key(failed_path),
                        wtm._norm_key(out_dir).rstrip("\\/")):
                    raise ValueError(
                        "e2e_report.failed.json resolve ra ngoài thư mục "
                        f"đích: {failed_path}")
                with open(failed_path, "xb") as handle:
                    handle.write(payload)
    return report


def _file_key(stat) -> tuple[int, int]:
    """(volume, inode) — cặp mà ``unlink`` rồi tạo file khác tại cùng path
    có thể tái dùng trên Linux; vì vậy nơi gọi so thêm ĐÚNG bytes run ghi."""
    return (stat.st_dev, stat.st_ino)


def _rewrite_owned_report(
        report_path: Path, identity: tuple[int, int],
        expected: bytes, payload: bytes) -> bool:
    """Ghi lại report QUA fd ràng buộc inode run tạo — không ``unlink``.

    Mở ``r+b``, ``fstat`` trên chính fd đó chứng minh fd trỏ tới inode run
    đã đăng ký rồi mới truncate: file bị thay tại cùng path ngay trước/sau
    mở đều bị từ chối (trả False) thay vì ghi đè file ngoài (review
    NAIA-12: ``stat``->``unlink``->``xb`` từng xoá report cạnh tranh)."""
    import os

    try:
        handle = open(report_path, "r+b")
    except OSError:
        return False
    try:
        if _file_key(os.fstat(handle.fileno())) != identity:
            return False
        if handle.read() != expected:
            return False
        handle.seek(0)
        handle.truncate(0)
        handle.write(payload)
        handle.flush()
    except OSError:
        return False
    finally:
        handle.close()
    return True


def main(argv=None) -> int:
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("src_dir")
    parser.add_argument("out_dir")
    args = parser.parse_args(argv)

    try:
        report = run(Path(args.src_dir), Path(args.out_dir))
    except (ValueError, FileExistsError) as exc:
        # Preflight từ chối TRƯỚC mọi ghi — thoát rõ ràng, không traceback.
        print(json.dumps({
            "ok": False,
            "preflight_refused": f"{type(exc).__name__}: {exc}",
        }, ensure_ascii=False, indent=2))
        return 2
    summary = {
        "ok": report["ok"],
        "file_count": (report["migration"] or {}).get("file_count"),
        "batch_saved": report["batch"]["saved"],
        "batch_failed": report["batch"]["failed"],
        "negative": report["negative"]["status"],
        "format_status": report.get("format_status"),
        "word_com_available": report["word_com_available"],
        "sources_unchanged": report["sources_unchanged"],
        "ui": report["ui"]["status"],
        "failures": report["failures"][:10],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
