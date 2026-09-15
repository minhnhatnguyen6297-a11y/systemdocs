from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.table import Table

from services.fast_audit.models import WordAuditDoc, WordField
from services.fast_audit.rules import (
    AREA_RE,
    DOC_NO_RE,
    MAP_SHEET_RE,
    PLOT_RE,
    cut_at_notary_anchor,
    detect_word_template_type,
    extract_dates,
    extract_id_numbers,
    has_placeholders,
)


class WordParser:
    """Parse .docx files into normalized fields for audit."""

    MIN_NAME_LEN = 6

    def parse(self, path: str | Path) -> WordAuditDoc:
        path = Path(path)
        doc = Document(str(path))
        paragraphs = []
        for block in doc.iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        paragraphs.append(" | ".join(c for c in cells if c))
            else:
                text = block.text.strip()
                if text:
                    paragraphs.append(text)

        raw_text = "\n".join(paragraphs)
        body_text, _ = cut_at_notary_anchor(raw_text)
        body_paragraphs = [p for p in body_text.splitlines() if p.strip()]

        title = body_paragraphs[0] if body_paragraphs else ""
        template_type = detect_word_template_type(path.name, title, body_text)
        mode = "placeholder_template" if has_placeholders(body_text) else "filled_document"
        anchors_found = []

        fields: list[WordField] = []
        if mode == "placeholder_template":
            fields = self._parse_placeholder_fields(path.name, body_text, body_paragraphs)
        else:
            fields = self._parse_filled_fields(path.name, template_type, body_text, body_paragraphs)

        return WordAuditDoc(
            word_file=str(path),
            template_type=template_type,
            raw_text=raw_text,
            paragraphs=body_paragraphs,
            fields=fields,
            anchors_found=anchors_found,
            mode=mode,
        )

    def _parse_placeholder_fields(
        self, file_name: str, body_text: str, paragraphs: list[str]
    ) -> list[WordField]:
        fields = []
        for match in re.finditer(r"\[([^\[\]]+)\]", body_text):
            label = match.group(1).strip()
            value = self._extract_value_after_placeholder(label, body_text)
            fields.append(WordField(name=label, value=value, context="placeholder"))
        # Generic scans for important numbers even in placeholder mode.
        for doc_no in extract_id_numbers(body_text):
            fields.append(WordField(name="cccd_or_doc_no", value=doc_no, context="generic"))
        for date in extract_dates(body_text):
            fields.append(WordField(name="date", value=date, context="generic"))
        return fields

    @staticmethod
    def _extract_value_after_placeholder(label: str, text: str) -> str:
        # Naive: look for label and grab following line or sentence.
        pattern = re.escape(f"[{label}]") + r"\s*[:\-]?\s*(.+?)(?:\n|$)"
        m = re.search(pattern, text)
        if m:
            return m.group(1).strip()
        return ""

    def _parse_filled_fields(
        self, file_name: str, template_type: str, body_text: str, paragraphs: list[str]
    ) -> list[WordField]:
        fields: list[WordField] = []
        full = "\n".join(paragraphs)

        # Names: try common anchors.
        name_anchors = [
            ("declarant_name", r"(?:^|\n)\s*Tôi\s+là:?\s*([A-ZÀ-Ỹ][A-Za-zÀ-ỹ\s]+?)(?:,|;|\.|\n)"),
            ("declarant_name", r"(?:^|\n)\s*Chúng\s+tôi\s+gồm:?\s*([A-ZÀ-Ỹ][A-Za-zÀ-ỹ\s]+?)(?:,|;|\.|\n)"),
            ("mother_name", r"(?:^|\n)\s*Mẹ\s+đẻ\s+tôi\s+là:?\s*([A-ZÀ-Ỹ][A-Za-zÀ-ỹ\s]+?)(?:,|;|\.|\n)"),
            ("father_name", r"(?:^|\n)\s*Bố\s+đẻ\s+tôi\s+là:?\s*([A-ZÀ-Ỹ][A-Za-zÀ-ỹ\s]+?)(?:,|;|\.|\n)"),
            ("spouse_name", r"(?:^|\n)\s*Vợ/Chồng\s+tôi\s+là:?\s*([A-ZÀ-Ỹ][A-Za-zÀ-ỹ\s]+?)(?:,|;|\.|\n)"),
            ("deceased_name", r"(?:^|\n)\s*Ngườ\s+để\s+lại\s+di\s+sản\s+là:?\s*([A-ZÀ-Ỹ][A-Za-zÀ-ỹ\s]+?)(?:,|;|\.|\n)"),
        ]
        seen: set[tuple[str, str]] = set()
        for name_field, pattern in name_anchors:
            for match in re.finditer(pattern, full, re.IGNORECASE):
                value = match.group(1).strip()
                key = (name_field, value)
                if key not in seen:
                    seen.add(key)
                    fields.append(WordField(name=name_field, value=value, context="anchor"))

        # Generic Vietnamese names (table rows, headings, free text).
        for match in re.finditer(r"\b([A-ZÀ-Ỹ][a-zà-ỹ]*(?:\s+[A-ZÀ-Ỹ][a-zà-ỹ]*){1,5})\b", full):
            value = match.group(1).strip()
            if len(value) >= self.MIN_NAME_LEN:
                key = ("name", value)
                if key not in seen:
                    seen.add(key)
                    fields.append(WordField(name="name", value=value, context="generic"))

        # IDs and dates.
        for doc_no in extract_id_numbers(full):
            fields.append(WordField(name="id_number", value=doc_no, context="regex"))
        for date in extract_dates(full):
            fields.append(WordField(name="date", value=date, context="regex"))

        # Land fields.
        for match in PLOT_RE.finditer(full):
            fields.append(WordField(name="plot_no", value=match.group(1).strip(), context="regex"))
        for match in MAP_SHEET_RE.finditer(full):
            fields.append(WordField(name="map_sheet_no", value=match.group(1).strip(), context="regex"))
        for match in AREA_RE.finditer(full):
            fields.append(WordField(name="area", value=match.group(1).strip(), context="regex"))

        # Document number.
        for match in DOC_NO_RE.finditer(full):
            fields.append(WordField(name="document_no", value=match.group(1).strip(), context="regex"))

        # File name declarant hint.
        file_stem = Path(file_name).stem
        if template_type == "van_ban_tu_choi" and "_" in file_stem:
            parts = file_stem.split("_")
            if parts:
                hint = parts[-1].strip()
                if hint:
                    fields.append(WordField(name="file_declarant_hint", value=hint, context="filename"))

        return fields
