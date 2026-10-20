from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import date
from functools import lru_cache
from pathlib import Path
from types import SimpleNamespace
from typing import Any


HIDDEN_BUILTIN_TEMPLATE_NAMES = {
    "system_placeholder_reference.docx",
    "system_template_chuan_v1.docx",
}


ROLE_LABELS = {
    "nguoi_duoc_uy_quyen": "người được ủy quyền",
    "chu_dat": "chủ đất",
    "vo_chong": "vợ/chồng",
    "cha": "cha",
    "me": "mẹ",
    "cha_vo_chong": "cha vợ/chồng",
    "me_vo_chong": "mẹ vợ/chồng",
    "con": "con",
    "anh_chi_em": "anh/chị/em",
    "chau": "cháu",
    "vo_chong_nhanh": "vợ/chồng nhánh",
}

MAX_WORD_ASSETS = 5
MAX_WORD_LAND_ROWS = 20

# Token loại đất trong bộ mẫu cũ là VỊ TRÍ dòng, không phải tra theo mã: mẫu
# DK18/thuế in cứng "Loại đất 1: ONT; ... Diện tích: [ONT] m2" nên "[ONT]" là
# diện tích dòng loại đất 1, "[CLN]" dòng 2, "[NTS]" dòng 3, "[LUC]" dòng 4;
# hậu tố " N" là chỉ số tài sản ("[ONT 2]" = tài sản 2, dòng 1). Quyết định
# Owner (review NAIA-12): dòng 1 lưu CLN/120 thì "[ONT]" phải ra 120.
LAND_ROW_SLOTS = {
    "ONT": 1,
    "CLN": 2,
    "NTS": 3,
    "LUC": 4,
}


class WordExportValidationError(Exception):
    """Lỗi nghiệp vụ khi dữ liệu không đủ để xuất Word đúng."""


@dataclass
class WordPerson:
    id: Any = None
    row_id: str = ""          # row_id Stage — node two_party tham chiếu bằng row_id
    ho_ten: str = ""
    gioi_tinh: str = ""
    ngay_sinh: Any = None
    ngay_chet: Any = None
    so_giay_to: str = ""
    loai_giay_to: str = ""
    ngay_cap: Any = None
    noi_cap: str = ""
    dia_chi: str = ""
    loai_dia_chi: str = ""
    role: str = ""
    relation: str = ""
    share: str = ""
    will_receive: bool = False
    is_land_owner: bool = False
    is_in_diagram: bool = True
    is_hidden: bool = False
    is_deleted: bool = False
    side: str = ""            # two_party: "a" | "b" | ""
    slot: int = 0             # two_party: vị trí slot p1..p30 (0 = chưa gán)

    @property
    def is_deceased(self) -> bool:
        return bool(self.ngay_chet or self.loai_dia_chi == "Nơi chết"
                    or self.loai_giay_to in ("Trích lục khai tử",
                                            "Trích lục khai tử (bản sao)", "Giấy chứng tử"))

    @property
    def is_alive(self) -> bool:
        return not self.is_deceased

    @property
    def gender_normalized(self) -> str:
        raw = _normalize_token(self.gioi_tinh)
        if raw in {"nam", "nam"}:
            return "nam"
        if raw in {"nu", "nữ", "nu"}:
            return "nữ"
        return ""

    @property
    def address_label(self) -> str:
        return _safe_text(self.loai_dia_chi) or "Địa chỉ"


@dataclass
class WordExportContext:
    case: Any
    today: date
    persons: list[WordPerson]            # toàn bộ Stage, đúng thứ tự Stage
    assets: list[Any]
    all_people: list[WordPerson]         # active (có node, không ẩn/xóa)
    landowners: list[WordPerson]
    living_landowners: list[WordPerson]
    deceased_landowners: list[WordPerson]
    receivers: list[WordPerson]
    legal_refusal_people: list[WordPerson]
    case_type: str = "inheritance"
    has_diagram_state: bool = False      # diagram.state.nodes đã lưu tồn tại
    party_a: list[WordPerson] = None     # two_party: slot p1..p15
    party_b: list[WordPerson] = None     # two_party: slot p16..p30
    signers: list[WordPerson] = None     # người ký hợp lệ theo loại hồ sơ
    authorized_people: list[WordPerson] = None  # vai trò "người được ủy quyền"

    def __post_init__(self):
        if self.party_a is None:
            self.party_a = []
        if self.party_b is None:
            self.party_b = []
        if self.signers is None:
            self.signers = []
        if self.authorized_people is None:
            self.authorized_people = []


def _fmt_date(d: Any) -> str:
    if not d:
        return ""
    if isinstance(d, str):
        return _fmt_snapshot_date(d)
    return d.strftime("%d/%m/%Y")


def _fmt_snapshot_date(value: Any) -> str:
    """Ngày trong snapshot payload.case (ISO 'YYYY-MM-DD' hoặc năm
    'YYYY' — `_emit_date_or_year`) → dd/mm/YYYY cho placeholder Word."""
    text = _safe_text(value)
    if not text:
        return ""
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    return text


def _fmt_birth_or_year(d: Any) -> str:
    """Ngày/năm sinh cho token và câu dựng sẵn.

    Chuỗi snapshot mang đúng độ đầy đủ đã lưu: `YYYY` → năm, ISO →
    dd/mm/yyyy (không rò ISO thô), dạng khác giữ nguyên. `date` chỉ đến
    từ fallback master (hồ sơ cũ thiếu key) → in ĐẦY ĐỦ `dd/mm/yyyy`,
    không hạ ngày Jan-1 thành chỉ-năm trên đường case.
    """
    if not d:
        return ""
    if isinstance(d, str):
        text = d.strip()
        if re.fullmatch(r"\d{4}", text):
            return text
        iso = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", text)
        if iso:
            return f"{iso.group(3)}/{iso.group(2)}/{iso.group(1)}"
        return text
    return _fmt_date(d)


def _fmt_year(d: Any) -> str:
    """Chỉ lấy phần NĂM của mốc ngày — cho token `[Năm sinh N]`/`[Năm chết N]`.

    Quyết định Owner (NAIA-12): "giữ định dạng năm của token" — `[Năm chết]`
    in ra "2024", không phải ngày đầy đủ "03/02/2024".
    """
    if not d:
        return ""
    if isinstance(d, str):
        text = d.strip()
        iso = re.fullmatch(r"(\d{4})-\d{2}-\d{2}", text)
        if iso:
            return iso.group(1)
        # "2024" -> "2024"; "03/02/2024" -> "2024"; chuỗi khác giữ nguyên.
        match = re.search(r"(\d{4})\s*$", text)
        return match.group(1) if match else text
    return str(getattr(d, "year", "") or "")


def _land_row_area(land_rows: list[dict[str, Any]], position: int) -> str:
    """Diện tích dòng loại đất tại VỊ TRÍ ``position`` (1-based); vượt số dòng -> ""."""
    if position < 1 or position > len(land_rows):
        return ""
    return _safe_text(land_rows[position - 1].get("dien_tich", ""))


def _safe_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _serial_text(v: Any) -> str:
    """Serial GCN là nguyên văn (entities.md §2): giữ đúng khoảng trắng,
    chỉ đổi None → chuỗi rỗng cho khớp wire/placeholder."""
    return "" if v is None else str(v)


def _coerce_bool(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return value != 0
    normalized = _normalize_token(str(value))
    if normalized in {"1", "true", "yes", "co"}:
        return True
    if normalized in {"0", "false", "no", "khong", ""}:
        return False
    return default


def _so_thanh_chu(so: float) -> str:
    try:
        so = float(so)
    except (TypeError, ValueError):
        return ""

    don_vi = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]

    def _doc_ba_chu_so(n: int, *, full: bool = False) -> str:
        """Đọc nhóm 3 chữ số. ``full`` buộc đọc đủ "không trăm linh X" cho nhóm sau
        nhóm đầu (1005 -> "một nghìn không trăm linh năm", không phải "một nghìn năm")."""
        tram = n // 100
        chuc = (n % 100) // 10
        dv = n % 10
        parts: list[str] = []
        if tram or full:
            parts.append(don_vi[tram] + " trăm")
        if chuc == 0:
            if dv:
                parts.append(("linh " if (tram or full) else "") + don_vi[dv])
        else:
            parts.append(don_vi[chuc] + " mươi" if chuc > 1 else "mười")
            if dv == 1 and chuc > 1:
                parts.append("mốt")
            elif dv == 5:
                parts.append("lăm")
            elif dv:
                parts.append(don_vi[dv])
        return " ".join(parts)

    phan_nguyen = int(so)
    phan_le_str = ""
    if so != phan_nguyen:
        le = round(so - phan_nguyen, 6)
        dec_s = f"{le:.6f}".split(".")[1].rstrip("0")
        if dec_s:
            phan_le_str = " phẩy " + " ".join(don_vi[int(d)] for d in dec_s)

    if phan_nguyen == 0:
        return ("không" + phan_le_str).strip()

    parts = []
    n = phan_nguyen
    ty = n // 1_000_000_000
    n %= 1_000_000_000
    tr = n // 1_000_000
    n %= 1_000_000
    ng = n // 1_000
    n %= 1_000

    if ty:
        parts.append(_doc_ba_chu_so(ty) + " tỷ")
    if tr:
        parts.append(_doc_ba_chu_so(tr, full=bool(ty)) + " triệu")
    if ng:
        parts.append(_doc_ba_chu_so(ng, full=bool(ty or tr)) + " nghìn")
    if n:
        parts.append(_doc_ba_chu_so(n, full=bool(ty or tr or ng)))

    return (" ".join(parts) + phan_le_str).strip()


def _ngay_chu(day: Any) -> str:
    """Ngày đọc thành chữ; ngày 1-9 thêm "mùng" (catalog §2: `mùng ba`)."""
    try:
        value = int(day)
    except (TypeError, ValueError):
        return ""
    if value <= 0:
        return ""
    if value < 10:
        return "mùng " + _so_thanh_chu(value)
    return _so_thanh_chu(value)


def _normalize_token(s: str) -> str:
    s = (s or "").strip().lower()
    s = s.replace("đ", "d")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"\s+", " ", s)
    return s


def _role_key(role: str) -> str:
    normalized = _normalize_token(role)
    normalized = normalized.replace("/", " ").replace("_", " ")
    normalized = re.sub(r"\s+", " ", normalized).strip()
    if normalized in {"owner", "chu dat"}:
        return "chu_dat"
    if normalized in {"vo chong"}:
        return "vo_chong"
    if normalized == "cha":
        return "cha"
    if normalized in {"me", "me"}:
        return "me"
    if normalized in {"cha vc", "cha vo chong"}:
        return "cha_vo_chong"
    if normalized in {"me vc", "me vo chong"}:
        return "me_vo_chong"
    if normalized == "con":
        return "con"
    if normalized in {"anh chi em"}:
        return "anh_chi_em"
    if normalized in {"chau"}:
        return "chau"
    if normalized in {"con dau re", "vo chong nhanh"}:
        return "vo_chong_nhanh"
    return normalized.replace(" ", "_")


# Mirror `services.case_workspace._case_meta` (không import — case_workspace
# kéo sqlalchemy/models, word_engine phải chạy được ở web cũ).
TWO_PARTY_DOCUMENT_TYPES = ("chuyen_nhuong", "tang_cho",
                          "cho_thue", "dat_coc")

_DOC_TYPE_LABELS = {
    "thoa_thuan": "Thỏa thuận phân chia di sản",
    "khai_nhan": "Văn bản khai nhận di sản thừa kế",
    "chuyen_nhuong": "Hợp đồng chuyển nhượng quyền sử dụng đất",
    "tang_cho": "Hợp đồng tặng cho quyền sử dụng đất",
    "cho_thue": "Hợp đồng cho thuê quyền sử dụng đất",
    "dat_coc": "Hợp đồng đặt cọc",
}


def _word_case_type(case: Any) -> str:
    """case_type từ loai_van_ban — cùng rule `_case_meta`; loại lạ
    (ngoài CASE_TYPES) đi nhánh inheritance như flow cũ."""
    loai = _safe_text(getattr(case, "loai_van_ban", ""))
    if loai in TWO_PARTY_DOCUMENT_TYPES:
        return "two_party"
    return "inheritance"


def _display_document_type(raw: str) -> str:
    raw = _safe_text(raw)
    return _DOC_TYPE_LABELS.get(raw, raw)


def _format_share(value: Any) -> str:
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        return ""
    if number.is_integer():
        return str(int(number))
    return f"{number:g}"


def _pick_core_people(case: Any) -> tuple[Any, Any, Any, list[Any]]:
    owner = getattr(case, "nguoi_chet", None)
    spouse = None
    for p in getattr(case, "participants", []) or []:
        if _role_key(getattr(p, "vai_tro", "")) == "vo_chong":
            spouse = getattr(p, "customer", None)
            break

    pair = [c for c in [owner, spouse] if c is not None]
    nam = [c for c in pair if _normalize_token(getattr(c, "gioi_tinh", "")) == "nam"]
    nu = [c for c in pair if _normalize_token(getattr(c, "gioi_tinh", "")) in {"nu", "nữ"}]

    if len(nam) == 1 and len(nu) == 1:
        person1 = nam[0]
        person2 = nu[0]
    elif len(pair) == 2:
        person1 = pair[0]
        person2 = pair[1]
    elif len(pair) == 1:
        person1 = pair[0]
        person2 = None
    else:
        person1 = None
        person2 = None

    excluded_ids = {getattr(c, "id", None) for c in [person1, person2] if c is not None}
    receivers = [
        p for p in getattr(case, "participants", []) or []
        if getattr(p, "co_nhan_tai_san", False) and getattr(p, "customer_id", None) not in excluded_ids
    ]
    receivers = sorted(receivers, key=lambda p: (-(getattr(p, "ty_le", 0) or 0), getattr(p, "customer_id", 0) or 0))
    non_receivers = [
        p for p in getattr(case, "participants", []) or []
        if not getattr(p, "co_nhan_tai_san", False) and getattr(p, "customer_id", None) not in excluded_ids
    ]
    non_receivers = sorted(non_receivers, key=lambda p: getattr(p, "customer_id", 0) or 0)

    person3 = getattr(receivers[0], "customer", None) if receivers else None
    people_4_plus = [p.customer for p in receivers[1:]] + [p.customer for p in non_receivers]
    return person1, person2, person3, people_4_plus


def _parse_land_rows(raw: str) -> list[dict[str, Any]]:
    if not raw:
        return []
    try:
        rows = json.loads(raw)
    except Exception:
        return []
    if not isinstance(rows, list):
        return []
    return [row for row in rows if isinstance(row, dict)]


def _fmt_area(v: Any) -> str:
    """Diện tích → text; số nguyên bỏ phần thập phân thừa (80.0 → '80')."""
    if v is None or v == "" or isinstance(v, bool):
        return ""
    if isinstance(v, (int, float)):
        return f"{v:g}"
    return _safe_text(v)


def _land_row_values(row: Any) -> tuple[str, str, str]:
    """Một cụm đất → (loai_dat, dien_tich, thoi_han) dạng text.

    Đọc được dict snapshot/JSON — key canonical loaidat/dientich/thoihan
    hoặc key wire legacy loai_dat/dien_tich/thoi_han (kể cả lỗi chính tả
    'dien_tích' của data cũ) — và object bảng con PropertyLandRow."""
    if not row:
        return "", "", ""
    if isinstance(row, dict):
        loai = row.get("loaidat")
        if loai is None:
            loai = row.get("loai_dat")
        area = row.get("dientich")
        if area is None:
            area = row.get("dien_tích", row.get("dien_tich"))
        term = row.get("thoihan")
        if term is None:
            term = row.get("thoi_han")
    else:
        loai = getattr(row, "loaidat", getattr(row, "loai_dat", ""))
        area = getattr(row, "dientich", getattr(row, "dien_tich", ""))
        term = getattr(row, "thoihan", getattr(row, "thoi_han", ""))
    return _safe_text(loai), _fmt_area(area), _safe_text(term)


def _norm_land_dict(row: Any) -> dict[str, str]:
    loai, area, term = _land_row_values(row)
    return {"loai_dat": loai, "dien_tich": area, "thoi_han": term}


def _committed_land_rows(prop: Any) -> list[dict[str, str]]:
    """Cụm đất đã commit của một asset view → list dict đã normalize
    {loai_dat, dien_tich, thoi_han}. Nguồn theo thứ tự: `.land_rows`
    (bảng con ORM / snapshot view) → `land_rows_json` legacy → [].
    KHÔNG fallback trường lẻ — dùng _land_rows_for khi cần pseudo-row."""
    if prop is None:
        return []
    rows = getattr(prop, "land_rows", None)
    if rows:
        return [_norm_land_dict(r) for r in rows]
    parsed = _parse_land_rows(_safe_text(getattr(prop, "land_rows_json", "")))
    return [_norm_land_dict(r) for r in parsed]


def _land_rows_for(prop: Any) -> list[dict[str, str]]:
    """Cụm đất để render: committed rows, hoặc pseudo-row từ trường tổng
    hợp loai_dat/dien_tich khi chưa có cụm nào.

    Thời hạn lẻ (properties.thoi_han) KHÔNG được đắp vào cụm — thời hạn
    thuộc từng cụm (MIN-141 đợt 2); dữ liệu cũ chỉ có thời hạn lẻ được
    workspace_get báo orphan_thoi_han để đối chiếu, không tự đoán."""
    rows = _committed_land_rows(prop)
    if rows:
        return rows
    return [{
        "loai_dat": _safe_text(getattr(prop, "loai_dat", "")),
        "dien_tich": _fmt_area(getattr(prop, "dien_tich", "")),
        "thoi_han": "",
    }]


# ---------------------------------------------------------------------------
# Word export V2: case_state_json-aware helpers
# ---------------------------------------------------------------------------


def _load_case_state(case: Any) -> dict[str, Any]:
    raw = _safe_text(getattr(case, "case_state_json", ""))
    if not raw:
        return {}
    try:
        payload = json.loads(raw)
    except Exception:
        return {}
    if not isinstance(payload, dict):
        return {}
    return payload


def _word_person_from_stage(stage_person: dict[str, Any],
                            customers: Any = None) -> WordPerson:
    # Stage v3 rows: `row_id` uuid4 (khóa node diagram.state) + `entity_id`
    # int ORM (khóa projection legacy engineInput/engineState). Bản cũ
    # (schemaVersion 1/2) chỉ có `id` = entity id.
    entity = stage_person.get("entity_id")
    if entity is None:
        entity = stage_person.get("id")
    customer = (customers or {}).get(_safe_text(entity))

    def date_field(key: str) -> Any:
        """Ngày Người — cùng quyền key-presence với `_snapshot_person_view`:
        key có trong snapshot (kể cả `null`) → chặn fallback; key vắng
        (hồ sơ cũ) → master của hồ sơ, `date` render đầy đủ dd/mm/yyyy."""
        if key in stage_person:
            return stage_person.get(key) or None
        return getattr(customer, key, None) if customer is not None else None

    return WordPerson(
        id=_safe_text(entity) or None,
        row_id=_safe_text(stage_person.get("row_id")),
        ho_ten=_safe_text(stage_person.get("ho_ten")),
        gioi_tinh=_safe_text(stage_person.get("gioi_tinh")),
        ngay_sinh=date_field("ngay_sinh"),
        ngay_chet=date_field("ngay_chet"),
        so_giay_to=_safe_text(stage_person.get("so_giay_to")),
        loai_giay_to=_safe_text(stage_person.get("loai_giay_to")),
        ngay_cap=date_field("ngay_cap"),
        noi_cap=_safe_text(stage_person.get("noi_cap")),
        dia_chi=_safe_text(stage_person.get("dia_chi")),
        loai_dia_chi=_safe_text(stage_person.get("loai_dia_chi", stage_person.get("place_of_origin"))),
    )


def _word_person_from_customer(customer: Any) -> WordPerson:
    return WordPerson(
        id=getattr(customer, "id", None),
        ho_ten=_safe_text(getattr(customer, "ho_ten", "")),
        gioi_tinh=_safe_text(getattr(customer, "gioi_tinh", "")),
        ngay_sinh=getattr(customer, "ngay_sinh", None),
        ngay_chet=getattr(customer, "ngay_chet", None),
        so_giay_to=_safe_text(getattr(customer, "so_giay_to", "")),
        loai_giay_to=_safe_text(getattr(customer, "loai_giay_to", "")),
        ngay_cap=getattr(customer, "ngay_cap", None),
        noi_cap=_safe_text(getattr(customer, "noi_cap", "")),
        dia_chi=_safe_text(getattr(customer, "dia_chi", "")),
        loai_dia_chi=_safe_text(getattr(customer, "loai_dia_chi", "")),
    )


def _apply_diagram_nodes(
    persons: list[WordPerson],
    nodes: list[dict[str, Any]],
    allocations: dict[str, Any] | None = None,
) -> list[WordPerson]:
    """Gắn cờ/role/share từ node Diagram (engineInput projection —
    `personId` là entity id). KHÔNG đổi thứ tự `persons`: `[Người N]`
    theo thứ tự Stage đã lưu (NAIA-8) — thao tác Diagram không được đổi
    số Người."""
    if not nodes:
        return persons
    for person in persons:
        person.is_in_diagram = False
    # personId trên node có thể là entity_id (projection legacy
    # engineInput/engineState) hoặc row_id Stage (diagram.state v3) —
    # index cả hai khóa.
    by_id: dict[str, WordPerson] = {}
    for p in persons:
        if p.id is not None:
            by_id[str(p.id)] = p
        if p.row_id:
            by_id[p.row_id] = p
    for node in nodes:
        if not isinstance(node, dict):
            continue
        person_id = _safe_text(node.get("personId") or (node.get("person") or {}).get("id"))
        if not person_id or person_id not in by_id:
            continue
        p = by_id[person_id]
        allocation = (allocations or {}).get(person_id, {})
        if not isinstance(allocation, dict):
            allocation = {}
        p.is_in_diagram = True
        p.role = _safe_text(node.get("role")) or p.role
        p.relation = _safe_text(node.get("relationLabel")) or p.relation
        p.share = _safe_text(
            allocation.get("finalShare")
            or node.get("finalShare")
            or node.get("share")
            or node.get("sharePercent")
        )
        legacy_decision = _safe_text(node.get("inheritanceDecision"))
        p.will_receive = _coerce_bool(
            node.get("willReceive"),
            bool(node.get("receivePositions"))
            or legacy_decision == "accept")
        p.is_land_owner = _coerce_bool(
            node.get("isLandOwner"), bool(node.get("ownPositions")))
        p.is_hidden = _coerce_bool(node.get("hidden"), False)
        p.is_deleted = _coerce_bool(node.get("deleted"), False)
    return persons


def _apply_two_party_slots(
    persons: list[WordPerson],
    nodes: list[dict[str, Any]],
) -> list[WordPerson]:
    """Gắn slot/bên từ `diagram.state.nodes` canonical p1..p30
    (`personId` = row_id Stage — §13.5). Node không gán người = slot
    trống; người Stage không có slot giữ `is_in_diagram=False` (vẫn đứng
    trong nhóm `[Người N]` theo thứ tự Stage)."""
    for person in persons:
        person.is_in_diagram = False
    by_row = {p.row_id: p for p in persons if p.row_id}
    for node in nodes or []:
        if not isinstance(node, dict):
            continue
        nid = _safe_text(node.get("id"))
        if not re.fullmatch(r"p([1-9]|[12][0-9]|30)", nid):
            continue
        row_id = _safe_text(node.get("personId"))
        p = by_row.get(row_id)
        if not row_id or p is None:
            continue
        slot = int(nid[1:])
        p.is_in_diagram = True
        p.slot = slot
        p.side = "a" if slot <= 15 else "b"
        p.role = "Bên A" if p.side == "a" else "Bên B"
        p.is_hidden = _coerce_bool(node.get("hidden"), False)
        p.is_deleted = _coerce_bool(node.get("deleted"), False)
    return persons


@dataclass
class _WordPersonsBuild:
    """Kết quả `_build_word_persons` — persons giữ nguyên thứ tự Stage
    (NAIA-8: `[Người N]` = số thứ tự trong danh sách Stage đã lưu)."""
    persons: list[WordPerson]
    case_type: str
    has_diagram_state: bool


def _build_word_persons(case: Any) -> _WordPersonsBuild:
    state = _load_case_state(case)
    has_state = bool(state)
    case_type = _word_case_type(case)
    stage = state.get("stage", []) if has_state else []
    diagram = state.get("diagram", {}) if has_state and isinstance(state.get("diagram"), dict) else {}
    diagram_state = diagram.get("state", {}) if isinstance(diagram.get("state"), dict) else {}
    state_nodes = diagram_state.get("nodes", [])
    has_diagram_state = isinstance(state_nodes, list) and bool(state_nodes)
    engine_input = diagram.get("engineInput", {}) if isinstance(diagram.get("engineInput"), dict) else {}
    legacy_engine_state = diagram.get("engineState", {}) if isinstance(diagram.get("engineState"), dict) else {}
    engine_result = diagram.get("engineResult", {}) if isinstance(diagram.get("engineResult"), dict) else {}
    node_source = engine_input if isinstance(engine_input.get("nodes"), list) else legacy_engine_state
    nodes = node_source.get("nodes", []) if isinstance(node_source.get("nodes"), list) else []
    allocations = engine_result.get("allocations", {})
    if not isinstance(allocations, dict):
        allocations = {}

    persons: list[WordPerson] = []
    owner = getattr(case, "nguoi_chet", None)

    if has_state and isinstance(stage, list):
        # Khách hàng của hồ sơ theo `id` entity (nguồn giống
        # `_pick_core_people`) — chỉ dùng làm fallback 3 field ngày khi
        # snapshot cũ thiếu key, không đổi thứ tự/row_id/nhóm Người.
        customers_by_id: dict[str, Any] = {}
        for candidate in [owner] + [
                getattr(p, "customer", None)
                for p in getattr(case, "participants", []) or []]:
            if candidate is not None:
                customers_by_id.setdefault(
                    _safe_text(getattr(candidate, "id", "")), candidate)
        for item in stage:
            if isinstance(item, dict):
                persons.append(_word_person_from_stage(item, customers_by_id))

    # Đảm bảo ngườ chết của hồ sơ luôn có mặt trong danh sách để xử lý chủ đất.
    # two_party: nguoi_chet chỉ là anchor kỹ thuật (không có ngữ nghĩa
    # chủ đất) → không gắn cờ/bổ sung.
    if owner is not None and case_type != "two_party":
        owner_id = str(getattr(owner, "id", ""))
        existing = next((p for p in persons if str(p.id) == owner_id), None)
        if existing is None:
            owner_wp = _word_person_from_customer(owner)
            owner_wp.is_land_owner = True
            persons.append(owner_wp)
        else:
            existing.is_land_owner = True

    if has_state:
        if case_type == "two_party":
            # Node v3 two_party tham chiếu Stage bằng row_id, có slot p1..p30.
            persons = _apply_two_party_slots(persons, state_nodes)
        else:
            # Inheritance: projection legacy engineInput/engineState trỏ
            # entity_id; fallback diagram.state v3 (personId = row_id) khi
            # projection thiếu.
            persons = _apply_diagram_nodes(
                persons, nodes or state_nodes, allocations)
        _merge_participant_roles(persons, case)

    # Fallback về participants cũ khi không có case_state_json hợp lệ.
    if not has_state:
        for participant in getattr(case, "participants", []) or []:
            customer = getattr(participant, "customer", None)
            if customer is None:
                continue
            wp = _word_person_from_customer(customer)
            wp.role = _safe_text(getattr(participant, "vai_tro", ""))
            wp.will_receive = bool(getattr(participant, "co_nhan_tai_san", False))
            wp.share = _format_share(getattr(participant, "ty_le", ""))
            persons.append(wp)

    deduped: list[WordPerson] = []
    seen_ids: set[str] = set()
    for index, person in enumerate(persons):
        # Stage rows không có entity id thì dedupe bằng row_id; cả hai
        # đều trống → giữ nguyên từng dòng (không gộp nhầm người lạ).
        person_id = person.row_id or _safe_text(person.id) or f"#{index}"
        if person_id in seen_ids:
            continue
        seen_ids.add(person_id)
        deduped.append(person)
    return _WordPersonsBuild(
        persons=deduped, case_type=case_type,
        has_diagram_state=has_diagram_state)


def _merge_participant_roles(persons: list[WordPerson], case: Any) -> None:
    """Bổ sung vai trò đã gán TAY trên participants (``vai_tro``) cho Stage persons.

    Node v3 của diagram không lưu role tự do, nên "Người được ủy quyền" — vai trò
    tác nghiệp xác nhận — chỉ còn sống ở `InheritanceParticipant.vai_tro`. Merge
    theo customer id: vai trò đã lưu chỉ ghi đè khi person chưa có role hoặc chính
    vai trò đó là vai trò ủy quyền đã xác nhận (không suy từ vị trí/topology).
    """
    stored: dict[str, str] = {}
    for participant in getattr(case, "participants", []) or []:
        role = _safe_text(getattr(participant, "vai_tro", ""))
        customer_id = getattr(participant, "customer_id", None)
        if role and customer_id is not None:
            stored[str(customer_id)] = role
    if not stored:
        return
    for person in persons:
        role = stored.get(_safe_text(person.id))
        if role and (
                not person.role
                or _role_key(role).startswith("nguoi_duoc_uy_quyen")):
            person.role = role


def _active_persons(persons: list[WordPerson]) -> list[WordPerson]:
    return [p for p in persons if p.is_in_diagram and not p.is_hidden and not p.is_deleted]


def _children(persons: list[WordPerson]) -> list[WordPerson]:
    return [p for p in _active_persons(persons) if _role_key(p.role) == "con"]


def _dedupe_people(people: list[WordPerson]) -> list[WordPerson]:
    """Dedupe theo danh tính Stage (row_id → id) — dùng cho signers."""
    out: list[WordPerson] = []
    seen: set[str] = set()
    for index, person in enumerate(people):
        key = person.row_id or _safe_text(person.id) or f"#{index}"
        if key in seen:
            continue
        seen.add(key)
        out.append(person)
    return out


def committed_landowner_names(payload: dict) -> str | None:
    """Case field from assigned, active Diagram owners only; no master fallback."""
    diagram = payload.get("diagram", {})
    state = diagram.get("state", {}) if isinstance(diagram, dict) else {}
    if state.get("domain") != "inheritance":
        return None
    stage = payload.get("stage", [])
    persons = [_word_person_from_stage(row, {}) for row in stage]
    asset_count = len(payload.get("assets", []))
    nodes = []
    for node in state.get("nodes", []):
        positions = node.get("ownPositions", [])
        # Committed/pruned positions only: stale legacy positions are not assets.
        owns = any(isinstance(pos, int) and not isinstance(pos, bool)
                   and 1 <= pos <= min(3, asset_count) for pos in positions)
        nodes.append({**node, "isLandOwner": owns})
    _apply_diagram_nodes(persons, nodes)
    owners = _dedupe_people([
        person for person in _active_persons(persons) if person.is_land_owner
    ])
    return "\n".join(person.ho_ten for person in owners if person.ho_ten) or None


def build_word_context(case: Any, today: date | None = None) -> WordExportContext:
    built = _build_word_persons(case)
    persons = built.persons
    case_type = built.case_type
    assets = _get_property_list(case)

    if case_type == "two_party":
        # Không có projection thừa kế — active = người đã gán slot còn
        # hiện lực; bên A/B theo thứ tự slot p1..p15 / p16..p30.
        all_people = sorted(_active_persons(persons),
                            key=lambda p: p.slot)
        party_a = [p for p in all_people if p.side == "a"]
        party_b = [p for p in all_people if p.side == "b"]
        return WordExportContext(
            case=case, today=today or date.today(),
            persons=persons, assets=assets, all_people=all_people,
            landowners=[], living_landowners=[], deceased_landowners=[],
            receivers=[], legal_refusal_people=[],
            case_type=case_type,
            has_diagram_state=built.has_diagram_state,
            party_a=party_a, party_b=party_b,
            signers=[*party_a, *party_b],
            authorized_people=_authorized_people(persons))

    all_people = _active_persons(persons)
    landowners = [person for person in all_people if person.is_land_owner]
    receivers = [person for person in all_people if person.will_receive]
    living_landowners = [p for p in landowners if p.is_alive]
    signers = _dedupe_people([*living_landowners, *receivers])
    return WordExportContext(
        case=case,
        today=today or date.today(),
        persons=persons,
        assets=assets,
        all_people=all_people,
        landowners=landowners,
        living_landowners=living_landowners,
        deceased_landowners=[person for person in landowners if person.is_deceased],
        receivers=receivers,
        # Diagram has no legally confirmed refusal input. Compatibility
        # placeholders therefore stay empty instead of inferring a refusal.
        legal_refusal_people=[],
        case_type=case_type,
        has_diagram_state=built.has_diagram_state,
        signers=signers,
        authorized_people=_authorized_people(persons),
    )


def _authorized_people(people: list[WordPerson]) -> list[WordPerson]:
    """Người có vai trò "người được ủy quyền" ĐÃ LƯU trên hồ sơ.

    NAIA-12 quyết định: KHÔNG suy từ vị trí dòng vì luật không cho phép; người
    được ủy quyền chỉ tính khi tác nghiệp đã gán vai trò. Nguồn xác nhận là
    ``InheritanceParticipant.vai_tro`` (merge vào Stage person bởi
    ``_merge_participant_roles``) hoặc ``role`` trên node legacy — node Diagram
    v3 không lưu role tự do và người được ủy quyền là đại diện, có thể hợp lệ
    mà KHÔNG nằm trên sơ đồ thừa kế, nên ``is_in_diagram`` không được lọc ở
    đây. Người có node bị ẩn/xoá vẫn loại (đã bị gỡ khỏi sơ đồ).
    """
    def _order(person: WordPerson) -> tuple[int, str]:
        key = _role_key(getattr(person, "role", ""))
        suffix = key[len("nguoi_duoc_uy_quyen"):].strip()
        return (int(suffix) if suffix.isdigit() else 0, getattr(person, "ho_ten", ""))

    found = [
        p for p in people
        if _role_key(getattr(p, "role", "")).startswith("nguoi_duoc_uy_quyen")
        and not p.is_hidden and not p.is_deleted
    ]
    return sorted(found, key=_order)


# Mọi dạng registry đặt slot NGAY SAU tên nhóm (hoặc sau từ khoá "inline"):
# [Người được ủy quyền 2 - ...], [Dòng/Họ tên/... người được ủy quyền 2],
# [Người được ủy quyền inline 2]. Token không mang số ([Danh sách ...],
# [Danh sách ... inline]) -> slot 1 ở nhánh elif bên dưới.
_AUTHORIZED_TOKEN_RE = re.compile(
    r"nguoi duoc uy quyen(?:\s+inline)?\s*(\d+)")


def required_authorized_slots(doc: Any) -> set[int]:
    """Chỉ số N của nhóm "người được ủy quyền" mà mẫu THỰC SỰ dùng.

    ``[Dòng người được ủy quyền 2]``/``[Họ tên người được ủy quyền 2]``/
    ``[Người được ủy quyền inline 2]`` -> {2}; token không mang số
    (``[Danh sách người được ủy quyền]``) cần ít nhất một người -> {1}.
    Batch export dùng danh sách này để báo thiếu dữ liệu bắt buộc
    thay vì để slot trống trôi thành văn bản rỗng (NAIA-12)."""
    slots: set[int] = set()
    for token in _collect_placeholder_occurrences(doc):
        normalized = _normalize_token(token[1:-1])
        match = _AUTHORIZED_TOKEN_RE.search(normalized)
        if match:
            slots.add(int(match.group(1)))
        elif "nguoi duoc uy quyen" in normalized:
            slots.add(1)  # token không mang số -> cần ít nhất người 1
    return slots


def _spouses_of(persons: list[WordPerson], target: WordPerson) -> list[WordPerson]:
    return [p for p in _active_persons(persons) if _role_key(p.role) == "vo_chong" and p.id != target.id]


def _title_for_person(person: WordPerson) -> str:
    if person.gender_normalized == "nam":
        return "ông"
    if person.gender_normalized == "nữ":
        return "bà"
    return "ông/bà"


def _format_name_list(people: list[WordPerson]) -> str:
    if not people:
        return ""
    parts = [f"{_title_for_person(p)} {p.ho_ten}".strip() for p in people if p.ho_ten]
    if len(parts) == 1:
        return parts[0]
    if len(parts) == 2:
        return f"{parts[0]} và {parts[1]}"
    return ", ".join(parts[:-1]) + f" và {parts[-1]}"


def _format_simple_name_list(people: list[WordPerson]) -> str:
    if not people:
        return ""
    names = [p.ho_ten for p in people if p.ho_ten]
    if len(names) == 1:
        return names[0]
    if len(names) == 2:
        return f"{names[0]} và {names[1]}"
    return ", ".join(names[:-1]) + f" và {names[-1]}"


def _relation_for_person(person: WordPerson) -> str:
    if person.relation:
        return person.relation
    role_key = _role_key(person.role)
    return ROLE_LABELS.get(role_key, _safe_text(person.role))


def _status_for_person(context: WordExportContext, person: WordPerson) -> str:
    if context.case_type == "two_party":
        return {"a": "Bên A", "b": "Bên B"}.get(person.side, "")
    person_id = str(person.id)
    if any(str(item.id) == person_id for item in context.receivers):
        return "Người nhận"
    if any(str(item.id) == person_id for item in context.deceased_landowners):
        return "Chủ đất chết"
    if any(str(item.id) == person_id for item in context.living_landowners):
        return "Chủ đất sống"
    return "Người không nhận"


def _person_placeholder_values(context: WordExportContext, person: WordPerson | None) -> dict[str, str]:
    if person is None:
        return {
            "Xưng hô": "",
            "Họ tên": "",
            "Giới tính": "",
            "Ngày sinh": "",
            "Năm sinh": "",
            "Ngày chết": "",
            "Năm chết": "",
            "Loại giấy tờ": "",
            "Số giấy tờ": "",
            "Ngày cấp": "",
            "Nơi cấp": "",
            "Nhãn địa chỉ": "",
            "Địa chỉ": "",
            "Quan hệ": "",
            "Vai trò": "",
            "Trạng thái": "",
            "Tỷ lệ": "",
        }
    return {
        "Xưng hô": _title_for_person(person),
        "Họ tên": person.ho_ten,
        "Giới tính": person.gioi_tinh,
        "Ngày sinh": _fmt_date(person.ngay_sinh),
        "Năm sinh": _fmt_year(person.ngay_sinh),
        "Ngày chết": _fmt_date(person.ngay_chet),
        "Năm chết": _fmt_year(person.ngay_chet),
        "Loại giấy tờ": person.loai_giay_to,
        "Số giấy tờ": person.so_giay_to,
        "Ngày cấp": _fmt_date(person.ngay_cap),
        "Nơi cấp": person.noi_cap,
        "Nhãn địa chỉ": person.address_label,
        "Địa chỉ": person.dia_chi,
        "Quan hệ": _relation_for_person(person),
        "Vai trò": person.role,
        "Trạng thái": _status_for_person(context, person),
        "Tỷ lệ": person.share,
    }


def _person_detail_block(index: int, person: WordPerson, *, refused: bool = False) -> str:
    lines = [f"{index}. {_title_for_person(person)} {person.ho_ten}".strip()]
    if person.ngay_sinh:
        lines[0] += f"; Sinh ngày: {_fmt_birth_or_year(person.ngay_sinh)}"
    lines[0] += "."

    document_parts = []
    if person.loai_giay_to:
        document_parts.append(person.loai_giay_to)
    if person.so_giay_to:
        document_parts.append(f"số: {person.so_giay_to}")
    document_line = " ".join(document_parts)
    if person.noi_cap:
        document_line += f" do {person.noi_cap} cấp"
    if person.ngay_cap:
        document_line += f" ngày {_fmt_date(person.ngay_cap)}"
    if document_line:
        lines.append(document_line + ".")

    if person.dia_chi:
        lines.append(f"{person.address_label}: {person.dia_chi}.")
    relation = _relation_for_person(person)
    if relation:
        lines.append(f"Là {relation}.")
    if refused:
        lines.append(
            f"{_title_for_person(person).capitalize()} {person.ho_ten} đã từ chối di sản theo Văn bản "
            "từ chối nhận di sản số ......................... ."
        )
    return "\n".join(lines)


def _add_word_person_group(
    mapping: dict[str, str],
    context: WordExportContext,
    group_name: str,
    people: list[WordPerson],
    *,
    limit: int = 20,
) -> None:
    group_alias = group_name.lower()
    slot_people = None
    if group_name in ("Bên A", "Bên B"):
        offset = 0 if group_name == "Bên A" else 15
        slot_people = {p.slot - offset: p for p in people}
    for index in range(1, limit + 1):
        person = (slot_people.get(index) if slot_people is not None
                  else people[index - 1] if index <= len(people) else None)
        values = _person_placeholder_values(context, person)
        for field, value in values.items():
            mapping[f"[{group_name} {index} - {field}]"] = value
            mapping[f"[{field} {group_alias} {index}]"] = value
        if group_name == "Hàng thừa kế":
            note = _combined_heir_note(context, person) if person else ""
            mapping[f"[Hàng thừa kế {index} - Ghi chú]"] = note
            mapping[f"[Ghi chú hàng thừa kế {index}]"] = note

        block = ""
        inline = ""
        if person is not None:
            block = _person_detail_block(index, person, refused=group_name == "Người từ chối")
            inline_name = f"{_title_for_person(person)} {person.ho_ten}".strip()
            inline_index = people.index(person) + 1 if slot_people is not None else index
            if inline_index == 1:
                inline = inline_name
            elif inline_index == len(people):
                inline = f" và {inline_name}"
            else:
                inline = f", {inline_name}"
        mapping[f"[Dòng {group_alias} {index}]"] = block
        mapping[f"[{group_name} inline {index}]"] = inline

    mapping[f"[Danh sách {group_alias} inline]"] = _format_name_list(people)
    mapping[f"[Danh sách {group_alias}]"] = _format_simple_name_list(people)


def _deceased_landowner_clauses(people: list[WordPerson]) -> str:
    """Dòng khai tử chủ đất chết — ghép từ field đã commit của người
    (MIN-141 đợt 3): loại giấy tờ/nơi cấp lấy bằng chứng xác nhận trong
    snapshot; phần nào chưa có thì lược bỏ, KHÔNG gán cứng "Trích lục
    khai tử (Bản sao)" hay suy tên UBND xã."""
    lines = []
    for p in people:
        title = _title_for_person(p)
        born = _fmt_birth_or_year(p.ngay_sinh)
        died = _fmt_date(p.ngay_chet)
        doc_no = _safe_text(p.so_giay_to)
        doc_type = _safe_text(p.loai_giay_to)
        issuer = _safe_text(p.noi_cap)
        issued = _fmt_date(p.ngay_cap)
        address = _safe_text(p.dia_chi)
        line = f"{title} {p.ho_ten}"
        if born:
            line += f"; Sinh năm: {born}"
        line += f"; chết ngày {died}"
        if doc_type or doc_no or issuer or issued:
            tail = " theo"
            tail += f" {doc_type}" if doc_type else " giấy tờ khai tử"
            if doc_no:
                tail += f" số {doc_no}"
            if issuer:
                tail += f" do {issuer}"
            if issued:
                tail += f" ký ngày {issued}"
            line += tail
        line += "."
        if address:
            line += f" Nơi chết: {address}."
        lines.append(line)
    return "\n".join(lines)


def _family_relation_clauses(
    deceased_landowners: list[WordPerson],
    persons: list[WordPerson],
    legal_refusal_people: list[WordPerson],
) -> str:
    paragraphs = []
    refused_ids = {str(person.id) for person in legal_refusal_people}
    for owner in deceased_landowners:
        title = _title_for_person(owner)
        spouses = _spouses_of(persons, owner)
        if spouses:
            spouse = spouses[0]
            spouse_title = _title_for_person(spouse)
            paragraphs.append(f"{title} {owner.ho_ten} có vợ/chồng là {spouse_title} {spouse.ho_ten}.")
        children = _children(persons)
        if children:
            child_lines = []
            for child in children:
                if str(child.id) in refused_ids:
                    child_lines.append(
                        f"{child.ho_ten} đã từ chối di sản theo Văn bản từ chối nhận di sản số ......................... ."
                    )
                else:
                    child_lines.append(child.ho_ten)
            child_text = "; ".join(child_lines)
            paragraphs.append(f"{title} {owner.ho_ten} có {len(children)} ngườ con là: {child_text}.")
    return "\n".join(paragraphs)


_SNAPSHOT_ASSET_META_KEYS = {"id", "entity_id", "row_id", "is_primary"}


def _safe_int(v: Any) -> int | None:
    if v is None or isinstance(v, bool):
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _snapshot_asset_fields(item: Any) -> bool:
    """Snapshot asset đợt 2 trở đi mang full field + land_rows canonical;
    snapshot cũ chỉ có {id, entity_id, row_id, is_primary} → fieldless."""
    return isinstance(item, dict) and any(
        key not in _SNAPSHOT_ASSET_META_KEYS for key in item)


def _snapshot_asset_view(item: dict[str, Any]) -> Any:
    """Asset snapshot trong case_state_json → view giống ORM Property cho
    các hàm placeholder (MIN-141 đợt 2 — Word đọc đúng bản đã commit của
    hồ sơ, không đổi theo master properties của hồ sơ khác).

    `land_rows` giữ nguyên dict canonical (loaidat/dientich/thoihan) —
    _land_row_values đọc được cả hai spelling; `land_rows_json` = None vì
    snapshot đã là nguồn committed."""
    ngay_cap = item.get("ngay_cap")
    if isinstance(ngay_cap, str) and ngay_cap:
        try:
            ngay_cap = date.fromisoformat(ngay_cap[:10])
        except ValueError:
            pass  # giữ chuỗi gốc — _fmt_date pass-through chuỗi
    entity_id = item.get("entity_id")
    return SimpleNamespace(
        id=entity_id if isinstance(entity_id, int)
            else _safe_int(item.get("id")),
        entity_id=entity_id,
        so_serial=item.get("so_serial"),
        so_vao_so=item.get("so_vao_so"),
        so_thua_dat=item.get("so_thua_dat"),
        so_to_ban_do=item.get("so_to_ban_do"),
        dia_chi=item.get("dia_chi"),
        loai_so=item.get("loai_so"),
        hinh_thuc_su_dung=item.get("hinh_thuc_su_dung"),
        thoi_han=item.get("thoi_han"),
        nguon_goc=item.get("nguon_goc"),
        ngay_cap=ngay_cap,
        co_quan_cap=item.get("co_quan_cap"),
        loai_dat=item.get("loai_dat"),
        dien_tich=item.get("dien_tich"),
        land_rows=(item.get("land_rows")
                   if isinstance(item.get("land_rows"), list) else []),
        land_rows_json=None,
    )


def _snapshot_person_view(customer: Any, row: dict) -> Any:
    """View Người cho slot `[Tên N]`…`[noicapN]` — đọc bằng chứng đã
    commit trong snapshot `stage[]` của chính hồ sơ thay vì master
    `customers` live (MIN-141 review: master dùng chung có thể bị hồ sơ
    khác đổi sau commit; vd loai_giay_to/ngay_cap của người chết model
    trả None nên đường live làm mất 'Giấy chứng tử' đã xác nhận).

    Per-key presence như `_snap_meta`: key có trong snapshot → dùng giá
    trị snapshot (kể cả null chủ ý → rỗng); key vắng (snapshot cũ) →
    fallback master live.

    Ngày giữ nguyên giá trị đã commit (`YYYY` | `YYYY-MM-DD`) — formatter
    `_fmt_date`/`_fmt_year`/`_fmt_birth_or_year` render theo đúng độ đầy
    đủ; không ép về `date(y,1,1)` vì sẽ làm `2010` ra `01/01/2010`.
    Fallback master là `date` → render đầy đủ dd/mm/yyyy."""

    def pick(key: str) -> Any:
        return row[key] if key in row else getattr(customer, key, None)

    return SimpleNamespace(
        id=getattr(customer, "id", None),
        ho_ten=_safe_text(pick("ho_ten")),
        gioi_tinh=_safe_text(pick("gioi_tinh")),
        ngay_sinh=pick("ngay_sinh"),
        ngay_chet=pick("ngay_chet"),
        so_giay_to=_safe_text(pick("so_giay_to")),
        ngay_cap=pick("ngay_cap"),
        noi_cap=_safe_text(pick("noi_cap")),
        dia_chi=_safe_text(pick("dia_chi")),
        loai_giay_to=_safe_text(pick("loai_giay_to")),
        loai_dia_chi=_safe_text(pick("loai_dia_chi")),
    )


def _snapshot_people_by_id(state: dict[str, Any]) -> dict[str, dict]:
    """{entity_id str → row} từ snapshot `stage[]` — đối chiếu slot
    Người legacy (Customer live) với bản đã commit của hồ sơ."""
    out: dict[str, dict] = {}
    stage = state.get("stage") if isinstance(state, dict) else None
    for item in stage or []:
        if isinstance(item, dict) and item.get("id") is not None:
            out[str(item.get("id"))] = item
    return out


def _live_props_by_id(case: Any) -> dict[int, Any]:
    """{property_id: Property live} — fallback cho snapshot item chỉ có
    pointer (payload trước đợt 2)."""
    out: dict[int, Any] = {}
    for link in getattr(case, "property_links", None) or []:
        prop = getattr(link, "property", None)
        pid = getattr(link, "property_id", None)
        if prop is not None and isinstance(pid, int):
            out[pid] = prop
    ts = getattr(case, "tai_san", None)
    ts_id = getattr(ts, "id", None)
    if ts is not None and isinstance(ts_id, int):
        out.setdefault(ts_id, ts)
    return out


def _get_property_list(case: Any) -> list[Any]:
    """Tài sản để render — snapshot đã commit trong case_state_json trước
    (MIN-141 đợt 2): mỗi hồ sơ xuất đúng bản commit của mình, kể cả khi
    master properties bị hồ sơ khác dùng chung sửa. Item snapshot chỉ có
    pointer (payload cũ) → tra master live theo entity_id; không có
    snapshot assets → đường live cũ (links/tai_san)."""
    snap_assets = _load_case_state(case).get("assets")
    if isinstance(snap_assets, list) and snap_assets:
        live = _live_props_by_id(case)
        views = []
        for item in snap_assets:
            if _snapshot_asset_fields(item):
                views.append(_snapshot_asset_view(item))
                continue
            pid = None
            if isinstance(item, dict):
                pid = item.get("entity_id")
                if not isinstance(pid, int) or isinstance(pid, bool):
                    pid = _safe_int(item.get("id"))
            prop = live.get(pid) if isinstance(pid, int) else None
            if prop is not None:
                views.append(prop)
        if views:
            return views
    links = getattr(case, "property_links", None)
    if links:
        ordered = sorted(links, key=lambda x: (not bool(getattr(x, "is_primary", False)), getattr(x, "id", 0)))
        props = [getattr(link, "property", None) for link in ordered]
        props = [p for p in props if p is not None]
        if props:
            return props
    ts = getattr(case, "tai_san", None)
    if ts is not None:
        return [ts]
    return []


def _property_description(properties: list[Any]) -> str:
    if not properties:
        return ""
    paragraphs = ["Các quyền sử dụng đất như sau:"] if len(properties) > 1 else []
    for idx, prop in enumerate(properties, start=1):
        loai_so = _safe_text(getattr(prop, "loai_so", "")) or "Giấy chứng nhận quyền sử dụng đất"
        serial = _serial_text(getattr(prop, "so_serial", ""))
        so_vao_so = _safe_text(getattr(prop, "so_vao_so", ""))
        so_thua = _safe_text(getattr(prop, "so_thua_dat", ""))
        so_to = _safe_text(getattr(prop, "so_to_ban_do", ""))
        dia_chi = _safe_text(getattr(prop, "dia_chi", ""))
        co_quan_cap = _safe_text(getattr(prop, "co_quan_cap", ""))
        ngay_cap = _fmt_date(getattr(prop, "ngay_cap", None))

        header = f"{idx}. Quyền sử dụng đất tại: {dia_chi} theo {loai_so} số: {serial}; Số vào sổ cấp GCN: {so_vao_so} do {co_quan_cap} cấp ngày {ngay_cap}."
        paragraphs.append(header)

        detail = f"Thửa đất số: {so_thua}, tờ bản đồ số: {so_to}."
        paragraphs.append(detail)

        land_rows = _land_rows_for(prop)

        total_area = 0.0
        for row in land_rows:
            try:
                total_area += float(row.get("dien_tich") or 0)
            except (TypeError, ValueError):
                pass
        area_str = f"{total_area:g}" if total_area else _safe_text(getattr(prop, "dien_tich", ""))
        paragraphs.append(f"Diện tích: {area_str} m2")
        paragraphs.append(f"Hình thức sử dụng: {_safe_text(getattr(prop, 'hinh_thuc_su_dung', ''))}")

        land_type_lines = []
        for m_idx, row in enumerate(land_rows, start=1):
            loai_dat = _safe_text(row.get("loai_dat", ""))
            dien_tich = _safe_text(row.get("dien_tich", ""))
            thoi_han = _safe_text(row.get("thoi_han", ""))
            if not loai_dat and not dien_tich:
                continue
            land_type_lines.append(f"{idx}.{m_idx}. {loai_dat}: {dien_tich} m2 ({thoi_han})")
        if land_type_lines:
            paragraphs.append("\n".join(land_type_lines))

        paragraphs.append(f"Thời hạn sử dụng: {_safe_text(getattr(prop, 'thoi_han', ''))}")
        paragraphs.append(f"Nguồn gốc sử dụng đất: {_safe_text(getattr(prop, 'nguon_goc', ''))}.")
    return "\n".join(paragraphs)


def _division_clause(context: WordExportContext) -> str:
    if not context.receivers:
        raise WordExportValidationError("Không có người nhận để xuất văn bản phân chia.")

    deceased_names = _format_name_list(context.deceased_landowners)
    receiver_names = _format_name_list(context.receivers)
    simple_receiver_names = _format_simple_name_list(context.receivers)
    paragraphs = [
        f"Chúng tôi gồm: {receiver_names} thống nhất phân chia di sản của "
        f"{deceased_names} như sau:"
    ]
    paragraphs.append(
        f"{simple_receiver_names} đồng ý nhận phần di sản theo nội dung thỏa thuận "
        "phân chia di sản thừa kế ở trên."
    )
    return "\n".join(paragraphs)


def _heir_note(context: WordExportContext, person: WordPerson) -> str:
    return _status_for_person(context, person)


def _role_note(person: WordPerson) -> str:
    role = _role_key(person.role)
    if role == "vo_chong":
        if person.gender_normalized == "nam":
            return "Là chồng"
        return "Là vợ"
    if role == "con":
        return "Là con"
    return ""


def _combined_heir_note(context: WordExportContext, person: WordPerson) -> str:
    parts = [p for p in [_role_note(person), _heir_note(context, person)] if p]
    return "; ".join(parts)


# Nhóm dữ liệu → gate `word.*` mà việc mẫu dùng token của nhóm kéo theo.
# Key nhóm do CHÍNH `_add_word_person_group` sinh (nguồn duy nhất định nghĩa
# "token của nhóm"): phân loại so khớp chính xác key đã chuẩn hoá — không
# đoán bằng chuỗi con, nên token metadata/ủy quyền (vd "[Người nhận ủy
# quyền]") không kéo nhầm gate Người nhận, còn alias không dấu/thiếu dấu của
# key nhóm vẫn nhận đúng như consumer `_lookup_placeholder`. Nhóm "Người",
# "Người từ chối", "Hàng thừa kế" và các nhóm vai trò chỉ có giới hạn kỹ
# thuật — không nằm trong bảng này.
_GROUP_GATE_NEEDS: tuple[tuple[str, int, tuple[str, ...]], ...] = (
    ("Chủ đất sống", 20, ("landowner",)),
    ("Chủ đất chết", 20, ("landowner", "deceased_landowner")),
    ("Người nhận", 20, ("receiver",)),
    ("Bên A", 15, ("two_party_slots",)),
    ("Bên B", 15, ("two_party_slots",)),
)

# Token ghế hai bên ghi TRỰC TIẾP trong `_add_two_party_placeholders` (không
# qua `_add_word_person_group`): tên từng bên cũng đọc từ slot diagram nên
# cùng cần `two_party_slots`.
_TWO_PARTY_SLOT_EXTRA_KEYS = (
    "Bên A - Họ tên", "Bên B - Họ tên", "Tên bên A", "Tên bên B",
)


@lru_cache(maxsize=None)
def _group_key_registry(group_name: str, limit: int) -> frozenset[str]:
    """Tập key CHUẨN HOÁ mà `_add_word_person_group` sinh cho một nhóm.

    Gọi chính hàm ghi mapping với danh sách rỗng (không chạm context) để
    registry luôn khớp định nghĩa thật — không chép lại công thức key."""
    probe: dict[str, str] = {}
    _add_word_person_group(probe, None, group_name, [], limit=limit)
    return frozenset(_normalize_token(key[1:-1]) for key in probe)


def template_group_needs(doc: Any) -> set[str]:
    """Nhóm dữ liệu mẫu THỰC SỰ dùng — quyết định gate nhóm (owner 15/10/2026).

    So khớp token mẫu với đúng registry key engine sinh cho từng nhóm
    (`_group_key_registry` + token ghế hai bên viết trực tiếp), sau chuẩn
    hoá không dấu như consumer `_lookup_placeholder` — không đoán theo chuỗi
    con. Trả về tập con của {'landowner', 'deceased_landowner', 'receiver',
    'two_party_slots'}; nhóm ủy quyền đã có resolver riêng
    (`required_authorized_slots`, NAIA-12)."""
    needs: set[str] = set()
    extra_keys = {_normalize_token(key) for key in _TWO_PARTY_SLOT_EXTRA_KEYS}
    for token in _collect_placeholder_occurrences(doc):
        normalized = _normalize_token(token[1:-1])
        if normalized in extra_keys:
            needs.add("two_party_slots")
            continue
        for group_name, limit, group_needs in _GROUP_GATE_NEEDS:
            if normalized in _group_key_registry(group_name, limit):
                needs.update(group_needs)
    return needs


def word_block_reason(context: WordExportContext, *, static: bool = False,
                      needed: set[str] | None = None) -> str | None:
    """Contract `word.*` block_reason đầu tiên cho một context đã build.

    Cùng thứ tự kiểm tra với `_add_block_placeholders` bên dưới (contract
    notary.case-drafting.v2 §13.10 mapping). Trả None khi dữ liệu đủ để xuất.
    Dùng bởi `services.word_batch_export` cho word_export_options +
    per-document failure của word_export_batch (MIN-110).

    `needed` = tập nhu cầu của CHÍNH mẫu đang xét
    (`template_group_needs`): gate nhóm `no_landowner`/`no_deceased_landowner`/
    `no_receiver` (thừa kế) và `no_diagram` hai bên chỉ áp khi mẫu thực sự
    dùng nhóm đó (owner 15/10/2026). `needed=None` chỉ dành cho caller
    KHÔNG có mẫu (vd đối chiếu word gốc) — giữ đủ gate cũ.

    NAIA-8: readiness theo loại hồ sơ — two_party không áp điều kiện
    riêng thừa kế (chủ đất/người nhận); `[Người N]` theo Stage nên giới
    hạn người đếm trên Stage (`persons`), không phải node Diagram."""
    if static:
        if not context.persons:
            return "word.no_people"
        if not context.assets:
            return "word.no_assets"
        if len(context.assets) > MAX_WORD_ASSETS:
            return "word.too_many_assets"
        if len(context.persons) > (30 if context.case_type == "two_party" else 20):
            return "word.too_many_people"
        if context.case_type == "two_party" \
                and (needed is None or "two_party_slots" in needed) \
                and not context.has_diagram_state:
            return "word.no_diagram"
        return None
    if context.case_type == "two_party":
        if not context.persons:
            return "word.no_people"
        if (needed is None or "two_party_slots" in needed) \
                and not context.has_diagram_state:
            return "word.no_diagram"
        if not context.assets:
            return "word.no_assets"
        if len(context.assets) > MAX_WORD_ASSETS:
            return "word.too_many_assets"
        if len(context.persons) > 30:
            return "word.too_many_people"
        return None
    if not context.assets:
        return "word.no_assets"
    if len(context.assets) > MAX_WORD_ASSETS:
        return "word.too_many_assets"
    # Gate nhóm chỉ khi mẫu THỰC SỰ dùng nhóm (needed=None = caller không
    # có mẫu → giữ đủ gate cũ; caller có mẫu luôn truyền needs thật).
    if (needed is None or "landowner" in needed) and not context.landowners:
        return "word.no_landowner"
    if (needed is None or "deceased_landowner" in needed) \
            and not context.deceased_landowners:
        return "word.no_deceased_landowner"
    if (needed is None or "receiver" in needed) and not context.receivers:
        return "word.no_receiver"
    if len(context.persons) > 20:
        return "word.too_many_people"
    if len(context.signers) > 20:
        return "word.too_many_signers"
    return None


def _add_two_party_placeholders(mapping: dict[str, str],
                                context: WordExportContext,
                                *, needed: set[str] | None = None) -> None:
    """Block placeholder cho hồ sơ hai bên (§13.5 + NAIA-8).

    `[Người N - ...]` = toàn bộ Stage theo thứ tự đã lưu (kể cả người
    chưa gán slot). `[Bên A N - ...]`/`[Bên B N - ...]` theo slot
    p1..p15 / p16..p30 của diagram.state đã lưu. Nhóm riêng thừa kế
    (chủ đất, người nhận, hàng thừa kế, đoạn phân chia) resolve rỗng —
    template cũ dùng chung không rớt unresolved_placeholders.

    `needed` (`template_group_needs`): chỉ đòi diagram.state đã lưu khi
    mẫu THỰC SỰ dùng ghế hai bên; `None` = caller không có mẫu → giữ
    gate cũ (luôn đòi)."""
    if not context.persons:
        raise WordExportValidationError("Stage chưa có người nào.")
    if (needed is None or "two_party_slots" in needed) \
            and not context.has_diagram_state:
        raise WordExportValidationError(
            "Chưa có sơ đồ hai bên đã lưu (diagram.state.nodes).")
    if not context.assets:
        raise WordExportValidationError("Không có tài sản để xuất Word.")
    if len(context.assets) > MAX_WORD_ASSETS:
        raise WordExportValidationError(
            f"Vượt quá {MAX_WORD_ASSETS} tài sản (hiện có {len(context.assets)}). "
            "Vui lòng giảm số lượng tài sản trước khi xuất Word."
        )
    if len(context.persons) > 30:
        raise WordExportValidationError(
            f"Stage vượt quá 30 người (hiện có {len(context.persons)})."
        )

    _add_word_person_group(mapping, context, "Người",
                           context.persons, limit=30)
    _add_word_person_group(mapping, context, "Bên A",
                           context.party_a, limit=15)
    _add_word_person_group(mapping, context, "Bên B",
                           context.party_b, limit=15)
    for group_name in ("Chủ đất sống", "Chủ đất chết", "Người nhận",
                       "Người từ chối", "Hàng thừa kế"):
        _add_word_person_group(mapping, context, group_name, [])

    mapping["[Bên A - Họ tên]"] = _format_name_list(context.party_a)
    mapping["[Bên B - Họ tên]"] = _format_name_list(context.party_b)
    mapping["[Tên bên A]"] = _format_simple_name_list(context.party_a)
    mapping["[Tên bên B]"] = _format_simple_name_list(context.party_b)

    # Block/đoạn riêng thừa kế → rỗng toàn bộ (spec word-output: tầng 2
    # trả rỗng khi không áp dụng, tránh dấu câu rác).
    mapping["[Cụm chủ đất chết]"] = ""
    mapping["[Cụm người để lại di sản]"] = ""
    mapping["[Cụm ngườ để lại di sản]"] = ""
    for index in range(1, 3):
        mapping[f"[Chủ đất chết mở đầu {index}]"] = ""
        mapping[f"[Dòng khai tử chủ đất chết {index}]"] = ""
        mapping[f"[Cụm ngày chết chủ đất chết {index}]"] = ""
    mapping["[Nối chủ đất chết 2]"] = ""
    mapping["[Cụm UBND niêm yết]"] = ""
    for index in range(1, 21):
        mapping[f"[Dòng chủ đất sống tặng cho {index}]"] = ""
    mapping["[Danh sách người nhận và chưa chọn]"] = ""
    mapping["[Danh sách ngườ nhận và chưa chọn]"] = ""
    mapping["[Đoạn người chết là chủ đất]"] = ""
    mapping["[Đoạn ngườ chết là chủ đất]"] = ""
    mapping["[Đoạn quan hệ gia đình]"] = ""
    mapping["[Đoạn phân chia di sản]"] = ""
    mapping["[Danh sách hàng thừa kế]"] = ""
    for idx in range(1, 21):
        mapping[f"[STT hàng thừa kế {idx}]"] = ""
        mapping[f"[Họ tên hàng thừa kế {idx}]"] = ""
        mapping[f"[Ngày sinh hàng thừa kế {idx}]"] = ""
        mapping[f"[Năm sinh hàng thừa kế {idx}]"] = ""
        mapping[f"[Địa chỉ hàng thừa kế {idx}]"] = ""
        mapping[f"[Ghi chú hàng thừa kế {idx}]"] = ""
        mapping[f"[Dòng hàng thừa kế {idx}]"] = ""

    mapping["[Đoạn mô tả di sản]"] = _property_description(context.assets)

    signers = context.signers
    if len(signers) > 30:
        raise WordExportValidationError(
            f"Danh sách người ký vượt quá 30 người (hiện có {len(signers)})."
        )
    mapping["[Danh sách người ký]"] = _format_simple_name_list(signers)
    mapping["[Danh sách ngườ ký]"] = mapping["[Danh sách người ký]"]
    for idx in range(1, 31):
        name = signers[idx - 1].ho_ten if idx <= len(signers) else ""
        mapping[f"[Họ tên người ký {idx}]"] = name
        mapping[f"[Họ tên ngườ ký {idx}]"] = name


def _add_block_placeholders(mapping: dict[str, str], context: WordExportContext,
                            *, static: bool = False,
                            needed: set[str] | None = None) -> None:
    """Block placeholder thừa kế/hai bên — mirror `word_block_reason`.

    `needed` (`template_group_needs`) gate các raise nhóm theo đúng nhu
    cầu mẫu; `None` = caller không có mẫu → giữ đủ gate cũ."""
    if context.case_type == "two_party":
        _add_two_party_placeholders(mapping, context, needed=needed)
        return
    if not context.assets:
        raise WordExportValidationError("Không có tài sản để xuất Word.")
    if len(context.assets) > MAX_WORD_ASSETS:
        raise WordExportValidationError(
            f"Vượt quá {MAX_WORD_ASSETS} tài sản (hiện có {len(context.assets)}). "
            "Vui lòng giảm số lượng tài sản trước khi xuất Word."
        )
    if not static and (needed is None or "landowner" in needed) \
            and not context.landowners:
        raise WordExportValidationError("Không xác định được chủ đất trên Diagram.")
    if not static and (needed is None or "deceased_landowner" in needed) \
            and not context.deceased_landowners:
        raise WordExportValidationError("Không xác định được chủ đất đã chết để xuất văn bản.")
    if not static and (needed is None or "receiver" in needed) \
            and not context.receivers:
        raise WordExportValidationError("Không có người nhận để xuất văn bản phân chia.")
    if len(context.persons) > 20:
        raise WordExportValidationError(
            f"Danh sách Stage vượt quá 20 người (hiện có {len(context.persons)})."
        )

    heir_rows = [person for person in context.all_people if not person.is_land_owner]
    groups = {
        # [Người N] = thứ tự Stage đã lưu, gồm cả người chưa có node
        # (NAIA-8) — đổi thứ tự/xóa node không đổi số Người.
        "Người": context.persons,
        "Chủ đất sống": context.living_landowners,
        "Chủ đất chết": context.deceased_landowners,
        "Người nhận": context.receivers,
        "Người từ chối": context.legal_refusal_people,
        # Người được ủy quyền: chỉ lấy từ vai trò đã xác nhận/lưu trên hồ sơ
        # (participants.vai_tro), không suy từ vị trí dòng và không bắt buộc
        # node Diagram — đại diện ủy quyền hợp lệ có thể không nằm trên sơ đồ
        # thừa kế (quyết định Owner trong NAIA-12).
        "Người được ủy quyền": context.authorized_people,
        "Hàng thừa kế": heir_rows,
    }
    for group_name, people in groups.items():
        _add_word_person_group(mapping, context, group_name, people)

    deceased_cluster = _format_name_list(context.deceased_landowners)
    mapping["[Cụm chủ đất chết]"] = deceased_cluster
    mapping["[Cụm người để lại di sản]"] = deceased_cluster
    mapping["[Cụm ngườ để lại di sản]"] = deceased_cluster
    for index in range(1, 3):
        person = context.deceased_landowners[index - 1] if index <= len(context.deceased_landowners) else None
        opening = f"{_title_for_person(person)} {person.ho_ten}".strip() if person else ""
        mapping[f"[Chủ đất chết mở đầu {index}]"] = opening
        mapping[f"[Dòng khai tử chủ đất chết {index}]"] = (
            _deceased_landowner_clauses([person]) if person else ""
        )
        mapping[f"[Cụm ngày chết chủ đất chết {index}]"] = (
            f"; Chết ngày: {_fmt_date(person.ngay_chet)}" if person and person.ngay_chet else ""
        )
    mapping["[Nối chủ đất chết 2]"] = " và " if len(context.deceased_landowners) >= 2 else ""
    mapping["[Cụm UBND niêm yết]"] = f"UBND xã {mapping.get('[Niêm Yết]', '')}".strip()

    for index in range(1, 21):
        mapping[f"[Dòng chủ đất sống tặng cho {index}]"] = ""

    # Alias cũ: không còn nhóm "chưa chọn", giá trị chỉ còn người nhận.
    mapping["[Danh sách người nhận và chưa chọn]"] = _format_name_list(context.receivers)
    mapping["[Danh sách ngườ nhận và chưa chọn]"] = mapping["[Danh sách người nhận và chưa chọn]"]

    mapping["[Đoạn người chết là chủ đất]"] = _deceased_landowner_clauses(context.deceased_landowners)
    mapping["[Đoạn ngườ chết là chủ đất]"] = mapping["[Đoạn người chết là chủ đất]"]

    mapping["[Đoạn quan hệ gia đình]"] = _family_relation_clauses(
        context.deceased_landowners,
        context.all_people,
        context.legal_refusal_people,
    )

    mapping["[Đoạn mô tả di sản]"] = _property_description(context.assets)
    mapping["[Đoạn phân chia di sản]"] = (
        _division_clause(context) if context.receivers else "")

    signers = context.signers
    mapping["[Danh sách người ký]"] = _format_simple_name_list(signers)
    mapping["[Danh sách ngườ ký]"] = mapping["[Danh sách người ký]"]
    mapping["[Danh sách hàng thừa kế]"] = _format_simple_name_list(heir_rows)

    for idx in range(1, 21):
        if idx <= len(heir_rows):
            p = heir_rows[idx - 1]
            mapping[f"[STT hàng thừa kế {idx}]"] = str(idx)
            mapping[f"[Họ tên hàng thừa kế {idx}]"] = p.ho_ten
            mapping[f"[Ngày sinh hàng thừa kế {idx}]"] = _fmt_birth_or_year(p.ngay_sinh)
            mapping[f"[Năm sinh hàng thừa kế {idx}]"] = _fmt_year(p.ngay_sinh)
            mapping[f"[Địa chỉ hàng thừa kế {idx}]"] = p.dia_chi
            mapping[f"[Ghi chú hàng thừa kế {idx}]"] = _combined_heir_note(context, p)
            mapping[f"[Dòng hàng thừa kế {idx}]"] = (
                f"{idx}. {p.ho_ten} - {_fmt_birth_or_year(p.ngay_sinh)} - {p.dia_chi} - "
                f"{_combined_heir_note(context, p)}"
            ).strip(" -")
        else:
            mapping[f"[STT hàng thừa kế {idx}]"] = ""
            mapping[f"[Họ tên hàng thừa kế {idx}]"] = ""
            mapping[f"[Ngày sinh hàng thừa kế {idx}]"] = ""
            mapping[f"[Năm sinh hàng thừa kế {idx}]"] = ""
            mapping[f"[Địa chỉ hàng thừa kế {idx}]"] = ""
            mapping[f"[Ghi chú hàng thừa kế {idx}]"] = ""
            mapping[f"[Dòng hàng thừa kế {idx}]"] = ""

    if len(signers) > 20:
        raise WordExportValidationError(
            f"Danh sách người ký vượt quá 20 người (hiện có {len(signers)})."
        )
    for idx in range(1, 21):
        if idx <= len(signers):
            mapping[f"[Họ tên người ký {idx}]"] = signers[idx - 1].ho_ten
        else:
            mapping[f"[Họ tên người ký {idx}]"] = ""
        mapping[f"[Họ tên ngườ ký {idx}]"] = mapping[f"[Họ tên người ký {idx}]"]


def _add_property_placeholders(mapping: dict[str, str], properties: list[Any]) -> None:
    asset_fields = {
        "Địa chỉ": "Địa chỉ đất",
        "Loại sổ": "Loại sổ",
        "Serial": "Serial",
        "Số vào sổ": "Số vào sổ",
        "Số thửa": "Số thửa",
        "Số tờ": "Số tờ",
        "Diện tích": "Diện tích",
        "Diện tích chữ": "Diện tích chữ",
        "Hình thức sử dụng": "Hình thức sử dụng",
        "Mục đích sử dụng": "Mục đích sử dụng",
        "Thời hạn": "Thời hạn",
        "Nguồn gốc": "Nguồn gốc",
        "Ngày cấp sổ": "Ngày cấp sổ",
        "Cơ quan cấp sổ": "Cơ quan cấp sổ",
    }

    for idx in range(1, MAX_WORD_ASSETS + 1):
        prop = properties[idx - 1] if idx <= len(properties) else None
        land_rows: list[dict[str, Any]] = []
        values = {field: "" for field in asset_fields}
        if prop is not None:
            land_rows = _land_rows_for(prop)
            total_area = 0.0
            for row in land_rows:
                try:
                    total_area += float(row.get("dien_tich") or 0)
                except (TypeError, ValueError):
                    pass
            area_str = f"{total_area:g}" if total_area else _safe_text(getattr(prop, "dien_tich", ""))
            purpose = "; ".join(
                f"{_safe_text(row.get('loai_dat'))}: {_safe_text(row.get('dien_tich'))} m2"
                for row in land_rows
                if _safe_text(row.get("loai_dat")) or _safe_text(row.get("dien_tich"))
            )
            values = {
                "Địa chỉ": _safe_text(getattr(prop, "dia_chi", "")),
                "Loại sổ": _safe_text(getattr(prop, "loai_so", "")) or "Giấy chứng nhận quyền sử dụng đất",
                "Serial": _serial_text(getattr(prop, "so_serial", "")),
                "Số vào sổ": _safe_text(getattr(prop, "so_vao_so", "")),
                "Số thửa": _safe_text(getattr(prop, "so_thua_dat", "")),
                "Số tờ": _safe_text(getattr(prop, "so_to_ban_do", "")),
                "Diện tích": area_str,
                "Diện tích chữ": _so_thanh_chu(total_area).capitalize() if total_area else "",
                "Hình thức sử dụng": _safe_text(getattr(prop, "hinh_thuc_su_dung", "")),
                "Mục đích sử dụng": purpose,
                "Thời hạn": _safe_text(getattr(prop, "thoi_han", "")),
                "Nguồn gốc": _safe_text(getattr(prop, "nguon_goc", "")),
                "Ngày cấp sổ": _fmt_date(getattr(prop, "ngay_cap", None)),
                "Cơ quan cấp sổ": _safe_text(getattr(prop, "co_quan_cap", "")),
            }

            # [ONT]/[CLN]/[NTS]/[LUC] là diện tích theo VỊ TRÍ dòng loại đất 1..4
            # (nhãn "Loại đất N: <mã>" in cứng trong mẫu), hậu tố " idx" = tài sản.
            # Dòng ở vị trí đó vắng -> rỗng; TÀI SẢN không tồn tại -> không map,
            # để export báo placeholder sót thay vì xuất thầm thiếu dữ liệu.
            for code, position in LAND_ROW_SLOTS.items():
                row_area = _land_row_area(land_rows, position)
                mapping[f"[{code} {idx}]"] = row_area
                if idx == 1:
                    mapping[f"[{code}]"] = row_area

        for field, legacy_name in asset_fields.items():
            value = values[field]
            mapping[f"[Tài sản {idx} - {field}]"] = value
            mapping[f"[{legacy_name} {idx}]"] = value
            mapping[f"[{legacy_name} tài sản {idx}]"] = value
            if idx == 1:
                mapping[f"[Tài sản - {field}]"] = value

        # Loại đất CẤP TÀI SẢN (prop.loai_dat — nhãn "Loại đất/Mục đích sử
        # dụng" tổng của thửa) emit riêng, KHÔNG qua asset_fields vì alias
        # "[Loại đất N]" đã mang nghĩa "dòng loại đất N của tài sản 1" —
        # ghi đè sẽ đổi ngữ nghĩa token cũ (quyết định Owner, review NAIA-12).
        prop_loai_dat = _safe_text(getattr(prop, "loai_dat", "")) if prop else ""
        mapping[f"[Tài sản {idx} - Loại đất]"] = prop_loai_dat
        mapping[f"[Loại đất tài sản {idx}]"] = prop_loai_dat
        if idx == 1:
            mapping["[Tài sản - Loại đất]"] = prop_loai_dat

        if prop is None:
            mapping[f"[Dòng tài sản {idx}]"] = ""
            mapping[f"[Dòng thửa đất {idx}]"] = ""
            mapping[f"[Dòng diện tích {idx}]"] = ""
            mapping[f"[Dòng hình thức sử dụng {idx}]"] = ""
            mapping[f"[Dòng mục đích sử dụng {idx}]"] = ""
            mapping[f"[Dòng thời hạn {idx}]"] = ""
            mapping[f"[Dòng nguồn gốc {idx}]"] = ""
        else:
            mapping[f"[Dòng tài sản {idx}]"] = (
                f"{idx}. Quyền sử dụng đất tại: {values['Địa chỉ']} theo {values['Loại sổ']} số: "
                f"{values['Serial']}; Số vào sổ cấp GCN: {values['Số vào sổ']} do "
                f"{values['Cơ quan cấp sổ']} cấp ngày {values['Ngày cấp sổ']}."
            )
            mapping[f"[Dòng thửa đất {idx}]"] = (
                f"Thửa đất số: {values['Số thửa']}, tờ bản đồ số: {values['Số tờ']}."
            )
            mapping[f"[Dòng diện tích {idx}]"] = f"Diện tích: {values['Diện tích']} m2."
            mapping[f"[Dòng hình thức sử dụng {idx}]"] = (
                f"Hình thức sử dụng: {values['Hình thức sử dụng']}."
            )
            mapping[f"[Dòng mục đích sử dụng {idx}]"] = (
                f"Mục đích sử dụng: {values['Mục đích sử dụng']}."
            )
            mapping[f"[Dòng thời hạn {idx}]"] = f"Thời hạn sử dụng: {values['Thời hạn']}."
            mapping[f"[Dòng nguồn gốc {idx}]"] = (
                f"Nguồn gốc sử dụng đất: {values['Nguồn gốc']}."
            )

        for line_name in ("thửa đất", "diện tích", "hình thức sử dụng", "mục đích sử dụng", "thời hạn", "nguồn gốc"):
            canonical = f"[Dòng {line_name} {idx}]"
            mapping[f"[Dòng {line_name} tài sản {idx}]"] = mapping[canonical]

        for m_idx in range(1, 11):
            row = land_rows[m_idx - 1] if m_idx <= len(land_rows) else None
            land_type, area, term = _land_row_values(row)
            mapping[f"[Loại đất {idx}.{m_idx} - Loại đất]"] = land_type
            mapping[f"[Loại đất {idx}.{m_idx} - Diện tích]"] = area
            mapping[f"[Loại đất {idx}.{m_idx} - Thời hạn]"] = term
            mapping[f"[Loại đất tài sản {idx}.{m_idx}]"] = land_type
            mapping[f"[Diện tích loại đất tài sản {idx}.{m_idx}]"] = area
            mapping[f"[Thời hạn loại đất tài sản {idx}.{m_idx}]"] = term
            mapping[f"[Thờ hạn loại đất tài sản {idx}.{m_idx}]"] = term
            mapping[f"[Dòng loại đất {idx}.{m_idx}]"] = (
                f"{idx}.{m_idx}. {land_type}: {area} m2; Thời hạn: {term}." if row else ""
            )
            mapping[f"[Dòng loại đất tài sản {idx}.{m_idx}]"] = mapping[f"[Dòng loại đất {idx}.{m_idx}]"]

        # Tên chuẩn (contracts/entities.md §9.2): [trường][loại đất M][tài sản N],
        # viết liền không dấu, không "_". Không sinh alias rút gọn [loaidatM].
        for m_idx in range(1, MAX_WORD_LAND_ROWS + 1):
            row = land_rows[m_idx - 1] if m_idx <= len(land_rows) else None
            land_type, area, term = _land_row_values(row)
            mapping[f"[loaidat{m_idx}{idx}]"] = land_type
            mapping[f"[dientich{m_idx}{idx}]"] = area
            mapping[f"[thoihan{m_idx}{idx}]"] = term


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def build_template_mapping(case: Any, today: date | None = None,
                           *, static: bool = False,
                           needed: set[str] | None = None) -> dict[str, str]:
    """Mapping placeholder cho hồ sơ.

    `static`/`needed` mirror `_add_block_placeholders`: `needed` = nhu cầu
    của mẫu đang xuất (`template_group_needs`); `None` = caller không có
    mẫu → giữ đủ gate cũ."""
    context = build_word_context(case, today=today)
    # Tài sản vị trí 1 (primary) — snapshot đã commit nếu có, không thì
    # master live case.tai_san (đường cũ, hồ sơ chưa qua commit_stage).
    ts = context.assets[0] if context.assets else getattr(case, "tai_san", None)
    today = context.today
    person1, person2, person3, people_4_plus = _pick_core_people(case)

    people_slots = [None] * 21
    people_slots[1] = person1
    people_slots[2] = person2
    people_slots[3] = person3
    for idx, c in enumerate(people_4_plus[:17], start=4):
        people_slots[idx] = c

    # Metadata hồ sơ (MIN-141 đợt 3): snapshot `payload.case` đã commit
    # là nguồn đọc chính — cột live chỉ fallback khi snapshot KHÔNG có
    # key (hồ sơ legacy chưa qua commit_stage sau khi nâng cấp). Key đã
    # commit mà null/rỗng là giá trị chủ ý → xuất rỗng, không đọc master.
    # `[Nơi niêm yết]` KHÔNG còn đắp từ địa chỉ đất: chưa tra cứu được →
    # để trống, không tự khẳng định.
    _state_doc = _load_case_state(case)
    snap_case = _state_doc.get("case")
    if not isinstance(snap_case, dict):
        # Phòng snapshot lồng `payload.case` (một số doc legacy/wire).
        _inner = _state_doc.get("payload")
        snap_case = (_inner.get("case") if isinstance(_inner, dict)
                     else None)
    snap_case = snap_case if isinstance(snap_case, dict) else {}

    def _snap_meta(key: str, live_attr: str) -> str:
        if key in snap_case:
            return _safe_text(snap_case.get(key))
        return _safe_text(getattr(case, live_attr, ""))

    noi_niem_yet = _snap_meta("noiniemyet", "noi_niem_yet")
    noi_dung_viec = _snap_meta("noidungviec", "noi_dung_viec")
    ghi_chu = _snap_meta("ghichu", "ghi_chu")
    if "nguoinhanuyquyen" in snap_case:
        _uq_snap = snap_case.get("nguoinhanuyquyen")
        nguoi_nhan_uy_quyen = (
            _safe_text(_uq_snap.get("ten"))
            if isinstance(_uq_snap, dict) else _safe_text(_uq_snap))
    else:
        nguoi_nhan_uy_quyen = _safe_text(
            getattr(case, "nguoi_nhan_uy_quyen", ""))
    ngay_lap_ho_so = (
        _fmt_snapshot_date(snap_case.get("ngaylaphoso"))
        if "ngaylaphoso" in snap_case
        else _fmt_date(getattr(case, "ngay_lap_ho_so", None)))
    land_rows = _committed_land_rows(ts)

    if land_rows:
        total = 0.0
        for row in land_rows:
            try:
                total += float(row.get("dien_tich") or 0)
            except (TypeError, ValueError):
                pass
        dien_tich_so = total if total > 0 else getattr(ts, "dien_tich", None)
    else:
        dien_tich_so = getattr(ts, "dien_tich", None)
    dien_tich_str = f"{dien_tich_so:g}" if dien_tich_so else ""
    dien_tich_chu = _so_thanh_chu(dien_tich_so).capitalize() if dien_tich_so else ""
    loai_so_val = _safe_text(getattr(ts, "loai_so", "")) or "Giấy chứng nhận quyền sử dụng đất"
    loai_van_ban = _display_document_type(getattr(case, "loai_van_ban", ""))

    mapping: dict[str, str] = {
        "[Tên file]": f"ho_so_thua_ke_{getattr(case, 'id', '')}",
        "[Loại văn bản]": loai_van_ban,
        "[documenttype]": loai_van_ban,
        "[Ngày lập hồ sơ]": ngay_lap_ho_so,
        "[ngaylaphoso]": ngay_lap_ho_so,
        "[Nơi niêm yết]": noi_niem_yet,
        "[noiniemyet]": noi_niem_yet,
        "[Ghi chú]": ghi_chu,
        "[ghichu]": ghi_chu,
        "[Niêm Yết]": noi_niem_yet,
        "[NIÊM YẾT]": noi_niem_yet.upper() if noi_niem_yet else "",
        "[Loại sổ]": loai_so_val,
        "[Địa chỉ đất]": _safe_text(getattr(ts, "dia_chi", "")),
        "[Số serial]": _serial_text(getattr(ts, "so_serial", "")),
        "[Serial]": _serial_text(getattr(ts, "so_serial", "")),
        "[Số vào sổ]": _safe_text(getattr(ts, "so_vao_so", "")),
        "[Số thửa]": _safe_text(getattr(ts, "so_thua_dat", "")),
        "[Số tờ]": _safe_text(getattr(ts, "so_to_ban_do", "")),
        "[Số tờ bản đồ]": _safe_text(getattr(ts, "so_to_ban_do", "")),
        "[Diện tích]": dien_tich_str,
        "[Diện tích chữ]": dien_tich_chu,  # alias: tổng diện tích tài sản 1
        "[Đơn vị diện tích]": "m2",
        "[Hình thức sử dụng]": _safe_text(getattr(ts, "hinh_thuc_su_dung", "")),
        "[Loại đất]": _safe_text(getattr(ts, "loai_dat", "")),
        "[Thờ hạn]": _safe_text(getattr(ts, "thoi_han", "")),
        "[Nguồn gốc]": _safe_text(getattr(ts, "nguon_goc", "")),
        "[Ngày cấp sổ]": _fmt_date(getattr(ts, "ngay_cap", None)),
        "[Cơ quan cấp sổ]": _safe_text(getattr(ts, "co_quan_cap", "")),
        # Catalog §2: [Ngày]/[Tháng] pad 2 số; [Ngày chữ] thêm "mùng" khi ngày < 10.
        "[Ngày]": f"{today.day:02d}",
        "[Tháng]": f"{today.month:02d}",
        "[Năm]": str(today.year),
        "[Ngày chữ]": _ngay_chu(today.day),
        "[Tháng chữ]": _so_thanh_chu(today.month),
        "[Ngườ ủy quyền]": nguoi_nhan_uy_quyen,
        "[nguoinhanuyquyen]": nguoi_nhan_uy_quyen,
        "[Nội dung việc]": noi_dung_viec,
        "[noidungviec]": noi_dung_viec,
        # Marker loại đất tài sản 1: giá trị thật do _add_property_placeholders() ghi
        # đè; giữ ở đây làm giá trị mặc định khi hồ sơ không có tài sản nào.

        "[ONT]": "",
        "[CLN]": "",
        "[NTS]": "",
        "[LUC]": "",
    }

    # Slot Người: thứ tự giữ nguyên legacy (_pick_core_people trên
    # participants live), nhưng FIELD đọc bằng chứng đã commit trong
    # snapshot stage[] — đối chiếu theo entity id (MIN-141 review:
    # master customers dùng chung đổi sau commit không được kéo giá trị
    # slot đi; người chết có bằng chứng 'Giấy chứng tử' không bị model
    # trả None xóa mất). Snapshot cũ thiếu key → fallback master.
    _snap_people = _snapshot_people_by_id(_state_doc)

    def _slot_person(c):
        if c is None:
            return None
        row = _snap_people.get(str(getattr(c, "id", "")))
        return _snapshot_person_view(c, row) if row is not None else c

    # Token dữ liệu nghiệp vụ chưa có nguồn trong engine (catalog §8.2/§9: "Ngoại lệ:
    # chưa có nguồn"). Cố ý KHÔNG đăng ký: thay bằng chuỗi rỗng là xuất thầm thiếu dữ
    # liệu, nên chúng phải rơi vào find_unresolved_placeholders() và làm file thất bại.
    # Danh sách này chỉ để báo cáo/migration, không dùng để thay thế.
    for pending in NO_SOURCE_PLACEHOLDERS:
        if pending == "[Ngườ ủy quyền]" and "nguoinhanuyquyen" in snap_case:
            # MIN-141 saved metadata is confirmed evidence, including null.
            continue
        mapping.pop(pending, None)


    for i in range(1, 21):
        c = _slot_person(people_slots[i])
        mapping[f"[Tên {i}]"] = _safe_text(getattr(c, "ho_ten", "") if c else "")
        mapping[f"[Năm sinh {i}]"] = _fmt_year(getattr(c, "ngay_sinh", None) if c else None)
        mapping[f"[CCCD {i}]"] = _safe_text(getattr(c, "so_giay_to", "") if c else "")
        mapping[f"[Ngày cấp {i}]"] = _fmt_date(getattr(c, "ngay_cap", None) if c else None)
        mapping[f"[Địa chỉ {i}]"] = _safe_text(getattr(c, "dia_chi", "") if c else "")
        mapping[f"[Loại CC {i}]"] = _safe_text(getattr(c, "loai_giay_to", "") if c else "")
        mapping[f"[Nơi cấp CC {i}]"] = _safe_text(getattr(c, "noi_cap", "") if c else "")
        mapping[f"[Thường trú {i}]"] = _safe_text(getattr(c, "loai_dia_chi", "") if c else "")
        mapping[f"[Năm chết {i}]"] = _fmt_year(getattr(c, "ngay_chet", None) if c else None)
        # Tên chuẩn Người — viết liền không dấu không "_" (entities.md
        # §9.1): mirror đúng giá trị alias tiếng Việt cũ ở trên.
        mapping[f"[ten{i}]"] = mapping[f"[Tên {i}]"]
        mapping[f"[gioitinh{i}]"] = _safe_text(getattr(c, "gioi_tinh", "") if c else "")
        mapping[f"[ngaysinh{i}]"] = _fmt_date(getattr(c, "ngay_sinh", None) if c else None)
        mapping[f"[ngaychet{i}]"] = _fmt_date(getattr(c, "ngay_chet", None) if c else None)
        mapping[f"[sogiayto{i}]"] = mapping[f"[CCCD {i}]"]
        mapping[f"[ngaycap{i}]"] = mapping[f"[Ngày cấp {i}]"]
        mapping[f"[diachi{i}]"] = mapping[f"[Địa chỉ {i}]"]
        mapping[f"[loaigiayto{i}]"] = mapping[f"[Loại CC {i}]"]
        mapping[f"[noicap{i}]"] = mapping[f"[Nơi cấp CC {i}]"]
        mapping[f"[loaicutru{i}]"] = mapping[f"[Thường trú {i}]"]
        mapping[f"[loaidiachi{i}]"] = mapping[f"[Thường trú {i}]"]

    mapping["[Năm chết]"] = mapping.get("[Năm chết 1]", "")

    for i, row in enumerate(land_rows[:10], start=1):
        mapping[f"[Loại đất {i}]"] = _safe_text(row.get("loai_dat", ""))
        mapping[f"[Diện tích {i}]"] = _safe_text(row.get("dien_tich", ""))
        mapping[f"[Thờ hạn {i}]"] = _safe_text(row.get("thoi_han", ""))
    for i in range(len(land_rows) + 1, 11):
        mapping[f"[Loại đất {i}]"] = ""
        mapping[f"[Diện tích {i}]"] = ""
        mapping[f"[Thờ hạn {i}]"] = ""
    # Không còn đắp thời hạn lẻ (properties.thoi_han) vào [Thờ hạn 1] —
    # thời hạn thuộc từng cụm (MIN-141 đợt 2); thời hạn lẻ không gắn được
    # cụm được workspace_get báo orphan_thoi_han để đối chiếu.

    role_groups: dict[str, list[WordPerson]] = {key: [] for key in ROLE_LABELS}
    for person in context.all_people:
        role_key = "chu_dat" if person.is_land_owner else _role_key(person.role)
        if role_key in role_groups:
            role_groups[role_key].append(person)
    # Vai trò ủy quyền lưu ở participants.vai_tro, không bắt buộc node Diagram —
    # dùng authorized_people (đã merge vai trò đã lưu) thay vì lọc trên all_people.
    role_groups["nguoi_duoc_uy_quyen"] = list(context.authorized_people)
    for role_key, label in ROLE_LABELS.items():
        role_people = role_groups[role_key]
        _add_word_person_group(mapping, context, label.capitalize(), role_people)
        first = role_people[0] if role_people else None
        for field, value in _person_placeholder_values(context, first).items():
            mapping[f"[{field} {label}]"] = value

    # -----------------------------------------------------------------------
    # Word export V2 placeholders based on case_state_json
    # -----------------------------------------------------------------------
    _add_block_placeholders(mapping, context, static=static, needed=needed)
    _add_property_placeholders(mapping, context.assets)
    mapping["[Danh sách người thừa kế]"] = mapping["[Danh sách hàng thừa kế]"]
    mapping["[Đoạn mô tả quan hệ]"] = mapping["[Đoạn quan hệ gia đình]"]

    # AC3 - alias chính tả: bộ mẫu thừa kế viết đúng chính tả "[Thời hạn N]" còn
    # registry cũ chỉ có "[Thờ hạn N]". Mirror sang chính tả đúng để hai cách viết
    # cùng một nguồn dữ liệu (loại đất tài sản 1 / thoi_han), không trôi lệch.
    for key, value in list(mapping.items()):
        if "Thờ hạn" in key:
            mapping[key.replace("Thờ hạn", "Thời hạn")] = value

    return mapping


def _build_normalized_mapping(mapping: dict[str, str]) -> dict[str, str]:
    normalized = {}
    for k, v in mapping.items():
        if not (k.startswith("[") and k.endswith("]")):
            continue
        normalized[_normalize_token(k[1:-1])] = v
    return normalized


def _replace_text_placeholders(text: str, mapping: dict[str, str], normalized_mapping: dict[str, str]) -> str:
    new_text = text
    for k, v in mapping.items():
        if k in new_text:
            new_text = new_text.replace(k, v)

    def _token_repl(match):
        token = match.group(1)
        direct = mapping.get(f"[{token}]")
        if direct is not None:
            return direct
        norm = _normalize_token(token)
        if norm in normalized_mapping:
            return normalized_mapping[norm]
        return match.group(0)

    new_text = re.sub(r"\[([^\[\]]+)\]", _token_repl, new_text)
    new_text = re.sub(r"[ \t]{2,}", " ", new_text)
    new_text = re.sub(r"[ \t]+([,.;:])", r"\1", new_text)
    return new_text


# Placeholder dữ liệu nghiệp vụ đã biết nhưng CHƯA có nguồn trong engine
# (catalog §8.2/§9). Chúng không được map về chuỗi rỗng: phải báo thiếu để file
# export thất bại thay vì ra văn bản khuyết dữ liệu.
NO_SOURCE_PLACEHOLDERS = (
    "[Người ủy quyền]",
    "[Người ủy quyền2]",
    "[Ngườ ủy quyền]",  # alias sai chính tả của [Người ủy quyền]
    "[Ngườ ủy quyền2]",
    "[Số công chứng]",
    "[Giá chuyển nhượng]",
    "[SĐT]",
)

# Mã mẫu biểu cố định ("[01]", "[06.1]", ...) là nhãn in cứng CHỈ trên các mẫu
# đã đối chiếu thủ công. Việc miễn phải gắn với từng mẫu (NAIA-12 review): một
# file bất kỳ chứa "[01]"/"[54.99]" KHÔNG được miễn — token số dạng này ngoài
# danh sách đã xác minh của mẫu vẫn tính là placeholder chưa map và chặn export.
FIXED_FORM_CODE_RE = re.compile(r"^\[\d{1,2}(?:\.\d{1,2})?\]$")

# Ba mẫu thuế DK18 của Owner dùng cùng bộ nhãn mục in cứng [01]..[54]/[NN.M].
# Danh sách khớp đúng manifest migration (status "fixed_code") của từng file.
_DK18_FORM_CODES = frozenset(
    [
        "[01]", "[02]", "[03]", "[04]", "[05]", "[06]", "[06.1]", "[06.2]",
        "[07]", "[07.1]", "[07.2]", "[08]", "[09]", "[10]", "[10.1]", "[10.2]",
        "[10.3]", "[10.4]", "[10.5]", "[11]", "[12]", "[13]", "[14]", "[15]",
        "[16]", "[17]", "[17.1]", "[17.2]", "[17.3]", "[17.4]", "[17.5]",
        "[18]", "[19]", "[20]", "[20.1]", "[20.2]", "[20.3]", "[20.4]",
        "[20.5]", "[20.6]", "[21]", "[21.1]", "[21.2]", "[21.3]", "[21.4]",
        "[22]", "[22.1]", "[22.2]", "[23]", "[24]", "[25]", "[25.1]", "[25.2]",
        "[25.3]", "[26]", "[26.1]", "[26.2]", "[26.3]", "[26.4]", "[27]",
        "[28]", "[29]", "[30]", "[30.1]", "[30.2]", "[30.3]", "[31]", "[31.1]",
        "[31.2]", "[32]", "[32.1]", "[32.2]", "[32.3]", "[32.4]", "[32.5]",
        "[33]", "[33.1]", "[33.2]", "[33.3]", "[33.4]", "[33.5]", "[34]",
        "[34.1]", "[34.2]", "[34.3]", "[34.4]", "[34.5]", "[34.6]", "[34.7]",
        "[34.8]", "[35]", "[35.1]", "[35.2]", "[36]", "[36.1]", "[36.2]",
        "[36.3]", "[36.4]", "[37]", "[38]", "[39]", "[40]", "[41]", "[41.1]",
        "[41.10]", "[41.2]", "[41.3]", "[41.4]", "[41.5]", "[41.6]", "[41.7]",
        "[41.8]", "[41.9]", "[42]", "[42.1]", "[42.10]", "[42.11]", "[42.12]",
        "[42.13]", "[42.14]", "[42.15]", "[42.16]", "[42.17]", "[42.18]",
        "[42.19]", "[42.2]", "[42.20]", "[42.21]", "[42.22]", "[42.23]",
        "[42.24]", "[42.25]", "[42.26]", "[42.27]", "[42.28]", "[42.3]",
        "[42.4]", "[42.5]", "[42.6]", "[42.7]", "[42.8]", "[42.9]", "[43]",
        "[43.1]", "[43.2]", "[44]", "[44.1]", "[44.2]", "[45]", "[46]", "[47]",
        "[48]", "[49]", "[50]", "[51]", "[51.1]", "[51.2]", "[51.3]", "[51.4]",
        "[51.5]", "[51.6]", "[51.7]", "[51.8]", "[52]", "[53]", "[54]",
    ]
)


def _template_key(template_ref: Any) -> str:
    """Khóa mẫu: stem file chuẩn hóa (NFC, hạ, gộp khoảng trắng) để so registry."""
    if not template_ref:
        return ""
    stem = Path(str(template_ref)).stem
    text = unicodedata.normalize("NFC", stem).lower()
    return re.sub(r"\s+", " ", text).strip()


# Danh sách mã in cứng đã xác minh, GẮN VỚI TỪNG MẪU (khóa = stem file chuẩn hóa).
# Mẫu không có trong bảng này -> không mã nào được miễn.
VERIFIED_FIXED_CODES: dict[str, frozenset[str]] = {
    _key: _DK18_FORM_CODES
    for _key in (
        "2. dk18, thuế 1",
        "2. dk18, thuế 2",
        "2. dk18, thuế 3",
    )
}


def verified_fixed_codes(template_ref: Any) -> frozenset[str]:
    """Bộ mã in cứng đã xác minh của mẫu; mẫu chưa xác minh -> rỗng."""
    return VERIFIED_FIXED_CODES.get(_template_key(template_ref), frozenset())


# Các part XML có thể chứa placeholder: thân văn bản, header/footer mọi loại
# (default/first/even), footnote và endnote.
_WORD_XML_PART_RE = re.compile(
    r"/word/(?:document|header\d*|footer\d*|footnotes|endnotes)\.xml$"
)

_PLACEHOLDER_RE = re.compile(r"\[[^\[\]]+\]")


def is_fixed_form_code(token: str) -> bool:
    """True nếu token có DẠNG mã mẫu biểu in cứng ("[01]"/"[NN.M]").

    Đây chỉ là nhận dạng HÌNH DẠNG: token dạng này chỉ được miễn khỏi danh sách
    placeholder sót khi nó nằm trong ``verified_fixed_codes`` của ĐÚNG mẫu đang
    export — mẫu khác/chưa xác minh vẫn báo thiếu.
    """
    text = (token or "").strip()
    if not FIXED_FORM_CODE_RE.match(text):
        return False
    body = text[1:-1]
    head = body.split(".", 1)[0]
    return len(head) == 2 and 1 <= int(head) <= 54


class _WordPartRoot:
    """Một part XML quét được + cách ghi ngược khi part chỉ có ``blob``.

    python-docx chỉ gắn ``.element`` cho part nó hiểu (XmlPart); footnotes /
    endnotes của file thật là ``Part`` thường — phải parse blob ra cây, sửa,
    rồi ``commit()`` đẩy bytes mới về part trước khi ``Document.save``.
    """

    __slots__ = ("element", "_part")

    def __init__(self, element: Any, part: Any = None):
        self.element = element
        self._part = part

    def commit(self) -> None:
        if self._part is None:
            return
        from lxml import etree  # lazy: python-docx đã kéo lxml

        self._part._blob = etree.tostring(
            self.element, xml_declaration=True,
            encoding="UTF-8", standalone=True)


def _iter_word_part_roots(doc: Any):
    """Sinh ``_WordPartRoot`` của mọi part Word có thể chứa text placeholder.

    Part có ``.element`` sửa trực tiếp trên cây; part chỉ có ``.blob`` được
    parse ra cây mới và phải ``commit()`` sau khi thay text.
    """
    from docx.oxml import parse_xml  # lazy: giữ module import được khi thiếu python-docx

    package = getattr(getattr(doc, "part", None), "package", None)
    if package is None:
        return
    seen: set[int] = set()
    for part in package.iter_parts():
        name = str(getattr(part, "partname", ""))
        if not _WORD_XML_PART_RE.search(name):
            continue
        element = getattr(part, "element", None)
        blob_part = None
        if element is None:
            blob = getattr(part, "blob", None)
            try:
                element = parse_xml(blob)
            except Exception as exc:
                raise WordExportValidationError(
                    f"Mẫu Word có XML không hợp lệ tại {name}.") from exc
            blob_part = part
        if id(element) in seen:
            continue
        seen.add(id(element))
        yield _WordPartRoot(element, part=blob_part)


def _iter_part_paragraphs(root: Any, qn) -> list[Any]:
    """Mọi w:p trong part, gồm cả bảng lồng nhau và textbox trong bảng/header."""
    return list(root.iter(qn("w:p")))


def _paragraph_text_nodes(paragraph: Any, qn) -> list[Any]:
    """Các node w:t thuộc trực tiếp paragraph (bỏ qua w:p lồng trong textbox)."""
    nodes: list[Any] = []
    for run in paragraph.iter(qn("w:r")):
        owner = None
        for ancestor in run.iterancestors():
            if ancestor.tag == qn("w:p"):
                owner = ancestor
                break
        if owner is not paragraph:
            continue
        for node in run.iter(qn("w:t")):
            nodes.append(node)
    return nodes


def _lookup_placeholder(token: str, mapping: dict[str, str], normalized_mapping: dict[str, str]):
    """Trả về (found, value) cho token, ưu tiên khớp chính xác rồi khớp chuẩn hóa."""
    direct = mapping.get(f"[{token}]")
    if direct is not None:
        return True, direct
    norm = _normalize_token(token)
    if norm in normalized_mapping:
        return True, normalized_mapping[norm]
    return False, ""


def _replace_pass_in_text_nodes(nodes: list[Any], mapping: dict[str, str], normalized_mapping: dict[str, str]) -> int:
    texts = [node.text or "" for node in nodes]
    full = "".join(texts)
    if "[" not in full or "]" not in full:
        return 0

    boundaries: list[tuple[int, int, Any]] = []
    offset = 0
    for node, text in zip(nodes, texts):
        boundaries.append((offset, offset + len(text), node))
        offset += len(text)

    def _segments(start: int, end: int):
        for seg_start, seg_end, node in boundaries:
            if seg_end <= start or seg_start >= end:
                continue
            yield node, max(start, seg_start) - seg_start, min(end, seg_end) - seg_start

    replaced = 0
    # Duyệt từ phải sang trái: các vị trí bên trái vùng vừa sửa không đổi.
    for match in reversed(list(_PLACEHOLDER_RE.finditer(full))):
        found, value = _lookup_placeholder(match.group(0)[1:-1], mapping, normalized_mapping)
        if not found:
            continue
        segments = list(_segments(match.start(), match.end()))
        if not segments:
            continue
        first_node, first_start, first_end = segments[0]
        original = first_node.text or ""
        first_node.text = original[:first_start] + value + original[first_end:]
        for node, seg_start, seg_end in segments[1:]:
            text = node.text or ""
            node.text = text[:seg_start] + text[seg_end:]
        replaced += 1
    return replaced


# Giá trị của một số placeholder được engine tự soạn và có thể nhúng placeholder khác
# (ví dụ dòng khai tử nhúng "[Nơi niêm yết]"). Lặp tới điểm cố định để các placeholder
# lồng nhau cũng được thay, có giới hạn để tránh vòng lặp vô hạn.
_MAX_REPLACE_PASSES = 8


def _write_word_text(node: Any) -> None:
    """Encode edited text inside its existing run: real breaks and edge spaces."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    lines = re.split(r"\r\n|\r|\n", node.text or "")
    node.text = lines[0]
    current = node
    text_nodes = [node]
    for line in lines[1:]:
        br = OxmlElement("w:br")
        current.addnext(br)
        text = OxmlElement("w:t")
        text.attrib.update(node.attrib)
        text.text = line
        br.addnext(text)
        current = text
        text_nodes.append(text)
    for text_node in text_nodes:
        text = text_node.text or ""
        if text and (text[0].isspace() or text[-1].isspace()):
            text_node.set(qn("xml:space"), "preserve")


def _replace_in_text_nodes(nodes: list[Any], mapping: dict[str, str], normalized_mapping: dict[str, str]) -> int:
    """Thay placeholder trên chuỗi w:t, giữ nguyên định dạng của từng run.

    Placeholder bị cắt qua nhiều run vẫn được ghép lại để khớp, nhưng giá trị
    thay thế chỉ ghi vào run đầu tiên của vùng khớp và phần text còn lại của
    các run đó bị cắt bỏ - các run khác trong paragraph không bị đụng tới.

    Trả về số lần thay (manifest NAIA-12 cần "số lần thay" theo token).
    """
    if not nodes:
        return 0
    before = [node.text for node in nodes]
    replaced = 0
    for _ in range(_MAX_REPLACE_PASSES):
        count = _replace_pass_in_text_nodes(nodes, mapping, normalized_mapping)
        if not count:
            break
        replaced += count
    # Resolve nested tokens before splitting w:t so all passes see the same text.
    for node, original in zip(nodes, before):
        if node.text != original:
            _write_word_text(node)
    return replaced


def _replace_in_paragraph(paragraph: Any, mapping: dict[str, str], normalized_mapping: dict[str, str]) -> int:
    """Tương thích ngược: thay placeholder trong 1 paragraph, giữ định dạng run."""
    from docx.oxml.ns import qn

    return _replace_in_text_nodes(_paragraph_text_nodes(paragraph, qn), mapping, normalized_mapping)


def replace_in_doc(doc: Any, mapping: dict[str, str]) -> int:
    """Thay placeholder trong toàn bộ part Word, giữ nguyên định dạng run.

    Trả về tổng số lần thay trên mọi part (body, bảng lồng, header/footer
    mọi loại, footnote/endnote, textbox). Part chỉ có ``blob`` (footnotes /
    endnotes do python-docx chưa model hoá) được ``commit()`` ghi ngược bytes.
    """
    from docx.oxml.ns import qn

    normalized_mapping = _build_normalized_mapping(mapping)
    total = 0
    for part_root in _iter_word_part_roots(doc):
        changed = 0
        for paragraph in _iter_part_paragraphs(part_root.element, qn):
            nodes = _paragraph_text_nodes(paragraph, qn)
            if nodes:
                changed += _replace_in_text_nodes(nodes, mapping, normalized_mapping)
        if changed:
            part_root.commit()
        total += changed
    return total


def _collect_placeholder_occurrences(doc: Any) -> list[str]:
    """Mọi LẦN XUẤT HIỆN của token trên mọi part — không khử trùng lặp.

    Migration dùng danh sách này để manifest ghi đúng số lần xuất hiện thật
    của từng token (kể cả nhãn in cứng), thay vì đếm trên tập đã dedupe.
    """
    from docx.oxml.ns import qn

    found: list[str] = []
    for part_root in _iter_word_part_roots(doc):
        for paragraph in _iter_part_paragraphs(part_root.element, qn):
            # Ghép text theo paragraph để bắt cả placeholder bị cắt qua nhiều run.
            text = "".join(node.text or "" for node in _paragraph_text_nodes(paragraph, qn))
            if "[" not in text:
                continue
            found.extend(match.group(0) for match in _PLACEHOLDER_RE.finditer(text))
    return found


def _collect_placeholders(
    doc: Any,
    *,
    include_fixed_codes: bool = False,
    verified_codes: frozenset[str] | None = None,
) -> set[str]:
    verified = verified_codes if verified_codes is not None else frozenset()
    found: set[str] = set()
    for token in _collect_placeholder_occurrences(doc):
        if include_fixed_codes:
            found.add(token)  # kiểm kê đầy đủ (migration)
        elif not is_fixed_form_code(token) or token not in verified:
            # Token không có dạng mã biểu, hoặc dạng mã biểu nhưng KHÔNG nằm
            # trong danh sách đã xác minh của đúng mẫu này -> placeholder sót.
            found.add(token)
    return found


def collect_placeholders(doc: Any, *, include_fixed_codes: bool = False) -> list[str]:
    """Token placeholder có trong mọi part của file, sắp xếp, đã khử trùng lặp."""
    return sorted(_collect_placeholders(doc, include_fixed_codes=include_fixed_codes))


def collect_placeholder_occurrences(doc: Any) -> list[str]:
    """Mọi lần xuất hiện placeholder (kể cả mã in cứng), giữ nguyên thứ tự quét."""
    return _collect_placeholder_occurrences(doc)


def find_unresolved_placeholders(doc: Any, template_ref: Any = None) -> list[str]:
    """Placeholder chưa được thay, quét cùng phạm vi với replace_in_doc.

    ``template_ref`` (path/tên file mẫu) chọn danh sách mã in cứng đã xác minh
    của mẫu đó — chỉ mã trong danh sách mới được miễn; mẫu không xác minh hoặc
    token số ngoài danh sách vẫn báo thiếu và chặn export.
    """
    return sorted(_collect_placeholders(
        doc, verified_codes=verified_fixed_codes(template_ref)))


def list_public_builtin_templates(root: str | Path = "word_templates") -> list[dict[str, Any]]:
    root_path = Path(root)
    items = []
    for path in sorted(root_path.glob("*.docx"), key=lambda p: p.name.lower()):
        if path.name in HIDDEN_BUILTIN_TEMPLATE_NAMES:
            continue
        if path.name.startswith("~$"):
            continue
        items.append({"id": f"builtin:{path.name}", "ten_mau": path.stem, "is_active": False, "builtin": True})
    return items
