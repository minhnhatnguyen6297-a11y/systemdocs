"""Batch Word export theo contract notary.case-drafting.v2 §13.10 (MIN-110).

Sở hữu phần nghiệp vụ lớp catalog/orchestration cho
`notary.word_export_options` + `notary.word_export_batch`:

- document catalog: `document_key` → display_name / filename_stem ASCII
  do backend sở hữu (KHÔNG derive từ display_name — §8.3);
- readiness `word.*` block_reason qua `word_engine.word_block_reason`
  (cùng thứ tự check với `_add_block_placeholders`);
- render từng DOCX độc lập vào destination người dùng chọn — không ZIP,
  không output mặc định, KHÔNG BAO GIỜ ghi đè (open "xb" exclusive +
  reservation nội batch); file lỗi dở chỉ xóa đúng item đó;
- per-document status saved|failed|skipped + breakdown đủ ba list;
- cancel giữa batch: `check_cancel(pending_data)` — caller (sidecar)
  bọc thành `job.check_cancel(result)` của MIN-115; doc chưa bắt đầu
  đi vào breakdown.skipped, file đã lưu giữ nguyên;
- MIN-116: probe writability destination MỘT lần đầu batch bằng file
  tạo/xóa thật (`os.access` trên Windows misreport ACL deny-write) và
  KHÔNG dùng `tempfile.mkstemp` — stdlib retry PermissionError
  `TMP_MAX` = 2**31-1 lần trên nt → job hang vô hạn. Destination deny
  → mọi doc per-file `file_locked`, job failed trong giây.

Module này không import DB/FastAPI — `case` là ORM object duck-typed
giống `word_engine.build_word_context`. Sidecar `notary_adapter` chịu
trách session, FileRef/destination validation và map CommandError.
"""
from __future__ import annotations

import logging
import re
import shutil
import unicodedata
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Optional

from services import word_engine

_logger = logging.getLogger("word_batch_export")

SCHEMA_VERSION = "notary.case-drafting.v2"

_DOC_KEY_RE = re.compile(r"^[a-z][a-z0-9_]*$")


class WordBatchError(Exception):
    """Lỗi job-level → CommandError(code) ở sidecar.

    `result_data` (khi có) là `result.data` của batch — dùng cho
    `word_batch_failed` để breakdown/per-file errors lên wire (§8.4)."""

    def __init__(self, code: str, message: str, *,
                 details: Optional[dict] = None,
                 result_data: Optional[dict] = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.details = details
        self.result_data = result_data


class _DocFailed(Exception):
    """Lỗi một văn bản — các văn bản khác tiếp tục (§8.4)."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


# ---------------------------------------------------------------- catalog


@dataclass(frozen=True)
class WordDocumentSpec:
    """Một loại văn bản xuất được — registry mở (contract §8.1).

    `has_template=False` = v1 chưa có template cho loại này → luôn
    `word.template_missing` (như NO_TEMPLATE_KEYS của mock)."""

    document_key: str
    display_name: str
    filename_stem: str
    has_template: bool = True


DOC_CATALOG: tuple[WordDocumentSpec, ...] = (
    WordDocumentSpec(
        document_key="khai_nhan_di_san",
        display_name="Văn bản khai nhận di sản",
        filename_stem="Van_ban_khai_nhan_di_san"),
    WordDocumentSpec(
        document_key="thoa_thuan_phan_chia",
        display_name="Thỏa thuận phân chia di sản",
        filename_stem="Thoa_thuan_phan_chia_di_san"),
    WordDocumentSpec(
        document_key="niem_yet",
        display_name="Thông báo niêm yết",
        filename_stem="Thong_bao_niem_yet",
        has_template=False),
)
DOC_CATALOG_BY_KEY: dict[str, WordDocumentSpec] = {
    spec.document_key: spec for spec in DOC_CATALOG}

_BLOCK_MESSAGES = {
    "word.too_many_assets": "Quá 5 tài sản",
    "word.too_many_people": "Quá 20 người trong Stage",
    "word.too_many_signers": "Quá 20 người ký",
    "word.template_missing": "Văn bản chưa có template",
}

# ------------------------------------------------------------- validation


def validate_document_keys(document_keys: Any,
                           allowed_keys: Optional[set[str]] = None
                           ) -> list[str]:
    """Contract §8.2 — validate TRƯỚC khi tạo file đầu tiên.

    Thiếu/null/không list → validation_error (word_no_documents_selected
    CHỈ cho list rỗng — parity với mock oracle).

    `allowed_keys` (NAIA-8 folder mode): tập key được phép = các `.docx`
    vừa quét từ thư mục mẫu người dùng chọn — key ngoài tập này (kể cả
    path traversal giả dạng filename) → word_unknown_document_key. Không
    truyền → catalog mode cũ: key phải là snake_case trong DOC_CATALOG."""
    if not isinstance(document_keys, list):
        raise WordBatchError(
            "validation_error", "document_keys phải là danh sách")
    if not document_keys:
        raise WordBatchError(
            "word_no_documents_selected", "document_keys rỗng")
    seen: list = []                     # list: dict/list keys unhashable
    for key in document_keys:
        if key in seen:
            raise WordBatchError(
                "word_duplicate_document_key",
                f"document_key lặp: {key!r}",
                details={"document_key": key})
        seen.append(key)
    for key in document_keys:
        if allowed_keys is not None:
            if not (isinstance(key, str) and key in allowed_keys):
                raise WordBatchError(
                    "word_unknown_document_key",
                    f"document_key ngoài thư mục mẫu: {key!r}",
                    details={"document_key": key})
        elif not (isinstance(key, str) and _DOC_KEY_RE.match(key)
                  and key in DOC_CATALOG_BY_KEY):
            raise WordBatchError(
                "word_unknown_document_key",
                f"document_key ngoài catalog: {key!r}",
                details={"document_key": key})
    return list(document_keys)


# ------------------------------------------------- template folder (NAIA-8)

_DOCX_SUFFIX = ".docx"


def _contained_in(path: Path, real_root: Path) -> bool:
    """`path` còn nằm trong `real_root` (đã resolve) sau khi đuổi theo
    symlink/JUNCTION mọi cấp?

    Windows: junction KHÔNG phải `is_symlink()` — chỉ `resolve()` mới lộ
    đích thật. Entry trỏ ra ngoài thư mục được cấp quyền → False.
    Lỗi resolve (vòng lặp, reparse hỏng) → False (fail-closed)."""
    try:
        return path.resolve().is_relative_to(real_root)
    except (OSError, RuntimeError):
        return False


def _ascii_stem(filename: str) -> str:
    """Tên file → filename_stem ASCII-safe cho output `<stem>_HS-<id>.docx`.

    Bỏ dấu tiếng Việt (NFD + strip Mn), đ→d/Đ→D, ký tự lạ → `_`.
    Stem rỗng/reserved Windows → "Mau"."""
    stem = unicodedata.normalize("NFD", Path(filename).stem)
    stem = "".join(
        ch for ch in stem if unicodedata.category(ch) != "Mn")
    stem = stem.replace("đ", "d").replace("Đ", "D")
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_")
    if not stem or stem.upper() in _WIN_RESERVED_NAMES:
        stem = "Mau"
    return stem[:80]


def scan_template_dir(template_dir: Any) -> list[WordDocumentSpec]:
    """Quét `.docx` NGAY trong thư mục người dùng chọn (không đệ quy).

    Bỏ file khóa `~$*` và entry không phải file; sort ổn định
    (casefold + tên gốc) — §AC thư mục mẫu. `document_key` = tên file
    gốc (basename — không bao giờ chứa separator nên key lookup chặn
    traversal tự nhiên); `display_name` = stem gốc.

    Entry là symlink/junction trỏ ra NGOÀI thư mục mẫu bị loại —
    `is_file()` đuổi theo link nên phải kiểm containment sau resolve
    (đọc mẫu từ ngoài thư mục được cấp quyền là vi phạm scope)."""
    root = Path(template_dir)
    if not root.is_dir():
        raise WordBatchError(
            "file_not_found",
            "thư mục mẫu không tồn tại hoặc không đọc được")
    try:
        root_real = root.resolve()
        entries = sorted(
            root.iterdir(), key=lambda p: (p.name.casefold(), p.name))
    except (OSError, RuntimeError) as exc:
        raise WordBatchError(
            "file_not_found", f"không đọc được thư mục mẫu: {exc}")
    specs: list[WordDocumentSpec] = []
    for path in entries:
        name = path.name
        if name.startswith("~$") or not name.lower().endswith(_DOCX_SUFFIX):
            continue
        if not path.is_file() or not _contained_in(path, root_real):
            continue
        specs.append(WordDocumentSpec(
            document_key=name,
            display_name=path.stem,
            filename_stem=_ascii_stem(name)))
    return specs


def template_path_resolver(template_dir: Any
                           ) -> Callable[[str], Optional[Path]]:
    """Resolver `document_key → Path` gắn chặt với thư mục mẫu đã quét:
    chỉ đúng filename đã scan (không nhận đường dẫn tùy ý, không join
    path — key không khớp scan map trả None → word.template_missing).

    Kiểm tra lại containment TẠI THỜI ĐIỂM MỞ: mẫu bị xóa hoặc bị đổi
    thành symlink/junction trỏ ra ngoài sau khi scan → None (per-file
    word.template_missing), không bao giờ đọc file ngoài thư mục mẫu."""
    root = Path(template_dir)
    scanned = {spec.document_key for spec in scan_template_dir(root)}
    try:
        root_real = root.resolve()
    except (OSError, RuntimeError) as exc:
        raise WordBatchError(
            "file_not_found", f"không đọc được thư mục mẫu: {exc}")

    def _resolve(document_key: str) -> Optional[Path]:
        if document_key not in scanned:
            return None
        path = root / document_key
        if not path.is_file() or not _contained_in(path, root_real):
            return None
        return path

    return _resolve


# ------------------------------------------------------- filename safety

_WIN_RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}
_WIN_FORBIDDEN_CHARS = frozenset('<>:"/\\|?*')


def is_safe_windows_filename(name: Any) -> bool:
    """Tên file Windows hợp lệ cho output — defense-in-depth trên
    `actual_filename` (stem catalog đã sạch, nhưng không bao giờ tin
    dữ liệu đi ra ổ đĩa)."""
    if not isinstance(name, str) or not name:
        return False
    if ".." in name:                              # traversal
        return False
    if name != name.strip() or name.startswith("."):
        return False
    if name.endswith("."):
        return False
    if any(ch in _WIN_FORBIDDEN_CHARS or ord(ch) < 32 for ch in name):
        return False
    if name.split(".")[0].upper() in _WIN_RESERVED_NAMES:
        return False
    return True


def _assert_safe_filename(name: str) -> None:
    # `..`/separator dùng data-code đã register của contract (§8.3);
    # các dạng unsafe khác → word.invalid_filename.
    if isinstance(name, str) and (
            ".." in name or "/" in name or "\\" in name):
        raise _DocFailed(
            "word_path_traversal",
            f"tên file sinh ra thoát khỏi destination: {name!r}")
    if not is_safe_windows_filename(name):
        raise _DocFailed(
            "word.invalid_filename",
            f"tên file sinh ra không hợp lệ: {name!r}")


# ---------------------------------------------------------------- helpers


def _block_message(code: str) -> str:
    return _BLOCK_MESSAGES.get(code, code)


def _skipped_entry(spec: WordDocumentSpec) -> dict:
    return {
        "document_key": spec.document_key,
        "display_name": spec.display_name,
        "status": "skipped",
        "actual_filename": None,
        "output_file": None,
        "error": None,
        "missing": [],
    }


def _failed_entry(spec: WordDocumentSpec, code: str, message: str) -> dict:
    return {
        "document_key": spec.document_key,
        "display_name": spec.display_name,
        "status": "failed",
        "actual_filename": None,
        "output_file": None,
        "error": {"code": code, "message": message},
        "missing": [],
    }


def _batch_data(dest_dir: Path, documents: list[dict], *,
                succeeded: list[str], failed: list[str],
                skipped: list[str]) -> dict:
    return {
        "schema_version": SCHEMA_VERSION,
        "destination": {
            "path": str(dest_dir),
            "scope": "machine_local",
            "is_dir": True,
        },
        "documents": documents,
        "breakdown": {
            "succeeded": succeeded,
            "failed": failed,
            "skipped": skipped,
        },
    }


def _doc_error_from(exc: BaseException,
                    context: Optional[word_engine.WordExportContext], *,
                    static: bool = False) -> tuple[str, str]:
    """Exception → (code, message) cho per-file error — data-code §2.5."""
    if isinstance(exc, _DocFailed):
        return exc.code, exc.message
    if isinstance(exc, word_engine.WordExportValidationError):
        # Trần số lượng là lý do block duy nhất còn lại của mapping; hồ sơ
        # chưa lưu (context None) chỉ có template-level readiness.
        reason = (word_engine.word_block_reason(context, static=static)
                  if context is not None else None)
        return (reason or "word.render_failed"), str(exc)
    if isinstance(exc, FileNotFoundError):
        return "file_not_found", f"không tìm thấy file/thư mục: {exc}"
    if isinstance(exc, PermissionError):
        return "file_locked", f"không ghi được file: {exc}"
    return "word.render_failed", f"{type(exc).__name__}: {exc}"


_TEMP_CREATE_ATTEMPTS = 8


def _temp_file_in(dest_dir: Path, *, prefix: str, suffix: str) -> Path:
    """Tạo file rỗng exclusive trong `dest_dir` (uuid name + open "xb"),
    trả path đã tạo.

    MIN-116: KHÔNG dùng tempfile.mkstemp — `_mkstemp_inner` của stdlib
    retry PermissionError trên Windows khi `os.access(dir, W_OK)` trả
    True (ACL deny-write bị misreport), và TMP_MAX = 2**31-1 → vòng
    lặp thực tế vô hạn, job hang không cancel được. open("xb") raise
    PermissionError NGAY ở lần đầu — chỉ retry trên FileExistsError
    (va chạm uuid, gần như không thể)."""
    for _ in range(_TEMP_CREATE_ATTEMPTS):
        candidate = dest_dir / f"{prefix}{uuid.uuid4().hex}{suffix}"
        try:
            with open(candidate, "xb"):
                pass
        except FileExistsError:
            continue
        return candidate
    raise FileExistsError(
        f"không tạo được file tạm trong destination {dest_dir}")


def _probe_dest_writable(dest_dir: Path) -> None:
    """MIN-116: verify destination THỰC SỰ ghi được bằng một lần tạo/xóa
    file ở đầu batch — `os.access()` trên Windows không đáng tin với
    ACL (chỉ đọc attribute read-only bit, không eval ACL thật)."""
    probe = _temp_file_in(
        dest_dir, prefix=".word_export_probe_", suffix=".tmp")
    try:
        probe.unlink()
    except OSError:
        pass


def _render_temp(doc: Any, dest_dir: Path) -> Path:
    """Render docx ra file tạm NGAY TRONG destination (cùng filesystem —
    publish sau chỉ là copy nội bộ). Lỗi → xóa file tạm.

    MIN-116: `_temp_file_in` thay tempfile.mkstemp — PermissionError
    propagate ngay thành per-doc `file_locked` thay vì hang trong
    stdlib retry loop."""
    tmp_path = _temp_file_in(
        dest_dir, prefix=".word_export_", suffix=".docx")
    try:
        doc.save(str(tmp_path))
    except BaseException:
        try:
            tmp_path.unlink()
        except OSError:
            pass
        raise
    return tmp_path


def _publish(tmp_path: Path, dest_dir: Path, spec: WordDocumentSpec,
             case_id: int, taken: set[str]) -> str:
    """Đặt tên `<stem>_HS-<case_id>[_n].docx` (n ≥ 2) và ghi exclusive.

    `taken` = reservation nội batch (hai document cùng stem không va);
    `open("xb")` chống cả file đã tồn tại lẫn race — KHÔNG BAO GIỜ ghi đè.
    Copy lỗi → xóa đúng file dở của item này. Trả actual_filename."""
    n = 1
    while True:
        suffix = "" if n == 1 else f"_{n}"
        name = f"{spec.filename_stem}_HS-{case_id}{suffix}.docx"
        n += 1
        _assert_safe_filename(name)
        if name in taken:
            continue
        target = dest_dir / name
        if target.is_dir():
            continue                       # directory cùng tên → _n kế
        try:
            fh = open(target, "xb")          # exclusive create — atomic
        except FileExistsError:
            continue
        taken.add(name)
        try:
            with open(tmp_path, "rb") as src:
                shutil.copyfileobj(src, fh)
            fh.close()
        except BaseException:
            taken.discard(name)            # trả reservation khi fail
            try:
                fh.close()
            finally:
                try:
                    target.unlink()          # không để file nửa vời lại
                except OSError:
                    pass
            raise
        return name


def _export_one_document(*, spec: WordDocumentSpec, case: Any,
                         context: word_engine.WordExportContext,
                         resolve_template: Callable[[str], Any],
                         mapping_holder: dict, dest_dir: Path,
                         taken: set[str], case_id: int, today: Any,
                         check: Callable[[], None], static: bool = False) -> dict:
    """Xuất một văn bản — trả word_document_result entry.

    Mẫu được mở ĐÚNG MỘT LẦN trong try: needs nhóm cho readiness lấy từ
    chính document đang render (owner 15/10/2026 — không quét riêng rồi
    mở lại). Mọi exception trong try → status:failed với data-code;
    exception cancel (do `check()` raise) vẫn thoát qua — `check()` trong
    except re-raise khi flag đang set, nên cancel không bao giờ nuốt
    thành failed.

    Thiếu dữ kiện KHÔNG làm file lỗi (owner 23/10/2026 phần W): slot
    trống → placeholder rỗng, danh sách thiếu (dữ kiện nhóm, người được
    ủy quyền mẫu đòi, token chưa hỗ trợ) đi vào `missing` của entry.
    """
    needed: Optional[set] = None
    missing: list[str] = []
    try:
        # Thứ tự check như mock: template → data validation → render.
        if not spec.has_template:
            raise _DocFailed("word.template_missing",
                             _block_message("word.template_missing"))
        template_path = resolve_template(spec.document_key)
        if template_path is None:
            raise _DocFailed("word.template_missing",
                             _block_message("word.template_missing"))

        import docx  # lazy — giữ module import được khi thiếu python-docx
        doc = docx.Document(str(template_path))
        # Nhu cầu nhóm của chính mẫu đang render (owner 15/10/2026):
        # quyết định ghi thiếu nhóm nào + cache mapping theo needs.
        needed = word_engine.template_group_needs(doc)
        data_reason = word_engine.word_block_reason(context, static=static)
        if data_reason is not None:
            raise _DocFailed(data_reason, _block_message(data_reason))
        # Dữ kiện thiếu (kể cả người được ủy quyền mẫu đòi) chỉ ghi nhận.
        missing_auth = sorted(
            slot for slot in word_engine.required_authorized_slots(doc)
            if slot > len(context.authorized_people))
        # Mapping cache theo (static, needs): gate nhóm phụ thuộc từng mẫu
        # nên không dùng chung một mapping cho mọi file (tránh kết quả phụ
        # thuộc thứ tự file trong batch); mẫu cùng needs vẫn tái dùng —
        # ghi chú thiếu đi kèm cache vì chỉ sinh lúc build.
        cache_key = (static, frozenset(needed or ()))
        cached = mapping_holder.get(cache_key)
        if cached is None:
            data_missing: list[str] = []
            mapping = word_engine.build_template_mapping(
                case, today=today, static=static, needed=needed,
                missing=data_missing)
            cached = (mapping, data_missing)
            mapping_holder[cache_key] = cached
        else:
            mapping, data_missing = cached
        missing.extend(data_missing)
        for slot in missing_auth:
            word_engine.note_missing(
                missing, f"người được ủy quyền {slot}")
        word_engine.replace_in_doc(doc, mapping)
        unresolved = word_engine.find_unresolved_placeholders(
            doc, template_ref=template_path)
        if unresolved:
            # Token lạ còn trong văn bản: xuất file như hiện trạng (owner
            # thấy chỗ hỏng ngay trong Word) + ghi rõ vào kết quả.
            missing.append(
                f"{len(unresolved)} trường chưa hỗ trợ: "
                + ", ".join(unresolved))

        check()                        # cancel trước khi render
        tmp_path = _render_temp(doc, dest_dir)
        try:
            check()                        # cancel trước khi publish
            name = _publish(tmp_path, dest_dir, spec, case_id, taken)
        finally:
            try:
                tmp_path.unlink()
            except OSError:
                pass
        return {
            "document_key": spec.document_key,
            "display_name": spec.display_name,
            "status": "saved",
            "actual_filename": name,
            "output_file": {
                "path": str(dest_dir / name),
                "scope": "machine_local",
            },
            "error": None,
            "missing": missing,
        }
    except ImportError as exc:
        # Thiếu python-docx = lỗi hạ tầng, không phải lỗi của 1 văn bản —
        # thoát khỏi per-doc boundary thành job-level engine_unavailable.
        raise WordBatchError(
            "engine_unavailable",
            f"thiếu dependency python-docx: {exc}") from exc
    except Exception as exc:  # noqa: BLE001 — per-document boundary
        check()   # cancel pending → raise tại đây (doc tính skipped)
        code, message = _doc_error_from(exc, context, static=static)
        _logger.info("word_doc_failed %s",
                     {"document_key": spec.document_key, "code": code})
        return _failed_entry(spec, code, message)


# ---------------------------------------------------------------- public


def export_options(case: Any, *, resolve_template: Callable[[str], Any],
                   today: Any = None,
                   doc_specs: Optional[list[WordDocumentSpec]] = None
                   ) -> dict:
    """`result.data` của notary.word_export_options — read-only.

    Liệt kê TOÀN BỘ catalog (không chỉ template active); block_reason
    theo `word_engine.word_block_reason` — cùng semantics với render.

    `doc_specs` (NAIA-8): danh sách spec động từ `scan_template_dir` —
    thay DOC_CATALOG khi nguồn mẫu là thư mục người dùng chọn.

    `case=None` = hồ sơ chưa lưu (owner 23/10/2026 phần W): chỉ còn
    readiness mức mẫu (thiếu/mở không được); dữ kiện hồ sơ chưa có gì
    để chặn — chính lượt xuất tự tạo + lưu trước."""
    context = (word_engine.build_word_context(case, today=today)
               if case is not None else None)
    documents = []
    for spec in (doc_specs if doc_specs is not None else DOC_CATALOG):
        block = None
        if not spec.has_template:
            block = "word.template_missing"
        else:
            template_path = resolve_template(spec.document_key)
            if template_path is None:
                block = "word.template_missing"
            else:
                try:
                    import docx  # lazy — giữ module import được khi thiếu python-docx
                    docx.Document(str(template_path))
                except ImportError as exc:
                    raise WordBatchError(
                        "engine_unavailable",
                        f"thiếu dependency python-docx: {exc}") from exc
                except Exception as exc:
                    # Mẫu hiện diện nhưng không mở/parse được: readiness báo
                    # ĐÚNG data-code per-file mà export_batch sẽ trả — không
                    # nuốt lỗi rồi fallback gate blanket cũ (AC2 mẫu hỏng).
                    block = _doc_error_from(exc, context)[0]
                else:
                    if context is not None:
                        block = word_engine.word_block_reason(
                            context, static=doc_specs is not None)
        documents.append({
            "document_key": spec.document_key,
            "display_name": spec.display_name,
            "ready": block is None,
            "block_reason": block,
        })
    return {"schema_version": SCHEMA_VERSION, "documents": documents}


def export_batch(case: Any, *, case_id: int, document_keys: Any,
                 dest_dir: Any, resolve_template: Callable[[str], Any],
                 check_cancel: Optional[Callable[[dict], None]] = None,
                 report_progress: Optional[Callable] = None,
                 today: Any = None,
                 doc_specs: Optional[list[WordDocumentSpec]] = None,
                 create_case_subdir: bool = False) -> dict:
    """`result.data` của notary.word_export_batch — read snapshot + ghi file.

    `check_cancel(data)` được gọi GIỮA các văn bản với result.data hiện
    tại (doc chưa bắt đầu → skipped) — caller bọc job.check_cancel(result)
    của jobstore để canceled job vẫn mang breakdown (MIN-115).

    `doc_specs` (NAIA-8): spec động từ `scan_template_dir` — khi truyền,
    `document_keys` được đối chiếu lại với ĐÚNG tập file của thư mục
    được cấp quyền (folder-authorized), không phải catalog cố định.
    `create_case_subdir=True` → ghi vào `<dest>/HS-<case_id>/`
    (tạo nếu thiếu; file bên trong vẫn không bao giờ ghi đè);
    `data.destination` trong result vẫn là thư mục người dùng chọn
    (contract §8.2).

    Tất cả failed → WordBatchError("word_batch_failed", result_data=...).
    """
    # Catalog map (không phải tuple) — test/caller có thể patch entry
    # theo key; folder mode override bằng doc_specs của thư mục mẫu.
    specs = ({spec.document_key: spec for spec in doc_specs}
             if doc_specs is not None else DOC_CATALOG_BY_KEY)
    keys = validate_document_keys(
        document_keys,
        allowed_keys=set(specs) if doc_specs is not None else None)
    granted_dir = Path(dest_dir)     # thư mục người dùng chọn (result)
    dest_dir = granted_dir
    if create_case_subdir:
        dest_dir = granted_dir / f"HS-{case_id}"
        try:
            dest_dir.mkdir(exist_ok=True)
        except FileExistsError:
            raise WordBatchError(
                "file_locked",
                "thư mục con hồ sơ bị chặn bởi file cùng tên") from None
        except OSError as exc:
            raise WordBatchError(
                "file_locked",
                f"không tạo được thư mục con hồ sơ: {exc}") from exc
        # `HS-<case_id>` tồn tại sẵn dạng junction/symlink trỏ ra ngoài:
        # mkdir(exist_ok=True) chấp nhận vì "đã có dir" và mọi ghi sau đó
        # thoát khỏi thư mục được cấp quyền. Kiểm tra đích THẬT (đuổi
        # theo reparse point mọi cấp) phải nằm trong granted dir —
        # TRƯỚC probe/ghi/publish.
        try:
            granted_real = granted_dir.resolve()
        except (OSError, RuntimeError) as exc:
            raise WordBatchError(
                "file_not_found",
                f"không xác định được destination: {exc}") from exc
        if not _contained_in(dest_dir, granted_real):
            raise WordBatchError(
                "word_path_traversal",
                "thư mục con hồ sơ trỏ ra ngoài thư mục đích được chọn")
    if not dest_dir.is_dir():
        raise WordBatchError(
            "file_not_found",
            "destination không phải thư mục đang tồn tại")

    taken: set[str] = set()
    docs: list[dict] = []
    saved: list[str] = []
    failed: list[str] = []
    total = len(keys)

    def _pending_data() -> dict:
        """Snapshot result.data: đã xong giữ status; chưa bắt đầu → skipped."""
        entries = list(docs)
        remaining = keys[len(entries):]
        for key in remaining:
            entries.append(_skipped_entry(specs[key]))
        return _batch_data(
            granted_dir, entries, succeeded=list(saved),
            failed=list(failed), skipped=list(remaining))

    def _check() -> None:
        if check_cancel is not None:
            check_cancel(_pending_data())

    # MIN-116: probe writability destination MỘT lần đầu batch, trước khi
    # build context/render — destination deny-write (ACL) thì mọi doc
    # nhận per-file error chung, job kết thúc trong giây, KHÔNG hang.
    _check()                                       # cancel trước khi probe
    dest_error: Optional[tuple[str, str]] = None
    try:
        _probe_dest_writable(dest_dir)
    except FileNotFoundError as exc:
        dest_error = ("file_not_found",
                      f"destination không còn tồn tại: {exc}")
    except PermissionError as exc:
        dest_error = ("file_locked",
                      f"destination không ghi được: {exc}")
    except OSError as exc:
        dest_error = ("word.render_failed",
                      f"destination không tạo được file: {exc}")

    # Chỉ build context khi dest ghi được — dest_error set thì mọi doc
    # nhận per-file error chung, context build vừa thừa vừa có thể ném
    # raw exception ngoài shape per-doc (review MIN-116 minor).
    context = None
    if dest_error is None:
        context = word_engine.build_word_context(case, today=today)
    mapping_holder: dict = {}

    for key in keys:
        _check()                                   # cancel giữa batch
        spec = specs[key]
        if dest_error is not None:
            code, message = dest_error
            entry = _failed_entry(spec, code, message)
        else:
            entry = _export_one_document(
                spec=spec, case=case, context=context,
                resolve_template=resolve_template,
                mapping_holder=mapping_holder, dest_dir=dest_dir,
                taken=taken, case_id=case_id, today=today, check=_check,
                static=doc_specs is not None)
        docs.append(entry)
        (saved if entry["status"] == "saved" else failed).append(key)
        if report_progress is not None:
            try:
                report_progress(len(docs), total,
                                f"văn bản {len(docs)}/{total}")
            except Exception:
                # jobstore.report_progress cung la cancel checkpoint va
                # raise CancelledByUser() KHONG kem result — _check() de
                # cancel van mang pending snapshot (doc -> skipped) len
                # wire (contract §8.4 canceled example, MIN-115); exception
                # khac thi _check() no-op roi re-raise nguyen ban.
                _check()
                raise

    data = _batch_data(granted_dir, docs, succeeded=saved, failed=failed,
                       skipped=[])
    if len(failed) == total:
        raise WordBatchError(                      # §8.4 — job failed
            "word_batch_failed",
            "Toàn bộ văn bản trong lượt xuất đều lỗi",
            details={"documents": [
                {"document_key": d["document_key"],
                 "code": d["error"]["code"],
                 "message": d["error"]["message"]} for d in docs]},
            result_data=data)
    return data


__all__ = [
    "DOC_CATALOG",
    "DOC_CATALOG_BY_KEY",
    "WordBatchError",
    "WordDocumentSpec",
    "export_batch",
    "export_options",
    "is_safe_windows_filename",
    "scan_template_dir",
    "template_path_resolver",
    "validate_document_keys",
]
