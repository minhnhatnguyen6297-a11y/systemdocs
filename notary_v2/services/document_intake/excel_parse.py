"""Shared Excel person-sheet parsing for customer import + document intake.

Moved verbatim from routers/customers.py (MIN-108) so the legacy web upload
route and the multi-source intake service share one implementation — intake
emits suggestions for review instead of writing customers directly.
"""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime
from typing import Any, Optional

EXCEL_PERSON_FIELDS = [
    "ho_ten",
    "gioi_tinh",
    "ngay_sinh",
    "ngay_chet",
    "so_giay_to",
    "ngay_cap",
    "dia_chi",
]

# Từ khóa nhận diện cột — giữ nguyên thứ tự/keywords của upload_excel gốc.
PERSON_COLUMN_KEYWORDS = {
    "ho_ten": ["ho_ten", "ho ten", "ho va ten", "ten", "full_name"],
    "gioi_tinh": ["gioi_tinh", "gioi", "gender"],
    "ngay_sinh": ["ngay_sinh", "ngay sinh", "birth"],
    "ngay_chet": ["ngay_chet", "ngay mat", "death", "chet"],
    "so_giay_to": ["so_giay_to", "cccd", "giay to", "id_number", "khai tu"],
    "ngay_cap": ["ngay_cap", "ngay cap", "issue"],
    "dia_chi": ["dia_chi", "dia chi", "a ch", "a chi", "address"],
}

DATE_FIELDS = {"ngay_sinh", "ngay_chet", "ngay_cap"}


class ExcelParseError(Exception):
    """Lỗi mở/đọc cấu trúc file Excel (không phải lỗi từng dòng dữ liệu)."""


def normalize_excel_header(value: Any) -> str:
    s = str(value or "").strip()
    if not s:
        return ""
    decomposed = unicodedata.normalize("NFD", s)
    without_marks = "".join(ch for ch in decomposed if unicodedata.category(ch) != "Mn")
    without_marks = without_marks.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^a-z0-9]+", " ", without_marks.lower()).strip()


def consonant_skeleton(value: Any) -> str:
    normalized = normalize_excel_header(value)
    if not normalized:
        return ""
    compact = normalized.replace(" ", "")
    return re.sub(r"[aeiouy]", "", compact)


def header_matches_keyword(header: Any, keyword: Any) -> bool:
    norm_header = normalize_excel_header(header)
    norm_keyword = normalize_excel_header(keyword)
    if not norm_header or not norm_keyword:
        return False
    if norm_keyword in norm_header:
        return True
    header_skeleton = consonant_skeleton(norm_header)
    return len(header_skeleton) >= 3 and header_skeleton == consonant_skeleton(norm_keyword)


def parse_date(value: Any, allow_year_only: bool = True) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        v = int(value)
        # Người dùng gõ năm trực tiếp (1900-2099) vào ô Excel chưa format Text
        # → Excel lưu integer, không phải serial ngày. Ưu tiên xử lý như năm.
        if allow_year_only and 1900 <= v <= 2099:
            return date(v, 1, 1)
        try:
            from openpyxl.utils.datetime import from_excel
            dt = from_excel(value)
            if isinstance(dt, datetime):
                return dt.date()
            if isinstance(dt, date):
                return dt
        except Exception:
            pass

    s = str(value).strip()
    if not s or s.lower() in ("nan", "none"):
        return None

    if allow_year_only and len(s) == 4 and s.isdigit():
        y = int(s)
        if 1 <= y <= 9999:
            return date(y, 1, 1)

    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def format_date_display(d: Optional[date]) -> str:
    if not d:
        return ""
    if d.day == 1 and d.month == 1:
        return f"{d.year:04d}"
    return d.strftime("%d/%m/%Y")


def normalize_gender(value: str) -> str:
    s = (value or "").strip().lower()
    if not s:
        return ""
    if s in ("nam", "male", "m"):
        return "Nam"
    if s in ("nữ", "nu", "female", "f"):
        return "Nữ"
    return ""


def as_input_value(value: Any, is_date: bool = False) -> str:
    if value is None:
        return ""
    if is_date:
        d = parse_date(value, allow_year_only=True)
        if d:
            return format_date_display(d)
    return str(value).strip()


def _find_col(normalized_headers: list[str], keywords: list[str]) -> Optional[int]:
    for kw in keywords:
        if not normalize_excel_header(kw):
            continue
        for i, h in enumerate(normalized_headers):
            if header_matches_keyword(h, kw):
                return i
    return None


def _cell_ref(col_idx: int, row_num: int) -> str:
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(col_idx + 1)}{row_num}"


def parse_people_workbook(content: bytes) -> dict:
    """Parse workbook định dạng danh sách người (sheet đầu tiên / active).

    Trả {"sheet": <tên sheet>, "rows": [{"row_num", "form", "cells"}]}:
      - form: 7 trường EXCEL_PERSON_FIELDS đã qua as_input_value
        (ngày ở dạng "dd/mm/YYYY" hoặc "YYYY")
      - cells: {field: "<col><row>"} — provenance cho source_refs.

    Raises ExcelParseError khi file không mở được, không có dữ liệu, hoặc
    thiếu cột bắt buộc (ho_ten) — cùng message với upload_excel gốc.
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise ExcelParseError(f"Khong the mo file: {e}") from e

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ExcelParseError("File khong co du lieu.")

    raw_headers = [str(h).strip() if h else "" for h in rows[0]]
    normalized_headers = [normalize_excel_header(h) for h in raw_headers]

    col = {field: _find_col(normalized_headers, kws) for field, kws in PERSON_COLUMN_KEYWORDS.items()}

    missing = [f for f in ["ho_ten"] if col[f] is None]
    if missing:
        raise ExcelParseError(f"Khong nhan dien duoc cot: {', '.join(missing)}")

    def get_raw(row_vals, field):
        idx = col.get(field)
        if idx is None or idx >= len(row_vals):
            return None
        return row_vals[idx]

    parsed_rows = []
    for row_num, row in enumerate(rows[1:], start=2):
        form = {}
        cells = {}
        for field in EXCEL_PERSON_FIELDS:
            form[field] = as_input_value(get_raw(row, field), is_date=field in DATE_FIELDS)
            idx = col.get(field)
            if idx is not None:
                cells[field] = _cell_ref(idx, row_num)
        parsed_rows.append({"row_num": row_num, "form": form, "cells": cells})

    return {"sheet": str(ws.title or "Sheet1"), "rows": parsed_rows}
