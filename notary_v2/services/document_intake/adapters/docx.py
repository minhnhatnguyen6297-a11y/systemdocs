"""DOCX source adapter — python-docx text extraction → common parser."""

from __future__ import annotations

import io

from .. import ocr_pipeline
from ..models import AdapterContext, SourceFailed, SourceSpec
from ..normalization import (
    doc_to_suggestion,
    make_source_ref,
    unsupported_target_code,
)


async def extract(spec: SourceSpec, ctx: AdapterContext) -> list[dict]:
    content = ctx.read_file(spec)
    try:
        import docx

        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise SourceFailed("intake.parse_failed", f"không mở được DOCX: {exc}") from exc

    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    lines.append(text)

    if not lines:
        raise SourceFailed("intake.parse_failed", "DOCX không có nội dung text")

    parsed = ocr_pipeline._normalize_native_ocr_doc(lines, spec.filename)
    suggestion = doc_to_suggestion(
        parsed, spec.source_id, make_source_ref(filename=spec.filename))
    if not suggestion:
        raise SourceFailed(
            unsupported_target_code(parsed),
            "không trích được thực thể nào từ DOCX")
    return [suggestion]
