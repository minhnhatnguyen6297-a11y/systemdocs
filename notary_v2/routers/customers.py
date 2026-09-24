"""Customer router: CRUD + Excel import."""

import io
from typing import Any, Dict, Optional
from sqlalchemy.exc import IntegrityError

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import or_
from sqlalchemy.orm import Session

from database import get_db
from models import Customer
from services.document_intake.excel_parse import (  # noqa: F401 — re-export cho tests/consumers
    ExcelParseError,
    as_input_value,
    consonant_skeleton,
    format_date_display,
    header_matches_keyword,
    normalize_excel_header,
    normalize_gender,
    parse_date,
    parse_people_workbook,
)

router = APIRouter()
templates = Jinja2Templates(directory="frontend/templates")

EXCEL_COLUMNS = ["ho_ten", "gioi_tinh", "ngay_sinh", "ngay_chet", "so_giay_to", "ngay_cap", "dia_chi"]
DATE_FIELDS = {"ngay_sinh", "ngay_chet", "ngay_cap"}


def to_customer_json(c: Customer) -> Dict[str, Any]:
    return {
        "id": c.id,
        "ho_ten": c.ho_ten,
        "gioi_tinh": c.gioi_tinh,
        "ngay_sinh": format_date_display(c.ngay_sinh),
        "ngay_chet": format_date_display(c.ngay_chet),
        "so_giay_to": c.so_giay_to,
        "ngay_cap": format_date_display(c.ngay_cap),
        "dia_chi": c.dia_chi,
    }


def validate_customer_form(form: Dict[str, str], db: Session, current_id: Optional[int] = None, duplicate_as_error: bool = True):
    field_errors: Dict[str, str] = {}

    required = ["ho_ten"]
    for key in required:
        if not (form.get(key) or "").strip():
            field_errors[key] = "Bat buoc"

    gioi_tinh = normalize_gender(form.get("gioi_tinh", ""))
    if form.get("gioi_tinh") and not gioi_tinh:
        field_errors["gioi_tinh"] = "Chon Nam hoac Nu"

    ngay_sinh = parse_date(form.get("ngay_sinh"), allow_year_only=True)
    ngay_chet = parse_date(form.get("ngay_chet"), allow_year_only=True)
    ngay_cap = parse_date(form.get("ngay_cap"), allow_year_only=True)

    if form.get("ngay_sinh") and ngay_sinh is None:
        field_errors["ngay_sinh"] = "Ngay khong hop le (yyyy hoac dd/mm/yyyy)"
    if form.get("ngay_chet") and ngay_chet is None:
        field_errors["ngay_chet"] = "Ngay khong hop le (yyyy hoac dd/mm/yyyy)"
    if form.get("ngay_cap") and ngay_cap is None:
        field_errors["ngay_cap"] = "Ngay khong hop le (yyyy hoac dd/mm/yyyy)"

    so_gt = (form.get("so_giay_to") or "").strip()
    if so_gt and duplicate_as_error:
        q = db.query(Customer).filter(Customer.so_giay_to == so_gt)
        if current_id is not None:
            q = q.filter(Customer.id != current_id)
        if q.first():
            field_errors["so_giay_to"] = "So giay to da ton tai"

    cleaned = {
        "ho_ten": (form.get("ho_ten") or "").strip(),
        "gioi_tinh": gioi_tinh,
        "ngay_sinh": ngay_sinh,
        "ngay_chet": ngay_chet,
        "so_giay_to": so_gt or None,  # NULL thay vì "" để tránh unique constraint khi không có số giấy tờ
        "ngay_cap": ngay_cap,
        "dia_chi": (form.get("dia_chi") or "").strip(),
    }
    return cleaned, field_errors


def result_message(field_errors: Dict[str, str]) -> str:
    order = ["ho_ten", "gioi_tinh", "ngay_sinh", "ngay_chet", "so_giay_to", "ngay_cap", "dia_chi"]
    labels = {
        "ho_ten": "Thieu ho ten",
        "gioi_tinh": "Gioi tinh khong hop le",
        "ngay_sinh": "Ngay sinh khong hop le",
        "ngay_chet": "Ngay chet khong hop le",
        "so_giay_to": "So giay to khong hop le",
        "ngay_cap": "Ngay cap khong hop le",
        "dia_chi": "Thieu dia chi",
    }
    msg = [labels[k] for k in order if k in field_errors]
    return " | ".join(msg) if msg else "Du lieu khong hop le"


@router.get("/")
def list_customers(request: Request, db: Session = Depends(get_db), q: str = ""):
    query = db.query(Customer)
    if q:
        query = query.filter(or_(Customer.ho_ten.contains(q), Customer.so_giay_to.contains(q), Customer.dia_chi.contains(q)))
    customers = query.order_by(Customer.ho_ten).all()
    return templates.TemplateResponse("customers/list.html", {"request": request, "customers": customers, "q": q})


@router.get("/api/search")
def search_customers(db: Session = Depends(get_db), q: str = "", limit: int = 10):
    query = db.query(Customer)
    if q:
        query = query.filter(or_(
            Customer.ho_ten.contains(q), 
            Customer.so_giay_to.contains(q), 
            Customer.dia_chi.contains(q)
        ))
    customers = query.order_by(Customer.ho_ten).limit(limit).all()
    return JSONResponse({
        "ok": True, 
        "data": [to_customer_json(c) for c in customers]
    })


@router.get("/download-template")
def download_template():
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sach nguoi"

    # Các cột ngày tháng — cần format Text để Excel không tự chuyển năm thành serial number
    DATE_COLS = {"ngay_sinh", "ngay_chet", "ngay_cap"}

    widths = [24, 14, 16, 16, 22, 16, 40]
    for idx, field in enumerate(EXCEL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=field)
        cell.font = Font(bold=True)
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = widths[idx - 1]
        if field in DATE_COLS:
            # Format @ = Text: người dùng gõ "1995" hay "15/06/1995" đều giữ nguyên text
            for row in range(2, 1002):
                ws.cell(row=row, column=idx).number_format = "@"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=mau_nhap_nguoi.xlsx"},
    )


@router.post("/upload-excel")
async def upload_excel(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        return templates.TemplateResponse("customers/upload_result.html", {
            "request": request, "error_global": "Chi chap nhan file .xlsx",
            "results": [], "added": 0, "skipped": 0, "errors": 0, "total": 0
        })

    content = await file.read()
    try:
        sheet = parse_people_workbook(content)
    except ExcelParseError as e:
        return templates.TemplateResponse("customers/upload_result.html", {
            "request": request, "error_global": str(e),
            "results": [], "added": 0, "skipped": 0, "errors": 0, "total": 0
        })

    results = []
    added_customers = []
    added = skipped = errors = 0

    for parsed_row in sheet["rows"]:
        row_num = parsed_row["row_num"]
        raw_form = parsed_row["form"]

        if not raw_form["ho_ten"] and not raw_form["so_giay_to"]:
            continue

        if raw_form["so_giay_to"] and db.query(Customer).filter(Customer.so_giay_to == raw_form["so_giay_to"]).first():
            results.append({"row": row_num, "name": raw_form["ho_ten"] or "?", "status": "skip", "message": "Trung so giay to, bo qua"})
            skipped += 1
            continue

        cleaned, field_errors = validate_customer_form(raw_form, db, duplicate_as_error=False)
        if field_errors:
            results.append({
                "row": row_num,
                "name": raw_form["ho_ten"] or "?",
                "status": "error",
                "message": result_message(field_errors),
                "raw": raw_form,
                "field_errors": field_errors,
            })
            errors += 1
            continue

        try:
            c = Customer(**cleaned)
            db.add(c)
            db.commit()
            db.refresh(c)
            added_customers.append(to_customer_json(c))
            results.append({"row": row_num, "name": cleaned["ho_ten"], "status": "ok", "message": "Them thanh cong"})
            added += 1
        except Exception as e:
            db.rollback()
            results.append({
                "row": row_num,
                "name": raw_form["ho_ten"] or "?",
                "status": "error",
                "message": f"Loi: {e}",
                "raw": raw_form,
                "field_errors": {},
            })
            errors += 1

    return templates.TemplateResponse("customers/upload_result.html", {
        "request": request,
        "error_global": None,
        "results": results,
        "added_customers": added_customers,
        "added": added,
        "skipped": skipped,
        "errors": errors,
        "total": added + skipped + errors,
    })


@router.post("/upload-excel/save-row")
def upload_excel_save_row(
    ho_ten: Optional[str] = Form(None),
    gioi_tinh: Optional[str] = Form(None),
    ngay_sinh: Optional[str] = Form(None),
    ngay_chet: Optional[str] = Form(None),
    so_giay_to: Optional[str] = Form(None),
    ngay_cap: Optional[str] = Form(None),
    dia_chi: Optional[str] = Form(None),
    db: Session = Depends(get_db),
):
    form = {
        "ho_ten": (ho_ten or "").strip(),
        "gioi_tinh": (gioi_tinh or "").strip(),
        "ngay_sinh": (ngay_sinh or "").strip(),
        "ngay_chet": (ngay_chet or "").strip(),
        "so_giay_to": (so_giay_to or "").strip(),
        "ngay_cap": (ngay_cap or "").strip(),
        "dia_chi": (dia_chi or "").strip(),
    }
    cleaned, field_errors = validate_customer_form(form, db, duplicate_as_error=True)
    if field_errors:
        return JSONResponse({"ok": False, "errors": field_errors, "message": result_message(field_errors)}, status_code=400)

    c = Customer(**cleaned)
    db.add(c)
    db.commit()
    db.refresh(c)
    return JSONResponse({"ok": True, "message": "Da luu thanh cong", "customer": to_customer_json(c)})


@router.get("/create")
def create_form(request: Request):
    form = {k: "" for k in EXCEL_COLUMNS}
    return templates.TemplateResponse("customers/form.html", {
        "request": request, "obj": None, "errors": [], "field_errors": {}, "form": form
    })


@router.post("/inline-create")
def inline_create(
    ho_ten: Optional[str] = Form(None), gioi_tinh: Optional[str] = Form(None),
    ngay_sinh: Optional[str] = Form(None), ngay_chet: Optional[str] = Form(None),
    so_giay_to: Optional[str] = Form(None), ngay_cap: Optional[str] = Form(None),
    dia_chi: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    name = (ho_ten or "").strip()
    if not name:
        return JSONResponse({"ok": False, "errors": {"ho_ten": "Bat buoc"}}, status_code=400)

    so_giay_to_val = (so_giay_to or "").strip() or None
    # Upsert: nếu CCCD đã tồn tại → cập nhật với dữ liệu mới rồi trả về
    if so_giay_to_val:
        existing = db.query(Customer).filter(Customer.so_giay_to == so_giay_to_val).first()
        if existing:
            if name:
                existing.ho_ten = name
            if gioi_tinh is not None:
                gt = normalize_gender(gioi_tinh)
                if gt:
                    existing.gioi_tinh = gt
                elif not gioi_tinh.strip():
                    existing.gioi_tinh = None
            if ngay_sinh is not None:
                ns = parse_date(ngay_sinh, allow_year_only=True)
                if ns:
                    existing.ngay_sinh = ns
                elif not ngay_sinh.strip():
                    existing.ngay_sinh = None
            if ngay_chet is not None:
                nd = parse_date(ngay_chet, allow_year_only=True)
                if nd:
                    existing.ngay_chet = nd
                elif not ngay_chet.strip():
                    existing.ngay_chet = None
            if ngay_cap is not None:
                nc = parse_date(ngay_cap, allow_year_only=True)
                if nc:
                    existing.ngay_cap = nc
                elif not ngay_cap.strip():
                    existing.ngay_cap = None
            dc = (dia_chi or "").strip()
            if dc or dia_chi is not None:
                existing.dia_chi = dc or None
            db.commit()
            db.refresh(existing)
            return JSONResponse({"ok": True, "customer": to_customer_json(existing), "updated": True})

    c = Customer(
        ho_ten=name,
        gioi_tinh=normalize_gender(gioi_tinh or "") or None,
        ngay_sinh=parse_date(ngay_sinh or "", allow_year_only=True),
        ngay_chet=parse_date(ngay_chet or "", allow_year_only=True),
        so_giay_to=so_giay_to_val,
        ngay_cap=parse_date(ngay_cap or "", allow_year_only=True),
        dia_chi=(dia_chi or "").strip() or None,
    )
    db.add(c)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        existing = db.query(Customer).filter(Customer.so_giay_to == so_giay_to_val).first()
        if existing:
            return JSONResponse({"ok": True, "customer": to_customer_json(existing)})
        return JSONResponse({"ok": False, "errors": {"so_giay_to": "Trùng"}}, status_code=400)
    db.refresh(c)
    return JSONResponse({"ok": True, "customer": to_customer_json(c)})


@router.post("/{cid}/quick-update")
def quick_update(
    cid: int,
    ho_ten: Optional[str] = Form(None), gioi_tinh: Optional[str] = Form(None),
    ngay_sinh: Optional[str] = Form(None), ngay_chet: Optional[str] = Form(None), so_giay_to: Optional[str] = Form(None),
    ngay_cap: Optional[str] = Form(None), dia_chi: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    c = db.get(Customer, cid)
    if not c:
        raise HTTPException(status_code=404, detail="Not found")
    if ho_ten and ho_ten.strip():
        c.ho_ten = ho_ten.strip()
    if gioi_tinh is not None:
        gt = normalize_gender(gioi_tinh)
        if gt:
            c.gioi_tinh = gt
        elif not gioi_tinh.strip():
            c.gioi_tinh = None
    if ngay_sinh is not None:
        ns = parse_date(ngay_sinh, allow_year_only=True)
        if ns:
            c.ngay_sinh = ns
        elif not ngay_sinh.strip():
            c.ngay_sinh = None
    if ngay_chet is not None:
        nd = parse_date(ngay_chet, allow_year_only=True)
        if nd:
            c.ngay_chet = nd
        elif not ngay_chet.strip():
            c.ngay_chet = None
    if so_giay_to is not None:
        so = (so_giay_to or "").strip()
        c.so_giay_to = so or None
    if ngay_cap is not None:
        nc = parse_date(ngay_cap, allow_year_only=True)
        if nc:
            c.ngay_cap = nc
        elif not ngay_cap.strip():
            c.ngay_cap = None
    dc = (dia_chi or "").strip()
    if dc or dia_chi is not None:
        c.dia_chi = dc or None
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return JSONResponse({"ok": False, "error": "so_giay_to trùng"}, status_code=400)
    db.refresh(c)
    return JSONResponse({"ok": True, "customer": to_customer_json(c)})


@router.post("/create")
def create(
    request: Request,
    ho_ten: Optional[str] = Form(None), gioi_tinh: Optional[str] = Form(None),
    ngay_sinh: Optional[str] = Form(None), ngay_chet: Optional[str] = Form(None),
    so_giay_to: Optional[str] = Form(None), ngay_cap: Optional[str] = Form(None),
    dia_chi: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    form = {
        "ho_ten": (ho_ten or "").strip(),
        "gioi_tinh": (gioi_tinh or "").strip(),
        "ngay_sinh": (ngay_sinh or "").strip(),
        "ngay_chet": (ngay_chet or "").strip(),
        "so_giay_to": (so_giay_to or "").strip(),
        "ngay_cap": (ngay_cap or "").strip(),
        "dia_chi": (dia_chi or "").strip(),
    }

    cleaned, field_errors = validate_customer_form(form, db)
    errors = []
    if "so_giay_to" in field_errors:
        errors.append(f"So giay to '{form['so_giay_to']}' da ton tai")

    if field_errors:
        return templates.TemplateResponse("customers/form.html", {
            "request": request, "obj": None,
            "errors": errors, "field_errors": field_errors, "form": form
        })

    c = Customer(**cleaned)
    db.add(c)
    db.commit()
    return RedirectResponse("/customers", status_code=302)


@router.get("/{cid}")
def detail(cid: int, request: Request, db: Session = Depends(get_db)):
    c = db.query(Customer).filter(Customer.id == cid).first()
    if not c:
        raise HTTPException(404)
    return templates.TemplateResponse("customers/detail.html", {"request": request, "obj": c})


@router.get("/{cid}/edit")
def edit_form(cid: int, request: Request, db: Session = Depends(get_db)):
    c = db.query(Customer).filter(Customer.id == cid).first()
    if not c:
        raise HTTPException(404)
    form = {
        "ho_ten": c.ho_ten or "",
        "gioi_tinh": c.gioi_tinh or "",
        "ngay_sinh": format_date_display(c.ngay_sinh),
        "ngay_chet": format_date_display(c.ngay_chet),
        "so_giay_to": c.so_giay_to or "",
        "ngay_cap": format_date_display(c.ngay_cap),
        "dia_chi": c.dia_chi or "",
    }
    return templates.TemplateResponse("customers/form.html", {
        "request": request, "obj": c, "errors": [], "field_errors": {}, "form": form
    })


@router.post("/{cid}/edit")
def edit(
    cid: int, request: Request,
    ho_ten: Optional[str] = Form(None), gioi_tinh: Optional[str] = Form(None),
    ngay_sinh: Optional[str] = Form(None), ngay_chet: Optional[str] = Form(None),
    so_giay_to: Optional[str] = Form(None), ngay_cap: Optional[str] = Form(None),
    dia_chi: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    c = db.query(Customer).filter(Customer.id == cid).first()
    if not c:
        raise HTTPException(404)

    form = {
        "ho_ten": (ho_ten or "").strip(),
        "gioi_tinh": (gioi_tinh or "").strip(),
        "ngay_sinh": (ngay_sinh or "").strip(),
        "ngay_chet": (ngay_chet or "").strip(),
        "so_giay_to": (so_giay_to or "").strip(),
        "ngay_cap": (ngay_cap or "").strip(),
        "dia_chi": (dia_chi or "").strip(),
    }

    cleaned, field_errors = validate_customer_form(form, db, current_id=cid)
    errors = []
    if "so_giay_to" in field_errors:
        errors.append(f"So giay to '{form['so_giay_to']}' da ton tai")

    if field_errors:
        return templates.TemplateResponse("customers/form.html", {
            "request": request, "obj": c, "errors": errors, "field_errors": field_errors, "form": form
        })

    c.ho_ten = cleaned["ho_ten"]
    c.gioi_tinh = cleaned["gioi_tinh"]
    c.ngay_sinh = cleaned["ngay_sinh"]
    c.ngay_chet = cleaned["ngay_chet"]
    c.so_giay_to = cleaned["so_giay_to"]
    c.ngay_cap = cleaned["ngay_cap"]
    c.dia_chi = cleaned["dia_chi"]
    db.commit()
    return RedirectResponse(f"/customers/{cid}", status_code=302)


@router.post("/{cid}/delete")
def delete(cid: int, db: Session = Depends(get_db)):
    c = db.query(Customer).filter(Customer.id == cid).first()
    if c:
        db.delete(c)
        db.commit()
    return RedirectResponse("/customers", status_code=302)
