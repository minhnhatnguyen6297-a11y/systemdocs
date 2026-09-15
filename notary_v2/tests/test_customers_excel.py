import asyncio
from datetime import date
from io import BytesIO
import json
import unittest

import openpyxl
from fastapi import UploadFile
from sqlalchemy.exc import IntegrityError
from starlette.requests import Request

from models import Customer
from routers.customers import (
    EXCEL_COLUMNS,
    as_input_value,
    consonant_skeleton,
    download_template,
    header_matches_keyword,
    inline_create,
    normalize_excel_header,
    parse_date,
    quick_update,
    upload_excel,
    validate_customer_form,
)


class DummyQuery:
    def __init__(self, result=None):
        self.result = result

    def filter(self, *args, **kwargs):
        return self

    def first(self):
        return self.result


class DummyDb:
    def __init__(self, customer=None, query_result=None, commit_error=None):
        self.customer = customer
        self.query_result = query_result
        self.commit_error = commit_error
        self.added = []

    def query(self, *args, **kwargs):
        return DummyQuery(self.query_result)

    def get(self, model, cid):
        return self.customer if self.customer and self.customer.id == cid else None

    def add(self, obj):
        self.added.append(obj)

    def commit(self):
        if self.commit_error:
            raise self.commit_error

    def refresh(self, obj):
        if getattr(obj, "id", None) is None:
            obj.id = len(self.added)

    def rollback(self):
        return None


class CustomerExcelImportTests(unittest.TestCase):
    def test_numeric_year_is_not_treated_as_excel_serial(self):
        self.assertEqual(parse_date(1995), date(1995, 1, 1))
        self.assertEqual(as_input_value(1995, is_date=True), "1995")

    def test_excel_serial_date_still_imports_as_real_date(self):
        self.assertEqual(parse_date(44927), date(2023, 1, 1))

    def test_normalize_excel_header_matches_vietnamese_titles(self):
        self.assertIn("ho va ten", normalize_excel_header("Họ và tên"))
        self.assertIn("ngay sinh", normalize_excel_header("Ngày sinh"))
        self.assertIn("so giay to", normalize_excel_header("Số giấy tờ"))

    def test_garbled_headers_match_without_false_positive(self):
        self.assertEqual(consonant_skeleton("H? v? t?n"), "hvtn")
        self.assertTrue(header_matches_keyword("H? v? t?n", "ho va ten"))
        self.assertTrue(header_matches_keyword("Gi?i t?nh", "gioi_tinh"))
        self.assertTrue(header_matches_keyword("Ng?y c?p", "ngay_cap"))
        self.assertFalse(header_matches_keyword("??a ch?", "dia_chi"))
        self.assertFalse(header_matches_keyword("G", "gioi"))
        self.assertFalse(header_matches_keyword("Ch", "chi"))

    def test_upload_excel_accepts_garbled_vietnamese_headers(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["H? v? t?n", "Gi?i t?nh", "Ng?y sinh", "S? gi?y t?", "Ng?y c?p", "??a ch?"])
        ws.append(["Tran Van A", "Nam", "1988", "", "2021-05-20", "Nam Dinh"])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        async def run_upload():
            request = Request({"type": "http", "method": "POST", "path": "/customers/upload-excel", "headers": []})
            response = await upload_excel(request, UploadFile(filename="test.xlsx", file=buf), DummyDb())
            return response.context

        context = asyncio.run(run_upload())
        self.assertIsNone(context["error_global"])
        self.assertEqual(context["added"], 1)
        self.assertEqual(context["added_customers"][0]["ho_ten"], "Tran Van A")

    def test_missing_document_number_is_cleaned_to_none(self):
        cleaned, errors = validate_customer_form(
            {
                "ho_ten": "Nguyen Van A",
                "gioi_tinh": "Nam",
                "ngay_sinh": "1995",
                "ngay_chet": "",
                "so_giay_to": "",
                "ngay_cap": "",
                "dia_chi": "",
            },
            DummyDb(),
        )

        self.assertEqual(errors, {})
        self.assertIsNone(cleaned["so_giay_to"])

    def test_invalid_dates_are_reported_for_each_date_field(self):
        for field in ("ngay_sinh", "ngay_chet", "ngay_cap"):
            form = {
                "ho_ten": "Nguyen Van A",
                "gioi_tinh": "Nam",
                "ngay_sinh": "",
                "ngay_chet": "",
                "so_giay_to": "",
                "ngay_cap": "",
                "dia_chi": "",
            }
            form[field] = "not-a-date"
            _, errors = validate_customer_form(form, DummyDb())
            self.assertIn(field, errors)

    def test_duplicate_document_is_rejected_by_validation(self):
        duplicate = Customer(id=2, ho_ten="Existing", so_giay_to="123")
        cleaned, errors = validate_customer_form(
            {
                "ho_ten": "Nguyen Van A",
                "gioi_tinh": "Nam",
                "ngay_sinh": "1995",
                "ngay_chet": "",
                "so_giay_to": "123",
                "ngay_cap": "",
                "dia_chi": "",
            },
            DummyDb(query_result=duplicate),
        )

        self.assertEqual(cleaned["so_giay_to"], "123")
        self.assertEqual(errors["so_giay_to"], "So giay to da ton tai")

    def test_quick_update_omitted_fields_preserve_and_empty_clears_nullable_fields(self):
        customer = Customer(
            id=1,
            ho_ten="Existing",
            gioi_tinh="Nam",
            ngay_sinh=date(1980, 1, 1),
            ngay_chet=date(2020, 2, 3),
            so_giay_to="123",
            ngay_cap=date(2021, 4, 5),
            dia_chi="Ha Noi",
        )
        db = DummyDb(customer=customer)

        quick_update(
            cid=1,
            ho_ten=None,
            gioi_tinh=None,
            ngay_sinh=None,
            ngay_chet=None,
            so_giay_to=None,
            ngay_cap=None,
            dia_chi=None,
            db=db,
        )
        self.assertEqual(customer.ho_ten, "Existing")
        self.assertEqual(customer.gioi_tinh, "Nam")
        self.assertEqual(customer.ngay_sinh, date(1980, 1, 1))
        self.assertEqual(customer.ngay_chet, date(2020, 2, 3))
        self.assertEqual(customer.so_giay_to, "123")
        self.assertEqual(customer.ngay_cap, date(2021, 4, 5))
        self.assertEqual(customer.dia_chi, "Ha Noi")

        quick_update(
            cid=1,
            ho_ten=None,
            gioi_tinh=" ",
            ngay_sinh=" ",
            ngay_chet=" ",
            so_giay_to=" ",
            ngay_cap=" ",
            dia_chi=" ",
            db=db,
        )
        self.assertEqual(customer.ho_ten, "Existing")
        self.assertIsNone(customer.gioi_tinh)
        self.assertIsNone(customer.ngay_sinh)
        self.assertIsNone(customer.ngay_chet)
        self.assertIsNone(customer.so_giay_to)
        self.assertIsNone(customer.ngay_cap)
        self.assertIsNone(customer.dia_chi)

    def test_quick_update_invalid_nonempty_values_preserve_existing_fields(self):
        customer = Customer(
            id=1,
            ho_ten="Existing",
            gioi_tinh="Nam",
            ngay_sinh=date(1980, 1, 1),
            ngay_chet=date(2020, 2, 3),
            so_giay_to="123",
            ngay_cap=date(2021, 4, 5),
            dia_chi="Ha Noi",
        )
        db = DummyDb(customer=customer)

        response = quick_update(
            cid=1,
            ho_ten=None,
            gioi_tinh="Unknown",
            ngay_sinh="not-a-date",
            ngay_chet="not-a-date",
            so_giay_to=None,
            ngay_cap="not-a-date",
            dia_chi=None,
            db=db,
        )

        self.assertTrue(json.loads(response.body)["ok"])
        self.assertEqual(customer.gioi_tinh, "Nam")
        self.assertEqual(customer.ngay_sinh, date(1980, 1, 1))
        self.assertEqual(customer.ngay_chet, date(2020, 2, 3))
        self.assertEqual(customer.ngay_cap, date(2021, 4, 5))

    def test_quick_update_duplicate_document_returns_existing_error_shape(self):
        db = DummyDb(
            customer=Customer(id=1, ho_ten="Existing", so_giay_to="123"),
            commit_error=IntegrityError("", {}, Exception()),
        )

        response = quick_update(
            cid=1,
            ho_ten=None,
            gioi_tinh=None,
            ngay_sinh=None,
            ngay_chet=None,
            so_giay_to="456",
            ngay_cap=None,
            dia_chi=None,
            db=db,
        )

        self.assertEqual(response.status_code, 400)
        payload = json.loads(response.body)
        self.assertFalse(payload["ok"])
        self.assertIn("error", payload)

    def test_inline_create_duplicate_document_preserves_omitted_fields_and_clears_empty_nullable_fields(self):
        customer = Customer(
            id=1,
            ho_ten="Existing",
            gioi_tinh="Nam",
            ngay_sinh=date(1980, 1, 1),
            ngay_chet=date(2020, 2, 3),
            so_giay_to="123",
            ngay_cap=date(2021, 4, 5),
            dia_chi="Ha Noi",
        )
        db = DummyDb(query_result=customer)

        response = inline_create(
            ho_ten="Updated",
            gioi_tinh=None,
            ngay_sinh=None,
            ngay_chet=None,
            so_giay_to="123",
            ngay_cap=None,
            dia_chi=None,
            db=db,
        )
        payload = json.loads(response.body)
        self.assertTrue(payload["ok"])
        self.assertTrue(payload["updated"])
        self.assertEqual(customer.ho_ten, "Updated")
        self.assertEqual(customer.gioi_tinh, "Nam")
        self.assertEqual(customer.ngay_sinh, date(1980, 1, 1))
        self.assertEqual(customer.ngay_chet, date(2020, 2, 3))
        self.assertEqual(customer.so_giay_to, "123")
        self.assertEqual(customer.ngay_cap, date(2021, 4, 5))
        self.assertEqual(customer.dia_chi, "Ha Noi")
        self.assertEqual(db.added, [])

        response = inline_create(
            ho_ten="Updated",
            gioi_tinh=" ",
            ngay_sinh=" ",
            ngay_chet=" ",
            so_giay_to="123",
            ngay_cap=" ",
            dia_chi=" ",
            db=db,
        )
        self.assertTrue(json.loads(response.body)["updated"])
        self.assertIsNone(customer.gioi_tinh)
        self.assertIsNone(customer.ngay_sinh)
        self.assertIsNone(customer.ngay_chet)
        self.assertEqual(customer.so_giay_to, "123")
        self.assertIsNone(customer.ngay_cap)
        self.assertIsNone(customer.dia_chi)

    def test_inline_create_invalid_nonempty_values_preserve_existing_fields(self):
        customer = Customer(
            id=1,
            ho_ten="Existing",
            gioi_tinh="Nam",
            ngay_sinh=date(1980, 1, 1),
            ngay_chet=date(2020, 2, 3),
            so_giay_to="123",
            ngay_cap=date(2021, 4, 5),
            dia_chi="Ha Noi",
        )
        db = DummyDb(query_result=customer)

        response = inline_create(
            ho_ten="Updated",
            gioi_tinh="Unknown",
            ngay_sinh="not-a-date",
            ngay_chet="not-a-date",
            so_giay_to="123",
            ngay_cap="not-a-date",
            dia_chi=None,
            db=db,
        )

        self.assertTrue(json.loads(response.body)["updated"])
        self.assertEqual(customer.ho_ten, "Updated")
        self.assertEqual(customer.gioi_tinh, "Nam")
        self.assertEqual(customer.ngay_sinh, date(1980, 1, 1))
        self.assertEqual(customer.ngay_chet, date(2020, 2, 3))
        self.assertEqual(customer.ngay_cap, date(2021, 4, 5))

    def test_inline_create_still_requires_name(self):
        response = inline_create(ho_ten=None, db=DummyDb())

        self.assertEqual(response.status_code, 400)
        self.assertEqual(json.loads(response.body), {"ok": False, "errors": {"ho_ten": "Bat buoc"}})

    def test_template_formats_date_columns_as_text(self):
        async def collect_response_body():
            chunks = []
            response = download_template()
            async for chunk in response.body_iterator:
                chunks.append(chunk)
            return b"".join(chunks)

        body = asyncio.run(collect_response_body())
        wb = openpyxl.load_workbook(BytesIO(body))
        ws = wb.active

        for idx, field in enumerate(EXCEL_COLUMNS, start=1):
            if field in {"ngay_sinh", "ngay_chet", "ngay_cap"}:
                self.assertEqual(ws.cell(row=2, column=idx).number_format, "@")


if __name__ == "__main__":
    unittest.main()
