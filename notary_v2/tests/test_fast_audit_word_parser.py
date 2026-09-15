from pathlib import Path

import pytest
from docx import Document

from services.fast_audit.word_parser import WordParser


@pytest.fixture
def parser() -> WordParser:
    return WordParser()


def _make_docx(tmp_path: Path, paragraphs: list[str], tables: list[list[str]] | None = None) -> Path:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if tables:
        for rows in tables:
            table = doc.add_table(rows=len(rows), cols=1)
            for idx, cell_text in enumerate(rows):
                table.cell(idx, 0).text = cell_text
    path = tmp_path / "test.docx"
    doc.save(str(path))
    return path


def test_parser_extracts_declarant_and_relative_names(parser: WordParser, tmp_path: Path):
    path = _make_docx(tmp_path, [
        "VĂN BẢN TỪ CHỐI NHẬN DI SẢN",
        "Tôi là: Nguyễn Văn Sơn, sinh năm 1980.",
        "Mẹ đẻ tôi là: Nguyễn Thị Ba.",
        "Ngườ để lại di sản là: Nguyễn Văn A.",
    ])
    doc = parser.parse(path)
    assert doc.template_type == "van_ban_tu_choi"
    values = {f.name: f.value for f in doc.fields}
    assert values.get("declarant_name") == "Nguyễn Văn Sơn"
    assert values.get("mother_name") == "Nguyễn Thị Ba"
    assert values.get("deceased_name") == "Nguyễn Văn A"


def test_parser_drops_notary_clause(parser: WordParser, tmp_path: Path):
    path = _make_docx(tmp_path, [
        "Tôi là: Nguyễn Văn Sơn.",
        "LỜI CHỨNG CỦA CÔNG CHỨNG VIÊN",
        "Tôi, công chứng viên Nguyễn Văn X, xác nhận ...",
    ])
    doc = parser.parse(path)
    full = " ".join(doc.paragraphs)
    assert "LỜI CHỨNG" not in full
    # No fields should be extracted from the notary clause.
    assert "Nguyễn Văn X" not in {f.value for f in doc.fields}


def test_placeholder_mode_detected(parser: WordParser, tmp_path: Path):
    path = _make_docx(tmp_path, [
        "[Họ tên ngườ từ chối]",
        "Tôi là: [Họ tên ngườ từ chối]",
    ])
    doc = parser.parse(path)
    assert doc.mode == "placeholder_template"
    assert any(f.name == "Họ tên ngườ từ chối" for f in doc.fields)


def test_parser_reads_tables(parser: WordParser, tmp_path: Path):
    path = _make_docx(tmp_path, [], tables=[["Họ tên"], ["Nguyễn Văn B"]])
    doc = parser.parse(path)
    assert any("Nguyễn Văn B" in f.value for f in doc.fields)


def test_parser_preserves_table_order_before_notary_clause(parser: WordParser, tmp_path: Path):
    docx = Document()
    table = docx.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Tôi là: Nguyễn Văn Sơn."
    docx.add_paragraph("LỜI CHỨNG CỦA CÔNG CHỨNG VIÊN")
    docx.add_paragraph("Tôi là: Nguyễn Văn X.")
    path = tmp_path / "ordered.docx"
    docx.save(str(path))

    doc = parser.parse(path)

    assert any(field.value == "Nguyễn Văn Sơn" for field in doc.fields)
    assert "Nguyễn Văn X" not in {field.value for field in doc.fields}


def test_parser_empty_and_heading_safe(parser: WordParser, tmp_path: Path):
    docx = Document()
    docx.add_paragraph("")
    docx.add_heading("Tiêu đề", level=1)
    docx.add_paragraph("Tôi là: Trần Thị C.")
    path = tmp_path / "heading.docx"
    docx.save(str(path))
    doc = parser.parse(path)
    assert any(f.value == "Trần Thị C" for f in doc.fields)
