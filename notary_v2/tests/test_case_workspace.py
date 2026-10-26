"""Tests for services.case_workspace — Storage B cutover (issue 13).

Rows case-owned (`case_person_rows`/`case_asset_rows`/`case_asset_land_rows`)
là SOT; `case_state_json` schemaVersion 4 chỉ còn owner_row_id/case scalar/
Diagram/docx_import. Trạng thái chỉ draft/done, không khóa soạn.

Wire shapes follow contracts/notary-case-drafting/*.schema.json §13:
stage `{owner_row_id?, people[], assets[]}`, diagram_state v3
`{version:3, domain, nodes[]}` — is_primary/isLandOwner/willReceive bị cấm
trên wire v2 (vị trí = index mảng; owner = stage.owner_row_id).
"""
import json
import re
import sqlite3
import uuid
from datetime import date
from pathlib import Path

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import database
from models import (
    CaseAssetRow,
    CasePersonRow,
    Customer,
    InheritanceCase,
    InheritanceCaseProperty,
    InheritanceParticipant,
    Property,
)
from services.case_workspace import (
    SCHEMA_VERSION,
    CaseWorkspaceService,
    WorkspaceError,
)


UUID4_RE = re.compile(
    r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-4[0-9a-fA-F]{3}"
    r"-[89abAB][0-9a-fA-F]{3}-[0-9a-fA-F]{12}$"
)
TP_IDS = [f"p{i}" for i in range(1, 31)]


# ---------------------------------------------------------------- fixtures


@pytest.fixture()
def session_factory(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'workspace.db'}")
    database.Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    yield factory
    engine.dispose()


@pytest.fixture()
def db(session_factory):
    session = session_factory()
    yield session
    session.close()


def _make_case(db, *, loai_van_ban="khai_nhan", trang_thai="draft",
               with_participant=False):
    """Case đã lưu với rows Storage B (master Customer/Property là identity)."""
    deceased = Customer(
        ho_ten="Nguyễn Văn An", gioi_tinh="Nam",
        ngay_sinh=date(1950, 1, 1), ngay_chet=date(2011, 5, 15),
        so_giay_to="001234567890", dia_chi="Số 1 phố Huế, Hà Nội",
        name_anchor="nguyen van an",
        document_anchor="001234567890",
    )
    prop = Property(
        so_serial="DD123456", dia_chi="Thửa 123, tờ 45, phường Bạch Mai",
        so_vao_so="CS 12345", so_thua_dat="123", so_to_ban_do="45",
    )
    db.add_all([deceased, prop])
    db.flush()
    case = InheritanceCase(
        nguoi_chet_id=deceased.id, tai_san_id=prop.id,
        ngay_lap_ho_so=date(2026, 9, 1), loai_van_ban=loai_van_ban,
        trang_thai=trang_thai,
    )
    db.add(case)
    db.flush()
    owner_row_id = str(uuid.uuid4())
    asset_row_id = str(uuid.uuid4())
    db.add(CasePersonRow(
        row_id=owner_row_id, case_id=case.id, customer_id=deceased.id,
        vitri=1, ho_ten=deceased.ho_ten, gioi_tinh="Nam",
        ngay_sinh="1950", ngay_chet="2011-05-15",
        so_giay_to="001234567890", dia_chi="Số 1 phố Huế, Hà Nội"))
    db.add(CaseAssetRow(
        row_id=asset_row_id, case_id=case.id, property_id=prop.id,
        vitri=1, so_serial="DD123456", so_vao_so="CS 12345",
        so_thua_dat="123", so_to_ban_do="45",
        dia_chi="Thửa 123, tờ 45, phường Bạch Mai",
        loai_so="GCN QSDĐ", hinh_thuc_su_dung="Sử dụng riêng",
        thoi_han="Lâu dài", nguon_goc="Cấp đổi",
        ngay_cap=date(2005, 9, 30),
        co_quan_cap="UBND quận Hai Bà Trưng", land_rows_present=True))
    db.add(InheritanceCaseProperty(
        case_id=case.id, case_asset_row_id=asset_row_id,
        property_id=prop.id, is_primary=True))
    payload = {"schemaVersion": 4,
               "case": {key: None for key in (
                   "tenvanban", "nhomhopdong", "loaivanbannhandien",
                   "socongchung", "ngaycongchung", "congchungvien", "thuky")}}
    payload["diagram"] = {
        "state": {"version": 3,
                  "domain": ("two_party"
                             if loai_van_ban in ("chuyen_nhuong",
                                                 "tang_cho", "cho_thue",
                                                 "dat_coc")
                             else "inheritance"),
                  "nodes": []},
        "render_model": None}
    if payload["diagram"]["state"]["domain"] == "inheritance":
        payload["owner_row_id"] = owner_row_id
    case.case_state_json = json.dumps(payload, ensure_ascii=False)
    heirs = []
    if with_participant:
        heir = Customer(
            ho_ten="Nguyễn Thị Bình", gioi_tinh="Nữ",
            ngay_sinh=date(1980, 3, 2), dia_chi="Số 1 phố Huế, Hà Nội",
            name_anchor="nguyen thi binh",
            document_anchor="001234567891",
        )
        db.add(heir)
        db.flush()
        heir_row_id = str(uuid.uuid4())
        db.add(CasePersonRow(
            row_id=heir_row_id, case_id=case.id, customer_id=heir.id,
            vitri=2, ho_ten=heir.ho_ten, gioi_tinh="Nữ",
            ngay_sinh="1980-03-02", dia_chi="Số 1 phố Huế, Hà Nội"))
        db.add(InheritanceParticipant(
            ho_so_id=case.id, case_person_row_id=heir_row_id,
            customer_id=heir.id, vai_tro="Con"))
        heirs.append(heir)
    db.commit()
    return case, deceased, prop, heirs


def _person_row(**overrides):
    row = {
        "row_id": str(uuid.uuid4()),
        "entity_id": None,
        "ho_ten": "Người Mới",
        "gioi_tinh": "Nam",
        "ngay_sinh": "1980-01-01",
        "ngay_chet": None,
        "so_giay_to": None,
        "ngay_cap": None,
        "noi_cap": None,
        "dia_chi": "xã test",
        "place_of_origin": None,
    }
    row.update(overrides)
    return row


def _asset_row(**overrides):
    """asset_row v2 — KHÔNG is_primary (vị trí = index mảng)."""
    row = {
        "row_id": str(uuid.uuid4()),
        "entity_id": None,
        "so_serial": "EE123456",
        "so_vao_so": "CS 99999",
        "so_thua_dat": "99",
        "so_to_ban_do": "12",
        "dia_chi": "Thửa 99, xã Yên Sở",
        "loai_so": "GCN QSDĐ",
        "hinh_thuc_su_dung": "Sử dụng riêng",
        "thoi_han": "Lâu dài",
        "nguon_goc": "Cấp đổi",
        "ngay_cap": "2005-09-30",
        "co_quan_cap": "UBND quận Hai Bà Trưng",
        "land_rows": [{"loai_dat": "ODT", "dien_tich": 85.5,
                       "thoi_han": "Lâu dài"}],
    }
    row.update(overrides)
    return row


def _service(db):
    return CaseWorkspaceService(db)


def _stage(people, assets, owner_row_id="__omit__"):
    """stage_v2 — inheritance cần owner_row_id; "__omit__" = không emit."""
    st = {"people": people, "assets": assets}
    if owner_row_id != "__omit__":
        st["owner_row_id"] = owner_row_id
    return st


def _commit(db, case_id, base_revision, people, assets, owner_row_id=None):
    """Commit stage inheritance — owner mặc định people[0].row_id."""
    if owner_row_id is None:
        owner_row_id = people[0]["row_id"]
    return _service(db).commit_stage(
        case_id, base_revision, _stage(people, assets, owner_row_id))


def _node(nid, person_id=None, parents=(), spouse=None, own=(),
          receive=(), hidden=False, deleted=False):
    """diagram node v3 (§13.4) — ownPositions/receivePositions ⊆ {1,2,3}."""
    return {"id": nid, "personId": person_id,
            "parentSlotIds": list(parents),
            "spouseSlotIds": [spouse] if spouse else [],
            "spouseSlotId": spouse,
            "spouseSuggestionId": None, "spouseUndetermined": False,
            "ownPositions": list(own), "receivePositions": list(receive),
            "hidden": hidden, "deleted": deleted}


def _v3_state(nodes, domain="inheritance"):
    return {"version": 3, "domain": domain, "nodes": nodes}


def _write_state(case, db, state, *, owner_row_id=None):
    payload = json.loads(case.case_state_json or "{}")
    payload["diagram"] = {"state": state, "render_model": None}
    if owner_row_id is not None or "owner_row_id" in payload:
        payload["owner_row_id"] = owner_row_id
    case.case_state_json = json.dumps(payload, ensure_ascii=False)
    db.commit()


# ------------------------------------------------------------------- get()


def test_get_reads_rows_and_case_status(db):
    case, deceased, prop, heirs = _make_case(db, with_participant=True)

    data = _service(db).get(case.id)

    assert data["schema_version"] == SCHEMA_VERSION
    assert data["backend_mode"] == "real"
    assert data["case"] == {
        "id": case.id,
        "case_type": "inheritance",
        "document_type": "khai_nhan",
        "status": "draft",
        "revision": 1,
        "ngay_lap_ho_so": "2026-09-01",
        "noi_niem_yet": None,
        "nguoi_nhan_uy_quyen": None,
        "nguoi_nhan_uy_quyen_id": None,
        "noi_dung_viec": None,
        **{key: None for key in (
            "tenvanban", "nhomhopdong", "loaivanbannhandien", "socongchung",
            "ngaycongchung", "congchungvien", "thuky", "chudat")},
        "ghi_chu": None,
    }
    assert "locked" not in data["case"]
    people = data["stage"]["people"]
    assert [p["entity_id"] for p in people] == [deceased.id, heirs[0].id]
    assert people[0]["ho_ten"] == "Nguyễn Văn An"
    assert people[0]["ngay_chet"] == "2011-05-15"
    assert people[0]["ngay_sinh"] == "1950"
    assert people[1]["ho_ten"] == "Nguyễn Thị Bình"
    for row in people:
        assert UUID4_RE.match(row["row_id"])
    assets = data["stage"]["assets"]
    assert len(assets) == 1
    assert assets[0]["entity_id"] == prop.id
    assert "is_primary" not in assets[0]
    assert assets[0]["so_serial"] == "DD123456"
    assert assets[0]["ngay_cap"] == "2005-09-30"
    assert data["stage"]["owner_row_id"] == case.person_rows[0].row_id
    assert data["diagram"]["domain"] == "inheritance"
    assert data["diagram"]["state"] == {
        "version": 3, "domain": "inheritance", "nodes": []}
    assert data["diagram"]["render_model"] is None
    assert data["capabilities"] == {
        "intake": ["image", "pdf", "docx", "xlsx", "text"],
        "diagram": True,
        "word_export": True,
    }


def test_get_does_not_read_or_write_legacy_stage_json(db):
    """Legacy JSON stage/engineState không còn là nguồn; GET không migrate."""
    case, _deceased, _prop, _h = _make_case(db)
    legacy = json.loads(case.case_state_json)
    legacy["stage"] = [{"id": "1", "row_id": str(uuid.uuid4()),
                        "ho_ten": "Legacy"}]
    legacy["diagram"]["engineState"] = {
        "version": 2, "nodes": [{"id": "owner", "personId": "1"}]}
    case.case_state_json = json.dumps(legacy)
    db.commit()
    before = db.get(InheritanceCase, case.id).case_state_json

    data = _service(db).get(case.id)

    assert [p["ho_ten"] for p in data["stage"]["people"]] == [
        "Nguyễn Văn An"]
    assert data["diagram"]["state"]["nodes"] == []
    assert db.get(InheritanceCase, case.id).case_state_json == before


def test_get_row_ids_stable_across_reloads(db, session_factory):
    case, _deceased, _prop, _h = _make_case(db, with_participant=True)

    first = _service(db).get(case.id)
    first_ids = [p["row_id"] for p in first["stage"]["people"]]
    first_asset_ids = [a["row_id"] for a in first["stage"]["assets"]]

    db2 = session_factory()
    try:
        second = _service(db2).get(case.id)
    finally:
        db2.close()

    assert [p["row_id"] for p in second["stage"]["people"]] == first_ids
    assert [a["row_id"] for a in second["stage"]["assets"]] == first_asset_ids


def test_get_reads_second_asset_row(db):
    case, _deceased, _prop, _h = _make_case(db)
    other = Property(so_serial="EE654321")
    db.add(other)
    db.flush()
    row_id = str(uuid.uuid4())
    db.add(CaseAssetRow(
        row_id=row_id, case_id=case.id, property_id=other.id, vitri=2,
        so_serial="EE654321", dia_chi="Thửa 77, tờ 12"))
    db.add(InheritanceCaseProperty(
        case_id=case.id, case_asset_row_id=row_id, property_id=other.id,
        is_primary=False))
    db.commit()

    assets = {a["so_serial"]: a for a in _service(db).get(case.id)["stage"]["assets"]}
    assert set(assets) == {"DD123456", "EE654321"}
    assert assets["EE654321"]["land_rows"] is None
    for asset in assets.values():
        assert "is_primary" not in asset


def test_get_two_party_case_state_and_capabilities(db):
    case, _d, _p, _h = _make_case(db, loai_van_ban="chuyen_nhuong")

    data = _service(db).get(case.id)

    assert data["case"]["case_type"] == "two_party"
    assert data["case"]["document_type"] == "chuyen_nhuong"
    assert "owner_row_id" not in data["stage"]
    st = data["diagram"]["state"]
    assert st["version"] == 3 and st["domain"] == "two_party"
    assert [n["id"] for n in st["nodes"]] == TP_IDS
    assert all(set(n) == {"id", "personId", "hidden", "deleted"}
               for n in st["nodes"])
    assert data["capabilities"]["word_export"] is True


def test_get_unsupported_case_type_disables_capabilities(db):
    case, _d, _p, _h = _make_case(db, loai_van_ban="khac")

    data = _service(db).get(case.id)

    assert data["case"]["case_type"] not in ("inheritance", "two_party")
    assert data["capabilities"]["intake"] == []
    assert data["capabilities"]["diagram"] is False
    assert data["capabilities"]["word_export"] is False
    assert "owner_row_id" in data["stage"]


def test_get_supports_thoa_thuan_document_type(db):
    case, _d, _p, _h = _make_case(db, loai_van_ban="thoa_thuan")

    data = _service(db).get(case.id)

    assert data["case"]["case_type"] == "inheritance"
    assert data["case"]["document_type"] == "thoa_thuan"


def test_get_missing_case_raises_case_not_found(db):
    with pytest.raises(WorkspaceError) as exc:
        _service(db).get(9999)
    assert exc.value.code == "case_not_found"


def test_get_done_case_readable_without_locked_key(db):
    case, _d, _p, _h = _make_case(db, trang_thai="done")

    data = _service(db).get(case.id)

    assert data["case"]["status"] == "done"
    assert "locked" not in data["case"]


# ------------------------------------------------------------- commit_stage


def test_commit_stage_increments_revision_and_resolves_identity(db):
    case, deceased, _prop, _h = _make_case(db)
    people = [
        _person_row(row_id=case.person_rows[0].row_id,
                    entity_id=deceased.id,
                    ho_ten="Nguyễn Văn An", ngay_sinh="1950",
                    ngay_chet="2011-05-15",
                    so_giay_to="001234567890",
                    dia_chi="Số 1 phố Huế, Hà Nội"),
        _person_row(ho_ten="Nguyễn Văn Cường", gioi_tinh="Nam",
                    ngay_sinh="1980-07-20",
                    so_giay_to="001080012345"),
    ]
    assets = [_asset_row(so_serial="DD123456")]

    data = _commit(db, case.id, 1, people, assets)

    assert data["schema_version"] == SCHEMA_VERSION
    assert data["revision"] == 2
    assert db.get(InheritanceCase, case.id).workspace_revision == 2
    assert data["stage"]["owner_row_id"] == people[0]["row_id"]
    assert data["diagram"]["render_model"] is not None
    assert data["diagram"]["render_model"]["engineVersion"] == 2
    committed_people = data["stage"]["people"]
    assert committed_people[0]["entity_id"] == deceased.id
    assert committed_people[0]["row_id"] == people[0]["row_id"]
    assert committed_people[1]["entity_id"] is not None
    assert committed_people[1]["row_id"] == people[1]["row_id"]
    assert "is_primary" not in data["stage"]["assets"][0]
    stored = db.query(Customer).filter_by(
        so_giay_to="001080012345").one()
    assert stored.ho_ten == "Nguyễn Văn Cường"


def test_commit_stage_owner_pointer_required(db):
    case, deceased, prop, _h = _make_case(db)
    people = [_person_row(entity_id=deceased.id)]
    assets = [_asset_row(entity_id=prop.id, so_serial="DD123456")]

    for bad in (None, str(uuid.uuid4()), "not-uuid"):
        with pytest.raises(WorkspaceError) as exc:
            _service(db).commit_stage(
                case.id, 1, _stage(people, assets, owner_row_id=bad))
        assert exc.value.code == "workspace_owner_required", bad
    with pytest.raises(WorkspaceError) as exc:
        _service(db).commit_stage(case.id, 1, _stage(people, assets))
    assert exc.value.code == "workspace_owner_required"


def test_commit_stage_owner_syncs_node(db):
    case, deceased, prop, heirs = _make_case(db, with_participant=True)
    heir = heirs[0]
    owner_row = case.person_rows[0].row_id
    heir_row = case.person_rows[1].row_id
    _write_state(case, db, _v3_state([
        _node("owner", owner_row, own=[1]),
        _node("spouse", heir_row, receive=[1])]), owner_row_id=owner_row)

    people = [
        _person_row(row_id=owner_row, entity_id=deceased.id,
                    ho_ten="Nguyễn Văn An", ngay_chet="2011-05-15"),
        _person_row(row_id=heir_row, entity_id=heir.id,
                    ho_ten="Nguyễn Thị Bình", gioi_tinh="Nữ",
                    ngay_sinh="1980-03-02"),
    ]
    data = _service(db).commit_stage(
        case.id, 1,
        _stage(people, [_asset_row(so_serial="DD123456")],
               owner_row_id=heir_row))
    owner = next(n for n in data["diagram"]["state"]["nodes"]
                 if n["id"] == "owner")
    assert owner["personId"] == heir_row
    assert data["stage"]["owner_row_id"] == heir_row


def test_commit_stage_new_rows_persist_then_reload(db, session_factory):
    case, deceased, _prop, _h = _make_case(db)
    people = [_person_row(row_id=case.person_rows[0].row_id,
                          entity_id=deceased.id, ho_ten="Nguyễn Văn An"),
              _person_row(ho_ten="Người Hoàn Toàn Mới")]
    assets = [_asset_row(so_serial="DD123456")]

    committed = _commit(db, case.id, 1, people, assets)

    db2 = session_factory()
    try:
        reloaded = _service(db2).get(case.id)
    finally:
        db2.close()

    assert reloaded["case"]["revision"] == 2
    assert reloaded["stage"]["owner_row_id"] == \
        committed["stage"]["owner_row_id"]
    assert [p["row_id"] for p in reloaded["stage"]["people"]] == [
        p["row_id"] for p in committed["stage"]["people"]]
    assert reloaded["stage"]["people"][1]["ho_ten"] == "Người Hoàn Toàn Mới"
    assert reloaded["diagram"]["render_model"] == \
        committed["diagram"]["render_model"]


def test_commit_stage_invalid_row_rolls_back_everything(db, session_factory):
    case, deceased, _prop, _h = _make_case(db)
    before_state = case.case_state_json
    customers_before = db.query(Customer).count()

    people = [
        _person_row(row_id=case.person_rows[0].row_id,
                    entity_id=deceased.id, ho_ten="Nguyễn Văn An"),
        _person_row(ho_ten="Dòng Hợp Lệ", so_giay_to="009999999999"),
        _person_row(ho_ten="Dòng Sai Giới Tính", gioi_tinh="Khác"),
    ]
    bad_row_id = people[2]["row_id"]

    with pytest.raises(WorkspaceError) as exc:
        _commit(db, case.id, 1, people,
                [_asset_row(so_serial="DD123456")])

    assert exc.value.code == "stage_validation_error"
    field_errors = exc.value.details["field_errors"]
    assert any(
        e["row_id"] == bad_row_id and e["field"] == "gioi_tinh"
        and e["code"] == "invalid_enum"
        for e in field_errors
    )

    db2 = session_factory()
    try:
        assert db2.query(Customer).count() == customers_before
        reloaded = db2.get(InheritanceCase, case.id)
        assert reloaded.workspace_revision == 1
        assert reloaded.case_state_json == before_state
    finally:
        db2.close()


def test_commit_stage_collects_multiple_field_errors(db):
    case, deceased, _prop, _h = _make_case(db)
    bad = _person_row(row_id=case.person_rows[0].row_id,
                      entity_id=deceased.id, ho_ten="",
                      ngay_sinh="15/05/2011")

    with pytest.raises(WorkspaceError) as exc:
        _commit(db, case.id, 1, [bad], [])

    errors = exc.value.details["field_errors"]
    by_field = {e["field"]: e for e in errors}
    assert by_field["ho_ten"]["code"] == "required"
    assert by_field["ngay_sinh"]["code"] == "invalid_date"
    assert all(e["row_id"] == bad["row_id"] for e in errors)


def test_commit_stage_rejects_duplicate_row_id(db):
    case, deceased, _prop, _h = _make_case(db)
    shared = str(uuid.uuid4())
    people = [_person_row(row_id=shared, ho_ten="A"),
              _person_row(row_id=shared, ho_ten="B")]

    with pytest.raises(WorkspaceError) as exc:
        _commit(db, case.id, 1, people, [])

    assert exc.value.code == "stage_validation_error"
    assert any(e["code"] == "duplicate_row_id" and e["row_id"] == shared
               for e in exc.value.details["field_errors"])


def test_commit_stage_rejects_is_primary_and_asset_limit(db):
    case, deceased, _prop, _h = _make_case(db)
    people = [_person_row(row_id=case.person_rows[0].row_id,
                          entity_id=deceased.id)]

    bad_asset = _asset_row(so_serial="DD123456")
    bad_asset["is_primary"] = True
    with pytest.raises(WorkspaceError) as exc:
        _commit(db, case.id, 1, people, [bad_asset])
    assert exc.value.code == "validation_error"

    # S7 owner 23/10/2026: trần server = 5 tài sản.
    five = [_asset_row(so_serial=f"GG{i:06d}") for i in range(5)]
    data = _commit(db, case.id, 1, people, five)
    assert len(data["stage"]["assets"]) == 5

    six = five + [_asset_row(so_serial="GG999999")]
    with pytest.raises(WorkspaceError) as exc2:
        _commit(db, case.id, data["revision"], people, six)
    assert any(e["code"] == "asset_limit"
               for e in exc2.value.details["field_errors"])


@pytest.mark.parametrize("serial", ["TEST-RAW/001", "", "  RAW  "])
def test_commit_stage_accepts_verbatim_so_serial(db, session_factory, serial):
    case, deceased, _p, _h = _make_case(db)
    data = _commit(db, case.id, 1,
                   [_person_row(row_id=case.person_rows[0].row_id,
                                entity_id=deceased.id)],
                   [_asset_row(so_serial=serial)])

    assert data["stage"]["assets"][0]["so_serial"] == serial
    fresh = session_factory()
    try:
        reopened = CaseWorkspaceService(fresh).get(case.id)
    finally:
        fresh.close()
    assert reopened["stage"]["assets"][0]["so_serial"] == serial
    # serial rỗng/toàn whitespace = thiếu neo → identity null
    if not serial.strip():
        assert data["stage"]["assets"][0]["entity_id"] is None


def test_commit_stage_same_serial_shares_identity_keeps_two_rows(db):
    """Cùng serial nguyên văn → cùng Property identity, hai dòng độc lập."""
    case, deceased, _prop, _h = _make_case(db)
    first = _asset_row(so_serial="SYN B 001", so_thua_dat="11",
                       so_to_ban_do="2", dia_chi="Thửa 11")
    second = _asset_row(so_serial="SYN B 001", so_thua_dat="22",
                        so_to_ban_do="2", dia_chi="Thửa 22")

    data = _commit(db, case.id, 1,
                   [_person_row(row_id=case.person_rows[0].row_id,
                                entity_id=deceased.id)],
                   [first, second])

    assets = data["stage"]["assets"]
    assert [a["so_serial"] for a in assets] == ["SYN B 001", "SYN B 001"]
    assert assets[0]["entity_id"] == assets[1]["entity_id"]
    assert db.query(Property).filter_by(so_serial="SYN B 001").count() == 1


def test_commit_stage_serial_is_verbatim_key(db):
    """Khác space/case = identity khác (không chuẩn hóa serial)."""
    case, deceased, _prop, _h = _make_case(db)
    first = _asset_row(so_serial="AB 001", dia_chi="A")
    second = _asset_row(so_serial="AB001", dia_chi="B")

    data = _commit(db, case.id, 1,
                   [_person_row(row_id=case.person_rows[0].row_id,
                                entity_id=deceased.id)],
                   [first, second])

    assets = data["stage"]["assets"]
    assert assets[0]["entity_id"] != assets[1]["entity_id"]
    assert db.query(Property).filter_by(so_serial="AB 001").count() == 1
    assert db.query(Property).filter_by(so_serial="AB001").count() == 1


def test_commit_stage_rejects_non_string_so_serial(db):
    case, deceased, _p, _h = _make_case(db)
    asset = _asset_row(so_serial=None)

    with pytest.raises(WorkspaceError) as exc:
        _commit(db, case.id, 1,
                [_person_row(row_id=case.person_rows[0].row_id,
                             entity_id=deceased.id)],
                [asset])

    assert any(e["field"] == "so_serial" and e["code"] == "invalid_type"
               for e in exc.value.details["field_errors"])


def test_commit_stage_done_case_remains_done(db):
    case, deceased, _prop, _h = _make_case(db, trang_thai="done")

    data = _commit(db, case.id, 1,
                   [_person_row(row_id=case.person_rows[0].row_id,
                                entity_id=deceased.id)],
                   [_asset_row(so_serial="DD123456")])

    assert data["revision"] == 2
    assert db.get(InheritanceCase, case.id).trang_thai == "done"
    assert "locked" not in _service(db).get(case.id)["case"]


def test_commit_stage_stale_and_ahead_base_revision(db):
    case, deceased, _prop, _h = _make_case(db)

    with pytest.raises(WorkspaceError) as ahead:
        _service(db).commit_stage(case.id, 99, {})
    assert ahead.value.code == "workspace_conflict"
    assert ahead.value.details["server_revision"] == 1

    _commit(db, case.id, 1,
            [_person_row(row_id=case.person_rows[0].row_id,
                         entity_id=deceased.id)],
            [_asset_row(so_serial="DD123456")])

    with pytest.raises(WorkspaceError) as stale:
        _service(db).commit_stage(case.id, 1, {})
    assert stale.value.code == "workspace_conflict"
    assert stale.value.details["server_revision"] == 2


def test_commit_stage_unsupported_case_type(db):
    case, deceased, prop, _h = _make_case(db, loai_van_ban="khac")

    with pytest.raises(WorkspaceError) as exc:
        _commit(db, case.id, 1,
                [_person_row(row_id=case.person_rows[0].row_id,
                             entity_id=deceased.id)],
                [_asset_row(entity_id=prop.id, so_serial="DD123456")])

    assert exc.value.code == "case_type_unsupported"
    assert exc.value.details["case_type"] not in ("inheritance", "two_party")


def test_commit_stage_two_party_canonical_and_null_owner(db):
    case, deceased, _prop, _h = _make_case(db, loai_van_ban="chuyen_nhuong")
    people = [_person_row(row_id=case.person_rows[0].row_id,
                          entity_id=deceased.id, ho_ten="Bên A",
                          ngay_chet=None)]
    data = _service(db).commit_stage(
        case.id, 1,
        _stage(people, [_asset_row(so_serial="DD123456")]))

    assert data["revision"] == 2
    assert "owner_row_id" not in data["stage"]
    st = data["diagram"]["state"]
    assert st["domain"] == "two_party"
    assert [n["id"] for n in st["nodes"]] == TP_IDS
    assert data["diagram"]["render_model"]["status"] == "unsupported"
    assert data["diagram"]["render_model"]["allocations"] == {}
    db.refresh(case)
    assert case.nguoi_chet_id is None


def test_commit_stage_two_party_owner_row_id_forbidden(db):
    case, deceased, prop, _h = _make_case(db, loai_van_ban="tang_cho")
    people = [_person_row(row_id=case.person_rows[0].row_id,
                          entity_id=deceased.id)]
    stage = _stage(people, [_asset_row(entity_id=prop.id)],
                   owner_row_id=None)

    with pytest.raises(WorkspaceError) as exc:
        _service(db).commit_stage(case.id, 1, stage)
    assert exc.value.code == "validation_error"


def test_commit_stage_missing_case(db):
    with pytest.raises(WorkspaceError) as exc:
        _service(db).commit_stage(9999, 1, _stage([], []))
    assert exc.value.code == "case_not_found"


def test_commit_stage_prunes_diagram_and_reevaluates(db):
    case, deceased, _prop, heirs = _make_case(db, with_participant=True)
    heir = heirs[0]
    owner_row = case.person_rows[0].row_id
    heir_row = case.person_rows[1].row_id
    _write_state(case, db, _v3_state([
        _node("owner", owner_row, spouse="spouse", own=[1]),
        _node("spouse", heir_row, spouse="owner", receive=[1])]),
        owner_row_id=owner_row)

    data = _commit(db, case.id, 1, [
        _person_row(row_id=owner_row, entity_id=deceased.id,
                    ho_ten="Nguyễn Văn An", ngay_chet="2011-05-15"),
    ], [_asset_row(so_serial="DD123456")])

    nodes = data["diagram"]["state"]["nodes"]
    assert [n["id"] for n in nodes] == ["owner"]
    assert nodes[0]["spouseSlotId"] is None
    assert heir_row not in data["diagram"]["render_model"]["allocations"]
    persisted = json.loads(db.get(InheritanceCase, case.id).case_state_json)
    assert set(persisted) == {"schemaVersion", "case", "diagram",
                              "owner_row_id"}
    assert "stage" not in persisted and "assets" not in persisted
    assert db.get(InheritanceCase, case.id).engine_state_json is None


def test_commit_stage_prunes_out_of_range_positions(db):
    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    _write_state(case, db, _v3_state([
        _node("owner", owner_row, own=[1, 2, 3])]),
        owner_row_id=owner_row)

    data = _commit(db, case.id, 1,
                   [_person_row(row_id=owner_row, entity_id=deceased.id,
                                ho_ten="Nguyễn Văn An",
                                ngay_chet="2011-05-15")],
                   [_asset_row(so_serial="DD123456")])

    owner = next(n for n in data["diagram"]["state"]["nodes"]
                 if n["id"] == "owner")
    assert owner["ownPositions"] == [1]
    codes = [w["code"] for w in data["diagram"].get("warnings") or []]
    assert "diagram.selection_pruned" in codes


def test_commit_stage_rebuilds_participants_by_row_id(db):
    """Participant projection keyed case_person_row_id, identity nullable."""
    case, deceased, _prop, heirs = _make_case(db, with_participant=True)
    # Không còn participant cũ: test dựng lại projection từ Diagram; vai trò
    # nhập tay đã lưu (nếu có) được giữ theo §2.4, nên xóa seed để kiểm derived.
    db.query(InheritanceParticipant).filter_by(ho_so_id=case.id).delete()
    db.commit()
    db.expire(case)
    heir = heirs[0]
    owner_row = case.person_rows[0].row_id
    heir_row = case.person_rows[1].row_id
    _write_state(case, db, _v3_state([
        _node("owner", owner_row, spouse="spouse", own=[1]),
        _node("spouse", heir_row, spouse="owner", receive=[1])]),
        owner_row_id=owner_row)

    _commit(db, case.id, 1, [
        _person_row(row_id=owner_row, entity_id=deceased.id,
                    ho_ten="Nguyễn Văn An", ngay_chet="2011-05-15",
                    so_giay_to="001234567890"),
        _person_row(row_id=heir_row, entity_id=heir.id,
                    ho_ten="Nguyễn Thị Bình", gioi_tinh="Nữ",
                    ngay_sinh="1980-03-02",
                    so_giay_to="001234567891"),
    ], [_asset_row(so_serial="DD123456")])

    db.refresh(case)
    assert case.nguoi_chet_id == deceased.id
    parts = {p.case_person_row_id: p for p in case.participants}
    assert owner_row not in parts
    assert parts[heir_row].customer_id == heir.id
    assert parts[heir_row].vai_tro == "Vợ/Chồng"
    assert parts[heir_row].co_nhan_tai_san is True


def test_commit_stage_reprojects_derived_roles_to_latest_diagram(
        db, session_factory):
    """#45 TC1: đổi quan hệ sau lần Lưu đầu → Lưu lại → participants khớp
    sơ đồ mới và Word điền placeholder vợ/chồng theo người mới."""
    from services import word_engine

    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    first_spouse = str(uuid.uuid4())
    child = str(uuid.uuid4())
    people = [
        _person_row(row_id=owner_row, entity_id=deceased.id,
                    ho_ten="Nguyễn Văn An", ngay_chet="2011-05-15",
                    so_giay_to="001234567890"),
        _person_row(row_id=first_spouse, ho_ten="Trần Thị Vợ",
                    gioi_tinh="Nữ", ngay_sinh="1955-04-05"),
        _person_row(row_id=child, ho_ten="Nguyễn Văn Con",
                    gioi_tinh="Nam", ngay_sinh="1980-08-09"),
    ]

    _write_state(case, db, _v3_state([
        _node("owner", owner_row, spouse="spouse", own=[1]),
        _node("spouse", first_spouse, spouse="owner", receive=[1]),
        _node("child_1", child, parents=("owner", "spouse"))]),
        owner_row_id=owner_row)
    _commit(db, case.id, 1, people, [_asset_row(so_serial="DD123456")])

    db.refresh(case)
    parts = {p.case_person_row_id: p for p in case.participants}
    assert parts[first_spouse].vai_tro == "Vợ/Chồng"
    assert parts[child].vai_tro == "Con"
    assert parts[child].parent_case_person_row_id == owner_row

    # Đổi quan hệ: đưa người con lên vai trò vợ/chồng, người vợ cũ xuống con.
    _write_state(case, db, _v3_state([
        _node("owner", owner_row, spouse="spouse", own=[1]),
        _node("spouse", child, spouse="owner", receive=[1]),
        _node("child_1", first_spouse, parents=("owner", "spouse"))]),
        owner_row_id=owner_row)
    _commit(db, case.id, 2, people, [_asset_row(so_serial="DD123456")])

    db.refresh(case)
    parts = {p.case_person_row_id: p for p in case.participants}
    assert parts[child].vai_tro == "Vợ/Chồng"
    assert parts[child].hang_thua_ke == 1
    assert parts[first_spouse].vai_tro == "Con"
    assert parts[first_spouse].hang_thua_ke == 1
    assert parts[first_spouse].parent_case_person_row_id == owner_row

    fresh = session_factory()
    try:
        db_case = fresh.get(InheritanceCase, case.id)
        mapping = word_engine.build_template_mapping(db_case)
        assert mapping["[Họ tên chủ đất]"] == "Nguyễn Văn An"
        assert mapping["[Họ tên vợ/chồng]"] == "Nguyễn Văn Con"
    finally:
        fresh.close()


def test_commit_stage_preserves_manual_participant_role(db):
    """#45 TC2: vai trò gán tay (ngoài bộ suy-từ-sơ-đồ) không bị ghi đè."""
    from services import word_engine

    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    spouse = str(uuid.uuid4())
    child = str(uuid.uuid4())
    people = [
        _person_row(row_id=owner_row, entity_id=deceased.id,
                    ho_ten="Nguyễn Văn An", ngay_chet="2011-05-15",
                    so_giay_to="001234567890"),
        _person_row(row_id=spouse, ho_ten="Trần Thị Vợ", gioi_tinh="Nữ"),
    ]

    _write_state(case, db, _v3_state([
        _node("owner", owner_row, spouse="spouse", own=[1]),
        _node("spouse", spouse, spouse="owner", receive=[1])]),
        owner_row_id=owner_row)
    _commit(db, case.id, 1, people, [_asset_row(so_serial="DD123456")])

    db.refresh(case)
    manual = next(p for p in case.participants
                  if p.case_person_row_id == spouse)
    # Tác nghiệp gán tay trên dòng đã có (routers/participants edit).
    manual.vai_tro = "Người được ủy quyền"
    db.commit()

    # Thay đổi sơ đồ không liên quan: thêm một người con.
    _write_state(case, db, _v3_state([
        _node("owner", owner_row, spouse="spouse", own=[1]),
        _node("spouse", spouse, spouse="owner", receive=[1]),
        _node("child_1", child, parents=("owner", "spouse"), receive=[1])]),
        owner_row_id=owner_row)
    _commit(db, case.id, 2, [
        *people,
        _person_row(row_id=child, ho_ten="Nguyễn Văn Con",
                    gioi_tinh="Nam"),
    ], [_asset_row(so_serial="DD123456")])

    db.refresh(case)
    parts = {p.case_person_row_id: p for p in case.participants}
    assert parts[spouse].vai_tro == "Người được ủy quyền"
    assert parts[child].vai_tro == "Con"

    db.expire(case)
    context = word_engine.build_word_context(case)
    assert [p.ho_ten for p in context.authorized_people] == [
        "Trần Thị Vợ"]


def test_commit_stage_projects_full_family_with_parent_anchor_and_word_slots(
        db):
    """#45 TC4: vợ/chồng owner, neo cha/mẹ và slot Word legacy giữ như trước."""
    from services import word_engine

    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    rows = {
        "spouse": str(uuid.uuid4()),
        "child": str(uuid.uuid4()),
        "father": str(uuid.uuid4()),
        "mother": str(uuid.uuid4()),
        "spouse_father": str(uuid.uuid4()),
        "spouse_mother": str(uuid.uuid4()),
        "child_spouse": str(uuid.uuid4()),
        "grandchild": str(uuid.uuid4()),
    }
    names = {
        "spouse": "Trần Thị Vợ", "child": "Nguyễn Văn Con",
        "father": "Nguyễn Văn Cha", "mother": "Lê Thị Mẹ",
        "spouse_father": "Trần Văn Cha Vợ",
        "spouse_mother": "Phạm Thị Mẹ Vợ",
        "child_spouse": "Đỗ Thị Dâu", "grandchild": "Nguyễn Văn Cháu",
    }
    genders = {
        "spouse": "Nữ", "child": "Nam", "father": "Nam", "mother": "Nữ",
        "spouse_father": "Nam", "spouse_mother": "Nữ",
        "child_spouse": "Nữ", "grandchild": "Nam",
    }
    people = [
        _person_row(row_id=owner_row, entity_id=deceased.id,
                    ho_ten="Nguyễn Văn An", ngay_chet="2011-05-15",
                    so_giay_to="001234567890"),
        *[_person_row(row_id=rows[key], ho_ten=names[key],
                      gioi_tinh=genders[key]) for key in rows],
    ]
    _write_state(case, db, _v3_state([
        _node("father", rows["father"]),
        _node("mother", rows["mother"]),
        _node("spouse_father", rows["spouse_father"]),
        _node("spouse_mother", rows["spouse_mother"]),
        _node("owner", owner_row, parents=("father", "mother"),
              spouse="spouse", own=[1]),
        _node("spouse", rows["spouse"],
              parents=("spouse_father", "spouse_mother"), spouse="owner"),
        _node("child_1", rows["child"], parents=("owner", "spouse"),
              receive=[1]),
        _node("child_spouse", rows["child_spouse"], spouse="child_1"),
        _node("grandchild", rows["grandchild"], parents=("child_1",))]),
        owner_row_id=owner_row)
    _commit(db, case.id, 1, people, [_asset_row(so_serial="DD123456")])

    db.refresh(case)
    parts = {p.case_person_row_id: p for p in case.participants}
    assert owner_row not in parts
    assert parts[rows["spouse"]].vai_tro == "Vợ/Chồng"
    assert parts[rows["child"]].vai_tro == "Con"
    assert parts[rows["child"]].parent_case_person_row_id == owner_row
    assert parts[rows["father"]].vai_tro == "Cha"
    assert parts[rows["mother"]].vai_tro == "Mẹ"
    assert parts[rows["spouse_father"]].vai_tro == "Cha_vc"
    assert parts[rows["spouse_mother"]].vai_tro == "Me_vc"
    assert parts[rows["child_spouse"]].vai_tro == "Con_dau_re"
    assert parts[rows["child_spouse"]].parent_case_person_row_id is None
    assert parts[rows["grandchild"]].vai_tro == "Cháu"
    assert parts[rows["grandchild"]].parent_case_person_row_id == rows["child"]

    db.expire(case)
    mapping = word_engine.build_template_mapping(case)
    assert mapping["[Họ tên chủ đất]"] == "Nguyễn Văn An"
    assert mapping["[Họ tên vợ/chồng]"] == "Trần Thị Vợ"
    assert mapping["[Người nhận 1 - Họ tên]"] == "Nguyễn Văn Con"


def test_commit_stage_concurrent_same_base_conflicts(db, session_factory):
    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    s1 = session_factory()
    s2 = session_factory()
    try:
        people = [_person_row(row_id=owner_row, entity_id=deceased.id)]
        _commit(s1, case.id, 1, people, [_asset_row(so_serial="DD123456")])
        with pytest.raises(WorkspaceError) as exc:
            _commit(s2, case.id, 1, people,
                    [_asset_row(so_serial="DD123456")])
        assert exc.value.code == "workspace_conflict"
        assert exc.value.details["server_revision"] == 2
    finally:
        s1.close()
        s2.close()
    s3 = session_factory()
    try:
        assert s3.get(InheritanceCase, case.id).workspace_revision == 2
    finally:
        s3.close()


def test_get_after_concurrent_commit_reads_new_rows(db, session_factory):
    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    assert case.workspace_revision == 1
    s2 = session_factory()
    try:
        committed = _commit(
            s2, case.id, 1,
            [_person_row(row_id=owner_row, entity_id=deceased.id)],
            [_asset_row(so_serial="DD123456")])
    finally:
        s2.close()

    data = _service(db).get(case.id)

    assert data["case"]["revision"] == 2
    assert [p["row_id"] for p in data["stage"]["people"]] == [
        p["row_id"] for p in committed["stage"]["people"]]


def test_commit_stage_does_not_merge_duplicate_names_without_anchor(db):
    """Thiếu số giấy → identity null, không tạo Customer giả."""
    case, _d, _prop, _h = _make_case(db)
    people = [
        _person_row(ho_ten="Trùng Tên"),
        _person_row(ho_ten="Trùng Tên"),
    ]

    data = _commit(db, case.id, 1, people,
                   [_asset_row(so_serial="DD123456")])

    assert [p["entity_id"] for p in data["stage"]["people"]] == [None, None]
    assert db.query(Customer).filter_by(ho_ten="Trùng Tên").count() == 0


def test_commit_stage_same_anchor_shares_identity_across_rows(db):
    case, _d, _prop, _h = _make_case(db)
    people = [
        _person_row(ho_ten="Nguyễn Đăng An", so_giay_to="07, quyển 1"),
        _person_row(ho_ten="NGUYEN DANG AN", so_giay_to="07, quyển 1"),
    ]

    data = _commit(db, case.id, 1, people,
                   [_asset_row(so_serial="DD123456")])

    first, second = data["stage"]["people"]
    assert first["entity_id"] is not None
    assert first["entity_id"] == second["entity_id"]
    assert first["ho_ten"] == "Nguyễn Đăng An"
    assert second["ho_ten"] == "NGUYEN DANG AN"
    assert db.query(Customer).filter_by(ho_ten="Nguyễn Đăng An").count() == 1


def test_commit_stage_updates_links_by_row_and_primary_position(
        db, session_factory):
    case, deceased, prop, _h = _make_case(db)
    other = Property(so_serial="HH654321")
    db.add(other)
    db.commit()

    people = [_person_row(row_id=case.person_rows[0].row_id,
                          entity_id=deceased.id)]
    data = _commit(
        db, case.id, 1, people,
        [_asset_row(so_serial="HH654321", dia_chi="Thửa khác"),
         _asset_row(so_serial="DD123456",
                    dia_chi="Thửa 123, tờ 45, phường Bạch Mai")])

    assert [a["so_serial"] for a in data["stage"]["assets"]] == [
        "HH654321", "DD123456"]
    db2 = session_factory()
    try:
        links = {l.case_asset_row_id: (l.property_id, l.is_primary)
                 for l in db2.query(InheritanceCaseProperty).filter_by(
                     case_id=case.id)}
        assert list(links.values()) == [
            (data["stage"]["assets"][0]["entity_id"], True),
            (data["stage"]["assets"][1]["entity_id"], False)]
        assert db2.get(InheritanceCase, case.id).tai_san_id == \
            data["stage"]["assets"][0]["entity_id"]
    finally:
        db2.close()
    assert prop.id in [l.property_id for l in db.query(
        InheritanceCaseProperty).filter_by(case_id=case.id)]


def test_commit_stage_rejects_row_id_from_another_case(db):
    case, _d, _prop, _h = _make_case(db)
    other_case, _d2, _p2, _h2 = _make_case(db)

    with pytest.raises(WorkspaceError) as exc:
        _commit(db, case.id, 1,
                [_person_row(
                    row_id=other_case.person_rows[0].row_id)],
                [_asset_row()])

    assert exc.value.code == "stage_validation_error"
    assert any(e["field"] == "row_id"
               for e in exc.value.details["field_errors"])


def test_commit_stage_identity_hint_is_advisory(db):
    """`entity_id` là hint nội bộ (owner 23/10/2026): hint lệch identity
    đang gắn dòng bị bỏ, identity resolve lại theo neo của chính dòng —
    commit thành công, không có field_error `entity_id`."""
    case, deceased, _prop, _h = _make_case(db)
    other = Customer(ho_ten="Người Khác", so_giay_to="000000000001",
                     name_anchor="nguoi khac",
                     document_anchor="000000000001")
    db.add(other)
    db.flush()

    owner_row = case.person_rows[0]
    data = _commit(db, case.id, 1, [_person_row(
        row_id=owner_row.row_id, ho_ten=owner_row.ho_ten,
        so_giay_to=owner_row.so_giay_to, entity_id=other.id)],
        [_asset_row()])

    saved = data["stage"]["people"][0]
    assert saved["entity_id"] == deceased.id


# -------------------------------------------------------------- migration


def test_migration_adds_workspace_columns_to_legacy_db(tmp_path):
    db_path = Path(tmp_path) / "legacy.db"
    con = sqlite3.connect(db_path)
    con.execute(
        """CREATE TABLE inheritance_cases (
               id INTEGER PRIMARY KEY,
               nguoi_chet_id INTEGER NOT NULL,
               tai_san_id INTEGER NOT NULL,
               ngay_lap_ho_so DATE NOT NULL)""")
    con.commit()
    con.close()

    original = database.DB_PATH
    try:
        database.DB_PATH = db_path
        database.migrate_inheritance_cases_schema()
        database.migrate_inheritance_cases_schema()
    finally:
        database.DB_PATH = original

    con = sqlite3.connect(db_path)
    info = {row[1]: row for row in con.execute(
        "PRAGMA table_info(inheritance_cases)")}
    con.close()

    assert "workspace_revision" in info
    assert info["workspace_revision"][4] == "1"
    assert "updated_at" in info


def test_model_declares_workspace_columns_and_no_locked():
    cols = InheritanceCase.__table__.columns
    assert "workspace_revision" in cols
    assert "updated_at" in cols
    assert "workspace_idempotency_key" in cols
    assert cols["workspace_revision"].nullable is False
    assert not hasattr(InheritanceCase, "is_locked")
    check_sql = " ".join(str(c.sqltext) for c in
                         InheritanceCase.__table__.constraints
                         if c.__class__.__name__ == "CheckConstraint")
    assert "'draft'" in check_sql and "'done'" in check_sql


# ------------------------------------------------------------------ create()


def _create_args():
    """(people, assets, state) hợp lệ tối thiểu — owner đã chết + vợ nhận."""
    owner = _person_row(
        ho_ten="Nguyễn Văn An", ngay_sinh="1950", ngay_chet="2011-05-15",
        so_giay_to="001234567890")
    spouse = _person_row(
        ho_ten="Trần Thị Bình", gioi_tinh="Nữ", ngay_sinh="1955-03-02",
        so_giay_to="009876543210")
    state = _v3_state([
        _node("owner", owner["row_id"], spouse="spouse", own=[1]),
        _node("spouse", spouse["row_id"], spouse="owner", receive=[1]),
    ])
    return [owner, spouse], [_asset_row()], state


_UNSET = object()


def _create(db, people, assets, state, key=_UNSET, meta=_UNSET,
            owner_row_id=_UNSET):
    if owner_row_id is _UNSET:
        owner_row_id = people[0]["row_id"] if people else None
    stage = _stage(people, assets, owner_row_id=owner_row_id)
    return _service(db).create(
        str(uuid.uuid4()) if key is _UNSET else key,
        {"document_type": "khai_nhan"} if meta is _UNSET else meta,
        stage, state)


def test_create_persists_case_rows_and_diagram_one_transaction(db):
    people, assets, state = _create_args()
    meta = {"document_type": "khai_nhan", "ngay_lap_ho_so": "2026-09-26",
            "noi_niem_yet": "xã Yên Sở", "ghi_chu": "Hồ sơ nháp"}

    data = _create(db, people, assets, state, meta=meta)

    assert data["schema_version"] == SCHEMA_VERSION
    assert data["backend_mode"] == "real"
    assert data["created"] is True
    case = data["case"]
    assert case["case_type"] == "inheritance"
    assert case["document_type"] == "khai_nhan"
    assert case["status"] == "draft"
    assert "locked" not in case
    assert case["revision"] == 1
    assert case["ngay_lap_ho_so"] == "2026-09-26"
    assert case["noi_niem_yet"] == "xã Yên Sở"
    assert case["ghi_chu"] == "Hồ sơ nháp"

    wire_people = data["stage"]["people"]
    assert [p["ho_ten"] for p in wire_people] == [
        "Nguyễn Văn An", "Trần Thị Bình"]
    assert all(isinstance(p["entity_id"], int) for p in wire_people)
    owner_entity = wire_people[0]["entity_id"]
    spouse_entity = wire_people[1]["entity_id"]
    assert data["stage"]["owner_row_id"] == people[0]["row_id"]

    db_case = db.get(InheritanceCase, case["id"])
    assert db_case.nguoi_chet_id == owner_entity
    assert db_case.tai_san_id == data["stage"]["assets"][0]["entity_id"]
    assert db_case.ngay_lap_ho_so == date(2026, 9, 26)
    assert db_case.loai_van_ban == "khai_nhan"
    assert db_case.noi_niem_yet == "xã Yên Sở"
    assert db_case.ghi_chu == "Hồ sơ nháp"
    assert db_case.workspace_idempotency_key
    assert db_case.workspace_creation_hash
    assert db_case.engine_state_json is None
    assert len(db_case.person_rows) == 2
    assert len(db_case.asset_rows) == 1

    assert data["diagram"]["state"]["version"] == 3
    assert data["diagram"]["state"]["domain"] == "inheritance"
    assert data["diagram"]["domain"] == "inheritance"
    assert data["diagram"]["render_model"] is not None
    parts = db.query(InheritanceParticipant).filter_by(
        ho_so_id=case["id"]).all()
    assert [p.customer_id for p in parts] == [spouse_entity]
    assert parts[0].vai_tro == "Vợ/Chồng"
    assert parts[0].case_person_row_id == people[1]["row_id"]

    again = _service(db).get(case["id"])
    assert again["case"]["revision"] == 1
    assert again["stage"]["owner_row_id"] == people[0]["row_id"]
    assert [(p["row_id"], p["entity_id"], p["ho_ten"])
            for p in again["stage"]["people"]] == [
        (p["row_id"], p["entity_id"], p["ho_ten"]) for p in wire_people]


def test_create_rows_without_anchors_save_null_identity(db):
    people = [_person_row(ho_ten="Người thử chưa đủ neo")]
    assets = [_asset_row(so_serial="")]
    data = _create(db, people, assets, None)

    assert data["stage"]["people"][0]["entity_id"] is None
    assert data["stage"]["assets"][0]["entity_id"] is None
    assert db.query(Customer).count() == 0
    assert db.query(Property).count() == 0


def test_noi_niem_yet_saved_value_flows_into_word_export(
        db, session_factory, tmp_path):
    from docx import Document

    from services import word_engine

    people, assets, state = _create_args()
    typed = "Trụ sở UBND xã Đại An, huyện Vụ Bản"
    data = _create(db, people, assets, state, meta={
        "document_type": "khai_nhan",
        "noi_niem_yet": typed,
    })

    fresh = session_factory()
    try:
        db_case = fresh.get(InheritanceCase, data["case"]["id"])
        assert db_case.noi_niem_yet == typed

        template = tmp_path / "mau_niem_yet.docx"
        doc = Document()
        doc.add_paragraph("Niêm yết tại: [Nơi niêm yết]")
        doc.save(str(template))

        out = Document(str(template))
        mapping = word_engine.build_template_mapping(db_case)
        word_engine.replace_in_doc(out, mapping)
        assert word_engine.find_unresolved_placeholders(
            out, template_ref=template) == []
        out_path = tmp_path / "out.docx"
        out.save(str(out_path))

        exported = Document(str(out_path))
        assert exported.paragraphs[0].text == f"Niêm yết tại: {typed}"
        assert "Thửa 99" not in exported.paragraphs[0].text
    finally:
        fresh.close()


def test_noi_niem_yet_missing_stays_empty_no_property_fallback(
        db, session_factory):
    from services import word_engine

    people, assets, state = _create_args()
    data = _create(db, people, assets, state, meta={
        "document_type": "khai_nhan"})

    fresh = session_factory()
    try:
        db_case = fresh.get(InheritanceCase, data["case"]["id"])
        assert db_case.noi_niem_yet is None
        mapping = word_engine.build_template_mapping(db_case)
        assert mapping["[Nơi niêm yết]"] == ""
        assert mapping["[Niêm Yết]"] == ""
        assert mapping["[Địa chỉ đất]"] == "Thửa 99, xã Yên Sở"
    finally:
        fresh.close()


def test_create_without_diagram_seeds_owner(db):
    people, assets, _state = _create_args()
    data = _create(db, people, assets, None)

    assert data["created"] is True
    st = data["diagram"]["state"]
    owner = next((n for n in st["nodes"] if n["id"] == "owner"), None)
    assert owner is not None
    assert owner["personId"] == people[0]["row_id"]
    assert owner["ownPositions"] == [1]
    assert data["diagram"]["render_model"] is None


def test_create_two_party_seeds_canonical_slots(db):
    people = [_person_row(ho_ten="Bên A", ngay_chet=None),
              _person_row(ho_ten="Bên B", gioi_tinh="Nữ",
                          ngay_chet=None)]
    meta = {"case_type": "two_party", "document_type": "chuyen_nhuong"}
    data = _service(db).create(
        str(uuid.uuid4()), meta, _stage(people, [_asset_row()]), None)

    assert data["created"] is True
    assert data["case"]["case_type"] == "two_party"
    assert data["case"]["document_type"] == "chuyen_nhuong"
    assert "owner_row_id" not in data["stage"]
    st = data["diagram"]["state"]
    assert st["domain"] == "two_party"
    assert [n["id"] for n in st["nodes"]] == TP_IDS
    assert data["capabilities"]["word_export"] is True
    db_case = db.get(InheritanceCase, data["case"]["id"])
    assert db_case.nguoi_chet_id is None


def test_create_idempotent_replay_same_key(db):
    people, assets, state = _create_args()
    key = str(uuid.uuid4())

    first = _create(db, people, assets, state, key=key)
    second = _create(db, people, assets, state, key=key)

    assert first["created"] is True
    assert second["created"] is False
    assert second["case"]["id"] == first["case"]["id"]
    assert db.query(InheritanceCase).count() == 1
    assert db.get(InheritanceCase, first["case"]["id"]).workspace_revision == 1

    # key khác + row_id mới = nháp khác → case mới
    people2, assets2, state2 = _create_args()
    third = _create(db, people2, assets2, state2, key=str(uuid.uuid4()))
    assert third["created"] is True
    assert third["case"]["id"] != first["case"]["id"]
    assert db.query(InheritanceCase).count() == 2


def test_create_same_key_different_hash_conflicts(db):
    people, assets, state = _create_args()
    key = str(uuid.uuid4())
    _create(db, people, assets, state, key=key)

    with pytest.raises(WorkspaceError) as exc:
        _create(db, people, assets, state, key=key,
                meta={"document_type": "khai_nhan", "ghi_chu": "khác"})

    assert exc.value.code == "idempotency_conflict"
    assert db.query(InheritanceCase).count() == 1


def test_create_requires_uuid4_idempotency_key(db):
    people, assets, state = _create_args()
    for bad in (None, "not-a-uuid", 42, str(uuid.uuid1())):
        with pytest.raises(WorkspaceError) as exc:
            _create(db, people, assets, state, key=bad)
        assert exc.value.code == "validation_error", bad
    assert db.query(InheritanceCase).count() == 0


def test_create_case_meta_validation(db):
    people, assets, state = _create_args()
    for bad_meta in (None, [], {"document_type": "khac"},
                     {"document_type": None}, {"bogus": 1},
                     {"document_type": "khai_nhan",
                      "ngay_lap_ho_so": "2026"},
                     {"document_type": "khai_nhan", "ghi_chu": "  "},
                     {"document_type": "khai_nhan", "noi_niem_yet": 5}):
        with pytest.raises(WorkspaceError) as exc:
            _create(db, people, assets, state, meta=bad_meta)
        assert exc.value.code == "validation_error", bad_meta
    assert db.query(InheritanceCase).count() == 0


def test_create_missing_owner_maps_workspace_owner_required(db):
    people, assets, _state = _create_args()
    # Vòng 2: create nháp cho owner absent/null; non-null sai/ngoài Stage
    # vẫn workspace_owner_required (§13.6).
    partial = _create(db, people, assets, None, owner_row_id=None)
    assert partial["stage"]["owner_row_id"] is None
    assert partial["diagram"]["state"]["nodes"] == []
    for bad_owner in (str(uuid.uuid4()), "not-uuid"):
        with pytest.raises(WorkspaceError) as exc:
            _create(db, people, assets, None, owner_row_id=bad_owner)
        assert exc.value.code == "workspace_owner_required", bad_owner
    assert db.query(InheritanceCase).count() == 1


def test_create_owner_node_mismatch_rejected(db):
    people, assets, _state = _create_args()
    state = _v3_state([
        _node("owner", people[1]["row_id"], spouse="spouse", own=[1]),
        _node("spouse", None, spouse="owner", receive=[1]),
    ])
    with pytest.raises(WorkspaceError) as exc:
        _create(db, people, assets, state)
    assert exc.value.code == "diagram_owner_mismatch"
    assert exc.value.details["owner_row_id"] == people[0]["row_id"]
    assert db.query(InheritanceCase).count() == 0


def test_create_stage_rules_apply(db):
    people, assets, state = _create_args()

    # `entity_id` là hint nội bộ — không phải lỗi field (owner 23/10/2026).
    hinted = [_person_row(ho_ten=p["ho_ten"], so_giay_to=p["so_giay_to"],
                          entity_id=5) for p in people]
    hinted_state = _v3_state([
        _node("owner", hinted[0]["row_id"], spouse="spouse", own=[1]),
        _node("spouse", hinted[1]["row_id"], spouse="owner", receive=[1]),
    ])
    created = _create(db, hinted, assets, hinted_state)
    hinted_row = created["stage"]["people"][0]
    assert hinted_row["entity_id"] is not None
    assert hinted_row["entity_id"] != 5

    # Vòng 2 — nháp ít dữ liệu: 0 Tài sản, 0 Người, tên rỗng đều lưu được.
    assert _create(db, people, [], state)["stage"]["assets"] == []
    assert _create(db, [], [], None)["stage"]["people"] == []
    blank = _create(db, [_person_row(ho_ten="  ")], [], None)
    assert blank["stage"]["people"][0]["ho_ten"] is None

    # Sai kiểu tên vẫn là lỗi field — không bị nới thành null.
    for bad_name in (123, ["Nguyễn"], {"ho_ten": "Nguyễn"}):
        with pytest.raises(WorkspaceError) as exc:
            _create(db, [_person_row(ho_ten=bad_name)], [], None)
        assert exc.value.code == "stage_validation_error"
        assert any(e["field"] == "ho_ten"
                   for e in exc.value.details["field_errors"])


def test_create_is_primary_rejected(db):
    people, assets, state = _create_args()
    assets[0]["is_primary"] = True
    with pytest.raises(WorkspaceError) as exc:
        _create(db, people, assets, state)
    assert exc.value.code == "validation_error"
    assert db.query(InheritanceCase).count() == 0


def test_create_domain_mismatch_rejected(db):
    people, assets, _state = _create_args()
    tp_state = {"version": 3, "domain": "two_party",
                "nodes": [{"id": f"p{i}", "personId": None,
                           "hidden": False, "deleted": False}
                          for i in range(1, 31)]}
    with pytest.raises(WorkspaceError) as exc:
        _create(db, people, assets, tp_state)
    assert exc.value.code == "diagram_domain_mismatch"
    assert exc.value.details["expected"] == "inheritance"


def test_create_diagram_invalid_and_outside_stage(db):
    people, assets, _state = _create_args()

    dangling = _v3_state([
        _node("owner", people[0]["row_id"], parents=("ghost",),
              own=[1])])
    with pytest.raises(WorkspaceError) as exc:
        _create(db, people, assets, dangling)
    assert exc.value.code == "diagram_invalid_state"

    outside = _v3_state([
        _node("owner", people[0]["row_id"], own=[1]),
        _node("child_1", str(uuid.uuid4()), parents=("owner",),
              receive=[1])])
    with pytest.raises(WorkspaceError) as exc:
        _create(db, people, assets, outside)
    assert exc.value.code == "diagram_reference_outside_stage"
    assert db.query(InheritanceCase).count() == 0


def test_update_meta_keeps_rows_and_word_snapshot(db):
    from services.word_engine import build_template_mapping

    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    _write_state(case, db, _v3_state([
        _node("owner", owner_row, own=[1], receive=[1])]),
        owner_row_id=owner_row)
    revision = case.workspace_revision

    result = _service(db).update_meta(
        case.id, revision,
        {"noi_niem_yet": "xã Yên Sở", "ghi_chu": "Ghi chú mới"})

    assert result["revision"] == revision + 1
    assert result["case"]["noi_niem_yet"] == "xã Yên Sở"
    db.refresh(case)
    assert [r.row_id for r in case.person_rows] == [
        case.person_rows[0].row_id]
    mapping = build_template_mapping(db.get(InheritanceCase, case.id))
    assert mapping["[Tên 1]"] == "Nguyễn Văn An"


def test_year_only_dates_survive_commit_reopen_and_unrelated_edit(db):
    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    people = [_person_row(row_id=owner_row, entity_id=deceased.id,
                          ho_ten="Nguyễn Văn An", ngay_sinh="1950",
                          ngay_chet="2011-05-15",
                          so_giay_to="001234567890")]
    assets = [_asset_row(so_serial="DD123456")]
    _commit(db, case.id, 1, people, assets)

    people[0]["dia_chi"] = "Địa chỉ mới"
    _commit(db, case.id, 2, people, assets)

    db.expire_all()
    data = _service(db).get(case.id)
    assert data["stage"]["people"][0]["ngay_sinh"] == "1950"
    assert data["stage"]["people"][0]["dia_chi"] == "Địa chỉ mới"


def test_saved_type_conversion_keeps_rows_and_resets_domain(db):
    case, deceased, _prop, _h = _make_case(db)
    owner_row = case.person_rows[0].row_id
    _commit(db, case.id, 1,
            [_person_row(row_id=owner_row, entity_id=deceased.id)],
            [_asset_row(so_serial="DD123456")])
    db.refresh(case)

    result = _service(db).convert_type(
        case.id, 2,
        {"case_type": "two_party", "document_type": "chuyen_nhuong",
         "confirm_diagram_reset": True})

    assert result["case"]["case_type"] == "two_party"
    assert result["case"]["revision"] == 3
    assert [p["row_id"] for p in result["stage"]["people"]] == [owner_row]
    assert [a["so_serial"] for a in result["stage"]["assets"]] == [
        "DD123456"]
    assert result["diagram"]["state"]["domain"] == "two_party"
    db.refresh(case)
    assert case.nguoi_chet_id is None
    assert case.participants == []


def test_saved_type_conversion_invalid_leaves_snapshot_unchanged(db):
    case, _deceased, _prop, _h = _make_case(db)
    before = case.case_state_json

    with pytest.raises(WorkspaceError) as exc:
        _service(db).convert_type(
            case.id, 1,
            {"case_type": "two_party", "document_type": "khai_nhan",
             "confirm_diagram_reset": True})

    assert exc.value.code == "validation_error"
    db.expire_all()
    assert db.get(InheritanceCase, case.id).case_state_json == before


def test_saved_type_conversion_stale_leaves_case_usable(db):
    case, _deceased, _prop, _h = _make_case(db)

    with pytest.raises(WorkspaceError) as exc:
        _service(db).convert_type(
            case.id, 99,
            {"case_type": "two_party", "document_type": "chuyen_nhuong",
             "confirm_diagram_reset": True})

    assert exc.value.code == "workspace_conflict"
    db.expire_all()
    assert db.get(InheritanceCase, case.id).trang_thai == "draft"
    assert db.get(InheritanceCase, case.id).workspace_revision == 1


@pytest.fixture()
def many_people_stage():
    return {
        "owner_row_id": None,
        "people": [{"row_id": str(uuid.uuid4()), "ho_ten": f"Người thử {i:02}"}
                   for i in range(1, 26)],
        "assets": [{"row_id": str(uuid.uuid4()), "so_serial": "",
                    "so_thua_dat": "901", "so_to_ban_do": "09"}],
    }


@pytest.fixture()
def many_people_word_source(tmp_path, many_people_stage):
    import hashlib
    import io
    from docx import Document
    from services.document_intake.docx_source import read_docx
    from services.document_intake.docx_store import DocxStore

    name = many_people_stage["people"][-1]["ho_ten"]
    doc = Document()
    doc.add_paragraph("Người: " + name)
    stream = io.BytesIO()
    doc.save(stream)
    content = stream.getvalue()
    text = read_docx(content).text
    candidate_id, occurrence_id = str(uuid.uuid4()), str(uuid.uuid4())
    suggestion = {
        "suggestion_id": candidate_id, "source_id": "", "target": "person",
        "fields": {"ho_ten": {
            "raw_value": name, "normalized_value": name,
            "observation_state": "normalized", "confidence": None,
            "source_refs": [{
                "occurrence_id": occurrence_id, "source_id": "",
                "source_label": "", "raw_snippet": name,
                "span": [text.index(name), text.index(name) + len(name)],
            }],
        }}, "warnings": [],
    }
    store = DocxStore(tmp_path / "sources")
    source_id = store.save(content, [suggestion], "Nguồn tổng hợp.docx")
    context = {"version": 1, "sources": [{
        "source_id": source_id, "sha256": hashlib.sha256(content).hexdigest(),
        "records": [{"candidate_id": candidate_id,
                     "row_id": many_people_stage["people"][-1]["row_id"]}],
        "bindings": [{"candidate_id": candidate_id,
                      "occurrence_id": occurrence_id, "field_key": "ho_ten",
                      "slot": "[Người 25 - Họ tên]"}],
    }]}
    return store, context


def test_save_case_inheritance_25_names_only_reopens_updates_replays(
        db, session_factory, many_people_stage):
    service = _service(db)
    key = str(uuid.uuid4())
    args = {
        "case_id": None, "base_revision": None, "idempotency_key": key,
        "case_meta": {"case_type": "inheritance", "document_type": "khai_nhan"},
        "stage": many_people_stage,
        "state": {"version": 3, "domain": "inheritance", "nodes": []},
        "docx_import": {"version": 1, "sources": []},
    }
    saved = service.save_case(**args)
    case_id = saved["case"]["id"]
    assert saved["created"] is True
    assert saved["case"]["status"] == "done"
    assert saved["case"]["revision"] == 1
    replay = service.save_case(**args)
    assert replay["created"] is False
    assert replay["case"]["id"] == case_id
    assert replay["case"]["revision"] == 1
    assert db.query(InheritanceCase).count() == 1

    with session_factory() as fresh:
        reopened = CaseWorkspaceService(fresh).get(case_id)
        assert [(p["row_id"], p["ho_ten"]) for p in reopened["stage"]["people"]] == [
            (p["row_id"], p["ho_ten"]) for p in many_people_stage["people"]]
        assert all(p["so_giay_to"] is None and p["ngay_sinh"] is None
                   for p in reopened["stage"]["people"])
        asset, = reopened["stage"]["assets"]
        assert (asset["so_thua_dat"], asset["so_to_ban_do"]) == ("901", "09")
        assert reopened["docx_import"] == {"version": 1, "sources": []}
        reopened["stage"]["people"][-1]["ho_ten"] = "Người thử đã sửa"
        updated = CaseWorkspaceService(fresh).save_case(
            case_id=case_id, base_revision=1, case_meta={},
            stage=reopened["stage"], state=reopened["diagram"]["state"])
        assert updated["case"]["revision"] == 2

    with session_factory() as fresh:
        reopened = CaseWorkspaceService(fresh).get(case_id)
        assert len(reopened["stage"]["people"]) == 25
        assert reopened["stage"]["people"][-1]["ho_ten"] == "Người thử đã sửa"
        assert reopened["case"]["revision"] == 2


def test_save_case_inheritance_25_with_reviewed_binding(
        db, session_factory, many_people_stage, many_people_word_source, tmp_path):
    from services.word_original import binding_proposals, export_original
    from services.word_batch_export import WordBatchError

    store, context = many_people_word_source
    saved = CaseWorkspaceService(db, store).save_case(
        case_id=None, base_revision=None, idempotency_key=str(uuid.uuid4()),
        case_meta={"case_type": "inheritance", "document_type": "khai_nhan"},
        stage=many_people_stage,
        state={"version": 3, "domain": "inheritance", "nodes": []},
        docx_import=context)
    case_id = saved["case"]["id"]
    source_id = context["sources"][0]["source_id"]
    with session_factory() as fresh:
        service = CaseWorkspaceService(fresh, store)
        reopened = service.get(case_id)
        assert len(reopened["stage"]["people"]) == 25
        binding, = reopened["docx_import"]["sources"][0]["bindings"]
        binding = dict(binding)
        assert binding["slot"] == "[Người 25 - Họ tên]"
        case = fresh.get(InheritanceCase, case_id)
        proposal, = binding_proposals(case, source_id, store)
        assert proposal["supported"] is True
        assert proposal["selected"] is True
        assert proposal["value"] == "Người thử 25"
        with pytest.raises(WordBatchError) as failure:
            export_original(case, source_id, store, tmp_path)
        assert failure.value.code == "word_batch_failed"
        entry, = failure.value.result_data["documents"]
        assert entry["status"] == "failed"
        assert entry["error"] == {
            "code": "word.too_many_people", "message": "Quá 20 người trong Stage"}
        assert list(tmp_path.glob("*_HS-*.docx")) == []

        reopened["docx_import"]["sources"][0]["bindings"][0]["slot"] = "[Người 24 - Họ tên]"
        with pytest.raises(WorkspaceError) as failure:
            service.save_case(
                case_id=case_id, base_revision=1, case_meta={},
                stage=reopened["stage"], state=reopened["diagram"]["state"],
                docx_import=reopened["docx_import"])
        assert failure.value.code == "validation_error"
        unchanged = service.get(case_id)
        assert unchanged["case"]["revision"] == 1
        assert unchanged["docx_import"]["sources"][0]["bindings"] == [binding]


def test_save_case_two_party_people_cap_stays_30(db, many_people_stage):
    stage = many_people_stage
    stage.pop("owner_row_id")
    stage["people"].extend(
        {"row_id": str(uuid.uuid4()), "ho_ten": f"Người thử {i:02}"}
        for i in range(26, 31))
    service = _service(db)
    args = {
        "case_id": None, "base_revision": None,
        "idempotency_key": str(uuid.uuid4()),
        "case_meta": {"case_type": "two_party", "document_type": "chuyen_nhuong"},
        "stage": stage,
        "state": {"version": 3, "domain": "two_party", "nodes": [
            {"id": pid, "personId": None, "hidden": False, "deleted": False}
            for pid in TP_IDS]},
        "docx_import": {"version": 1, "sources": []},
    }
    saved = service.save_case(**args)
    assert len(saved["stage"]["people"]) == 30
    stage["people"].append(
        {"row_id": str(uuid.uuid4()), "ho_ten": "Người thử 31"})
    args["idempotency_key"] = str(uuid.uuid4())
    with pytest.raises(WorkspaceError) as failure:
        service.save_case(**args)
    assert failure.value.code == "stage_validation_error"
    assert any(e["code"] == "people_limit"
               for e in failure.value.details["field_errors"])
    assert len(service.get(saved["case"]["id"])["stage"]["people"]) == 30
    assert db.query(InheritanceCase).count() == 1


def test_save_case_inheritance_25_rejects_binding_without_saved_value(
        db, many_people_stage, many_people_word_source):
    store, context = many_people_word_source
    source = context["sources"][0]
    content, package = store.load(source["source_id"])
    candidate, = package["suggestions"]
    candidate["fields"]["so_giay_to"] = candidate["fields"].pop("ho_ten")
    source["source_id"] = store.save(
        content, [candidate], "Nguồn giấy tờ tổng hợp.docx")
    binding, = source["bindings"]
    binding.update(field_key="so_giay_to", slot="[Người 25 - Số giấy tờ]")
    with pytest.raises(WorkspaceError) as failure:
        CaseWorkspaceService(db, store).save_case(
            case_id=None, base_revision=None, idempotency_key=str(uuid.uuid4()),
            case_meta={"case_type": "inheritance", "document_type": "khai_nhan"},
            stage=many_people_stage,
            state={"version": 3, "domain": "inheritance", "nodes": []},
            docx_import=context)
    assert failure.value.code == "validation_error"
    assert "chưa có giá trị" in str(failure.value)
    assert db.query(InheritanceCase).count() == 0
    assert db.query(CasePersonRow).count() == 0


@pytest.mark.parametrize("schema_version", [3, 4])
def test_v3_spouse_migration_keeps_candidates_order_and_living_suggestion(
        db, schema_version):
    from services.case_workspace import _v3_to_v2_nodes
    from services.inheritance_engine import run_inheritance_case
    from services.inheritance_workspace import _validate_diagram_wire

    deceased, first, second = (str(uuid.uuid4()) for _ in range(3))
    raw_nodes = [
        _node("dead", deceased, spouse="first", own=[1]),
        _node("first", first, spouse="second", receive=[1]),
        _node("second", second, spouse="dead", receive=[1]),
    ]
    for node in raw_nodes:
        for key in ("spouseSlotIds", "spouseSuggestionId", "spouseUndetermined"):
            node.pop(key)
    payload = {
        "schemaVersion": schema_version,
        "diagram": {"state": _v3_state(raw_nodes)},
    }
    before = json.dumps(payload)
    state = _service(db)._state_v3(
        payload, domain="inheritance",
        valid_row_ids={deceased, first, second},
        person_rows={
            deceased: {"row_id": deceased, "ngay_chet": "2020-01-01"},
            first: {"row_id": first, "ngay_chet": None},
            second: {"row_id": second, "ngay_chet": None},
        })
    nodes = {node["id"]: node for node in state["nodes"]}

    assert json.dumps(payload) == before
    assert nodes["dead"]["spouseSlotIds"] == ["first", "second"]
    assert nodes["dead"]["spouseSlotId"] == "second"
    assert nodes["dead"]["spouseUndetermined"] is True
    assert nodes["first"]["spouseSlotIds"] == []
    assert nodes["first"]["spouseSuggestionId"] == "second"
    assert _validate_diagram_wire(state) == []
    engine_nodes = _v3_to_v2_nodes(state["nodes"])
    assert engine_nodes[0]["spouseUndetermined"] is True
    rendered = run_inheritance_case(
        {"version": 2, "nodes": engine_nodes},
        {deceased: {"ngay_chet": "2020-01-01"},
         first: {"ngay_chet": None}, second: {"ngay_chet": None}},
    )
    assert rendered["status"] == "incomplete", rendered["errors"]
    assert any(warning["code"] == "spouse_undetermined"
               for warning in rendered["warnings"])


def test_get_normalizes_singular_spouse_without_writing_then_commit_persists(db):
    case, _deceased, _prop, _heirs = _make_case(db, with_participant=True)
    owner_row, spouse_row = (row.row_id for row in case.person_rows)
    legacy_nodes = [
        _node("owner", owner_row, spouse="spouse", own=[1]),
        _node("spouse", spouse_row, spouse="owner", receive=[1]),
    ]
    for node in legacy_nodes:
        for key in ("spouseSlotIds", "spouseSuggestionId", "spouseUndetermined"):
            node.pop(key)
    _write_state(case, db, _v3_state(legacy_nodes), owner_row_id=owner_row)
    before = case.case_state_json
    service = _service(db)

    loaded = service.get(case.id)

    nodes = {node["id"]: node for node in loaded["diagram"]["state"]["nodes"]}
    assert nodes["owner"]["spouseSlotIds"] == ["spouse"]
    assert nodes["spouse"]["spouseSlotIds"] == []
    assert nodes["spouse"]["spouseSuggestionId"] == "owner"
    assert case.case_state_json == before
    assert service.get(case.id)["diagram"]["state"] == loaded["diagram"]["state"]
    assert case.workspace_revision == 1

    committed = service.commit_stage(case.id, 1, loaded["stage"])

    persisted = json.loads(case.case_state_json)
    assert persisted["diagram"]["state"] == committed["diagram"]["state"]
    assert persisted["diagram"]["state"] == loaded["diagram"]["state"]
    assert "stage" not in persisted and "assets" not in persisted
    assert case.engine_state_json is None
    assert case.workspace_revision == 2


@pytest.mark.parametrize("spouse_slots,lawful", [
    (["first", "second"], "second"),
    ([], None),
])
def test_spouse_role_projection_uses_ordered_outgoing_links(
        spouse_slots, lawful):
    from services.case_workspace import _project_role_nodes

    owner = _node("owner", spouse="first")
    owner["spouseSlotIds"] = spouse_slots
    nodes = [
        owner,
        _node("first", spouse="owner"),
        _node("second"),
        _node("child", parents=("owner",), spouse="child_spouse"),
        _node("child_spouse"),
    ]

    roles = {node["id"]: node for node in _project_role_nodes(nodes, [])}

    assert roles["first"]["relationType"] == ""
    assert roles["second"]["relationType"] == ("spouse" if lawful else "")
    assert roles["child_spouse"]["relationType"] == "branchSpouse"
    assert roles["child_spouse"]["spouseOf"] == "child"
