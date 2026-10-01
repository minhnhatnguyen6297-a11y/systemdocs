"""MIN-141 đợt 3 — Person fields + case metadata xuyên
backend → snapshot → Word.

- Người: 6 trường cơ sở round-trip; trường suy ra theo mốc ngày
  (01/10/2024); người chết không gán cứng — loại giấy tờ/nơi cấp lấy
  bằng chứng đã xác nhận trong snapshot; thiếu ngày/bằng chứng →
  None (chưa xác định), không suy cơ quan cụ thể.
- Hồ sơ: noiniemyet / nguoinhanuyquyen / noidungviec đi trọn
  create → commit → reload → export; metadata nằm trong payload.case;
  commit meta cùng Stage trong một transaction (revision/idempotency).
"""
import json
import sqlite3
import uuid
from datetime import date
from pathlib import Path

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import database
from models import Customer, InheritanceCase, Property, PropertyLandRow
from services import word_engine
from services.case_workspace import CaseWorkspaceService, WorkspaceError


@pytest.fixture()
def db(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'meta.db'}")
    database.Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    yield session
    session.close()
    engine.dispose()


def _service(db):
    return CaseWorkspaceService(db)


def _person_row(**overrides):
    row = {
        "row_id": str(uuid.uuid4()),
        "entity_id": None,
        "ho_ten": "Người Mới",
        "gioi_tinh": "Nam",
        "ngay_sinh": "1980-01-01",
        "ngay_chet": None,
        "so_giay_to": None,
        "ngay_cap": None,
        "noi_cap": None,
        "dia_chi": "xã test",
        "place_of_origin": None,
    }
    row.update(overrides)
    return row


def _asset_row(**overrides):
    row = {
        "row_id": str(uuid.uuid4()),
        "entity_id": None,
        "so_serial": "EE123456",
        "so_vao_so": "CS 99999",
        "so_thua_dat": "99",
        "so_to_ban_do": "12",
        "dia_chi": "Thửa 99, xã Yên Sở",
        "loai_so": "GCN QSDĐ",
        "hinh_thuc_su_dung": "Sử dụng riêng",
        "thoi_han": "Lâu dài",
        "nguon_goc": "Cấp đổi",
        "ngay_cap": "2005-09-30",
        "co_quan_cap": "UBND quận Hai Bà Trưng",
        "land_rows": [],
    }
    row.update(overrides)
    return row


def _node(nid, person_id=None, own=(), receive=()):
    return {"id": nid, "personId": person_id, "parentSlotIds": [],
            "spouseSlotId": None, "ownPositions": list(own),
            "receivePositions": list(receive), "hidden": False,
            "deleted": False}


def _create_args(people_overrides=(), case_meta=None, with_state=False,
                 **kw):
    owner_fields = {
        "ho_ten": "Nguyễn Văn Chết", "ngay_chet": "2011-05-15",
        "so_giay_to": "TLK 12/2011", "ngay_cap": "2011-05-16",
        "dia_chi": "xã X, huyện Y",
    }
    owner_fields.update(people_overrides[0] if people_overrides else {})
    owner = _person_row(**owner_fields)
    heir_fields = {"ho_ten": "Trần Thị Sống", "ngay_sinh": "1980-03-02"}
    heir_fields.update(
        people_overrides[1] if len(people_overrides) > 1 else {})
    heir = _person_row(**heir_fields)
    stage = {"people": [owner, heir], "assets": [_asset_row()],
             "owner_row_id": owner["row_id"]}
    # Word export cần chủ đất chết + người nhận trên Diagram (đợt 3 test
    # [Dòng khai tử]/[Nơi niêm yết]/[Ngườ ủy quyền] đọc từ snapshot).
    state = ({"version": 3, "domain": "inheritance",
              "nodes": [_node("owner", owner["row_id"], own=[1]),
                        _node("heir", heir["row_id"], receive=[1])]}
             if with_state else None)
    return {
        "idempotency_key": str(uuid.uuid4()),
        "case_meta": dict(case_meta or {
            "case_type": "inheritance",
            "document_type": "khai_nhan",
            "ngay_lap_ho_so": "2026-09-01",
        }),
        "stage": stage,
        "state": state,
        **kw,
    }


def _create(db, **kwargs):
    return _service(db).create(**_create_args(**kwargs))


def _commit(db, case_id, base_revision, stage, case_meta=None):
    return _service(db).commit_stage(
        case_id, base_revision, stage, case_meta)


def _snapshot(db, case_id):
    case = db.get(InheritanceCase, case_id)
    return json.loads(case.case_state_json)


# ============================================================ Người


def test_living_person_derived_fields_before_threshold(db):
    """Người sống, ngay_cap trước 01/10/2024 → CCCD cũ + Cục CSQLHC."""
    data = _create(db, people_overrides=({},
                                         {"ngay_cap": "2020-05-20"}))
    heir = data["stage"]["people"][1]
    assert heir["loai_giay_to"] == "Căn cước công dân"
    assert (heir["noi_cap"]
            == "Cục cảnh sát quản lý hành chính về trật tự xã hội")
    assert heir["loai_dia_chi"] == "Thường trú tại"


def test_living_person_derived_fields_at_and_after_threshold(db):
    """Đúng/sau 01/10/2024 → Căn cước + Bộ Công an + Cư trú."""
    for ngay in ("2024-10-01", "2025-01-15"):
        data = _create(db, people_overrides=({}, {"ngay_cap": ngay}))
        heir = data["stage"]["people"][1]
        assert heir["loai_giay_to"] == "Căn cước", ngay
        assert heir["noi_cap"] == "Bộ Công an", ngay
        assert heir["loai_dia_chi"] == "Cư trú tại", ngay


def test_living_person_missing_ngay_cap_not_inferred(db):
    """Thiếu ngaycap → để trống/chưa xác định, không tự suy."""
    data = _create(db, people_overrides=({}, {}))
    heir = data["stage"]["people"][1]
    assert heir["loai_giay_to"] is None
    assert heir["noi_cap"] is None
    assert heir["loai_dia_chi"] is None


def test_deceased_person_not_hardcoded(db):
    """Người chết: model/service KHÔNG gán cứng 'Trích lục khai tử
    (Bản sao)' hay 'Ủy ban nhân dân cấp xã' khi chưa có bằng chứng."""
    data = _create(db)
    owner = data["stage"]["people"][0]
    # owner đã khai ngay_cap 2011 — đợt 3 không được suy UBND xã từ ngày.
    assert owner["loai_giay_to"] is None
    assert owner["noi_cap"] is None
    assert owner["loai_dia_chi"] == "Nơi chết"


def test_deceased_person_confirmed_evidence_persists(db):
    """Loại giấy tờ/nơi cấp người chết = lựa chọn/bằng chứng trong
    payload → giữ trong snapshot, emit lại trên wire + Word."""
    data = _create(db, people_overrides=(
        {"loai_giay_to": "Giấy chứng tử",
         "noi_cap": "UBND xã A"}, {}))
    owner = data["stage"]["people"][0]
    assert owner["loai_giay_to"] == "Giấy chứng tử"
    assert owner["noi_cap"] == "UBND xã A"

    snap_person = _snapshot(db, data["case"]["id"])["stage"][0]
    assert snap_person["loai_giay_to"] == "Giấy chứng tử"
    assert snap_person["noi_cap"] == "UBND xã A"
    # Reopen — wire phải trả lại đúng giá trị snapshot đã commit.
    reopened = _service(db).get(data["case"]["id"])
    assert reopened["stage"]["people"][0]["loai_giay_to"] == "Giấy chứng tử"


def test_person_row_canonical_keys_accepted(db):
    """person_row nhận key canonical liền không dấu (ten/ngaysinh/
    loaigiayto/…) — fold về legacy trước validate."""
    heir = _person_row(ho_ten=None)
    heir.update({
        "ho_ten": None,  # legacy key null
        "ten": "Nguyễn Canonical",
        "ngaycap": "2024-10-05",
        "loaigiayto": None,
        "noicap": None,
    })
    stage = {"people": [_person_row(ho_ten="Chết"), heir],
             "assets": [_asset_row()],
             "owner_row_id": None}
    stage["owner_row_id"] = stage["people"][0]["row_id"]
    data = _service(db).create(
        str(uuid.uuid4()),
        {"case_type": "inheritance", "document_type": "khai_nhan",
         "ngay_lap_ho_so": "2026-09-01"},
        stage, None)
    wire_heir = data["stage"]["people"][1]
    assert wire_heir["ho_ten"] == "Nguyễn Canonical"
    assert wire_heir["ngay_cap"] == "2024-10-05"
    assert wire_heir["loai_giay_to"] == "Căn cước"  # suy theo mốc


def test_person_row_alias_conflict_rejected(db):
    """Cùng field ở cả canonical lẫn legacy mà khác giá trị → lỗi."""
    heir = _person_row()
    heir["ten"] = "Tên Khác"  # canonical mâu thuẫn ho_ten
    stage = {"people": [_person_row(ho_ten="Chết"), heir],
             "assets": [_asset_row()]}
    stage["owner_row_id"] = stage["people"][0]["row_id"]
    with pytest.raises(WorkspaceError) as exc:
        _service(db).create(
            str(uuid.uuid4()),
            {"case_type": "inheritance", "document_type": "khai_nhan",
             "ngay_lap_ho_so": "2026-09-01"},
            stage, None)
    assert exc.value.code == "validation_error"


# ============================================================ Hồ sơ


def _catalog_person(db, name="Người Được Ủy Quyền"):
    c = Customer(ho_ten=name, gioi_tinh="Nữ",
                 ngay_sinh=date(1985, 6, 6), dia_chi="xã Catalog")
    db.add(c)
    db.flush()
    return c


def test_case_meta_create_canonical_round_trip(db):
    """payload.case nhận key canonical; get() emit snake wire; snapshot
    payload.case giữ canonical — đủ 3 trường đợt 3."""
    uq = _catalog_person(db)
    data = _create(db, case_meta={
        "case_type": "inheritance", "document_type": "khai_nhan",
        "ngaylaphoso": "2026-09-02",
        "noiniemyet": "UBND xã B",
        "nguoinhanuyquyenid": uq.id,
        "noidungviec": "Đính chính năm sinh",
        "ghichu": "ghi chú mẫu",
    })
    case = data["case"]
    assert case["noi_niem_yet"] == "UBND xã B"
    assert case["nguoi_nhan_uy_quyen"] == uq.ho_ten  # resolve theo danh bạ
    assert case["nguoi_nhan_uy_quyen_id"] == uq.id
    assert case["noi_dung_viec"] == "Đính chính năm sinh"
    assert case["ghi_chu"] == "ghi chú mẫu"
    assert case["ngay_lap_ho_so"] == "2026-09-02"

    snap_case = _snapshot(db, case["id"])["case"]
    assert snap_case["noiniemyet"] == "UBND xã B"
    assert snap_case["nguoinhanuyquyen"] == {"id": uq.id, "ten": uq.ho_ten}
    assert snap_case["noidungviec"] == "Đính chính năm sinh"
    assert snap_case["ghichu"] == "ghi chú mẫu"


def test_case_meta_commit_same_transaction(db):
    """Commit Stage + payload.case một transaction: revision +1,
    cột + snapshot đều đổi; reload đọc lại đúng."""
    data = _create(db)
    case_id = data["case"]["id"]
    uq = _catalog_person(db, "Ông Ủy Quyền")
    stage = _service(db).get(case_id)["stage"]

    result = _commit(db, case_id, 1, stage, case_meta={
        "noiniemyet": "UBND xã C",
        "nguoinhanuyquyen": uq.ho_ten,
        "nguoinhanuyquyenid": uq.id,
        "noidungviec": "Đính chính hộ",
        "ghichu": None,
    })
    assert result["revision"] == 2

    case = db.get(InheritanceCase, case_id)
    assert case.noi_niem_yet == "UBND xã C"
    assert case.nguoi_nhan_uy_quyen == "Ông Ủy Quyền"
    assert case.nguoi_nhan_uy_quyen_id == uq.id
    assert case.noi_dung_viec == "Đính chính hộ"
    snap_case = json.loads(case.case_state_json)["case"]
    assert snap_case["noiniemyet"] == "UBND xã C"
    assert snap_case["nguoinhanuyquyen"] == {
        "id": uq.id, "ten": "Ông Ủy Quyền"}

    reopened = _service(db).get(case_id)["case"]
    assert reopened["noi_niem_yet"] == "UBND xã C"
    assert reopened["nguoi_nhan_uy_quyen_id"] == uq.id
    assert reopened["noi_dung_viec"] == "Đính chính hộ"


def test_case_meta_commit_invalid_rolls_back_everything(db):
    """Meta lỗi (uy quyền id không tồn tại) → cả Stage lẫn meta
    rollback — revision và cột giữ nguyên."""
    data = _create(db)
    case_id = data["case"]["id"]
    stage = _service(db).get(case_id)["stage"]
    stage["people"][1]["ho_ten"] = "Đổi Tên Bị Hủy"
    with pytest.raises(WorkspaceError) as exc:
        _commit(db, case_id, 1, stage,
                case_meta={"nguoinhanuyquyenid": 999999})
    assert exc.value.code == "validation_error"
    case = db.get(InheritanceCase, case_id)
    assert case.workspace_revision == 1
    assert case.nguoi_nhan_uy_quyen_id is None


def test_case_meta_conflicting_name_and_id_rejected(db):
    """nguoinhanuyquyen + id cùng gửi mà khác tên master → lỗi."""
    uq = _catalog_person(db)
    data = _service(db).create(**_create_args())
    cid = data["case"]["id"]
    st = _service(db).get(cid)["stage"]
    with pytest.raises(WorkspaceError):
        _commit(db, cid, 1, st, case_meta={
            "nguoinhanuyquyenid": uq.id,
            "nguoinhanuyquyen": "Tên Không Khớp"})


def test_case_meta_immutable_type_rejected(db):
    """case_type/document_type gửi khác giá trị đã lưu → lỗi."""
    data = _create(db)
    cid = data["case"]["id"]
    st = _service(db).get(cid)["stage"]
    with pytest.raises(WorkspaceError):
        _commit(db, cid, 1, st, case_meta={"document_type": "thoa_thuan"})
    with pytest.raises(WorkspaceError):
        _commit(db, cid, 1, st, case_meta={"casetype": "two_party"})


# ============================================================ Word


def test_word_case_meta_from_committed_snapshot(db):
    """[Nơi niêm yết]/[Ngườ ủy quyền]/[Nội dung việc] đọc từ snapshot
    đã commit của chính hồ sơ; không đắp từ địa chỉ đất."""
    uq = _catalog_person(db, "Người Quyền")
    data = _create(db, with_state=True, case_meta={
        "case_type": "inheritance", "document_type": "khai_nhan",
        "ngay_lap_ho_so": "2026-09-01",
        "noiniemyet": "xã Nơi Niêm Yết",
        "nguoinhanuyquyenid": uq.id,
        "noidungviec": "Đính chính sai chính tả",
    })
    case = db.get(InheritanceCase, data["case"]["id"])
    mapping = word_engine.build_template_mapping(case)
    assert mapping["[Nơi niêm yết]"] == "xã Nơi Niêm Yết"
    assert mapping["[Niêm Yết]"] == "xã Nơi Niêm Yết"
    assert mapping["[Ngườ ủy quyền]"] == "Người Quyền"
    assert mapping["[Ngườ ủy quyền2]"] == "Người Quyền"
    assert mapping["[Nội dung việc]"] == "Đính chính sai chính tả"


def test_word_niem_yet_empty_without_value_not_address(db):
    """Chưa tra cứu nơi niêm yết → placeholder rỗng, KHÔNG lấy địa chỉ
    đất đắp vào (AC đợt 3)."""
    data = _create(db, with_state=True)  # không set noi_niem_yet
    case = db.get(InheritanceCase, data["case"]["id"])
    assert case.noi_niem_yet is None
    mapping = word_engine.build_template_mapping(case)
    assert mapping["[Nơi niêm yết]"] == ""
    assert mapping["[Niêm Yết]"] == ""


def test_word_deceased_clause_uses_confirmed_evidence(db):
    """Dòng khai tử ghép theo bằng chứng đã xác nhận; không còn
    'Trích lục khai tử (Bản sao) ... UBND xã [Nơi niêm yết]' cứng."""
    data = _create(db, with_state=True, people_overrides=(
        {"loai_giay_to": "Giấy chứng tử",
         "noi_cap": "UBND xã Thật", "so_giay_to": "CT 44"}, {}))
    case = db.get(InheritanceCase, data["case"]["id"])
    mapping = word_engine.build_template_mapping(case)
    clause = mapping["[Dòng khai tử chủ đất chết 1]"]
    assert "Giấy chứng tử" in clause
    assert "UBND xã Thật" in clause
    assert "CT 44" in clause
    assert "Trích lục khai tử (Bản sao)" not in clause
    assert "[Nơi niêm yết]" not in clause

    # Không bằng chứng → không suy "Ủy ban nhân dân xã ..." hay cụm
    # trích lục cứng; chỉ ghép phần có.
    data2 = _create(db, with_state=True)
    case2 = db.get(InheritanceCase, data2["case"]["id"])
    clause2 = word_engine.build_template_mapping(case2)[
        "[Dòng khai tử chủ đất chết 1]"]
    assert "Trích lục" not in clause2
    assert "Ủy ban nhân dân xã" not in clause2
    assert "TLK 12/2011" in clause2  # số giấy tờ có → vẫn in


# ============================================================ Review fix
# (đổi input nguồn → tính lại trường suy ra; snapshot rỗng giữ nguyên;
# token chuẩn Hồ sơ/Người xuất thật vào Word)


def test_derived_fields_recompute_back_across_threshold(db):
    """Đổi ngay_cap ngược qua mốc 01/10/2024 → loại giấy tờ/nơi cấp/
    nhãn cư trú tính lại, không giữ giá trị suy ra cũ."""
    data = _create(db, people_overrides=({}, {"ngay_cap": "2024-10-01"}))
    heir = data["stage"]["people"][1]
    assert heir["loai_giay_to"] == "Căn cước"

    stage = data["stage"]
    stage["people"][1]["ngay_cap"] = "2024-09-30"
    updated = _commit(db, data["case"]["id"], data["case"]["revision"],
                      stage)
    heir = updated["stage"]["people"][1]
    assert heir["loai_giay_to"] == "Căn cước công dân"
    assert (heir["noi_cap"]
            == "Cục cảnh sát quản lý hành chính về trật tự xã hội")
    assert heir["loai_dia_chi"] == "Thường trú tại"


def test_derived_fields_recompute_when_issue_date_cleared(db):
    """Xóa ngay_cap → trường suy ra về None (chưa xác định), không giữ
    giá trị suy ra của lần commit trước."""
    data = _create(db, people_overrides=({}, {"ngay_cap": "2024-10-01"}))
    stage = data["stage"]
    stage["people"][1]["ngay_cap"] = None
    updated = _commit(db, data["case"]["id"], data["case"]["revision"],
                      stage)
    heir = updated["stage"]["people"][1]
    assert heir["loai_giay_to"] is None
    assert heir["noi_cap"] is None
    assert heir["loai_dia_chi"] is None


def test_derived_fields_recompute_alive_to_deceased(db):
    """Sống → chết: echo giá trị suy ra cũ không đè — người chết không
    suy loại giấy tờ/nơi cấp (None), nhãn địa chỉ thành Nơi chết."""
    data = _create(db, people_overrides=({}, {"ngay_cap": "2024-10-01"}))
    stage = data["stage"]
    stage["people"][1]["ngay_chet"] = "2025-02-10"
    updated = _commit(db, data["case"]["id"], data["case"]["revision"],
                      stage)
    heir = updated["stage"]["people"][1]
    assert heir["loai_giay_to"] is None
    assert heir["noi_cap"] is None
    assert heir["loai_dia_chi"] == "Nơi chết"


def test_deceased_confirmed_evidence_survives_recommit(db):
    """Bằng chứng khai tử đã xác nhận KHÔNG bị giá trị suy ra đè khi
    commit lại dù đổi field khác."""
    data = _create(db, people_overrides=(
        {"loai_giay_to": "Giấy chứng tử", "noi_cap": "UBND xã Thật"},
        {}))
    stage = _service(db).get(data["case"]["id"])["stage"]
    stage["people"][0]["dia_chi"] = "xã mới đổi"
    updated = _commit(db, data["case"]["id"], data["case"]["revision"],
                      stage)
    owner = updated["stage"]["people"][0]
    assert owner["loai_giay_to"] == "Giấy chứng tử"
    assert owner["noi_cap"] == "UBND xã Thật"
    assert owner["dia_chi"] == "xã mới đổi"


def test_user_confirmed_value_survives_source_change(db):
    """Người sống xác nhận giá trị khác suy ra → đổi ngay_cap vẫn giữ
    xác nhận; field không chạm vẫn tính lại theo mốc mới."""
    data = _create(db, people_overrides=(
        {}, {"ngay_cap": "2024-09-30",
             "loai_giay_to": "Chứng minh nhân dân"}))
    heir = data["stage"]["people"][1]
    assert heir["loai_giay_to"] == "Chứng minh nhân dân"

    stage = data["stage"]
    stage["people"][1]["ngay_cap"] = "2024-10-01"
    updated = _commit(db, data["case"]["id"], data["case"]["revision"],
                      stage)
    heir = updated["stage"]["people"][1]
    assert heir["loai_giay_to"] == "Chứng minh nhân dân"  # xác nhận giữ
    assert heir["noi_cap"] == "Bộ Công an"                # echo → tính lại


def test_commit_unchanged_keeps_derived_values_stable(db):
    """Commit không đổi input → trường suy ra giữ nguyên (idempotent)."""
    data = _create(db, people_overrides=({}, {"ngay_cap": "2024-10-01"}))
    stage = _service(db).get(data["case"]["id"])["stage"]
    updated = _commit(db, data["case"]["id"], data["case"]["revision"],
                      stage)
    heir = updated["stage"]["people"][1]
    assert heir["loai_giay_to"] == "Căn cước"
    assert heir["noi_cap"] == "Bộ Công an"
    assert heir["loai_dia_chi"] == "Cư trú tại"


def test_word_canonical_person_tokens_render(db):
    """Tên chuẩn của Người (ten/ngaysinh/sogiayto/loaigiayto/noicap/
    loaicutru/…) thay giá trị thật qua replace_in_doc, alias cũ giữ."""
    from docx import Document

    data = _create(db, with_state=True,
                   people_overrides=({}, {"ngay_cap": "2024-10-01"}))
    case = db.get(InheritanceCase, data["case"]["id"])
    mapping = word_engine.build_template_mapping(case)
    doc = Document()
    doc.add_paragraph(
        "[ten1]|[sogiayto1]|[ngaychet1]|[ten3]|[ngaysinh3]|[ngaycap3]"
        "|[loaigiayto3]|[noicap3]|[loaicutru3]|[diachi3]")
    word_engine.replace_in_doc(doc, mapping)
    text = doc.paragraphs[0].text
    assert text == ("Nguyễn Văn Chết|TLK 12/2011|15/05/2011|"
                    "Trần Thị Sống|02/03/1980|01/10/2024|"
                    "Căn cước|Bộ Công an|Cư trú tại|xã test")
    # Alias cũ vẫn hoạt động cùng giá trị.
    doc2 = Document()
    doc2.add_paragraph("[Tên 3]|[Loại CC 3]|[Thường trú 3]")
    word_engine.replace_in_doc(doc2, mapping)
    assert doc2.paragraphs[0].text == "Trần Thị Sống|Căn cước|Cư trú tại"


def test_word_canonical_case_meta_tokens_render(db):
    """Tên chuẩn Hồ sơ ngaylaphoso/ghichu/documenttype ngoài 3 trường
    đợt 3 — qua replace_in_doc thật."""
    from docx import Document

    uq = _catalog_person(db, "Bà Ủy Quyền")
    data = _create(db, with_state=True, case_meta={
        "document_type": "thoa_thuan", "ngay_lap_ho_so": "2026-09-05",
        "noiniemyet": "UBND xã Y", "nguoinhanuyquyenid": uq.id,
        "noidungviec": "Đính chính", "ghichu": "ghi chú A"})
    case = db.get(InheritanceCase, data["case"]["id"])
    doc = Document()
    doc.add_paragraph(
        "[ngaylaphoso]|[noiniemyet]|[nguoinhanuyquyen]|[noidungviec]"
        "|[ghichu]|[documenttype]")
    word_engine.replace_in_doc(
        doc, word_engine.build_template_mapping(case))
    assert doc.paragraphs[0].text == (
        "05/09/2026|UBND xã Y|Bà Ủy Quyền|Đính chính|ghi chú A|"
        "Thỏa thuận phân chia di sản")


def test_snapshot_explicit_empty_meta_kept_for_all_fields(db):
    """Snapshot commit null/rỗng chủ ý → Word xuất rỗng dù cột master
    sau đó đổi — áp dụng noiniemyet/nguoinhanuyquyen/noidungviec."""
    data = _create(db, with_state=True)  # meta để trống → snapshot null
    case = db.get(InheritanceCase, data["case"]["id"])
    case.noi_niem_yet = "UBND xã Đổi Sau"
    case.nguoi_nhan_uy_quyen = "Người Đổi Sau"
    case.noi_dung_viec = "Nội dung đổi sau"
    db.flush()
    mapping = word_engine.build_template_mapping(case)
    assert mapping["[Nơi niêm yết]"] == ""
    assert mapping["[Ngườ ủy quyền]"] == ""
    assert mapping["[Nội dung việc]"] == ""
    assert mapping["[noiniemyet]"] == ""
    assert mapping["[nguoinhanuyquyen]"] == ""
    assert mapping["[noidungviec]"] == ""


def test_snapshot_cleared_meta_via_commit_kept_empty(db):
    """Meta có giá trị → commit xóa (null) → master đổi sau → vẫn rỗng."""
    data = _create(db, with_state=True, case_meta={
        "document_type": "khai_nhan", "ngay_lap_ho_so": "2026-09-01",
        "noiniemyet": "UBND xã Cũ", "noidungviec": "Việc cũ"})
    case_id = data["case"]["id"]
    stage = _service(db).get(case_id)["stage"]
    _commit(db, case_id, data["case"]["revision"], stage,
            case_meta={"noiniemyet": None, "noidungviec": None})
    case = db.get(InheritanceCase, case_id)
    case.noi_niem_yet = "UBND xã Sửa Ngoài"
    case.noi_dung_viec = "Sửa ngoài"
    db.flush()
    mapping = word_engine.build_template_mapping(case)
    assert mapping["[Nơi niêm yết]"] == ""
    assert mapping["[Nội dung việc]"] == ""


def test_legacy_snapshot_without_case_block_reads_master(db):
    """Hồ sơ legacy: case_state_json không có block `case` → fallback
    cột master như trước (không nuốt giá trị)."""
    data = _create(db, with_state=True, case_meta={
        "document_type": "khai_nhan", "ngay_lap_ho_so": "2026-09-01",
        "noiniemyet": "UBND xã Snap", "noidungviec": "Việc snap"})
    case = db.get(InheritanceCase, data["case"]["id"])
    # Giả lập snapshot legacy: giữ stage/assets, bỏ block `case`.
    payload = json.loads(case.case_state_json)
    payload.pop("case", None)
    case.case_state_json = json.dumps(payload, ensure_ascii=False)
    case.noi_niem_yet = "UBND xã Master"
    case.noi_dung_viec = "Việc master"
    db.flush()
    mapping = word_engine.build_template_mapping(case)
    assert mapping["[Nơi niêm yết]"] == "UBND xã Master"
    assert mapping["[Nội dung việc]"] == "Việc master"
