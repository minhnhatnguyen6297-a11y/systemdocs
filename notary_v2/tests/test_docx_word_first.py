"""Synthetic labels for DOCX → review → saved Stage → original template."""
import asyncio
import hashlib
import io
import json
import unicodedata
import uuid
import base64
from zipfile import ZipFile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from services.document_intake.adapters.docx import extract
from services.document_intake.docx_source import read_docx, replace_spans
from services.document_intake.docx_store import DocxStore
from services.document_intake.models import AdapterContext, SourceSpec
from services.case_workspace import CaseWorkspaceService, WorkspaceError
from services.word_original import binding_proposals, reviewed_template, export_original, slot_for
from tests.test_case_workspace import db, session_factory, _person_row, _asset_row
from models import CaseAssetRow, InheritanceCase, Property, Customer


def fixture_bytes():
    doc = Document()
    doc.add_paragraph('HỢP ĐỒNG CHUYỂN NHƯỢNG')
    doc.add_paragraph('Bên A:')
    p = doc.add_paragraph('Ông: ')
    p.add_run('Nguyễn Văn ').bold = True
    p.add_run('Mẫu').bold = True
    p.add_run(', sinh ngày 01/02/1980; CCCD: 001234567890; thường trú: Nhà A')
    doc.add_paragraph('Bên B:')
    doc.add_paragraph('Bà: Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891; thường trú: Nhà B')
    doc.add_paragraph('Tài sản 1:')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Số serial: ZZ123456; có địa chỉ tại: Nhà đất A\nThửa đất số: 12; tờ bản đồ số: 5'
    doc.add_paragraph('Loại đất: ONT; Diện tích: 50; Thời hạn: Lâu dài')
    doc.add_paragraph('Loại đất: CLN; Diện tích: 20; Thời hạn: 2050')
    doc.add_paragraph('Tài sản 2:')
    doc.add_paragraph('Số serial: YY123456; có địa chỉ tại: Nhà đất B\nThửa đất số: 22; tờ bản đồ số: 7')
    doc.add_paragraph('Loại đất: ONT; Diện tích: 90; Thời hạn: Lâu dài')
    doc.add_paragraph('Điều 2: Nội dung giữ nguyên 😀')
    doc.add_paragraph('Lời chứng: Nguyễn Văn Mẫu — tên nhắc lại giữ nguyên')
    doc.sections[0].header.paragraphs[0].text = 'ĐẦU TRANG GIỮ NGUYÊN'
    doc.add_picture(io.BytesIO(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a3ioAAAAASUVORK5CYII=')))
    stream = io.BytesIO(); doc.save(stream)
    return stream.getvalue()


def _docx_bytes(lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    stream = io.BytesIO(); doc.save(stream)
    return stream.getvalue()


def extraction(content, store=None):
    spec = SourceSpec(str(uuid.uuid4()), 'docx', filename='synthetic.docx')
    ctx = AdapterContext('', '', None, None, read_file=lambda _: content,
                         store_docx=store.save if store else None)
    return asyncio.run(extract(spec, ctx))


@pytest.mark.parametrize('issuer', ['Bộ Coongan', 'Cơ quan Mẫu giữ chữ tự do'])
def test_literal_issuer_and_number_survive_commit_manual_edit_reload_with_exact_field_evidence(db, tmp_path, issuer):
    doc = Document()
    for line in ['HỢP ĐỒNG CHUYỂN NHƯỢNG', 'Bên A:',
                 f'Ông Lê Văn Mẫu, sinh ngày 01/02/1980; CCCD: 0123456789do {issuer} cấp; thường trú: Nhà Mẫu',
                 'Bên B:', 'Tài sản 1:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà đất Mẫu',
                 'Loại đất: ONT; Diện tích: 50; Thời hạn: Lâu dài', 'Điều 2: Nội dung']:
        doc.add_paragraph(line)
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue()
    suggestion, = [s for s in extraction(content) if s['target'] == 'person']
    for key, literal in [('so_giay_to', '0123456789'), ('noi_cap', issuer)]:
        field = suggestion['fields'][key]
        assert field['raw_value'] == field['normalized_value'] == literal
        ref, = field['source_refs']
        assert ref['raw_snippet'] == literal
        assert read_docx(content).text[slice(*ref['span'])] == literal
        assert ref['docx_refs']
    original, store, service, data = reviewed_case(db, tmp_path, content)
    cid = data['case']['id']
    person, = data['stage']['people']
    assert (person['so_giay_to'], person['noi_cap']) == ('0123456789', issuer)
    person['noi_cap'] = 'Cơ quan người dùng sửa tay'
    person['so_giay_to'] = 'Số do người dùng ghi tay'
    service.commit_stage(cid, data['case']['revision'], data['stage'], docx_import=data['docx_import'])
    restored, = service.get(cid)['stage']['people']
    assert (restored['so_giay_to'], restored['noi_cap']) == ('Số do người dùng ghi tay', 'Cơ quan người dùng sửa tay')
    sid = data['docx_import']['sources'][0]['source_id']
    assert store.load(sid)[0] == original == content


@pytest.mark.parametrize('raw', ['31/02/1980', '1980-02-31', '1980-2-31'])
def test_invalid_source_date_requires_correction_and_commits_exact_original_evidence(db, tmp_path, raw):
    doc = Document(io.BytesIO(fixture_bytes()))
    for p in doc.paragraphs:
        if '01/02/1980' in p.text:
            for run in p.runs:
                run.text = run.text.replace('01/02/1980', raw)
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue()
    store = DocxStore(tmp_path / 'sources')
    suggestions = extraction(content, store)
    person = next(s for s in suggestions if s['target'] == 'person')
    field = person['fields']['ngay_sinh']; ref = field['source_refs'][0]
    assert field['normalized_value'] is None
    assert ref['core_state'] == 'error' and ref['raw_snippet'] == raw
    assert read_docx(content).text[slice(*ref['span'])] == raw
    assert any(w['code'] == 'intake.invalid_date' for w in person['warnings'])
    _, _, service, data = reviewed_case(db, tmp_path / 'case', content)
    data['stage']['people'][0]['ngay_sinh'] = raw
    with pytest.raises(WorkspaceError):
        service.commit_stage(data['case']['id'], data['case']['revision'], data['stage'], docx_import=data['docx_import'])
    data['stage']['people'][0]['ngay_sinh'] = '1980-02-29'
    saved = service.commit_stage(data['case']['id'], data['case']['revision'], data['stage'], docx_import=data['docx_import'])
    assert service.get(data['case']['id'])['stage']['people'][0]['ngay_sinh'] == '1980-02-29'
    assert saved['revision'] == data['case']['revision'] + 1
    assert store.load(ref['source_id'])[0] == content


@pytest.mark.parametrize('raw,expected', [('1980-02-29', '1980-02-29'), ('1980-2-9', '1980-02-09'), ('1980', '1980')])
def test_valid_iso_and_year_only_source_dates_remain_accepted(raw, expected):
    doc = Document(io.BytesIO(fixture_bytes()))
    for p in doc.paragraphs:
        for run in p.runs:
            run.text = run.text.replace('01/02/1980', raw)
    stream = io.BytesIO(); doc.save(stream)
    field = extraction(stream.getvalue())[0]['fields']['ngay_sinh']
    assert field['normalized_value'] == expected
    assert field['source_refs'][0]['raw_snippet'] == raw
    assert field['source_refs'][0]['core_state'] == 'matched'


@pytest.mark.parametrize('raw', ['1980-02-31','31/02/1980','31.02.1980','1980/2/31'])
def test_all_recognized_complete_dates_validate_calendar(raw):
    from services.document_intake.normalization import norm_date
    assert norm_date(raw)[0] is None


@pytest.mark.parametrize('raw', ['-50', '−50', '1e3', '1E-3', '−1.234,5e+3', '-' + '9' * 400, '9' * 400])
@pytest.mark.parametrize('land', [False, True])
def test_invalid_source_area_keeps_complete_evidence_and_corrected_span(raw, land):
    doc = Document()
    for line in ['Bên A:', 'Ông: Nguyễn Văn Mẫu, CCCD: 001234567890', 'Tài sản 1:',
                 'Serial: ZZ123456; có địa chỉ tại: Nhà A',
                 ('Loại đất: ONT; ' if land else '') + f'Diện tích: {raw}; Thời hạn: Lâu dài',
                 'Điều 2: giữ nguyên']:
        doc.add_paragraph(line)
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue()
    suggestions = extraction(content)
    asset = next(s for s in suggestions if s['target'] == 'asset')
    field = asset['fields']['land_rows' if land else 'dien_tich']
    if land:
        assert json.loads(field['normalized_value'])[0]['dien_tich'] is None
        ref = next(r for r in field['source_refs'] if r['land_field'] == 'dien_tich')
    else:
        assert field['normalized_value'] is None
        ref = field['source_refs'][0]
    assert ref['raw_snippet'] == raw and ref['core_state'] == 'error'
    assert any(w['code'] == 'intake.invalid_number' for w in asset['warnings'])
    json.dumps(suggestions, allow_nan=False)
    source = read_docx(content)
    corrected = replace_spans(content, source, [(ref['span'], '1234.5')])
    assert 'Diện tích: 1234.5;' in read_docx(corrected).text
    assert raw not in read_docx(corrected).text


@pytest.mark.parametrize('raw,expected', [('1.234,5', 1234.5), ('+50', 50), ('5', 5)])
@pytest.mark.parametrize('land', [False, True])
def test_supported_source_area_keeps_whole_raw_span(raw, expected, land):
    doc = Document()
    for line in ['Bên A:', 'Ông: Nguyễn Văn Mẫu, CCCD: 001234567890', 'Tài sản 1:',
                 'Serial: ZZ123456; có địa chỉ tại: Nhà A',
                 ('Loại đất: ONT; ' if land else '') + f'Diện tích: {raw}; Thời hạn: Lâu dài',
                 'Điều 2: giữ nguyên']:
        doc.add_paragraph(line)
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue()
    asset = next(s for s in extraction(content) if s['target'] == 'asset')
    field = asset['fields']['land_rows' if land else 'dien_tich']
    value = json.loads(field['normalized_value'])[0]['dien_tich'] if land else field['normalized_value']
    assert value == expected
    ref = next(r for r in field['source_refs'] if not land or r.get('land_field') == 'dien_tich')
    assert ref['raw_snippet'] == raw and ref['core_state'] == 'matched'
    assert read_docx(content).text[slice(*ref['span'])] == raw
    changed = replace_spans(content,read_docx(content),[(ref['span'],'25')])
    assert 'Diện tích: 25; Thời hạn: Lâu dài' in read_docx(changed).text


def test_identifier_missing_assets_retain_address_and_land_but_not_empty_noise():
    doc = Document()
    for line in ['Bên A:', 'Ông: Nguyễn Văn Mẫu, CCCD: 001234567890',
                 'Tài sản 1:', 'có địa chỉ tại: Nhà A',
                 'Tài sản 2:', 'Loại đất: ONT; Diện tích: 50; Thời hạn: Lâu dài',
                 'Tài sản 3:', 'chỉ là nhãn trống', 'Điều 2: giữ nguyên']:
        doc.add_paragraph(line)
    stream = io.BytesIO(); doc.save(stream)
    assets = [s for s in extraction(stream.getvalue()) if s['target'] == 'asset']
    assert len(assets) == 2
    assert assets[0]['fields']['dia_chi']['normalized_value'] == 'Nhà A'
    assert json.loads(assets[1]['fields']['land_rows']['normalized_value'])[0]['dien_tich'] == 50
    assert all(any(w['code'] == 'intake.fragment' for w in a['warnings']) for a in assets)
    assert all('so_serial' not in a['fields'] for a in assets)
    # An unlabelled address in prose does not create another asset.
    doc = Document(); doc.add_paragraph('Bên A:'); doc.add_paragraph('Ông: Nguyễn Văn Mẫu, CCCD: 001234567890')
    doc.add_paragraph('có địa chỉ tại: ghi chú rời không có nhãn tài sản')
    stream = io.BytesIO(); doc.save(stream)
    assert not any(s['target'] == 'asset' for s in extraction(stream.getvalue()))


@pytest.mark.parametrize('limit', ['expanded', 'text'])
def test_actual_docx_analyze_limits_keep_valid_batch_source_and_reject_storage(tmp_path, limit):
    from services.document_intake.service import analyze
    from zipfile import ZIP_DEFLATED
    valid = fixture_bytes()
    if limit == 'text':
        doc = Document(io.BytesIO(valid)); doc.add_paragraph('X' * 100001)
        stream = io.BytesIO(); doc.save(stream); invalid = stream.getvalue()
        control = Document(io.BytesIO(valid)); control.add_paragraph('X' * 1000)
        stream = io.BytesIO(); control.save(stream)
        assert len(extraction(stream.getvalue())) == 4
    else:
        stream = io.BytesIO()
        with ZipFile(io.BytesIO(valid)) as src, ZipFile(stream, 'w', compression=ZIP_DEFLATED) as dst:
            for item in src.infolist(): dst.writestr(item, src.read(item.filename))
            dst.writestr('word/oversize.bin', b'0' * (100 * 1024 * 1024))
        invalid = stream.getvalue()
    bad_id, good_id = str(uuid.uuid4()), str(uuid.uuid4())
    sources = [{'source_id': sid, 'kind': 'docx', 'file_ref': {'path': f'C:/synthetic/{sid}.docx', 'scope': 'machine_local', 'is_dir': False,
                'size_bytes': len(invalid if sid == bad_id else valid)}}
               for sid in [bad_id, good_id]]
    store = DocxStore(tmp_path / 'sources'); stored = []
    def save(content, suggestions, label):
        stored.append(content); return store.save(content, suggestions, label)
    outcome = analyze(sources, read_file=lambda s: invalid if s.source_id == bad_id else valid, store_docx=save)
    assert outcome.status == 'partial' and outcome.failed == [bad_id] and outcome.succeeded == [good_id]
    assert len(outcome.suggestions) == 4 and len(outcome.errors) == 1
    assert outcome.errors[0]['source_id'] == bad_id and outcome.errors[0]['code'] == 'intake.parse_failed'
    assert stored == [valid] and len(list(store.root.iterdir())) == 1


def test_asset_heading_before_parties_never_becomes_an_asset_group():
    content = fixture_bytes()
    doc = Document(io.BytesIO(content))
    doc.paragraphs[0].insert_paragraph_before('Tài sản 99: bản mô tả ở đầu hợp đồng')
    stream = io.BytesIO(); doc.save(stream)
    suggestions = extraction(stream.getvalue())
    assets = [s for s in suggestions if s['target'] == 'asset']
    assert len(assets) == 2
    assert [s['fields']['so_serial']['normalized_value'] for s in assets] == ['ZZ123456', 'YY123456']
    source = read_docx(stream.getvalue())
    for asset in assets:
        assert all(ref['span'][0] > source.text.index('Bên B:') for f in asset['fields'].values() for ref in f['source_refs'])


def test_blank_land_row_does_not_drop_valid_people_or_assets(tmp_path):
    doc = Document(io.BytesIO(fixture_bytes()))
    for p in doc.paragraphs:
        if p.text.startswith('Loại đất:'):
            p.text = 'Loại đất:'
    stream = io.BytesIO(); doc.save(stream)
    store = DocxStore(tmp_path / 'sources')
    suggestions = extraction(stream.getvalue(), store)
    assert len(suggestions) == 4
    assert len([s for s in suggestions if s['target'] == 'person']) == 2
    assets = [s for s in suggestions if s['target'] == 'asset']
    assert [s['fields']['so_serial']['normalized_value'] for s in assets] == ['ZZ123456', 'YY123456']
    for asset in assets:
        assert 'land_rows' not in asset['fields']
        assert any('Cụm đất' in w['message'] and 'thiếu dữ liệu' in w['message'] for w in asset['warnings'])
    ref = suggestions[0]['fields']['ho_ten']['source_refs'][0]
    assert store.load(ref['source_id'])[0] == stream.getvalue()


def test_source_publication_failure_exposes_no_partial_source_and_retry_succeeds(tmp_path, monkeypatch):
    store = DocxStore(tmp_path / 'sources'); write_bytes = Path.write_bytes
    def fail_extraction(path, content):
        if path.name == 'extraction.json':
            write_bytes(path, content[:20])
            raise OSError('interrupted')
        return write_bytes(path, content)
    with monkeypatch.context() as patch:
        patch.setattr(Path, 'write_bytes', fail_extraction)
        with pytest.raises(OSError, match='interrupted'):
            extraction(fixture_bytes(), store)
        assert list(store.root.iterdir()) == []
    content = fixture_bytes(); candidates = extraction(content, store)
    ref = candidates[0]['fields']['ho_ten']['source_refs'][0]
    assert store.load(ref['source_id'])[0] == content


def test_template_publication_failure_retry_and_existing_immutable_file(tmp_path, monkeypatch):
    import services.document_intake.docx_store as storage
    store = DocxStore(tmp_path / 'sources'); content = fixture_bytes()
    candidates = extraction(content, store); sid = candidates[0]['fields']['ho_ten']['source_refs'][0]['source_id']
    template = b'synthetic template bytes'; digest = hashlib.sha256(template).hexdigest()
    folder = store.root / sid / 'templates'
    def fail_link(source, dest):
        assert Path(source).read_bytes() == template
        assert not Path(dest).exists()
        raise PermissionError('interrupted publication')
    with monkeypatch.context() as patch:
        patch.setattr(storage.os, 'link', fail_link)
        with pytest.raises(PermissionError): store.save_template(sid, template)
    assert list(folder.iterdir()) == []
    assert store.save_template(sid, template) == digest
    path = folder / (digest + '.docx'); before = path.stat().st_mtime_ns
    assert store.save_template(sid, template) == digest
    assert path.stat().st_mtime_ns == before and path.read_bytes() == template
    assert store.load(sid)[0] == content


def test_original_export_cancel_after_temp_write_leaves_no_output_and_preserves_source(db, tmp_path):
    original, store, service, data = reviewed_case(db, tmp_path)
    cid = data['case']['id']; source = data['docx_import']['sources'][0]
    proposals = binding_proposals(db.get(InheritanceCase, cid), source['source_id'], store)
    source['bindings'] = [{k:p[k] for k in ('candidate_id','occurrence_id','field_key','slot')} for p in proposals if p['supported']]
    service.commit_stage(cid, data['case']['revision'], data['stage'], docx_import=data['docx_import'])
    dest = tmp_path / 'exports'; dest.mkdir(); seen = []
    def cancel():
        files = list(dest.iterdir()); assert len(files) == 1
        assert files[0].name.startswith('.word_export_') and read_docx(files[0].read_bytes()).text
        seen.append(files[0]); raise RuntimeError('user_canceled')
    with pytest.raises(RuntimeError, match='user_canceled'):
        export_original(db.get(InheritanceCase, cid), source['source_id'], store, dest, check_cancel=cancel)
    assert seen and list(dest.iterdir()) == []
    assert store.load(source['source_id'])[0] == original


def test_labeled_group_ownership_and_raw_refs():
    content = fixture_bytes()
    source = read_docx(content)
    suggestions = extraction(content)
    people = [s for s in suggestions if s['target'] == 'person']
    assets = [s for s in suggestions if s['target'] == 'asset']
    assert len(people) == len(assets) == 2
    assert [p['fields']['ho_ten']['normalized_value'] for p in people] == ['Nguyễn Văn Mẫu', 'Trần Thị Mẫu']
    assert [p['fields']['so_giay_to']['normalized_value'] for p in people] == ['001234567890', '001234567891']
    assert [a['fields']['so_serial']['normalized_value'] for a in assets] == ['ZZ123456', 'YY123456']
    assert [len(json.loads(a['fields']['land_rows']['normalized_value'])) for a in assets] == [2, 1]
    for s in suggestions:
        for field in s['fields'].values():
            for ref in field['source_refs']:
                lo, hi = ref['span']
                assert source.text[lo:hi] == ref['raw_snippet']
                assert ref['docx_refs']


def test_ambiguity_retains_all_evidence_and_duplicates_never_merge():
    doc = Document()
    for line in ['Bên A:', 'Ông: Nguyễn Văn Mẫu, CCCD: 001234567890; CCCD: 001234567899',
                 'Bên B:', 'Bà: Trần Thị Mẫu, CCCD: 001234567890',
                 'Tài sản 1:', 'Serial: Z123456; có địa chỉ tại: Nhà A',
                 'Tài sản 2:', 'Serial: YY123456; có địa chỉ tại: Nhà B',
                 'Tài sản 3:', 'Serial: YY123456; có địa chỉ tại: Nhà C', 'Điều 2: Giữ nguyên']:
        doc.add_paragraph(line)
    stream = io.BytesIO(); doc.save(stream)
    suggestions = extraction(stream.getvalue())
    assert len(suggestions) == 5
    field = suggestions[0]['fields']['so_giay_to']
    assert field['normalized_value'] is None and len(field['source_refs']) == 2
    assert suggestions[2]['fields']['so_serial']['normalized_value'] is None
    assert all(any(w['code'] == 'intake.suspected_duplicate' for w in s['warnings']) for s in suggestions[3:])


def test_two_party_mapping_uses_saved_seat_not_source_or_stage_order():
    snapshot = {'case_type': 'two_party', 'stage': [{'row_id': 'first'}, {'row_id': 'second'}],
                'diagram': {'state': {'nodes': [{'id': 'p17', 'personId': 'first'}, {'id': 'p2', 'personId': 'second'}]}}}
    assert slot_for(snapshot, 'person', 'first', 'ho_ten') == '[Bên B 2 - Họ tên]'
    assert slot_for(snapshot, 'person', 'second', 'ho_ten') == '[Bên A 2 - Họ tên]'


def test_physical_merged_nested_tables_unicode_and_repeated_occurrences():
    doc = Document()
    doc.add_paragraph('TRƯỚC 😀')
    table = doc.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = 'GỘP'
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'LỒNG'
    text = unicodedata.normalize('NFD', 'Tên Đặng')
    p = doc.add_paragraph(); p.add_run(text[:4]).bold = True; p.add_run(text[4:]).bold = True
    doc.add_paragraph(text)
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue()
    source = read_docx(content)
    assert source.text.count('GỘP') == 1
    assert source.text.index('TRƯỚC') < source.text.index('GỘP') < source.text.index('LỒNG') < source.text.index(text)
    start = source.text.index(text)
    rendered = replace_spans(content, source, [([start, start + len(text)], 'Đã sửa')])
    assert read_docx(rendered).text.count(text) == 1
    assert 'Đã sửa' in read_docx(rendered).text
    with ZipFile(io.BytesIO(content)) as a, ZipFile(io.BytesIO(rendered)) as b:
        for name in a.namelist():
            if name != 'word/document.xml':
                assert a.read(name) == b.read(name)
        assert b'<w:b' in b.read('word/document.xml')


def test_unsupported_hyperlink_and_overlap_rejected():
    doc = Document(); p = doc.add_paragraph('A ')
    hyperlink = OxmlElement('w:hyperlink'); run = OxmlElement('w:r'); text = OxmlElement('w:t')
    text.text = 'Không sửa'; run.append(text); hyperlink.append(run); p._p.append(hyperlink)
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue(); source = read_docx(content)
    start = source.text.index('Không sửa')
    assert source.warnings
    with pytest.raises(ValueError, match='chưa hỗ trợ'):
        replace_spans(content, source, [([start, start + 3], 'X')])
    with pytest.raises(ValueError, match='chồng'):
        replace_spans(content, source, [([0, 1], 'X'), ([0, 2], 'Y')])


def reviewed_case(db, tmp_path, content=None):
    store = DocxStore(tmp_path / 'sources')
    content = content or fixture_bytes(); suggestions = extraction(content, store)
    people, assets, records = [], [], []
    source_ref = next(iter(suggestions[0]['fields'].values()))['source_refs'][0]
    for suggestion in suggestions:
        values = {key: field['normalized_value'] for key, field in suggestion['fields'].items()}
        # Stage stores areas in land rows; scalar source evidence maps to the saved total.
        values.pop('dien_tich', None)
        if 'land_rows' in values:
            values['land_rows'] = json.loads(values['land_rows'])
        row = _person_row(**values) if suggestion['target'] == 'person' else _asset_row(**values)
        (people if suggestion['target'] == 'person' else assets).append(row)
        records.append({'candidate_id': suggestion['suggestion_id'], 'row_id': row['row_id']})
    context = {'version': 1, 'sources': [{'source_id': source_ref['source_id'], 'sha256': source_ref['source_sha256'], 'records': records, 'bindings': []}]}
    service = CaseWorkspaceService(db, store)
    data = service.create(str(uuid.uuid4()), {'case_type': 'inheritance', 'document_type': 'khai_nhan'},
        {'owner_row_id': people[0]['row_id'], 'people': people, 'assets': assets}, None, context)
    return content, store, service, data


def reviewed_case_with_missing_asset_values(db, tmp_path, content):
    """Keep an absent asset field absent instead of applying fixture defaults."""
    store = DocxStore(tmp_path / 'sources')
    suggestions = extraction(content, store)
    people, assets, records = [], [], []
    source_ref = next(iter(suggestions[0]['fields'].values()))['source_refs'][0]
    for suggestion in suggestions:
        values = {key: field['normalized_value'] for key, field in suggestion['fields'].items()}
        values.pop('dien_tich', None)
        if 'land_rows' in values:
            values['land_rows'] = json.loads(values['land_rows'])
        if suggestion['target'] == 'asset':
            values.setdefault('nguon_goc', None)
            # S5b 23/10/2026: số thửa + số tờ là mức tối thiểu của dòng
            # Tài sản; giữ phép thử này tập trung vào nguồn gốc.
            values['so_thua_dat'] = values.get('so_thua_dat') or '0'
            values['so_to_ban_do'] = values.get('so_to_ban_do') or '0'
        row = _person_row(**values) if suggestion['target'] == 'person' else _asset_row(**values)
        (people if suggestion['target'] == 'person' else assets).append(row)
        records.append({'candidate_id': suggestion['suggestion_id'], 'row_id': row['row_id']})
    context = {'version': 1, 'sources': [{'source_id': source_ref['source_id'], 'sha256': source_ref['source_sha256'], 'records': records, 'bindings': []}]}
    service = CaseWorkspaceService(db, store)
    data = service.create(str(uuid.uuid4()), {'case_type': 'inheritance', 'document_type': 'khai_nhan'},
        {'owner_row_id': people[0]['row_id'], 'people': people, 'assets': assets}, None, context)
    return content, store, service, data


def reviewed_case_with_source_gender(db, tmp_path, content):
    """Reviewed case that keeps an absent gender absent (no _person_row default)."""
    store = DocxStore(tmp_path / 'sources')
    suggestions = extraction(content, store)
    people, assets, records = [], [], []
    source_ref = next(iter(suggestions[0]['fields'].values()))['source_refs'][0]
    for suggestion in suggestions:
        values = {key: field['normalized_value'] for key, field in suggestion['fields'].items()}
        values.pop('dien_tich', None)
        if 'land_rows' in values:
            values['land_rows'] = json.loads(values['land_rows'])
        if suggestion['target'] == 'person':
            values.setdefault('gioi_tinh', None)
        row = _person_row(**values) if suggestion['target'] == 'person' else _asset_row(**values)
        (people if suggestion['target'] == 'person' else assets).append(row)
        records.append({'candidate_id': suggestion['suggestion_id'], 'row_id': row['row_id']})
    context = {'version': 1, 'sources': [{'source_id': source_ref['source_id'], 'sha256': source_ref['source_sha256'], 'records': records, 'bindings': []}]}
    service = CaseWorkspaceService(db, store)
    data = service.create(str(uuid.uuid4()), {'case_type': 'inheritance', 'document_type': 'khai_nhan'},
        {'owner_row_id': people[0]['row_id'], 'people': people, 'assets': assets}, None, context)
    return content, store, service, data


def test_asset_origin_roundtrips_per_asset_not_broadcast(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:',
        'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Nguồn gốc sử dụng đất: Nguồn A',
        'Nguồn gốc sử dụng đất: Nguồn C',
        'Thửa đất số: 22; Tờ bản đồ số: 7',
        'Nguồn gốc sử dụng đất: Nguồn B',
        'Thửa đất số: 32; Tờ bản đồ số: 9',
    ])
    suggestions = extraction(content)
    assets = [s for s in suggestions if s['target'] == 'asset']
    assert [s['fields'].get('nguon_goc', {}).get('normalized_value') for s in assets] == [
        None, 'Nguồn B', None]
    assert 'nguon_goc' not in assets[2]['fields']
    ambiguous = assets[0]['fields']['nguon_goc']
    assert ambiguous['normalized_value'] is None
    assert [r['raw_snippet'] for r in ambiguous['source_refs']] == ['Nguồn A', 'Nguồn C']
    assert any(w['code'] == 'intake.ambiguous' and 'nguon_goc' in w['message']
               for w in assets[0]['warnings'])

    _, _, service, data = reviewed_case_with_missing_asset_values(db, tmp_path, content)
    cid = data['case']['id']
    assert [a['nguon_goc'] for a in data['stage']['assets']] == [None, 'Nguồn B', None]
    db.expire_all()
    reloaded = service.get(cid)
    assert [a['nguon_goc'] for a in reloaded['stage']['assets']] == [None, 'Nguồn B', None]
    rows = db.query(CaseAssetRow).filter_by(case_id=cid).order_by(
        CaseAssetRow.vitri).all()
    assert [row.nguon_goc for row in rows] == [None, 'Nguồn B', None]
    from services.word_engine import build_template_mapping
    text = build_template_mapping(
        db.get(InheritanceCase, cid), static=True)['[Đoạn mô tả di sản]']
    assert 'Nguồn B' in text and 'Nguồn A' not in text and 'Nguồn C' not in text
    assert text.count('Nguồn gốc sử dụng đất:') == 3


def test_group_asset_origins_stay_scoped_and_ambiguous_not_broadcast(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản 1:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5', 'Nguồn gốc sử dụng đất: Nguồn A',
        'Tài sản 2:', 'Số serial: YY123456; có địa chỉ tại: Nhà B',
        'Thửa đất số: 22; Tờ bản đồ số: 7', 'Nguồn gốc sử dụng đất: Nguồn B',
        'Tài sản 3:', 'Số serial: XX123456; có địa chỉ tại: Nhà C',
        'Thửa đất số: 32; Tờ bản đồ số: 9', 'Nguồn gốc sử dụng đất: Nguồn C',
        'Thửa đất số: 33; Tờ bản đồ số: 10', 'Nguồn gốc sử dụng đất: Nguồn D',
    ])
    suggestions = extraction(content)
    assets = [s for s in suggestions if s['target'] == 'asset']
    assert [s['fields'].get('nguon_goc', {}).get('normalized_value') for s in assets] == [
        'Nguồn A', 'Nguồn B', None]
    ambiguous = assets[2]['fields']['nguon_goc']
    assert ambiguous['normalized_value'] is None
    assert [r['raw_snippet'] for r in ambiguous['source_refs']] == ['Nguồn C', 'Nguồn D']
    assert any(w['code'] == 'intake.ambiguous' and 'nguon_goc' in w['message']
               for w in assets[2]['warnings'])

    _, _, service, data = reviewed_case_with_missing_asset_values(db, tmp_path, content)
    cid = data['case']['id']
    assert [a['nguon_goc'] for a in data['stage']['assets']] == ['Nguồn A', 'Nguồn B', None]
    db.expire_all()
    reloaded = service.get(cid)
    assert [a['nguon_goc'] for a in reloaded['stage']['assets']] == ['Nguồn A', 'Nguồn B', None]
    rows = db.query(CaseAssetRow).filter_by(case_id=cid).order_by(
        CaseAssetRow.vitri).all()
    assert [row.nguon_goc for row in rows] == ['Nguồn A', 'Nguồn B', None]
    from services.word_engine import build_template_mapping
    text = build_template_mapping(
        db.get(InheritanceCase, cid), static=True)['[Đoạn mô tả di sản]']
    assert 'Nguồn A' in text and 'Nguồn B' in text
    assert 'Nguồn C' not in text and 'Nguồn D' not in text


def reviewed_case_with_source_form(db, tmp_path, content):
    """Reviewed case that keeps an absent hinh_thuc_su_dung absent."""
    store = DocxStore(tmp_path / 'sources')
    suggestions = extraction(content, store)
    people, assets, records = [], [], []
    source_ref = next(iter(suggestions[0]['fields'].values()))['source_refs'][0]
    for suggestion in suggestions:
        values = {key: field['normalized_value'] for key, field in suggestion['fields'].items()}
        values.pop('dien_tich', None)
        if 'land_rows' in values:
            values['land_rows'] = json.loads(values['land_rows'])
        if suggestion['target'] == 'asset':
            values.setdefault('hinh_thuc_su_dung', None)
        row = _person_row(**values) if suggestion['target'] == 'person' else _asset_row(**values)
        (people if suggestion['target'] == 'person' else assets).append(row)
        records.append({'candidate_id': suggestion['suggestion_id'], 'row_id': row['row_id']})
    context = {'version': 1, 'sources': [{'source_id': source_ref['source_id'], 'sha256': source_ref['source_sha256'], 'records': records, 'bindings': []}]}
    service = CaseWorkspaceService(db, store)
    data = service.create(str(uuid.uuid4()), {'case_type': 'inheritance', 'document_type': 'khai_nhan'},
        {'owner_row_id': people[0]['row_id'], 'people': people, 'assets': assets}, None, context)
    return content, store, service, data


def test_person_gender_from_titles_roundtrips_with_title_provenance(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
    ])
    suggestions = extraction(content)
    people = [s for s in suggestions if s['target'] == 'person']
    assert [s['fields']['ho_ten']['normalized_value'] for s in people] == [
        'Nguyễn Văn Mẫu', 'Trần Thị Mẫu']
    assert [s['fields']['gioi_tinh']['normalized_value'] for s in people] == ['Nam', 'Nữ']
    assert [s['fields']['gioi_tinh']['raw_value'] for s in people] == ['Ông', 'Bà']
    assert [r['raw_snippet'] for s in people
            for r in s['fields']['gioi_tinh']['source_refs']] == ['Ông', 'Bà']
    assert all(r['rule_id'] == 'person.gioi_tinh'
               for s in people for r in s['fields']['gioi_tinh']['source_refs'])

    _, _, service, data = reviewed_case_with_source_gender(db, tmp_path, content)
    cid = data['case']['id']
    assert [p['gioi_tinh'] for p in data['stage']['people']] == ['Nam', 'Nữ']
    assert [p['ho_ten'] for p in data['stage']['people']] == ['Nguyễn Văn Mẫu', 'Trần Thị Mẫu']
    assert not {'Ông', 'Bà'} & {p['gioi_tinh'] for p in data['stage']['people']}
    db.expire_all()
    reloaded = service.get(cid)
    assert [p['gioi_tinh'] for p in reloaded['stage']['people']] == ['Nam', 'Nữ']
    # Danh tính dùng chung chỉ giữ neo tên+số giấy; giới tính là dữ kiện
    # của dòng Hồ sơ (đã assert ở trên), không ghi ngược master.
    assert {c.ho_ten: c.gioi_tinh for c in db.query(Customer).all()} == {
        'Nguyễn Văn Mẫu': None, 'Trần Thị Mẫu': None}
    from services.word_engine import build_word_context
    context = build_word_context(db.get(InheritanceCase, cid))
    assert [(p.ho_ten, p.gender_normalized) for p in context.persons] == [
        ('Nguyễn Văn Mẫu', 'nam'), ('Trần Thị Mẫu', 'nữ')]


def test_person_gender_neutral_title_stays_absent_per_person(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông/bà Nguyễn Thị Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
    ])
    suggestions = extraction(content)
    people = [s for s in suggestions if s['target'] == 'person']
    assert 'gioi_tinh' not in people[0]['fields']
    assert people[1]['fields']['gioi_tinh']['normalized_value'] == 'Nữ'

    _, _, service, data = reviewed_case_with_source_gender(db, tmp_path, content)
    cid = data['case']['id']
    assert [p['gioi_tinh'] for p in data['stage']['people']] == [None, 'Nữ']
    db.expire_all()
    assert [p['gioi_tinh'] for p in service.get(cid)['stage']['people']] == [None, 'Nữ']
    assert {c.ho_ten: c.gioi_tinh for c in db.query(Customer).all()} == {
        'Nguyễn Thị Mẫu': None, 'Trần Thị Mẫu': None}
    from services.word_engine import build_word_context
    context = build_word_context(db.get(InheritanceCase, cid))
    assert [(p.ho_ten, p.gender_normalized) for p in context.persons] == [
        ('Nguyễn Thị Mẫu', ''), ('Trần Thị Mẫu', 'nữ')]


def test_person_gender_representative_title_maps_to_nam(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:',
        'Đại diện: Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
    ])
    people = [s for s in extraction(content) if s['target'] == 'person']
    representative = people[0]
    assert representative['fields']['ho_ten']['normalized_value'] == 'Nguyễn Văn Mẫu'
    assert representative['fields']['gioi_tinh']['normalized_value'] == 'Nam'
    ref, = representative['fields']['gioi_tinh']['source_refs']
    assert ref['raw_snippet'] == 'Ông' and ref['rule_id'] == 'person.gioi_tinh'


def test_person_gender_table_row_title_evidence_without_raw_text(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'DANH SÁCH NGƯỜI SỬ DỤNG CHUNG',
        '1 | Ông Nguyễn Văn Mẫu | 1980 | CCCD 001234567890',
        '2 | Chị Trần Thị Mẫu | 1985 | CCCD 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
    ])
    people = [s for s in extraction(content) if s['target'] == 'person']
    assert [s['fields']['ho_ten']['normalized_value'] for s in people] == [
        'Nguyễn Văn Mẫu', 'Trần Thị Mẫu']
    assert [s['fields']['gioi_tinh']['normalized_value'] for s in people] == ['Nam', 'Nữ']
    # Lõi in Anh/Chị thành Ông/Bà; bằng chứng phải giữ đúng token nguồn.
    assert [r['raw_snippet'] for s in people
            for r in s['fields']['gioi_tinh']['source_refs']] == ['Ông', 'Chị']
    _, _, service, data = reviewed_case_with_source_gender(db, tmp_path, content)
    assert [p['gioi_tinh'] for p in data['stage']['people']] == ['Nam', 'Nữ']


def test_form_from_explicit_label_roundtrips_per_parcel_without_notes(db, tmp_path, session_factory):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Hình thức sử dụng: Sử dụng riêng; Ghi chú: không lấy phần này',
        'Thửa đất số: 22; Tờ bản đồ số: 7',
        'Hình thức sử dụng: Sử dụng chung',
    ])
    assets = [s for s in extraction(content) if s['target'] == 'asset']
    assert [a['fields']['hinh_thuc_su_dung']['normalized_value'] for a in assets] == [
        'Sử dụng riêng', 'Sử dụng chung']
    assert [a['fields']['hinh_thuc_su_dung']['raw_value'] for a in assets] == [
        'Sử dụng riêng', 'Sử dụng chung']
    assert [r['raw_snippet'] for a in assets
            for r in a['fields']['hinh_thuc_su_dung']['source_refs']] == [
                'Sử dụng riêng', 'Sử dụng chung']
    assert all('Ghi chú' not in a['fields']['hinh_thuc_su_dung']['raw_value'] for a in assets)

    _, _, service, data = reviewed_case_with_source_form(db, tmp_path, content)
    cid = data['case']['id']
    assert [a['hinh_thuc_su_dung'] for a in data['stage']['assets']] == [
        'Sử dụng riêng', 'Sử dụng chung']
    db.expire_all()
    assert [a['hinh_thuc_su_dung'] for a in service.get(cid)['stage']['assets']] == [
        'Sử dụng riêng', 'Sử dụng chung']
    fresh = session_factory()
    try:
        fresh_service = CaseWorkspaceService(fresh, None)
        assert [a['hinh_thuc_su_dung']
                for a in fresh_service.get(cid)['stage']['assets']] == [
                    'Sử dụng riêng', 'Sử dụng chung']
        from services.word_engine import build_template_mapping
        fresh_rows = fresh.query(CaseAssetRow).filter_by(
            case_id=cid).order_by(CaseAssetRow.vitri).all()
        assert [row.hinh_thuc_su_dung for row in fresh_rows] == [
            'Sử dụng riêng', 'Sử dụng chung']
        fresh_text = build_template_mapping(
            fresh.get(InheritanceCase, cid), static=True)['[Đoạn mô tả di sản]']
        assert 'Hình thức sử dụng: Sử dụng riêng' in fresh_text
        assert 'Hình thức sử dụng: Sử dụng chung' in fresh_text
    finally:
        fresh.close()
    rows = db.query(CaseAssetRow).filter_by(case_id=cid).order_by(
        CaseAssetRow.vitri).all()
    assert [row.hinh_thuc_su_dung for row in rows] == [
        'Sử dụng riêng', 'Sử dụng chung']
    from services.word_engine import build_template_mapping
    text = build_template_mapping(
        db.get(InheritanceCase, cid), static=True)['[Đoạn mô tả di sản]']
    assert 'Hình thức sử dụng: Sử dụng riêng' in text
    assert 'Hình thức sử dụng: Sử dụng chung' in text
    assert 'không lấy phần này' not in text


def test_form_label_scoped_to_asset_and_absent_stays_absent(db, tmp_path, session_factory):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Hình thức sử dụng: Sử dụng ngoài tài sản',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Thửa đất số: 22; Tờ bản đồ số: 7',
        'Hình thức sử dụng: Sử dụng chung',
    ])
    assets = [s for s in extraction(content) if s['target'] == 'asset']
    assert [a['fields'].get('hinh_thuc_su_dung', {}).get('normalized_value') for a in assets] == [
        None, 'Sử dụng chung']
    assert 'hinh_thuc_su_dung' not in assets[0]['fields']
    assert assets[1]['fields']['hinh_thuc_su_dung']['raw_value'] == 'Sử dụng chung'

    _, _, service, data = reviewed_case_with_source_form(db, tmp_path, content)
    assert [a['hinh_thuc_su_dung'] for a in data['stage']['assets']] == [None, 'Sử dụng chung']
    fresh = session_factory()
    try:
        assert [a['hinh_thuc_su_dung']
                for a in CaseWorkspaceService(fresh, None).get(data['case']['id'])['stage']['assets']] == [
                    None, 'Sử dụng chung']
        from services.word_engine import build_template_mapping
        fresh_text = build_template_mapping(
            fresh.get(InheritanceCase, data['case']['id']),
            static=True)['[Đoạn mô tả di sản]']
        assert 'Hình thức sử dụng: Sử dụng chung' in fresh_text
    finally:
        fresh.close()


def test_form_repeated_labels_duplicate_vs_conflict(db, tmp_path, session_factory):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Hình thức sử dụng: Sử dụng riêng',
        'Hình thức sử dụng: Sử dụng riêng',
        'Thửa đất số: 22; Tờ bản đồ số: 7',
        'Hình thức sử dụng: Sử dụng riêng',
        'Hình thức sử dụng: Sử dụng chung',
    ])
    assets = [s for s in extraction(content) if s['target'] == 'asset']
    same, conflict = assets
    assert same['fields']['hinh_thuc_su_dung']['normalized_value'] == 'Sử dụng riêng'
    assert len(same['fields']['hinh_thuc_su_dung']['source_refs']) == 2
    assert not any(w['code'] == 'intake.ambiguous' for w in same['warnings'])
    ambiguous = conflict['fields']['hinh_thuc_su_dung']
    assert ambiguous['normalized_value'] is None
    assert [r['raw_snippet'] for r in ambiguous['source_refs']] == [
        'Sử dụng riêng', 'Sử dụng chung']
    assert any(w['code'] == 'intake.ambiguous' and 'hinh_thuc_su_dung' in w['message']
               for w in conflict['warnings'])

    _, _, service, data = reviewed_case_with_source_form(db, tmp_path, content)
    assert [a['hinh_thuc_su_dung'] for a in data['stage']['assets']] == [
        'Sử dụng riêng', None]
    cid = data['case']['id']
    rows = db.query(CaseAssetRow).filter_by(case_id=cid).order_by(
        CaseAssetRow.vitri).all()
    assert [row.hinh_thuc_su_dung for row in rows] == [
        'Sử dụng riêng', None]
    fresh = session_factory()
    try:
        assert [a['hinh_thuc_su_dung']
                for a in CaseWorkspaceService(fresh, None).get(cid)['stage']['assets']] == [
                    'Sử dụng riêng', None]
        from services.word_engine import build_template_mapping
        fresh_text = build_template_mapping(
            fresh.get(InheritanceCase, cid), static=True)['[Đoạn mô tả di sản]']
        assert 'Hình thức sử dụng: Sử dụng riêng' in fresh_text
        assert 'Hình thức sử dụng: Sử dụng chung' not in fresh_text
    finally:
        fresh.close()


def test_form_header_only_or_unclear_cell_not_guessed(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12 | Tờ bản đồ số: 5 | Hình thức sử dụng: Sử dụng riêng | Ghi chú: không lấy',
        'Thửa đất số: 22 | Tờ bản đồ số: 7 | Hình thức sử dụng | Ghi chú: ô trống',
    ])
    assets = [s for s in extraction(content) if s['target'] == 'asset']
    first = next(a for a in assets
                 if a['fields']['so_thua_dat']['normalized_value'] == '12')
    assert first['fields']['hinh_thuc_su_dung']['normalized_value'] == 'Sử dụng riêng'
    assert first['fields']['hinh_thuc_su_dung']['raw_value'] == 'Sử dụng riêng'
    assert all(a['fields'].get('hinh_thuc_su_dung', {}).get('normalized_value')
               != 'Ghi chú' for a in assets)
    _, _, service, data = reviewed_case_with_source_form(db, tmp_path, content)
    assert data['stage']['assets'][0]['hinh_thuc_su_dung'] == 'Sử dụng riêng'


def test_form_nfd_bullet_and_comma_labels_preserve_source(db, tmp_path):
    nfd_label = unicodedata.normalize('NFD', 'Hình thức sử dụng')
    nfd_value = unicodedata.normalize('NFD', 'Sử dụng riêng')
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5', f'{nfd_label}: {nfd_value}',
        'Thửa đất số: 22; Tờ bản đồ số: 7', '- Hình thức sử dụng: Sử dụng chung',
        'Thửa đất số: 32, Tờ bản đồ số: 9, Hình thức sử dụng: Sử dụng riêng',
    ])
    assets = {a['fields']['so_thua_dat']['normalized_value']: a
              for a in extraction(content) if a['target'] == 'asset'}
    nfd_field = assets['12']['fields']['hinh_thuc_su_dung']
    assert unicodedata.normalize('NFC', nfd_field['normalized_value']) == 'Sử dụng riêng'
    assert nfd_field['raw_value'] == nfd_value
    assert [r['raw_snippet'] for r in nfd_field['source_refs']] == [nfd_value]
    assert assets['22']['fields']['hinh_thuc_su_dung']['normalized_value'] == 'Sử dụng chung'
    assert assets['32']['fields']['hinh_thuc_su_dung']['normalized_value'] == 'Sử dụng riêng'
    _, _, service, data = reviewed_case_with_source_form(db, tmp_path, content)
    assert [a['hinh_thuc_su_dung'] for a in data['stage']['assets']] == [
        'Sử dụng riêng', 'Sử dụng chung', 'Sử dụng riêng']


def test_form_header_newline_and_prose_tail_not_stolen(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Hình thức sử dụng:', 'Nội dung ngoài ô không được lấy',
        'Thửa đất số: 22; Tờ bản đồ số: 7',
        'Hình thức sử dụng: Sử dụng riêng theo Giấy chứng nhận AB',
    ])
    assets = {a['fields']['so_thua_dat']['normalized_value']: a
              for a in extraction(content) if a['target'] == 'asset'}
    assert 'hinh_thuc_su_dung' not in assets['12']['fields']
    assert assets['22']['fields']['hinh_thuc_su_dung']['normalized_value'] == (
        'Sử dụng riêng theo Giấy chứng nhận AB')


def test_form_semicolon_compound_value_is_one_explicit_value(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Hình thức sử dụng: Sử dụng riêng; Sử dụng chung',
    ])
    asset, = [a for a in extraction(content) if a['target'] == 'asset']
    field = asset['fields']['hinh_thuc_su_dung']
    assert field['normalized_value'] == 'Sử dụng riêng; Sử dụng chung'
    assert len(field['source_refs']) == 1
    assert not any(w['code'] == 'intake.ambiguous' for w in asset['warnings'])


def test_form_same_line_repeated_labels_are_ambiguous(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Hình thức sử dụng: Sử dụng riêng; Hình thức sử dụng: Sử dụng chung',
    ])
    asset, = [a for a in extraction(content) if a['target'] == 'asset']
    field = asset['fields']['hinh_thuc_su_dung']
    assert field['normalized_value'] is None
    assert [r['raw_snippet'] for r in field['source_refs']] == [
        'Sử dụng riêng', 'Sử dụng chung']
    assert any(w['code'] == 'intake.ambiguous' and 'hinh_thuc_su_dung' in w['message']
               for w in asset['warnings'])


def test_form_context_asset_does_not_borrow_pre_asset_label(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Hình thức sử dụng: Sử dụng ngoài tài sản',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:',
        'Hình thức sử dụng: Sử dụng riêng',
        'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Điều 2: Quyền và nghĩa vụ',
        'Hình thức sử dụng: Sử dụng ngoài tài sản',
    ])
    assets = [a for a in extraction(content) if a['target'] == 'asset']
    assert [a['fields']['hinh_thuc_su_dung']['normalized_value'] for a in assets] == [
        'Sử dụng riêng']
    assert all(a['fields']['hinh_thuc_su_dung']['normalized_value'] != 'Sử dụng ngoài tài sản'
               for a in assets)


def test_form_last_parcel_and_group_do_not_read_past_asset_block(db, tmp_path):
    parcel_doc = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Điều 2: Quyền và nghĩa vụ',
        'Hình thức sử dụng: Ngoài tài sản',
    ])
    parcel_asset, = [a for a in extraction(parcel_doc) if a['target'] == 'asset']
    assert 'hinh_thuc_su_dung' not in parcel_asset['fields']

    group_doc = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản 1:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Điều 2: Quyền và nghĩa vụ',
        'Hình thức sử dụng: Ngoài tài sản',
    ])
    group_asset, = [a for a in extraction(group_doc) if a['target'] == 'asset']
    assert 'hinh_thuc_su_dung' not in group_asset['fields']

    no_thua_doc = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Mẫu, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Mẫu, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Tờ bản đồ số: 5',
        'Điều 2: Quyền và nghĩa vụ',
        'Hình thức sử dụng: Ngoài tài sản',
    ])
    no_thua_asset, = [a for a in extraction(no_thua_doc) if a['target'] == 'asset']
    assert 'hinh_thuc_su_dung' not in no_thua_asset['fields']


def test_intentionally_empty_land_roundtrip_exports_remaining_asset_fields_without_rebinding(db, tmp_path):
    original, store, service, data = reviewed_case(db, tmp_path)
    cid = data['case']['id']; source_id = data['docx_import']['sources'][0]['source_id']
    data['stage']['assets'][0]['land_rows'] = []
    data['stage']['assets'][0]['so_serial'] = 'AA654321'
    data['docx_import']['sources'][0]['records'][2]['land_rows'] = []
    saved = service.commit_stage(cid, data['case']['revision'], data['stage'], docx_import=data['docx_import'])
    assert service.get(cid)['stage']['assets'][0]['land_rows'] == []
    proposals = binding_proposals(db.get(InheritanceCase,cid),source_id,store)
    first_candidate = saved['docx_import']['sources'][0]['records'][2]['candidate_id']
    assert all(not p['supported'] for p in proposals if p['candidate_id'] == first_candidate and p['field_key'] == 'land_rows')
    serial = next(p for p in proposals if p['raw_snippet'] == 'ZZ123456')
    context = saved['docx_import']; context['sources'][0]['bindings'] = [
        {k:serial[k] for k in ('candidate_id','occurrence_id','field_key','slot')}]
    service.commit_stage(cid,saved['revision'],saved['stage'],docx_import=context)
    result = export_original(db.get(InheritanceCase,cid),source_id,store,tmp_path)
    text = read_docx(Path(result['documents'][0]['output_file']['path']).read_bytes()).text
    assert 'AA654321' in text and 'ZZ123456' not in text
    assert 'Loại đất: ONT; Diện tích: 50; Thời hạn: Lâu dài' in text
    assert 'Loại đất: CLN; Diện tích: 20; Thời hạn: 2050' in text
    assert store.load(source_id)[0] == original


@pytest.mark.parametrize('land', [False, True])
def test_corrected_negative_area_exports_saved_value_over_whole_signed_occurrence(db, tmp_path, land):
    doc = Document(io.BytesIO(fixture_bytes()))
    for p in list(doc.paragraphs):
        if p.text.startswith('Loại đất: ONT; Diện tích: 50'):
            p.text = 'Loại đất: ONT; Diện tích: -50; Thời hạn: Lâu dài' if land else 'Diện tích: -50'
        elif not land and p.text.startswith('Loại đất: CLN;'):
            p._p.getparent().remove(p._p)
    stream = io.BytesIO(); doc.save(stream)
    original, store, service, data = reviewed_case(db, tmp_path, stream.getvalue())
    cid = data['case']['id']; source_id = data['docx_import']['sources'][0]['source_id']
    data['stage']['assets'][0]['land_rows'] = [{'loai_dat':'ONT','dien_tich':1234.5,'thoi_han':'Lâu dài'}]
    if land:
        data['docx_import']['sources'][0]['records'][2]['land_rows'] = [1]
    saved = service.commit_stage(cid, data['case']['revision'], data['stage'], docx_import=data['docx_import'])
    proposals = binding_proposals(db.get(InheritanceCase,cid), source_id, store)
    selected = next(p for p in proposals if p['raw_snippet'] == '-50')
    assert selected['supported']
    context = saved['docx_import']; context['sources'][0]['bindings'] = [
        {k:selected[k] for k in ('candidate_id','occurrence_id','field_key','slot')}]
    service.commit_stage(cid,saved['revision'],saved['stage'],docx_import=context)
    exported = export_original(db.get(InheritanceCase,cid),source_id,store,tmp_path)
    path = Path(exported['documents'][0]['output_file']['path'])
    text = read_docx(path.read_bytes()).text
    assert '-50' not in text and 'Diện tích: 1234.5' in text
    assert store.load(source_id)[0] == original


def test_end_to_end_review_commit_reload_original_preservation_and_edit(db, tmp_path):
    original, store, service, data = reviewed_case(db, tmp_path)
    cid = data['case']['id']; revision = data['case']['revision']
    source_id = data['docx_import']['sources'][0]['source_id']
    data['stage']['people'][0]['ho_ten'] = 'Tên đã kiểm tra'
    service.commit_stage(cid, revision, data['stage'])
    db.expire_all(); reloaded = service.get(cid)
    assert reloaded['docx_import']['sources'][0]['records'] == data['docx_import']['sources'][0]['records']
    case = db.get(InheritanceCase, cid)
    proposals = binding_proposals(case, source_id, store)
    bindings = [{key: p[key] for key in ('candidate_id', 'occurrence_id', 'field_key', 'slot')}
                for p in proposals if p['supported']]
    context = reloaded['docx_import']; context['sources'][0]['bindings'] = bindings
    committed = service.commit_stage(cid, reloaded['case']['revision'], reloaded['stage'], docx_import=context)
    with pytest.raises(WorkspaceError, match='Revision'):
        service.commit_stage(cid, reloaded['case']['revision'], reloaded['stage'], docx_import=context)
    db.expire_all(); case = db.get(InheritanceCase, cid)
    rendered, template_hash = reviewed_template(case, source_id, store)
    assert len(template_hash) == 64
    text = read_docx(rendered).text
    assert 'Tên đã kiểm tra' in text
    assert 'Lời chứng: Nguyễn Văn Mẫu — tên nhắc lại giữ nguyên' in text
    assert 'Điều 2: Nội dung giữ nguyên 😀' in text
    with ZipFile(io.BytesIO(original)) as a, ZipFile(io.BytesIO(rendered)) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != 'word/document.xml':
                assert a.read(name) == b.read(name)
        assert any(name.startswith('word/media/') for name in a.namelist())
    assert hashlib.sha256(original).hexdigest() == context['sources'][0]['sha256']
    result = export_original(case, source_id, store, tmp_path)
    assert result['documents'][0]['status'] == 'saved'
    reordered = committed['stage']; reordered['people'].reverse()
    service.commit_stage(cid, committed['revision'], reordered)
    db.expire_all()
    with pytest.raises(ValueError, match='Thứ tự'):
        reviewed_template(db.get(InheritanceCase, cid), source_id, store)
    path = store.root / source_id / 'original.docx'; path.write_bytes(b'changed')
    with pytest.raises(ValueError, match='đã đổi'):
        reviewed_template(db.get(InheritanceCase, cid), source_id, store)


def test_context_new_session_invalid_binding_rollback_and_row_removal(db, session_factory, tmp_path):
    _, store, _, data = reviewed_case(db, tmp_path)
    cid = data['case']['id']; source_id = data['docx_import']['sources'][0]['source_id']
    with session_factory() as fresh:
        service = CaseWorkspaceService(fresh, DocxStore(store.root))
        loaded = service.get(cid)
        assert loaded['docx_import'] == data['docx_import']
        bad = json.loads(json.dumps(loaded['docx_import']))
        bad['sources'][0]['bindings'] = [{'candidate_id': bad['sources'][0]['records'][0]['candidate_id'],
            'occurrence_id': str(uuid.uuid4()), 'field_key': 'ho_ten', 'slot': '[Người 1 - Họ tên]'}]
        loaded['stage']['people'][0]['ho_ten'] = 'Không được lưu'
        with pytest.raises(WorkspaceError, match='Vị trí'):
            service.commit_stage(cid, loaded['case']['revision'], loaded['stage'], docx_import=bad)
        unchanged = service.get(cid)
        assert unchanged['stage']['people'][0]['ho_ten'] != 'Không được lưu'
        assert unchanged['case']['revision'] == loaded['case']['revision']
        removed = unchanged['stage']['people'].pop(0)['row_id']
        unchanged['stage']['owner_row_id'] = unchanged['stage']['people'][0]['row_id']
        result = service.commit_stage(cid, unchanged['case']['revision'], unchanged['stage'])
        assert all(r['row_id'] != removed for r in result['docx_import']['sources'][0]['records'])
        with pytest.raises(ValueError, match='không thuộc'):
            binding_proposals(fresh.get(InheritanceCase, cid), str(uuid.uuid4()), store)


def test_actual_context_and_proposals_match_contract_schema(db, tmp_path):
    from jsonschema import Draft7Validator, FormatChecker
    from referencing import Registry, Resource
    _, store, service, data = reviewed_case(db, tmp_path)
    root = Path(__file__).resolve().parents[2] / 'contracts/notary-case-drafting'
    schemas = {p.name: json.loads(p.read_text(encoding='utf-8')) for p in root.glob('*.schema.json')}
    registry = Registry().with_resources(
        [((root / name).as_uri(), Resource.from_contents(schema)) for name, schema in schemas.items()]
        + [(schema['$id'], Resource.from_contents(schema)) for schema in schemas.values()])
    def validate(name, definition, value):
        Draft7Validator({'$ref': (root / name).as_uri()+'#/definitions/'+definition}, registry=registry,
                        format_checker=FormatChecker()).validate(value)
    validate('common.schema.json', 'docx_import', data['docx_import'])
    cid = data['case']['id']; case = db.get(InheritanceCase, cid)
    proposals = binding_proposals(case, data['docx_import']['sources'][0]['source_id'], store)
    validate('draft-v2.schema.json', 'word_options_result_v2', {'schema_version': data['schema_version'],
        'documents': [], 'binding_proposals': proposals})
    context = data['docx_import']; context['sources'][0]['bindings'] = [
        {k:p[k] for k in ('candidate_id','occurrence_id','field_key','slot')} for p in proposals if p['supported']]
    committed = service.commit_stage(cid, data['case']['revision'], data['stage'], docx_import=context)
    validate('common.schema.json', 'docx_import', committed['docx_import'])
    validate('draft-v2.schema.json', 'payload_commit_stage_v2', {'case_id': cid,
        'base_revision': committed['revision'], 'stage': committed['stage'], 'docx_import': committed['docx_import']})
    # Removed source land row produces unsupported proposals without a slot.
    asset = committed['stage']['assets'][0]
    asset['land_rows'].pop(0)
    record = next(r for r in committed['docx_import']['sources'][0]['records'] if r['row_id'] == asset['row_id'])
    record['land_rows'] = [2]
    committed['docx_import']['sources'][0]['bindings'] = [b for b in committed['docx_import']['sources'][0]['bindings']
        if b['candidate_id'] != record['candidate_id'] or b['field_key'] != 'land_rows']
    saved = service.commit_stage(cid, committed['revision'], committed['stage'], docx_import=committed['docx_import'])
    proposals = binding_proposals(db.get(InheritanceCase, cid), saved['docx_import']['sources'][0]['source_id'], store)
    assert any(not p['supported'] and p['slot'] == '' for p in proposals)
    validate('draft-v2.schema.json', 'word_options_result_v2', {'schema_version': data['schema_version'],
        'documents': [], 'binding_proposals': proposals})
    supported = next(p for p in proposals if p['supported'])
    from jsonschema import ValidationError
    with pytest.raises(ValidationError):
        validate('draft-v2.schema.json', 'word_options_result_v2', {'schema_version': data['schema_version'],
            'documents': [], 'binding_proposals': [{**supported, 'slot': ''}]})
    bad = json.loads(json.dumps(saved['docx_import']))
    bad['sources'][0]['bindings'] = [{k: supported[k] for k in ('candidate_id', 'occurrence_id', 'field_key', 'slot')}]
    bad['sources'][0]['bindings'][0]['slot'] = ''
    with pytest.raises(ValidationError):
        validate('common.schema.json', 'docx_import', bad)


def test_localized_area_complete_span_provenance_and_export(db, tmp_path):
    doc = Document(io.BytesIO(fixture_bytes()))
    next(p for p in doc.paragraphs if p.text.startswith('Loại đất:')).text = 'Loại đất: ONT; Diện tích: 1.234,5; Thời hạn: Lâu dài'
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue()
    _, store, service, data = reviewed_case(db, tmp_path, content)
    source = data['docx_import']['sources'][0]
    suggestions = store.load(source['source_id'])[1]['suggestions']
    asset = next(s for s in suggestions if s['target'] == 'asset')
    rows = json.loads(asset['fields']['land_rows']['normalized_value'])
    assert rows[0]['dien_tich'] == 1234.5
    ref = next(r for r in asset['fields']['land_rows']['source_refs'] if r['land_field'] == 'dien_tich')
    assert ref['raw_snippet'] == '1.234,5'
    assert ref['candidate_state'] == 'matched'
    rendered = replace_spans(content, read_docx(content), [(ref['span'], '2000')])
    assert 'Diện tích: 2000;' in read_docx(rendered).text
    assert '2000,5' not in read_docx(rendered).text
    cid = data['case']['id']
    data['stage']['assets'][0]['land_rows'][0]['dien_tich'] = 2000
    committed = service.commit_stage(cid, data['case']['revision'], data['stage'])
    proposals = binding_proposals(db.get(InheritanceCase, cid), source['source_id'], store)
    p = next(p for p in proposals if p['occurrence_id'] == ref['occurrence_id'])
    committed['docx_import']['sources'][0]['bindings'] = [{k: p[k] for k in ('candidate_id', 'occurrence_id', 'field_key', 'slot')}]
    service.commit_stage(cid, committed['revision'], committed['stage'], docx_import=committed['docx_import'])
    result = export_original(db.get(InheritanceCase, cid), source['source_id'], store, tmp_path)
    output = Path(result['documents'][0]['output_file']['path']).read_bytes()
    assert 'Diện tích: 2000;' in read_docx(output).text
    assert store.load(source['source_id'])[0] == content


@pytest.mark.parametrize('damage', ['tamper', 'missing', 'corrupt'])
def test_extraction_integrity_rejects_tamper_or_missing_checksum(tmp_path, damage):
    store = DocxStore(tmp_path / 'sources'); suggestions = extraction(fixture_bytes(), store)
    ref = suggestions[0]['fields']['ho_ten']['source_refs'][0]
    folder = store.root / ref['source_id']
    before = (folder / 'original.docx').read_bytes()
    if damage == 'tamper':
        package = json.loads((folder / 'extraction.json').read_text(encoding='utf-8'))
        package['suggestions'][0]['fields']['ho_ten']['source_refs'][0]['span'] = [0, 5]
        (folder / 'extraction.json').write_text(json.dumps(package), encoding='utf-8')
    elif damage == 'missing':
        (folder / 'extraction.sha256').unlink()
    else:
        (folder / 'extraction.sha256').write_text('broken', encoding='ascii')
    with pytest.raises(ValueError, match='dấu kiểm'):
        store.load(ref['source_id'], ref['source_sha256'])
    assert (folder / 'original.docx').read_bytes() == before


def test_nested_fields_across_paragraphs_reject_results_and_preserve_package():
    doc = Document()
    def marker(kind):
        p = doc.add_paragraph(); char = OxmlElement('w:fldChar')
        char.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType', kind)
        p.add_run()._r.append(char)
    marker('begin'); marker('separate'); doc.add_paragraph('Outer result')
    marker('begin'); marker('separate'); doc.add_paragraph('Inner result'); marker('end')
    doc.add_paragraph('Remaining outer result'); marker('end'); doc.add_paragraph('Safe text')
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue(); source = read_docx(content)
    for text in ('Outer result', 'Inner result', 'Remaining outer result'):
        lo = source.text.index(text)
        with pytest.raises(ValueError, match='chưa hỗ trợ'):
            replace_spans(content, source, [([lo, lo + len(text)], 'Changed')])
    lo = source.text.index('Safe text')
    result = replace_spans(content, source, [([lo, lo + 9], 'Reviewed')])
    assert all(t in read_docx(result).text for t in ('Outer result', 'Inner result', 'Remaining outer result'))
    with ZipFile(io.BytesIO(content)) as a, ZipFile(io.BytesIO(result)) as b:
        for name in a.namelist():
            if name != 'word/document.xml': assert a.read(name) == b.read(name)


def test_source_row_21_survives_saved_stage_and_schema(db, tmp_path):
    doc = Document(); doc.add_paragraph('Tài sản 1:')
    doc.add_paragraph('Số serial: ZZ123456; có địa chỉ tại: Nhà đất A')
    for n in range(1, 22): doc.add_paragraph(f'Loại đất: ONT; Diện tích: {n}; Thời hạn: Lâu dài')
    doc.add_paragraph('Điều 2: Giữ nguyên')
    stream = io.BytesIO(); doc.save(stream); store = DocxStore(tmp_path / 'sources')
    candidate = next(s for s in extraction(stream.getvalue(), store) if s['target'] == 'asset')
    rows = json.loads(candidate['fields']['land_rows']['normalized_value'])
    assert len(rows) == 21
    person = _person_row(); asset = _asset_row(land_rows=rows[1:])
    ref = candidate['fields']['so_serial']['source_refs'][0]
    context = {'version': 1, 'sources': [{'source_id': ref['source_id'], 'sha256': ref['source_sha256'],
        'records': [{'candidate_id': candidate['suggestion_id'], 'row_id': asset['row_id'], 'land_rows': list(range(2, 22))}], 'bindings': []}]}
    service = CaseWorkspaceService(db, store)
    data = service.create(str(uuid.uuid4()), {'case_type': 'inheritance', 'document_type': 'khai_nhan'},
        {'owner_row_id': person['row_id'], 'people': [person], 'assets': [asset]}, None, context)
    saved = service.commit_stage(data['case']['id'], data['case']['revision'], data['stage'], docx_import=data['docx_import'])
    db.expire_all(); loaded = service.get(data['case']['id'])
    assert loaded['docx_import']['sources'][0]['records'][0]['land_rows'][-1] == 21
    assert len(loaded['stage']['assets'][0]['land_rows']) == 20
    from jsonschema import Draft7Validator
    schema = json.loads((Path(__file__).resolve().parents[2] / 'contracts/notary-case-drafting/common.schema.json').read_text(encoding='utf-8'))
    Draft7Validator(schema['definitions']['docx_import']).validate(loaded['docx_import'])


def test_mixed_styles_rejected_and_original_runs_preserved(db, tmp_path):
    doc = Document(io.BytesIO(fixture_bytes()))
    doc.paragraphs[2].runs[2].bold = False
    doc.paragraphs[2].runs[2].italic = True
    stream = io.BytesIO(); doc.save(stream); original = stream.getvalue()
    source = read_docx(original); lo = source.text.index('Nguyễn Văn Mẫu')
    with pytest.raises(ValueError, match='nhiều kiểu chữ'):
        replace_spans(original, source, [([lo, lo + len('Nguyễn Văn Mẫu')], 'Tên khác')])
    # An unrelated supported replacement must preserve both original styled runs.
    at = source.text.index('Nhà B')
    result = replace_spans(original, source, [([at, at + 5], 'Nhà C')])
    parsed = Document(io.BytesIO(result))
    assert parsed.paragraphs[2].runs[1].bold is True
    assert parsed.paragraphs[2].runs[2].italic is True
    assert parsed.paragraphs[2].text == doc.paragraphs[2].text


def test_land_deletion_does_not_rebind_earlier_occurrence(db, tmp_path):
    _, store, service, data = reviewed_case(db, tmp_path)
    cid = data['case']['id']; context = data['docx_import']; source = context['sources'][0]
    asset = data['stage']['assets'][0]
    record = next(r for r in source['records'] if r['row_id'] == asset['row_id'])
    assert record['land_rows'] == [1, 2]
    asset['land_rows'].pop(0); asset['land_rows'][0]['dien_tich'] = 77
    record['land_rows'] = [2]
    service.commit_stage(cid, data['case']['revision'], data['stage'], docx_import=context)
    loaded = service.get(cid); case = db.get(InheritanceCase, cid)
    proposals = binding_proposals(case, source['source_id'], store)
    deleted = [p for p in proposals if p['field_key'] == 'land_rows' and p['raw_snippet'] in ('ONT', '50') and p['row_id'] == asset['row_id']]
    assert deleted and all(not p['supported'] for p in deleted)
    surviving = [p for p in proposals if p['field_key'] == 'land_rows' and p['row_id'] == asset['row_id'] and p['supported']]
    assert next(p for p in surviving if p['raw_snippet'] == '20')['slot'] == '[Loại đất 1.1 - Diện tích]'
    loaded['docx_import']['sources'][0]['bindings'] = [{k:p[k] for k in ('candidate_id','occurrence_id','field_key','slot')} for p in surviving]
    service.commit_stage(cid, loaded['case']['revision'], loaded['stage'], docx_import=loaded['docx_import'])
    db.expire_all(); rendered, _ = reviewed_template(db.get(InheritanceCase, cid), source['source_id'], store)
    text = read_docx(rendered).text
    assert 'Loại đất: ONT; Diện tích: 50; Thời hạn: Lâu dài' in text
    assert 'Loại đất: CLN; Diện tích: 77; Thời hạn: 2050' in text
    reloaded = service.get(cid)
    assert next(r for r in reloaded['docx_import']['sources'][0]['records'] if r['row_id'] == asset['row_id'])['land_rows'] == [2]


def test_empty_catalog_wrong_slot_rollback_and_corrupt_source_detach(db, tmp_path):
    _, store, service, data = reviewed_case(db, tmp_path)
    cid = data['case']['id']; source = data['docx_import']['sources'][0]
    data['stage']['people'][0]['dia_chi'] = None
    saved = service.commit_stage(cid, data['case']['revision'], data['stage'])
    proposals = binding_proposals(db.get(InheritanceCase, cid), source['source_id'], store)
    blank = next(p for p in proposals if p['row_id'] == data['stage']['people'][0]['row_id'] and p['field_key'] == 'dia_chi')
    assert not blank['supported'] and 'giá trị' in blank['reason']
    p = next(p for p in proposals if p['supported'] and p['field_key'] == 'ho_ten')
    bad = json.loads(json.dumps(saved['docx_import']))
    bad['sources'][0]['bindings'] = [{k:p[k] for k in ('candidate_id','occurrence_id','field_key','slot')}]
    bad['sources'][0]['bindings'][0]['slot'] = '[Người 2 - Họ tên]'
    saved['stage']['people'][0]['ho_ten'] = 'Không lưu tên này'
    with pytest.raises(WorkspaceError, match='Vị trí Word'):
        service.commit_stage(cid, saved['revision'], saved['stage'], docx_import=bad)
    unchanged = service.get(cid)
    assert unchanged['case']['revision'] == saved['revision']
    assert unchanged['stage']['people'][0]['ho_ten'] != 'Không lưu tên này'
    (store.root / source['source_id'] / 'original.docx').write_bytes(b'corrupt')
    detached = service.commit_stage(cid, unchanged['case']['revision'], unchanged['stage'], docx_import={'version':1,'sources':[]})
    assert detached['docx_import']['sources'] == []
    assert len(detached['stage']['people']) == 2


def test_binding_save_reopen_and_clear(db, tmp_path):
    _, store, service, data = reviewed_case(db, tmp_path)
    cid = data['case']['id']; source = data['docx_import']['sources'][0]
    proposals = binding_proposals(db.get(InheritanceCase, cid), source['source_id'], store)
    p = next(p for p in proposals if p['supported'])
    source['bindings'] = [{k:p[k] for k in ('candidate_id','occurrence_id','field_key','slot')}]
    saved = service.commit_stage(cid, data['case']['revision'], data['stage'], docx_import=data['docx_import'])
    db.expire_all(); case = db.get(InheritanceCase, cid)
    reopened = binding_proposals(case, source['source_id'], store)
    assert [p['occurrence_id'] for p in reopened if p['selected']] == [p['occurrence_id']]
    source = saved['docx_import']['sources'][0]; source['bindings'] = []
    service.commit_stage(cid, saved['revision'], saved['stage'], docx_import=saved['docx_import'])
    loaded = service.get(cid)
    assert loaded['docx_import']['sources'][0]['bindings'] == []
    assert 'template_sha256' not in loaded['docx_import']['sources'][0]
    assert not any(p['selected'] for p in binding_proposals(db.get(InheritanceCase, cid), source['source_id'], store))


def test_eight_sources_persist_ninth_rejects_transaction(db, tmp_path):
    content, store, service, data = reviewed_case(db, tmp_path)
    context = data['docx_import']; cid = data['case']['id']
    for _ in range(8):
        candidates = extraction(content, store)
        ref = candidates[0]['fields']['ho_ten']['source_refs'][0]
        context['sources'].append({'source_id':ref['source_id'],'sha256':ref['source_sha256'],
            'records':[{'candidate_id':candidates[0]['suggestion_id'],'row_id':data['stage']['people'][0]['row_id']}],'bindings':[]})
    ninth = context['sources'].pop()
    saved = service.commit_stage(cid, data['case']['revision'], data['stage'], docx_import=context)
    assert len(service.get(cid)['docx_import']['sources']) == 8
    context['sources'].append(ninth)
    data['stage']['people'][0]['ho_ten'] = 'Ninth must not save'
    with pytest.raises(WorkspaceError, match='8 nguồn'):
        service.commit_stage(cid, saved['revision'], data['stage'], docx_import=context)
    loaded = service.get(cid)
    assert loaded['case']['revision'] == saved['revision']
    assert loaded['stage']['people'][0]['ho_ten'] != 'Ninth must not save'


def test_year_only_source_date_survives_commit_and_reload(db, tmp_path):
    """Nguồn chỉ có năm sinh → gợi ý giữ `1980`; commit rồi mở lại vẫn
    `1980` (không bị đệm thành `01/01/1980`)."""
    doc = Document(io.BytesIO(fixture_bytes()))
    for p in doc.paragraphs:
        for run in p.runs:
            run.text = run.text.replace('01/02/1980', '1980')
    stream = io.BytesIO(); doc.save(stream); content = stream.getvalue()
    _, _, service, data = reviewed_case(db, tmp_path / 'case', content)
    person = data['stage']['people'][0]
    assert person['ngay_sinh'] == '1980'
    saved = service.commit_stage(
        data['case']['id'], data['case']['revision'], data['stage'],
        docx_import=data['docx_import'])
    assert saved['stage']['people'][0]['ngay_sinh'] == '1980'
    assert service.get(data['case']['id'])['stage']['people'][0]['ngay_sinh'] == '1980'


def test_saved_conversion_preserves_old_bindings_requires_reconfirmation_and_exports(db, tmp_path):
    from services.inheritance_workspace import InheritanceWorkspaceService
    from services.word_batch_export import WordBatchError
    original, store, service, data = reviewed_case(db, tmp_path)
    cid = data["case"]["id"]
    source = data["docx_import"]["sources"][0]
    sid = source["source_id"]
    proposals = binding_proposals(db.get(InheritanceCase, cid), sid, store)
    source["bindings"] = [
        {k: p[k] for k in ("candidate_id", "occurrence_id", "field_key", "slot")}
        for p in proposals if p["supported"]]
    assert any(b["slot"].startswith("[Người ") for b in source["bindings"])
    saved = service.commit_stage(cid, data["case"]["revision"], data["stage"],
                                 docx_import=data["docx_import"])
    old_context = json.loads(json.dumps(saved["docx_import"]))
    converted = service.convert_type(cid, saved["revision"], {
        "case_type": "two_party", "document_type": "tang_cho",
        "confirm_diagram_reset": True})
    assert converted["docx_import"] == old_context
    with pytest.raises(WordBatchError) as invalid:
        export_original(db.get(InheritanceCase, cid), sid, store, tmp_path)
    assert invalid.value.code == "word.original_mapping_invalid"
    state = converted["diagram"]["state"]
    for index, person in enumerate(converted["stage"]["people"]):
        state["nodes"][index * 15]["personId"] = person["row_id"]
    diagram = InheritanceWorkspaceService(db).save_diagram(
        cid, converted["case"]["revision"], state)
    loaded = service.get(cid)
    proposals = binding_proposals(db.get(InheritanceCase, cid), sid, store)
    bindings = [{k: p[k] for k in ("candidate_id", "occurrence_id", "field_key", "slot")}
                for p in proposals if p["supported"]]
    assert any(b["slot"].startswith("[Bên ") for b in bindings)
    loaded["docx_import"]["sources"][0]["bindings"] = bindings
    rebound = service.commit_stage(cid, diagram["revision"], loaded["stage"],
                                    docx_import=loaded["docx_import"])
    db.expire_all()
    assert service.get(cid)["docx_import"] == rebound["docx_import"]
    output = export_original(db.get(InheritanceCase, cid), sid, store, tmp_path)
    assert output["documents"][0]["status"] == "saved"
    assert store.load(sid)[0] == original


def test_deceased_person_maps_explicit_source_fields_without_identity_merge(db, tmp_path, session_factory):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Sống, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Sống, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Người để lại di sản: Ông Nguyễn Văn Cố, sinh năm 1940, chết ngày 05/06/2016 '
        'theo Trích lục khai tử số 95/2026/TLKT do UBND xã Mẫu cấp ngày 10/06/2016; '
        'Nơi thường trú cuối cùng trước khi chết: Thôn Mẫu, xã Mẫu.',
    ])
    suggestions = extraction(content)
    deceased = [s for s in suggestions if s['target'] == 'person'
                and s['fields'].get('ho_ten', {}).get('normalized_value') == 'Nguyễn Văn Cố']
    assert len(deceased) == 1
    fields = deceased[0]['fields']
    assert fields['gioi_tinh']['normalized_value'] == 'Nam'
    assert fields['ngay_sinh']['normalized_value'] == '1940'
    assert fields['ngay_chet']['normalized_value'] == '2016-06-05'
    assert fields['loai_giay_to']['normalized_value'] == 'Trích lục khai tử'
    assert fields['ngay_cap']['normalized_value'] == '2016-06-10'
    assert fields['noi_cap']['normalized_value'] == 'UBND xã Mẫu'
    assert fields['dia_chi']['normalized_value'] == 'Thôn Mẫu, xã Mẫu'
    assert fields['so_giay_to']['normalized_value'] == '95/2026/TLKT'
    assert fields['ngay_chet']['source_refs'][0]['raw_snippet'] == '05/06/2016'
    assert fields['loai_giay_to']['source_refs'][0]['raw_snippet'] == 'Trích lục khai tử'

    _, _, service, data = reviewed_case_with_source_form(db, tmp_path, content)
    cid = data['case']['id']
    stage_deceased = next(p for p in data['stage']['people'] if p['ho_ten'] == 'Nguyễn Văn Cố')
    assert stage_deceased['loai_giay_to'] == 'Trích lục khai tử'
    assert (stage_deceased['ngay_cap'], stage_deceased['noi_cap']) == (
        '2016-06-10', 'UBND xã Mẫu')
    assert stage_deceased['dia_chi'] == 'Thôn Mẫu, xã Mẫu'
    assert stage_deceased['so_giay_to'] == '95/2026/TLKT'
    db.expire_all()
    assert any(p['ho_ten'] == 'Nguyễn Văn Cố' and p['ngay_chet'] == '2016-06-05'
               for p in service.get(cid)['stage']['people'])
    from models import Customer
    deceased_customer = db.query(Customer).filter(Customer.ho_ten == 'Nguyễn Văn Cố').one()
    assert deceased_customer.so_giay_to == '95/2026/TLKT'
    # Ngày chết thuộc dòng Hồ sơ; danh tính dùng chung không bị ghi ngược.
    assert deceased_customer.ngay_chet is None
    fresh = session_factory()
    try:
        fresh_people = CaseWorkspaceService(fresh, None).get(cid)['stage']['people']
        fresh_deceased = next(p for p in fresh_people if p['ho_ten'] == 'Nguyễn Văn Cố')
        assert (fresh_deceased['ngay_sinh'], fresh_deceased['ngay_chet']) == ('1940', '2016-06-05')
        assert fresh_deceased['gioi_tinh'] == 'Nam'
        assert fresh_deceased['loai_giay_to'] == 'Trích lục khai tử'
        assert (fresh_deceased['ngay_cap'], fresh_deceased['noi_cap']) == (
            '2016-06-10', 'UBND xã Mẫu')
        assert fresh_deceased['dia_chi'] == 'Thôn Mẫu, xã Mẫu'
    finally:
        fresh.close()
    from services.word_engine import (
        build_template_mapping, build_word_context, replace_in_doc)
    case = db.get(InheritanceCase, cid)
    context = build_word_context(case)
    assert any(p.ho_ten == 'Nguyễn Văn Cố' and p.is_deceased for p in context.persons)
    assert (context.persons[2].ho_ten, context.persons[2].is_deceased) == (
        'Nguyễn Văn Cố', True)
    # Template tokens are fixed here BEFORE the engine mapping/output exists:
    # fixed Stage person slot 3, source-confirmed deceased identity.
    expected_by_token = {
        '[Người 3 - Họ tên]': 'Nguyễn Văn Cố',
        '[Người 3 - Ngày chết]': '05/06/2016',
        '[Người 3 - Năm sinh]': '1940',
        '[Người 3 - Giới tính]': 'Nam',
        '[Người 3 - Loại giấy tờ]': 'Trích lục khai tử',
        '[Người 3 - Ngày cấp]': '10/06/2016',
        '[Người 3 - Nơi cấp]': 'UBND xã Mẫu',
        '[Người 3 - Địa chỉ]': 'Thôn Mẫu, xã Mẫu',
    }
    template_text = 'Tài liệu thử: Người để lại di sản. ' + ' | '.join(expected_by_token)
    mapping = build_template_mapping(case, static=True)
    mismatches = {token: (mapping.get(token), value)
                  for token, value in expected_by_token.items()
                  if mapping.get(token) != value}
    assert not mismatches, mismatches
    template = Document()
    template.add_paragraph(template_text)
    replace_in_doc(template, mapping)
    rendered_path = tmp_path / 'deceased-rendered.docx'
    template.save(str(rendered_path))
    rendered_text = read_docx(rendered_path.read_bytes()).text
    assert not any(token in rendered_text for token in expected_by_token)
    assert 'Tài liệu thử: Người để lại di sản.' in rendered_text
    assert all(value in rendered_text for value in expected_by_token.values())


def test_deceased_namesake_living_not_merged_and_missing_type_stays_partial(db, tmp_path):
    content = _docx_bytes([
        'HỢP ĐỒNG CHUYỂN NHƯỢNG',
        'Bên A:', 'Ông Nguyễn Văn Trùng, sinh ngày 01/02/1980; CCCD: 001234567890',
        'Bên B:', 'Bà Trần Thị Sống, sinh ngày 03/04/1985; CCCD: 001234567891',
        'Tài sản:', 'Số serial: ZZ123456; có địa chỉ tại: Nhà A',
        'Thửa đất số: 12; Tờ bản đồ số: 5',
        'Người để lại di sản: Ông Nguyễn Văn Trùng, đã chết ngày 05/06/2016.',
    ])
    people = [s for s in extraction(content) if s['target'] == 'person']
    same_name = [s for s in people
                 if s['fields'].get('ho_ten', {}).get('normalized_value') == 'Nguyễn Văn Trùng']
    assert len(same_name) == 2
    deceased = next(s for s in same_name if 'ngay_chet' in s['fields'])
    assert deceased['fields']['ngay_chet']['normalized_value'] == '2016-06-05'
    assert 'loai_giay_to' not in deceased['fields']
    assert 'ngay_cap' not in deceased['fields']
    assert 'noi_cap' not in deceased['fields']
    assert 'so_giay_to' not in deceased['fields']

    _, _, service, data = reviewed_case_with_source_form(db, tmp_path, content)
    from models import Customer
    # Dòng người chết thiếu neo giấy tờ → chưa có identity dùng chung;
    # dòng còn sống có neo thì một identity. Không gộp hai dòng cùng tên.
    assert db.query(Customer).filter(Customer.ho_ten == 'Nguyễn Văn Trùng').count() == 1
    rows = [p for p in service.get(data['case']['id'])['stage']['people']
            if p['ho_ten'] == 'Nguyễn Văn Trùng']
    assert len(rows) == 2
    assert [row['ngay_chet'] for row in rows] == [None, '2016-06-05']
    assert [row['entity_id'] for row in rows][0] is not None
    assert rows[1]['entity_id'] is None
