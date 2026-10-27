"""E2 — Nhân bản toàn bộ Hồ sơ end-to-end (issue 13 vòng 2, lát 5).

Đáp án đọc từ `fixtures/issue13_storage_b.json` (prototype) + literal của
test; không lấy output `duplicate` làm oracle. Phạm vi:

- F1 đầy đủ: meta cột + 7 scalar, đất hàng trống/null/[], participants
  thủ công + parent, Diagram, docx provenance; map FK mới toàn bộ; giữ
  identity; draft/rev1.
- Sửa bản sao ↔ nguồn độc lập; Word đọc đúng hai bộ.
- Duplicate nháp thiếu dữ liệu (F3 shape) vẫn đủ nội dung.
- Replay cùng key: không ghi thêm/không đổi revision, trả bản sao hiện
  tại (đã sửa/done, kể cả nguồn đã bị xóa); key khác source/command →
  conflict; hai Session cùng key → đúng một bản sao.
- Fault sau land/participant/binding → rollback toàn bộ, nguồn nguyên vẹn.
- Nguồn không tồn tại → `case_not_found {source_case_id}`.
"""
import copy
import json
import threading
import uuid
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import database
from models import (
    CaseAssetLandRow,
    CaseAssetRow,
    CasePersonRow,
    Customer,
    InheritanceCase,
    InheritanceCaseProperty,
    InheritanceParticipant,
    Property,
)
from services import case_rows
from services import word_batch_export as wbe
from services.case_workspace import CaseWorkspaceService, WorkspaceError
from services.document_intake.docx_store import DocxStore

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "issue13_storage_b.json"

SCALARS = {
    "tenvanban": "Văn bản thử nhân bản",
    "nhomhopdong": "Nhóm giả lập",
    "loaivanbannhandien": "Khai nhận di sản",
    "socongchung": "12/2026",
    "ngaycongchung": "2026-10-20",
    "congchungvien": "Công chứng viên Thử",
    "thuky": "Thư ký Thử",
}


# ---------------------------------------------------------------- fixtures


@pytest.fixture(scope="module")
def dataset():
    with FIXTURE_PATH.open(encoding="utf-8") as handle:
        return json.load(handle)


@pytest.fixture()
def session_factory(tmp_path):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'workspace_duplicate.db'}",
        connect_args={"timeout": 15})
    database.enable_sqlite_foreign_keys(engine)
    database.Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    yield factory
    engine.dispose()


@pytest.fixture()
def db(session_factory):
    session = session_factory()
    yield session
    session.close()


@pytest.fixture()
def store(tmp_path):
    return DocxStore(tmp_path / "docx_sources")


# ------------------------------------------------------------------ helpers


def _case_spec(dataset, alias):
    matches = [case for case in dataset["cases"] if case["alias"] == alias]
    assert len(matches) == 1
    return matches[0]


def _proto(dataset, section, alias, override=None):
    data = dict(dataset[section][alias])
    data.pop("alias", None)
    data.update(override or {})
    return data


def _service(db, store=None):
    return CaseWorkspaceService(db, lambda: store) if store else \
        CaseWorkspaceService(db)


def _uuid4():
    return str(uuid.uuid4())


def _reopen(factory, case_id):
    session = factory()
    try:
        return CaseWorkspaceService(session).get(case_id), session
    except Exception:
        session.close()
        raise


def _land_wire(asset):
    land = asset.get("land_rows")
    if land is None:
        return None
    return [{key: row.get(key) for key in ("loai_dat", "dien_tich",
                                           "thoi_han")}
            for row in land]


def _row_free(stage):
    """Stage bỏ row_id để so nội dung hai bản."""
    out = {key: copy.deepcopy(value) for key, value in stage.items()
           if key not in ("people", "assets")}
    out["people"] = [{k: copy.deepcopy(v) for k, v in row.items()
                      if k != "row_id"} for row in stage.get("people") or []]
    out["assets"] = [{k: copy.deepcopy(v) for k, v in row.items()
                      if k != "row_id"} for row in stage.get("assets") or []]
    return out


def _stage_content(stage):
    """Nội dung Stage bỏ row_id + owner pointer (resolver cấp UUID mới)."""
    content = _row_free(stage)
    content.pop("owner_row_id", None)
    return content


def _owner_name(stage):
    owner = stage.get("owner_row_id")
    if owner is None:
        return None
    return next(row["ho_ten"] for row in stage["people"]
                if row["row_id"] == owner)


def _case_counts(db):
    return {
        "cases": db.query(InheritanceCase).count(),
        "people": db.query(CasePersonRow).count(),
        "assets": db.query(CaseAssetRow).count(),
        "lands": db.query(CaseAssetLandRow).count(),
        "participants": db.query(InheritanceParticipant).count(),
        "links": db.query(InheritanceCaseProperty).count(),
        "customers": db.query(Customer).count(),
        "properties": db.query(Property).count(),
    }


def _source_snapshot(factory, case_id):
    session = factory()
    try:
        case = session.get(InheritanceCase, case_id)
        people = sorted(case.person_rows, key=lambda r: r.vitri)
        assets = sorted(case.asset_rows, key=lambda r: r.vitri)
        return {
            "case_state_json": case.case_state_json,
            "revision": case.workspace_revision,
            "status": case.trang_thai,
            "updated_at": case.updated_at,
            "people": [(r.row_id, r.customer_id, r.vitri, r.ho_ten)
                       for r in people],
            "assets": [(r.row_id, r.property_id, r.vitri, r.so_serial,
                        r.land_rows_present) for r in assets],
            "lands": sorted(
                (l.id, l.case_asset_row_id, l.vitri, l.loaidat,
                 l.dientich, l.thoihan)
                for r in assets for l in r.land_rows),
            "participants": sorted(
                (p.id, p.case_person_row_id, p.customer_id,
                 p.parent_case_person_row_id, p.parent_customer_id,
                 p.vai_tro, p.hang_thua_ke, p.co_nhan_tai_san,
                 p.ty_le, p.ghi_chu)
                for p in case.participants),
        }
    finally:
        session.close()


# ------------------------------------------------- synthetic Word provenance


def _build_docx_source(store, person_row_id, asset_row_id):
    """Nguồn Word tổng hợp + suggestion/record/binding cho một Người + một
    Tài sản. Trả (context, source_id, candidate ids, binding)."""
    document = Document()
    document.add_paragraph("NGƯỜI: Nguyễn Đăng An")
    document.add_paragraph("SERIAL: SYN B 001")
    path = Path(store.root) / ".seed.docx"
    Path(store.root).mkdir(parents=True, exist_ok=True)
    document.save(path)
    content = path.read_bytes()
    path.unlink()

    from services.document_intake.docx_source import read_docx
    text = read_docx(content).text
    name_span = [text.index("Nguyễn Đăng An"),
                 text.index("Nguyễn Đăng An") + len("Nguyễn Đăng An")]
    serial_span = [text.index("SYN B 001"),
                   text.index("SYN B 001") + len("SYN B 001")]
    person_candidate = _uuid4()
    asset_candidate = _uuid4()
    person_occurrence = _uuid4()
    asset_occurrence = _uuid4()

    def suggestion(sid, target, field, value, span, occurrence):
        return {
            "suggestion_id": sid,
            "source_id": "",
            "target": target,
            "fields": {field: {
                "raw_value": value,
                "normalized_value": value,
                "observation_state": "normalized",
                "confidence": None,
                "source_refs": [{
                    "occurrence_id": occurrence,
                    "span": span,
                    "source_id": "",
                    "source_label": "",
                }],
            }},
            "warnings": [],
        }

    suggestions = [
        suggestion(person_candidate, "person", "ho_ten",
                   "Nguyễn Đăng An", name_span, person_occurrence),
        suggestion(asset_candidate, "asset", "so_serial",
                   "SYN B 001", serial_span, asset_occurrence),
    ]
    source_id = store.save(content, suggestions, "Nguồn tổng hợp E2.docx")
    import hashlib
    context = {
        "version": 1,
        "sources": [{
            "source_id": source_id,
            "sha256": hashlib.sha256(content).hexdigest(),
            "records": [
                {"candidate_id": person_candidate,
                 "row_id": person_row_id},
                {"candidate_id": asset_candidate,
                 "row_id": asset_row_id},
            ],
            "bindings": [
                {"candidate_id": person_candidate,
                 "occurrence_id": person_occurrence,
                 "field_key": "ho_ten",
                 "slot": "[Người 1 - Họ tên]"},
                {"candidate_id": asset_candidate,
                 "occurrence_id": asset_occurrence,
                 "field_key": "so_serial",
                 "slot": "[Tài sản 1 - Serial]"},
            ],
        }],
    }
    return context, source_id, person_candidate, asset_candidate


# ------------------------------------------------------------ F1 source


def _f1_stage(dataset):
    """F1 + hai dòng tài sản bổ sung để phủ land null (T203) và [] (T204)."""
    people = [_proto(dataset, "people_prototypes", alias)
              for alias in ("P101", "P102", "P103")]
    assets = [_proto(dataset, "asset_prototypes", alias)
              for alias in ("T101", "T204", "T203")]
    stage = {"owner_row_id": people[0]["row_id"],
             "people": people, "assets": assets}
    return stage


def _f1_state(stage):
    owner, spouse, child = stage["people"]
    return {"version": 3, "domain": "inheritance", "nodes": [
        {"id": "owner", "personId": owner["row_id"], "parentSlotIds": [],
         "spouseSlotIds": ["spouse"], "spouseSlotId": "spouse",
         "spouseSuggestionId": None, "ownPositions": [1],
         "receivePositions": [], "hidden": False, "deleted": False},
        {"id": "spouse", "personId": spouse["row_id"],
         "parentSlotIds": [], "spouseSlotIds": ["owner"],
         "spouseSlotId": "owner", "spouseSuggestionId": None,
         "ownPositions": [1], "receivePositions": [1],
         "hidden": False, "deleted": False},
        {"id": "child", "personId": child["row_id"],
         "parentSlotIds": ["owner", "spouse"], "spouseSlotIds": [],
         "spouseSlotId": None, "spouseSuggestionId": None,
         "ownPositions": [], "receivePositions": [1],
         "hidden": False, "deleted": False},
    ]}


def _save_f1_source(db, dataset, store):
    stage = _f1_stage(dataset)
    context, source_id, _pc, _ac = _build_docx_source(
        store, stage["people"][0]["row_id"], stage["assets"][0]["row_id"])
    meta = {
        "case_type": "inheritance",
        "document_type": "khai_nhan",
        "ngay_lap_ho_so": "2026-10-22",
        "noi_niem_yet": "Xã Giả Lập",
        "nguoi_nhan_uy_quyen": "Người Nhận Ủy Quyền Thử",
        "nguoi_nhan_uy_quyen_id": None,
        "noi_dung_viec": "Nội dung việc thử",
        "ghi_chu": "Ghi chú nguồn",
        **SCALARS,
    }
    saved = _service(db, store).save_case(
        case_id=None, base_revision=None, idempotency_key=_uuid4(),
        case_meta=meta, stage=stage, state=_f1_state(stage),
        docx_import=context)

    people = {row["ho_ten"]: row for row in saved["stage"]["people"]}
    owner_row = people["Nguyễn Đăng An"]["row_id"]
    child_row = people["Nguyễn Văn Con"]["row_id"]
    spouse_row = people["Trần Thị Bình"]["row_id"]
    source_case = db.get(InheritanceCase, saved["case"]["id"])
    rows = {r.row_id: r for r in source_case.person_rows}
    # Participant thủ công (không thuộc tập derived) + parent pointer.
    db.add(InheritanceParticipant(
        ho_so_id=source_case.id,
        case_person_row_id=owner_row,
        customer_id=rows[owner_row].customer_id,
        vai_tro="Khác",
        hang_thua_ke=3,
        co_nhan_tai_san=False,
        ty_le=12.5,
        ghi_chu="Participant thủ công",
        parent_case_person_row_id=child_row,
        parent_customer_id=rows[child_row].customer_id))
    # Participant legacy: pointer null, identity nguồn giữ nguyên.
    db.add(InheritanceParticipant(
        ho_so_id=source_case.id,
        case_person_row_id=None,
        customer_id=rows[spouse_row].customer_id,
        vai_tro="Khác",
        hang_thua_ke=2,
        co_nhan_tai_san=True,
        ty_le=0.0,
        ghi_chu="Legacy không pointer",
        parent_case_person_row_id=None,
        parent_customer_id=None))
    db.commit()
    return saved, source_id


def _duplicate(db, store, source_case_id, key=None):
    return _service(db, store).duplicate(
        source_case_id=source_case_id,
        idempotency_key=key or _uuid4())


# ------------------------------------------------------- export helper


def _export_text(factory, case_id, tmp_path, *, key):
    template_dir = tmp_path / f"templates_{key}"
    dest = tmp_path / f"out_{key}"
    template_dir.mkdir(parents=True)
    dest.mkdir(parents=True)
    template = template_dir / f"{key}.docx"
    document = Document()
    for token in ("[Tên 1]", "[Tên 2]", "[Tên 3]", "[Địa chỉ 1]",
                  "[Địa chỉ 2]", "[Địa chỉ 3]", "[Serial]",
                  "[Diện tích 1]"):
        document.add_paragraph(token)
    document.save(template)
    session = factory()
    try:
        case = session.get(InheritanceCase, case_id)
        spec = wbe.WordDocumentSpec(document_key=key, display_name=key,
                                    filename_stem=key)
        batch = wbe.export_batch(
            case, case_id=case_id, document_keys=[key], dest_dir=dest,
            resolve_template=lambda _k: template, doc_specs=[spec])
        entry = batch["documents"][0]
        assert entry["status"] == "saved", entry
        output = Document(entry["output_file"]["path"])
        return "\n".join(p.text for p in output.paragraphs)
    finally:
        session.close()


# ------------------------------------------------------------- F1 proof


def test_duplicate_f1_copies_everything_and_remaps_foreign_keys(
        dataset, db, session_factory, store):
    saved, docx_source_id = _save_f1_source(db, dataset, store)
    source_id = saved["case"]["id"]
    source = _reopen(session_factory, source_id)[0]
    before_snapshot = _source_snapshot(session_factory, source_id)
    before_counts = _case_counts(db)

    saved_templates = []
    original_save_template = store.save_template

    def spy_save_template(*args, **kwargs):
        saved_templates.append(args[0])
        return original_save_template(*args, **kwargs)

    store.save_template = spy_save_template
    try:
        key = _uuid4()
        copied = _duplicate(db, store, source_id, key)
    finally:
        store.save_template = original_save_template

    assert copied["created"] is True
    assert copied["source_case_id"] == source_id
    assert saved_templates == [], "duplicate không được save_template"
    copy_id = copied["case"]["id"]
    assert copy_id != source_id
    assert copied["case"]["status"] == "draft"
    assert copied["case"]["revision"] == 1

    # Meta cột + 7 scalar.
    for field in ("case_type", "document_type", "ngay_lap_ho_so",
                  "noi_niem_yet", "nguoi_nhan_uy_quyen",
                  "nguoi_nhan_uy_quyen_id", "noi_dung_viec", "ghi_chu"):
        assert copied["case"][field] == source["case"][field], field
    for field, value in SCALARS.items():
        assert copied["case"][field] == value

    # Rows: row_id mới toàn bộ, identity + dữ kiện/null/[] giữ nguyên.
    old_person_ids = [row["row_id"] for row in source["stage"]["people"]]
    old_asset_ids = [row["row_id"] for row in source["stage"]["assets"]]
    new_person_ids = [row["row_id"] for row in copied["stage"]["people"]]
    new_asset_ids = [row["row_id"] for row in copied["stage"]["assets"]]
    assert new_person_ids == [row["row_id"]
                              for row in copied["stage"]["people"]]
    assert not set(new_person_ids) & set(old_person_ids)
    assert not set(new_asset_ids) & set(old_asset_ids)
    for rid in new_person_ids + new_asset_ids:
        assert uuid.UUID(rid).version == 4
    assert [uuid.UUID(rid).version for rid in new_person_ids] == \
        [4] * len(new_person_ids)
    assert _row_free(copied["stage"])["people"] == \
        _row_free(source["stage"])["people"]
    assert _row_free(copied["stage"])["assets"] == \
        _row_free(source["stage"])["assets"]
    assert [row["entity_id"] for row in copied["stage"]["people"]] == \
        [row["entity_id"] for row in source["stage"]["people"]]
    assert [row["entity_id"] for row in copied["stage"]["assets"]] == \
        [row["entity_id"] for row in source["stage"]["assets"]]
    # Đất: ba hàng (có hàng trống), null, [] giữ đúng kiểu.
    assert _land_wire(copied["stage"]["assets"][0]) == [
        {"loai_dat": "ONT", "dien_tich": 100.0, "thoi_han": "Lâu dài"},
        {"loai_dat": None, "dien_tich": None, "thoi_han": None},
        {"loai_dat": "CLN", "dien_tich": 50.0, "thoi_han": "2063"},
    ]
    assert _land_wire(copied["stage"]["assets"][1]) == []
    assert _land_wire(copied["stage"]["assets"][2]) is None
    assert copied["stage"]["owner_row_id"] == new_person_ids[0]
    assert copied["stage"]["owner_row_id"] != source["stage"]["owner_row_id"]

    # Diagram: slot ids/vị trí giữ; personId map sang row mới.
    source_state = copy.deepcopy(source["diagram"]["state"])
    expected_state = copy.deepcopy(source_state)
    person_map = dict(zip(old_person_ids, new_person_ids))
    for node in expected_state["nodes"]:
        if node["personId"] is not None:
            node["personId"] = person_map[node["personId"]]
    assert copied["diagram"]["state"] == expected_state
    assert copied["diagram"]["domain"] == "inheritance"
    model_text = json.dumps(copied["diagram"]["render_model"],
                            ensure_ascii=False)
    for old_id in old_person_ids + old_asset_ids:
        assert old_id not in model_text
    assert copied["diagram"]["render_model"]["status"] == \
        source["diagram"]["render_model"]["status"]
    assert copied["diagram"].get("warnings") == \
        source["diagram"].get("warnings")

    # Participants: PK mới, pointer remap, identity theo rows mới, manual giữ.
    copy_rows = db.get(InheritanceCase, copy_id).person_rows
    identity_by_row = {r.row_id: r.customer_id for r in copy_rows}
    source_participants = sorted(
        db.get(InheritanceCase, source_id).participants,
        key=lambda p: p.id)
    copy_participants = sorted(
        db.get(InheritanceCase, copy_id).participants, key=lambda p: p.id)
    assert len(copy_participants) == len(source_participants) == 4
    assert not ({p.id for p in copy_participants}
                & {p.id for p in source_participants})
    manual = next(p for p in copy_participants
                  if p.ghi_chu == "Participant thủ công")
    legacy = next(p for p in copy_participants
                  if p.ghi_chu == "Legacy không pointer")
    legacy_source = next(p for p in source_participants
                         if p.ghi_chu == "Legacy không pointer")
    assert manual.case_person_row_id == new_person_ids[0]
    assert manual.customer_id == identity_by_row[new_person_ids[0]]
    assert manual.parent_case_person_row_id == new_person_ids[2]
    assert manual.parent_customer_id == identity_by_row[new_person_ids[2]]
    assert (manual.vai_tro, manual.hang_thua_ke, manual.co_nhan_tai_san,
            manual.ty_le) == ("Khác", 3, False, 12.5)
    assert legacy.case_person_row_id is None
    assert legacy.parent_case_person_row_id is None
    assert legacy.customer_id == legacy_source.customer_id

    # Property links: PK mới, trỏ asset row mới, primary theo vị trí.
    links = db.query(InheritanceCaseProperty).filter_by(
        case_id=copy_id).order_by(InheritanceCaseProperty.id).all()
    assert [link.case_asset_row_id for link in links] == new_asset_ids
    assert [link.is_primary for link in links] == [True, False, False]
    assert [link.property_id for link in links] == [
        row["entity_id"] for row in copied["stage"]["assets"]]
    after_counts = _case_counts(db)
    assert after_counts["cases"] == before_counts["cases"] + 1
    assert after_counts["participants"] == \
        before_counts["participants"] + 4
    assert after_counts["links"] == before_counts["links"] + 3
    assert after_counts["people"] == before_counts["people"] + 3
    assert after_counts["assets"] == before_counts["assets"] + 3
    assert after_counts["lands"] == before_counts["lands"] + 3

    # docx provenance: ngoài records.row_id, mọi thứ bằng nguồn.
    expected_docx = copy.deepcopy(source["docx_import"])
    asset_map = dict(zip(old_asset_ids, new_asset_ids))
    for source_block in expected_docx["sources"]:
        for record in source_block["records"]:
            row_id = record["row_id"]
            record["row_id"] = (person_map[row_id] if row_id in person_map
                                else asset_map[row_id])
    assert copied["docx_import"] == expected_docx
    assert copied["docx_import"]["sources"][0]["source_id"] == \
        docx_source_id
    assert copied["docx_import"]["sources"][0]["bindings"] == \
        source["docx_import"]["sources"][0]["bindings"]
    assert copied["docx_import"]["sources"][0]["template_sha256"] == \
        source["docx_import"]["sources"][0]["template_sha256"]

    # Nguồn không bị UPDATE/stamp; mọi identity dùng chung giữ nguyên.
    after = _source_snapshot(session_factory, source_id)
    assert after == before_snapshot
    assert _case_counts(db)["customers"] == before_counts["customers"]
    assert _case_counts(db)["properties"] == before_counts["properties"]


# ------------------------------------ hai chiều độc lập + Word hai bộ

def test_edits_are_independent_and_word_reads_both_copies(
        dataset, db, session_factory, store, tmp_path):
    saved, _source = _save_f1_source(db, dataset, store)
    source_id = saved["case"]["id"]
    copy_id = _duplicate(db, store, source_id)["case"]["id"]
    source_before = _reopen(session_factory, source_id)[0]
    copy_before = _reopen(session_factory, copy_id)[0]

    service = _service(db, store)
    copy_stage = service.get(copy_id)["stage"]
    copy_owner = copy_stage["owner_row_id"]
    owner_row = next(row for row in copy_stage["people"]
                     if row["row_id"] == copy_owner)
    owner_row["ho_ten"] = "Tên Bản Sao E2"
    owner_row["dia_chi"] = "Địa chỉ Bản Sao E2"
    service.commit_stage(copy_id, 1, copy_stage)

    assert _reopen(session_factory, source_id)[0] == source_before
    copy_changed = _reopen(session_factory, copy_id)[0]
    assert copy_changed["case"]["revision"] == 2
    assert next(p for p in copy_changed["stage"]["people"]
                if p["row_id"] == copy_owner)["ho_ten"] == "Tên Bản Sao E2"

    # Sửa nguồn: bản sao vẫn giữ giá trị riêng.
    source_stage = service.get(source_id)["stage"]
    spouse = next(row for row in source_stage["people"]
                  if row["ho_ten"] == "Trần Thị Bình")
    spouse["dia_chi"] = "Địa chỉ Nguồn E2"
    service.commit_stage(source_id, 1, source_stage)

    copy_after_source_edit = _reopen(session_factory, copy_id)[0]
    assert copy_after_source_edit["case"]["revision"] == 2
    assert "Địa chỉ Nguồn E2" not in json.dumps(
        copy_after_source_edit, ensure_ascii=False)
    assert copy_after_source_edit["stage"]["assets"] == \
        copy_before["stage"]["assets"]

    source_text = _export_text(session_factory, source_id, tmp_path,
                               key="nguon")
    copy_text = _export_text(session_factory, copy_id, tmp_path,
                             key="bansao")
    assert "Địa chỉ Nguồn E2" in source_text
    assert "Địa chỉ Bản Sao E2" in copy_text
    assert "Tên Bản Sao E2" in copy_text
    assert "Tên Bản Sao E2" not in source_text
    assert "Địa chỉ Nguồn E2" not in copy_text


# --------------------------------------------------- nháp thiếu dữ liệu

def test_duplicate_partial_draft_keeps_exact_content(dataset, db, store):
    spec = _case_spec(dataset, "F3")
    person = _proto(dataset, "people_prototypes", "P301")
    created = _service(db).create(
        _uuid4(),
        {"case_type": "inheritance", "document_type": "khai_nhan",
         "ngay_lap_ho_so": spec["metadata"]["ngay_lap_ho_so"]},
        {"owner_row_id": None, "people": [person], "assets": []}, None)
    source_id = created["case"]["id"]
    assert created["case"]["status"] == "draft"

    copied = _duplicate(db, store, source_id)
    assert copied["created"] is True
    assert copied["case"]["status"] == "draft"
    assert copied["case"]["revision"] == 1
    assert copied["stage"]["owner_row_id"] is None
    assert _row_free(copied["stage"]) == _row_free(created["stage"])
    assert copied["stage"]["people"][0]["ho_ten"] == \
        person["ho_ten"]
    assert copied["stage"]["people"][0]["entity_id"] is None
    assert copied["stage"]["people"][0]["row_id"] != person["row_id"]
    assert copied["stage"]["assets"] == []
    assert copied["diagram"]["state"]["nodes"] == []
    assert db.query(InheritanceParticipant).filter_by(
        ho_so_id=copied["case"]["id"]).count() == 0
    assert db.query(Customer).count() == 0
    assert db.query(Property).count() == 0


# ------------------------------------------------------------- idempotency

def _mutate_copy_to_done(db, store, copy_id):
    service = _service(db, store)
    service.update_meta(copy_id, 1, {"ghi_chu": "Bản sao đã sửa"})
    workspace = service.get(copy_id)
    case = workspace["case"]
    result = service.save_case(
        case_id=copy_id,
        base_revision=case["revision"],
        case_meta={
            "case_type": case["case_type"],
            "document_type": case["document_type"],
            "ngay_lap_ho_so": case["ngay_lap_ho_so"],
            "noi_niem_yet": case["noi_niem_yet"],
            "nguoi_nhan_uy_quyen": case["nguoi_nhan_uy_quyen"],
            "nguoi_nhan_uy_quyen_id": case["nguoi_nhan_uy_quyen_id"],
            "noi_dung_viec": case["noi_dung_viec"],
            "ghi_chu": case["ghi_chu"],
        },
        stage=workspace["stage"],
        state=workspace["diagram"]["state"])
    assert result["case"]["status"] == "done"
    return result


def test_replay_same_key_returns_current_copy_without_write(
        dataset, db, session_factory, store):
    saved, _docx = _save_f1_source(db, dataset, store)
    source_id = saved["case"]["id"]
    key = _uuid4()
    first = _duplicate(db, store, source_id, key)
    copy_id = first["case"]["id"]
    assert first["created"] is True

    done = _mutate_copy_to_done(db, store, copy_id)
    counts = _case_counts(db)
    snapshot = _source_snapshot(session_factory, copy_id)

    replay = _duplicate(db, store, source_id, key)
    assert replay["created"] is False
    assert replay["source_case_id"] == source_id
    assert replay["case"]["id"] == copy_id
    assert replay["case"]["status"] == "done"
    assert replay["case"]["revision"] == done["case"]["revision"] == 3
    assert replay["case"]["ghi_chu"] == "Bản sao đã sửa"
    assert _case_counts(db) == counts
    assert _source_snapshot(session_factory, copy_id) == snapshot

    # Nguồn bị xóa sau lần đầu → replay vẫn trả bản sao đã tạo.
    db.query(InheritanceParticipant).filter_by(
        ho_so_id=source_id).delete(synchronize_session=False)
    db.query(InheritanceCaseProperty).filter_by(
        case_id=source_id).delete(synchronize_session=False)
    db.query(CasePersonRow).filter_by(
        case_id=source_id).delete(synchronize_session=False)
    source_asset_ids = [row.row_id for row in db.query(CaseAssetRow).filter_by(
        case_id=source_id)]
    if source_asset_ids:
        db.query(CaseAssetLandRow).filter(
            CaseAssetLandRow.case_asset_row_id.in_(
                source_asset_ids)).delete(synchronize_session=False)
    db.query(CaseAssetRow).filter_by(
        case_id=source_id).delete(synchronize_session=False)
    db.query(InheritanceCase).filter_by(id=source_id).delete(
        synchronize_session=False)
    db.commit()
    assert db.query(InheritanceCase).filter_by(id=source_id).count() == 0
    counts_after_delete = _case_counts(db)
    again = _duplicate(db, store, source_id, key)
    assert again["created"] is False
    assert again["case"]["id"] == copy_id
    assert again["case"]["revision"] == 3
    assert _case_counts(db) == counts_after_delete


def test_same_key_other_source_or_command_conflicts(dataset, db, store):
    saved, _docx = _save_f1_source(db, dataset, store)
    source_id = saved["case"]["id"]
    key = _uuid4()
    _duplicate(db, store, source_id, key)

    person = _proto(dataset, "people_prototypes", "P301")
    other = _service(db).create(
        _uuid4(),
        {"case_type": "inheritance", "document_type": "khai_nhan",
         "ngay_lap_ho_so": None},
        {"owner_row_id": None, "people": [person], "assets": []}, None)
    with pytest.raises(WorkspaceError) as exc:
        _duplicate(db, store, other["case"]["id"], key)
    assert exc.value.code == "idempotency_conflict"

    # Cùng key nhưng command create → hash khác command/intent.
    key2 = _uuid4()
    _duplicate(db, store, source_id, key2)
    with pytest.raises(WorkspaceError) as exc:
        _service(db).create(
            key2,
            {"case_type": "inheritance", "document_type": "khai_nhan",
             "ngay_lap_ho_so": None},
            {"owner_row_id": None, "people": [person], "assets": []}, None)
    assert exc.value.code == "idempotency_conflict"


def test_validation_and_missing_source(db, store):
    for kwargs in (
            {"source_case_id": True, "idempotency_key": _uuid4()},
            {"source_case_id": 0, "idempotency_key": _uuid4()},
            {"source_case_id": "1", "idempotency_key": _uuid4()},
            {"source_case_id": None, "idempotency_key": _uuid4()},
            {"source_case_id": 1, "idempotency_key": None},
            {"source_case_id": 1,
             "idempotency_key": str(uuid.uuid1())},
    ):
        with pytest.raises(WorkspaceError) as exc:
            _service(db, store).duplicate(**kwargs)
        assert exc.value.code == "validation_error"
    assert _case_counts(db)["cases"] == 0

    with pytest.raises(WorkspaceError) as exc:
        _duplicate(db, store, 424242)
    assert exc.value.code == "case_not_found"
    assert exc.value.details["source_case_id"] == 424242


# ---------------------------------------------------------- fault rollback

def test_faults_after_land_participant_binding_rollback_fully(
        dataset, db, session_factory, store, monkeypatch):
    saved, _docx = _save_f1_source(db, dataset, store)
    source_id = saved["case"]["id"]
    before = _source_snapshot(session_factory, source_id)
    counts = _case_counts(db)

    def boom(*_args, **_kwargs):
        raise RuntimeError("injected")

    faults = [
        ("after_land", "_copy_participants"),
        ("after_participant", "_copy_docx_import"),
        ("after_binding", "_validate_docx_bindings"),
    ]
    for label, attribute in faults:
        key = _uuid4()
        monkeypatch.setattr(CaseWorkspaceService, attribute, boom)
        with pytest.raises(RuntimeError):
            _duplicate(db, store, source_id, key)
        monkeypatch.undo()
        assert _case_counts(db) == counts, label
        assert db.query(InheritanceCase).filter(
            InheritanceCase.workspace_idempotency_key == key).count() == 0
        assert _source_snapshot(session_factory, source_id) == before, label

    # Key sau rollback dùng lại được (không mồ côi).
    recovered = _duplicate(db, store, source_id, _uuid4())
    assert recovered["created"] is True


def test_two_sessions_same_key_create_exactly_one_copy(
        dataset, session_factory, store):
    setup = session_factory()
    try:
        saved, _docx = _save_f1_source(setup, dataset, store)
        source_id = saved["case"]["id"]
        before_cases = setup.query(InheritanceCase).count()
    finally:
        setup.close()

    key = _uuid4()
    barrier = threading.Barrier(2)
    results = []
    lock = threading.Lock()

    def worker():
        session = session_factory()
        try:
            barrier.wait()
            data = CaseWorkspaceService(session, lambda: store).duplicate(
                source_case_id=source_id, idempotency_key=key)
            outcome = ("ok", data["created"], data["case"]["id"])
        except WorkspaceError as exc:
            outcome = ("error", exc.code, exc.details)
        except Exception as exc:  # pragma: no cover - báo lỗi thật
            outcome = ("crash", repr(exc), None)
        finally:
            session.close()
        with lock:
            results.append(outcome)

    threads = [threading.Thread(target=worker) for _ in range(2)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert len(results) == 2
    created = [r for r in results if r[0] == "ok" and r[1] is True]
    replayed = [r for r in results if r[0] == "ok" and r[1] is False]
    busy = [r for r in results
            if r[0] == "error" and r[1] == "engine_unavailable"]
    assert len(created) == 1, results
    assert len(created) + len(replayed) + len(busy) == 2, results
    if busy:
        # retryable: retry nguyên key trả bản sao, không tạo thêm.
        session = session_factory()
        try:
            retry = CaseWorkspaceService(
                session, lambda: store).duplicate(
                    source_case_id=source_id, idempotency_key=key)
            assert retry["created"] is False
        finally:
            session.close()
    assert created[0][2] == (replayed[0][2] if replayed else created[0][2])

    check = session_factory()
    try:
        assert check.query(InheritanceCase).filter(
            InheritanceCase.workspace_idempotency_key == key).count() == 1
        assert check.query(InheritanceCase).count() == before_cases + 1
    finally:
        check.close()


def test_duplicate_two_party_copies_slots_without_owner(dataset, db, store):
    people = [_proto(dataset, "people_prototypes", alias)
              for alias in ("P102", "P103")]
    assets = [_proto(dataset, "asset_prototypes", "T101")]
    nodes = [{"id": f"p{i}", "personId": None, "hidden": False,
              "deleted": False} for i in range(1, 31)]
    nodes[0]["personId"] = people[0]["row_id"]
    nodes[15]["personId"] = people[1]["row_id"]
    created = _service(db).create(
        _uuid4(),
        {"case_type": "two_party", "document_type": "chuyen_nhuong",
         "ngay_lap_ho_so": "2026-10-22"},
        {"people": people, "assets": assets},
        {"version": 3, "domain": "two_party", "nodes": nodes})

    copied = _duplicate(db, store, created["case"]["id"])
    assert "owner_row_id" not in copied["stage"]
    assert copied["stage"]["people"][0]["row_id"] != \
        people[0]["row_id"]
    assert _row_free(copied["stage"]) == _row_free(created["stage"])
    old_ids = [row["row_id"] for row in created["stage"]["people"]]
    new_ids = [row["row_id"] for row in copied["stage"]["people"]]
    expected_nodes = copy.deepcopy(created["diagram"]["state"]["nodes"])
    person_map = dict(zip(old_ids, new_ids))
    for node in expected_nodes:
        if node["personId"] is not None:
            node["personId"] = person_map[node["personId"]]
    assert copied["diagram"]["state"]["nodes"] == expected_nodes
    assert copied["diagram"]["render_model"]["status"] == "unsupported"
    assert db.query(InheritanceCase).filter_by(
        id=copied["case"]["id"]).one().nguoi_chet_id is None


# ------------------------------ E1: hai khóa, độc lập, nguồn Done theo mốc


def test_two_keys_two_independent_copies_and_draft_never_source(
        dataset, db, session_factory, store):
    """Hai khóa khác nhau → hai bản sao độc lập; chỉ bản đã lưu Done làm nguồn.

    Bản sao 2 còn draft dù có mốc ghi mới hơn không bao giờ xuất hiện làm
    nguồn. Replay cùng khóa bản sao 1 trả đúng bản sao 1, không tạo bản sao 3.
    """
    saved, _docx = _save_f1_source(db, dataset, store)
    source_id = saved["case"]["id"]
    key1, key2 = _uuid4(), _uuid4()
    first = _duplicate(db, store, source_id, key1)
    second = _duplicate(db, store, source_id, key2)
    copy1, copy2 = first["case"]["id"], second["case"]["id"]

    assert first["created"] is True and second["created"] is True
    assert len({source_id, copy1, copy2}) == 3

    rows1 = [row["row_id"] for row in first["stage"]["people"]] + \
        [row["row_id"] for row in first["stage"]["assets"]]
    rows2 = [row["row_id"] for row in second["stage"]["people"]] + \
        [row["row_id"] for row in second["stage"]["assets"]]
    source_rows = [row["row_id"] for row in saved["stage"]["people"]] + \
        [row["row_id"] for row in saved["stage"]["assets"]]
    assert not set(rows1) & set(rows2)
    assert not set(rows1) & set(source_rows)
    assert not set(rows2) & set(source_rows)

    service = _service(db, store)
    source_before = _source_snapshot(session_factory, source_id)
    copy2_before = _source_snapshot(session_factory, copy2)

    # Sửa Stage bản sao 1: bản sao 2 và nguồn gốc không đổi.
    stage1 = service.get(copy1)["stage"]
    edited = next(row for row in stage1["people"]
                  if row["row_id"] == stage1["owner_row_id"])
    edited["ho_ten"] = "Người Bản Sao Một"
    service.commit_stage(copy1, first["case"]["revision"], stage1)
    assert _source_snapshot(session_factory, source_id) == source_before
    assert _source_snapshot(session_factory, copy2) == copy2_before

    # Lưu bản sao 1 lên Done → mốc lưu mới nhất.
    workspace1 = service.get(copy1)
    case1 = workspace1["case"]
    saved1 = service.save_case(
        case_id=copy1, base_revision=case1["revision"],
        case_meta={field: case1[field] for field in (
            "case_type", "document_type", "ngay_lap_ho_so",
            "noi_niem_yet", "nguoi_nhan_uy_quyen",
            "nguoi_nhan_uy_quyen_id", "noi_dung_viec", "ghi_chu")},
        stage=workspace1["stage"],
        state=workspace1["diagram"]["state"])
    assert saved1["case"]["status"] == "done"

    resolved = service.source_latest()
    assert resolved["source_case_id"] == copy1
    assert resolved["source_status"] == "done"
    assert _owner_name(resolved["stage"]) == "Người Bản Sao Một"
    assert _stage_content(resolved["stage"]) == \
        _stage_content(service.get(copy1)["stage"])
    assert resolved["stage"]["owner_row_id"] not in rows1

    # Bản sao 2 vẫn draft nhưng có mốc ghi mới hơn bản sao 1 → vẫn không
    # được làm nguồn; resolver tiếp tục trả bản sao 1.
    stage2 = service.get(copy2)["stage"]
    stage2["people"][0]["dia_chi"] = "Địa chỉ bản sao hai"
    service.commit_stage(copy2, second["case"]["revision"], stage2)
    assert service.get(copy2)["case"]["status"] == "draft"

    again = service.source_latest()
    assert again["source_case_id"] == copy1
    assert again["source_status"] == "done"
    assert _owner_name(again["stage"]) == "Người Bản Sao Một"
    assert "Địa chỉ bản sao hai" not in json.dumps(
        again["stage"], ensure_ascii=False)

    # Replay cùng khóa bản sao 1 → đúng bản sao 1, không tạo bản sao 3.
    counts = _case_counts(db)
    replay = _duplicate(db, store, source_id, key1)
    assert replay["created"] is False
    assert replay["case"]["id"] == copy1
    assert replay["case"]["status"] == "done"
    assert db.query(InheritanceCase).count() == counts["cases"]
    assert _case_counts(db) == counts
    assert service.source_latest()["source_case_id"] == copy1
