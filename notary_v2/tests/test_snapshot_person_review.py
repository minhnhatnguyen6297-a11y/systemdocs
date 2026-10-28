import json
from datetime import date
from docx import Document
from tests import test_case_metadata as t
from tests.test_case_metadata import db
from services import word_batch_export as wbe


def test_canonical_person_token_preserves_confirmed_death_evidence(db):
    data = t._create(db, with_state=True, people_overrides=({
        'loai_giay_to': 'Giấy chứng tử', 'noi_cap': 'UBND X'},))
    case = db.get(t.InheritanceCase, data['case']['id'])
    mapping = t.word_engine.build_template_mapping(case)
    assert mapping['[loaigiayto1]'] == 'Giấy chứng tử'
    assert mapping['[noicap1]'] == 'UBND X'


def test_derived_fields_ignore_other_case_master_edits(db):
    data = t._create(db, people_overrides=({}, {'ngay_cap': '2024-09-30'}))
    stage = data['stage']
    # Another consumer updates the shared master after this case committed.
    person = db.get(t.Customer, stage['people'][1]['entity_id'])
    person.ngay_cap = date(2024, 10, 1)
    db.commit()
    stage['people'][1]['ngay_cap'] = '2024-10-02'
    result = t._commit(db, data['case']['id'], data['case']['revision'], stage)
    assert result['stage']['people'][1]['loai_giay_to'] == 'Căn cước'


# ------------------------------- độ đầy đủ ngày: Word (Owner 15/10/2026)


def _export_docx(db, case_id, tokens, tmp_path):
    """Export thật (engine + python-docx) một mẫu tạm; trả text paragraph."""
    template_dir = tmp_path / "templates"
    dest = tmp_path / "out"
    template_dir.mkdir(parents=True)
    dest.mkdir(parents=True)
    template = template_dir / "probe.docx"
    doc = Document()
    for token in tokens:
        doc.add_paragraph(token)
    doc.save(template)
    case = db.get(t.InheritanceCase, case_id)
    spec = wbe.WordDocumentSpec(document_key="probe", display_name="Probe",
                                filename_stem="Probe")
    batch = wbe.export_batch(
        case, case_id=case_id, document_keys=["probe"], dest_dir=dest,
        resolve_template=lambda _k: template, doc_specs=[spec])
    entry = batch["documents"][0]
    assert entry["status"] == "saved", entry
    out = dest / entry["actual_filename"]
    return [p.text for p in Document(str(out)).paragraphs if p.text.strip()]


def test_year_only_and_full_dates_stay_distinct_through_export(db, tmp_path):
    """`2010` (chỉ năm) vs `01/01/2010` (đầy đủ) đi trọn lưu → mở lại →
    sửa trường khác → xuất Word thật; token văn xuôi không rò ISO."""
    data = t._create(db, with_state=True, people_overrides=(
        {"ngay_sinh": "2010", "ngay_chet": "2015", "ngay_cap": "2018"},
        {"ngay_sinh": "2010-01-01"}))
    case_id = data["case"]["id"]
    got = t._service(db).get(case_id)
    stage = got["stage"]
    stage["people"][0]["dia_chi"] = "xã đã sửa"     # sửa trường khác
    t._commit(db, case_id, got["case"]["revision"], stage)

    mapping = t.word_engine.build_template_mapping(
        db.get(t.InheritanceCase, case_id))
    assert mapping["[Người 1 - Ngày sinh]"] == "2010"
    assert mapping["[Người 2 - Ngày sinh]"] == "01/01/2010"
    assert mapping["[Người 1 - Năm sinh]"] == "2010"
    assert mapping["[Người 2 - Năm sinh]"] == "2010"
    assert mapping["[ngaysinh1]"] == "2010"
    assert mapping["[ngaychet1]"] == "2015"
    assert mapping["[ngaycap1]"] == "2018"
    assert "Sinh ngày: 2010." in mapping["[Dòng người 1]"]

    paragraphs = _export_docx(db, case_id, [
        "[Người 1 - Ngày sinh]", "[Người 2 - Ngày sinh]",
        "[ngaysinh1]", "[Năm sinh 1]", "[Dòng người 1]",
        "[ngaychet1]"], tmp_path)
    assert paragraphs[0] == "2010"
    assert paragraphs[1] == "01/01/2010"
    assert paragraphs[2] == "2010"
    assert paragraphs[3] == "2010"
    assert "Sinh ngày: 2010." in paragraphs[4]
    assert "2010-01-01" not in paragraphs[4]
    assert paragraphs[5] == "2015"


def test_row_dates_are_sot_and_json_stage_is_ignored(db, tmp_path):
    """Ngày Người đọc từ dòng case-owned; JSON không còn là Stage.

    - Mất Diagram (chỉ còn owner_row_id) → ngày vẫn đủ từ rows.
    - Chèn `stage` giả vào JSON → xuất vẫn đúng rows (không đọc JSON/master).
    """
    data = t._create(db, with_state=True, people_overrides=(
        {"ngay_sinh": "1950-01-01"}, {}))
    case_id = data["case"]["id"]
    case = db.get(t.InheritanceCase, case_id)
    payload = json.loads(case.case_state_json)
    case.case_state_json = json.dumps({
        "schemaVersion": 4, "owner_row_id": payload["owner_row_id"]})
    db.commit()

    mapping = t.word_engine.build_template_mapping(
        db.get(t.InheritanceCase, case_id))
    assert "Sinh ngày: 01/01/1950" in mapping["[Dòng người 1]"]
    assert mapping["[ngaysinh1]"] == "01/01/1950"

    data2 = t._create(db, with_state=True, people_overrides=(
        {"ngay_sinh": "1950-01-01", "ngay_chet": "2011-01-01",
         "ngay_cap": "2011-01-02"}, {}))
    case2_id = data2["case"]["id"]
    case2 = db.get(t.InheritanceCase, case2_id)
    payload = json.loads(case2.case_state_json)
    owner_row_id = payload["owner_row_id"]
    payload["stage"] = [{
        "row_id": owner_row_id, "ho_ten": "Tên bịa từ JSON",
        "ngay_sinh": "1900", "ngay_chet": None, "ngay_cap": None}]
    case2.case_state_json = json.dumps(payload, ensure_ascii=False)
    db.commit()

    owner_name = db.get(
        t.Customer, db.get(t.InheritanceCase, case2_id).person_rows[0].customer_id
    ).ho_ten
    mapping2 = t.word_engine.build_template_mapping(
        db.get(t.InheritanceCase, case2_id))
    assert mapping2["[Người 1 - Họ tên]"] == owner_name
    assert mapping2["[Người 1 - Ngày sinh]"] == "01/01/1950"
    assert mapping2["[Người 1 - Ngày chết]"] == "01/01/2011"
    assert mapping2["[Người 1 - Ngày cấp]"] == "02/01/2011"
    assert mapping2["[ngaysinh1]"] == "01/01/1950"
    assert mapping2["[Chủ đất chết 1 - Họ tên]"] == owner_name

    paragraphs = _export_docx(db, case2_id, [
        "[Người 1 - Ngày sinh]", "[Người 1 - Ngày chết]",
        "[Người 1 - Ngày cấp]", "[Dòng người 1]",
        "[Chủ đất chết 1 - Họ tên]"], tmp_path)
    assert paragraphs[0] == "01/01/1950"
    assert paragraphs[1] == "01/01/2011"
    assert paragraphs[2] == "02/01/2011"
    assert "Sinh ngày: 01/01/1950" in paragraphs[3]
    assert paragraphs[4] == owner_name
