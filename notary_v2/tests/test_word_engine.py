from datetime import date
from pathlib import Path
from types import SimpleNamespace

_PCDS_TEMPLATE = (Path(__file__).resolve().parent.parent
                  / "word_templates" / "1. PCDS .docx")

import pytest
from docx import Document

from services.word_engine import (
    NO_SOURCE_PLACEHOLDERS,
    WordExportValidationError,
    _collect_placeholders,
    build_word_context,
    build_template_mapping,
    find_unresolved_placeholders,
    is_fixed_form_code,
    list_public_builtin_templates,
    replace_in_doc,
)


def _person(
    cid,
    name,
    *,
    gender="Nam",
    born=None,
    dead=None,
    doc_no="",
    issued=None,
    address="",
    role="",
):
    return SimpleNamespace(
        id=cid,
        ho_ten=name,
        gioi_tinh=gender,
        ngay_sinh=born,
        ngay_chet=dead,
        so_giay_to=doc_no,
        ngay_cap=issued,
        dia_chi=address,
        vai_tro=role,
        loai_giay_to="Căn cước công dân",
        noi_cap="Cục cảnh sát quản lý hành chính về trật tự xã hội",
        loai_dia_chi="Thường trú tại",
    )


def _participant(person, role, *, receive=True, share=0, inheritance_order=1):
    return SimpleNamespace(
        customer=person,
        customer_id=person.id,
        vai_tro=role,
        co_nhan_tai_san=receive,
        ty_le=share,
        hang_thua_ke=inheritance_order,
        parent_customer_id=None,
    )


def _prop(**kwargs):
    defaults = dict(
        so_serial="BM 1451111",
        so_vao_so="CS123",
        so_thua_dat="45",
        so_to_ban_do="12",
        dia_chi="Thửa đất tại xã Minh Tân",
        loai_dat="Đất ở tại nông thôn",
        dien_tich=120.5,
        loai_so="Giấy chứng nhận quyền sử dụng đất",
        land_rows_json='[{"loai_dat":"ONT","dien_tich":"80","thoi_han":"Lâu dài"}]',
        hinh_thuc_su_dung="Sử dụng riêng",
        thoi_han="Lâu dài",
        nguon_goc="Nhà nước công nhận quyền sử dụng đất",
        ngay_cap=date(2020, 7, 8),
        co_quan_cap="UBND huyện Ý Yên",
    )
    defaults.update(kwargs)
    return SimpleNamespace(**defaults)


def _case_state(*people, nodes=None):
    stage = []
    for p in people:
        stage.append({
            "id": str(p.id),
            "ho_ten": p.ho_ten,
            "gioi_tinh": p.gioi_tinh,
            "ngay_sinh": p.ngay_sinh.isoformat() if p.ngay_sinh else None,
            "ngay_chet": p.ngay_chet.isoformat() if p.ngay_chet else None,
            "so_giay_to": p.so_giay_to,
            "ngay_cap": p.ngay_cap.isoformat() if p.ngay_cap else None,
            "noi_cap": p.noi_cap,
            "dia_chi": p.dia_chi,
            "loai_dia_chi": p.loai_dia_chi,
        })
    payload = {"schemaVersion": 1, "stage": stage, "diagram": {}}
    if nodes:
        payload["diagram"] = {"engineState": {"nodes": nodes}}
    import json
    return json.dumps(payload, ensure_ascii=False)


def _case_state_v2(*people, nodes, allocations):
    """Payload v4 với Diagram state row-keyed + render_model allocations."""
    import json

    payload = json.loads(_case_state(*people, nodes=nodes))
    payload["diagram"] = {
        "state": {"version": 3, "domain": "inheritance", "nodes": nodes},
        "render_model": {"status": "complete", "allocations": allocations},
    }
    return json.dumps(payload, ensure_ascii=False)


def _node(person_id, *, role="Khac", will_receive=False, is_land_owner=False, legacy_decision=None):
    node = {
        "id": f"node_{person_id}",
        "personId": str(person_id),
        "role": role,
        "isLandOwner": is_land_owner,
        "hidden": False,
        "deleted": False,
    }
    if will_receive is not None:
        node["willReceive"] = will_receive
    if legacy_decision is not None:
        node["inheritanceDecision"] = legacy_decision
    return node


def _person_row(stage_row, vitri):
    raw_id = stage_row.get("entity_id")
    if raw_id is None:
        legacy = str(stage_row.get("id") or "")
        raw_id = int(legacy) if legacy.isdigit() else None
    return SimpleNamespace(
        row_id=stage_row.get("row_id") or f"row-{raw_id}-{vitri}",
        customer_id=raw_id,
        vitri=vitri,
        ho_ten=stage_row.get("ho_ten"),
        gioi_tinh=stage_row.get("gioi_tinh"),
        ngay_sinh=stage_row.get("ngay_sinh"),
        ngay_chet=stage_row.get("ngay_chet"),
        so_giay_to=stage_row.get("so_giay_to"),
        ngay_cap=stage_row.get("ngay_cap"),
        noi_cap=stage_row.get("noi_cap"),
        dia_chi=stage_row.get("dia_chi"),
        place_of_origin=stage_row.get("place_of_origin"),
        loai_giay_to=stage_row.get("loai_giay_to"),
        loai_dia_chi=stage_row.get("loai_dia_chi"),
    )


def _asset_row(prop, vitri):
    import json

    raw = prop.land_rows_json
    rows = []
    if raw:
        try:
            rows = json.loads(raw)
        except ValueError:
            rows = []
    land_rows = [
        SimpleNamespace(vitri=index + 1,
                        loaidat=row.get("loai_dat"),
                        dientich=row.get("dien_tich"),
                        thoihan=row.get("thoi_han"))
        for index, row in enumerate(rows if isinstance(rows, list) else [])
        if isinstance(row, dict)
    ]
    return SimpleNamespace(
        row_id=f"asset-{vitri}",
        property_id=getattr(prop, "id", None),
        vitri=vitri,
        so_serial=prop.so_serial,
        so_vao_so=prop.so_vao_so,
        so_thua_dat=prop.so_thua_dat,
        so_to_ban_do=prop.so_to_ban_do,
        dia_chi=prop.dia_chi,
        loai_so=prop.loai_so,
        hinh_thuc_su_dung=prop.hinh_thuc_su_dung,
        thoi_han=prop.thoi_han,
        nguon_goc=prop.nguon_goc,
        ngay_cap=prop.ngay_cap,
        co_quan_cap=prop.co_quan_cap,
        land_rows_present=bool(raw is not None),
        land_rows=land_rows,
    )


def _stage_fields(person):
    return {
        "id": str(person.id),
        "ho_ten": person.ho_ten,
        "gioi_tinh": person.gioi_tinh,
        "ngay_sinh": person.ngay_sinh.isoformat() if person.ngay_sinh else None,
        "ngay_chet": person.ngay_chet.isoformat() if person.ngay_chet else None,
        "so_giay_to": person.so_giay_to,
        "ngay_cap": person.ngay_cap.isoformat() if person.ngay_cap else None,
        "noi_cap": person.noi_cap,
        "dia_chi": person.dia_chi,
        "place_of_origin": getattr(person, "place_of_origin", None),
        "loai_giay_to": getattr(person, "loai_giay_to", None),
        "loai_dia_chi": getattr(person, "loai_dia_chi", None),
    }


def _make_case(owner, participants=None, *, properties=None, case_state=None):
    """Fake case Storage B: dòng Người/Tài sản + JSON v4 (không Stage).

    `case_state` chỉ là nguồn dựng Diagram/rows cho test; sau khi dựng,
    Word chỉ đọc `person_rows`/`asset_rows`/`case_state_json` giống ORM.
    """
    import json

    participants = participants or []
    if properties is None:
        properties = [_prop()]
    primary = properties[0] if properties else _prop()
    payload = json.loads(case_state) if case_state else {}
    domain = payload.get("case_type") or "inheritance"
    stage_items = payload.get("stage") if isinstance(payload.get("stage"), list) else None
    if stage_items is None:
        people = ([owner] if owner is not None else []) + [
            p.customer for p in participants]
        stage_items = [_stage_fields(person) for person in people]

    person_rows = [_person_row(item, index + 1)
                   for index, item in enumerate(stage_items)]
    entity_to_row = {str(row.customer_id): row.row_id
                     for row in person_rows if row.customer_id is not None}
    asset_rows = [_asset_row(prop, index + 1)
                  for index, prop in enumerate(properties)]

    diagram = payload.get("diagram")
    diagram = diagram if isinstance(diagram, dict) else {}
    state = diagram.get("state")
    legacy_engine = diagram.get("engineState")
    if not isinstance(state, dict) and isinstance(legacy_engine, dict):
        state = {"version": 3, "domain": domain,
                 "nodes": legacy_engine.get("nodes") or []}
    if isinstance(state, dict):
        nodes = []
        for node in state.get("nodes") or []:
            if isinstance(node, dict):
                mapped = entity_to_row.get(str(node.get("personId")))
                nodes.append({**node, "personId": mapped or node.get("personId")})
        state = {**state, "nodes": nodes}
    render_model = diagram.get("render_model")
    if isinstance(render_model, dict) and isinstance(render_model.get("allocations"), dict):
        render_model = {
            **render_model,
            "allocations": {
                entity_to_row.get(str(key), key): value
                for key, value in render_model["allocations"].items()},
        }

    case_type = {"case_type": domain,
                 "diagram": {"state": state or {},
                             **({"render_model": render_model}
                                if render_model else {})}}
    if domain == "inheritance":
        case_type["owner_row_id"] = entity_to_row.get(
            str(getattr(owner, "id", None)))
    new_state = json.dumps(case_type, ensure_ascii=False)

    case = SimpleNamespace(
        id=99,
        nguoi_chet=owner,
        tai_san=primary,
        ngay_lap_ho_so=date(2026, 7, 3),
        loai_van_ban="thoa_thuan",
        trang_thai="draft",
        workspace_revision=1,
        noi_niem_yet="UBND xã Minh Tân",
        nguoi_nhan_uy_quyen=None,
        nguoi_nhan_uy_quyen_id=None,
        noi_dung_viec=None,
        ghi_chu="Ghi chú nội bộ",
        participants=participants,
        person_rows=person_rows,
        asset_rows=asset_rows,
        case_state_json=new_state,
    )
    by_customer = {row.customer_id: row.row_id for row in person_rows}
    for participant in participants:
        participant.case_person_row_id = by_customer.get(participant.customer_id)
    return case


def test_build_template_mapping_exposes_standard_vietnamese_placeholders_and_legacy_aliases():
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    receiver = _person(2, "Nguyễn Văn Con", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    prop = _prop()
    case = _make_case(
        owner,
        participants=[_participant(receiver, "Con")],
        properties=[prop],
    )
    mapping = build_template_mapping(case, today=date(2026, 7, 3))

    assert mapping["[Họ tên chủ đất]"] == "Nguyễn Văn Chủ"
    assert mapping["[Chủ đất chết 1 - Họ tên]"] == "Nguyễn Văn Chủ"
    assert mapping["[Người nhận 1 - Họ tên]"] == "Nguyễn Văn Con"
    assert mapping["[Số serial]"] == "BM 1451111"
    assert mapping["[Ngày lập hồ sơ]"] == "03/07/2026"
    assert mapping["[Tên 1]"] == "Nguyễn Văn Chủ"
    # Owner NAIA-12: [Năm chết] giữ định dạng NĂM (= ngaychet1), không in ngày đầy đủ.
    assert mapping["[Năm chết]"] == "2024"
    assert mapping["[Năm chết 1]"] == "2024"
    assert "[ho_ten_1_chu_dat]" not in mapping


def test_serial_keeps_surrounding_spaces_verbatim():
    """#40: serial nguyên văn — không bị _safe_text strip khi map/export."""
    prop = _prop(so_serial="  RAW  ")
    case = _simple_export_case(properties=[prop])
    mapping = build_template_mapping(case)

    assert mapping["[Serial]"] == "  RAW  "
    assert mapping["[Số serial]"] == "  RAW  "
    assert mapping["[Serial 1]"] == "  RAW  "

    doc = Document()
    doc.add_paragraph("S=[Serial]")
    replace_in_doc(doc, mapping)
    assert doc.paragraphs[0].text == "S=  RAW  "


def test_build_template_mapping_resolves_role_and_data_only_snippets():
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    spouse = _person(2, "Trần Thị Vợ", gender="Nữ", born=date(1955, 4, 5), doc_no="001055000002")
    child = _person(3, "Nguyễn Văn Con", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    case = _make_case(
        owner,
        participants=[_participant(spouse, "Vợ/Chồng"), _participant(child, "Con")],
        properties=[_prop()],
    )
    mapping = build_template_mapping(case, today=date(2026, 7, 3))

    assert mapping["[Họ tên vợ/chồng]"] == "Trần Thị Vợ"
    assert "Nguyễn Văn Con" in mapping["[Danh sách người thừa kế]"]


def test_replace_in_doc_replaces_split_runs_tables_headers_and_footers():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("[Họ tên")
    para.add_run(" chủ đất]")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "[Số serial]"
    doc.sections[0].header.paragraphs[0].text = "[Địa chỉ đất]"
    doc.sections[0].footer.paragraphs[0].text = "[Tên 1]"

    replace_in_doc(
        doc,
        {
            "[Họ tên chủ đất]": "Nguyễn Văn Chủ",
            "[Số serial]": "BM 1451111",
            "[Địa chỉ đất]": "Thửa đất tại xã Minh Tân",
            "[Tên 1]": "Nguyễn Văn Chủ",
        },
    )

    assert doc.paragraphs[0].text == "Nguyễn Văn Chủ"
    assert table.cell(0, 0).text == "BM 1451111"
    assert doc.sections[0].header.paragraphs[0].text == "Thửa đất tại xã Minh Tân"
    assert doc.sections[0].footer.paragraphs[0].text == "Nguyễn Văn Chủ"


from docx.oxml.ns import qn

TEXT_TAG = qn("w:t")


def _simple_export_case(properties=None):
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    receiver = _person(2, "Nguyễn Văn Con", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    if properties is None:
        properties = [_prop()]
    return _make_case(owner, participants=[_participant(receiver, "Con")], properties=properties)


def test_correct_spelling_thoi_han_alias_shares_the_same_source():
    case = _simple_export_case()
    mapping = build_template_mapping(case)

    assert mapping["[Thời hạn 1]"] == mapping["[Thờ hạn 1]"]

    doc = Document()
    doc.add_paragraph("Thời hạn sử dụng đất: [Thời hạn 1]")
    replace_in_doc(doc, mapping)
    assert find_unresolved_placeholders(doc) == []


def test_no_source_placeholders_are_reported_not_silently_emptied():
    mapping = build_template_mapping(_simple_export_case())

    # `[Ngườ ủy quyền]` có nguồn metadata cột; null tường minh → chuỗi rỗng.
    assert mapping["[Ngườ ủy quyền]"] == ""
    for token in NO_SOURCE_PLACEHOLDERS:
        if token == "[Ngườ ủy quyền]":
            continue
        assert token not in mapping, f"{token} không được map về rỗng"

    doc = Document()
    doc.add_paragraph("[Số công chứng] [Giá chuyển nhượng] [SĐT]")
    replace_in_doc(doc, mapping)
    assert find_unresolved_placeholders(doc) == [
        "[Giá chuyển nhượng]",
        "[SĐT]",
        "[Số công chứng]",
    ]


def _summary_export_case(properties=None, *, role=""):
    """Ca đã có Diagram: 1 chủ đất đã chết + 1 người nhận — dùng cho test chuỗi/mở rộng."""
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001", role=role)
    receiver = _person(2, "Nguyễn Văn Con", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    case = _make_case(
        owner,
        participants=[_participant(receiver, "Con")],
        properties=properties if properties is not None else [_prop()],
        case_state=_case_state(
            owner, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    return case


def test_so_thanh_chu_reads_vietnamese_numbers_correctly():
    """Bộ mẫu cần "Diện tích chữ" đúng tiếng Việt (mười hai, không phải "mườ hai")."""
    from services.word_engine import _so_thanh_chu

    assert _so_thanh_chu(12) == "mười hai"
    assert _so_thanh_chu(15) == "mười lăm"
    assert _so_thanh_chu(21) == "hai mươi mốt"
    assert _so_thanh_chu(25) == "hai mươi lăm"
    assert _so_thanh_chu(101) == "một trăm linh một"
    assert _so_thanh_chu(1005) == "một nghìn không trăm linh năm"
    assert _so_thanh_chu(1500) == "một nghìn năm trăm"


def test_export_date_tokens_pad_digits_and_read_mung():
    """Catalog §2: [Ngày] pad 2 số, [Ngày chữ] thêm "mùng" cho ngày < 10, [Năm] tồn tại."""
    mapping = build_template_mapping(_simple_export_case(), today=date(2026, 7, 3))
    assert mapping["[Ngày]"] == "03"
    assert mapping["[Tháng]"] == "07"
    assert mapping["[Năm]"] == "2026"
    assert mapping["[Ngày chữ]"] == "mùng ba"
    assert mapping["[Tháng chữ]"] == "bảy"

    mapping = build_template_mapping(_simple_export_case(), today=date(2026, 12, 15))
    assert mapping["[Ngày]"] == "15"
    assert mapping["[Ngày chữ]"] == "mười lăm"
    assert mapping["[Tháng chữ]"] == "mười hai"


def test_nam_sinh_and_nam_chet_print_year_only():
    """Token `[Năm sinh N]`/`[Năm chết N]` luôn in năm (YYYY), kể cả khi DB là ngày đủ."""
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    receiver = _person(2, "Nguyễn Văn Con", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    case = _make_case(owner, participants=[_participant(receiver, "Con")], properties=[_prop()])
    mapping = build_template_mapping(case)

    assert mapping["[Năm sinh 1]"] == "1950"
    # Người nhận rơi vào slot 3 theo layout hồ sơ thừa kế.
    assert mapping["[Năm sinh 3]"] == "1980"
    assert mapping["[Năm chết 1]"] == "2024"
    assert mapping["[Năm chết người 1]"] == "2024"
    # Ngày đầy đủ vẫn in đủ dd/mm/yyyy qua trường "Ngày ...".
    assert mapping["[Ngày chết người 1]"] == "03/02/2024"


def test_year_tokens_extract_iso_dates_from_saved_snapshot():
    owner = _person(1, "Chủ", born=date(1950, 4, 5), dead=date(2024, 2, 3))
    heir = _person(2, "Con", born=date(1980, 8, 9))
    saved = _case_state(owner, heir, nodes=[
        _node(1, role="Owner", is_land_owner=True),
        _node(2, role="Con", will_receive=True)])
    # Master changes later: Word must keep the persisted ISO dates.
    owner.ngay_sinh = date(1960, 6, 7)
    owner.ngay_chet = date(2025, 6, 7)
    case = _make_case(owner, participants=[_participant(heir, "Con")],
                      properties=[_prop()], case_state=saved)
    mapping = build_template_mapping(case)
    assert mapping["[Năm sinh 1]"] == "1950"
    assert mapping["[Năm sinh 3]"] == "1980"
    assert mapping["[Năm chết 1]"] == "2024"
    assert mapping["[Người 1 - Năm sinh]"] == "1950"
    assert mapping["[Ngày chết người 1]"] == "03/02/2024"


def test_replace_in_doc_returns_replacement_count():
    doc = Document()
    doc.add_paragraph("[Tên 1] và [Tên 1] lần nữa")
    doc.add_paragraph("[Chưa map]")
    assert replace_in_doc(doc, {"[Tên 1]": "X"}) == 2


def test_dien_tich_chu_exposes_asset_one_total_area_in_words_and_legacy_alias():
    """[Diện tích chữ] / [Diện tích chữ tài sản 1] = tổng diện tích đọc thành chữ."""
    case = _summary_export_case(
        properties=[_prop(land_rows_json='[{"loai_dat":"ONT","dien_tich":"80"}]')]
    )
    mapping = build_template_mapping(case)
    assert mapping["[Diện tích]"] == "80"
    assert mapping["[Diện tích chữ]"] == "Tám mươi"
    assert mapping["[Diện tích chữ tài sản 1]"] == "Tám mươi"


def test_dien_tich_chu_is_empty_when_asset_has_no_area_instead_of_zero():
    """Không có diện tích -> rỗng, không được ghi "Không" vào văn bản thật."""
    case = _summary_export_case(
        properties=[_prop(dien_tich=None, land_rows_json="[]", so_thua_dat="", so_to_ban_do="")]
    )
    assert build_template_mapping(case)["[Diện tích chữ]"] == ""


def test_authorized_person_group_only_uses_confirmed_diagram_role():
    """"Người được ủy quyền" chỉ lấy từ vai trò đã xác nhận, không suy từ vị trí dòng."""
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    proxy = _person(2, "Nguyễn Văn Ủy", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    case = _make_case(
        owner,
        participants=[_participant(proxy, "Khác")],
        properties=[_prop()],
        case_state=_case_state(
            owner, proxy,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Người được ủy quyền", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)
    assert mapping["[Họ tên người được ủy quyền 1]"] == "Nguyễn Văn Ủy"
    assert mapping["[Danh sách người được ủy quyền]"] == "Nguyễn Văn Ủy"


def test_authorized_person_group_stays_empty_without_confirmed_role():
    """Không có vai trò được ủy quyền -> rỗng, tuyệt đối không mượn người nhận tài sản."""
    mapping = build_template_mapping(_summary_export_case())
    assert mapping["[Họ tên người được ủy quyền 1]"] == ""
    assert mapping["[Danh sách người được ủy quyền]"] == ""


def test_authorized_person_group_orders_suffixed_roles():
    """Role "Người được ủy quyền2" phải được nhận diện và thay bằng người đầu tiên."""
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    first = _person(2, "Nguyễn Văn Một", gender="Nam", born=date(1975, 1, 1), doc_no="001070000002")
    case = _make_case(
        owner,
        participants=[_participant(first, "Khác")],
        properties=[_prop()],
        case_state=_case_state(
            owner, first,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Người được ủy quyền2", will_receive=True),
            ],
        ),
    )
    assert build_template_mapping(case)["[Họ tên người được ủy quyền 1]"] == "Nguyễn Văn Một"


def test_authorized_role_merged_from_persisted_participant_vai_tro():
    """Node diagram v3 KHÔNG lưu role tự do — vai trò "người được ủy quyền" đã gán
    trên participant (cột vai_tro, lưu DB) phải được merge lại vào Stage person."""
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    proxy = _person(2, "Nguyễn Văn Ủy", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    case = _make_case(
        owner,
        participants=[_participant(proxy, "Người được ủy quyền")],
        properties=[_prop()],
        case_state=_case_state(
            owner, proxy,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                # Node không mang role ủy quyền — chỉ participant lưu vai_tro.
                _node(2, role="Khac", will_receive=True),
            ],
        ),
    )
    context = build_word_context(case)
    assert [p.ho_ten for p in context.authorized_people] == ["Nguyễn Văn Ủy"]
    mapping = build_template_mapping(case)
    assert mapping["[Họ tên người được ủy quyền 1]"] == "Nguyễn Văn Ủy"


def test_authorized_person_resolves_from_saved_role_without_diagram_node():
    """Đại diện ủy quyền hợp lệ có thể KHÔNG có node trên sơ đồ thừa kế —
    vai trò lưu ở participants.vai_tro vẫn phải resolve, không bị lọc mất
    khỏi _active_persons (lỗi Leader tái hiện trên _synthetic_case NAIA-12)."""
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    receiver = _person(2, "Nguyễn Văn Con", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    proxy = _person(3, "Trần Thị Ủy", gender="Nữ", born=date(1975, 5, 5), doc_no="001075000005")
    case = _make_case(
        owner,
        participants=[
            _participant(receiver, "Con"),
            _participant(proxy, "Người được ủy quyền", receive=False),
        ],
        properties=[_prop()],
        case_state=_case_state(
            owner, receiver, proxy,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
                # proxy cố tình KHÔNG có node — chỉ participant lưu vai_tro.
            ],
        ),
    )
    context = build_word_context(case)
    assert [p.ho_ten for p in context.authorized_people] == ["Trần Thị Ủy"]
    mapping = build_template_mapping(case)
    assert mapping["[Họ tên người được ủy quyền 1]"] == "Trần Thị Ủy"
    assert mapping["[Danh sách người được ủy quyền]"] == "Trần Thị Ủy"
    assert "Trần Thị Ủy" in mapping["[Dòng người được ủy quyền 1]"]
    assert "001075000005" in mapping["[Dòng người được ủy quyền 1]"]


def test_authorized_person_hidden_or_deleted_node_stays_excluded():
    """Participant có vai_tro ủy quyền nhưng node đã bị ẩn/xoá khỏi Diagram ->
    không tính (đã bị gỡ khỏi sơ đồ)."""
    owner = _person(1, "Nguyễn Văn Chủ", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001050000001")
    receiver = _person(2, "Nguyễn Văn Con", gender="Nam", born=date(1980, 8, 9), doc_no="001080000003")
    proxy = _person(3, "Trần Thị Ủy", gender="Nữ", doc_no="001075000005")
    case = _make_case(
        owner,
        participants=[
            _participant(receiver, "Con"),
            _participant(proxy, "Người được ủy quyền", receive=False),
        ],
        properties=[_prop()],
        case_state=_case_state(
            owner, receiver, proxy,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
                {**_node(3, role="Khac"), "deleted": True},
            ],
        ),
    )
    assert build_word_context(case).authorized_people == []
    assert build_template_mapping(case)["[Họ tên người được ủy quyền 1]"] == ""


def test_required_authorized_slots_reads_template_demand():
    """Mẫu dùng slot N nào của nhóm ủy quyền thì báo đúng slot đó."""
    from services.word_engine import required_authorized_slots

    doc = Document()
    doc.add_paragraph("tôi uỷ quyền cho: [Dòng người được ủy quyền 1]")
    doc.add_paragraph("và [Họ tên người được ủy quyền 2]")
    assert required_authorized_slots(doc) == {1, 2}

    doc = Document()
    doc.add_paragraph("danh sách: [Danh sách người được ủy quyền]")
    assert required_authorized_slots(doc) == {1}

    doc = Document()
    doc.add_paragraph("không liên quan [Người nhận 1 - Họ tên]")
    assert required_authorized_slots(doc) == set()


def test_required_authorized_slots_inline_and_every_registry_form():
    """Phản ví dụ review NAIA-12: [Người được ủy quyền inline 2] phải ra slot
    2, KHÔNG trôi thành slot 1. Mọi dạng registry đều đọc đúng slot."""
    from services.word_engine import required_authorized_slots

    doc = Document()
    doc.add_paragraph("đại diện: [Người được ủy quyền inline 2]")
    assert required_authorized_slots(doc) == {2}

    doc = Document()
    doc.add_paragraph("[Người được ủy quyền 2 - Họ tên]")
    doc.add_paragraph("[Người được ủy quyền inline 3]")
    doc.add_paragraph("[Dòng người được ủy quyền 4]")
    doc.add_paragraph("[Địa chỉ người được ủy quyền 5]")
    assert required_authorized_slots(doc) == {2, 3, 4, 5}

    # Không hậu tố -> chỉ cần người 1; "inline" không số cũng vậy.
    doc = Document()
    doc.add_paragraph("[Danh sách người được ủy quyền]")
    doc.add_paragraph("[Danh sách người được ủy quyền inline]")
    doc.add_paragraph("[Họ tên người được ủy quyền]")
    assert required_authorized_slots(doc) == {1}

    # Token khác nhóm (người nhận/chủ đất có cụm 'uy quyen' ở field) không tính.
    doc = Document()
    doc.add_paragraph("[Người ủy quyền] và [Người nhận 2 - Họ tên]")
    assert required_authorized_slots(doc) == set()


def test_noi_niem_yet_uses_saved_value_only_no_property_address_fallback():
    """[Nơi niêm yết] chỉ lấy trường người dùng nhập — thiếu thì rỗng,
    TUYỆT ĐỐI không mượn địa chỉ tài sản để biến ca thiếu dữ liệu thành PASS."""
    case = _summary_export_case()
    case.noi_niem_yet = None  # người dùng chưa nhập nơi niêm yết

    mapping = build_template_mapping(case)
    assert mapping["[Nơi niêm yết]"] == ""
    assert mapping["[Niêm Yết]"] == ""
    assert mapping["[NIÊM YẾT]"] == ""
    # Bằng chứng ngược: địa chỉ tài sản có sẵn nhưng không bị dùng làm fallback.
    assert mapping["[Địa chỉ đất]"] == "Thửa đất tại xã Minh Tân"

    case.noi_niem_yet = "Trụ sở UBND xã Đại An"
    mapping = build_template_mapping(case)
    assert mapping["[Nơi niêm yết]"] == "Trụ sở UBND xã Đại An"
    assert mapping["[Niêm Yết]"] == "Trụ sở UBND xã Đại An"
    assert mapping["[NIÊM YẾT]"] == "TRỤ SỞ UBND XÃ ĐẠI AN"


def test_land_type_tokens_follow_row_position_not_land_code():
    """[ONT]/[CLN]/[NTS]/[LUC] là diện tích DÒNG 1..4, không tra theo mã loại đất.

    Quyết định Owner (review NAIA-12): mẫu DK18 in cứng nhãn "Loại đất 1: ONT"
    rồi đặt "[ONT]" làm diện tích dòng đó — nên dòng 1 lưu CLN/120 thì "[ONT]"
    phải ra "120", không phải rỗng vì "không có dòng ONT".
    """
    case = _summary_export_case(
        properties=[
            _prop(land_rows_json='[{"loai_dat":"CLN","dien_tich":"120"},'
                                 '{"loai_dat":"ONT","dien_tich":"80"},'
                                 '{"loai_dat":"LUC","dien_tich":"30"}]'),
            _prop(land_rows_json='[{"loai_dat":"NTS","dien_tich":"55"}]'),
        ]
    )
    mapping = build_template_mapping(case)
    # Tài sản 1: token = vị trí dòng, kể cả khi mã khớp dòng khác.
    assert mapping["[ONT]"] == "120"   # dòng 1 (CLN)
    assert mapping["[CLN]"] == "80"    # dòng 2 (ONT)
    assert mapping["[NTS]"] == "30"    # dòng 3 (LUC)
    assert mapping["[LUC]"] == ""      # dòng 4 vắng
    # Tài sản 2: hậu tố " 2" là chỉ số tài sản, vẫn theo vị trí dòng.
    assert mapping["[ONT 2]"] == "55"  # tài sản 2, dòng 1 (NTS)
    assert mapping["[CLN 2]"] == ""    # tài sản 2, dòng 2 vắng


def test_land_type_row_position_absent_stays_empty_without_blocking_export():
    """Vị trí dòng không có trong land_rows -> rỗng, vẫn không còn placeholder sót."""
    case = _summary_export_case(
        properties=[_prop(land_rows_json='[{"loai_dat":"CLN","dien_tich":"900"}]')]
    )
    mapping = build_template_mapping(case)
    assert mapping["[ONT]"] == "900"   # dòng 1 (CLN)
    assert mapping["[CLN]"] == ""      # dòng 2 vắng

    doc = Document()
    doc.add_paragraph("Diện tích: [ONT] m2; [CLN] m2")
    replace_in_doc(doc, mapping)

    assert find_unresolved_placeholders(doc) == []
    assert doc.paragraphs[0].text == "Diện tích: 900 m2;  m2"


def test_asset_two_token_fails_loudly_when_case_has_one_asset():
    """Mẫu của tài sản 2 chỉ xuất được khi hồ sơ có tài sản 2 (không im lặng bỏ qua)."""
    doc = Document()
    doc.add_paragraph("Diện tích: [ONT 2] m2")
    replace_in_doc(doc, build_template_mapping(_summary_export_case()))

    assert find_unresolved_placeholders(doc) == ["[ONT 2]"]


def _all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def _textbox_paragraph_xml():
    from docx.oxml import parse_xml

    return parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>[Trong textbox]</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )


def test_find_unresolved_placeholders_reports_only_unknown_tokens():
    doc = Document()
    doc.add_paragraph("[Đã map] [Chưa map]")

    replace_in_doc(doc, {"[Đã map]": "x"})

    assert find_unresolved_placeholders(doc) == ["[Chưa map]"]


def test_replace_in_doc_keeps_run_formatting_and_only_touches_token_runs():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Kính gửi: ").bold = True
    para.add_run("[Họ tên chủ đất]")
    para.add_run(" - nơi cư trú:").bold = True

    doc.add_paragraph().add_run("[Chưa map]")

    replace_in_doc(
        doc,
        {
            "[Họ tên chủ đất]": "Nguyễn Văn Chủ",
            "[Chưa map]": "x",
        },
    )

    assert doc.paragraphs[0].text == "Kính gửi: Nguyễn Văn Chủ - nơi cư trú:"
    # Ba run vẫn còn: không hợp nhất text vào run đầu rồi xóa các run sau.
    assert [run.text for run in doc.paragraphs[0].runs] == [
        "Kính gửi: ",
        "Nguyễn Văn Chủ",
        " - nơi cư trú:",
    ]
    assert doc.paragraphs[0].runs[0].bold is True
    assert doc.paragraphs[0].runs[1].bold is not True


def test_replace_in_doc_handles_token_split_across_runs_without_clobbering_siblings():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("[Họ tên")
    para.add_run(" chủ đất]")

    replace_in_doc(doc, {"[Họ tên chủ đất]": "Nguyễn Văn Chủ"})

    assert doc.paragraphs[0].text == "Nguyễn Văn Chủ"

    # Sibling runs nằm ngoài vùng token phải còn nguyên vẹn.
    doc2 = Document()
    para2 = doc2.add_paragraph()
    para2.add_run("A")
    para2.add_run("[Tên")
    para2.add_run(" 1]")
    para2.add_run("B")
    replace_in_doc(doc2, {"[Tên 1]": "X"})
    assert [run.text for run in doc2.paragraphs[0].runs] == ["A", "X", "", "B"]
    assert doc2.paragraphs[0].text == "AXB"


def test_find_unresolved_placeholders_scans_nested_tables_headers_textboxes():
    doc = Document()
    header_para = doc.sections[0].header.paragraphs[0]
    header_para.add_run("[Tháng] ")
    header_para.add_run("[Năm]")

    first_page = doc.sections[0].first_page_header.paragraphs[0]
    first_page.add_run("[Số trang đầu]")

    table = doc.add_table(rows=1, cols=1)
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "[Số vào sổ]"

    doc.element.body.append(_textbox_paragraph_xml())

    assert find_unresolved_placeholders(doc) == [
        "[Năm]",
        "[Số trang đầu]",
        "[Số vào sổ]",
        "[Tháng]",
        "[Trong textbox]",
    ]


def test_replace_in_doc_reaches_nested_tables_and_textboxes():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "[Số vào sổ]"
    doc.element.body.append(_textbox_paragraph_xml())

    replace_in_doc(doc, {"[Số vào sổ]": "123", "[Trong textbox]": "đã thay"})

    assert find_unresolved_placeholders(doc) == []
    assert "đã thay" in "".join(t.text for t in doc.element.body.iter(TEXT_TAG))


def test_replace_in_doc_resolves_placeholder_nested_in_replacement_value():
    doc = Document()
    doc.add_paragraph("[Dòng khai tử]")

    replace_in_doc(
        doc,
        {"[Dòng khai tử]": "do xã [Nơi niêm yết] ký", "[Nơi niêm yết]": "Minh Tân"},
    )

    assert doc.paragraphs[0].text == "do xã Minh Tân ký"
    assert find_unresolved_placeholders(doc) == []


def test_split_nested_replacement_uses_word_breaks_and_preserves_spaces(tmp_path):
    from docx.oxml.ns import qn

    doc = Document()
    para = doc.add_paragraph()
    sibling = para.add_run("Không sửa")
    sibling.italic = True
    sibling_xml = sibling._r.xml
    first = para.add_run("[Đoạn")
    first.bold = True
    last = para.add_run(" dài]")
    last.underline = True
    first_properties = first._r.rPr.xml
    last_properties = last._r.rPr.xml
    replace_in_doc(doc, {"[Đoạn dài]": " [Tên]\r\n dòng hai \n", "[Tên]": "An "})
    path = tmp_path / "multiline.docx"
    doc.save(path)
    result = Document(path).paragraphs[0]
    assert result.text == "Không sửa An \n dòng hai \n"
    assert len(list(result._p.iter(qn("w:br")))) == 2
    texts = list(result.runs[1]._r.iter(qn("w:t")))
    assert [t.text or "" for t in texts] == [" An ", " dòng hai ", ""]
    assert all(t.get(qn("xml:space")) == "preserve" for t in texts[:2])
    assert all("\n" not in (t.text or "") for t in result._p.iter(qn("w:t")))
    assert result.runs[0]._r.xml == sibling_xml
    assert result.runs[1]._r.rPr.xml == first_properties
    assert result.runs[2]._r.rPr.xml == last_properties
    assert result.runs[2].text == ""
    assert find_unresolved_placeholders(Document(path)) == []


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _write_docx_with_notes(path, *, footnote_token=None, endnote_token=None):
    """Tạo DOCX THẬT có footnotes.xml/endnotes.xml, save + reopen qua python-docx.

    python-docx không model hai part này (part chỉ có ``blob``, không có
    ``.element``) — test này chứng minh engine đọc/ghi ngược blob được, và file
    sau thay vẫn mở lại được như DOCX hợp lệ.
    """
    import zipfile
    from xml.sax.saxutils import escape

    doc = Document()
    doc.add_paragraph("Thân văn bản chính.")
    doc.save(str(path))

    def note_xml(root, marker_type, token):
        return (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:{root} xmlns:w="{_W_NS}">'
            f'<w:{marker_type} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{marker_type}>'
            f'<w:{marker_type} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{marker_type}>'
            f'<w:{marker_type} w:id="1"><w:p><w:pPr><w:pStyle w:val="{"Foot" if marker_type == "footnote" else "End"}noteText"/></w:pPr>'
            f'<w:r><w:rPr><w:rStyle w:val="{"Foot" if marker_type == "footnote" else "End"}noteReference"/></w:rPr><w:{marker_type}Ref/></w:r>'
            f'<w:r><w:t xml:space="preserve"> {escape(token)}</w:t></w:r></w:p></w:{marker_type}>'
            f'</w:{root}>'
        )

    with zipfile.ZipFile(path, "r") as zin:
        entries = {item.filename: zin.read(item.filename) for item in zin.infolist()}

    overrides = entries["[Content_Types].xml"].decode("utf-8")
    for part, ctype in (
        ("footnotes", "wordprocessingml.footnotes+xml"),
        ("endnotes", "wordprocessingml.endnotes+xml"),
    ):
        overrides = overrides.replace(
            "</Types>",
            f'<Override PartName="/word/{part}.xml" '
            f'ContentType="application/vnd.openxmlformats-officedocument.{ctype}"/></Types>',
        )
    entries["[Content_Types].xml"] = overrides.encode("utf-8")

    rels = entries["word/_rels/document.xml.rels"].decode("utf-8")
    rel_add = (
        '<Relationship Id="rIdFt" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
        '<Relationship Id="rIdEn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes" Target="endnotes.xml"/>'
    )
    rels = rels.replace("</Relationships>", rel_add + "</Relationships>")
    entries["word/_rels/document.xml.rels"] = rels.encode("utf-8")

    entries["word/footnotes.xml"] = note_xml("footnotes", "footnote", footnote_token or "").encode("utf-8")
    entries["word/endnotes.xml"] = note_xml("endnotes", "endnote", endnote_token or "").encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)


def _part_blob(doc, part_name):
    for part in doc.part.package.iter_parts():
        if str(part.partname) == part_name:
            return part.blob
    return None


def test_replace_in_doc_reaches_real_footnote_and_endnote_parts(tmp_path):
    """DOCX thật có footnote/endnote: token trong hai part đó phải được thay,
    file save + mở lại vẫn giữ kết quả (python-docx chỉ đọc blob của chúng)."""
    path = tmp_path / "notes.docx"
    _write_docx_with_notes(
        path, footnote_token="ghi chú chân [Họ tên chủ đất]",
        endnote_token="ghi chú cuối [Số vào sổ]")

    doc = Document(str(path))
    # Tiền đề: python-docx không gắn .element cho footnotes/endnotes của file thật.
    note_parts = [
        p for p in doc.part.package.iter_parts()
        if str(p.partname) in ("/word/footnotes.xml", "/word/endnotes.xml")]
    assert len(note_parts) == 2
    assert all(getattr(p, "element", None) is None for p in note_parts)
    assert "[Họ tên chủ đất]" in find_unresolved_placeholders(doc)

    replace_in_doc(doc, {"[Họ tên chủ đất]": "Nguyễn Văn Chủ", "[Số vào sổ]": "CS123"})
    doc.save(str(path))

    # Mở lại file đã save: kết quả phải sống trong blob part, không chỉ trong RAM.
    reloaded = Document(str(path))
    assert "Nguyễn Văn Chủ" in _part_blob(reloaded, "/word/footnotes.xml").decode("utf-8")
    assert "CS123" in _part_blob(reloaded, "/word/endnotes.xml").decode("utf-8")
    assert find_unresolved_placeholders(reloaded) == []


def test_unresolved_placeholder_in_real_footnote_blocks_export(tmp_path):
    """Token sót nằm trong footnote của file thật phải vẫn bị báo — không bị
    bỏ qua chỉ vì part đó không có .element."""
    path = tmp_path / "only_note.docx"
    _write_docx_with_notes(path, footnote_token="xem [Chưa map]")

    doc = Document(str(path))
    assert find_unresolved_placeholders(doc) == ["[Chưa map]"]


@pytest.mark.parametrize("part", ["footnotes", "endnotes"])
def test_malformed_real_note_xml_blocks_export(tmp_path, part):
    import zipfile
    from services import word_batch_export as wbe

    path = tmp_path / "malformed.docx"
    _write_docx_with_notes(path)
    with zipfile.ZipFile(path) as archive:
        entries = {n: archive.read(n) for n in archive.namelist()}
    entries[f"word/{part}.xml"] = b'<w:broken xmlns:w="urn:broken">'
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    with pytest.raises(WordExportValidationError, match=part + r"\.xml"):
        replace_in_doc(Document(path), {})
    with pytest.raises(WordExportValidationError, match=part + r"\.xml"):
        find_unresolved_placeholders(Document(path))
    dest = tmp_path / "out"
    dest.mkdir()
    with pytest.raises(wbe.WordBatchError) as failure:
        wbe.export_batch(_summary_export_case(), case_id=99,
            document_keys=["khai_nhan_di_san"], dest_dir=dest,
            resolve_template=lambda _key: path)
    assert failure.value.code == "word_batch_failed"
    result = failure.value.result_data
    assert result["documents"][0]["status"] == "failed"
    assert part + ".xml" in result["documents"][0]["error"]["message"]
    assert list(dest.iterdir()) == []


def _field_paragraph_xml():
    from docx.oxml import parse_xml

    return parse_xml(
        f'<w:p xmlns:w="{_W_NS}">'
        '<w:fldSimple w:instr=" DOCVARIABLE  ho_ten  \\* MERGEFORMAT ">'
        "<w:r><w:rPr><w:b/></w:rPr><w:t>[Họ tên chủ đất]</w:t></w:r>"
        "</w:fldSimple>"
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        "<w:r><w:instrText xml:space=\"preserve\"> DOCVARIABLE so </w:instrText></w:r>"
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>[Số vào sổ]</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        "</w:p>"
    )


def test_replace_in_doc_reaches_field_result_runs(tmp_path):
    """Token nằm trong kết quả field (fldSimple + fldChar complex) phải được thay;
    cấu trúc field (instr) và định dạng run giữ nguyên."""
    from docx.oxml.ns import qn

    doc = Document()
    doc.element.body.append(_field_paragraph_xml())

    assert find_unresolved_placeholders(doc) == ["[Họ tên chủ đất]", "[Số vào sổ]"]

    replace_in_doc(doc, {"[Họ tên chủ đất]": "Nguyễn Văn Chủ", "[Số vào sổ]": "CS123"})

    assert find_unresolved_placeholders(doc) == []
    para = doc.paragraphs[0]
    # fldSimple bọc run con — ghép mọi w:t để kiểm cả kết quả field.
    all_text = "".join(t.text or "" for t in para._element.iter(TEXT_TAG))
    assert "Nguyễn Văn Chủ" in all_text and "CS123" in all_text
    # Cấu trúc field còn nguyên sau khi thay.
    assert para._element.find(qn("w:fldSimple")) is not None
    assert b"DOCVARIABLE" in etree_tostring(para._element)

    path = tmp_path / "field.docx"
    doc.save(str(path))
    reloaded = Document(str(path))
    assert find_unresolved_placeholders(reloaded) == []


def _nested_textbox_paragraph_xml():
    from docx.oxml import parse_xml

    return parse_xml(
        f'<w:p xmlns:w="{_W_NS}" xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>[Textbox ngoài]</w:t></w:r>"
        "<w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>[Textbox trong]</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
        "</w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )


def test_replace_in_doc_reaches_textbox_nested_inside_textbox(tmp_path):
    """Textbox LỒNG trong textbox: token tầng sâu nhất vẫn được thay/đếm đúng."""
    doc = Document()
    doc.element.body.append(_nested_textbox_paragraph_xml())

    assert find_unresolved_placeholders(doc) == ["[Textbox ngoài]", "[Textbox trong]"]

    replace_in_doc(doc, {"[Textbox ngoài]": "NGOÀI", "[Textbox trong]": "TRONG"})

    assert find_unresolved_placeholders(doc) == []
    text = "".join(t.text for t in doc.element.body.iter(TEXT_TAG))
    assert "NGOÀI" in text and "TRONG" in text

    path = tmp_path / "nested_txbx.docx"
    doc.save(str(path))
    assert find_unresolved_placeholders(Document(str(path))) == []


def test_split_token_around_textbox_preserves_paragraph_isolation(tmp_path):
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    from services.word_engine import collect_placeholder_occurrences

    doc = Document()
    doc.element.body.append(parse_xml(
        f'<w:p xmlns:w="{_W_NS}" xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:r><w:rPr><w:b/></w:rPr><w:t>[Tên </w:t></w:r>"
        "<w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:rPr><w:i/></w:rPr><w:t>[Năm sinh 1]</w:t>"
        "</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
        "<w:r><w:t>1]</w:t></w:r></w:p>"
    ))
    source = tmp_path / "source.docx"
    doc.save(source)
    doc = Document(source)
    assert collect_placeholder_occurrences(doc) == ["[Tên 1]", "[Năm sinh 1]"]
    assert replace_in_doc(
        doc, {"[Tên 1]": "NGƯỜI TỔNG HỢP", "[Năm sinh 1]": "1980"}
    ) == 2

    output = tmp_path / "output.docx"
    doc.save(output)
    saved = Document(output)
    outer = saved.element.body.find(qn("w:p"))
    assert [node.text or "" for node in outer.iter(TEXT_TAG)] == [
        "NGƯỜI TỔNG HỢP", "1980", "",
    ]
    assert outer.find("w:r/w:rPr/w:b", outer.nsmap) is not None
    inner = outer.find(".//w:txbxContent/w:p", outer.nsmap)
    assert [node.text for node in inner.iter(TEXT_TAG)] == ["1980"]
    assert inner.find("w:r/w:rPr/w:i", inner.nsmap) is not None
    assert collect_placeholder_occurrences(saved) == []
    assert collect_placeholder_occurrences(Document(source)) == [
        "[Tên 1]", "[Năm sinh 1]",
    ]


def etree_tostring(element):
    from lxml import etree

    return etree.tostring(element)


def test_fixed_codes_on_unverified_template_are_unresolved():
    """Mẫu KHÔNG nằm trong danh sách xác minh: "[01]"/"[54.99]" vẫn là token sót.

    Quy tắc mã số KHÔNG còn toàn cục (review NAIA-12): file giả bất kỳ chứa
    token dạng mã biểu phải bị chặn, không được miễn ngầm.
    """
    doc = Document()
    doc.add_paragraph("[01] [02] [06.1] [10] [54.99] [Đã map]")

    replace_in_doc(doc, {"[Đã map]": "x"})

    assert find_unresolved_placeholders(doc) == [
        "[01]", "[02]", "[06.1]", "[10]", "[54.99]"]
    # Không truyền template_ref tương đương mẫu chưa xác minh.
    assert find_unresolved_placeholders(doc, template_ref=None) == [
        "[01]", "[02]", "[06.1]", "[10]", "[54.99]"]
    assert find_unresolved_placeholders(doc, template_ref="mau khac.docx") == [
        "[01]", "[02]", "[06.1]", "[10]", "[54.99]"]


def test_fixed_form_code_shape_is_narrow_but_still_template_scoped():
    """Dạng không khớp mã biểu luôn là sót; dạng khớp nhưng mẫu chưa xác minh cũng sót."""
    doc = Document()
    doc.add_paragraph("[6] [N6] [Tên 6] [Ngườ 6] [99]")

    replace_in_doc(doc, {"x": "y"})

    assert find_unresolved_placeholders(doc) == ["[6]", "[99]", "[N6]", "[Ngườ 6]", "[Tên 6]"]
    # "[99]" có dạng mã biểu: được miễn chỉ trên mẫu DK18 đã xác minh chứa nó?
    # [99] không nằm trong bộ mã DK18 nên vẫn sót kể cả trên mẫu DK18.
    assert "[99]" in find_unresolved_placeholders(
        doc, template_ref="2. DK18, thuế  1.docx")


def test_verified_template_keeps_fixed_codes_and_still_flags_real_placeholders():
    doc = Document()
    doc.add_paragraph("Mục [01] Họ tên: [Họ tên chủ đất]")
    doc.add_paragraph("Mục [06.1] Diện tích: [Diện tích đất]")

    replace_in_doc(doc, {"[Họ tên chủ đất]": "Nguyễn Văn A", "[Diện tích đất]": "100 m2"})

    # Mẫu DK18 đã xác minh: [01]/[06.1] là nhãn in cứng, không báo sót.
    assert find_unresolved_placeholders(
        doc, template_ref="2. DK18, thuế  1.docx") == []
    # Mã NGOÀI danh sách xác minh của mẫu vẫn bị chặn.
    doc2 = Document()
    doc2.add_paragraph("Mục [54.99] ngoài bộ DK18")
    assert find_unresolved_placeholders(
        doc2, template_ref="2. DK18, thuế  1.docx") == ["[54.99]"]

    full_text = _all_text(doc)
    assert "[01]" in full_text and "[06.1]" in full_text




def test_list_public_builtin_templates_hides_reference_and_snake_case_templates(tmp_path):
    (tmp_path / "xa_PCDS_template.docx").write_bytes(b"doc")
    (tmp_path / "system_placeholder_reference.docx").write_bytes(b"doc")
    (tmp_path / "system_template_chuan_v1.docx").write_bytes(b"doc")
    (tmp_path / "~$ PCDS .docx").write_bytes(b"doc")

    items = list_public_builtin_templates(tmp_path)

    assert [item["id"] for item in items] == ["builtin:xa_PCDS_template.docx"]


def test_placeholder_catalog_documents_vietnamese_contract_without_snake_case_public_contract():
    catalog = (Path(__file__).resolve().parent.parent
               / "word_templates" / "placeholder_mapping.md").read_text(encoding="utf-8")

    assert "[Người N - Họ tên]" in catalog
    assert "[Người nhận N - Họ tên]" in catalog
    assert "[Dòng người từ chối N]" in catalog
    assert "người chưa chọn" not in catalog.lower()
    assert "[ho_ten_1_chu_dat]" not in catalog


# ---------------------------------------------------------------------------
# Word export V2 tests
# ---------------------------------------------------------------------------


def test_one_deceased_landowner_spouse_alive_not_landowner():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001")
    spouse = _person(2, "Trần Thị B", gender="Nữ", born=date(1955, 4, 5), doc_no="002")
    child = _person(3, "Nguyễn Văn C", gender="Nam", born=date(1980, 8, 9), doc_no="003")
    case = _make_case(
        owner,
        properties=[_prop()],
        case_state=_case_state(
            owner, spouse, child,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Vợ/Chồng"),
                _node(3, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)
    assert mapping["[Cụm chủ đất chết]"] == "ông Nguyễn Văn A"
    assert mapping["[Người nhận 1 - Họ tên]"] == "Nguyễn Văn C"
    assert mapping["[Người từ chối 1 - Họ tên]"] == ""
    assert mapping["[Người 2 - Trạng thái]"] == "Người không nhận"


def test_word_context_reads_v2_engine_input_and_exact_engine_result_shares():
    owner = _person(1, "X", dead=date(2024, 2, 3))
    receiver = _person(2, "M")
    case = _make_case(
        owner,
        properties=[_prop()],
        case_state=_case_state_v2(
            owner,
            receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
            allocations={
                "1": {"finalShare": "0"},
                "2": {"finalShare": "1/3"},
            },
        ),
    )

    context = build_word_context(case)
    mapping = build_template_mapping(case)

    assert [person.ho_ten for person in context.receivers] == ["M"]
    assert context.receivers[0].share == "1/3"
    assert mapping["[Người nhận 1 - Tỷ lệ]"] == "1/3"


def test_two_deceased_landowners_both_in_estate_cluster():
    owner1 = _person(1, "Nguyễn Văn A", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001")
    owner2 = _person(2, "Trần Thị B", gender="Nữ", born=date(1952, 3, 4), dead=date(2024, 5, 6), doc_no="002")
    child = _person(3, "Nguyễn Văn C", gender="Nam", born=date(1980, 8, 9), doc_no="003")
    case = _make_case(
        owner1,
        properties=[_prop()],
        case_state=_case_state(
            owner1, owner2, child,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Vợ/Chồng", is_land_owner=True),
                _node(3, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)
    assert mapping["[Chủ đất chết 1 - Họ tên]"] == "Nguyễn Văn A"
    assert mapping["[Chủ đất chết 2 - Họ tên]"] == "Trần Thị B"
    assert mapping["[Nối chủ đất chết 2]"] == " và "


def test_living_landowner_nonreceiver_keeps_base_without_implicit_gift():
    owner1 = _person(1, "Nguyễn Văn A", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001")
    owner2 = _person(2, "Trần Thị B", gender="Nữ", born=date(1952, 3, 4), doc_no="002")
    child = _person(3, "Nguyễn Văn C", gender="Nam", born=date(1980, 8, 9), doc_no="003")
    case = _make_case(
        owner1,
        properties=[_prop()],
        case_state=_case_state(
            owner1, owner2, child,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Vợ/Chồng", is_land_owner=True),
                _node(3, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)
    assert mapping["[Cụm chủ đất chết]"] == "ông Nguyễn Văn A"
    assert mapping["[Chủ đất sống 1 - Họ tên]"] == "Trần Thị B"
    assert mapping["[Người từ chối 1 - Họ tên]"] == ""
    assert "Trần Thị B" not in mapping["[Đoạn phân chia di sản]"]
    assert "tặng cho" not in mapping["[Đoạn phân chia di sản]"]
    assert mapping["[Dòng chủ đất sống tặng cho 1]"] == ""


def test_nonreceiver_does_not_create_a_legal_refusal_statement():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001")
    child_accept = _person(2, "Nguyễn Văn B", gender="Nam", born=date(1980, 8, 9), doc_no="002")
    child_refuse = _person(3, "Nguyễn Thị C", gender="Nữ", born=date(1982, 10, 11), doc_no="003")
    case = _make_case(
        owner,
        properties=[_prop()],
        case_state=_case_state(
            owner, child_accept, child_refuse,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
                _node(3, role="Con"),
            ],
        ),
    )
    mapping = build_template_mapping(case)
    assert mapping["[Người từ chối 1 - Họ tên]"] == ""
    assert mapping["[Dòng người từ chối 1]"] == ""
    assert "Nguyễn Thị C" in mapping["[Đoạn quan hệ gia đình]"]
    assert "từ chối di sản" not in mapping["[Đoạn quan hệ gia đình]"]


def test_nonreceiver_is_not_in_division_clause_or_legal_refusal_group():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001")
    unset = _person(2, "Trần Thị B", gender="Nữ", born=date(1955, 4, 5), doc_no="002")
    receiver = _person(3, "Nguyễn Văn C", gender="Nam", born=date(1980, 8, 9), doc_no="003")
    case = _make_case(
        owner,
        properties=[_prop()],
        case_state=_case_state(
            owner, unset, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Vợ/Chồng"),
                _node(3, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)
    assert mapping["[Người từ chối 1 - Họ tên]"] == ""
    assert "Trần Thị B" not in mapping["[Đoạn phân chia di sản]"]
    assert "chưa chọn" not in mapping["[Đoạn phân chia di sản]"].lower()


def test_multiple_receivers_listed_in_accept_paragraph():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001")
    r1 = _person(2, "Nguyễn Văn B", gender="Nam", born=date(1980, 8, 9), doc_no="002")
    r2 = _person(3, "Nguyễn Thị C", gender="Nữ", born=date(1982, 10, 11), doc_no="003")
    case = _make_case(
        owner,
        properties=[_prop()],
        case_state=_case_state(
            owner, r1, r2,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
                _node(3, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)
    assert "Nguyễn Văn B" in mapping["[Đoạn phân chia di sản]"]
    assert "Nguyễn Thị C" in mapping["[Đoạn phân chia di sản]"]
    assert "đồng ý nhận phần di sản" in mapping["[Đoạn phân chia di sản]"]


def test_word_context_keeps_legal_refusal_empty_and_allows_owner_receiver_overlap():
    owner_dead = _person(1, "Nguyễn Văn A", dead=date(2024, 2, 3))
    owner_alive = _person(2, "Trần Thị B", gender="Nữ")
    receiver = _person(3, "Nguyễn Văn C")
    refused = _person(4, "Nguyễn Thị D", gender="Nữ")
    case = _make_case(
        owner_dead,
        properties=[_prop()],
        case_state=_case_state(
            owner_dead,
            owner_alive,
            receiver,
            refused,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Vợ/Chồng", is_land_owner=True, will_receive=True),
                _node(3, role="Con", will_receive=True),
                _node(4, role="Con"),
            ],
        ),
    )

    context = build_word_context(case)

    assert [person.ho_ten for person in context.landowners] == ["Nguyễn Văn A", "Trần Thị B"]
    assert [person.ho_ten for person in context.receivers] == ["Trần Thị B", "Nguyễn Văn C"]
    assert context.legal_refusal_people == []


def test_legacy_accept_is_read_only_when_will_receive_is_missing():
    owner = _person(1, "Nguyễn Văn A", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B")
    case = _make_case(
        owner,
        properties=[_prop()],
        case_state=_case_state(
            owner,
            receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=None, legacy_decision="accept"),
            ],
        ),
    )

    mapping = build_template_mapping(case)

    assert mapping["[Người nhận 1 - Họ tên]"] == "Nguyễn Văn B"


def test_empty_numbered_slots_clear_the_whole_tier_two_block():
    owner = _person(1, "Nguyễn Văn A", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B")
    case = _make_case(
        owner,
        properties=[_prop()],
        case_state=_case_state(
            owner,
            receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )

    mapping = build_template_mapping(case)

    assert mapping["[Người nhận 2 - Họ tên]"] == ""
    assert mapping["[Dòng người nhận 2]"] == ""
    assert mapping["[Nối chủ đất chết 2]"] == ""
    assert mapping["[Chủ đất chết 2 - Họ tên]"] == ""
    assert mapping["[Dòng tài sản 2]"] == ""


def test_five_properties_numbering():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", born=date(1950, 1, 1), dead=date(2024, 2, 3), doc_no="001")
    receiver = _person(2, "Nguyễn Văn B", gender="Nam", born=date(1980, 8, 9), doc_no="002")
    props = [
        _prop(dia_chi="Đất 1", so_serial="S1", land_rows_json='[{"loai_dat":"ONT","dien_tich":"80"},{"loai_dat":"CLN","dien_tich":"20"}]'),
        _prop(dia_chi="Đất 2", so_serial="S2", land_rows_json='[{"loai_dat":"NTS","dien_tich":"50"}]'),
        _prop(dia_chi="Đất 3", so_serial="S3"),
        _prop(dia_chi="Đất 4", so_serial="S4"),
        _prop(dia_chi="Đất 5", so_serial="S5"),
    ]
    case = _make_case(
        owner,
        properties=props,
        case_state=_case_state(
            owner, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)
    desc = mapping["[Đoạn mô tả di sản]"]
    assert desc.startswith("Các quyền sử dụng đất như sau:\n1. Quyền sử dụng đất tại: Đất 1")
    assert "1. Quyền sử dụng đất tại: Đất 1" in desc
    assert "1.1. ONT: 80" in desc
    assert "1.2. CLN: 20" in desc
    assert "2.1. NTS: 50" in desc
    assert mapping["[Serial tài sản 5]"] == "S5"
    assert mapping["[Tài sản 5 - Serial]"] == "S5"
    assert mapping["[Dòng tài sản 5]"].startswith("5. Quyền sử dụng đất tại: Đất 5")


def test_multiple_properties_follow_asset_row_order():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B", gender="Nam")
    primary = _prop(dia_chi="Tài sản chính", so_serial="PRIMARY")
    linked_first = _prop(dia_chi="Tài sản liên kết trước", so_serial="LINK-1")
    linked_second = _prop(dia_chi="Tài sản liên kết sau", so_serial="LINK-2")
    case = _make_case(
        owner,
        properties=[primary, linked_first, linked_second],
        case_state=_case_state(
            owner,
            receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    context = build_word_context(case)
    mapping = build_template_mapping(case)

    # Thứ tự tài sản = thứ tự dòng (vitri); link chỉ là projection.
    assert [asset.so_serial for asset in context.assets] == [
        "PRIMARY", "LINK-1", "LINK-2"]
    assert mapping["[Tài sản 1 - Serial]"] == "PRIMARY"
    assert mapping["[Tài sản 2 - Serial]"] == "LINK-1"
    assert mapping["[Tài sản 3 - Serial]"] == "LINK-2"


def test_real_pcds_template_resolves_two_properties_into_one_document():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B", gender="Nam")
    case = _make_case(
        owner,
        properties=[
            _prop(dia_chi="Thửa đất thứ nhất", so_serial="S1"),
            _prop(dia_chi="Thửa đất thứ hai", so_serial="S2"),
        ],
        case_state=_case_state(
            owner,
            receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    doc = Document(_PCDS_TEMPLATE)

    replace_in_doc(doc, build_template_mapping(case))

    assert find_unresolved_placeholders(doc) == []
    full_text = "\n".join(
        [paragraph.text for paragraph in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )
    assert "Các quyền sử dụng đất như sau:" in full_text
    assert "1. Quyền sử dụng đất tại: Thửa đất thứ nhất" in full_text
    assert "2. Quyền sử dụng đất tại: Thửa đất thứ hai" in full_text


def test_more_than_five_properties_raises():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B", gender="Nam")
    props = [_prop(dia_chi=f"Đất {i}", so_serial=f"S{i}") for i in range(1, 7)]
    case = _make_case(
        owner,
        properties=props,
        case_state=_case_state(
            owner, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    with pytest.raises(WordExportValidationError, match="5 tài sản"):
        build_template_mapping(case)


def test_more_than_twenty_heirs_raises():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    people = [owner]
    nodes = [_node(1, role="Owner", is_land_owner=True)]
    for i in range(2, 24):
        p = _person(i, f"Ngườ {i}", gender="Nam")
        people.append(p)
        nodes.append(_node(i, role="Con", will_receive=True))
    case = _make_case(
        owner,
        properties=[_prop()],
        case_state=_case_state(*people, nodes=nodes),
    )
    with pytest.raises(WordExportValidationError, match="20 người"):
        build_template_mapping(case)


def test_pcds_template_no_old_critical_placeholders():
    doc = Document(_PCDS_TEMPLATE)
    all_texts = []
    for p in doc.paragraphs:
        all_texts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                all_texts.append(cell.text)
    full = "\n".join(all_texts)
    import re
    placeholders = set(re.findall(r"\[[^\[\]]+\]", full))
    old_critical = {
        "[Tên 1]", "[Tên 2]", "[Tên 3]", "[ONT]", "[CLN]", "[NTS]",
        "[Cụm ngườ để lại di sản]",
        "[Danh sách ngườ nhận và chưa chọn]",
        "[Danh sách ngườ ký]",
    }
    assert not (old_critical & placeholders), f"Old placeholders remain: {old_critical & placeholders}"
    assert "[Chủ đất chết 1 - Xưng hô]" in placeholders
    assert "[Chủ đất chết 2 - Họ tên]" in placeholders
    assert "[Dòng người 20]" in placeholders
    assert "[Dòng người từ chối 20]" in placeholders
    assert "[Đoạn phân chia di sản]" in placeholders


# ---------------------------------------------------------------------------
# Canonical land placeholders: [loaidatMN]/[dientichMN]/[thoihanMN] (§9.2)
# ---------------------------------------------------------------------------


def _dense_land_rows(asset_idx, count=20):
    return [
        {
            "loai_dat": f"LD{m}-TS{asset_idx}",
            "dien_tich": f"{asset_idx * 100 + m}.5",
            "thoi_han": f"TH{m}-TS{asset_idx}",
        }
        for m in range(1, count + 1)
    ]


def _json_rows(rows):
    import json
    return json.dumps(rows, ensure_ascii=False)


def _three_asset_case():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B", gender="Nam")
    props = [
        _prop(dia_chi=f"Đất {n}", so_serial=f"S{n}", land_rows_json=_json_rows(_dense_land_rows(n)))
        for n in range(1, 4)
    ]
    return _make_case(
        owner,
        properties=props,
        case_state=_case_state(
            owner, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )


def test_canonical_land_names_cover_three_assets_and_twenty_rows():
    mapping = build_template_mapping(_three_asset_case())

    for n in range(1, 4):
        for m in range(1, 21):
            assert mapping[f"[loaidat{m}{n}]"] == f"LD{m}-TS{n}"
            assert mapping[f"[dientich{m}{n}]"] == f"{n * 100 + m}.5"
            assert mapping[f"[thoihan{m}{n}]"] == f"TH{m}-TS{n}"

    assert mapping["[loaidat12]"] == "LD1-TS2"
    assert mapping["[loaidat102]"] == "LD10-TS2"
    assert mapping["[loaidat203]"] == "LD20-TS3"

    for m in range(1, 21):
        assert mapping[f"[loaidat{m}4]"] == ""
        assert mapping[f"[dientich{m}4]"] == ""
        assert mapping[f"[thoihan{m}5]"] == ""


def test_canonical_land_names_single_asset_use_full_form_without_short_aliases():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B", gender="Nam")
    case = _make_case(
        owner,
        properties=[_prop(land_rows_json=_json_rows(_dense_land_rows(1)))],
        case_state=_case_state(
            owner, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)

    for m in range(1, 21):
        assert mapping[f"[loaidat{m}1]"] == f"LD{m}-TS1"
        assert mapping[f"[dientich{m}1]"] == f"{100 + m}.5"
        assert mapping[f"[thoihan{m}1]"] == f"TH{m}-TS1"

    # Không sinh alias rút gọn: [loaidat11] là tên đầy đủ loại đất 1/tài sản 1,
    # KHÔNG phải rút gọn của loại đất 11.
    assert mapping["[loaidat11]"] == "LD1-TS1"
    assert mapping["[dientich11]"] == "101.5"
    for m in (1, 2, 5, 9, 10, 16, 20):
        assert f"[loaidat{m}]" not in mapping
        assert f"[dientich{m}]" not in mapping
        assert f"[thoihan{m}]" not in mapping


def test_canonical_land_names_keep_position_when_a_row_is_empty():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B", gender="Nam")
    rows = [
        {"loai_dat": "ONT", "dien_tich": "80", "thoi_han": "Lâu dài"},
        {},
        {"loai_dat": "CLN", "dien_tich": "30", "thoi_han": "Đến 2043"},
    ]
    case = _make_case(
        owner,
        properties=[_prop(land_rows_json=_json_rows(rows))],
        case_state=_case_state(
            owner, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)

    assert mapping["[loaidat11]"] == "ONT"
    assert mapping["[loaidat21]"] == ""
    assert mapping["[dientich21]"] == ""
    assert mapping["[thoihan21]"] == ""
    assert mapping["[loaidat31]"] == "CLN"
    assert mapping["[dientich31]"] == "30"
    assert mapping["[thoihan31]"] == "Đến 2043"


def test_canonical_land_names_zero_area_and_missing_term():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B", gender="Nam")
    rows = [
        {"loai_dat": "ONT", "dien_tich": 0, "thoi_han": "Lâu dài"},
        {"loai_dat": "CLN", "dien_tich": "25"},
    ]
    case = _make_case(
        owner,
        properties=[_prop(land_rows_json=_json_rows(rows))],
        case_state=_case_state(
            owner, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)

    assert mapping["[dientich11]"] == "0"
    assert mapping["[thoihan11]"] == "Lâu dài"
    assert mapping["[loaidat21]"] == "CLN"
    assert mapping["[dientich21]"] == "25"
    assert mapping["[thoihan21]"] == ""


def test_canonical_land_names_resolve_through_real_doc_replacement():
    case = _three_asset_case()
    doc = Document()
    doc.add_paragraph("Đất: [loaidat12]; DT: [dientich102]; TH: [thoihan203].")
    doc.add_paragraph("[LoaiDat31] / [THOIHAN11] / [dientich201]")

    replace_in_doc(doc, build_template_mapping(case))

    assert doc.paragraphs[0].text == "Đất: LD1-TS2; DT: 210.5; TH: TH20-TS3."
    assert doc.paragraphs[1].text == "LD3-TS1 / TH1-TS1 / 120.5"
    assert find_unresolved_placeholders(doc) == []


def test_legacy_land_placeholders_unchanged_and_cover_assets_4_and_5():
    owner = _person(1, "Nguyễn Văn A", gender="Nam", dead=date(2024, 2, 3))
    receiver = _person(2, "Nguyễn Văn B", gender="Nam")
    props = [
        _prop(dia_chi=f"Đất {n}", so_serial=f"S{n}", land_rows_json=_json_rows(_dense_land_rows(n)))
        for n in range(1, 6)
    ]
    case = _make_case(
        owner,
        properties=props,
        case_state=_case_state(
            owner, receiver,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ],
        ),
    )
    mapping = build_template_mapping(case)

    # Legacy indexed keys (asset.row) unchanged, consistent values with canonical.
    assert mapping["[Loại đất 2.1 - Loại đất]"] == "LD1-TS2"
    assert mapping["[Loại đất tài sản 2.1]"] == "LD1-TS2"
    assert mapping["[Diện tích loại đất tài sản 2.1]"] == "201.5"
    assert mapping["[Dòng loại đất 1.2]"] == "1.2. LD2-TS1: 102.5 m2; Thời hạn: TH2-TS1."
    # Legacy primary-asset block keeps its 10-row cap.
    assert mapping["[Loại đất 1]"] == "LD1-TS1"
    assert mapping["[Loại đất 10]"] == "LD10-TS1"
    assert "[Loại đất 11]" not in mapping
    # Canonical names exist for assets 4-5 as well.
    assert mapping["[loaidat14]"] == "LD1-TS4"
    assert mapping["[dientich205]"] == "520.5"
    # Typo fix in rendered paragraph (consumer: [Đoạn mô tả di sản]).
    assert "Thời hạn sử dụng:" in mapping["[Đoạn mô tả di sản]"]

# ================================================================ NAIA-8
# `[Người N]` = thứ tự Stage đã lưu (ổn định khi đảo/xóa node Diagram);
# hồ sơ hai bên map Bên A/B theo slot p1..p30 (personId = row_id Stage).

import json
import uuid as _uuid

from services.word_engine import word_block_reason


def _stage_row_v3(row_id, entity_id, name, *, gender="Nam", dead=None):
    return {
        "row_id": row_id,
        "entity_id": entity_id,
        "id": str(entity_id) if entity_id is not None else row_id,
        "ho_ten": name,
        "gioi_tinh": gender,
        "ngay_sinh": None,
        "ngay_chet": dead.isoformat() if dead else None,
        "so_giay_to": "",
        "ngay_cap": None,
        "noi_cap": "",
        "dia_chi": "",
        "place_of_origin": "",
    }


def _case_state_v3(rows, *, domain="inheritance", nodes=None,
                   extra_diagram=None):
    diagram = {"state": {"version": 3, "domain": domain,
                         "nodes": nodes or []}}
    if extra_diagram:
        diagram.update(extra_diagram)
    return json.dumps(
        {"schemaVersion": 3, "case_type": domain, "stage": rows,
         "assets": [], "diagram": diagram},
        ensure_ascii=False)


def _two_party_state(rows, slots):
    """slots: {slot_number: row_id} — build canonical p1..p30 nodes."""
    nodes = []
    for i in range(1, 31):
        nodes.append({
            "id": f"p{i}",
            "personId": slots.get(i),
            "hidden": False,
            "deleted": False,
        })
    return _case_state_v3(rows, domain="two_party", nodes=nodes)


def _two_party_case(rows, slots, *, properties=None):
    case = _make_case(
        None,
        properties=properties if properties is not None else [_prop()],
        case_state=_two_party_state(rows, slots))
    case.loai_van_ban = "chuyen_nhuong"
    return case


def _rid():
    return str(_uuid.uuid4())


def test_nguoi_numbers_follow_stage_order_not_node_order():
    a = _person(1, "An Stage 1", dead=date(2024, 1, 1))
    b = _person(2, "Bích Stage 2")
    c = _person(3, "Cường Stage 3")
    rows = [_stage_row_v3(_rid(), 1, a.ho_ten, dead=date(2024, 1, 1)),
            _stage_row_v3(_rid(), 2, b.ho_ten),
            _stage_row_v3(_rid(), 3, c.ho_ten)]
    # Node order ngược Stage: C trước B trước A.
    nodes = [
        _node(3, role="Con", will_receive=True),
        _node(2, role="Con", will_receive=True),
        _node(1, role="Owner", is_land_owner=True),
    ]
    case = _make_case(
        a, properties=[_prop()],
        case_state=_case_state_v3(
            rows, nodes=nodes,
            extra_diagram={"engineInput": {"version": 2, "nodes": nodes},
                           "engineResult": {"status": "complete",
                                            "allocations": {}}}))
    mapping = build_template_mapping(case)
    assert mapping["[Người 1 - Họ tên]"] == "An Stage 1"
    assert mapping["[Người 2 - Họ tên]"] == "Bích Stage 2"
    assert mapping["[Người 3 - Họ tên]"] == "Cường Stage 3"


def test_removing_diagram_node_does_not_renumber_nguoi():
    a = _person(1, "An Stage 1", dead=date(2024, 1, 1))
    b = _person(2, "Bích Stage 2")
    c = _person(3, "Cường Stage 3")
    r1, r2, r3 = _rid(), _rid(), _rid()
    rows = [_stage_row_v3(r1, 1, a.ho_ten, dead=date(2024, 1, 1)),
            _stage_row_v3(r2, 2, b.ho_ten),
            _stage_row_v3(r3, 3, c.ho_ten)]
    # B không có node — số Người vẫn giữ theo Stage.
    nodes = [
        _node(1, role="Owner", is_land_owner=True),
        _node(3, role="Con", will_receive=True),
    ]
    case = _make_case(
        a, properties=[_prop()],
        case_state=_case_state_v3(
            rows, nodes=nodes,
            extra_diagram={"engineInput": {"version": 2, "nodes": nodes},
                           "engineResult": {"status": "complete",
                                            "allocations": {}}}))
    mapping = build_template_mapping(case)
    assert mapping["[Người 1 - Họ tên]"] == "An Stage 1"
    assert mapping["[Người 2 - Họ tên]"] == "Bích Stage 2"
    assert mapping["[Người 3 - Họ tên]"] == "Cường Stage 3"


def test_two_party_side_names_include_every_person_across_empty_seats():
    a1, a2, b1, b2, unassigned = [_rid() for _ in range(5)]
    rows = [_stage_row_v3(b2, 14, "B Hai"),
            _stage_row_v3(a2, 12, "A Hai"),
            _stage_row_v3(unassigned, 15, "Chưa gán"),
            _stage_row_v3(a1, 11, "A Một"),
            _stage_row_v3(b1, 13, "B Một")]
    mapping = build_template_mapping(_two_party_case(
        rows, {2: a1, 5: a2, 17: b1, 20: b2}))
    assert mapping["[Tên bên A]"] == "A Một và A Hai"
    assert mapping["[Tên bên B]"] == "B Một và B Hai"
    assert mapping["[Bên A 1 - Họ tên]"] == ""
    assert mapping["[Bên A 2 - Họ tên]"] == "A Một"
    assert mapping["[Bên A 3 - Họ tên]"] == ""
    assert mapping["[Bên A 5 - Họ tên]"] == "A Hai"
    assert mapping["[Bên B 1 - Họ tên]"] == ""
    assert mapping["[Bên B 2 - Họ tên]"] == "B Một"
    assert mapping["[Bên B 5 - Họ tên]"] == "B Hai"
    assert mapping["[Người 3 - Họ tên]"] == "Chưa gán"


def test_two_party_ben_a_b_slots_and_names():
    ra, rb, rc = _rid(), _rid(), _rid()
    rows = [
        _stage_row_v3(ra, 11, "Nguyễn Bên A"),
        _stage_row_v3(rb, 12, "Trần Bên B"),
        _stage_row_v3(rc, 13, "Lê Không Slot"),
    ]
    # p2 = bên A (thứ tự slot ≠ thứ tự stage), p16 = bên B; rc không slot.
    case = _two_party_case(rows, {2: ra, 16: rb})
    mapping = build_template_mapping(case)

    assert mapping["[Người 1 - Họ tên]"] == "Nguyễn Bên A"
    assert mapping["[Người 2 - Họ tên]"] == "Trần Bên B"
    assert mapping["[Người 3 - Họ tên]"] == "Lê Không Slot"
    assert mapping["[Bên A 1 - Họ tên]"] == ""       # p1 trống, không dồn p2 lên
    assert mapping["[Bên A 2 - Họ tên]"] == "Nguyễn Bên A"
    assert mapping["[Bên B 1 - Họ tên]"] == "Trần Bên B"
    assert mapping["[Bên B 2 - Họ tên]"] == ""
    assert mapping["[Tên bên A]"] == "Nguyễn Bên A"
    assert mapping["[Tên bên B]"] == "Trần Bên B"
    assert mapping["[Bên A - Họ tên]"].endswith("Nguyễn Bên A")
    assert mapping["[Bên B - Họ tên]"].endswith("Trần Bên B")
    # Người Stage không có slot: không vào Bên A/B nhưng vẫn trong
    # nhóm Người và alias ký giữ hoạt động.
    assert "Nguyễn Bên A" in mapping["[Danh sách bên a]"]
    assert "Lê Không Slot" not in mapping["[Danh sách bên a]"]
    assert mapping["[Danh sách ngườ ký]"] == mapping["[Danh sách người ký]"]
    assert mapping["[Họ tên ngườ ký 1]"] == mapping["[Họ tên người ký 1]"]
    # Placeholder riêng thừa kế render rỗng trên hồ sơ hai bên.
    assert mapping["[Đoạn phân chia di sản]"] == ""
    assert mapping["[Cụm chủ đất chết]"] == ""
    assert mapping["[Người nhận 1 - Họ tên]"] == ""


def test_two_party_word_block_reasons():
    """W6: thiếu Người/sơ đồ/Tài sản KHÔNG còn là block — chỉ ghi thiếu."""
    ra = _rid()
    rows = [_stage_row_v3(ra, 11, "Nguyễn Bên A")]

    empty = _two_party_case([], {})
    ctx = build_word_context(empty)
    assert word_block_reason(ctx) is None
    missing = []
    mapping = build_template_mapping(empty, needed={"two_party_slots"},
                                     missing=missing)
    assert "người trong hồ sơ" in missing
    assert mapping["[Bên A 1 - Họ tên]"] == ""

    # Có người nhưng chưa lưu diagram.state → ghế trống + ghi thiếu.
    no_diagram = _make_case(
        None, properties=[_prop()],
        case_state=json.dumps(
            {"schemaVersion": 3, "case_type": "two_party",
             "stage": rows, "assets": [], "diagram": {"state": {}}},
            ensure_ascii=False))
    no_diagram.loai_van_ban = "chuyen_nhuong"
    ctx = build_word_context(no_diagram)
    assert word_block_reason(ctx) is None
    missing = []
    mapping = build_template_mapping(no_diagram, needed={"two_party_slots"},
                                     missing=missing)
    assert "sơ đồ hai bên đã lưu" in missing
    assert mapping["[Bên A 1 - Họ tên]"] == ""

    # Không áp điều kiện thừa kế: đủ người + slot + tài sản → ready.
    ok = _two_party_case(rows, {1: ra})
    ctx = build_word_context(ok)
    assert word_block_reason(ctx) is None


def test_two_party_no_assets_not_blocked_and_noted():
    ra = _rid()
    rows = [_stage_row_v3(ra, 11, "Nguyễn Bên A")]
    case = _two_party_case(rows, {1: ra}, properties=[])
    case.tai_san = None
    ctx = build_word_context(case)
    assert word_block_reason(ctx) is None
    missing = []
    build_template_mapping(case, missing=missing)
    assert "tài sản" in missing


def test_word_spouse_role_uses_only_last_slot_member_and_preserves_other_roles():
    from services.word_engine import (
        WordPerson,
        _apply_spouse_slot_roles,
    )

    deceased = WordPerson(row_id="dead", ngay_chet="2020-01-01")
    faded = WordPerson(row_id="faded", role="Vợ/Chồng",
                       relation="Vợ/Chồng")
    faded_with_role = WordPerson(row_id="faded-role", role="Con")
    lawful = WordPerson(row_id="lawful")
    persons = [deceased, faded, faded_with_role, lawful]
    nodes = [
        {"id": "dead", "personId": "dead",
         "spouseSlotIds": ["faded", "faded-role", "lawful"]},
        {"id": "faded", "personId": "faded"},
        {"id": "faded-role", "personId": "faded-role"},
        {"id": "lawful", "personId": "lawful"},
    ]

    _apply_spouse_slot_roles(persons, nodes, "dead")

    assert faded.role == "" and faded.relation == ""
    assert faded_with_role.role == "Con"
    assert lawful.role == "Vợ/Chồng"


def test_word_context_applies_saved_one_way_spouse_slots():
    owner = _person(1, "Chủ đất", dead=date(2024, 1, 1))
    faded = _person(2, "Vợ trước")
    lawful = _person(3, "Vợ sau")
    representative = _person(4, "Đại diện")
    owner_node = _node(1, role="Owner", is_land_owner=True)
    owner_node["spouseSlotIds"] = ["node_2", "node_3"]
    nodes = [
        _node(3, will_receive=True),
        _node(2, role="Vợ/Chồng"),
        owner_node,
    ]
    case = _make_case(
        owner,
        participants=[_participant(representative, "Người được ủy quyền")],
        case_state=_case_state_v2(
            owner, faded, lawful, representative, nodes=nodes,
            allocations={"3": {"finalShare": "1/2"}}))

    context = build_word_context(case)
    by_name = {person.ho_ten: person for person in context.persons}
    assert [person.ho_ten for person in context.persons] == [
        "Chủ đất", "Vợ trước", "Vợ sau", "Đại diện"]
    assert by_name["Vợ trước"].role == ""
    assert by_name["Vợ sau"].role == "Vợ/Chồng"
    assert by_name["Vợ sau"].share == "1/2"
    assert context.receivers == [by_name["Vợ sau"]]
    assert context.authorized_people == [by_name["Đại diện"]]


def test_branch_spouse_never_becomes_owners_spouse_in_family_output():
    owner = _person(1, "Chủ đất", dead=date(2020, 1, 1))
    branch_spouse = _person(2, "Con dâu")
    child = _person(3, "Con trai", dead=date(2021, 1, 1))
    owner_spouse = _person(4, "Vợ chủ đất")
    nodes = [
        {"id": "owner", "personId": "1", "isLandOwner": True,
         "spouseSlotIds": ["spouse_owner"]},
        {"id": "child", "personId": "3", "parentSlotIds": ["owner"],
         "spouseSlotIds": ["spouse_branch"]},
        {"id": "spouse_branch", "personId": "2"},
        {"id": "spouse_owner", "personId": "4"},
    ]
    case = _make_case(
        owner, participants=[
            _participant(branch_spouse, "Con_dau_re"),
            _participant(owner_spouse, "Vợ/Chồng"),
        ], case_state=_case_state_v2(
            branch_spouse, owner, child, owner_spouse,
            nodes=nodes, allocations={}))

    context = build_word_context(case)
    assert next(p for p in context.persons if p.ho_ten == "Con dâu").role == "Con_dau_re"
    family = build_template_mapping(case)["[Đoạn quan hệ gia đình]"]
    spouse_lines = [line for line in family.splitlines() if "có vợ/chồng là" in line]
    assert len(spouse_lines) == 1
    assert "Vợ chủ đất" in spouse_lines[0]
    assert "Con dâu" not in spouse_lines[0]
