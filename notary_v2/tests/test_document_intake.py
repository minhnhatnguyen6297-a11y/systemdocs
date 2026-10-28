"""Tests for services.document_intake (MIN-108 — multi-source intake).

Contract: notary.case-drafting.v2 — suggestions carry observation metadata,
never `confirmed`; per-source failure isolation; hard limits §5.2.
"""
import json
import unicodedata
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

import openpyxl

from services.document_intake import (
    IntakeError,
    IntakeOutcome,
    MAX_FILE_BYTES,
    MAX_TEXT_CHARS,
    analyze,
)
from services.document_intake.docx_source import read_docx
from services.document_intake import service as document_intake_service
from services.document_intake.service import _is_abs, _is_unc


def _sid(n):
    # uuid4 hợp lệ theo contract (version=4, variant=8..b), deterministic cho test
    return f"00000000-0000-4000-8000-{n:012d}"


def _file_source(n, kind, path, size=None):
    return {
        "source_id": _sid(n),
        "kind": kind,
        "file_ref": {
            "path": str(path),
            "scope": "machine_local",
            "size_bytes": Path(path).stat().st_size if size is None else size,
        },
    }


def _text_source(n, text):
    return {"source_id": _sid(n), "kind": "text", "text": text}


CCCD_FRONT_LINES = [
    "CĂN CƯỚC CÔNG DÂN",
    "Họ và tên: NGUYỄN VĂN AN",
    "Số: 001234567890",
    "Ngày sinh: 20/03/1980",
    "Giới tính: Nam",
    "Nơi thường trú: Số 1 Đường ABC, Hà Nội",
]

CCCD_BACK_LINES = [
    "CĂN CƯỚC CÔNG DÂN",
    "IDVNA001234567890<<<<<<<<<<<<<",
    "8003201M2503201VNA<<<<<<<<<<<4",
    "NGUYEN<<VAN<AN<<<<<<<<<<<<<<<<",
    "Ngày cấp: 15/01/2021",
]

PROPERTY_LINES = [
    "GIẤY CHỨNG NHẬN",
    "QUYỀN SỬ DỤNG ĐẤT",
    "Số phát hành (serial): DD 123456",
    "Số vào sổ: VP00166",
    "Thửa đất số: 10",
    "Tờ bản đồ số: 5",
    "Diện tích: 447,0 m2",
    "Địa chỉ: Thôn A, Xã B, Huyện C, Tỉnh D",
    "Ngày cấp: 20/05/2010",
]

# Bản ASCII cho PDF text layer — PyMuPDF insert_text chỉ encode được latin-1
# với font built-in; parser fold dấu nên bản không dấu parse tương đương.
PROPERTY_LINES_ASCII = [
    "GIAY CHUNG NHAN",
    "QUYEN SU DUNG DAT",
    "So phat hanh (serial): DD 123456",
    "So vao so: VP00166",
    "Thua dat so: 10",
    "To ban do so: 5",
    "Dien tich: 447,0 m2",
    "Dia chi: Thon A, Xa B, Huyen C, Tinh D",
    "Ngay cap: 20/05/2010",
]

DEATH_LINES = [
    "GIẤY KHAI TỬ",
    "Họ, chữ đệm, tên người chết: TRẦN THỊ BÌNH",
    "Ngày sinh: 01/02/1940",
    "Giấy tờ tùy thân: 009876543210",
    "Nơi cư trú cuối cùng: Số 2 Đường XYZ, Hà Nội",
    "Đã chết vào ngày: 10/11/2020",
]

# GCN nhiều thửa — parser thật trích land_rows=[{ONT,200.0,Lâu dài},{CLN,247.0,12/2043}]
# và fill flat fields: dien_tich=447 (sum), loai_dat=ONT (thửa đầu).
MULTI_PARCEL_LINES = [
    "GIẤY CHỨNG NHẬN",
    "QUYỀN SỬ DỤNG ĐẤT",
    "Số phát hành (serial): DD 123456",
    "Số vào sổ: VP00166",
    "Thửa đất số: 10",
    "Tờ bản đồ số: 5",
    "c. Loại đất: Đất ở tại nông thôn 200,0m²; Đất trồng cây lâu năm 247,0m²",
    "d. Thời hạn sử dụng: Đất ở tại nông thôn: Lâu dài; Đất trồng cây lâu năm: 12/2043",
    "Địa chỉ: Thôn A, Xã B",
    "Ngày cấp: 20/05/2010",
]


async def _fake_ocr_image(client, *, api_key, model, image_b64, filename, enable_rotate=False):
    """Trả lines theo tên file — fixture CCCD 2 mặt / GCN / giấy báo tử."""
    name = (filename or "").lower()
    if "back" in name:
        return list(CCCD_BACK_LINES)
    if "front" in name or "cccd" in name:
        return list(CCCD_FRONT_LINES)
    if "death" in name:
        return list(DEATH_LINES)
    if "gcn" in name or "property" in name or "so_do" in name:
        return list(PROPERTY_LINES)
    return ["không đọc được gì"]


async def _fake_ocr_by_page(client, *, api_key, model, image_b64, filename, enable_rotate=False):
    name = (filename or "").lower()
    if "#p1" in name:
        return list(CCCD_FRONT_LINES)
    if "#p2" in name:
        return list(CCCD_BACK_LINES)
    return list(PROPERTY_LINES)


def _walk_strings(obj):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield k
            yield from _walk_strings(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from _walk_strings(v)
    elif isinstance(obj, str):
        yield obj


def _make_pdf(tmp_path, name: str, page_texts: list) -> Path:
    import fitz

    pdf = fitz.open()
    for text in page_texts:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    out = Path(tmp_path) / name
    pdf.save(str(out))
    pdf.close()
    return out


def _make_docx(tmp_path, name: str, lines: list) -> Path:
    import docx

    document = docx.Document()
    for line in lines:
        document.add_paragraph(line)
    out = Path(tmp_path) / name
    document.save(str(out))
    return out


def _make_xlsx(tmp_path, name: str, headers: list, rows: list) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DanhSach"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    out = Path(tmp_path) / name
    wb.save(str(out))
    return out


class SourceValidationTests(unittest.TestCase):
    """Job-level limits → IntakeError với contract codes."""

    def test_too_many_sources(self):
        sources = [_text_source(i, "x") for i in range(1, 10)]
        with self.assertRaises(IntakeError) as cm:
            analyze(sources)
        self.assertEqual(cm.exception.code, "intake_too_many_sources")

    def test_empty_sources(self):
        with self.assertRaises(IntakeError) as cm:
            analyze([])
        self.assertEqual(cm.exception.code, "validation_error")

    def test_unsupported_kind(self):
        with self.assertRaises(IntakeError) as cm:
            analyze([{"source_id": _sid(1), "kind": "zip", "text": "x"}])
        self.assertEqual(cm.exception.code, "intake_unsupported_source")

    def test_duplicate_source_id(self):
        src = _text_source(1, "abc")
        with self.assertRaises(IntakeError) as cm:
            analyze([src, dict(src)])
        self.assertEqual(cm.exception.code, "validation_error")

    def test_text_over_limit(self):
        with self.assertRaises(IntakeError) as cm:
            analyze([_text_source(1, "x" * (MAX_TEXT_CHARS + 1))])
        self.assertEqual(cm.exception.code, "intake_text_too_long")

    def test_text_with_file_ref_rejected(self):
        src = _text_source(1, "abc")
        src["file_ref"] = {"path": "D:/x.txt", "scope": "machine_local", "size_bytes": 3}
        with self.assertRaises(IntakeError) as cm:
            analyze([src])
        self.assertEqual(cm.exception.code, "validation_error")

    def test_declared_size_over_limit(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "a.png"
            p.write_bytes(b"img")
            src = _file_source(1, "image", str(p), size=MAX_FILE_BYTES + 1)
            with self.assertRaises(IntakeError) as cm:
                analyze([src])
            self.assertEqual(cm.exception.code, "intake_source_too_large")

    def test_actual_size_over_limit(self):
        # Declared size đã > limit → job-level validation trước khi đọc file.
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "big.png"
            p.write_bytes(b"x" * (MAX_FILE_BYTES + 1))
            src = _file_source(1, "image", str(p))  # declared = actual > limit
            with self.assertRaises(IntakeError) as cm:
                analyze([src])
            self.assertEqual(cm.exception.code, "intake_source_too_large")

    def test_misdeclared_size_per_source_error(self):
        # Declared nhỏ nhưng file thật > limit → per-source error (stat trước
        # read nên file > limit không bao giờ được đọc vào RAM).
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "big.png"
            p.write_bytes(b"x" * (MAX_FILE_BYTES + 1))
            src = _file_source(1, "image", str(p), size=100)
            outcome = analyze([src, _text_source(2, "abc")], api_key="k")
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.errors[0]["code"], "intake_source_too_large")
        self.assertEqual(outcome.errors[0]["source_id"], _sid(1))

    def test_is_dir_true_rejected(self):
        src = {
            "source_id": _sid(1),
            "kind": "pdf",
            "file_ref": {
                "path": "D:/folder",
                "scope": "machine_local",
                "size_bytes": 1,
                "is_dir": True,
            },
        }
        with self.assertRaises(IntakeError) as cm:
            analyze([src])
        self.assertEqual(cm.exception.code, "validation_error")

    def test_missing_size_bytes_rejected(self):
        src = {
            "source_id": _sid(1),
            "kind": "pdf",
            "file_ref": {"path": "D:/x.pdf", "scope": "machine_local"},
        }
        with self.assertRaises(IntakeError) as cm:
            analyze([src])
        self.assertEqual(cm.exception.code, "validation_error")

    def test_bad_scope_rejected(self):
        src = {
            "source_id": _sid(1),
            "kind": "pdf",
            "file_ref": {"path": "D:/x.pdf", "scope": "onedrive", "size_bytes": 1},
        }
        with self.assertRaises(IntakeError) as cm:
            analyze([src])
        self.assertEqual(cm.exception.code, "file_scope_not_supported")

    def test_posix_absolute_path_accepted_only_on_non_windows(self):
        with patch.object(document_intake_service.os, "name", "posix"):
            source = {
                "source_id": _sid(1),
                "kind": "pdf",
                "file_ref": {
                    "path": "/tmp/x.docx",
                    "scope": "machine_local",
                    "size_bytes": 1,
                },
            }
            self.assertEqual(
                document_intake_service._validate_sources([source])[0].path,
                "/tmp/x.docx",
            )

            for path in ("tmp/x.docx", "//server/share/x.docx"):
                source["file_ref"]["path"] = path
                with self.subTest(path=path):
                    with self.assertRaises(IntakeError) as cm:
                        document_intake_service._validate_sources([source])
                    self.assertEqual(cm.exception.code, "file_scope_not_supported")

    def test_relative_path_rejected(self):
        src = {
            "source_id": _sid(1),
            "kind": "pdf",
            "file_ref": {"path": "docs/x.pdf", "scope": "machine_local", "size_bytes": 1},
        }
        with self.assertRaises(IntakeError) as cm:
            analyze([src])
        self.assertEqual(cm.exception.code, "file_scope_not_supported")

    def test_missing_file_ref_for_file_kind(self):
        with self.assertRaises(IntakeError) as cm:
            analyze([{"source_id": _sid(1), "kind": "pdf"}])
        self.assertEqual(cm.exception.code, "validation_error")

    def test_text_on_file_kind_rejected(self):
        src = _text_source(1, "abc")
        src["kind"] = "pdf"
        with self.assertRaises(IntakeError) as cm:
            analyze([src])
        self.assertEqual(cm.exception.code, "validation_error")


class PathHelperTests(unittest.TestCase):
    r"""Helper absolute/UNC local cua document_intake — dung cho Windows
    drive, \\?\ long path, POSIX absolute; tu choi relative va UNC."""

    def test_is_abs_windows_drive(self):
        for p in ("D:\\intake\\f.pdf", "X:/tmp/a.jpg", "z:\\x"):
            self.assertTrue(_is_abs(p), p)

    def test_is_abs_long_prefix(self):
        self.assertTrue(_is_abs("\\\\?\\D:\\duong\\dan.pdf"))

    def test_is_abs_posix_absolute(self):
        self.assertTrue(_is_abs("/home/user/a.pdf", windows=False))
        self.assertFalse(_is_abs("/home/user/a.pdf", windows=True))

    def test_is_abs_relative_rejected(self):
        for p in ("docs/x.pdf", "a.pdf", "./x", "..\\y"):
            self.assertFalse(_is_abs(p), p)

    def test_is_unc_forms(self):
        for p in ("\\\\server\\share\\f.pdf",
                  "//server/share/f.pdf",
                  "\\\\?\\UNC\\server\\share\\f.pdf"):
            self.assertTrue(_is_unc(p), p)

    def test_is_unc_long_prefix_local_not_unc(self):
        self.assertFalse(_is_unc("\\\\?\\D:\\duong\\dan.pdf"))
        self.assertFalse(_is_unc("D:\\a.pdf"))
        self.assertFalse(_is_unc("/home/u/a.pdf"))



class ShapeInvariantTests(unittest.TestCase):
    """Contract §4 invariants trên mọi suggestion."""

    def test_no_confirmed_no_empty_raw(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            img = Path(td) / "cccd_front.png"
            img.write_bytes(b"fake-image-bytes")
            outcome = analyze(
                [_file_source(1, "image", str(img))],
                api_key="k", ocr_call=_fake_ocr_image)
        self.assertEqual(outcome.status, "succeeded")
        data = outcome.result_data
        blob = list(_walk_strings(data))
        self.assertNotIn("confirmed", blob)
        for s in data["suggestions"]:
            uuid.UUID(s["suggestion_id"])  # uuid4
            self.assertIn(s["target"], ("person", "asset"))
            for name, f in s["fields"].items():
                self.assertIn(f["observation_state"], ("observed", "normalized", "inferred"))
                if f["observation_state"] != "inferred":
                    self.assertIsNone(f["confidence"])
                self.assertTrue(f["source_refs"], f"{name} thiếu source_refs")
                self.assertNotEqual(f["raw_value"], "")

    def test_raw_differs_from_normalized(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            img = Path(td) / "cccd_front.png"
            img.write_bytes(b"fake-image-bytes")
            outcome = analyze(
                [_file_source(1, "image", str(img))],
                api_key="k", ocr_call=_fake_ocr_image)
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "person")
        fields = sug["fields"]
        # name uppercase → title case normalized
        self.assertEqual(fields["ho_ten"]["raw_value"], "NGUYỄN VĂN AN")
        self.assertEqual(fields["ho_ten"]["normalized_value"], "Nguyễn Văn An")
        self.assertEqual(fields["ho_ten"]["observation_state"], "normalized")
        self.assertIsNone(fields["ho_ten"]["confidence"])
        # date dd/mm/YYYY → YYYY-MM-DD
        self.assertEqual(fields["ngay_sinh"]["normalized_value"], "1980-03-20")


class ImageAdapterTests(unittest.TestCase):
    def test_cccd_front_person(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            img = Path(td) / "cccd_front.png"
            img.write_bytes(b"fake-image")
            outcome = analyze(
                [_file_source(1, "image", str(img))],
                api_key="k", ocr_call=_fake_ocr_image)
        self.assertEqual(outcome.status, "succeeded")
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "person")
        self.assertEqual(sug["fields"]["so_giay_to"]["normalized_value"], "001234567890")
        self.assertEqual(sug["fields"]["gioi_tinh"]["normalized_value"], "Nam")

    def test_death_certificate_person(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            img = Path(td) / "death.png"
            img.write_bytes(b"fake-image")
            outcome = analyze(
                [_file_source(1, "image", str(img))],
                api_key="k", ocr_call=_fake_ocr_image)
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "person")
        self.assertEqual(sug["fields"]["ngay_chet"]["normalized_value"], "2020-11-10")

    def test_property_asset(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            img = Path(td) / "gcn.png"
            img.write_bytes(b"fake-image")
            outcome = analyze(
                [_file_source(1, "image", str(img))],
                api_key="k", ocr_call=_fake_ocr_image)
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "asset")
        self.assertEqual(
            sug["fields"]["so_serial"]["normalized_value"], "DD123456")
        self.assertEqual(sug["fields"]["so_serial"]["observation_state"], "inferred")
        self.assertIsNotNone(sug["fields"]["so_serial"]["confidence"])
        self.assertEqual(sug["fields"]["dien_tich"]["normalized_value"], 447.0)

    def test_missing_api_key_is_per_source_error(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            img = Path(td) / "cccd_front.png"
            img.write_bytes(b"fake-image")
            outcome = analyze(
                [_file_source(1, "image", str(img)), _text_source(2, "ngẫu nhiên")],
                api_key="", ocr_call=_fake_ocr_image)
        self.assertEqual(outcome.status, "partial")
        err = outcome.errors[0]
        self.assertEqual(err["code"], "ocr.engine_unavailable")
        self.assertEqual(err["source_id"], _sid(1))
        # text source vẫn được xử lý (parse_failed ở đây vì text vô nghĩa)
        self.assertEqual(set(outcome.failed + outcome.succeeded), {_sid(1), _sid(2)})

    def test_ocr_failure_does_not_stop_other_sources(self):
        import tempfile
        async def _boom(client, **kw):
            raise RuntimeError("engine down")
        with tempfile.TemporaryDirectory() as td:
            img = Path(td) / "cccd_front.png"
            img.write_bytes(b"fake-image")
            xlsx = _make_xlsx(td, "people.xlsx",
                              ["Họ tên", "Giới tính", "Ngày sinh", "Số giấy tờ"],
                              [["Nguyen Van A", "Nam", "1988", "111222333444"]])
            outcome = analyze(
                [_file_source(1, "image", str(img)), _file_source(2, "xlsx", str(xlsx))],
                api_key="k", ocr_call=_boom)
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.failed, [_sid(1)])
        self.assertEqual(outcome.succeeded, [_sid(2)])
        self.assertEqual(outcome.errors[0]["code"], "ocr.engine_unavailable")
        self.assertEqual(len(outcome.suggestions), 1)

    def test_wrong_extension_parse_failed(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = Path(td) / "evil.exe"
            f.write_bytes(b"MZ")
            outcome = analyze([_file_source(1, "image", str(f))], api_key="k")
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.errors[0]["code"], "intake.parse_failed")

    def test_missing_file_per_source_error(self):
        outcome = analyze(
            [{"source_id": _sid(1), "kind": "image",
              "file_ref": {"path": "D:/nonexistent_dir_xyz/no.png",
                           "scope": "machine_local", "size_bytes": 5}},
             _text_source(2, "tối thiểu text")],
            api_key="k")
        self.assertIn(_sid(1), outcome.failed)
        self.assertEqual(outcome.errors[0]["code"], "file_not_found")


class PdfAdapterTests(unittest.TestCase):
    def test_pdf_text_layer_property(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            pdf = _make_pdf(td, "gcn.pdf", ["\n".join(PROPERTY_LINES_ASCII)])
            outcome = analyze([_file_source(1, "pdf", str(pdf))], api_key="k")
        self.assertEqual(outcome.status, "succeeded")
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "asset")
        self.assertEqual(sug["fields"]["so_serial"]["normalized_value"], "DD123456")
        refs = sug["fields"]["so_serial"]["source_refs"]
        self.assertTrue(any(r.get("page") == 1 for r in refs))

    def test_pdf_scanned_pages_use_ocr_and_pair(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            # 2 trang không text layer → OCR path
            pdf = _make_pdf(td, "cccd.pdf", ["", ""])
            outcome = analyze(
                [_file_source(1, "pdf", str(pdf))],
                api_key="k", ocr_call=_fake_ocr_by_page)
        self.assertEqual(outcome.status, "succeeded")
        sugs = outcome.result_data["suggestions"]
        self.assertEqual(len(sugs), 1)
        sug = sugs[0]
        self.assertEqual(sug["target"], "person")
        self.assertEqual(sug["fields"]["so_giay_to"]["normalized_value"], "001234567890")
        # ghép mặt sau: ngay_cap lấy từ trang back
        self.assertEqual(sug["fields"]["ngay_cap"]["normalized_value"], "2021-01-15")
        pages = {r.get("page") for f in sug["fields"].values() for r in f["source_refs"]}
        self.assertEqual(pages, {1, 2})

    def test_pdf_scanned_no_api_key(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            pdf = _make_pdf(td, "scan.pdf", [""])
            outcome = analyze([_file_source(1, "pdf", str(pdf))], api_key="")
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.errors[0]["code"], "ocr.engine_unavailable")

    def test_pdf_too_many_pages(self):
        # Limit §5.2 phát hiện lúc mở file → per-source error, không hủy batch.
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            pdf = _make_pdf(td, "big.pdf", [f"page {i}" for i in range(51)])
            ok = _make_xlsx(td, "people.xlsx", ["Họ tên"], [["Nguyen A"]])
            outcome = analyze(
                [_file_source(1, "pdf", str(pdf)), _file_source(2, "xlsx", str(ok))],
                api_key="k")
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.failed, [_sid(1)])
        self.assertEqual(outcome.succeeded, [_sid(2)])
        self.assertEqual(outcome.errors[0]["code"], "intake_source_too_large")


class ReviewFindingAdapterTests(unittest.TestCase):
    """7 phát hiện Quick review #22 — hành vi của Document Intake (consumer).

    Viết TRƯỚC khi vá; log đỏ ở `evidence/q-fixes/before-adapter.txt`.
    """

    TRANSFER = [
        "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
        "BÊN CHUYỂN NHƯỢNG (Bên A):",
        "Ông Nguyễn  Văn Mẫu; Sinh ngày: 01/01/1980; CCCD số: 0010 7500 0001;",
        "BÊN NHẬN CHUYỂN NHƯỢNG (Bên B):",
        "Bà Trần Thị Mẫu; Sinh năm: 1985; CCCD số: 000000000002;",
        "TÀI SẢN:",
        "Tài sản thứ 1:",
        "Giấy chứng nhận số: AB123456; Thửa đất số: 12; Tờ bản đồ số: 3;",
        "Địa chỉ thửa đất: xã Mẫu, tỉnh Mẫu;",
        "Loại đất: Đất ở; Diện tích: 120 m2; Thời hạn: Lâu dài",
        "Loại đất: Đất vườn; Diện tích: 120 m2; Thời hạn: Lâu dài",
        "ĐIỀU 1: Nội dung thỏa thuận.",
        "LỜI CHỨNG CỦA CÔNG CHỨNG VIÊN",
        "Ông Nhân Chứng Mẫu; Sinh năm: 1970;",
    ]

    def _run(self, lines):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = _make_docx(td, "transfer.docx", lines)
            content = f.read_bytes()
            outcome = analyze([_file_source(1, "docx", str(f))], api_key="k")
        return outcome, content

    def test_q2_normalized_values_keep_source_span(self):
        """Q2: tên gộp khoảng trắng + CCCD có khoảng trắng vẫn ra trường với span nguồn."""
        outcome, content = self._run(self.TRANSFER)
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        people = [s for s in outcome.result_data["suggestions"] if s["target"] == "person"]
        person = next(p for p in people if "Văn Mẫu" in p["fields"]["ho_ten"]["raw_value"])
        ho_ten = person["fields"]["ho_ten"]
        ref = ho_ten["source_refs"][0]
        self.assertEqual(text[slice(*ref["span"])], "Nguyễn  Văn Mẫu")
        self.assertEqual(ref["raw_snippet"], "Nguyễn  Văn Mẫu")
        self.assertEqual(ho_ten["normalized_value"], "Nguyễn Văn Mẫu")
        giay_to = person["fields"]["so_giay_to"]
        ref_id = giay_to["source_refs"][0]
        self.assertEqual(text[slice(*ref_id["span"])], "0010 7500 0001")
        self.assertEqual(giay_to["normalized_value"], "001075000001")

    def test_q2_mixed_unicode_keeps_source_span(self):
        """Q2: nguồn trộn NFC/NFD vẫn định vị được span (không bỏ trường)."""
        lines = list(self.TRANSFER)
        lines[2] = "Ông Nguyễn Văn Mẫu; Sinh ngày: 01/01/1980; CCCD số: 000000000001;"
        nfd_line = "Bà Trần Thị Bi\u0300nh; Sinh năm: 1985; CCCD số: 000000000002;"
        lines[4] = nfd_line
        outcome, content = self._run(lines)
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        people = [s for s in outcome.result_data["suggestions"] if s["target"] == "person"]
        names = {p["fields"]["ho_ten"]["normalized_value"] for p in people}
        self.assertIn("Trần Thị Bình", names)
        person = next(p for p in people if p["fields"]["ho_ten"]["normalized_value"] == "Trần Thị Bình")
        ref = person["fields"]["ho_ten"]["source_refs"][0]
        self.assertEqual(unicodedata.normalize("NFC", text[slice(*ref["span"])]), "Trần Thị Bình")

    def test_q3_land_row_refs_stay_in_their_own_cluster(self):
        """Q3: cụm đất thứ hai nhận bằng chứng của chính nó, không mượn cụm đầu."""
        outcome, content = self._run(self.TRANSFER)
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        refs = asset["fields"]["land_rows"]["source_refs"]
        row1 = next(r for r in refs if r["land_row"] == 1 and r["land_field"] == "dien_tich")
        row2 = next(r for r in refs if r["land_row"] == 2 and r["land_field"] == "dien_tich")
        self.assertEqual(text[slice(*row1["span"])], "120 m2")
        self.assertEqual(text[slice(*row2["span"])], "120 m2")
        self.assertGreater(row2["span"][0], row1["span"][0])
        before_row2 = text[max(0, row2["span"][0] - 60):row2["span"][0]]
        self.assertIn("Đất vườn", before_row2)
        before_row1 = text[max(0, row1["span"][0] - 60):row1["span"][0]]
        self.assertIn("Đất ở", before_row1)
        self.assertNotIn("Đất vườn", before_row1)

    def test_q3_editing_one_row_does_not_touch_other_rows(self):
        """Q3: sửa một dòng đất chỉ đổi ref của dòng đó (không đụng dòng khác)."""
        base_outcome, base_content = self._run(self.TRANSFER)
        edited = list(self.TRANSFER)
        edited[10] = "Loại đất: Đất vườn; Diện tích: 150 m2; Thời hạn: Lâu dài"
        edited_outcome, edited_content = self._run(edited)
        base_text = read_docx(base_content).text
        edited_text = read_docx(edited_content).text

        def row_refs(outcome, content, text):
            asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
            refs = asset["fields"]["land_rows"]["source_refs"]
            return {
                (r["land_row"], r["land_field"]): (text[slice(*r["span"])], r["raw_snippet"])
                for r in refs
            }

        base_refs = row_refs(base_outcome, base_content, base_text)
        edited_refs = row_refs(edited_outcome, edited_content, edited_text)
        self.assertEqual(edited_refs[(2, "dien_tich")][0], "150 m2")
        for key in ((1, "loai_dat"), (1, "dien_tich"), (1, "thoi_han"), (2, "loai_dat"), (2, "thoi_han")):
            self.assertEqual(edited_refs[key], base_refs[key], key)

    def test_q4_multiple_serials_kept_for_reviewer(self):
        """Q4: hai serial trong một nhóm -> giữ mọi lần thấy, normalized null + cảnh báo."""
        lines = list(self.TRANSFER)
        lines[7] = "Giấy chứng nhận số: AB 123456; Thửa đất số: 12; Tờ bản đồ số: 3;"
        lines.insert(8, "Giấy chứng nhận số: CD 654321;")
        outcome, content = self._run(lines)
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        serial = asset["fields"]["so_serial"]
        self.assertIsNone(serial["normalized_value"])
        snippets = {r["raw_snippet"] for r in serial["source_refs"]}
        self.assertEqual(snippets, {"AB 123456", "CD 654321"})
        for ref in serial["source_refs"]:
            self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])
            self.assertEqual(ref["candidate_state"], "ambiguous_group")
        self.assertTrue(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))

    def test_q5_blank_land_cluster_warns_and_next_cluster_kept(self):
        """Q5: cụm `Loại đất:` trống không bị nuốt; vẫn cảnh báo thiếu + cụm sau đủ dữ liệu."""
        lines = list(self.TRANSFER)
        lines[9] = "Loại đất:"
        outcome, content = self._run(lines)
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        rows = json.loads(asset["fields"]["land_rows"]["normalized_value"])
        self.assertEqual(len(rows), 2)
        self.assertIsNone(rows[0]["loai_dat"])
        self.assertEqual(rows[1]["loai_dat"], "Đất vườn")
        refs = asset["fields"]["land_rows"]["source_refs"]
        self.assertTrue(any(r["land_row"] == 2 and r["land_field"] == "loai_dat" for r in refs))
        row2_dien_tich = next(r for r in refs if r["land_row"] == 2 and r["land_field"] == "dien_tich")
        self.assertEqual(row2_dien_tich["raw_snippet"], "120 m2")


class DocxAdapterTests(unittest.TestCase):
    def test_docx_property(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = _make_docx(td, "gcn.docx", PROPERTY_LINES)
            outcome = analyze([_file_source(1, "docx", str(f))], api_key="k")
        self.assertEqual(outcome.status, "succeeded")
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "asset")
        self.assertEqual(sug["fields"]["so_serial"]["normalized_value"], "DD123456")

    def test_docx_transfer_serial_raw_and_labeled_parcel_address(self):
        # Bug #7: serial nguyên văn TEST-001 + nhãn "Địa chỉ thửa đất" phải vào gợi ý tài sản.
        import tempfile
        lines = [
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "BÊN CHUYỂN NHƯỢNG (Bên A):",
            "Ông Nguyễn Văn Mẫu; Sinh ngày: 01/01/1980; CCCD số: 000000000001;",
            "BÊN NHẬN CHUYỂN NHƯỢNG (Bên B):",
            "Bà Trần Thị Mẫu; Sinh năm: 1985; CCCD số: 000000000002;",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: TEST-001; Thửa đất số: 12; Tờ bản đồ số: 3;",
            "Địa chỉ thửa đất: xã Mẫu, tỉnh Mẫu;",
            "Loại đất: Đất ở; Diện tích: 123,5 m2; Thời hạn: Lâu dài;",
            "ĐIỀU 1: Nội dung thỏa thuận.",
            "LỜI CHỨNG CỦA CÔNG CHỨNG VIÊN",
            "Ông Nhân Chứng Mẫu; Sinh năm: 1970;",
        ]
        with tempfile.TemporaryDirectory() as td:
            f = _make_docx(td, "transfer.docx", lines)
            content = f.read_bytes()
            outcome = analyze([_file_source(1, "docx", str(f))], api_key="k")
        self.assertEqual(outcome.status, "succeeded")
        suggestions = outcome.result_data["suggestions"]

        assets = [s for s in suggestions if s["target"] == "asset"]
        self.assertEqual(len(assets), 1)
        asset = assets[0]

        serial = asset["fields"]["so_serial"]
        self.assertEqual(serial["raw_value"], "TEST-001")
        self.assertIsNone(serial["normalized_value"])
        self.assertEqual(serial["observation_state"], "observed")
        ref = serial["source_refs"][0]
        self.assertEqual(ref["rule_id"], "asset.so_serial")
        self.assertEqual(ref["raw_snippet"], "TEST-001")
        self.assertEqual(read_docx(content).text[slice(*ref["span"])], "TEST-001")
        self.assertTrue(ref["docx_refs"])
        self.assertTrue(any(w["code"] == "intake.warning_nonstandard" for w in asset["warnings"]))
        self.assertFalse(any(w["code"] == "intake.fragment" for w in asset["warnings"]))

        dia_chi = asset["fields"]["dia_chi"]
        self.assertEqual(dia_chi["raw_value"], "xã Mẫu, tỉnh Mẫu")
        self.assertEqual(dia_chi["normalized_value"], "xã Mẫu, tỉnh Mẫu")
        self.assertEqual(dia_chi["observation_state"], "observed")
        ref2 = dia_chi["source_refs"][0]
        self.assertEqual(ref2["rule_id"], "asset.dia_chi")
        self.assertEqual(ref2["raw_snippet"], "xã Mẫu, tỉnh Mẫu")
        self.assertEqual(read_docx(content).text[slice(*ref2["span"])], ref2["raw_snippet"])

        people = [s for s in suggestions if s["target"] == "person"]
        self.assertEqual(len(people), 2)
        self.assertEqual(
            {p["fields"]["ho_ten"]["raw_value"] for p in people},
            {"Nguyễn Văn Mẫu", "Trần Thị Mẫu"},
        )
        self.assertFalse(any("Nhân Chứng" in s for s in _walk_strings(outcome.result_data)))

    def test_docx_transfer_blank_serial_and_address_stay_missing(self):
        # Nhãn trống (dòng mới/';'/nhãn kế cùng dòng) không được mượn "Loại đất"/"Thửa đất số".
        import tempfile
        base = [
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "BÊN CHUYỂN NHƯỢNG (Bên A):",
            "Ông Nguyễn Văn Mẫu; Sinh ngày: 01/01/1980; CCCD số: 000000000001;",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số:",
            "Thửa đất số: 12; Tờ bản đồ số: 3;",
        ]
        variants = {
            "newline": base + [
                "Địa chỉ thửa đất:",
                "Loại đất: Đất ở; Diện tích: 123,5 m2; Thời hạn: Lâu dài;",
                "ĐIỀU 1: Nội dung thỏa thuận.",
            ],
            "same-line label": base + [
                "Địa chỉ thửa đất: Loại đất: Đất ở; Diện tích: 123,5 m2; Thời hạn: Lâu dài;",
                "ĐIỀU 1: Nội dung thỏa thuận.",
            ],
        }
        for name, lines in variants.items():
            with self.subTest(variant=name):
                with tempfile.TemporaryDirectory() as td:
                    f = _make_docx(td, "blank.docx", lines)
                    outcome = analyze([_file_source(1, "docx", str(f))], api_key="k")
                self.assertEqual(outcome.status, "succeeded")
                assets = [s for s in outcome.result_data["suggestions"] if s["target"] == "asset"]
                self.assertEqual(len(assets), 1)
                asset = assets[0]
                self.assertNotIn("so_serial", asset["fields"])
                self.assertNotIn("dia_chi", asset["fields"])
                self.assertEqual(asset["fields"]["so_thua_dat"]["raw_value"], "12")
                self.assertTrue(any(w["code"] == "intake.fragment" for w in asset["warnings"]))
                for field in asset["fields"].values():
                    for ref in field["source_refs"]:
                        self.assertNotIn("Thửa đất số", ref["raw_snippet"])
                        self.assertNotIn("Loại đất", ref["raw_snippet"])

    def test_docx_corrupt_parse_failed(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = Path(td) / "bad.docx"
            f.write_bytes(b"not a docx")
            outcome = analyze([_file_source(1, "docx", str(f))], api_key="k")
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.errors[0]["code"], "intake.parse_failed")


class ExcelAdapterTests(unittest.TestCase):
    def test_person_rows_with_cell_refs(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = _make_xlsx(
                td, "people.xlsx",
                ["Họ tên", "Giới tính", "Ngày sinh", "Số giấy tờ", "Địa chỉ"],
                [
                    ["NGUYỄN VĂN A", "Nam", "1988", "111 222 333 444", "Hà Nội"],
                    ["Trần Thị B", "Nữ", "1990-05-10", "", "Nam Định"],
                ])
            outcome = analyze([_file_source(1, "xlsx", str(f))], api_key="k")
        self.assertEqual(outcome.status, "succeeded")
        sugs = outcome.result_data["suggestions"]
        self.assertEqual(len(sugs), 2)
        first = sugs[0]
        self.assertEqual(first["target"], "person")
        self.assertEqual(first["fields"]["ho_ten"]["normalized_value"], "Nguyễn Văn A")
        self.assertEqual(
            first["fields"]["so_giay_to"]["normalized_value"], "111222333444")
        ref = first["fields"]["ho_ten"]["source_refs"][0]
        self.assertEqual(ref["sheet"], "DanhSach")
        self.assertEqual(ref["cell"], "A2")
        # row 2 (sheet row 3)
        second = sugs[1]
        self.assertEqual(second["fields"]["gioi_tinh"]["normalized_value"], "Nữ")
        self.assertEqual(second["fields"]["ngay_sinh"]["normalized_value"], "1990-05-10")

    def test_missing_required_column_parse_failed(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = _make_xlsx(td, "bad.xlsx", ["A", "B"], [["1", "2"]])
            outcome = analyze([_file_source(1, "xlsx", str(f))], api_key="k")
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.errors[0]["code"], "intake.parse_failed")


class TextAdapterTests(unittest.TestCase):
    def test_pasted_property_text(self):
        outcome = analyze([_text_source(1, "\n".join(PROPERTY_LINES))], api_key="k")
        self.assertEqual(outcome.status, "succeeded")
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "asset")
        self.assertTrue(any("span" in r for f in sug["fields"].values() for r in f["source_refs"]))

    def test_cccd_qr_text(self):
        qr = "001234567890|012345678|Nguyen Van An|20031980|Nam|So 1 Duong ABC, Ha Noi|15012021"
        outcome = analyze([_text_source(1, qr)], api_key="k")
        self.assertEqual(outcome.status, "succeeded")
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "person")
        self.assertEqual(sug["fields"]["so_giay_to"]["normalized_value"], "001234567890")
        self.assertEqual(sug["fields"]["ngay_sinh"]["normalized_value"], "1980-03-20")

    def test_garbage_text_parse_failed(self):
        outcome = analyze([_text_source(1, "12345")], api_key="k")
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.errors[0]["code"], "intake.parse_failed")

    def test_land_rows_multi_parcel(self):
        # I-1: GCN nhiều thửa → field land_rows (JSON normalized_value) +
        # warning intake.multi_parcel; flat fields vẫn đúng (sum / thửa đầu).
        import json
        outcome = analyze(
            [_text_source(1, "\n".join(MULTI_PARCEL_LINES))], api_key="k")
        self.assertEqual(outcome.status, "succeeded")
        sug = outcome.result_data["suggestions"][0]
        self.assertEqual(sug["target"], "asset")

        land = sug["fields"].get("land_rows")
        self.assertIsNotNone(land, "thiếu field land_rows cho GCN nhiều thửa")
        rows = json.loads(land["normalized_value"])
        self.assertEqual(len(rows), 2)
        self.assertEqual(
            rows[0],
            {"loai_dat": "ONT", "dien_tich": 200, "thoi_han": "Lâu dài"})
        self.assertEqual(
            rows[1],
            {"loai_dat": "CLN", "dien_tich": 247, "thoi_han": "12/2043"})
        self.assertEqual(land["observation_state"], "normalized")
        self.assertIsNone(land["confidence"])
        self.assertTrue(land["raw_value"])
        self.assertTrue(any("span" in r for r in land["source_refs"]))

        codes = [w["code"] for w in sug["warnings"]]
        self.assertIn("intake.multi_parcel", codes)
        # Flat fields không đổi: dien_tich = tổng, loai_dat = thửa đầu.
        self.assertEqual(sug["fields"]["dien_tich"]["normalized_value"], 447)
        self.assertEqual(sug["fields"]["loai_dat"]["normalized_value"], "ONT")

    def test_single_parcel_no_land_rows_field(self):
        # 1 thửa → flat fields đủ, không emit land_rows (field chỉ cho ≥2).
        lines = [ln for ln in MULTI_PARCEL_LINES
                 if not ln.startswith(("c.", "d."))]
        lines.insert(6, "Loại đất: Đất ở tại nông thôn 200,0m²")
        outcome = analyze([_text_source(1, "\n".join(lines))], api_key="k")
        sug = outcome.result_data["suggestions"][0]
        self.assertNotIn("land_rows", sug["fields"])
        codes = [w["code"] for w in sug["warnings"]]
        self.assertNotIn("intake.multi_parcel", codes)

    def test_unsupported_target_vs_parse_failed(self):
        # M-1: doc_type nhận diện được nhưng chưa map person/asset
        # → intake.unsupported_target (unknown/missing → parse_failed).
        from services.document_intake import ocr_pipeline
        orig = ocr_pipeline._normalize_native_ocr_doc

        def _fake_marriage(lines, filename):
            return {"doc_type": "marriage",
                    "data": {"vo": "A", "chong": "B"},
                    "filename": filename}

        ocr_pipeline._normalize_native_ocr_doc = _fake_marriage
        try:
            outcome = analyze([_text_source(1, "giay chung nhan ket hon")],
                              api_key="k")
        finally:
            ocr_pipeline._normalize_native_ocr_doc = orig
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(
            outcome.errors[0]["code"], "intake.unsupported_target")


class AggregationTests(unittest.TestCase):
    def test_partial_breakdown(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            bad = Path(td) / "bad.docx"
            bad.write_bytes(b"junk")
            ok = _make_xlsx(td, "people.xlsx", ["Họ tên"], [["Nguyen A"]])
            outcome = analyze(
                [_file_source(1, "docx", str(bad)), _file_source(2, "xlsx", str(ok))],
                api_key="k")
        self.assertEqual(outcome.status, "partial")
        data = outcome.result_data
        self.assertEqual(data["breakdown"]["succeeded"], [_sid(2)])
        self.assertEqual(data["breakdown"]["failed"], [_sid(1)])
        self.assertEqual(len(data["suggestions"]), 1)
        self.assertEqual(len(data["errors"]), 1)

    def test_all_failed_status(self):
        outcome = analyze([_text_source(1, "12345")], api_key="k")
        self.assertEqual(outcome.status, "partial")
        self.assertEqual(outcome.succeeded, [])

    def test_all_succeeded_status(self):
        outcome = analyze([_text_source(1, "\n".join(PROPERTY_LINES))], api_key="k")
        self.assertEqual(outcome.status, "succeeded")
        self.assertNotIn("breakdown", outcome.result_data)

    def test_cancel_between_sources(self):
        calls = {"n": 0}
        def _cancel():
            calls["n"] += 1
            if calls["n"] > 1:
                raise RuntimeError("user_canceled")
        with self.assertRaises(RuntimeError):
            analyze(
                [_text_source(1, "\n".join(PROPERTY_LINES)),
                 _text_source(2, "\n".join(PROPERTY_LINES))],
                check_cancel=_cancel, api_key="k")

    def test_result_schema_version(self):
        outcome = analyze([_text_source(1, "\n".join(PROPERTY_LINES))], api_key="k")
        self.assertEqual(outcome.result_data["schema_version"], "notary.case-drafting.v2")


def test_spans_slice_source_for_nfd_and_non_bmp_text(tmp_path):
    """Span/`raw_snippet` phải cắt đúng NGUỒN khi nguồn ở dạng NFD hoặc có emoji.

    Lõi profile cũ có lớp test SpanUnicodeFidelity (CRLF/NFD/non-BMP); sau #22
    span do adapter Document Intake dựng nên hành vi đó được khoá ở đây: NFD
    không được mất trường, và mọi ref phải cắt đúng chuỗi trong nguồn gốc.
    """
    lines = [
        "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
        "I. BÊN CHUYỂN NHƯỢNG (Bên A):",
        "Ông: Nguyễn Văn An, sinh ngày: 10/05/1975;",
        "Căn cước công dân số: 001075000001 do Bộ Công an cấp ngày 15/03/2021;",
        "ĐIỀU 1: ĐỐI TƯỢNG",
        "- Thửa đất số: 215; Tờ bản đồ số: 34; Giấy chứng nhận số: AB 123456",
    ]

    def collect(variant_lines, name):
        path = _make_docx(tmp_path, f"{name}.docx", variant_lines)
        outcome = analyze([_file_source(1, "docx", str(path))], api_key="k")
        assert outcome.status == "succeeded", (name, outcome.errors)
        text = read_docx(path.read_bytes()).text
        fields, refs = {}, []
        for suggestion in outcome.result_data["suggestions"]:
            for key, field in suggestion["fields"].items():
                fields[(suggestion["target"], key)] = field["raw_value"]
                refs.extend(field["source_refs"])
        return fields, refs, text

    nfc_fields, nfc_refs, nfc_text = collect(lines, "nfc")
    nfd_fields, nfd_refs, nfd_text = collect(
        [unicodedata.normalize("NFD", line) for line in lines], "nfd"
    )
    emoji_fields, emoji_refs, emoji_text = collect(
        ["😀 ghi chú ngoài văn bản"] + lines, "non_bmp"
    )

    assert nfc_fields == nfd_fields == emoji_fields
    assert {target for target, _ in nfc_fields} == {"person", "asset"}
    for refs, text, name in ((nfc_refs, nfc_text, "nfc"), (nfd_refs, nfd_text, "nfd"),
                             (emoji_refs, emoji_text, "non_bmp")):
        assert refs, name
        for ref in refs:
            assert text[slice(*ref["span"])] == ref["raw_snippet"], (name, ref["rule_id"])


def test_refs_stay_candidate_local_when_values_repeat(tmp_path):
    """Giá trị lặp ở nhiều người: ref phải trỏ đúng vùng từng người.

    Cùng ngày sinh ở hai người → hai span khác nhau; adapter không mượn lần
    xuất hiện đầu tiên trong toàn văn bản (provenance cục bộ theo ứng viên).
    """
    lines = [
        "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
        "I. BÊN CHUYỂN NHƯỢNG (Bên A):",
        "Ông: Nguyễn Văn Mẫu, sinh ngày: 10/05/1975;",
        "Căn cước công dân số: 001075000001 do Bộ Công an cấp ngày 15/03/2021;",
        "II. BÊN NHẬN CHUYỂN NHƯỢNG (Bên B):",
        "Bà: Trần Thị Mẫu, sinh ngày: 10/05/1975;",
        "Căn cước công dân số: 001182000002 do Công an Hà Nội cấp ngày 20/04/2022;",
        "ĐIỀU 1: ĐỐI TƯỢNG",
        "- Thửa đất số: 215; Tờ bản đồ số: 34; Giấy chứng nhận số: AB 123456",
    ]
    path = _make_docx(tmp_path, "repeat.docx", lines)
    outcome = analyze([_file_source(1, "docx", str(path))], api_key="k")
    assert outcome.status == "succeeded", outcome.errors
    text = read_docx(path.read_bytes()).text
    refs = [ref for s in outcome.result_data["suggestions"] for f in s["fields"].values()
            for ref in f["source_refs"]]
    birth_refs = [ref for ref in refs if ref["raw_snippet"] == "10/05/1975"]
    spans = {tuple(ref["span"]) for ref in birth_refs}
    assert len(spans) == 2, f"ref ngày sinh phải cục bộ theo từng người: {sorted(spans)}"
    for ref in birth_refs:
        assert text[slice(*ref["span"])] == "10/05/1975"


class AdjudicationConsumerTests(unittest.TestCase):
    """Hành vi consumer sau adjudication 16/10/2026 (Q2/Q4 residual + MB3/MB4).

    Viết trước khi vá: Q2 gắn sai provenance, Q4 chọn hộ địa chỉ, MB3 bỏ địa
    chỉ của thửa khi câu mở nằm ngoài vùng tìm.
    """

    def _run(self, lines):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = _make_docx(td, "case.docx", lines)
            content = f.read_bytes()
            outcome = analyze([_file_source(1, "docx", str(f))], api_key="k")
        return outcome, content

    def test_q2_same_id_two_people_keep_own_positions(self):
        """Q2: hai người cùng số giấy tờ (một người ghi có khoảng trắng) -> mỗi người đúng span."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "BÊN A:",
            "Ông Nguyễn Văn An; Sinh năm: 1980; CCCD số: 0010 7500 0001;",
            "BÊN B:",
            "Bà Trần Thị Bình; Sinh năm: 1985; CCCD số: 001075000001;",
            "TÀI SẢN:",
            "Giấy chứng nhận số: AB123456; Thửa đất số: 12; Tờ bản đồ số: 3;",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        people = [s for s in outcome.result_data["suggestions"] if s["target"] == "person"]
        by_name = {p["fields"]["ho_ten"]["normalized_value"]: p for p in people}
        an_ref = by_name["Nguyễn Văn An"]["fields"]["so_giay_to"]["source_refs"][0]
        binh_ref = by_name["Trần Thị Bình"]["fields"]["so_giay_to"]["source_refs"][0]
        self.assertEqual(text[slice(*an_ref["span"])], "0010 7500 0001")
        self.assertEqual(text[slice(*binh_ref["span"])], "001075000001")
        self.assertLess(an_ref["span"][0], binh_ref["span"][0])

    def test_q4_repeated_address_group_ambiguous(self):
        """Q4: địa chỉ lặp trong nhóm -> normalized null + ref mọi lần thấy + cảnh báo."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: AB123456;",
            "Địa chỉ thửa đất: xã Một;",
            "Địa chỉ thửa đất: xã Hai;",
            "Thửa đất số: 12; Tờ bản đồ số: 3;",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        dia_chi = asset["fields"]["dia_chi"]
        self.assertIsNone(dia_chi["normalized_value"])
        self.assertEqual({r["raw_snippet"] for r in dia_chi["source_refs"]}, {"xã Một", "xã Hai"})
        for ref in dia_chi["source_refs"]:
            self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])
        self.assertTrue(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))

    def test_mb3_opener_address_keeps_ref(self):
        """MB3: câu mở thửa nằm TRƯỚC nhãn thửa -> địa chỉ vẫn có ref (không bị bỏ)."""
        outcome, content = self._run([
            "VĂN BẢN PHÂN CHIA DI SẢN",
            "Quyền sử dụng đất tại địa chỉ: xã Mẫu, tỉnh Mẫu theo Giấy chứng nhận quyền sử "
            "dụng đất số AK 732800 do UBND huyện Mẫu cấp ngày 01/01/2016.",
            "- Thửa đất số: 180; Tờ bản đồ số: 9",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        ref = asset["fields"]["dia_chi"]["source_refs"][0]
        self.assertEqual(text[slice(*ref["span"])], "xã Mẫu, tỉnh Mẫu")
        self.assertEqual(asset["fields"]["dia_chi"]["normalized_value"], "xã Mẫu, tỉnh Mẫu")


class Round4ConsumerTests(unittest.TestCase):
    """Hành vi consumer vòng 4 (16/10/2026): MB9 diện tích land_rows, Q4 khoá
    metadata theo shape, loại giấy giữ nguyên văn, serial nguyên văn, nhãn
    trống không thành dấu `:`. Viết trước khi vá: đỏ ở MB9/Q4/loại giấy."""

    def _run(self, lines):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = _make_docx(td, "case.docx", lines)
            content = f.read_bytes()
            outcome = analyze([_file_source(1, "docx", str(f))], api_key="k")
        return outcome, content

    def test_mb9_land_rows_area_typed_with_whole_raw_span(self):
        """MB9: diện tích cụm đất có đơn vị -> số đúng (200/247), nguyên văn giữ đủ."""
        outcome, content = self._run([
            "GIẤY CHỨNG NHẬN",
            "QUYỀN SỬ DỤNG ĐẤT",
            "Số phát hành (serial): DD 123456",
            "Thửa đất số: 10",
            "Tờ bản đồ số: 5",
            "c. Loại đất: Đất ở tại nông thôn 200,0m²; Đất trồng cây lâu năm 247,0m²",
            "d. Thời hạn sử dụng: Đất ở tại nông thôn: Lâu dài; Đất trồng cây lâu năm: 12/2043",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        rows = json.loads(asset["fields"]["land_rows"]["normalized_value"])
        self.assertEqual([row["dien_tich"] for row in rows], [200, 247])
        self.assertEqual([row["loai_dat"] for row in rows],
                         ["Đất ở tại nông thôn", "Đất trồng cây lâu năm"])
        self.assertEqual([row["thoi_han"] for row in rows], ["Lâu dài", "12/2043"])
        area_refs = [r for r in asset["fields"]["land_rows"]["source_refs"]
                     if r.get("land_field") == "dien_tich"]
        self.assertEqual([r["raw_snippet"] for r in area_refs], ["200,0m²", "247,0m²"])
        for ref in area_refs:
            self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])
        term_refs = [r for r in asset["fields"]["land_rows"]["source_refs"]
                     if r.get("land_field") == "thoi_han"]
        self.assertEqual([r["raw_snippet"] for r in term_refs], ["Lâu dài", "12/2043"])
        self.assertFalse(any(w["code"] == "intake.invalid_number" for w in asset["warnings"]))
        self.assertNotIn("447", asset["fields"]["land_rows"]["normalized_value"])
        self.assertNotIn("dien_tich", asset["fields"])  # không suy tổng cấp thửa

    def test_q4_serial_without_group_ambiguous_with_both_refs(self):
        """Q4: hai serial cùng thửa KHÔNG có nhãn nhóm -> null + đủ hai ref + cảnh báo."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Giấy chứng nhận số: AB123456;",
            "Giấy chứng nhận số: CD654321;",
            "Thửa đất số: 215; Tờ bản đồ số: 34;",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        field = asset["fields"]["so_serial"]
        self.assertIsNone(field["normalized_value"])
        self.assertEqual({r["raw_snippet"] for r in field["source_refs"]},
                         {"AB123456", "CD654321"})
        for ref in field["source_refs"]:
            self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])
        self.assertTrue(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))

    def test_loai_so_keeps_uppercase_source_slice(self):
        """MB1: loại giấy giữ NGUYÊN VĂN nguồn (tiêu đề hoa 2 dòng) + giá trị chuẩn."""
        outcome, content = self._run([
            "GIẤY CHỨNG NHẬN",
            "QUYỀN SỬ DỤNG ĐẤT",
            "Số phát hành (serial): DD 123456",
            "Thửa đất số: 10",
            "Tờ bản đồ số: 5",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        field = asset["fields"]["loai_so"]
        raw = "GIẤY CHỨNG NHẬN\nQUYỀN SỬ DỤNG ĐẤT"
        self.assertEqual(field["normalized_value"], "Giấy chứng nhận quyền sử dụng đất")
        self.assertEqual(field["raw_value"], raw)
        ref = field["source_refs"][0]
        self.assertEqual(ref["raw_snippet"], raw)
        self.assertEqual(text[slice(*ref["span"])], raw)

    def test_serial_raw_keeps_source_spacing(self):
        """Serial nguyên văn: hai dấu cách trong nguồn giữ đúng ở raw + span."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: AB  123456;",
            "Thửa đất số: 215;",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        field = asset["fields"]["so_serial"]
        self.assertEqual(field["raw_value"], "AB  123456")
        ref = field["source_refs"][0]
        self.assertEqual(ref["raw_snippet"], "AB  123456")
        self.assertEqual(text[slice(*ref["span"])], "AB  123456")
        self.assertFalse(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))

    def test_serial_same_value_spacing_not_ambiguous(self):
        """Cùng một serial ghi hai lần khác khoảng trắng: nguyên văn khác nhau nhưng KHÔNG mơ hồ giả."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: AB  123456;",
            "Giấy chứng nhận số: AB 123456;",
            "Thửa đất số: 215;",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        field = asset["fields"]["so_serial"]
        self.assertIsNotNone(field["normalized_value"])
        self.assertFalse(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))
        ref = field["source_refs"][0]
        self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])

    def test_blank_label_does_not_create_colon_address(self):
        """Nhãn trống: dấu `:` không thành địa chỉ giả, không cảnh báo mơ hồ."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: AB123456;",
            "Địa chỉ thửa đất:",
            "Địa chỉ thửa đất: xã Mẫu;",
            "Thửa đất số: 215;",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")
        field = asset["fields"]["dia_chi"]
        self.assertEqual(field["normalized_value"], "xã Mẫu")
        self.assertEqual([r["raw_snippet"] for r in field["source_refs"]], ["xã Mẫu"])
        self.assertEqual(text[slice(*field["source_refs"][0]["span"])], "xã Mẫu")
        self.assertFalse(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))


class Round5ConsumerTests(unittest.TestCase):
    """Hành vi consumer vòng 5 (16/10/2026): mọi lần thấy trên cùng một dòng,
    hợp nhất metadata khi gộp thửa, và KHÔNG báo mơ hồ giả khi hai dòng đất
    hợp lệ khác nhau. Viết trước khi vá: đỏ ở cả ba."""

    def _run(self, lines):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            f = _make_docx(td, "case.docx", lines)
            content = f.read_bytes()
            outcome = analyze([_file_source(1, "docx", str(f))], api_key="k")
        return outcome, content

    def _asset(self, outcome):
        return next(s for s in outcome.result_data["suggestions"] if s["target"] == "asset")

    def test_same_line_two_areas_ambiguous_with_both_refs(self):
        """Cùng dòng ghi diện tích hai lần khác nhau: null + đủ hai ref + cảnh báo."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Thửa đất số: 215; Tờ bản đồ số: 34;",
            "Diện tích: 120 m2; Diện tích: 130 m2",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = self._asset(outcome)
        field = asset["fields"]["dien_tich"]
        self.assertIsNone(field["normalized_value"])
        self.assertEqual([r["raw_snippet"] for r in field["source_refs"]], ["120 m2", "130 m2"])
        for ref in field["source_refs"]:
            self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])
        self.assertTrue(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))

    def test_repeated_parcel_term_ambiguity_after_merge(self):
        """Cùng thửa/tờ/diện tích nhắc lại với thời hạn khác: gộp thửa vẫn giữ đủ ứng viên."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Thửa đất số: 215; Tờ bản đồ số: 34; Diện tích: 120 m2; Thời hạn sử dụng: Lâu dài;",
            "Thửa đất số: 215; Tờ bản đồ số: 34; Diện tích: 120 m2; Thời hạn sử dụng: 12/2043;",
            "Giá chuyển nhượng: 100.000.000 đồng",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = self._asset(outcome)
        field = asset["fields"]["thoi_han"]
        self.assertIsNone(field["normalized_value"])
        self.assertEqual([r["raw_snippet"] for r in field["source_refs"]], ["Lâu dài", "12/2043"])
        for ref in field["source_refs"]:
            self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])
        self.assertTrue(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))

    def test_group_two_rows_no_false_ambiguous(self):
        """Nhóm có serial + thửa/tờ + hai dòng cụm đất khác diện tích/thời hạn:
        hai dòng đúng, đủ ref, KHÔNG mơ hồ giả."""
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: AB123456;",
            "Thửa đất số: 215; Tờ bản đồ số: 34;",
            "Loại đất: Đất ở; Diện tích: 120 m2; Thời hạn: Lâu dài",
            "Loại đất: Đất vườn; Diện tích: 130 m2; Thời hạn: 12/2043",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = self._asset(outcome)
        rows = json.loads(asset["fields"]["land_rows"]["normalized_value"])
        self.assertEqual(
            [(row["loai_dat"], row["dien_tich"], row["thoi_han"]) for row in rows],
            [("Đất ở", 120, "Lâu dài"), ("Đất vườn", 130, "12/2043")])
        refs = asset["fields"]["land_rows"]["source_refs"]
        area_refs = [r for r in refs if r.get("land_field") == "dien_tich"]
        term_refs = [r for r in refs if r.get("land_field") == "thoi_han"]
        self.assertEqual([r["raw_snippet"] for r in area_refs], ["120 m2", "130 m2"])
        self.assertEqual([r["raw_snippet"] for r in term_refs], ["Lâu dài", "12/2043"])
        for ref in area_refs + term_refs:
            self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])
        self.assertNotIn("dien_tich", asset["fields"])
        self.assertNotIn("thoi_han", asset["fields"])
        self.assertFalse(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))
        self.assertEqual(asset["fields"]["so_serial"]["raw_value"], "AB123456")
        self.assertEqual(asset["fields"]["so_thua_dat"]["raw_value"], "215")

    def test_multiline_group_two_rows_no_false_ambiguous(self):
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: AB123456;",
            "Thửa đất số: 215; Tờ bản đồ số: 34;",
            "Loại đất: Đất ở",
            "Diện tích: 120 m2",
            "Thời hạn: Lâu dài",
            "Loại đất: Đất vườn",
            "Diện tích: 130 m2",
            "Thời hạn: 12/2043",
        ])
        self.assertEqual(outcome.status, "succeeded")
        text = read_docx(content).text
        asset = self._asset(outcome)
        field = asset["fields"]["land_rows"]
        rows = json.loads(field["normalized_value"])
        self.assertEqual(
            [(row["loai_dat"], row["dien_tich"], row["thoi_han"]) for row in rows],
            [("Đất ở", 120, "Lâu dài"), ("Đất vườn", 130, "12/2043")])
        for key, expected in (
            ("loai_dat", ["Đất ở", "Đất vườn"]),
            ("dien_tich", ["120 m2", "130 m2"]),
            ("thoi_han", ["Lâu dài", "12/2043"]),
        ):
            refs = [r for r in field["source_refs"] if r.get("land_field") == key]
            self.assertEqual([r["raw_snippet"] for r in refs], expected)
            self.assertEqual([r["land_row"] for r in refs], [1, 2])
            for ref in refs:
                self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])
        self.assertNotIn("dien_tich", asset["fields"])
        self.assertNotIn("thoi_han", asset["fields"])
        self.assertFalse(any(w["code"] == "intake.ambiguous" for w in asset["warnings"]))
        self.assertEqual(asset["fields"]["so_serial"]["raw_value"], "AB123456")
        self.assertEqual(asset["fields"]["so_thua_dat"]["raw_value"], "215")
        self.assertEqual(asset["fields"]["so_to_ban_do"]["raw_value"], "34")

    def test_next_group_flat_areas_remain_ambiguous(self):
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: AB123456;",
            "Thửa đất số: 215; Tờ bản đồ số: 34;",
            "Loại đất: Đất ở",
            "Diện tích: 90 m2",
            "Thời hạn: Lâu dài",
            "Tài sản thứ 2:",
            "Giấy chứng nhận số: CD123456;",
            "Thửa đất số: 300; Tờ bản đồ số: 40;",
            "Diện tích: 120 m2",
            "Diện tích: 130 m2",
        ])
        self.assertEqual(outcome.status, "succeeded")
        assets = [s for s in outcome.result_data["suggestions"] if s["target"] == "asset"]
        self.assertEqual(len(assets), 2)
        self.assertFalse(any(w["code"] == "intake.ambiguous" for w in assets[0]["warnings"]))
        field = assets[1]["fields"]["dien_tich"]
        self.assertIsNone(field["normalized_value"])
        self.assertTrue(any(w["code"] == "intake.ambiguous" for w in assets[1]["warnings"]))
        self.assertEqual([r["raw_snippet"] for r in field["source_refs"]], ["120 m2", "130 m2"])
        text = read_docx(content).text
        for ref in field["source_refs"]:
            self.assertEqual(text[slice(*ref["span"])], ref["raw_snippet"])

    def test_bare_group_negative_flat_area_still_invalid(self):
        outcome, content = self._run([
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "TÀI SẢN:",
            "Tài sản thứ 1:",
            "Giấy chứng nhận số: AB123456;",
            "Diện tích: -50",
        ])
        self.assertEqual(outcome.status, "succeeded")
        asset = self._asset(outcome)
        field = asset["fields"]["dien_tich"]
        self.assertIsNone(field["normalized_value"])
        self.assertTrue(any(w["code"] == "intake.invalid_number" for w in asset["warnings"]))
        text = read_docx(content).text
        self.assertEqual(field["raw_value"], "-50")
        for ref in field["source_refs"]:
            self.assertEqual(text[slice(*ref["span"])], "-50")


if __name__ == "__main__":
    unittest.main()
