from docx import Document
from tests import test_case_metadata as m
from tests import test_draft_cleanup as c
from tests.test_case_metadata import db


def test_canonical_case_tokens_render(db):
    data = m._create(db, with_state=True, case_meta={
        'document_type': 'khai_nhan', 'noiniemyet': 'UBND X',
        'nguoinhanuyquyen': 'Nguoi UQ', 'noidungviec': 'Dinh chinh'})
    case = db.get(m.InheritanceCase, data['case']['id'])
    doc = Document()
    doc.add_paragraph('[noiniemyet]|[nguoinhanuyquyen]|[noidungviec]')
    m.word_engine.replace_in_doc(doc, m.word_engine.build_template_mapping(case))
    assert doc.paragraphs[0].text == 'UBND X|Nguoi UQ|Dinh chinh'


def test_changing_issue_date_recomputes_derived_fields(db):
    data = m._create(db, people_overrides=({}, {'ngay_cap': '2024-09-30'}))
    stage = data['stage']
    stage['people'][1]['ngay_cap'] = '2024-10-01'
    updated = m._commit(db, data['case']['id'], data['case']['revision'], stage)
    person = updated['stage']['people'][1]
    assert person['loai_giay_to'] == 'Căn cước'
    assert person['noi_cap'] == 'Bộ Công an'


def test_explicit_empty_snapshot_meta_does_not_read_master(db):
    data = m._create(db, with_state=True)
    case = db.get(m.InheritanceCase, data['case']['id'])
    case.noi_dung_viec = 'Changed outside snapshot'
    db.flush()
    assert m.word_engine.build_template_mapping(case)['[Nội dung việc]'] == ''


def test_catalog_recipient_also_participant_survives_delete(db):
    owner = c._customer(db, 'Owner')
    recipient = c._customer(db, 'Recipient')
    recipient_id = recipient.id
    prop = c._property(db, 'AA123456')
    case = c._case(db, owner, prop, uq_id=recipient_id, participants=[(recipient, None)])
    c.cases_router.delete(case.id, db)
    db.expire_all()
    assert db.get(c.Customer, recipient_id) is not None
