"""Tests cho services.word_batch_export — batch Word export theo
contract notary.case-drafting.v2 §8 (MIN-110).

Case duck-typed như test_word_engine (SimpleNamespace + case_state_json
theo persisted workspace shape). Chạy từ repo root notary_v2.
"""
from __future__ import annotations

import dataclasses
import hashlib
import json
import os
import re
import time
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from services import word_batch_export as wbe
from services import word_engine


BUILTIN_TEMPLATE = Path("word_templates/1. PCDS .docx")


# ---------------------------------------------------------------- fixtures


def _person(cid, name, *, gender="Nam", born=None, dead=None, doc_no="",
            issued=None, address=""):
    return SimpleNamespace(
        id=cid, ho_ten=name, gioi_tinh=gender, ngay_sinh=born,
        ngay_chet=dead, so_giay_to=doc_no, ngay_cap=issued,
        dia_chi=address, loai_giay_to="CCCD", noi_cap="CA",
        loai_dia_chi="",
    )


def _participant(person, role, *, receive=True):
    return SimpleNamespace(
        customer=person, customer_id=person.id, vai_tro=role,
        co_nhan_tai_san=receive, ty_le=0.0, hang_thua_ke=1,
        parent_customer_id=None,
    )


def _prop(**kwargs):
    defaults = dict(
        so_serial="BM 1451111", so_vao_so="CS123", so_thua_dat="45",
        so_to_ban_do="12", dia_chi="Thửa đất tại xã Minh Tân",
        loai_dat="Đất ở", dien_tich=120.5,
        loai_so="Giấy chứng nhận quyền sử dụng đất",
        land_rows_json='[{"loai_dat":"ONT","dien_tich":"80","thoi_han":"Lâu dài"}]',
        hinh_thuc_su_dung="Sử dụng riêng", thoi_han="Lâu dài",
        nguon_goc="Nhà nước công nhận", ngay_cap=date(2020, 7, 8),
        co_quan_cap="UBND huyện Ý Yên",
    )
    defaults.update(kwargs)
    return SimpleNamespace(**defaults)


def _node(person_id, *, role="Khac", will_receive=False, is_land_owner=False,
          hidden=False, deleted=False):
    return {
        "id": f"node_{person_id}", "personId": str(person_id),
        "role": role, "isLandOwner": is_land_owner,
        "willReceive": will_receive,
        "hidden": hidden, "deleted": deleted,
    }


def _case_state(*people, nodes=None):
    """Persisted workspace shape: stage list + engineInput projection (V2)."""
    stage = []
    for p in people:
        stage.append({
            "id": str(p.id), "ho_ten": p.ho_ten, "gioi_tinh": p.gioi_tinh,
            "ngay_sinh": p.ngay_sinh.isoformat() if p.ngay_sinh else None,
            "ngay_chet": p.ngay_chet.isoformat() if p.ngay_chet else None,
            "so_giay_to": p.so_giay_to,
            "ngay_cap": p.ngay_cap.isoformat() if p.ngay_cap else None,
            "noi_cap": p.noi_cap, "dia_chi": p.dia_chi,
            "loai_dia_chi": p.loai_dia_chi,
        })
    payload = {"schemaVersion": 2, "stage": stage, "assets": [],
               "diagram": {}}
    if nodes is not None:
        payload["diagram"] = {
            "engineInput": {"version": 2, "nodes": nodes},
            "engineResult": {"allocations": {}},
        }
    return json.dumps(payload, ensure_ascii=False)


def _person_row(stage_row, vitri):
    raw_id = stage_row.get("entity_id")
    if raw_id is None:
        legacy = str(stage_row.get("id") or "")
        raw_id = int(legacy) if legacy.isdigit() else None
    return SimpleNamespace(
        row_id=stage_row.get("row_id") or f"row-{raw_id}-{vitri}",
        customer_id=raw_id, vitri=vitri,
        ho_ten=stage_row.get("ho_ten"),
        gioi_tinh=stage_row.get("gioi_tinh"),
        ngay_sinh=stage_row.get("ngay_sinh"),
        ngay_chet=stage_row.get("ngay_chet"),
        so_giay_to=stage_row.get("so_giay_to"),
        ngay_cap=stage_row.get("ngay_cap"),
        noi_cap=stage_row.get("noi_cap"),
        dia_chi=stage_row.get("dia_chi"),
        place_of_origin=stage_row.get("place_of_origin"),
        loai_giay_to=stage_row.get("loai_giay_to"),
        loai_dia_chi=stage_row.get("loai_dia_chi"),
    )


def _asset_row(prop, vitri):
    raw = prop.land_rows_json
    rows = json.loads(raw) if raw else []
    return SimpleNamespace(
        row_id=f"asset-{vitri}", property_id=getattr(prop, "id", None),
        vitri=vitri, so_serial=prop.so_serial, so_vao_so=prop.so_vao_so,
        so_thua_dat=prop.so_thua_dat, so_to_ban_do=prop.so_to_ban_do,
        dia_chi=prop.dia_chi, loai_so=prop.loai_so,
        hinh_thuc_su_dung=prop.hinh_thuc_su_dung, thoi_han=prop.thoi_han,
        nguon_goc=prop.nguon_goc, ngay_cap=prop.ngay_cap,
        co_quan_cap=prop.co_quan_cap, land_rows_present=bool(raw),
        land_rows=[SimpleNamespace(
            vitri=index + 1, loaidat=row.get("loai_dat"),
            dientich=row.get("dien_tich"), thoihan=row.get("thoi_han"))
            for index, row in enumerate(rows) if isinstance(row, dict)],
    )


def _make_case(owner, *, participants=None, properties=None,
               case_state=None, cid=99):
    """Case Storage B: rows case-owned + JSON v4 (Diagram/meta)."""
    participants = participants or []
    props = properties or []
    payload = json.loads(case_state) if case_state else {}
    stage_items = payload.get("stage") if isinstance(payload.get("stage"), list) else None
    if stage_items is None:
        stage_items = []
        if owner is not None:
            stage_items.append({
                "id": str(owner.id), "ho_ten": owner.ho_ten,
                "gioi_tinh": owner.gioi_tinh,
                "ngay_sinh": owner.ngay_sinh.isoformat() if owner.ngay_sinh else None,
                "ngay_chet": owner.ngay_chet.isoformat() if owner.ngay_chet else None,
                "so_giay_to": owner.so_giay_to,
                "ngay_cap": owner.ngay_cap.isoformat() if owner.ngay_cap else None,
                "noi_cap": owner.noi_cap, "dia_chi": owner.dia_chi,
                "loai_dia_chi": owner.loai_dia_chi})
        for participant in participants:
            person = participant.customer
            stage_items.append({
                "id": str(person.id), "ho_ten": person.ho_ten,
                "gioi_tinh": person.gioi_tinh,
                "ngay_sinh": person.ngay_sinh.isoformat() if person.ngay_sinh else None,
                "ngay_chet": person.ngay_chet.isoformat() if person.ngay_chet else None,
                "so_giay_to": person.so_giay_to,
                "ngay_cap": person.ngay_cap.isoformat() if person.ngay_cap else None,
                "noi_cap": person.noi_cap, "dia_chi": person.dia_chi,
                "loai_dia_chi": person.loai_dia_chi})
    person_rows = [_person_row(item, index + 1)
                   for index, item in enumerate(stage_items)]
    entity_to_row = {str(row.customer_id): row.row_id
                     for row in person_rows if row.customer_id is not None}

    diagram = payload.get("diagram") or {}
    nodes = None
    engine = diagram.get("engineInput")
    if isinstance(engine, dict):
        nodes = engine.get("nodes") or []
    elif isinstance(diagram.get("state"), dict):
        nodes = diagram["state"].get("nodes") or []
    remapped = []
    for node in nodes or []:
        if isinstance(node, dict):
            mapped = entity_to_row.get(str(node.get("personId")))
            remapped.append({**node,
                             "personId": mapped or node.get("personId")})
    state = json.dumps({
        "schemaVersion": 4,
        "owner_row_id": entity_to_row.get(str(getattr(owner, "id", None))),
        "diagram": {"state": {"version": 3, "domain": "inheritance",
                              "nodes": remapped}},
    }, ensure_ascii=False)
    for participant in participants:
        participant.case_person_row_id = entity_to_row.get(
            str(participant.customer_id))
    return SimpleNamespace(
        id=cid,
        nguoi_chet=owner,
        tai_san=props[0] if props else None,
        ngay_lap_ho_so=date(2026, 7, 3),
        loai_van_ban="khai_nhan",
        trang_thai="draft",
        workspace_revision=1,
        noi_niem_yet="UBND xã Minh Tân",
        nguoi_nhan_uy_quyen=None,
        nguoi_nhan_uy_quyen_id=None,
        noi_dung_viec=None,
        ghi_chu="",
        participants=participants,
        person_rows=person_rows,
        asset_rows=[_asset_row(prop, index + 1)
                    for index, prop in enumerate(props)],
        case_state_json=state,
    )


def _ready_case(cid=99):
    owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3),
                    doc_no="001")
    heir = _person(2, "Nguyễn Văn Con", doc_no="002")
    return _make_case(
        owner, properties=[_prop()], cid=cid,
        case_state=_case_state(
            owner, heir,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=True),
            ]))


def _resolve_builtin(_key):
    return BUILTIN_TEMPLATE


def _resolve_none(_key):
    return None


def _batch(case, keys, dest_dir, **kw):
    kw.setdefault("resolve_template", _resolve_builtin)
    return wbe.export_batch(
        case, case_id=case.id, document_keys=keys, dest_dir=dest_dir, **kw)


def _docs_by_key(data):
    return {d["document_key"]: d for d in data["documents"]}


# ------------------------------------------------------------- options


class TestExportOptions:
    def test_ready_case_all_documents_listed(self):
        data = wbe.export_options(
            _ready_case(), resolve_template=_resolve_builtin)
        assert data["schema_version"] == "notary.case-drafting.v2"
        docs = _docs_by_key(data)
        assert set(docs) == {"khai_nhan_di_san", "thoa_thuan_phan_chia",
                             "niem_yet"}
        assert docs["khai_nhan_di_san"]["ready"] is True
        assert docs["khai_nhan_di_san"]["block_reason"] is None
        assert docs["thoa_thuan_phan_chia"]["ready"] is True
        assert docs["niem_yet"]["ready"] is False
        assert docs["niem_yet"]["block_reason"] == "word.template_missing"

    def test_empty_case_ready_and_exports_with_missing_note(self, tmp_path):
        # W6: thiếu dữ kiện không chặn — readiness chỉ còn mức kỹ thuật.
        owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3))
        case = _make_case(owner, properties=[], case_state=None)
        data = wbe.export_options(case, resolve_template=_resolve_builtin)
        docs = _docs_by_key(data)
        assert docs["khai_nhan_di_san"]["ready"] is True
        assert docs["khai_nhan_di_san"]["block_reason"] is None
        # niem_yet: template check trước data validation (mock parity)
        assert docs["niem_yet"]["block_reason"] == "word.template_missing"
        # Xuất thật: file ra + ghi thiếu tài sản/chủ đất/người nhận.
        batch = _batch(case, ["khai_nhan_di_san"], tmp_path)
        entry = _docs_by_key(batch)["khai_nhan_di_san"]
        assert entry["status"] == "saved"
        assert (tmp_path / entry["actual_filename"]).is_file()
        assert "tài sản" in entry["missing"]

    def test_missing_template_blocks_templated_docs(self):
        data = wbe.export_options(_ready_case(),
                                  resolve_template=_resolve_none)
        docs = _docs_by_key(data)
        assert docs["khai_nhan_di_san"]["block_reason"] == (
            "word.template_missing")
        assert docs["thoa_thuan_phan_chia"]["block_reason"] == (
            "word.template_missing")

    def test_no_receiver(self):
        owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3))
        child = _person(2, "Nguyễn Văn Con")
        case = _make_case(
            owner, properties=[_prop()],
            case_state=_case_state(
                owner, child,
                nodes=[
                    _node(1, role="Owner", is_land_owner=True),
                    _node(2, role="Con", will_receive=False),
                ]))
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=_resolve_builtin))
        # Mẫu builtin không dùng nhóm "Người nhận" → gate nhóm bỏ qua theo
        # template (owner 15/10/2026: readiness theo dữ liệu MẪU cần).
        assert docs["khai_nhan_di_san"]["block_reason"] is None

    def test_receiver_gate_only_when_template_needs_group(self, tmp_path):
        owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3))
        child = _person(2, "Nguyễn Văn Con")
        case = _make_case(
            owner, properties=[_prop()],
            case_state=_case_state(
                owner, child,
                nodes=[
                    _node(1, role="Owner", is_land_owner=True),
                    _node(2, role="Con", will_receive=False),
                ]))
        # Template TỰ tạo có token nhóm "Người nhận" → gate vẫn báo thiếu;
        # alias KHÔNG DẤU cùng key cũng phải nhận đúng như consumer
        # (`_lookup_placeholder` chuẩn hoá token trước khi khớp).
        needs_receiver = tmp_path / "needs_receiver.docx"
        doc = Document()
        doc.add_paragraph("[Người nhận 1 - Họ tên]")
        doc.save(str(needs_receiver))
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=lambda _k: needs_receiver))
        # W6: nhóm mẫu cần mà thiếu → file vẫn xuất, ghi thiếu người nhận.
        assert docs["khai_nhan_di_san"]["block_reason"] is None
        dest = tmp_path / "recv_out"
        dest.mkdir()
        batch = _batch(case, ["khai_nhan_di_san"], dest,
                       resolve_template=lambda _k: needs_receiver)
        entry = _docs_by_key(batch)["khai_nhan_di_san"]
        assert entry["status"] == "saved"
        assert "người nhận" in entry["missing"]

        needs_receiver_unaccented = tmp_path / "needs_receiver_ascii.docx"
        doc = Document()
        doc.add_paragraph("[Nguoi nhan 3 - Ho ten]")
        doc.save(str(needs_receiver_unaccented))
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=lambda _k: needs_receiver_unaccented))
        assert docs["khai_nhan_di_san"]["block_reason"] is None

    def test_stage_only_template_two_party_without_saved_diagram(self, tmp_path):
        a = _person(1, "Lê Văn A")
        b = _person(2, "Phạm Thị B")
        case = _make_case(
            a, properties=[_prop()],
            participants=[_participant(b, 'Ben B')],
            case_state=None)
        case.loai_van_ban = "chuyen_nhuong"
        stage_only = tmp_path / "stage_only.docx"
        doc = Document()
        doc.add_paragraph("[Người 1 - Họ tên] và [Người 2 - Họ tên]")
        doc.save(str(stage_only))
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=lambda _k: stage_only))
        # Mẫu chỉ cần Stage → không bắt buộc sơ đồ hai bên đã lưu.
        assert docs["khai_nhan_di_san"]["block_reason"] is None
        # Và xuất THẬT được: render cùng needs của mẫu, không rơi lại gate cũ.
        dest = tmp_path / "tp_out"
        dest.mkdir()
        data = _batch(case, ["khai_nhan_di_san"], dest,
                      resolve_template=lambda _k: stage_only)
        assert data["breakdown"]["succeeded"] == ["khai_nhan_di_san"]
        text = Document(
            data["documents"][0]["output_file"]["path"]).paragraphs[0].text
        assert "Lê Văn A" in text and "Phạm Thị B" in text

    def test_no_deceased_landowner(self):
        owner = _person(1, "Nguyễn Văn Chủ")          # còn sống
        child = _person(2, "Nguyễn Văn Con")
        case = _make_case(
            owner, properties=[_prop()],
            case_state=_case_state(
                owner, child,
                nodes=[
                    _node(1, role="Owner", is_land_owner=True),
                    _node(2, role="Con", will_receive=True),
                ]))
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=_resolve_builtin))
        # W6: chủ đất còn sống → file vẫn xuất, ghi thiếu chủ đất đã mất.
        assert docs["khai_nhan_di_san"]["block_reason"] is None

    def test_no_landowner_when_owner_hidden_in_diagram(self):
        owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3))
        child = _person(2, "Nguyễn Văn Con")
        case = _make_case(
            owner, properties=[_prop()],
            case_state=_case_state(
                owner, child,
                nodes=[
                    _node(1, role="Owner", is_land_owner=True,
                          hidden=True),          # owner ẩn → không active
                    _node(2, role="Con", will_receive=True),
                ]))
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=_resolve_builtin))
        # W6: chủ đất ẩn → file vẫn xuất, ghi thiếu chủ đất.
        assert docs["khai_nhan_di_san"]["block_reason"] is None

    def test_too_many_assets(self):
        owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3))
        heir = _person(2, "Nguyễn Văn Con")
        props = [_prop(dia_chi=f"Đất {i}", so_serial=f"S{i}")
                 for i in range(1, 7)]
        case = _make_case(
            owner, properties=props,
            case_state=_case_state(
                owner, heir,
                nodes=[
                    _node(1, role="Owner", is_land_owner=True),
                    _node(2, role="Con", will_receive=True),
                ]))
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=_resolve_builtin))
        assert docs["khai_nhan_di_san"]["block_reason"] == (
            "word.too_many_assets")

    def test_too_many_people(self):
        owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3))
        people = [owner]
        nodes = [_node(1, role="Owner", is_land_owner=True)]
        for i in range(2, 23):                       # 22 người active
            p = _person(i, f"Người {i}")
            people.append(p)
            nodes.append(_node(i, role="Con", will_receive=True))
        case = _make_case(owner, properties=[_prop()],
                          case_state=_case_state(*people, nodes=nodes))
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=_resolve_builtin))
        assert docs["khai_nhan_di_san"]["block_reason"] == (
            "word.too_many_people")


# ------------------------------------------------------------- batch


class TestExportBatch:
    def test_writes_independent_docx_per_document(self, tmp_path):
        data = _batch(_ready_case(),
                      ["khai_nhan_di_san", "thoa_thuan_phan_chia"],
                      tmp_path)
        assert data["schema_version"] == "notary.case-drafting.v2"
        assert data["destination"]["is_dir"] is True
        assert data["destination"]["scope"] == "machine_local"
        assert data["breakdown"] == {
            "succeeded": ["khai_nhan_di_san", "thoa_thuan_phan_chia"],
            "failed": [], "skipped": []}
        names = {d["actual_filename"] for d in data["documents"]}
        assert names == {
            "Van_ban_khai_nhan_di_san_HS-99.docx",
            "Thoa_thuan_phan_chia_di_san_HS-99.docx"}
        for d in data["documents"]:
            assert d["status"] == "saved"
            assert d["error"] is None
            assert re.match(r"^[A-Za-z0-9_-]+_HS-99(_\d+)?\.docx$",
                            d["actual_filename"])
            assert ".." not in d["actual_filename"]
            out = Path(d["output_file"]["path"])
            assert out.parent == tmp_path
            assert out.is_file()
            doc = Document(str(out))                 # DOCX mở được
            assert doc.paragraphs

    def test_collision_with_existing_files_keeps_old_bytes(self, tmp_path):
        old1 = tmp_path / "Van_ban_khai_nhan_di_san_HS-99.docx"
        old2 = tmp_path / "Van_ban_khai_nhan_di_san_HS-99_2.docx"
        old1.write_bytes(b"cu-1")
        old2.write_bytes(b"cu-2")
        data = _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path)
        doc = data["documents"][0]
        assert doc["actual_filename"] == (
            "Van_ban_khai_nhan_di_san_HS-99_3.docx")
        # file cũ nguyên byte — KHÔNG BAO GIỜ ghi đè
        assert old1.read_bytes() == b"cu-1"
        assert old2.read_bytes() == b"cu-2"

    def test_intra_batch_same_filename_stem_no_clash(
            self, tmp_path, monkeypatch):
        """Hai document_key trỏ cùng filename_stem → reservation nội
        batch → file thứ hai hậu tố _2."""
        spec = dataclasses.replace(
            wbe.DOC_CATALOG_BY_KEY["thoa_thuan_phan_chia"],
            filename_stem="Van_ban_khai_nhan_di_san")
        monkeypatch.setitem(
            wbe.DOC_CATALOG_BY_KEY, "thoa_thuan_phan_chia", spec)
        data = _batch(_ready_case(),
                      ["khai_nhan_di_san", "thoa_thuan_phan_chia"],
                      tmp_path)
        names = sorted(d["actual_filename"] for d in data["documents"])
        assert names == [
            "Van_ban_khai_nhan_di_san_HS-99.docx",
            "Van_ban_khai_nhan_di_san_HS-99_2.docx"]

    def test_partial_saved_and_failed(self, tmp_path):
        data = _batch(_ready_case(),
                      ["khai_nhan_di_san", "niem_yet"], tmp_path)
        docs = _docs_by_key(data)
        assert docs["khai_nhan_di_san"]["status"] == "saved"
        assert docs["niem_yet"]["status"] == "failed"
        assert docs["niem_yet"]["error"]["code"] == "word.template_missing"
        assert data["breakdown"]["succeeded"] == ["khai_nhan_di_san"]
        assert data["breakdown"]["failed"] == ["niem_yet"]
        assert data["breakdown"]["skipped"] == []
        # file của doc thành công tồn tại dù doc kia lỗi
        assert (tmp_path
                / "Van_ban_khai_nhan_di_san_HS-99.docx").is_file()

    def test_all_failed_raises_batch_failed_with_result(self, tmp_path):
        # W6: thiếu dữ kiện không còn làm file lỗi — dùng trần kỹ thuật
        # (6 tài sản) + văn bản chưa có template để có lượt all-failed.
        owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3))
        heir = _person(2, "Nguyễn Văn Con")
        props = [_prop(dia_chi=f"Đất {i}", so_serial=f"S{i}")
                 for i in range(1, 7)]
        over = _make_case(
            owner, properties=props,
            case_state=_case_state(
                owner, heir,
                nodes=[
                    _node(1, role="Owner", is_land_owner=True),
                    _node(2, role="Con", will_receive=True),
                ]))
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(over, ["khai_nhan_di_san", "niem_yet"], tmp_path)
        err = exc.value
        assert err.code == "word_batch_failed"
        rd = err.result_data
        assert rd["breakdown"] == {
            "succeeded": [],
            "failed": ["khai_nhan_di_san", "niem_yet"],
            "skipped": []}
        codes = {d["error"]["code"] for d in rd["documents"]}
        assert codes == {"word.too_many_assets", "word.template_missing"}
        # không file nào được tạo
        assert list(tmp_path.iterdir()) == []

    def test_template_missing_when_no_template(self, tmp_path):
        # Doc duy nhat fail -> all-failed -> word_batch_failed kem result.
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path,
                   resolve_template=_resolve_none)
        assert exc.value.code == "word_batch_failed"
        doc = exc.value.result_data["documents"][0]
        assert doc["status"] == "failed"
        assert doc["error"]["code"] == "word.template_missing"
        assert list(tmp_path.iterdir()) == []

    def test_unresolved_placeholders_export_with_note(self, tmp_path):
        # W8: token lạ không chặn file — xuất bản còn token + ghi kết quả.
        bad_template = tmp_path / "bad_template.docx"
        doc = Document()
        doc.add_paragraph("Văn bản thử — [Trường Chưa Hỗ Trợ]")
        doc.save(str(bad_template))
        data = _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: bad_template)
        result = data["documents"][0]
        assert result["status"] == "saved"
        assert result["error"] is None
        assert any("[Trường Chưa Hỗ Trợ]" in note
                   for note in result["missing"])
        out = Document(str(tmp_path / result["actual_filename"]))
        assert "[Trường Chưa Hỗ Trợ]" in out.paragraphs[0].text

    def test_unresolved_placeholder_in_header_export_with_note(self, tmp_path):
        """Placeholder ngoài body (header) cũng chỉ ghi thiếu, không giấu."""
        bad_template = tmp_path / "bad_header_template.docx"
        doc = Document()
        doc.add_paragraph("Thân văn bản")
        doc.sections[0].header.paragraphs[0].add_run("[Trường Chưa Hỗ Trợ]")
        doc.save(str(bad_template))
        data = _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: bad_template)
        result = data["documents"][0]
        assert result["status"] == "saved"
        assert any("[Trường Chưa Hỗ Trợ]" in note
                   for note in result["missing"])

    def test_verified_fixed_form_codes_do_not_block_export(self, tmp_path):
        """Mã in cứng đã xác minh của ĐÚNG mẫu DK18 không bị coi là thiếu."""
        case = _ready_case()
        mapping = word_engine.build_template_mapping(case)
        real_token = next(
            token for token, value in sorted(mapping.items())
            if value and token.startswith("[") and not token.startswith("[Tài sản")
        )
        template = tmp_path / "2. DK18, thuế  1.docx"
        doc = Document()
        doc.add_paragraph("[01] Nội dung mục một")
        doc.add_paragraph("[06.1] Nội dung mục sáu")
        doc.add_paragraph(real_token)
        doc.save(str(template))
        data = _batch(case, ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: template)
        docs = _docs_by_key(data)
        assert docs["khai_nhan_di_san"]["status"] == "saved"
        assert list(tmp_path.glob("*_HS-*.docx"))

    def test_numeric_tokens_on_unverified_template_export_with_note(
            self, tmp_path):
        """Mẫu bất kỳ chứa "[01]"/"[54.99]" KHÔNG được miễn — ghi vào kết quả."""
        template = tmp_path / "codes_template.docx"
        doc = Document()
        doc.add_paragraph("[01] Mục chưa xác minh")
        doc.add_paragraph("[54.99] Ngoài bộ DK18")
        doc.save(str(template))
        data = _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: template)
        result = data["documents"][0]
        assert result["status"] == "saved"
        note = next(n for n in result["missing"] if "chưa hỗ trợ" in n)
        assert "[01]" in note
        assert "[54.99]" in note

    def test_fixed_code_outside_verified_list_reported_but_exported(
            self, tmp_path):
        """Mẫu DK18 đã xác minh nhưng chứa mã NGOÀI danh sách → ghi đúng mã."""
        case = _ready_case()
        mapping = word_engine.build_template_mapping(case)
        real_token = next(
            token for token, value in sorted(mapping.items())
            if value and token.startswith("[") and not token.startswith("[Tài sản")
        )
        template = tmp_path / "2. DK18, thuế  2.docx"
        doc = Document()
        doc.add_paragraph("[01] Mục một")
        doc.add_paragraph("[54.99] Mã lạ")
        doc.add_paragraph(real_token)
        doc.save(str(template))
        data = _batch(case, ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: template)
        result = data["documents"][0]
        assert result["status"] == "saved"
        note = next(n for n in result["missing"] if "chưa hỗ trợ" in n)
        assert "[54.99]" in note
        assert "[01]" not in note

    def _uq_template(self, path, slot=1):
        doc = Document()
        doc.add_paragraph(f"tôi uỷ quyền cho: [Dòng người được ủy quyền {slot}]")
        doc.save(str(path))
        return path

    def _case_with_authorized_via_saved_role(self):
        """Người được ủy quyền chỉ gán qua participants.vai_tro (không node
        Diagram) — đúng tình huống luồng nhập lưu vai trò đã xác nhận."""
        owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3), doc_no="001")
        heir = _person(2, "Nguyễn Văn Con", doc_no="002")
        proxy = _person(3, "Trần Thị Ủy", gender="Nữ", doc_no="003")
        return _make_case(
            owner,
            participants=[
                _participant(heir, "Con"),
                _participant(proxy, "Người được ủy quyền", receive=False),
            ],
            properties=[_prop()],
            case_state=_case_state(
                owner, heir, proxy,
                nodes=[
                    _node(1, role="Owner", is_land_owner=True),
                    _node(2, role="Con", will_receive=True),
                ]))

    def test_template_requiring_authorized_person_exports_with_note(
            self, tmp_path):
        """W8: mẫu đòi người được ủy quyền mà hồ sơ chưa gán → file vẫn ra,
        phần đó trống, kết quả ghi thiếu slot."""
        template = self._uq_template(tmp_path / "3. UQ.docx")
        data = _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: template)
        result = data["documents"][0]
        assert result["status"] == "saved"
        assert "người được ủy quyền 1" in result["missing"]

    def test_template_authorized_slot_2_exports_with_note(self, tmp_path):
        """Mẫu đòi slot 2 nhưng hồ sơ chỉ gán một → ghi thiếu slot 2."""
        template = self._uq_template(tmp_path / "3. UQ.docx", slot=2)
        case = self._case_with_authorized_via_saved_role()
        data = _batch(case, ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: template)
        result = data["documents"][0]
        assert result["status"] == "saved"
        assert "người được ủy quyền 2" in result["missing"]

    def test_template_authorized_inline_slot_2_exports_with_note(
            self, tmp_path):
        """Mẫu dùng "[Người được ủy quyền inline 2]" với 1 đại diện → ghi
        thiếu slot 2 (trước đây inline 2 bị nhận nhầm slot 1)."""
        template = tmp_path / "3. UQ.docx"
        doc = Document()
        doc.add_paragraph(
            "đại diện thứ hai: [Người được ủy quyền inline 2]")
        doc.save(str(template))
        case = self._case_with_authorized_via_saved_role()
        data = _batch(case, ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: template)
        result = data["documents"][0]
        assert result["status"] == "saved"
        assert "người được ủy quyền 2" in result["missing"]

    def test_template_authorized_person_saved_from_persisted_role(
            self, tmp_path):
        """Vai trò ủy quyền đã lưu (không node) -> export được và đầu ra chứa
        ĐÚNG người đó, không chỉ qua được vòng kiểm unresolved."""
        template = self._uq_template(tmp_path / "3. UQ.docx")
        case = self._case_with_authorized_via_saved_role()
        data = _batch(case, ["khai_nhan_di_san"], tmp_path,
                      resolve_template=lambda _k: template)
        entry = _docs_by_key(data)["khai_nhan_di_san"]
        assert entry["status"] == "saved"
        out = Document(str(tmp_path / entry["actual_filename"]))
        assert "Trần Thị Ủy" in out.paragraphs[0].text

    def test_cancel_mid_batch_keeps_saved_and_marks_skipped(self, tmp_path):
        class _Cancel(Exception):
            def __init__(self, data):
                super().__init__("cancel")
                self.data = data

        def _check(data):
            # Raise khi đã có doc saved mà còn doc pending → tương đương
            # cancel flag của jobstore bật giữa văn bản 1 và 2.
            if data["breakdown"]["succeeded"] \
                    and data["breakdown"]["skipped"]:
                raise _Cancel(data)

        keys = ["khai_nhan_di_san", "thoa_thuan_phan_chia", "niem_yet"]
        with pytest.raises(_Cancel) as exc:
            _batch(_ready_case(), keys, tmp_path, check_cancel=_check)
        data = exc.value.data
        docs = _docs_by_key(data)
        assert docs["khai_nhan_di_san"]["status"] == "saved"
        assert docs["thoa_thuan_phan_chia"]["status"] == "skipped"
        assert docs["niem_yet"]["status"] == "skipped"
        assert data["breakdown"]["succeeded"] == ["khai_nhan_di_san"]
        assert data["breakdown"]["skipped"] == [
            "thoa_thuan_phan_chia", "niem_yet"]
        # file đã lưu giữ nguyên; file chưa bắt đầu không được tạo
        files = sorted(p.name for p in tmp_path.glob("*.docx"))
        assert files == ["Van_ban_khai_nhan_di_san_HS-99.docx"]

    def test_traversal_generated_filename_fails_document(
            self, tmp_path, monkeypatch):
        """Defense-in-depth: stem chứa `..`/separator → word_path_traversal
        (code đã register §8.3), không file nào được tạo."""
        spec = dataclasses.replace(
            wbe.DOC_CATALOG_BY_KEY["khai_nhan_di_san"],
            filename_stem="..\\evil")
        monkeypatch.setitem(
            wbe.DOC_CATALOG_BY_KEY, "khai_nhan_di_san", spec)
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path)
        assert exc.value.code == "word_batch_failed"
        result = exc.value.result_data["documents"][0]
        assert result["status"] == "failed"
        assert result["error"]["code"] == "word_path_traversal"
        assert list(tmp_path.iterdir()) == []

    def test_unsafe_generated_filename_fails_document(
            self, tmp_path, monkeypatch):
        """Stem xấu nhưng không traversal → word.invalid_filename."""
        spec = dataclasses.replace(
            wbe.DOC_CATALOG_BY_KEY["khai_nhan_di_san"],
            filename_stem="bad:name")
        monkeypatch.setitem(
            wbe.DOC_CATALOG_BY_KEY, "khai_nhan_di_san", spec)
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path)
        assert exc.value.code == "word_batch_failed"
        result = exc.value.result_data["documents"][0]
        assert result["status"] == "failed"
        assert result["error"]["code"] == "word.invalid_filename"
        assert list(tmp_path.iterdir()) == []

    def test_same_name_directory_in_destination_increments(
            self, tmp_path, monkeypatch):
        """Directory trùng tên file đích → tăng `_2`, không fail doc."""
        (tmp_path / "Van_ban_khai_nhan_di_san_HS-99.docx").mkdir()
        data = _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path)
        assert data["documents"][0]["status"] == "saved"
        assert data["documents"][0]["actual_filename"].endswith("_2.docx")

    def test_progress_reported_per_document(self, tmp_path):
        calls = []
        _batch(_ready_case(),
               ["khai_nhan_di_san", "thoa_thuan_phan_chia"], tmp_path,
               report_progress=lambda d, t, label: calls.append((d, t)))
        assert calls == [(1, 2), (2, 2)]

    def test_cancel_at_progress_checkpoint_keeps_pending_result(
            self, tmp_path):
        """jobstore.report_progress cung la cancel checkpoint va raise
        CancelledByUser() KHONG kem result — batch phai re-check de
        pending snapshot (doc chua bat dau -> skipped) van len wire
        (contract §8.4 canceled example, MIN-115)."""
        class _Cancel(Exception):
            def __init__(self, data=None):
                super().__init__("cancel")
                self.data = data

        flag = {"canceled": False}

        def _check(data):
            if flag["canceled"]:
                raise _Cancel(data)

        def _report(done, _total, _label):
            # Gia lap jobstore: progress quan sat cancel flag -> raise
            # exception KHONG kem result (bare).
            if done == 2:
                flag["canceled"] = True
                raise _Cancel()

        keys = ["khai_nhan_di_san", "thoa_thuan_phan_chia", "niem_yet"]
        with pytest.raises(_Cancel) as exc:
            _batch(_ready_case(), keys, tmp_path,
                   check_cancel=_check, report_progress=_report)
        data = exc.value.data
        assert data is not None           # pending snapshot len wire
        docs = _docs_by_key(data)
        assert docs["khai_nhan_di_san"]["status"] == "saved"
        assert docs["thoa_thuan_phan_chia"]["status"] == "saved"
        assert docs["niem_yet"]["status"] == "skipped"
        assert data["breakdown"]["skipped"] == ["niem_yet"]


# ------------------------------------------- MIN-116 dest deny-write


def _deny_writes_under(monkeypatch, root, *, name_filter=None):
    """Giả lập destination bị ACL deny-write (Windows): mọi open/create
    ghi dưới `root` raise PermissionError ngay lần đầu — deterministic
    thay vì icacls thật.

    Vá cả ba seam: builtins.open (probe/temp/publish), io.open (zipfile
    của python-docx save) và os.open (tempfile.mkstemp — đường code cũ).
    `name_filter` (tuỳ chọn) giới hạn deny theo basename — dùng cho
    kịch bản TOCTOU (probe pass nhưng publish/render bị deny).
    """
    import builtins
    import io as _io
    root = str(root)
    real_open = open
    real_io_open = _io.open
    real_os_open = os.open

    def _is_write_mode(mode):
        return any(c in mode for c in ("w", "a", "x", "+"))

    def _hit(file):
        s = str(file)
        return s.startswith(root) and (
            name_filter is None or name_filter(os.path.basename(s)))

    def _open_guard(file, mode="r", *args, **kwargs):
        if _hit(file) and _is_write_mode(mode):
            raise PermissionError(13, "Access is denied", str(file))
        return real_open(file, mode, *args, **kwargs)

    def _io_open_guard(file, mode="r", *args, **kwargs):
        if _hit(file) and _is_write_mode(mode):
            raise PermissionError(13, "Access is denied", str(file))
        return real_io_open(file, mode, *args, **kwargs)

    def _os_open_guard(path, flags, *args, **kwargs):
        if _hit(path) and (flags & (os.O_WRONLY | os.O_RDWR)):
            raise PermissionError(13, "Access is denied", str(path))
        return real_os_open(path, flags, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", _open_guard)
    monkeypatch.setattr(_io, "open", _io_open_guard)
    monkeypatch.setattr(os, "open", _os_open_guard)


class TestDestWriteDenied:
    """MIN-116: destination ACL deny-write → job fail fast với per-doc
    file_locked, KHÔNG hang trong tempfile.mkstemp retry loop."""

    def test_denied_destination_fails_fast_all_docs_file_locked(
            self, tmp_path, monkeypatch):
        """Probe writability fail → mọi doc (kể cả template_missing)
        nhận per-doc file_locked; all-failed → word_batch_failed kèm
        result.data; hoàn tất <5s — trước fix path này hang >60s."""
        _deny_writes_under(monkeypatch, tmp_path)
        keys = ["khai_nhan_di_san", "thoa_thuan_phan_chia", "niem_yet"]
        t0 = time.monotonic()
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), keys, tmp_path)
        assert time.monotonic() - t0 < 5
        assert exc.value.code == "word_batch_failed"
        rd = exc.value.result_data
        assert rd["breakdown"] == {
            "succeeded": [], "failed": keys, "skipped": []}
        for d in rd["documents"]:
            assert d["status"] == "failed"
            assert d["error"]["code"] == "file_locked"
            assert d["actual_filename"] is None
            assert d["output_file"] is None
        # không file output/probe/temp nào sót lại trong dest
        assert list(tmp_path.iterdir()) == []

    def test_denied_batch_cancel_marks_pending_skipped(
            self, tmp_path, monkeypatch):
        """Cancel giữa batch deny-write: doc đã xử lý giữ file_locked,
        doc chưa bắt đầu → skipped (pending-data semantics MIN-115)."""
        class _Cancel(Exception):
            def __init__(self, data):
                super().__init__("cancel")
                self.data = data

        _deny_writes_under(monkeypatch, tmp_path)
        keys = ["khai_nhan_di_san", "thoa_thuan_phan_chia", "niem_yet"]

        def _check(data):
            # Cancel flag bật sau doc đầu: failed=[khai], còn pending.
            if data["breakdown"]["failed"] \
                    and data["breakdown"]["skipped"]:
                raise _Cancel(data)

        with pytest.raises(_Cancel) as exc:
            _batch(_ready_case(), keys, tmp_path, check_cancel=_check)
        docs = _docs_by_key(exc.value.data)
        assert docs["khai_nhan_di_san"]["status"] == "failed"
        assert docs["khai_nhan_di_san"]["error"]["code"] == "file_locked"
        assert docs["thoa_thuan_phan_chia"]["status"] == "skipped"
        assert docs["niem_yet"]["status"] == "skipped"
        assert exc.value.data["breakdown"]["skipped"] == [
            "thoa_thuan_phan_chia", "niem_yet"]
        assert list(tmp_path.iterdir()) == []

    def test_cancel_before_probe_skips_all(self, tmp_path):
        """Cancel trước khi probe destination → mọi doc skipped, không
        file nào được tạo."""
        class _Cancel(Exception):
            def __init__(self, data):
                super().__init__("cancel")
                self.data = data

        def _check(data):
            raise _Cancel(data)

        with pytest.raises(_Cancel) as exc:
            _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path,
                   check_cancel=_check)
        data = exc.value.data
        assert data["breakdown"]["skipped"] == ["khai_nhan_di_san"]
        assert data["documents"][0]["status"] == "skipped"
        assert list(tmp_path.iterdir()) == []

    def test_toctou_publish_denied_maps_file_locked(
            self, tmp_path, monkeypatch):
        """Probe + tạo temp OK nhưng quyền bị thu hồi trước lúc publish
        (TOCTOU): chỉ file đích `*_HS-*.docx` bị deny → per-doc
        file_locked, file tạm được dọn, không hang."""
        _deny_writes_under(
            monkeypatch, tmp_path,
            name_filter=lambda name: "_HS-" in name)
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path)
        assert exc.value.code == "word_batch_failed"
        doc = exc.value.result_data["documents"][0]
        assert doc["status"] == "failed"
        assert doc["error"]["code"] == "file_locked"
        assert list(tmp_path.iterdir()) == []

    def test_toctou_render_denied_maps_file_locked(
            self, tmp_path, monkeypatch):
        """Probe OK nhưng `_render_temp` gặp PermissionError (quyền thu
        hồi sau probe) → per-doc file_locked, không propagate thành
        lỗi job khô hay hang."""
        def _denied(_doc, _dest):
            raise PermissionError(13, "Access is denied")
        monkeypatch.setattr(wbe, "_render_temp", _denied)
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path)
        assert exc.value.code == "word_batch_failed"
        doc = exc.value.result_data["documents"][0]
        assert doc["status"] == "failed"
        assert doc["error"]["code"] == "file_locked"
        assert list(tmp_path.iterdir()) == []


# ---------------------------------------------------- payload validation


class TestDocumentKeysValidation:
    @pytest.mark.parametrize("bad", [None, "khai_nhan_di_san", 3,
                                     {"a": 1}, (("khai_nhan_di_san"),)])
    def test_non_list_keys_validation_error(self, tmp_path, bad):
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), bad, tmp_path)
        assert exc.value.code == "validation_error"

    def test_empty_keys(self, tmp_path):
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), [], tmp_path)
        assert exc.value.code == "word_no_documents_selected"

    def test_duplicate_keys(self, tmp_path):
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(),
                   ["khai_nhan_di_san", "khai_nhan_di_san"], tmp_path)
        assert exc.value.code == "word_duplicate_document_key"
        assert exc.value.details["document_key"] == "khai_nhan_di_san"

    @pytest.mark.parametrize("key", ["khong_co", "BAD-KEY", 42])
    def test_unknown_or_malformed_keys(self, tmp_path, key):
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), [key], tmp_path)
        assert exc.value.code == "word_unknown_document_key"

    def test_validation_happens_before_any_file(self, tmp_path):
        for keys in ([], ["khai_nhan_di_san", "khai_nhan_di_san"],
                     ["khong_co"]):
            with pytest.raises(wbe.WordBatchError):
                _batch(_ready_case(), keys, tmp_path)
        assert list(tmp_path.iterdir()) == []

    def test_missing_destination_dir(self, tmp_path):
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["khai_nhan_di_san"],
                   tmp_path / "khong-ton-tai")
        assert exc.value.code == "file_not_found"


# ------------------------------------------------- filename safety unit


class TestFilenameSafety:
    @pytest.mark.parametrize("name", [
        "CON.docx", "NUL.txt", "LPT1.docx",
        "a\\b_HS-99.docx", "a/b_HS-99.docx", "a:b_HS-99.docx",
        "..\\evil_HS-99.docx", "a..b_HS-99.docx",
        "a?_HS-99.docx", "a*_HS-99.docx",
        "trail _HS-99.docx ", ".dot_HS-99.docx", "name_HS-99.docx.",
    ])
    def test_unsafe_windows_names_rejected(self, name):
        assert wbe.is_safe_windows_filename(name) is False

    @pytest.mark.parametrize("name", [
        "Van_ban_khai_nhan_di_san_HS-99.docx",
        "Thoa_thuan_phan_chia_di_san_HS-99_12.docx",
        # "CON_HS-99" khong phai reserved name (chi stem dung "CON" moi bi cam)
        "CON_HS-99.docx",
    ])
    def test_safe_names_accepted(self, name):
        assert wbe.is_safe_windows_filename(name) is True


# ================================================================ NAIA-8
# Folder mode — mẫu từ thư mục người dùng chọn (template_dir).

import shutil
import subprocess


def _folder_templates(tmp_path, names):
    """Copy template builtin vào tmp dưới các tên cho trước."""
    tdir = tmp_path / "templates"
    tdir.mkdir()
    for name in names:
        shutil.copyfile(BUILTIN_TEMPLATE, tdir / name)
    return tdir


def _dir_link(link, target):
    """Directory link `link` → `target`: os.symlink trước, fallback
    junction `mklink /J` trên Windows (không cần quyền admin)."""
    try:
        os.symlink(target, link, target_is_directory=True)
        return
    except OSError:
        pass
    if os.name == "nt":
        r = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(target)],
            capture_output=True)
        if r.returncode == 0:
            return
    pytest.skip("không tạo được directory link (thiếu quyền)")


def _file_link(link, target):
    """File symlink — Windows cần SeCreateSymbolicLinkPrivilege
    (WinError 1314 trên máy không có quyền) → skip."""
    try:
        os.symlink(target, link)
    except OSError as exc:
        pytest.skip(f"không tạo được file symlink: {exc}")


class TestScanTemplateDir:
    def test_scans_only_top_level_docx_sorted(self, tmp_path):
        tdir = _folder_templates(
            tmp_path, ["b mau.docx", "A mau.docx", "~$locked.docx"])
        (tdir / "readme.txt").write_text("x")
        (tdir / "sub").mkdir()
        (tdir / "sub" / "inner.docx").write_text("x")
        specs = wbe.scan_template_dir(tdir)
        assert [s.document_key for s in specs] == [
            "A mau.docx", "b mau.docx"]
        assert [s.display_name for s in specs] == ["A mau", "b mau"]

    def test_missing_dir_raises(self, tmp_path):
        with pytest.raises(wbe.WordBatchError) as exc:
            wbe.scan_template_dir(tmp_path / "nope")
        assert exc.value.code == "file_not_found"

    def test_ascii_stem_sanitized(self):
        assert wbe._ascii_stem("Mẫu đơn xin.docx") == "Mau_don_xin"
        assert wbe._ascii_stem("CON.docx") == "Mau"


class TestFolderModeExport:
    def test_static_folder_exports_saved_stage_without_receiver(self, tmp_path):
        case = _ready_case()
        case.case_state_json = None
        case.participants = []
        tdir = tmp_path / "static_templates"
        tdir.mkdir()
        doc = Document()
        doc.add_paragraph("[Người 1 - Họ tên] | [Người nhận 1 - Họ tên]")
        doc.save(tdir / "static.docx")
        specs = wbe.scan_template_dir(tdir)
        resolve = wbe.template_path_resolver(tdir)
        options = wbe.export_options(case, resolve_template=resolve, doc_specs=specs)
        assert options["documents"][0]["ready"] is True
        dest = tmp_path / "static_output"
        dest.mkdir()
        result = wbe.export_batch(case, case_id=case.id,
            document_keys=["static.docx"], dest_dir=dest,
            resolve_template=resolve, doc_specs=specs)
        assert result["breakdown"]["succeeded"] == ["static.docx"]
        output = Document(result["documents"][0]["output_file"]["path"])
        assert case.person_rows[0].ho_ten in output.paragraphs[0].text
        assert "[" not in output.paragraphs[0].text

    def test_options_lists_folder_files(self, tmp_path):
        tdir = _folder_templates(tmp_path, ["Mẫu thuê nhà.docx"])
        specs = wbe.scan_template_dir(tdir)
        data = wbe.export_options(
            _ready_case(), resolve_template=wbe.template_path_resolver(tdir),
            doc_specs=specs)
        docs = _docs_by_key(data)
        assert set(docs) == {"Mẫu thuê nhà.docx"}
        assert docs["Mẫu thuê nhà.docx"]["display_name"] == "Mẫu thuê nhà"
        assert docs["Mẫu thuê nhà.docx"]["ready"] is True

    def test_batch_writes_ascii_stem_outputs(self, tmp_path):
        tdir = _folder_templates(
            tmp_path, ["Mẫu đơn.docx", "b mau.docx"])
        dest = tmp_path / "out"
        dest.mkdir()
        specs = wbe.scan_template_dir(tdir)
        data = _batch(
            _ready_case(), ["Mẫu đơn.docx", "b mau.docx"], dest,
            resolve_template=wbe.template_path_resolver(tdir),
            doc_specs=specs)
        assert data["breakdown"]["succeeded"] == [
            "Mẫu đơn.docx", "b mau.docx"]
        names = sorted(p.name for p in dest.iterdir())
        assert names == ["Mau_don_HS-99.docx", "b_mau_HS-99.docx"]
        for d in data["documents"]:
            assert d["status"] == "saved"
            assert Path(d["output_file"]["path"]).is_file()

    def test_batch_rejects_unknown_or_traversal_key(self, tmp_path):
        tdir = _folder_templates(tmp_path, ["a.docx"])
        dest = tmp_path / "out"
        dest.mkdir()
        specs = wbe.scan_template_dir(tdir)
        resolver = wbe.template_path_resolver(tdir)
        for bad in ("../other.docx", "khai_nhan_di_san",
                    "chua_quet.docx", r"..\win.docx"):
            with pytest.raises(wbe.WordBatchError) as exc:
                _batch(_ready_case(), [bad], dest,
                       resolve_template=resolver, doc_specs=specs)
            assert exc.value.code == "word_unknown_document_key"

    def test_create_case_subdir(self, tmp_path):
        tdir = _folder_templates(tmp_path, ["mau.docx"])
        dest = tmp_path / "out"
        dest.mkdir()
        specs = wbe.scan_template_dir(tdir)
        data = _batch(
            _ready_case(), ["mau.docx"], dest,
            resolve_template=wbe.template_path_resolver(tdir),
            doc_specs=specs, create_case_subdir=True)
        subdir = dest / "HS-99"
        assert data["documents"][0]["status"] == "saved"
        assert (subdir / "mau_HS-99.docx").is_file()
        # contract §8.2: data.destination vẫn là thư mục người dùng chọn
        assert data["destination"]["path"] == str(dest)
        # Không ghi đè: chạy lại cùng mẫu → suffix _2
        data2 = _batch(
            _ready_case(), ["mau.docx"], dest,
            resolve_template=wbe.template_path_resolver(tdir),
            doc_specs=specs, create_case_subdir=True)
        assert data2["documents"][0]["actual_filename"] == "mau_HS-99_2.docx"

    def test_template_deleted_after_scan_fails_file(self, tmp_path):
        tdir = _folder_templates(tmp_path, ["mau.docx", "mau2.docx"])
        dest = tmp_path / "out"
        dest.mkdir()
        specs = wbe.scan_template_dir(tdir)
        resolver = wbe.template_path_resolver(tdir)
        (tdir / "mau.docx").unlink()
        data = _batch(
            _ready_case(), ["mau.docx", "mau2.docx"], dest,
            resolve_template=resolver, doc_specs=specs)
        docs = _docs_by_key(data)
        assert docs["mau.docx"]["status"] == "failed"
        assert docs["mau.docx"]["error"]["code"] == "word.template_missing"
        assert docs["mau2.docx"]["status"] == "saved"
        assert data["breakdown"]["failed"] == ["mau.docx"]

    def test_catalog_mode_still_works(self, tmp_path):
        dest = tmp_path / "out"
        dest.mkdir()
        data = _batch(_ready_case(), ["khai_nhan_di_san"], dest)
        assert data["documents"][0]["status"] == "saved"


class TestDestSubdirLinkEscape:
    """Review NAIA-8: `HS-<case_id>` tồn tại sẵn dạng junction/symlink
    trỏ ra ngoài — mkdir(exist_ok=True) chấp nhận vì "đã có dir" và mọi
    ghi sau đó thoát khỏi thư mục được cấp quyền. Phải kiểm tra đích
    THẬT và từ chối trước probe/ghi/publish."""

    def test_case_subdir_link_to_outside_rejected(self, tmp_path):
        dest = tmp_path / "selected"
        dest.mkdir()
        outside = tmp_path / "outside"
        outside.mkdir()
        _dir_link(dest / "HS-99", outside)
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["khai_nhan_di_san"], dest,
                   create_case_subdir=True)
        assert exc.value.code == "word_path_traversal"
        # không file nào được tạo — ngoài outside lẫn trong selected
        assert list(outside.iterdir()) == []
        assert [p.name for p in dest.iterdir()] == ["HS-99"]

    def test_case_subdir_link_inside_dest_allowed(self, tmp_path):
        """Link trỏ vào TRONG destination (không thoát) vẫn được chấp
        nhận — chỉ chặn liên kết thoát ra ngoài."""
        dest = tmp_path / "selected"
        dest.mkdir()
        inner = dest / "inner"
        inner.mkdir()
        _dir_link(dest / "HS-99", inner)
        data = _batch(_ready_case(), ["khai_nhan_di_san"], dest,
                      create_case_subdir=True)
        assert data["documents"][0]["status"] == "saved"
        assert (inner
                / "Van_ban_khai_nhan_di_san_HS-99.docx").is_file()
        # file đi qua link vào inner — vẫn nằm trong destination
        assert sorted(p.name for p in dest.iterdir()) == ["HS-99", "inner"]


class TestTemplateLinkEscape:
    """Mẫu qua liên kết: `.docx` trong thư mục mẫu là symlink trỏ ra
    ngoài — `is_file()` đuổi theo link nên basename check không đủ;
    phải kiểm containment khi quét lẫn khi resolve để mở."""

    def test_scan_skips_file_link_to_outside(self, tmp_path):
        tdir = tmp_path / "templates"
        tdir.mkdir()
        shutil.copyfile(BUILTIN_TEMPLATE, tdir / "mau.docx")
        outside = tmp_path / "outside.docx"
        shutil.copyfile(BUILTIN_TEMPLATE, outside)
        _file_link(tdir / "linked.docx", outside)
        specs = wbe.scan_template_dir(tdir)
        assert [s.document_key for s in specs] == ["mau.docx"]
        resolver = wbe.template_path_resolver(tdir)
        assert resolver("linked.docx") is None
        assert resolver("mau.docx") is not None
        # key của file liên kết cũng bị từ chối ở tầng validate
        dest = tmp_path / "out"
        dest.mkdir()
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["linked.docx"], dest,
                   resolve_template=resolver, doc_specs=specs)
        assert exc.value.code == "word_unknown_document_key"
        assert list(dest.iterdir()) == []

    def test_template_swapped_for_link_after_scan(self, tmp_path):
        """Mẫu quét được lúc scan bị thay bằng symlink ra ngoài trước khi
        mở — resolver kiểm tra lại containment → word.template_missing,
        không bao giờ đọc file ngoài thư mục được cấp quyền."""
        tdir = tmp_path / "templates"
        tdir.mkdir()
        target = tdir / "mau.docx"
        shutil.copyfile(BUILTIN_TEMPLATE, target)
        outside = tmp_path / "outside.docx"
        doc = Document()                       # nội dung khác hẳn —
        doc.add_paragraph("SECRET-OUTSIDE")    # nếu bị đọc lộ sẽ thấy
        doc.save(str(outside))
        specs = wbe.scan_template_dir(tdir)
        resolver = wbe.template_path_resolver(tdir)
        assert resolver("mau.docx") is not None
        # swap sau khi quét: xóa file thật, đặt symlink trỏ ra ngoài
        target.unlink()
        _file_link(target, outside)
        assert resolver("mau.docx") is None
        dest = tmp_path / "out"
        dest.mkdir()
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["mau.docx"], dest,
                   resolve_template=resolver, doc_specs=specs)
        assert exc.value.code == "word_batch_failed"
        d = exc.value.result_data["documents"][0]
        assert d["status"] == "failed"
        assert d["error"]["code"] == "word.template_missing"
        assert list(dest.iterdir()) == []

    def test_dir_link_in_templates_skipped(self, tmp_path):
        """Junction/symlink THƯ MỤC có đuôi .docx trỏ ra ngoài: không
        phải file → loại khỏi scan (và không đuổi link)."""
        tdir = tmp_path / "templates"
        tdir.mkdir()
        outside_dir = tmp_path / "outside_dir"
        outside_dir.mkdir()
        _dir_link(tdir / "linkdir.docx", outside_dir)
        specs = wbe.scan_template_dir(tdir)
        assert specs == []


def _folder(tmp_path, files):
    """Thư mục mẫu + (specs, resolver) như export_options/export_batch."""
    tdir = tmp_path / "templates"
    tdir.mkdir()
    for name, paragraphs in files.items():
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.save(tdir / name)
    return tdir, wbe.scan_template_dir(tdir), wbe.template_path_resolver(tdir)


def _no_receiver_case():
    owner = _person(1, "Nguyễn Văn Chủ", dead=date(2024, 2, 3))
    child = _person(2, "Nguyễn Văn Con")
    return _make_case(
        owner, properties=[_prop()],
        case_state=_case_state(
            owner, child,
            nodes=[
                _node(1, role="Owner", is_land_owner=True),
                _node(2, role="Con", will_receive=False),
            ]))


def _two_party_case():
    a = _person(1, "Lê Văn A")
    b = _person(2, "Phạm Thị B")
    case = _make_case(
        a, properties=[_prop()], case_state=None,
        participants=[_participant(b, 'Ben B')])
    case.loai_van_ban = "chuyen_nhuong"
    return case


class TestTemplateGroupClassification:
    """Gate nhóm theo token mẫu thực dùng (owner 15/10/2026) — kiểm ở
    biên consumer (export_options/export_batch), không pin helper nội bộ."""

    def test_metadata_token_shapes_do_not_demand_receiver(self, tmp_path):
        """Token metadata chứa "người nhận" nhưng KHÔNG phải key nhóm
        Người nhận thì không được kéo gate `word.no_receiver`;
        `[nguoinhanuyquyen]` là key metadata thật (xuất được), còn
        `[Người nhận ủy quyền]` chưa hỗ trợ → per-file unresolved."""
        tdir, specs, resolve = _folder(tmp_path, {
            "metadata.docx": ["[nguoinhanuyquyen] — [Người 1 - Họ tên]"],
            "metadata_unsupported.docx": ["[Người nhận ủy quyền]"],
        })
        docs = _docs_by_key(wbe.export_options(
            _no_receiver_case(), resolve_template=resolve, doc_specs=specs))
        assert docs["metadata.docx"]["block_reason"] is None
        assert docs["metadata_unsupported.docx"]["block_reason"] is None
        assert all(d["block_reason"] != "word.no_receiver"
                   for d in docs.values())

        dest = tmp_path / "out"
        dest.mkdir()
        case = _no_receiver_case()
        data = _batch(case, ["metadata.docx"], dest,
                      resolve_template=resolve, doc_specs=specs)
        assert data["breakdown"]["succeeded"] == ["metadata.docx"]
        assert "[" not in Document(
            data["documents"][0]["output_file"]["path"]).paragraphs[0].text

    def test_two_party_slots_template_without_diagram_exports_blank(
            self, tmp_path):
        """W6: mẫu dùng ghế `[Bên A N …]` khi chưa có diagram state vẫn
        xuất được — ghế trống, kết quả ghi thiếu sơ đồ hai bên."""
        tdir, specs, resolve = _folder(tmp_path, {
            "slots.docx": ["[Bên A 1 - Họ tên]"],
            "names.docx": ["[Tên bên A] và [Tên bên B]"],
        })
        case = _two_party_case()
        docs = _docs_by_key(wbe.export_options(
            case, resolve_template=resolve, doc_specs=specs))
        assert docs["slots.docx"]["block_reason"] is None
        assert docs["names.docx"]["block_reason"] is None

        dest = tmp_path / "out"
        dest.mkdir()
        data = _batch(case, ["slots.docx"], dest,
                      resolve_template=resolve, doc_specs=specs)
        entry = data["documents"][0]
        assert entry["status"] == "saved"
        assert "sơ đồ hai bên đã lưu" in entry["missing"]
        out = Document(str(dest / entry["actual_filename"]))
        assert out.paragraphs[0].text.strip() == ""

    def test_folder_stage_only_two_party_exports_without_diagram(self, tmp_path):
        """AC1 folder mode: mẫu chỉ cần Stage xuất được khi chưa có
        diagram state — render dùng đúng needs của mẫu."""
        tdir, specs, resolve = _folder(tmp_path, {
            "stage_only.docx": ["[Người 1 - Họ tên] và [Người 2 - Họ tên]"],
        })
        case = _two_party_case()
        options = wbe.export_options(
            case, resolve_template=resolve, doc_specs=specs)
        assert options["documents"][0]["ready"] is True
        dest = tmp_path / "out"
        dest.mkdir()
        data = _batch(case, ["stage_only.docx"], dest,
                      resolve_template=resolve, doc_specs=specs)
        assert data["breakdown"]["succeeded"] == ["stage_only.docx"]
        text = Document(
            data["documents"][0]["output_file"]["path"]).paragraphs[0].text
        assert "Lê Văn A" in text and "Phạm Thị B" in text

    def test_folder_group_missing_stays_empty_not_fail(self, tmp_path):
        """Folder mode giữ semantics §3.3: token nhóm thiếu dữ liệu render
        rỗng (không thêm gate/fail mới)."""
        tdir, specs, resolve = _folder(tmp_path, {
            "needs_receiver.docx": ["[Người nhận 1 - Họ tên] | [Người 1 - Họ tên]"],
        })
        case = _no_receiver_case()
        options = wbe.export_options(
            case, resolve_template=resolve, doc_specs=specs)
        assert options["documents"][0]["ready"] is True
        dest = tmp_path / "out"
        dest.mkdir()
        data = _batch(case, ["needs_receiver.docx"], dest,
                      resolve_template=resolve, doc_specs=specs)
        assert data["breakdown"]["succeeded"] == ["needs_receiver.docx"]
        text = Document(
            data["documents"][0]["output_file"]["path"]).paragraphs[0].text
        assert "[" not in text
        assert "Nguyễn Văn Chủ" in text


class TestFolderTemplateErrors:
    def test_malformed_template_reports_per_file_code(self, tmp_path):
        """Mẫu hỏng: readiness trả đúng data-code per-file mà batch trả —
        không nuốt lỗi thành gate dữ liệu cũ; nguồn không đổi."""
        tdir = tmp_path / "templates"
        tdir.mkdir()
        bad = tdir / "hong.docx"
        bad.write_bytes(b"khong-phai-zip-docx")
        before = bad.read_bytes()
        specs = wbe.scan_template_dir(tdir)
        resolve = wbe.template_path_resolver(tdir)
        options = wbe.export_options(
            _ready_case(), resolve_template=resolve, doc_specs=specs)
        assert options["documents"][0]["ready"] is False
        assert options["documents"][0]["block_reason"] == "word.render_failed"

        dest = tmp_path / "out"
        dest.mkdir()
        with pytest.raises(wbe.WordBatchError) as exc:
            _batch(_ready_case(), ["hong.docx"], dest,
                   resolve_template=resolve, doc_specs=specs)
        assert exc.value.code == "word_batch_failed"
        entry = exc.value.result_data["documents"][0]
        assert entry["status"] == "failed"
        assert entry["error"]["code"] == "word.render_failed"
        assert bad.read_bytes() == before
        assert list(dest.iterdir()) == []

    def test_mixed_batch_reports_per_file_states(self, tmp_path):
        """Batch trộn: MỌI file lưu; token lạ/thiếu ủy quyền/slot ngoài
        dải thành ghi chú `missing` riêng từng file; nguồn không đổi."""
        tdir, specs, resolve = _folder(tmp_path, {
            "good.docx": ["[Người 1 - Họ tên]"],
            "unknown.docx": ["Ký hiệu: [Ký hiệu lạ XYZ]"],
            "auth_short.docx": ["[Người được ủy quyền 2 - Họ tên]"],
            "beyond.docx": ["[Người 25 - Họ tên]"],
        })
        hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                  for p in tdir.iterdir()}
        dest = tmp_path / "out"
        dest.mkdir()
        case = _ready_case()
        data = _batch(case, [s.document_key for s in specs], dest,
                      resolve_template=resolve, doc_specs=specs)
        by_key = {d["document_key"]: d for d in data["documents"]}
        # W6/W8: cả 4 file đều xuất; thiếu dữ kiện/token lạ thành ghi chú.
        assert {d["status"] for d in data["documents"]} == {"saved"}
        assert by_key["good.docx"]["missing"] == []
        assert any("[Ký hiệu lạ XYZ]" in n
                   for n in by_key["unknown.docx"]["missing"])
        assert "người được ủy quyền 2" in by_key["auth_short.docx"]["missing"]
        assert any("[Người 25 - Họ tên]" in n
                   for n in by_key["beyond.docx"]["missing"])
        assert sorted(data["breakdown"]["succeeded"]) == [
            "auth_short.docx", "beyond.docx", "good.docx", "unknown.docx"]
        assert data["breakdown"]["failed"] == []
        assert {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                for p in tdir.iterdir()} == hashes
        assert sorted(p.name for p in dest.iterdir()) == [
            "auth_short_HS-99.docx", "beyond_HS-99.docx",
            "good_HS-99.docx", "unknown_HS-99.docx"]


@pytest.fixture()
def synthetic_limit_templates(tmp_path):
    folder = tmp_path / "templates"
    folder.mkdir()
    for key in ("khai_nhan_di_san", "thoa_thuan_phan_chia"):
        doc = Document()
        doc.add_paragraph("[Người 1 - Họ tên] | [Người 20 - Họ tên]")
        doc.save(folder / f"{key}.docx")
    return folder


@pytest.fixture()
def inheritance_people_case():
    def make(count):
        people = [_person(i, f"Người thử {i:02}") for i in range(1, count + 1)]
        return _make_case(
            people[0],
            properties=[_prop(so_serial="", so_thua_dat="901",
                              so_to_ban_do="09")],
            case_state=_case_state(*people, nodes=[]))
    return make


@pytest.mark.parametrize("count", [20, 21, 25])
@pytest.mark.parametrize("folder_mode", [False, True], ids=["catalog", "folder"])
def test_inheritance_people_export_limit_per_document(
        tmp_path, synthetic_limit_templates, inheritance_people_case,
        count, folder_mode):
    case = inheritance_people_case(count)
    folder = synthetic_limit_templates
    specs = wbe.scan_template_dir(folder) if folder_mode else None
    keys = ([spec.document_key for spec in specs] if folder_mode else
            ["khai_nhan_di_san", "thoa_thuan_phan_chia"])
    resolve = (wbe.template_path_resolver(folder) if folder_mode else
               lambda key: folder / f"{key}.docx")
    options = _docs_by_key(wbe.export_options(
        case, resolve_template=resolve, doc_specs=specs))
    for key in keys:
        assert options[key]["ready"] is (count == 20)
        assert options[key]["block_reason"] == (
            None if count == 20 else "word.too_many_people")

    destination = tmp_path / "output"
    destination.mkdir()
    if count == 20:
        result = _batch(case, keys, destination,
                        resolve_template=resolve, doc_specs=specs)
        assert result["breakdown"]["succeeded"] == keys
        for entry in result["documents"]:
            text = Document(entry["output_file"]["path"]).paragraphs[0].text
            assert text == "Người thử 01 | Người thử 20"
        return

    with pytest.raises(wbe.WordBatchError) as failure:
        _batch(case, keys, destination,
               resolve_template=resolve, doc_specs=specs)
    assert failure.value.code == "word_batch_failed"
    result = failure.value.result_data
    assert result["breakdown"] == {"succeeded": [], "failed": keys, "skipped": []}
    assert len(result["documents"]) == 2
    for entry in result["documents"]:
        assert entry["status"] == "failed"
        assert entry["output_file"] is None
        assert entry["error"] == {
            "code": "word.too_many_people", "message": "Quá 20 người trong Stage"}
    assert list(destination.iterdir()) == []


def test_export_limit_exception_uses_catalog_message(inheritance_people_case):
    context = word_engine.build_word_context(inheritance_people_case(25))
    code, message = wbe._doc_error_from(
        word_engine.WordExportValidationError("Danh sách Stage vượt quá 20 người"),
        context, static=True)
    assert (code, message) == ("word.too_many_people", "Quá 20 người trong Stage")


def test_export_error_message_has_no_python_exception_prefix(tmp_path):
    def resolve(_key):
        raise ValueError("Không thể đọc mẫu")

    with pytest.raises(wbe.WordBatchError) as failure:
        _batch(_ready_case(), ["khai_nhan_di_san"], tmp_path,
               resolve_template=resolve)
    entry, = failure.value.result_data["documents"]
    assert entry["error"] == {
        "code": "word.render_failed", "message": "Không thể đọc mẫu"}
