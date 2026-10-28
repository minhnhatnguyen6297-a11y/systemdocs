"""Tests cho services.word_template_migration (NAIA-12)."""
from __future__ import annotations

import json
import os
import shutil
from pathlib import Path

import pytest
from docx import Document

from services import word_template_migration as wtm

SRC = Path(__file__).resolve().parent.parent / "word_templates"


def _docx(path: Path, paragraphs: list[str]) -> Path:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def test_is_template_candidate_filters_temp_lock_and_workbook(tmp_path):
    assert wtm.is_template_candidate(_docx(tmp_path / "mau.docx", ["x"]))
    assert wtm.is_template_candidate(_docx(tmp_path / "mau.DOCX", ["x"]))
    assert wtm.is_template_candidate(_docx(tmp_path / "mau.doc", ["x"]))
    # AC1: chỉ loại lock/temp/workbook — " - Copy" hay "_old" vẫn là mẫu thật.
    assert wtm.is_template_candidate(_docx(tmp_path / "mau - Copy.docx", ["x"]))
    assert wtm.is_template_candidate(_docx(tmp_path / "mau_backup.docx", ["x"]))

    assert not wtm.is_template_candidate(_docx(tmp_path / "~$mau.docx", ["x"]))
    assert not wtm.is_template_candidate(tmp_path / "so_theo_doi.xlsx")
    assert not wtm.is_template_candidate(tmp_path / "du_lieu.xlsb")
    assert not wtm.is_template_candidate(tmp_path / "huong_dan.pdf")


def test_scan_source_templates_reports_excluded_files(tmp_path):
    _docx(tmp_path / "a.docx", ["x"])
    _docx(tmp_path / "nested" / "b.docx", ["y"])
    _docx(tmp_path / "~$lock.docx", ["z"])
    (tmp_path / "xlsx").mkdir()
    (tmp_path / "xlsx" / "data.xlsx").write_bytes(b"PK")

    kept, excluded = wtm.scan_source_templates(tmp_path)

    assert [p.name for p in kept] == ["a.docx", "b.docx"]
    assert "~$lock.docx" in excluded
    assert "xlsx/data.xlsx" in excluded


def test_real_docm_is_excluded_and_reported_before_any_write(tmp_path, monkeypatch):
    import zipfile

    src, out = tmp_path / "src", tmp_path / "out"
    macro = _docx(src / "macro.docm", ["macro template"])
    with zipfile.ZipFile(macro) as archive:
        entries = {n: archive.read(n) for n in archive.namelist()}
    entries["[Content_Types].xml"] = entries["[Content_Types].xml"].replace(
        b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        b"application/vnd.ms-word.document.macroEnabled.main+xml")
    with zipfile.ZipFile(macro, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    # Genuine macro-enabled main content, not DOCX with a changed suffix.
    with pytest.raises(ValueError):
        Document(macro)
    _docx(src / "supported.docx", ["DOCX"])
    (src / "legacy.doc").write_bytes(b"legacy-doc")
    before = _bytes(src)
    kept, excluded = wtm.scan_source_templates(src)
    assert {p.name for p in kept} == {"supported.docx", "legacy.doc"}
    assert excluded == ["macro.docm"]
    assert not wtm.is_template_candidate(macro)
    warn = wtm.warnings.warn
    warned = []

    def before_write_warning(message, *args, **kwargs):
        assert not out.exists()
        warned.append(str(message))
        warn(message, *args, **kwargs)

    monkeypatch.setattr(wtm.warnings, "warn", before_write_warning)
    converted = []

    def convert(source, target):
        converted.append(source.name)
        _docx(target, ["converted DOC"])

    with pytest.warns(UserWarning, match="macro.docm.*DOCM.*DOCX.*DOC"):
        manifest = wtm.migrate_templates(src, out, convert=convert)
    assert len(warned) == 1
    assert converted == ["legacy.doc"]
    assert manifest["excluded_files"] == ["macro.docm"]
    assert manifest["file_count"] == 2
    assert not (out / "macro.docx").exists()
    assert Document(out / "supported.docx").paragraphs[0].text == "DOCX"
    assert Document(out / "legacy.docx").paragraphs[0].text == "converted DOC"
    assert _bytes(src) == before


def test_diff_tokens_classifies_renamed_kept_fixed_no_source_and_unmapped():
    tokens = [
        "[Niêm Yết]",
        "[Thời hạn]",
        "[01]",
        "[06.1]",
        "[Số công chứng]",
        "[Trường Lạ Hoắc]",
        "[Họ tên chủ đất]",
    ]
    result = wtm.diff_tokens(
        tokens,
        registry=["[Họ tên chủ đất]", "[Thời hạn 1]"],
        no_source=["[Không có nguồn]"],
        verified_codes={"[01]", "[06.1]"},
    )
    by_old = {item.old: item for item in result}

    assert by_old["[Niêm Yết]"].status == "renamed"
    assert by_old["[Niêm Yết]"].new == "[Nơi niêm yết]"
    assert by_old["[Thời hạn]"].status == "renamed"
    assert by_old["[Thời hạn]"].new == "[Thời hạn 1]"
    assert by_old["[Họ tên chủ đất]"].status == "kept"
    assert by_old["[01]"].status == "fixed_code"
    assert by_old["[06.1]"].status == "fixed_code"
    # Owner đã gỡ token khỏi bộ mẫu -> "dropped", không phải chờ nguồn dữ liệu.
    assert by_old["[Số công chứng]"].status == "dropped"
    assert by_old["[Trường Lạ Hoắc]"].status == "unmapped"


def test_diff_tokens_numeric_code_unmapped_without_verified_allowlist():
    """Mã dạng số trên mẫu CHƯA xác minh là unmapped, không được gán fixed_code."""
    result = wtm.diff_tokens(
        ["[01]", "[06.1]", "[54.99]", "[Họ tên chủ đất]"],
        registry=["[Họ tên chủ đất]"],
        no_source=[],
        verified_codes=(),
    )
    by_old = {item.old: item for item in result}
    assert by_old["[01]"].status == "unmapped"
    assert by_old["[06.1]"].status == "unmapped"
    assert by_old["[54.99]"].status == "unmapped"
    assert by_old["[Họ tên chủ đất]"].status == "kept"


def test_diff_tokens_counts_each_occurrence():
    result = wtm.diff_tokens(
        ["[Niêm Yết]", "[Niêm Yết]", "[01]"], registry=[], no_source=[]
    )
    by_old = {item.old: item for item in result}
    assert by_old["[Niêm Yết]"].count == 2
    assert by_old["[01]"].count == 1


def test_diff_tokens_alias_map_is_normalization_insensitive():
    result = wtm.diff_tokens(
        ["[Niêm yết]", "[Niêm   Yết]"],
        registry=[],
        no_source=[],
    )
    assert all(item.status == "renamed" for item in result)
    assert {item.new for item in result} == {"[Nơi niêm yết]"}


def test_migrate_templates_never_mutates_source_and_writes_manifest(tmp_path):
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["[Niêm Yết] và [Thời hạn]"])
    before = wtm.sha256_file(src / "mau.docx")

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=["[Thời hạn 1]"], no_source=[]
    )

    assert wtm.sha256_file(src / "mau.docx") == before
    assert manifest["file_count"] == 1
    entry = manifest["files"][0]
    assert entry["source_sha256"] == before
    assert entry["status_counts"].get("renamed") == 2
    assert (tmp_path / "out" / "mau.docx").exists()
    assert (tmp_path / "out" / "migration_manifest.json").exists()


def test_migrate_templates_renames_old_tokens_to_catalog_names(tmp_path):
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["Mo so [Niêm Yết], thoi han [Thời hạn]"])

    wtm.migrate_templates(src, tmp_path / "out", registry=["[Thời hạn 1]"], no_source=[])

    text = "\n".join(p.text for p in Document(str(tmp_path / "out" / "mau.docx")).paragraphs)
    assert "[Nơi niêm yết]" in text
    assert "[Thời hạn 1]" in text
    assert "[Niêm Yết]" not in text
    assert "[Thời hạn]" not in text


def test_migrate_templates_keeps_fixed_form_codes_on_verified_template(tmp_path):
    """Đúng stem DK18 đã đối chiếu bộ mẫu thật -> mã in cứng được ghi fixed_code."""
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế  1.docx", ["Mục [01]", "Mục [06.1]"])

    manifest = wtm.migrate_templates(src, tmp_path / "out", registry=[], no_source=[])

    text = "\n".join(
        p.text for p in Document(str(tmp_path / "out" / "2. DK18, thuế  1.docx")).paragraphs
    )
    assert "[01]" in text and "[06.1]" in text
    assert manifest["files"][0]["status_counts"] == {"fixed_code": 2}


def test_migrate_templates_numeric_codes_unmapped_on_unverified_template(tmp_path):
    """Mẫu bất kỳ chứa [01]/[54.99] KHÔNG được miễn — manifest phải báo unmapped."""
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["Mục [01]", "Mục [54.99]"])

    manifest = wtm.migrate_templates(src, tmp_path / "out", registry=[], no_source=[])

    statuses = {t["old"]: t["status"] for t in manifest["files"][0]["tokens"]}
    assert statuses["[01]"] == "unmapped"
    assert statuses["[54.99]"] == "unmapped"


def test_migrate_templates_no_apply_only_inventories(tmp_path):
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["[Niêm Yết]"])

    wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[], apply_renames=False
    )

    text = "\n".join(p.text for p in Document(str(tmp_path / "out" / "mau.docx")).paragraphs)
    assert "[Niêm Yết]" in text


def _bytes(dirpath: Path) -> dict[str, bytes]:
    return {
        p.relative_to(dirpath).as_posix(): p.read_bytes()
        for p in sorted(dirpath.rglob("*")) if p.is_file()
    }


def test_migrate_templates_refuses_same_source_and_dest(tmp_path):
    """migrate_templates(src, src) phải từ chối TRƯỚC khi ghi — bytes nguồn giữ nguyên."""
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["[Niêm Yết]"])
    before = _bytes(src)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, src, registry=[], no_source=[])

    assert _bytes(src) == before
    assert not (src / "migration_manifest.json").exists()


def test_migrate_templates_refuses_dest_inside_source(tmp_path):
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["[Niêm Yết]"])
    before = _bytes(src)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, src / "out", registry=[], no_source=[])

    assert _bytes(src) == before


def test_migrate_templates_refuses_source_inside_dest(tmp_path):
    src = tmp_path / "out" / "src"
    _docx(src / "mau.docx", ["x"])
    out = tmp_path / "out"
    before_src = _bytes(src)
    before_out = _bytes(out)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, out, registry=[], no_source=[])

    assert _bytes(src) == before_src
    assert _bytes(out) == before_out


def test_migrate_templates_refuses_existing_dest_file(tmp_path):
    """File đích đã tồn tại -> FileExistsError, không ghi đè, nguồn nguyên vẹn."""
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["[Niêm Yết]"])
    out = tmp_path / "out"
    _docx(out / "mau.docx", ["dữ liệu cũ không được đụng"])
    before_src = _bytes(src)
    before_out = _bytes(out)

    with pytest.raises(FileExistsError):
        wtm.migrate_templates(src, out, registry=[], no_source=[])

    assert _bytes(src) == before_src
    assert _bytes(out) == before_out


def test_migrate_templates_refuses_doc_docx_stem_collision(tmp_path):
    """cu.doc + cu.docx cùng đổ về cu.docx ở đích -> chặn trước mọi ghi."""
    src = tmp_path / "src"
    src.mkdir(parents=True)
    (src / "cu.doc").write_bytes(b"legacy")
    _docx(src / "cu.docx", ["x"])
    before = _bytes(src)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, tmp_path / "out", registry=[], no_source=[])

    assert _bytes(src) == before
    assert not (tmp_path / "out").exists() or not _bytes(tmp_path / "out")


def _make_dir_link(link: Path, target: Path) -> bool:
    """Tạo junction (mklink /J — không cần quyền admin) rồi fallback symlink."""
    import os
    import subprocess

    link.parent.mkdir(parents=True, exist_ok=True)
    if os.name == "nt":
        result = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(target)],
            capture_output=True)
        if result.returncode == 0:
            return True
    try:
        os.symlink(str(target), str(link), target_is_directory=True)
        return True
    except (OSError, NotImplementedError):
        return False


def test_migrate_templates_refuses_symlinked_dest(tmp_path):
    """Đích qua symlink/junction trỏ vào nguồn -> realpath bắt trùng -> từ chối."""
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["[Niêm Yết]"])
    link = tmp_path / "linked"
    if not _make_dir_link(link, src):
        pytest.skip("môi trường không tạo được junction/symlink thư mục")
    before = _bytes(src)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, link, registry=[], no_source=[])

    assert _bytes(src) == before


def test_migrate_templates_refuses_output_escaping_via_junction(tmp_path):
    """out/sub là junction sang 'outside' -> đầu ra resolve ra ngoài đích ->
    từ chối TRƯỚC mọi ghi, 'outside' và 'out' giữ nguyên bytes.

    File đích chưa tồn tại nên ``exists()``/``is_symlink()`` không bắt được —
    chỉ containment trên realpath mới chặn được (Leader tái hiện NAIA-12).
    """
    src = tmp_path / "src"
    _docx(src / "sub" / "a.docx", ["[Niêm Yết]"])
    outside = tmp_path / "outside"
    outside.mkdir()
    out = tmp_path / "out"
    out.mkdir()
    if not _make_dir_link(out / "sub", outside):
        pytest.skip("môi trường không tạo được junction/symlink thư mục")
    before_src = _bytes(src)
    before_outside = _bytes(outside)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, out, registry=[], no_source=[])

    assert _bytes(src) == before_src
    assert _bytes(outside) == before_outside
    assert not (outside / "a.docx").exists()
    assert not (outside / "migration_manifest.json").exists()


def test_migrate_templates_refuses_existing_manifest(tmp_path):
    """migration_manifest.json đã có sẵn ở đích -> FileExistsError, không ghi đè."""
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["[Niêm Yết]"])
    out = tmp_path / "out"
    out.mkdir()
    manifest = out / "migration_manifest.json"
    manifest.write_text('{"old": true}', encoding="utf-8")
    before_out = _bytes(out)

    with pytest.raises(FileExistsError):
        wtm.migrate_templates(src, out, registry=[], no_source=[])

    assert _bytes(out) == before_out
    assert manifest.read_text(encoding="utf-8") == '{"old": true}'


def test_scan_source_templates_excludes_files_via_junctioned_dir(tmp_path):
    """src/sub là junction sang 'outside' -> file bên trong (file tự nó không
    phải link) không được đọc làm nguồn — liệt kê vào excluded."""
    outside = tmp_path / "outside"
    _docx(outside / "leak.docx", ["ngoài phạm vi"])
    src = tmp_path / "src"
    src.mkdir()
    _docx(src / "mau.docx", ["[Niêm Yết]"])
    if not _make_dir_link(src / "sub", outside):
        pytest.skip("môi trường không tạo được junction/symlink thư mục")

    kept, excluded = wtm.scan_source_templates(src)

    assert [p.name for p in kept] == ["mau.docx"]
    assert "sub/leak.docx" in excluded


def test_manifest_counts_real_occurrences_not_deduplicated(tmp_path):
    """[Niêm Yết] xuất hiện 2 lần trong file -> count=2 và tổng occurrence=2."""
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["Mo so [Niêm Yết], nhac lai [Niêm Yết]"])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    entry = manifest["files"][0]
    token = next(t for t in entry["tokens"] if t["old"] == "[Niêm Yết]")
    assert token["count"] == 2          # số LẦN XUẤT HIỆN, không phải 1
    assert token["applied"] == 2        # cả hai lần đều đã đổi tên
    assert entry["occurrence_counts"]["_occurrences"] == 2
    assert manifest["totals"]["occurrences"] == 2

    text = "\n".join(
        p.text for p in Document(str(tmp_path / "out" / "mau.docx")).paragraphs)
    assert text.count("[Nơi niêm yết]") == 2


def test_powershell_converter_disables_macros_and_closes_safely():
    """AC1: AutomationSecurity=3 + tắt UpdateLinks TRƯỚC Open; doc/Word đóng
    trong finally, kể cả khi lỗi giữa chừng."""
    script = wtm._POWERSHELL_CONVERT
    assert script.index("AutomationSecurity = 3") < script.index("Documents.Open")
    assert "UpdateLinksAtOpen = $false" in script
    assert "$true" in script  # mở read-only (tham số ReadOnly)
    assert "$doc.Close(0)" in script
    assert "finally" in script
    assert "$word.Quit()" in script
    assert "ReleaseComObject" in script


def test_convert_doc_to_docx_refuses_existing_dest(tmp_path):
    src = tmp_path / "cu.doc"
    src.write_bytes(b"x")
    dst = tmp_path / "cu.docx"
    dst.write_bytes("đích cũ".encode("utf-8"))

    with pytest.raises(FileExistsError):
        wtm.convert_doc_to_docx(src, dst)
    assert dst.read_bytes() == "đích cũ".encode("utf-8")


def test_migrate_templates_converts_doc_through_injected_converter(tmp_path):
    src = tmp_path / "src"
    src.mkdir(parents=True)
    (src / "cu.doc").write_bytes(b"legacy-binary")

    calls = []

    def fake_convert(source, target):
        calls.append((Path(source).name, Path(target).name))
        _docx(Path(target), ["[Niêm Yết]"])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[], convert=fake_convert
    )

    # Converter ghi tên tạm NGẪU NHIÊN trong scratch dir của run — run
    # chuyển bytes qua fd 'xb' của chính nó để chứng minh ownership đích
    # bằng fstat handle tạo (review NAIA-12), không để converter ghi thẳng
    # ra path đích dễ bị chen giữa.
    assert len(calls) == 1
    assert calls[0][0] == "cu.doc"
    assert calls[0][1].startswith(".") and calls[0][1].endswith(".conv.docx")
    assert manifest["files"][0]["converted_from_doc"] is True
    out = tmp_path / "out"
    assert (out / "cu.docx").exists()
    # Staging giữ lại bên trong `.naia12-scratch-*` (không unlink theo
    # path — repair-83e6d03) và không bao giờ lọt vào danh sách mẫu.
    scratch = [p for p in out.iterdir()
               if p.name.startswith(".naia12-scratch-")]
    assert len(scratch) == 1 and scratch[0].is_dir()
    staged = [p for p in scratch[0].iterdir() if p.name == calls[0][1]]
    assert len(staged) == 1
    from services import word_batch_export as wbe
    scanned = [s.document_key for s in wbe.scan_template_dir(out)]
    assert calls[0][1] not in scanned
    assert "cu.docx" in scanned


def test_convert_doc_to_docx_raises_when_word_missing(tmp_path, monkeypatch):
    src = tmp_path / "cu.doc"
    src.write_bytes(b"x")

    def boom(*args, **kwargs):
        raise FileNotFoundError("no powershell")

    monkeypatch.setattr("subprocess.run", boom)
    with pytest.raises(wtm.WordUnavailableError):
        wtm.convert_doc_to_docx(src, tmp_path / "cu.docx")


def test_main_cli_writes_manifest_and_prints_summary(tmp_path, capsys):
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["[Niêm Yết]"])
    registry = tmp_path / "registry.json"
    registry.write_text(json.dumps(["[Nơi niêm yết]"]), encoding="utf-8")

    code = wtm.main([str(src), str(tmp_path / "out"), "--registry-json", str(registry)])

    assert code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["mau.docx"]["renamed"] == 1


@pytest.mark.skipif(not SRC.exists(), reason="không có bộ mẫu thừa kế")
def test_real_legacy_templates_are_inventoried_without_touching_source(tmp_path):
    hashes = {p: wtm.sha256_file(p) for p in wtm.scan_source_templates(SRC)[0]}

    manifest = wtm.migrate_templates(
        SRC, tmp_path / "out", registry=[], no_source=wtm.NO_SOURCE_PLACEHOLDERS
    )

    assert manifest["file_count"] >= 1
    for path, digest in hashes.items():
        assert wtm.sha256_file(path) == digest, f"file nguồn bị sửa: {path}"
    for entry in manifest["files"]:
        assert entry["tokens"], entry["source_name"]


@pytest.mark.skipif(not SRC.exists(), reason="không có bộ mẫu thừa kế")
def test_real_legacy_templates_report_every_unmapped_token(tmp_path):
    """Mọi token chưa map phải xuất hiện trong manifest, không bị bỏ im lặng."""
    manifest = wtm.migrate_templates(
        SRC,
        tmp_path / "out",
        registry=[],
        no_source=wtm.NO_SOURCE_PLACEHOLDERS,
        apply_renames=False,
    )
    reported = {
        token["old"]
        for entry in manifest["files"]
        for token in entry["tokens"]
        if token["status"] in {"unmapped", "no_source"}
    }
    assert reported, "phải báo cáo token chưa map để Leader quyết định"
    assert all(token.startswith("[") and token.endswith("]") for token in reported)


@pytest.mark.skipif(not SRC.exists(), reason="không có bộ mẫu thừa kế")
def test_real_legacy_templates_apply_catalog_renames(tmp_path):
    """Token có alias catalog §8 phải được đổi tên trong file đích, không sửa nguồn."""
    src_files, _ = wtm.scan_source_templates(SRC)
    before = {p: wtm.sha256_file(p) for p in src_files}

    manifest = wtm.migrate_templates(
        SRC, tmp_path / "out", registry=[], no_source=wtm.NO_SOURCE_PLACEHOLDERS
    )

    renamed = [
        (entry["output_name"], token)
        for entry in manifest["files"]
        for token in entry["tokens"]
        if token["status"] == "renamed"
    ]
    assert renamed, "bộ mẫu phải có token khớp alias catalog §8"
    texts: dict[str, str] = {}
    for entry in manifest["files"]:
        path = tmp_path / "out" / entry["output_name"]
        texts[entry["output_name"]] = "\n".join(
            p.text for p in Document(str(path)).paragraphs
        )
    for output_name, token in renamed:
        assert token["old"] not in texts[output_name]
    for path, digest in before.items():
        assert wtm.sha256_file(path) == digest


def test_owner_removed_tokens_leave_no_trace_and_drop_redundant_line():
    """Owner gỡ [Giá chuyển nhượng]/[Số công chứng]/[Người ủy quyền2] khỏi bộ mẫu.

    Hai token đầu phải chừa ô điền tay ngay trên nhãn in sẵn của biểu mẫu hành chính;
    [Người ủy quyền2] nằm ở dòng "2. ..." nên cả dòng là chữ thừa -> xoá hẳn.
    """
    doc = Document()
    doc.add_paragraph("[32.1] Số: [Số công chứng] [32.2] Nơi lập: ………")
    doc.add_paragraph("[41.10] Giá trị đất thực tế chuyển giao (nếu có): [Giá chuyển nhượng] đồng")
    doc.add_paragraph("1. Ông [Tên 3];")
    doc.add_paragraph("2. [Người ủy quyền2]")

    counts, removed_lines = wtm.drop_tokens_in_doc(doc, wtm.DROP_TOKENS)
    doc.save(str(_tmp_docx(doc)))
    reloaded = Document(str(_tmp_docx(doc)))
    text = "\n".join(p.text for p in reloaded.paragraphs)

    assert counts["[Số công chứng]"] == 1
    assert counts["[Giá chuyển nhượng]"] == 1
    assert counts["[Người ủy quyền2]"] == 1
    assert removed_lines == 1
    assert "[Số công chứng]" not in text
    assert "[Giá chuyển nhượng]" not in text
    assert "[Người ủy quyền2]" not in text
    # Nhãn in sẵn + đơn vị "đồng" còn nguyên, chỉ chừa chỗ điền tay.
    assert "[32.1] Số: …………" in text
    assert "(nếu có): ………… đồng" in text
    assert "1. Ông" in text
    assert "2. " not in text


def test_migrate_templates_actually_removes_dropped_tokens_from_output(tmp_path):
    """Token Owner gỡ phải biến mất khỏi file đích — manifest 'dropped' mà file
    vẫn còn token là bug (lookup key chuẩn hoá lệch case đã từng xảy ra)."""
    src = tmp_path / "src"
    _docx(src / "mau.docx", ["Số chứng thực: [Số công chứng]", "Giá: [Giá chuyển nhượng] đồng"])

    manifest = wtm.migrate_templates(src, tmp_path / "out", registry=[], no_source=[])

    text = "\n".join(
        p.text for p in Document(str(tmp_path / "out" / "mau.docx")).paragraphs
    )
    assert "[Số công chứng]" not in text
    assert "[Giá chuyển nhượng]" not in text
    entry = manifest["files"][0]
    assert entry["status_counts"].get("dropped") == 2


def _tmp_docx(doc):
    import tempfile
    from pathlib import Path as _Path

    path = _Path(tempfile.mkdtemp()) / "probe.docx"
    doc.save(str(path))
    return path


# ------------------------------------------------- context fix (NAIA-12 review)


def _para_texts(path: Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs]


def test_migrate_dk18_hardcoded_land_labels_become_dynamic(tmp_path):
    """Nhãn in cứng 'Loại đất N: <mã>' kèm token vị trí -> placeholder động
    [Loại đất tài sản N.M] (in đúng loại đất THẬT của dòng, không còn nhãn sai
    cứng). Token vị trí [ONT N] giữ nguyên; mã hành chính [41.7] giữ nguyên."""
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế  3.docx", [
        "[41.7] Loại đất, vị trí thửa đất (1,2,3,4…)",
        "+ Loại đất 1: ONT; Vị trí:………… Diện tích: [ONT 3] m2",
        "+ Loại đất 4: LUC; Vị trí:………… Diện tích: [LUC 3] m2",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=["[ONT 3]", "[LUC 3]"], no_source=[])

    texts = _para_texts(tmp_path / "out" / "2. DK18, thuế  3.docx")
    assert "+ Loại đất 1: [Loại đất tài sản 3.1]; Vị trí:………… Diện tích: [ONT 3] m2" in texts
    assert "+ Loại đất 4: [Loại đất tài sản 3.4]; Vị trí:………… Diện tích: [LUC 3] m2" in texts
    # Mã in cứng đã xác minh giữ nguyên.
    assert "[41.7] Loại đất, vị trí thửa đất (1,2,3,4…)" in texts
    entry = manifest["files"][0]
    # Nhãn in cứng là sự kiện LITERAL — nằm ở literal_fixes, KHÔNG phải
    # placeholder nguồn nên không lọt vào tokens/occurrence (review NAIA-12).
    literal = {t["new"] for t in entry["literal_fixes"]}
    assert literal == {
        "[Loại đất tài sản 3.1]", "[Loại đất tài sản 3.4]"}
    assert not [t for t in entry["tokens"]
                if t["status"] == "context_fix"
                and t["new"] in literal]
    assert entry["status_counts"].get("fixed_code") == 1


def test_migrate_dk18_n_scopes_unsuffixed_tokens_to_asset_n(tmp_path):
    """Mẫu 'DK18, thuế 2' là tờ khai tài sản 2: [Diện tích]/[Địa chỉ đất]
    thiếu hậu tố phải chỉ tài sản 2 — không phải tài sản 1 (context_fix)."""
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế  2.docx", [
        "1.2. Địa chỉ thửa đất: [Địa chỉ đất]",
        "1.5. Diện tích (m2): [Diện tích 2]",
        "[20.5] Diện tích: [Diện tích] [20.6] Loại đất/Mục đích sử dụng: [Loại đất 2]",
        "[41.2] Địa chỉ: [Địa chỉ đất 2]",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    texts = _para_texts(tmp_path / "out" / "2. DK18, thuế  2.docx")
    joined = "\n".join(texts)
    assert "1.2. Địa chỉ thửa đất: [Địa chỉ đất tài sản 2]" in joined
    assert "1.5. Diện tích (m2): [Diện tích tài sản 2]" in joined
    assert "[20.5] Diện tích: [Diện tích tài sản 2] [20.6] Loại đất/Mục đích sử dụng: [Loại đất tài sản 2]" in joined
    # Token đã đúng hậu tố tài sản giữ nguyên.
    assert "[41.2] Địa chỉ: [Địa chỉ đất 2]" in joined
    statuses = {t["old"]: t["status"] for t in manifest["files"][0]["tokens"]}
    assert statuses["[Địa chỉ đất]"] == "context_fix"
    assert statuses["[Diện tích]"] == "context_fix"
    assert statuses["[Diện tích 2]"] == "context_fix"
    assert statuses["[Loại đất 2]"] == "context_fix"
    assert statuses["[Địa chỉ đất 2]"] == "unmapped"  # ngoài registry -> vẫn theo luồng cũ


def test_migrate_purpose_enumeration_becomes_purpose_token(tmp_path):
    """Liệt kê in cứng 'Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2; ...'
    -> [Mục đích sử dụng 1]: chỉ in các dòng đất THẬT của hồ sơ."""
    src = tmp_path / "src"
    _docx(src / "1. PCDS .docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2; NTS: [NTS] m2",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    texts = _para_texts(tmp_path / "out" / "1. PCDS .docx")
    assert texts == ["Mục đích sử dụng: [Mục đích sử dụng 1]"]
    fixes = {t["old"]: t["new"] for t in manifest["files"][0]["tokens"]
             if t["status"] == "context_fix"}
    assert fixes.get("[ONT]") == "[Mục đích sử dụng 1]"
    assert fixes.get("[CLN]") == "[Mục đích sử dụng 1]"
    assert fixes.get("[NTS]") == "[Mục đích sử dụng 1]"


def test_migrate_context_fix_leaves_other_files_untouched(tmp_path):
    """Mẫu KHÔNG thuộc danh sách DK18/PCDS: nhãn 'ONT' và token giữ nguyên —
    context fix chỉ áp trên file đã đối chiếu."""
    src = tmp_path / "src"
    _docx(src / "mau khac.docx", [
        "+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT] m2",
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2; NTS: [NTS] m2",
    ])

    wtm.migrate_templates(src, tmp_path / "out", registry=[], no_source=[])

    texts = _para_texts(tmp_path / "out" / "mau khac.docx")
    assert texts[0] == "+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT] m2"
    assert texts[1] == "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2; NTS: [NTS] m2"


# ------------------------------------- review NAIA-12 (đợt 81ab0ca) ---------


def test_context_fix_enum_stops_at_intervening_note(tmp_path):
    """Ghi chú chen giữa liệt kê KHÔNG được nuốt vào span thay thế —
    phản ví dụ Leader: enum gộp ăn luôn note/token ở giữa cụm."""
    src = tmp_path / "src"
    _docx(src / "1. PCDS .docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2; NTS: [NTS] m2 "
        "(ghi chú của tác nghiệp) LUC: [LUC] m2",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    texts = _para_texts(tmp_path / "out" / "1. PCDS .docx")
    # Cụm liên tiếp đầu vẫn gộp thành placeholder mục đích; phần sau note
    # chỉ có MỘT item nên giữ nguyên — quan trọng nhất: note còn nguyên.
    assert texts == [
        "Mục đích sử dụng: [Mục đích sử dụng 1] "
        "(ghi chú của tác nghiệp) LUC: [LUC] m2"]


def test_context_fix_enum_intervening_token_splits_run(tmp_path):
    """Token khác chen giữa liệt kê phá run tại đúng chỗ đó — mỗi cụm liên
    tiếp hợp lệ thay riêng, token chen giữa giữ nguyên không bị nuốt."""
    src = tmp_path / "src"
    _docx(src / "1. Khai nhận di sản Vũ Dương.docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2 theo [Nơi niêm yết]; "
        "NTS: [NTS] m2; LUC: [LUC] m2",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    texts = _para_texts(tmp_path / "out" / "1. Khai nhận di sản Vũ Dương.docx")
    assert texts == [
        "Mục đích sử dụng: [Mục đích sử dụng 1] theo [Nơi niêm yết]; "
        "[Mục đích sử dụng 1]"]


def test_context_fix_mixed_asset_enum_kept_with_warning(tmp_path):
    """Liệt kê trộn mã của NHIỀU tài sản (ONT asset1 + CLN asset2) không
    thể gom về một [Mục đích sử dụng N] -> giữ nguyên + ghi cảnh báo."""
    src = tmp_path / "src"
    _docx(src / "1. PCDS .docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN 2] m2; NTS: [NTS] m2",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    texts = _para_texts(tmp_path / "out" / "1. PCDS .docx")
    assert texts == [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN 2] m2; NTS: [NTS] m2"]
    entry = manifest["files"][0]
    assert any("trộn nhiều tài sản" in n and "giữ nguyên" in n
               for n in entry["notes"])


def test_context_fix_labels_pair_same_code_across_assets(tmp_path):
    """Hai nhãn cùng mã ONT nhưng kèm token tài sản khác nhau phải gán
    đúng hậu tố từng tài sản — phản ví dụ Leader: cả hai dính nhãn 1."""
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế  2.docx", [
        "+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT] m2",
        "+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT 2] m2",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    texts = _para_texts(tmp_path / "out" / "2. DK18, thuế  2.docx")
    assert texts == [
        "+ Loại đất 1: [Loại đất tài sản 1.1]; Vị trí:…… Diện tích: [ONT] m2",
        "+ Loại đất 1: [Loại đất tài sản 2.1]; Vị trí:…… Diện tích: [ONT 2] m2",
    ]
    literal = {t["new"] for t in manifest["files"][0]["literal_fixes"]}
    assert literal == {"[Loại đất tài sản 1.1]", "[Loại đất tài sản 2.1]"}


def test_write_guard_blocks_junction_swapped_after_copy(tmp_path, monkeypatch):
    """Thư mục con trong đích bị đổi thành junction SAU copy nhưng TRƯỚC
    context save -> guard phải raise, không ghi ra ngoài (review NAIA-12
    tìm 1: context save từng đuổi theo junction ghi vào sentinel)."""
    src = tmp_path / "src"
    _docx(src / "sub" / "2. DK18, thuế  3.docx", [
        "+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT 3] m2",
    ])
    out = tmp_path / "out"
    sentinel = tmp_path / "sentinel"
    sentinel.mkdir()

    # Sentinel giả file đích nằm ngoài — nếu save đuổi theo junction thì
    # nội dung này bị ghi đè bằng nội dung đã migrate.
    _docx(sentinel / "2. DK18, thuế  3.docx", ["SENTINEL — KHÔNG ĐƯỢC GHI"])
    sentinel_bytes = (
        sentinel / "2. DK18, thuế  3.docx").read_bytes()

    real_read_tokens = wtm._read_tokens
    swapped = {"done": False}

    def _swap_after_copy(path: Path) -> list[str]:
        result = real_read_tokens(path)
        if not swapped["done"]:
            swapped["done"] = True
            real_dir = path.parent
            moved = path.parent.parent / "sub_real"
            real_dir.rename(moved)
            os.system(f'cmd /c mklink /J "{real_dir}" "{sentinel}" >nul')
            assert real_dir.is_symlink() or wtm._real_path(real_dir) != real_dir
        return result

    monkeypatch.setattr(wtm, "_read_tokens", _swap_after_copy)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, out, registry=[], no_source=[])

    # Không byte nào lọt ra ngoài đích; file gốc vẫn nguyên trong sub_real.
    assert (sentinel / "2. DK18, thuế  3.docx").read_bytes() == sentinel_bytes
    assert (out / "sub_real" / "2. DK18, thuế  3.docx").exists()


def test_manifest_totals_exclude_literal_events(tmp_path):
    """occurrence trong manifest chỉ đếm placeholder NGUỒN — literal event
    (nhãn in cứng -> placeholder) tách riêng ở literal_events (review:
    tổng từng phình 2142 thay vì 2128)."""
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế  3.docx", [
        "+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT 3] m2",
        "+ Loại đất 4: LUC; Vị trí:…… Diện tích: [LUC 3] m2",
        "Địa chỉ: [Địa chỉ đất]",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    entry = manifest["files"][0]
    source_occurrences = sum(t["count"] for t in entry["tokens"])
    assert source_occurrences == 3  # [ONT 3], [LUC 3], [Địa chỉ đất]
    assert manifest["totals"]["occurrences"] == 3
    # Hai nhãn in cứng -> hai sự kiện literal, KHÔNG cộng vào occurrences.
    assert len(entry["literal_fixes"]) == 2
    assert manifest["totals"]["literal_events"] == 2
    # Token bị gỡ theo literal fix (không có ở đây) vẫn nằm trong tokens.
    removed = sum(
        sum(f["removed_tokens"].values()) for f in entry["literal_fixes"])
    assert removed == 0


def test_manifest_enum_removed_tokens_counted_as_token_occurrences(tmp_path):
    """Token bị gỡ cùng liệt kê enum vẫn là placeholder nguồn: count trong
    tokens, applied >0, literal event ghi nhận removed_tokens tương ứng."""
    src = tmp_path / "src"
    _docx(src / "1. PCDS .docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2; NTS: [NTS] m2",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    entry = manifest["files"][0]
    statuses = {t["old"]: t for t in entry["tokens"]}
    for old in ("[ONT]", "[CLN]", "[NTS]"):
        assert statuses[old]["status"] == "context_fix"
        assert statuses[old]["new"] == "[Mục đích sử dụng 1]"
        assert statuses[old]["applied"] >= 1
    # Source vẫn có 3 placeholder occurrences; literal_events ghi 1 span.
    assert sum(t["count"] for t in entry["tokens"]) == 3
    assert manifest["totals"]["occurrences"] == 3
    assert manifest["totals"]["literal_events"] == 1
    fix = entry["literal_fixes"][0]
    assert fix["removed_tokens"] == {"[ONT]": 1, "[CLN]": 1, "[NTS]": 1}


# ------------------------------------- review NAIA-12 (đợt e47ad1d) ---------


def _swap_sub_to_junction_inside_root(
        monkeypatch, out: Path, victim_dir: Path) -> dict:
    """Đổi ``out/sub`` thành junction trỏ sang ``victim_dir`` (VẪN TRONG
    root đích) NGAY SAU copy — junction đổi hướng giữa chừng tới file
    khác trong root (review NAIA-12 e47ad1d: realpath mới trong root
    từng qua guard và đè lên sentinel)."""
    state = {"done": False}
    real_read_tokens = wtm._read_tokens

    def _swap(path: Path) -> list[str]:
        tokens = real_read_tokens(path)
        if not state["done"]:
            state["done"] = True
            owned = out / "sub"
            owned.rename(out / "sub-original")
            if not _make_dir_link(owned, victim_dir):
                (out / "sub-original").rename(owned)
                pytest.skip("môi trường không tạo được junction/symlink")
        return tokens

    monkeypatch.setattr(wtm, "_read_tokens", _swap)
    return state


@pytest.mark.parametrize("kind", ["rename", "drop", "context"])
def test_write_guard_blocks_junction_to_other_dir_inside_root(
        tmp_path, monkeypatch, kind):
    """Junction trong đích đổi sang THƯ MỤC KHÁC VẪN TRONG ROOT sau copy:
    mọi write sau copy (tmp->replace rename, drop save, context save)
    phải từ chối — sentinel cùng root không bị ghi đè, nguồn nguyên vẹn.
    """
    paragraphs = {
        "rename": ["Mẫu: [Niêm Yết]"],
        "drop": ["2. [Người ủy quyền2]"],
        "context": ["+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT] m2"],
    }[kind]
    name = "2. DK18, thuế  1.docx" if kind == "context" else "m.docx"
    src = tmp_path / "src"
    _docx(src / "sub" / name, paragraphs)
    out = tmp_path / "out"
    victim_dir = out / "other"
    _docx(victim_dir / name, ["SENTINEL — KHÔNG ĐƯỢC GHI"])
    sentinel = victim_dir / name
    sentinel_bytes = sentinel.read_bytes()
    source_bytes = (src / "sub" / name).read_bytes()
    _swap_sub_to_junction_inside_root(monkeypatch, out, victim_dir)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, out, registry=[], no_source=[])

    assert sentinel.read_bytes() == sentinel_bytes
    assert (src / "sub" / name).read_bytes() == source_bytes


def test_migrate_templates_never_touches_legacy_temp_name(tmp_path):
    """``m.renamed.docx`` lạ có sẵn ở đích: publish mới không còn bước
    ``Path.replace`` qua file tạm — path đó KHÔNG còn là write target nên
    migrate thành công và file ngoài giữ nguyên bytes (không ai đụng tới).
    """
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Mẫu: [Niêm Yết]"])
    out = tmp_path / "out"
    out.mkdir()
    foreign = out / "m.renamed.docx"
    foreign.write_bytes(b"FOREIGN TEMP")

    manifest = wtm.migrate_templates(src, out, registry=[], no_source=[])

    assert manifest["file_count"] == 1
    assert foreign.read_bytes() == b"FOREIGN TEMP"


def test_write_guard_foreign_dest_planted_before_create(
        tmp_path, monkeypatch):
    """File lạ xuất hiện ở ĐÍCH SAU ``check`` nhưng TRƯỚC ``xb`` create ->
    ``xb`` phải raise FileExistsError — đích tạo exclusive, không ghi đè
    file chen giữa (review NAIA-12)."""
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Mẫu: [Niêm Yết]"])
    foreign = b"FOREIGN TEMP - KHONG DUOC GHI DE"
    real_check = wtm._WriteGuard.check

    def _plant(self, path, what="đầu ra"):
        real_check(self, path, what)
        if what.startswith("đầu ra"):
            Path(path).write_bytes(foreign)

    monkeypatch.setattr(wtm._WriteGuard, "check", _plant)

    with pytest.raises(FileExistsError):
        wtm.migrate_templates(src, tmp_path / "out",
                              registry=[], no_source=[])

    dest = tmp_path / "out" / "m.docx"
    assert dest.read_bytes() == foreign


def test_manifest_partial_context_fix_counts_kept_occurrence(tmp_path):
    """``[ONT]`` ×2: một occurrence trong liệt kê bị gỡ (context_fix),
    một đứng riêng giữ nguyên -> occurrence_counts tách context_fix/kept
    đúng số lần — không được gán cả 2 vào context_fix (review NAIA-12)."""
    src = tmp_path / "src"
    _docx(src / "1. PCDS .docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2",
        "Giữ riêng: [ONT]",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    entry = manifest["files"][0]
    counts = entry["occurrence_counts"]
    assert counts["context_fix"] == 2   # [ONT] + [CLN] trong cụm liệt kê
    assert counts["kept"] == 1          # [ONT] đứng riêng giữ nguyên
    assert counts["_occurrences"] == 3
    assert counts["_applied"] == 2
    token = next(t for t in entry["tokens"] if t["old"] == "[ONT]")
    assert token["count"] == 2 and token["applied"] == 1
    texts = _para_texts(tmp_path / "out" / "1. PCDS .docx")
    assert texts == [
        "Mục đích sử dụng: [Mục đích sử dụng 1]", "Giữ riêng: [ONT]"]


def test_manifest_normalized_context_alias_keeps_one_record(tmp_path):
    """Nguồn viết HOA ``[ĐỊA CHỈ ĐẤT]``; context rename khai báo ``[Địa
    chỉ đất]`` — manifest phải cập nhật ĐÚNG record nguồn (giữ tên gốc,
    count=1), KHÔNG thêm record thứ hai cho alias đã chuẩn hoá
    (review NAIA-12: 1 occurrence nguồn từng thành 2 record manifest)."""
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế  3.docx", ["1.2. Địa chỉ: [ĐỊA CHỈ ĐẤT]"])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    entry = manifest["files"][0]
    assert len(entry["tokens"]) == 1
    token = entry["tokens"][0]
    assert token["old"] == "[ĐỊA CHỈ ĐẤT]"      # lineage tên NGUỒN giữ nguyên
    assert token["new"] == "[Địa chỉ đất tài sản 3]"
    assert token["status"] == "context_fix"
    assert token["count"] == 1 and token["applied"] == 1
    assert manifest["totals"]["occurrences"] == 1
    assert "[Địa chỉ đất tài sản 3]" in "\n".join(
        _para_texts(tmp_path / "out" / "2. DK18, thuế  3.docx"))


def test_context_fix_same_paragraph_labels_pair_per_asset(tmp_path):
    """Hai mục cùng mã ONT TRONG CÙNG MỘT ĐOẠN nhưng kèm token tài sản
    khác nhau phải nhận nhãn đúng từng tài sản — pairing theo vị trí
    trong đoạn, không theo mã chung (review NAIA-12)."""
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế  1.docx", [
        "Loại đất 1: ONT; Diện tích: [ONT 1] m2; "
        "Loại đất 2: ONT; Diện tích: [ONT 2] m2",
    ])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    texts = _para_texts(tmp_path / "out" / "2. DK18, thuế  1.docx")
    assert texts == [
        "Loại đất 1: [Loại đất tài sản 1.1]; Diện tích: [ONT 1] m2; "
        "Loại đất 2: [Loại đất tài sản 2.1]; Diện tích: [ONT 2] m2"]
    literal = {t["new"] for t in manifest["files"][0]["literal_fixes"]}
    assert literal == {"[Loại đất tài sản 1.1]", "[Loại đất tài sản 2.1]"}


# ----------------------------- review NAIA-12 (repair-2dddc41) -------------


def test_manifest_mixed_normalized_spellings_attribute_per_occurrence(
        tmp_path):
    """``[ĐỊA CHỈ ĐẤT]`` + ``[Địa chỉ đất]`` cùng normalized key nhưng là
    HAI occurrence nguồn riêng: manifest phải có 2 record count=1, mỗi
    record context_fix applied=1 — không được gộp applied=2 vào record
    đầu trong khi record sau báo kept dù occurrence của nó cũng bị thay
    (review NAIA-12 repair-2dddc41: tổng count bảo toàn vẫn che lỗi)."""
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế  3.docx", ["[ĐỊA CHỈ ĐẤT]", "[Địa chỉ đất]"])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    entry = manifest["files"][0]
    assert len(entry["tokens"]) == 2
    by_old = {t["old"]: t for t in entry["tokens"]}
    for old in ("[ĐỊA CHỈ ĐẤT]", "[Địa chỉ đất]"):
        token = by_old[old]
        assert token["count"] == 1
        assert token["status"] == "context_fix"
        assert token["new"] == "[Địa chỉ đất tài sản 3]"
        assert token["applied"] == 1
    counts = entry["occurrence_counts"]
    assert counts["context_fix"] == 2
    assert counts.get("kept", 0) == 0
    assert counts["_occurrences"] == 2
    assert counts["_applied"] == 2
    texts = _para_texts(tmp_path / "out" / "2. DK18, thuế  3.docx")
    assert texts == ["[Địa chỉ đất tài sản 3]", "[Địa chỉ đất tài sản 3]"]


def test_write_guard_registration_rejects_swapped_file(
        tmp_path, monkeypatch):
    """File đích bị thay giữa đóng handle ``xb`` và ``register`` -> register
    đối chiếu danh tính đã ``fstat`` TỪ FD TẠO (không ``stat()`` path sau
    đóng handle) -> ValueError; file cạnh tranh giữ nguyên bytes, nguồn
    nguyên vẹn (review NAIA-12 repair-2dddc41: register sau đóng handle
    từng nhận nhầm file ngoài làm của run rồi ghi đè nó)."""
    src = tmp_path / "src"
    name = "2. DK18, thuế  1.docx"
    _docx(src / "sub" / name,
          ["+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT] m2"])
    out = tmp_path / "out"
    foreign = out / "other" / name
    _docx(foreign, ["FOREIGN_SENTINEL — KHÔNG ĐƯỢC GHI"])
    foreign_bytes = foreign.read_bytes()
    owned = out / "sub" / name
    backup = out / "owned-copy.docx"
    source_bytes = (src / "sub" / name).read_bytes()
    real_register = wtm._WriteGuard.register
    swapped = {"done": False}

    def _swap_before_register(self, path):
        if Path(path) == owned and not swapped["done"]:
            swapped["done"] = True
            owned.rename(backup)
            foreign.rename(owned)
        return real_register(self, path)

    monkeypatch.setattr(wtm._WriteGuard, "register", _swap_before_register)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, out, registry=[], no_source=[])

    # File cạnh tranh KHÔNG bị run nhận nhầm là của mình rồi ghi đè.
    assert owned.read_bytes() == foreign_bytes
    assert backup.exists()
    assert (src / "sub" / name).read_bytes() == source_bytes


def test_write_guard_blocks_same_path_swap_before_rewrite(
        tmp_path, monkeypatch):
    """File đích bị XOÁ + TẠO LẠI tại CÙNG logical/real path ngay trước
    ``open_bound`` -> inode mới khác danh tính đã fstat từ handle
    tạo -> ValueError trước đọc/ghi; sentinel cùng path giữ nguyên
    bytes, nguồn nguyên vẹn (repair-2dddc41: phải test thay inode tại
    cùng path, không chỉ junction đổi realpath)."""
    src = tmp_path / "src"
    name = "2. DK18, thuế  1.docx"
    _docx(src / "sub" / name,
          ["+ Loại đất 1: ONT; Vị trí:…… Diện tích: [ONT] m2"])
    out = tmp_path / "out"
    owned = out / "sub" / name
    sentinel = b"FOREIGN SAME-PATH SENTINEL - KHONG DUOC GHI DE"
    source_bytes = (src / "sub" / name).read_bytes()
    real_open = wtm._WriteGuard.open_bound
    swapped = {"done": False}

    def _swap_at_same_path(self, path, what="đầu ra"):
        if Path(path) == owned and not swapped["done"]:
            swapped["done"] = True
            # Thay inode TẠI CÙNG PATH — realpath không đổi nên chỉ fd
            # ràng buộc inode (fstat) mới phân biệt được file của run
            # với file cạnh tranh.
            Path(path).unlink()
            Path(path).write_bytes(sentinel)
        return real_open(self, path, what)

    monkeypatch.setattr(
        wtm._WriteGuard, "open_bound", _swap_at_same_path)

    with pytest.raises(ValueError):
        wtm.migrate_templates(src, out, registry=[], no_source=[])

    assert owned.read_bytes() == sentinel
    assert (src / "sub" / name).read_bytes() == source_bytes


# -------------------------------------- repair-f35c9f6 --------------------


def test_migrate_templates_rejects_swapped_converter_output(
        tmp_path, monkeypatch):
    """Converter staging bị thay bằng file cạnh tranh ngay trước khi run
    đọc -> ``fstat`` trên fd đọc lệch danh tính staging run tạo ->
    ValueError; file cạnh tranh giữ nguyên bytes tại path staging,
    KHÔNG bị copy sang đích hay xoá (repair-f35c9f6)."""
    import builtins
    src = tmp_path / "src"
    src.mkdir(parents=True)
    source = src / "legacy.doc"
    source_bytes = b"SYNTHETIC_LEGACY_SOURCE"
    source.write_bytes(source_bytes)
    foreign = tmp_path / "foreign.docx"
    _docx(foreign, ["FOREIGN CONVERSION SENTINEL"])
    foreign_bytes = foreign.read_bytes()
    real_open = builtins.open
    state = {"tmp": None, "swapped": False}

    def fake_convert(src_path, dst_path):
        # Converter ghi in-place vào đúng inode staging run đã tạo.
        state["tmp"] = Path(dst_path)
        _docx(state["tmp"], ["OWNED CONVERSION OUTPUT"])

    def swap_before_read(path, *args, **kwargs):
        mode = args[0] if args else kwargs.get("mode", "r")
        if (state["tmp"] is not None and Path(path) == state["tmp"]
                and mode == "rb" and not state["swapped"]):
            state["swapped"] = True
            state["tmp"].rename(tmp_path / "owned-backup.docx")
            foreign.rename(state["tmp"])
        return real_open(path, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", swap_before_read)
    out = tmp_path / "out"
    with pytest.raises(ValueError):
        wtm.migrate_templates(src, out, registry=[], no_source=[],
                              convert=fake_convert)

    assert state["swapped"]
    # File cạnh tranh nguyên bytes tại path staging — không copy/xoá.
    assert state["tmp"].read_bytes() == foreign_bytes
    assert not (out / "legacy.docx").exists()
    assert source.read_bytes() == source_bytes


def test_migrate_templates_accepts_receipt_converter(tmp_path):
    """Converter tạo inode MỚI (ví dụ Word COM SaveAs2) thay vì ghi
    in-place staging -> chứng minh quyền sở hữu bằng receipt dict
    ``{"sha256", "identity"}`` mà producer đọc trên handle file mới tạo —
    hash + danh tính producer báo đúng file đang đọc (repair-83e6d03:
    receipt chỉ-hash không còn đủ chứng minh khi inode khác staging).
    Danh tính dùng ``_file_identity`` — cùng biểu diễn NTFS với producer
    PowerShell (dwVolumeSerialNumber + file index), không phải
    ``st_dev`` của CRT (repair-1ae81e3)."""
    import hashlib
    src = tmp_path / "src"
    src.mkdir(parents=True)
    (src / "legacy.doc").write_bytes(b"SYNTHETIC_LEGACY_SOURCE")
    staged = {"tmp": None}

    def fake_convert(src_path, dst_path):
        staging_tmp = Path(dst_path).parent / (Path(dst_path).name + ".work")
        _docx(staging_tmp, ["[Niêm Yết]"])
        staging_tmp.replace(dst_path)            # inode mới tại staging
        staged["tmp"] = Path(dst_path)
        # Producer đọc hash + danh tính trên MỘT handle file nó vừa ghi.
        with open(staged["tmp"], "rb") as handle:
            identity = wtm._file_identity(handle)
            digest = hashlib.sha256(handle.read()).hexdigest()
        return {"sha256": digest, "identity": identity}

    out = tmp_path / "out"
    manifest = wtm.migrate_templates(
        src, out, registry=["[Niêm Yết]"], no_source=[],
        convert=fake_convert)

    item = manifest["files"][0]
    assert item.get("converted_from_doc") is True
    output = out / "legacy.docx"
    assert output.exists()
    doc = Document(str(output))
    assert any("[Nơi niêm yết]" in p.text for p in doc.paragraphs)
    # Staging giữ lại trong `.naia12-scratch-*` (không unlink theo path)
    # và nằm ngoài vùng scan.
    assert staged["tmp"].exists()
    assert staged["tmp"].parent.name.startswith(".naia12-scratch-")


def test_migrate_templates_rejects_wrong_receipt(tmp_path):
    """Converter trả receipt sha256 KHÔNG khớp payload staging -> từ chối,
    giữ nguyên staging, không tạo đích (không tin output không rõ chủ)."""
    src = tmp_path / "src"
    src.mkdir(parents=True)
    source = src / "legacy.doc"
    source_bytes = b"SYNTHETIC_LEGACY_SOURCE"
    source.write_bytes(source_bytes)
    staged = {"tmp": None}

    def fake_convert(src_path, dst_path):
        staged["tmp"] = Path(dst_path)
        _docx(staged["tmp"], ["[Niêm Yết]"])
        return "0" * 64                          # receipt sai

    out = tmp_path / "out"
    with pytest.raises(ValueError):
        wtm.migrate_templates(src, out, registry=[], no_source=[],
                              convert=fake_convert)

    assert staged["tmp"].exists()
    assert not (out / "legacy.docx").exists()
    assert source.read_bytes() == source_bytes


# -------------------------------------- repair-83e6d03 --------------------


def test_convert_doc_to_docx_refuses_zero_byte_foreign_dest(
        tmp_path, monkeypatch):
    """Đích RỖNG đã tồn tại mà không phải staging run tạo (gọi trực tiếp,
    không ``staging_identity``) -> FileExistsError TRƯỚC khi dispatch
    Word (repair-83e6d03: đích 0-byte lạ từng được chấp nhận rồi gọi
    converter ghi đè)."""
    src = tmp_path / "cu.doc"
    src.write_bytes(b"x")
    dst = tmp_path / "foreign.docx"
    dst.write_bytes(b"")                       # file ngoài, rỗng
    called = []

    def boom(*args, **kwargs):
        called.append(args)
        raise AssertionError("không được dispatch Word")

    monkeypatch.setattr("subprocess.run", boom)
    with pytest.raises(FileExistsError):
        wtm.convert_doc_to_docx(src, dst)
    assert called == []
    assert dst.read_bytes() == b""             # file ngoài nguyên vẹn


def test_convert_doc_to_docx_refuses_wrong_staging_identity(
        tmp_path, monkeypatch):
    """Caller khai staging_identity nhưng path đích đang là file KHÁC ->
    ValueError trước dispatch — staging bị thay sau khi capture không
    được đưa tới Word."""
    src = tmp_path / "cu.doc"
    src.write_bytes(b"x")
    foreign = tmp_path / "foreign.docx"
    foreign.write_bytes(b"FOREIGN")
    called = []

    def boom(*args, **kwargs):
        called.append(args)
        raise AssertionError("không được dispatch Word")

    monkeypatch.setattr("subprocess.run", boom)
    with pytest.raises(ValueError):
        wtm.convert_doc_to_docx(
            src, foreign, staging_identity=(999999, 999999))
    assert called == []
    assert foreign.read_bytes() == b"FOREIGN"


def test_convert_doc_to_docx_parses_bound_receipt(tmp_path, monkeypatch):
    """Script thật xuất ``NAIA12_CONV_SHA256`` + ``NAIA12_CONV_ID`` —
    helper trả dict receipt cả hash lẫn danh tính producer."""
    src = tmp_path / "cu.doc"
    src.write_bytes(b"x")
    dst = tmp_path / "cu.docx"
    seen = {}

    def fake_run(args, **kwargs):
        dst.write_bytes(b"CONVERTED")
        seen["env"] = kwargs["env"]
        return type("P", (), {"returncode": 0, "stderr": "",
                              "stdout": "NAIA12_CONV_SHA256=ab12\n"
                                        "NAIA12_CONV_ID=777:888\n"})()

    monkeypatch.setattr("subprocess.run", fake_run)
    receipt = wtm.convert_doc_to_docx(src, dst)
    assert receipt == {"sha256": "ab12", "identity": (777, 888)}
    # Danh tính staging run tạo được đưa qua env cho kiểm tra producer.
    assert "NAIA_DST_DEV" in seen["env"] and "NAIA_DST_INO" in seen["env"]


def test_migrate_templates_str_receipt_inplace_accepted(tmp_path):
    """Receipt chỉ-hash vẫn được chấp khi converter ghi IN-PLACE lên đúng
    inode staging run tạo — hash đúng + danh tính đúng = output owned
    (tương thích producer trả sha256 dạng chuỗi)."""
    import hashlib
    src = tmp_path / "src"
    src.mkdir(parents=True)
    (src / "legacy.doc").write_bytes(b"SYNTHETIC_LEGACY_SOURCE")

    def fake_convert(src_path, dst_path):
        _docx(Path(dst_path), ["[Niêm Yết]"])     # ghi in-place cùng inode
        with open(dst_path, "rb") as handle:
            return hashlib.sha256(handle.read()).hexdigest()

    out = tmp_path / "out"
    manifest = wtm.migrate_templates(
        src, out, registry=["[Niêm Yết]"], no_source=[],
        convert=fake_convert)
    assert (out / "legacy.docx").exists()
    assert manifest["files"][0]["converted_from_doc"] is True


def test_migrate_templates_rejects_str_receipt_new_inode(tmp_path):
    """Receipt chỉ-hash KHÔNG đủ chứng minh khi file ở staging là inode
    khác (SaveAs2/đổi file): hash đúng của file đang ngồi path vẫn bị
    từ chối vì không ràng buộc được producer với file (repair-83e6d03:
    hash theo path sau khi file bị thay từng được tin nhầm)."""
    import hashlib
    src = tmp_path / "src"
    src.mkdir(parents=True)
    source = src / "legacy.doc"
    source_bytes = b"SYNTHETIC_LEGACY_SOURCE"
    source.write_bytes(source_bytes)
    staged = {"tmp": None}

    def fake_convert(src_path, dst_path):
        staged["tmp"] = Path(dst_path)
        other = staged["tmp"].with_name("other.docx")
        _docx(other, ["OWNED VIA NEW INODE"])
        other.replace(staged["tmp"])             # inode mới tại path
        with open(staged["tmp"], "rb") as handle:
            return hashlib.sha256(handle.read()).hexdigest()

    out = tmp_path / "out"
    with pytest.raises(ValueError):
        wtm.migrate_templates(
            src, out, registry=[], no_source=[], convert=fake_convert)
    assert staged["tmp"].exists()                # giữ để audit, không xoá
    assert not (out / "legacy.docx").exists()
    assert source.read_bytes() == source_bytes


def test_migrate_templates_rejects_dict_receipt_wrong_identity(tmp_path):
    """Receipt dict có sha256 đúng payload nhưng ``identity`` trỏ file
    khác -> từ chối: hash đúng một mình không chứng minh file đang đọc
    là output producer ghi."""
    import hashlib
    src = tmp_path / "src"
    src.mkdir(parents=True)
    (src / "legacy.doc").write_bytes(b"SYNTHETIC_LEGACY_SOURCE")
    staged = {"tmp": None}

    def fake_convert(src_path, dst_path):
        staged["tmp"] = Path(dst_path)
        _docx(staged["tmp"], ["[Niêm Yết]"])
        return {"sha256": hashlib.sha256(
                    staged["tmp"].read_bytes()).hexdigest(),
                "identity": (1, 2)}              # danh tính bịa

    out = tmp_path / "out"
    with pytest.raises(ValueError):
        wtm.migrate_templates(
            src, out, registry=[], no_source=[], convert=fake_convert)
    assert staged["tmp"].exists()
    assert not (out / "legacy.docx").exists()


def test_migrate_templates_never_unlinks_staging_by_path(
        tmp_path, monkeypatch):
    """Probe cleanup-swap: hook ``Path.unlink`` hoán staging sang file
    cạnh tranh ngay khi được gọi -> migrate KHÔNG bao giờ unlink staging
    nên swap không xảy ra, sentinel ngoài nguyên bytes, output vẫn
    publish đúng (repair-83e6d03: stat->unlink từng xoá sentinel)."""
    src = tmp_path / "src"
    src.mkdir(parents=True)
    (src / "legacy.doc").write_bytes(b"SYNTHETIC_LEGACY_SOURCE")
    foreign = tmp_path / "foreign.docx"
    _docx(foreign, ["FOREIGN SENTINEL"])
    foreign_bytes = foreign.read_bytes()
    state = {"tmp": None, "unlink_calls": []}
    real_unlink = Path.unlink

    def fake_convert(src_path, dst_path):
        state["tmp"] = Path(dst_path)
        _docx(state["tmp"], ["OWNED OUTPUT"])

    def swap_on_unlink(self, *args, **kwargs):
        if state["tmp"] is not None and self == state["tmp"]:
            state["unlink_calls"].append(self)
            state["tmp"].rename(tmp_path / "owned-backup.docx")
            foreign.rename(state["tmp"])
        return real_unlink(self, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", swap_on_unlink)
    out = tmp_path / "out"
    manifest = wtm.migrate_templates(
        src, out, registry=[], no_source=[], convert=fake_convert)

    # Không một lần unlink nào vào staging -> sentinel ngoài nguyên.
    assert state["unlink_calls"] == []
    assert foreign.read_bytes() == foreign_bytes
    # Output publish đúng từ payload owned (inode staging không đổi).
    assert (out / "legacy.docx").exists()
    assert manifest["files"][0]["converted_from_doc"] is True


def test_migrate_templates_rejects_receipt_path_replacement(tmp_path):
    """Probe receipt-path: converter ghi output owned, sau đó path bị
    thay bằng file cạnh tranh và producer trả hash TÍNH THEO PATH của
    file ngoài -> hash đúng payload ngoài nhưng danh tính lệch staging
    run tạo -> từ chối; file ngoài giữ nguyên bytes (repair-83e6d03:
    receipt hash theo path từng chứng nhận nhầm file bị thay)."""
    import hashlib
    src = tmp_path / "src"
    src.mkdir(parents=True)
    source = src / "legacy.doc"
    source_bytes = b"SYNTHETIC_LEGACY_SOURCE"
    source.write_bytes(source_bytes)
    foreign = tmp_path / "foreign.docx"
    _docx(foreign, ["FOREIGN REPLACEMENT"])
    foreign_bytes = foreign.read_bytes()
    staged = {"tmp": None, "backup": None}

    def fake_convert(src_path, dst_path):
        staged["tmp"] = Path(dst_path)
        _docx(staged["tmp"], ["OWNED CONVERTED"])
        # Path bị thay sau khi converter ghi — như race sau Close() của
        # Word: owned giữ backup, file ngoài ngồi vào path staging.
        staged["backup"] = tmp_path / "owned-backup.docx"
        staged["tmp"].rename(staged["backup"])
        foreign.rename(staged["tmp"])
        # Producer "hash theo path" trả đúng sha của file NGOÀI.
        return hashlib.sha256(staged["tmp"].read_bytes()).hexdigest()

    out = tmp_path / "out"
    with pytest.raises(ValueError):
        wtm.migrate_templates(
            src, out, registry=[], no_source=[], convert=fake_convert)

    # File ngoài nguyên bytes tại path staging — không copy/xoá.
    assert staged["tmp"].read_bytes() == foreign_bytes
    assert staged["backup"].exists()
    assert not (out / "legacy.docx").exists()


# -------------------------------------- repair-1ae81e3 --------------------

def _powershell() -> str:
    import shutil
    return shutil.which("powershell.exe") or shutil.which("powershell")


def test_powershell_interop_compiles_and_reports_native_identity(
        tmp_path):
    """AC1: đoạn interop C# của script THẬT phải biên dịch được trên
    PowerShell thật — ``FILETIME`` từng là tham chiếu mơ hồ giữa
    ``System.Runtime.InteropServices`` và ``System.Runtime.InteropServices
    .ComTypes`` -> SOURCE_CODE_ERROR trước khi tới COM (review 1ae81e3).

    ``Get-Naia12ConvId`` trên file thật phải trả đúng danh tính NTFS
    ``dwVolumeSerialNumber:fileIndex`` mà ``_file_identity`` phía Python
    đọc được trên CÙNG file — cùng một biểu diễn hai phía (st_dev 64-bit
    của CRT KHÔNG khớp volume serial 32-bit), và phân biệt được hai file
    khác nhau. Không dùng Word/COM."""
    if os.name != "nt" or _powershell() is None:
        pytest.skip("probe PowerShell chỉ chạy trên Windows")
    import subprocess

    probe = tmp_path / "probe.bin"
    probe.write_bytes(b"NAIA12_IDENTITY_PROBE")
    other = tmp_path / "other.bin"
    other.write_bytes(b"NAIA12_OTHER_FILE")
    with open(probe, "rb") as handle:
        expected = _file_identity_for_test(handle)
    with open(other, "rb") as handle:
        other_identity = _file_identity_for_test(handle)
    assert expected != other_identity

    # Phần script thật trước ``New-Object Word`` — gồm Add-Type C#,
    # Get-Naia12ConvId và check staging — compile chạy THẬT.
    prefix = wtm._POWERSHELL_CONVERT.split(
        "$word = New-Object -ComObject Word.Application", 1)[0]
    script = prefix + """
$fs = [System.IO.File]::Open(
    $dst, [System.IO.FileMode]::Open, [System.IO.FileAccess]::Read,
    [System.IO.FileShare]::ReadWrite)
try { Write-Output ("NAIA12_CONV_ID=" + (Get-Naia12ConvId $fs)) }
finally { $fs.Close() }
"""
    env = dict(os.environ, NAIA_SRC=str(tmp_path / "x.doc"),
               NAIA_DST=str(probe))
    env.pop("NAIA_DST_DEV", None)
    env.pop("NAIA_DST_INO", None)
    proc = subprocess.run(
        [_powershell(), "-NoProfile", "-NonInteractive",
         "-ExecutionPolicy", "Bypass", "-Command", script],
        capture_output=True, text=True, env=env, timeout=60)
    assert "SOURCE_CODE_ERROR" not in (proc.stderr + proc.stdout), (
        proc.stderr or proc.stdout)
    assert proc.returncode == 0, proc.stderr
    line = next(
        item for item in proc.stdout.splitlines()
        if item.startswith("NAIA12_CONV_ID="))
    vol_s, ino_s = line.split("=", 1)[1].split(":")
    assert (int(vol_s), int(ino_s)) == expected

    # File KHÁC -> PowerShell báo danh tính khác (guard vẫn phân biệt
    # được file/volume khác — không hard-code volume number).
    env["NAIA_DST"] = str(other)
    proc = subprocess.run(
        [_powershell(), "-NoProfile", "-NonInteractive",
         "-ExecutionPolicy", "Bypass", "-Command", script],
        capture_output=True, text=True, env=env, timeout=60)
    assert proc.returncode == 0, proc.stderr
    line = next(
        item for item in proc.stdout.splitlines()
        if item.startswith("NAIA12_CONV_ID="))
    vol_s, ino_s = line.split("=", 1)[1].split(":")
    assert (int(vol_s), int(ino_s)) == other_identity != expected


def _file_identity_for_test(handle):
    """Alias đọc danh tính — giữ test đọc rõ nghĩa phía consumer."""
    return wtm._file_identity(handle)


# Word GIẢ qua override ``New-Object`` — prelude này + script sản phẩm
# nguyên vẹn (chỉ thay dòng ReleaseComObject vì object giả không phải
# COM). SaveAs2 ghi file OWNED; ``NAIA_FAKE_SWAP`` chọn THỜI ĐIỂM một
# tiến trình NGOÀI cố swap staging sang FOREIGN: ``prewrite`` = sau guard
# check, TRƯỚC SaveAs2 (đúng race review be750f7 mô tả), ``prereceipt``
# = sau khi SaveAs2 ghi xong nhưng trước receipt, ``1``/``close`` = trong
# Close(0) sau receipt. Swap được mô phỏng như hành vi của TÁC NHÂN NGOÀI
# chạy đồng thời: Move nằm trong try/catch — khi pin giữ inode tại path,
# Move của tác nhân bị sharing-violation TRÊN PHÍA NÓ, còn lời gọi COM
# của ta (SaveAs2/Close) vẫn trả về bình thường. Không phải race COM thật.
_FAKE_WORD_PRELUDE = r"""
function Invoke-NaiaSwap {
  try {
    [System.IO.File]::Move(
        $env:NAIA_DST, $env:NAIA_DST + '.owned-backup')
    [System.IO.File]::Move($env:NAIA_FAKE_FOREIGN, $env:NAIA_DST)
  } catch {}
}
function New-Object {
  param([Parameter(Position=0)][string]$TypeName, [string]$ComObject)
  if ($ComObject -eq 'Word.Application') {
    if ($env:NAIA_FAKE_SWAP -eq 'prewrite') { Invoke-NaiaSwap }
    $doc = [pscustomobject]@{}
    if ($env:NAIA_FAKE_FLAT) {
      $doc | Add-Member NoteProperty WordOpenXML (
          [System.IO.File]::ReadAllText($env:NAIA_FAKE_FLAT))
    }
    $doc | Add-Member ScriptMethod SaveAs2 {
      param($path, $format)
      if ($env:NAIA_FAKE_FLAT) { throw 'SaveAs2 rename refused by pin' }
      [System.IO.File]::WriteAllBytes(
          $path, [System.IO.File]::ReadAllBytes($env:NAIA_FAKE_OWNED))
      if ($env:NAIA_FAKE_SWAP -eq 'prereceipt') { Invoke-NaiaSwap }
    }
    $doc | Add-Member ScriptMethod Close {
      param($save)
      if ($env:NAIA_FAKE_SWAP -eq '1' -or $env:NAIA_FAKE_SWAP -eq 'close') {
        Invoke-NaiaSwap
      }
    }
    $global:naia_fake_doc = $doc
    $docs = [pscustomobject]@{}
    $docs | Add-Member ScriptMethod Open {
      param($p, $c, $r, $v) return $global:naia_fake_doc }
    $word = [pscustomobject]@{
      Documents = $docs; Visible = $false; DisplayAlerts = 0;
      AutomationSecurity = 0 }
    $word | Add-Member ScriptMethod Quit {}
    return $word
  }
  return Microsoft.PowerShell.Utility\New-Object -TypeName $TypeName
}
"""


def _fake_word_script() -> str:
    """Script sản phẩm + fake Word — NGUYÊN mã sản phẩm: lifecycle
    SaveAs2 -> receipt -> Close và interop danh tính chạy thật.
    ``ReleaseComObject`` nằm trong try/catch nên lỗi trên object giả
    bị nuốt; ``exit 0`` cuối script giữ rc=0 cho path thành công."""
    return _FAKE_WORD_PRELUDE + wtm._POWERSHELL_CONVERT


def test_flatopc_fallback_preserves_xml_binary_and_ownership(tmp_path, monkeypatch):
    import base64
    import zipfile
    from lxml import etree
    if os.name != "nt" or _powershell() is None:
        pytest.skip("native handle fallback requires Windows PowerShell")
    source = tmp_path / "source.doc"
    source.write_bytes(b"SYNTHETIC_LEGACY_SOURCE")
    owned = _docx(tmp_path / "owned.docx", ["FLAT OPC CONTENT"])
    ns = "http://schemas.microsoft.com/office/2006/xmlPackage"
    package = etree.Element(f"{{{ns}}}package", nsmap={"pkg": ns})
    with zipfile.ZipFile(owned) as archive:
        types = etree.fromstring(archive.read("[Content_Types].xml"))
        defaults = {e.get("Extension"): e.get("ContentType")
                    for e in types if e.tag.endswith("Default")}
        overrides = {e.get("PartName"): e.get("ContentType")
                     for e in types if e.tag.endswith("Override")}
        for name in archive.namelist():
            if name == "[Content_Types].xml":
                continue
            part = etree.SubElement(package, f"{{{ns}}}part")
            part.set(f"{{{ns}}}name", "/" + name)
            part.set(f"{{{ns}}}contentType", overrides.get(
                "/" + name, defaults.get(name.rsplit(".", 1)[-1])))
            data = archive.read(name)
            if name.endswith((".xml", ".rels")):
                etree.SubElement(part, f"{{{ns}}}xmlData").append(
                    etree.fromstring(data))
            else:
                etree.SubElement(part, f"{{{ns}}}binaryData").text = (
                    base64.b64encode(data).decode("ascii"))
    binary = b"\x00\x01\xffSYNTHETIC_BINARY"
    part = etree.SubElement(package, f"{{{ns}}}part")
    part.set(f"{{{ns}}}name", "/word/media/test.bin")
    part.set(f"{{{ns}}}contentType", "application/octet-stream")
    etree.SubElement(part, f"{{{ns}}}binaryData").text = (
        base64.b64encode(binary).decode("ascii"))
    flat = tmp_path / "flat.xml"
    flat.write_bytes(etree.tostring(package))
    foreign = tmp_path / "foreign.docx"
    foreign.write_bytes(b"FOREIGN MUST STAY")
    monkeypatch.setenv("NAIA_FAKE_FLAT", str(flat))
    monkeypatch.setenv("NAIA_FAKE_FOREIGN", str(foreign))
    monkeypatch.setenv("NAIA_FAKE_SWAP", "close")
    monkeypatch.setattr(wtm, "_POWERSHELL_CONVERT", _fake_word_script())
    output = tmp_path / "converted.docx"
    receipt = wtm.convert_doc_to_docx(source, output)
    assert receipt["sha256"] == wtm.sha256_file(output)
    assert foreign.read_bytes() == b"FOREIGN MUST STAY"
    assert source.read_bytes() == b"SYNTHETIC_LEGACY_SOURCE"
    assert not output.with_name(output.name + ".owned-backup").exists()
    assert Document(output).paragraphs[0].text == "FLAT OPC CONTENT"
    with zipfile.ZipFile(output) as archive:
        assert archive.read("word/media/test.bin") == binary


def test_fake_word_lifecycle_owned_output_accepted(
        tmp_path, monkeypatch):
    """AC1 lifecycle (synthetic, KHÔNG COM thật): Word giả SaveAs2 ghi
    owned vào staging rồi Close(0) không swap -> producer tính receipt
    trên handle mở TRƯỚC Close -> hash + danh tính NTFS đúng file ->
    migrate publish output owned; check NAIA_DST_DEV/INO cùng biểu diễn
    (dwVolumeSerialNumber + file index) nên không còn mismatch volume
    (repair-1ae81e3)."""
    if os.name != "nt" or _powershell() is None:
        pytest.skip("fake-Word lifecycle probe chạy PowerShell — Windows")
    src = tmp_path / "src"
    src.mkdir(parents=True)
    source = src / "legacy.doc"
    source.write_bytes(b"SYNTHETIC_LEGACY_SOURCE")
    owned = _docx(tmp_path / "owned.docx", ["OWNED CONVERSION OUTPUT"])
    monkeypatch.setenv("NAIA_FAKE_OWNED", str(owned))
    monkeypatch.delenv("NAIA_FAKE_SWAP", raising=False)
    monkeypatch.delenv("NAIA_FAKE_FOREIGN", raising=False)
    monkeypatch.setattr(wtm, "_POWERSHELL_CONVERT", _fake_word_script())

    out = tmp_path / "out"
    manifest = wtm.migrate_templates(
        src, out, registry=[], no_source=[])

    item = manifest["files"][0]
    assert item["converted_from_doc"] is True
    published = out / "legacy.docx"
    assert published.exists()
    assert any("OWNED CONVERSION OUTPUT" in p.text
               for p in Document(str(published)).paragraphs)
    assert source.read_bytes() == b"SYNTHETIC_LEGACY_SOURCE"


def test_fake_word_lifecycle_close_swap_blocked_by_pin(
        tmp_path, monkeypatch):
    """AC1 lifecycle: fake Close(0) cố swap staging sang FOREIGN sau khi
    receipt đã phát — nhưng guard handle giữ mở không share-delete đã
    khoá inode staging tại path, nên ``File.Move`` trong Close BỊ CHẶN:
    foreign giữ nguyên bytes TẠI PATH GỐC CỦA NÓ (không tới được staging),
    staging giữ owned bytes, receipt đứng -> publish output owned hợp lệ
    (repair-be750f7: pin từ guard tới receipt; trước đây receipt mở lại
    path SAU Close từng ký foreign bị swap vào)."""
    if os.name != "nt" or _powershell() is None:
        pytest.skip("fake-Word lifecycle probe chạy PowerShell — Windows")
    src = tmp_path / "src"
    src.mkdir(parents=True)
    source = src / "legacy.doc"
    source.write_bytes(b"SYNTHETIC_LEGACY_SOURCE")
    owned = _docx(tmp_path / "owned.docx", ["OWNED CONVERSION OUTPUT"])
    owned_bytes = owned.read_bytes()
    foreign = _docx(tmp_path / "foreign.docx", ["FOREIGN SENTINEL"])
    foreign_bytes = foreign.read_bytes()
    monkeypatch.setenv("NAIA_FAKE_OWNED", str(owned))
    monkeypatch.setenv("NAIA_FAKE_FOREIGN", str(foreign))
    monkeypatch.setenv("NAIA_FAKE_SWAP", "close")
    monkeypatch.setattr(wtm, "_POWERSHELL_CONVERT", _fake_word_script())

    out = tmp_path / "out"
    manifest = wtm.migrate_templates(src, out, registry=[], no_source=[])

    _assert_swap_blocked_contract(
        tmp_path, out, foreign, foreign_bytes, owned_bytes, source)
    # Receipt đứng sau Close -> publish conversion chứng nhận thật.
    assert manifest["files"][0]["converted_from_doc"] is True


def _assert_swap_blocked_contract(tmp_path, out, foreign, foreign_bytes,
                                  owned_bytes, source):
    """Bất biến khi tác nhân ngoài CỐ swap staging: pin giữ inode tại
    path nên Move của tác nhân thất bại TRÊN PHÍA NÓ — conversion của ta
    chạy tiếp bình thường (review be750f7)."""
    # Foreign nguyên bytes tại path GỐC — không bao giờ tới được staging
    # để bị ghi đè, cũng không bị publish.
    assert foreign.exists() and foreign.read_bytes() == foreign_bytes
    # Staging giữ owned bytes của fake SaveAs2 và KHÔNG bị di chuyển
    # ra '.owned-backup' (Move đầu của tác nhân đã bị pin chặn).
    scratch = [p for p in out.iterdir()
               if p.name.startswith(".naia12-scratch-")]
    staged = [p for p in scratch[0].iterdir()
              if p.name.endswith(".conv.docx")]
    assert staged and staged[0].read_bytes() == owned_bytes
    assert not staged[0].with_name(
        staged[0].name + ".owned-backup").exists()
    # Output owned được publish hợp lệ; nguồn nguyên vẹn.
    published = out / "legacy.docx"
    assert published.exists()
    assert any("OWNED CONVERSION OUTPUT" in p.text
               for p in Document(str(published)).paragraphs)
    assert source.read_bytes() == b"SYNTHETIC_LEGACY_SOURCE"


def test_fake_word_prewrite_swap_blocked_by_pin(
        tmp_path, monkeypatch):
    """AC1: tác nhân ngoài cố swap staging -> FOREIGN ngay sau guard
    check, TRƯỚC SaveAs2 — đúng race review be750f7 (guard từng đóng
    handle rồi SaveAs2 ghi theo path nên foreign bị đè mà receipt vẫn
    nhận). Guard pin giờ khoá inode staging suốt dispatch: Move của tác
    nhân bị sharing-violation -> swap không xảy ra -> conversion chạy
    bình thường, output owned publish hợp lệ, bytes foreign không bị
    chạm. Probe dùng PowerShell/NTFS thật; Word là mô phỏng New-Object —
    KHÔNG phải race COM thật."""
    if os.name != "nt" or _powershell() is None:
        pytest.skip("fake-Word lifecycle probe chạy PowerShell — Windows")
    src = tmp_path / "src"
    src.mkdir(parents=True)
    source = src / "legacy.doc"
    source.write_bytes(b"SYNTHETIC_LEGACY_SOURCE")
    owned = _docx(tmp_path / "owned.docx", ["OWNED CONVERSION OUTPUT"])
    owned_bytes = owned.read_bytes()
    foreign = _docx(tmp_path / "foreign.docx", ["FOREIGN SENTINEL"])
    foreign_bytes = foreign.read_bytes()
    monkeypatch.setenv("NAIA_FAKE_OWNED", str(owned))
    monkeypatch.setenv("NAIA_FAKE_FOREIGN", str(foreign))
    monkeypatch.setenv("NAIA_FAKE_SWAP", "prewrite")
    monkeypatch.setattr(wtm, "_POWERSHELL_CONVERT", _fake_word_script())

    out = tmp_path / "out"
    wtm.migrate_templates(src, out, registry=[], no_source=[])

    _assert_swap_blocked_contract(
        tmp_path, out, foreign, foreign_bytes, owned_bytes, source)


def test_fake_word_prereceipt_swap_blocked_by_pin(
        tmp_path, monkeypatch):
    """AC1: tác nhân ngoài cố swap staging -> FOREIGN sau khi SaveAs2 ghi
    xong nhưng TRƯỚC receipt — Move bị pin chặn nên receipt vẫn đọc đúng
    bytes owned trên inode đã pin; conversion thành công, foreign không
    bị chạm (review be750f7)."""
    if os.name != "nt" or _powershell() is None:
        pytest.skip("fake-Word lifecycle probe chạy PowerShell — Windows")
    src = tmp_path / "src"
    src.mkdir(parents=True)
    source = src / "legacy.doc"
    source.write_bytes(b"SYNTHETIC_LEGACY_SOURCE")
    owned = _docx(tmp_path / "owned.docx", ["OWNED CONVERSION OUTPUT"])
    owned_bytes = owned.read_bytes()
    foreign = _docx(tmp_path / "foreign.docx", ["FOREIGN SENTINEL"])
    foreign_bytes = foreign.read_bytes()
    monkeypatch.setenv("NAIA_FAKE_OWNED", str(owned))
    monkeypatch.setenv("NAIA_FAKE_FOREIGN", str(foreign))
    monkeypatch.setenv("NAIA_FAKE_SWAP", "prereceipt")
    monkeypatch.setattr(wtm, "_POWERSHELL_CONVERT", _fake_word_script())

    out = tmp_path / "out"
    wtm.migrate_templates(src, out, registry=[], no_source=[])

    _assert_swap_blocked_contract(
        tmp_path, out, foreign, foreign_bytes, owned_bytes, source)


# ----------------------------- review NAIA-12 (repair-039a994) -------------


def test_token_key_uses_engine_semantics_not_filename_normalize():
    """``_token_key`` (phân loại/match token) đi theo ``_normalize_token``
    của engine — bỏ dấu, đ->d — trong khi ``_normalize`` giữ dấu cho
    stem/lineage. Hai ngữ nghĩa không được trộn (review 039a994: spelling
    không dấu bị xếp unmapped dù producer đã gỡ)."""
    assert wtm._token_key("[Số công chứng]") == "so cong chung"
    assert wtm._token_key("[So cong chung]") == "so cong chung"
    assert wtm._token_key("[Địa chỉ đất]") == "dia chi dat"
    # `_normalize` vẫn giữ ngữ nghĩa stem: `đ` không đổi thành `d`.
    assert wtm._normalize("ĐK 08") == "đk 08"


def test_diff_tokens_unaccented_drop_spelling_is_dropped_not_unmapped():
    """`[So cong chung]` không dấu phải xếp ``dropped`` — producer/engine
    match occurrence theo khóa bỏ dấu nên nó bị gỡ cùng `[Số công chứng]`
    (review 039a994: producer gỡ cả hai nhưng record không dấu từng là
    ``unmapped``). Token lạ ngoài mọi catalog vẫn ``unmapped``."""
    result = wtm.diff_tokens(
        ["[So cong chung]", "[Nguoi uy quyen2]", "[Trường Lạ Hoắc]"],
        registry=[], no_source=[])
    by_old = {item.old: item for item in result}
    assert by_old["[So cong chung]"].status == "dropped"
    assert by_old["[Nguoi uy quyen2]"].status == "dropped"
    assert by_old["[Trường Lạ Hoắc]"].status == "unmapped"


def test_migrate_templates_mixed_diacritic_drop_spellings_per_record(
        tmp_path):
    """Nguồn có CẢ `[Số công chứng]` và `[So cong chung]`: producer gỡ cả
    hai occurrence qua lookup chuẩn hoá — manifest phải có HAI record theo
    spelling nguồn, mỗi record dropped count=1 applied=1 (review 039a994:
    record không dấu `unmapped` nhưng applied=1 -> replay reject oan)."""
    src = tmp_path / "src"
    _docx(src / "m.docx", ["A: [Số công chứng]; B: [So cong chung]"])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=[], no_source=[])

    entry = manifest["files"][0]
    assert len(entry["tokens"]) == 2
    by_old = {t["old"]: t for t in entry["tokens"]}
    for spelling in ("[Số công chứng]", "[So cong chung]"):
        token = by_old[spelling]
        assert token["status"] == "dropped"
        assert token["count"] == 1 and token["applied"] == 1
    assert manifest["totals"]["occurrences"] == 2
    assert manifest["totals"]["applied"] == 2
    text = "\n".join(_para_texts(tmp_path / "out" / "m.docx"))
    assert "[Số công chứng]" not in text
    assert "[So cong chung]" not in text


def test_migrate_templates_chained_alias_records_final_destination(
        tmp_path):
    """Alias nối tiếp A->B->C khi B cũng là token nguồn: producer áp các
    rename theo THỨ TỰ phase trên cùng tài liệu nên occurrence `[A]` ra
    `[C]` — record phải ghi ĐÍCH CUỐI `[C]` thay alias trực tiếp `[B]`
    (review 039a994: record `[A]`->`[B]` làm replay oracle từ chối oan vì
    budget `[B]` bị occurrence nguồn `[B]` tiêu hết trước)."""
    src = tmp_path / "src"
    _docx(src / "m.docx", ["First: [B]", "Second: [A]"])

    manifest = wtm.migrate_templates(
        src, tmp_path / "out", registry=["[C]"], no_source=[],
        alias_map={"[A]": "[B]", "[B]": "[C]"})

    entry = manifest["files"][0]
    by_old = {t["old"]: t for t in entry["tokens"]}
    assert by_old["[A]"]["status"] == "renamed"
    assert by_old["[A]"]["new"] == "[C]"      # đích cuối, không phải [B]
    assert by_old["[A]"]["count"] == 1 and by_old["[A]"]["applied"] == 1
    assert by_old["[B]"]["status"] == "renamed"
    assert by_old["[B]"]["new"] == "[C]"
    assert by_old["[B]"]["count"] == 1 and by_old["[B]"]["applied"] == 1
    assert _para_texts(tmp_path / "out" / "m.docx") == [
        "First: [C]", "Second: [C]"]
