"""Regression tests cho script E2E NAIA-12 — preflight write targets,
format gate theo bằng chứng, oracle fixture độc lập.

Các test này phủ đúng phản ví dụ Leader review tại 86967d6:
- run(src, src) từng tạo dirs/report trong nguồn dù migration từ chối;
- COM mô phỏng 1 trang vs 2 trang từng cho pages_match=false nhưng ok=true;
- expected nội dung từng đi qua cùng mapping engine đang test.
"""
from __future__ import annotations

import os
import subprocess
from pathlib import Path

import pytest
from docx import Document

from scripts import naia12_e2e_migrate_export as e2e
from services import word_template_migration as wtm


def _docx(path: Path, paragraphs: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))
    return path


def _make_dir_link(link: Path, target: Path) -> bool:
    """Tạo junction (``mklink /J``, không cần admin) trên Windows, symlink
    thư mục ở nơi khác. KHÔNG dùng ``cmd /c ... >nul``: trên Linux ``>nul``
    tạo file rác tên ``nul`` và mklink không chạy, guard không được kiểm."""
    link.parent.mkdir(parents=True, exist_ok=True)
    if os.name == "nt":
        result = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(target)],
            capture_output=True)
        if result.returncode == 0:
            return True
    try:
        os.symlink(str(target), str(link), target_is_directory=True)
        return True
    except (OSError, NotImplementedError):
        return False


# ----------------------------------------------------------- preflight


class TestPreflightWriteTargets:
    def test_src_equals_dst_refused(self, tmp_path):
        src = tmp_path / "src"
        _docx(src / "mau.docx", ["x"])
        targets = [
            src / "migrated", src / "exports", src / "exports_neg",
            src / "_scratch", src / "e2e_report.json"]
        with pytest.raises(ValueError):
            e2e._preflight_write_targets(
                src, src, targets, src / "e2e_report.json")
        # Không một target nào được tạo trong nguồn.
        assert {p.name for p in src.iterdir()} == {"mau.docx"}

    def test_dst_inside_src_refused(self, tmp_path):
        src = tmp_path / "src"
        _docx(src / "mau.docx", ["x"])
        out = src / "out"
        with pytest.raises(ValueError):
            e2e._preflight_write_targets(
                src, out, [out / "migrated"], out / "e2e_report.json")
        assert not out.exists()

    def test_src_inside_dst_refused(self, tmp_path):
        src = tmp_path / "out" / "src"
        _docx(src / "mau.docx", ["x"])
        out = tmp_path / "out"
        with pytest.raises(ValueError):
            e2e._preflight_write_targets(
                src, out, [out / "migrated"], out / "e2e_report.json")
        assert not (out / "migrated").exists()

    def test_junction_escape_target_refused(self, tmp_path):
        """Write target là junction trỏ ra ngoài đích -> từ chối (Windows
        junction không cần quyền admin, khác file symlink)."""
        src = tmp_path / "src"
        _docx(src / "mau.docx", ["x"])
        out = tmp_path / "out"
        out.mkdir()
        outside = tmp_path / "outside"
        outside.mkdir()
        escape = out / "migrated"
        if not _make_dir_link(escape, outside):
            pytest.skip("môi trường không tạo được junction/symlink thư mục")
        assert escape.is_symlink() or wtm._real_path(escape) != escape
        with pytest.raises(ValueError):
            e2e._preflight_write_targets(
                src, out, [escape], out / "e2e_report.json")

    def test_existing_report_refused(self, tmp_path):
        src = tmp_path / "src"
        _docx(src / "mau.docx", ["x"])
        out = tmp_path / "out"
        out.mkdir()
        report = out / "e2e_report.json"
        report.write_text("{}", encoding="utf-8")
        with pytest.raises(FileExistsError):
            e2e._preflight_write_targets(
                src, out, [out / "migrated"], report)
        assert report.read_text(encoding="utf-8") == "{}"

    def test_report_junction_refused(self, tmp_path):
        """e2e_report.json là link (junction tới thư mục) -> từ chối,
        không ghi đè/đuổi theo link."""
        src = tmp_path / "src"
        _docx(src / "mau.docx", ["x"])
        out = tmp_path / "out"
        out.mkdir()
        target_dir = tmp_path / "elsewhere"
        target_dir.mkdir()
        report = out / "e2e_report.json"
        if not _make_dir_link(report, target_dir):
            pytest.skip("môi trường không tạo được junction/symlink thư mục")
        with pytest.raises(FileExistsError):
            e2e._preflight_write_targets(
                src, out, [out / "migrated"], report)

    def test_run_src_equals_dst_creates_nothing(self, tmp_path):
        """Phản ví dụ Leader: run(src, src) từng tạo exports/report trong
        nguồn dù migration từ chối. Giờ phải raise TRƯỚC mọi mkdir/ghi."""
        src = tmp_path / "src"
        _docx(src / "mau.docx", ["x"])
        before = {p.name for p in src.iterdir()}
        with pytest.raises(ValueError):
            e2e.run(src, src)
        assert {p.name for p in src.iterdir()} == before


# ----------------------------------------------------------- format gate


class TestFormatStatus:
    """Trạng thái format theo bằng chứng — phản ví dụ Leader: mô phỏng COM
    nguồn 1 trang / đích 2 trang -> pages_match=false nhưng ok=true."""

    def test_pages_mismatch_is_failed(self):
        assert e2e._format_status(True, 1, 2) == "failed"

    def test_pages_equal_and_signature_equal_is_passed(self):
        assert e2e._format_status(True, 2, 2) == "passed"

    def test_com_unavailable_is_not_verified(self):
        assert e2e._format_status(True, None, None) == "not_verified"
        assert e2e._format_status(True, 1, None) == "not_verified"
        assert e2e._format_status(True, None, 1) == "not_verified"

    def test_signature_diff_is_failed(self):
        assert e2e._format_status(False, None, None) == "failed"
        assert e2e._format_status(False, 2, 2) == "failed"

    def test_doc_uncompared_signature_is_not_verified(self):
        """DOC không so được signature (None) -> không mặc định passed dù
        số trang bằng nhau."""
        assert e2e._format_status(None, 3, 3) == "not_verified"
        assert e2e._format_status(None, None, None) == "not_verified"
        assert e2e._format_status(None, 3, 4) == "failed"


# ----------------------------------------------------------- fixture oracle


def test_fixture_oracle_uses_real_row_types_not_engine():
    """Oracle phải lấy nhãn/diện tích từ FIXTURE — phản ví dụ Leader:
    tài sản 3 LUC 66/CLN 12/ONT 7/NTS 3 không được in thành nhãn cứng
    ONT/CLN/NTS/LUC. Nếu oracle tái dùng mapping engine, bug engine
    cũng pass -> test này khóa expected từ dữ liệu giả lập."""
    case = e2e._synthetic_case(with_proxy=True)
    fixture = e2e._fixture_case_data(case)

    # Nhãn động [Loại đất tài sản 3.M] -> loại đất THẬT theo vị trí dòng.
    assert e2e._fixture_lookup("[Loại đất tài sản 3.1]", fixture) == (True, "LUC")
    assert e2e._fixture_lookup("[Loại đất tài sản 3.2]", fixture) == (True, "CLN")
    assert e2e._fixture_lookup("[Loại đất tài sản 3.3]", fixture) == (True, "ONT")
    assert e2e._fixture_lookup("[Loại đất tài sản 3.4]", fixture) == (True, "NTS")

    # Token vị trí [ONT 3] = diện tích DÒNG 1 (= LUC 66), không phải "đất ONT".
    assert e2e._fixture_lookup("[ONT 3]", fixture) == (True, "66")
    assert e2e._fixture_lookup("[CLN 3]", fixture) == (True, "12")
    assert e2e._fixture_lookup("[NTS 3]", fixture) == (True, "7")
    assert e2e._fixture_lookup("[LUC 3]", fixture) == (True, "3")

    # Mục đích sử dụng liệt kê động theo dòng đất thật.
    found, purpose = e2e._fixture_lookup("[Mục đích sử dụng 3]", fixture)
    assert found and purpose == "LUC: 66 m2; CLN: 12 m2; ONT: 7 m2; NTS: 3 m2"

    # Đại diện lấy từ vai trò đã lưu (không node Diagram).
    assert fixture["authorized"] == [e2e.PROXY_NAME]


# ------------------------------------------------- oracle cứng (review 81ab0ca)


def test_oracle_never_falls_back_to_engine(tmp_path):
    """Token đáng kiểm mà oracle không phủ được phải thành oracle_missing —
    KHÔNG lặng lẽ lấy giá trị engine (đúng lỗi oracle cũ: expected = cái
    engine xuất ra nên sai gì cũng pass)."""
    fixture = e2e._fixture_case_data(e2e._synthetic_case())
    found, _value = e2e._fixture_lookup("[Diện tích lạ chưa có rule]", fixture)
    assert found is False
    template = _docx(tmp_path / "m.docx", ["Diện tích: [Diện tích lạ chưa có rule] m2"])
    _patterns, missing = e2e._expected_patterns(template, fixture)
    assert missing and "[Diện tích lạ chưa có rule]" in missing[0]


def test_fixed_form_code_lines_are_still_checked(tmp_path):
    """Dòng chứa mã biểu mẫu [01] KHÔNG bị bỏ qua: literal mã giữ nguyên và
    phần dữ liệu vẫn phải khớp oracle (bản cũ skip cả dòng)."""
    fixture = e2e._fixture_case_data(e2e._synthetic_case())
    template = _docx(tmp_path / "m.docx", [
        "[01] Diện tích: [Diện tích] m2",
    ])
    patterns, missing = e2e._expected_patterns(template, fixture)
    assert not missing and len(patterns) == 1
    import re
    # [01] literal trong output + diện tích đúng tài sản 1 = 120.5.
    assert re.search(patterns[0][1], "[01] Diện tích: 120.5 m2")
    # Sai giá trị -> không khớp.
    assert not re.search(patterns[0][1], "[01] Diện tích: 999 m2")
    # Mã biểu khác số -> không khớp (literal phải nguyên).
    assert not re.search(patterns[0][1], "[02] Diện tích: 120.5 m2")


def test_wrong_area_value_rejected_by_pattern(tmp_path):
    """Phản ví dụ Leader: diện tích sai 999 từng tạo expected 999 và pass.
    Oracle fixture kỳ vọng 88 (tổng tài sản 3) — output ghi 999 phải trượt."""
    fixture = e2e._fixture_case_data(e2e._synthetic_case())
    template = _docx(tmp_path / "m.docx", [
        "Diện tích: [Diện tích 3] m2",
        "Loại đất 1: [Loại đất tài sản 3.1]",
        "[ONT 3] m2",
    ])
    patterns, missing = e2e._expected_patterns(template, fixture)
    assert not missing
    import re
    good = "Diện tích: 88 m2 Loại đất 1: LUC 66 m2"
    bad = "Diện tích: 999 m2 Loại đất 1: LUC 66 m2"
    bad_label = "Diện tích: 88 m2 Loại đất 1: ONT 999 m2"
    assert all(re.search(p, good) for _i, p in patterns)
    assert not all(re.search(p, bad) for _i, p in patterns)
    assert not all(re.search(p, bad_label) for _i, p in patterns)


def test_visible_paragraph_reader_keeps_word_controls_and_nested_ownership(tmp_path):
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    doc = Document()
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    doc.element.body.append(parse_xml(
        f'<w:p {ns}><w:r><w:t>cha</w:t><w:br/><w:t>dòng</w:t>'
        '<w:cr/><w:tab/><w:t>cuối</w:t></w:r>'
        '<w:r><w:drawing><w:txbxContent><w:p><w:r><w:t>con</w:t>'
        '<w:br/><w:t>riêng</w:t></w:r></w:p></w:txbxContent></w:drawing></w:r></w:p>'))
    for part_name, rel_type, root, singular in (
        ('footnotes', RT.FOOTNOTES, 'footnotes', 'footnote'),
        ('endnotes', RT.ENDNOTES, 'endnotes', 'endnote'),
    ):
        blob = (f'<w:{root} {ns}><w:{singular} w:id="1"><w:p><w:r>'
                f'<w:t>{part_name}</w:t><w:br/><w:tab/><w:t>text</w:t>'
                f'</w:r></w:p></w:{singular}></w:{root}>').encode()
        part = Part(PackURI(f'/word/{part_name}.xml'),
            f'application/vnd.openxmlformats-officedocument.wordprocessingml.{part_name}+xml',
            blob, doc.part.package)
        doc.part.relate_to(part, rel_type)
    doc.sections[0].header.paragraphs[0].text = 'header\nline\tvalue'
    doc.sections[0].footer.paragraphs[0].text = 'footer\nline\tvalue'
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = 'nested\nline\tvalue'
    path = tmp_path / 'visible.docx'
    doc.save(path)
    texts = e2e._doc_paragraph_texts(Document(path))
    assert 'cha\ndòng\n\tcuối' in texts
    assert 'con\nriêng' in texts
    assert not any('cha' in t and 'con' in t for t in texts)
    for name in ('header', 'footer', 'nested', 'footnotes', 'endnotes'):
        assert (f'{name}\n\ttext' if name.endswith('notes') else f'{name}\nline\tvalue') in texts
    assert len(list(Document(path).element.body.iter(qn('w:cr')))) == 1


@pytest.mark.parametrize('mutation', [None, 'person', 'value'])
def test_authorized_block_with_word_breaks_keeps_exact_oracle(tmp_path, monkeypatch, mutation):
    import json
    from docx.oxml.ns import qn

    _fast_com(monkeypatch)
    src = tmp_path / 'src'
    _docx(src / 'uq.docx', ['Đại diện: [Dòng người được ủy quyền 1]'])
    if mutation:
        real_mapping = e2e.word_engine.build_template_mapping

        def mutant_mapping(case, **kwargs):
            mapping = dict(real_mapping(case, **kwargs))
            token = '[Dòng người được ủy quyền 1]'
            old = e2e.PROXY_NAME if mutation == 'person' else '1975'
            new = 'Người Sai' if mutation == 'person' else '1999'
            # Hồ sơ phản ví dụ (không người được ủy quyền) giờ cũng build
            # mapping (W6/W8) -> chỉ đột biến hồ sơ có đại diện; hiệu lực
            # đột biến vẫn được kiểm bằng failures batch.* bên dưới.
            if old not in mapping[token]:
                assert e2e.PROXY_NAME not in mapping[token]
                return mapping
            mapping[token] = mapping[token].replace(old, new)
            return mapping

        monkeypatch.setattr(e2e.word_engine, 'build_template_mapping', mutant_mapping)
    out = tmp_path / 'out'
    assert e2e.main([str(src), str(out)]) == (1 if mutation else 0)
    report = json.loads((out / 'e2e_report.json').read_text(encoding='utf-8'))
    output = next((out / 'exports').glob('*.docx'))
    assert list(Document(output).element.body.iter(qn('w:br')))
    failures = report['failures']
    assert bool([f for f in failures if f.startswith('batch.content[')]) == bool(mutation)
    assert bool([f for f in failures if f.startswith('batch.authorized[')]) == bool(mutation)
    assert report['negative']['status'] == 'passed'


def test_uq_block_validated_at_required_slot(tmp_path):
    """`[Dòng người được ủy quyền 1]` phải in đúng block đại diện đã lưu —
    tên người khác hoặc thiếu chi tiết đều phải trượt (bản cũ chỉ tìm tên
    ở toàn file)."""
    fixture = e2e._fixture_case_data(e2e._synthetic_case())
    template = _docx(tmp_path / "m.docx", [
        "Đại diện: [Dòng người được ủy quyền 1]",
    ])
    patterns, missing = e2e._expected_patterns(template, fixture)
    assert not missing and len(patterns) == 1
    import re
    block = fixture["authorized_block"](1)
    good = "Đại diện: " + block.replace("\n", " ")
    assert re.search(patterns[0][1], good)
    # Đổi tên đại diện -> trượt; đổi quan hệ -> trượt.
    assert not re.search(patterns[0][1], good.replace("Trần Thị Ủy", "Người Lạ"))
    assert not re.search(
        patterns[0][1], good.replace("người được ủy quyền", "con trai"))


def test_canonical_tai_san_tokens_are_interesting(tmp_path):
    """Dạng canonical `[Tài sản N - <trường>]` trước đây bị oracle bỏ qua —
    giờ phải được kiểm bằng fixture (diện tích sai vẫn trượt)."""
    fixture = e2e._fixture_case_data(e2e._synthetic_case())
    assert e2e._is_interesting("[Tài sản 3 - Diện tích]")
    template = _docx(tmp_path / "m.docx", ["Diện tích: [Tài sản 3 - Diện tích] m2"])
    patterns, missing = e2e._expected_patterns(template, fixture)
    assert not missing
    import re
    assert re.search(patterns[0][1], "Diện tích: 88 m2")
    assert not re.search(patterns[0][1], "Diện tích: 999 m2")


# ------------------------------------------- mutant trên đường run()/CLI


def _fast_com(monkeypatch):
    """Bỏ Word COM khỏi run-path test — không liên quan mutant đang kiểm,
    giữ test nhanh và chạy được cả khi máy không có Word."""
    monkeypatch.setattr(
        e2e, "_word_com_stats", lambda paths, work_dir: {})


def test_mutant_wrong_values_fail_run(tmp_path, monkeypatch):
    """Mutant ghi sai nội dung (diện tích 999, loại đất sai, niêm yết sai)
    phải làm kiểm nội dung trượt: run -> report fail, CLI exit != 0.
    Chỉ patch GIÁ TRỊ engine ghi; đường migrate/batch/oracle giữ nguyên."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Thửa đất số: [Số thửa 3]",
        "+ Loại đất 1: [Loại đất tài sản 3.1]; Diện tích: [ONT 3] m2",
        "Nơi niêm yết: [Nơi niêm yết]",
    ])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[ONT 3]"] = "999"
        mapping["[Loại đất tài sản 3.1]"] = "XXX"
        mapping["[Nơi niêm yết]"] = "UBND SAI"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    assert report["ok"] is False
    content_fail = [
        f for f in report["failures"] if f.startswith("batch.content[")]
    assert content_fail, report["failures"]


def test_mutant_wrong_uq_block_fails_run(tmp_path, monkeypatch):
    """Block UQ sai người — dù tên đó xuất hiện chỗ khác trong file — phải
    trượt check slot. Phản ví dụ: bản cũ tìm tên trên toàn file nên pass."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Người kê khai: Nguyễn Văn Chủ",  # tên sai vẫn "có mặt" trong file
        "Đại diện: [Dòng người được ủy quyền 1]",
    ])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[Dòng người được ủy quyền 1]"] = (
            "1. ông Nguyễn Văn Chủ; Sinh ngày: 1960-01-01.\n"
            "Là người được ủy quyền.")
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    uq_fail = [
        f for f in report["failures"]
        if f.startswith("batch.authorized[")]
    assert uq_fail, report["failures"]


def test_mutant_wrong_context_label_fails_run(tmp_path, monkeypatch):
    """Mutant `tài sản 2` trong block anchor `tài sản 3` phải trượt check
    context — DÙ token/engine tự nhất quán (engine điền đúng giá trị tài
    sản 2 cho nhãn sai). Oracle nội dung một mình không bắt được lỗi này."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Thửa đất số: [Số thửa 3]; Tờ bản đồ số: [Số tờ 3]",
        "+ Loại đất 1: [Loại đất tài sản 2.1]; Diện tích: [ONT 2] m2",
    ])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    label_fail = [
        f for f in report["failures"]
        if f.startswith("migration.context_labels[")]
    assert label_fail, report["failures"]
    # Token/engine nhất quán với nhãn sai -> content check KHÔNG bắt được;
    # chỉ context_labels đứng ra — chứng minh check này cần thiết riêng.
    assert not [
        f for f in report["failures"] if f.startswith("batch.content[")]


def test_mixed_niem_yet_and_uq_paragraph_still_checked(tmp_path):
    """Đoạn chứa đồng thời `[Nơi niêm yết]` và token UQ không được bỏ qua
    (bản cũ skip cả đoạn khi thấy token UQ)."""
    fixture = e2e._fixture_case_data(e2e._synthetic_case())
    template = _docx(tmp_path / "m.docx", [
        "Niêm yết tại [Nơi niêm yết]; đại diện [Dòng người được ủy quyền 1]",
    ])
    patterns, missing = e2e._expected_patterns(template, fixture)
    assert not missing and len(patterns) == 1
    import re
    good = ("Niêm yết tại UBND xã Minh Tân (giá trị synthetic); đại diện "
            + fixture["authorized_block"](1).replace("\n", " "))
    assert re.search(patterns[0][1], good)
    assert not re.search(
        patterns[0][1], good.replace("Minh Tân", "SAI XÃ"))


# ------------------------------------------- review NAIA-12 (đợt e47ad1d)


def test_mutant_migrated_context_all_refs_fail(tmp_path, monkeypatch):
    """Mutant đổi CẢ anchor lẫn rows tài sản 3 -> 2 trong file MIGRATED
    (tự nhất quán trong output) vẫn phải fail — expected context lấy từ
    file NGUỒN, không suy từ output (review NAIA-12: oracle cũ đọc file
    migrated nên mutant nhất quán pass exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", [
        "Thửa đất số: [Số thửa 3]; Tờ bản đồ: [Số tờ 3]",
        "Loại đất 1: [Loại đất tài sản 3.1]; Diện tích: [ONT 3] m2",
    ])
    real_migrate = wtm.migrate_templates

    def mutant_migrate(src_dir, out_dir, **kw):
        manifest = real_migrate(src_dir, out_dir, **kw)
        for item in manifest["files"]:
            path = Path(out_dir) / item["output_name"]
            doc = Document(str(path))
            for p in doc.paragraphs:
                p.text = (p.text
                          .replace("[Số thửa 3]", "[Số thửa 2]")
                          .replace("[Số tờ 3]", "[Số tờ 2]")
                          .replace("tài sản 3.1", "tài sản 2.1")
                          .replace("[ONT 3]", "[ONT 2]"))
            doc.save(str(path))
        return manifest

    monkeypatch.setattr(wtm, "migrate_templates", mutant_migrate)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    context_fail = [
        f for f in report["failures"]
        if f.startswith("migration.context_labels[")]
    assert context_fail, report["failures"]
    # Mutant tự nhất quán -> oracle nội dung pass; CHỈ context từ nguồn
    # đứng ra bắt — chứng minh hai check phải tồn tại riêng.
    assert not [
        f for f in report["failures"] if f.startswith("batch.content[")]


def test_mutant_second_occurrence_wrong_fails_run(tmp_path, monkeypatch):
    """Chỉ đổi occurrence THỨ HAI của đoạn lặp thành 999 -> phải fail:
    expected khớp trên ĐÚNG ĐOẠN tương ứng, không search toàn văn
    (review NAIA-12: đoạn đúng đầu tiên từng chứng nhận thay cho đoạn
    lặp bị sai -> exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Diện tích: [Tài sản 3 - Diện tích] m2",
        "Diện tích: [Tài sản 3 - Diện tích] m2",
    ])
    real_export = e2e.wbe.export_batch

    def mutant_export(case, **kw):
        result = real_export(case, **kw)
        for item in result["documents"]:
            if item["status"] != "saved":
                continue
            path = Path(kw["dest_dir"]) / item["actual_filename"]
            doc = Document(str(path))
            doc.paragraphs[1].text = "Diện tích: 999 m2"
            doc.save(str(path))
        return result

    monkeypatch.setattr(e2e.wbe, "export_batch", mutant_export)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    content_fail = [
        f for f in report["failures"] if f.startswith("batch.content[")]
    assert content_fail, report["failures"]


def test_run_canonical_rows_wrong_asset_fail(tmp_path, monkeypatch):
    """Canonical `[Loại đất 2.1 - …]` dưới anchor `[Số thửa 3]` phải fail
    context — canonical row CŨNG được đối chiếu ngữ cảnh nguồn, không
    chỉ token `[ONT N]` (review NAIA-12: canonical trước đây bị bỏ kiểm)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Thửa đất số: [Số thửa 3]",
        "Loại đất 1: [Loại đất 2.1 - Loại đất]; "
        "Diện tích: [Loại đất 2.1 - Diện tích] m2",
    ])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    context_fail = [
        f for f in report["failures"]
        if f.startswith("migration.context_labels[")]
    assert context_fail, report["failures"]


def test_run_valid_uq_field_token_passes(tmp_path, monkeypatch):
    """Mẫu chỉ có `[Họ tên người được ủy quyền 1]` (trường lẻ, không phải
    full block) -> oracle kiểm GIÁ TRỊ trường, không đòi cả block -> pass
    (review NAIA-12: trường lẻ từng bị đòi full block -> fail oan)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Đại diện: [Họ tên người được ủy quyền 1]"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_run_valid_legacy_thoi_han_row(tmp_path, monkeypatch):
    """`[Thời hạn 2]` = dòng đất 2 của tài sản 1 (legacy) — dưới anchor
    tài sản 1 là hợp lệ -> pass (review NAIA-12: từng bị báo sai
    "tài sản 2")."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Thửa đất số: [Số thửa 1]",
        "Thời hạn dòng 2: [Thời hạn 2]",
    ])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_run_valid_default_canonical_tai_san(tmp_path, monkeypatch):
    """`[Tài sản - Diện tích]` (canonical không số) = tài sản 1 — token
    hợp lệ không được rơi vào oracle_missing (review NAIA-12)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [Tài sản - Diện tích] m2"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    assert not [f for f in report["failures"] if "oracle_missing" in f]


def test_report_replaced_before_final_snapshot_is_preserved(
        tmp_path, monkeypatch):
    """e2e_report.json bị thay bằng file khác trước snapshot cuối -> run
    KHÔNG unlink/ghi đè file ngoài: sentinel giữ nguyên bytes, bản fail
    ghi ở e2e_report.failed.json, run vẫn báo lỗi (review NAIA-12: report
    cạnh tranh từng bị unlink xoá mất)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Mẫu tĩnh synthetic"])
    out = tmp_path / "out"
    report_path = out / "e2e_report.json"
    sentinel = b"COMPETING REPORT - KHONG DUOC GHI DE"
    real_snapshot = e2e._snapshot_source_tree
    swapped = {"done": False}

    def swap_report_and_mutate(path):
        if report_path.exists() and not swapped["done"]:
            swapped["done"] = True
            report_path.rename(out / "run-report.json")
            report_path.write_bytes(sentinel)
            (src / "external.txt").write_text(
                "external source change", encoding="utf-8")
        return real_snapshot(path)

    monkeypatch.setattr(e2e, "_snapshot_source_tree", swap_report_and_mutate)
    assert e2e.main([str(src), str(out)]) == 1
    assert report_path.read_bytes() == sentinel       # file ngoài nguyên vẹn
    # Bản report ok ban đầu của run vẫn còn ở tên probe dự phòng.
    assert (out / "run-report.json").exists()
    failed = out / "e2e_report.failed.json"
    assert failed.exists()
    import json
    report = json.loads(failed.read_text(encoding="utf-8"))
    assert report["ok"] is False
    assert any("e2e_report.ownership" in f for f in report["failures"])
    assert any("sources_unchanged_final" in f for f in report["failures"])


def test_report_deleted_before_final_snapshot_fails_cleanly(
        tmp_path, monkeypatch):
    """Report run tạo bị XOÁ trước snapshot cuối -> path trống thì 'xb'
    ghi bản fail vào đúng tên chuẩn (không đụng file ngoài nào), run
    vẫn báo lỗi ownership + nguồn đổi."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Mẫu tĩnh synthetic"])
    out = tmp_path / "out"
    report_path = out / "e2e_report.json"
    real_snapshot = e2e._snapshot_source_tree
    swapped = {"done": False}

    def delete_report_and_mutate(path):
        if report_path.exists() and not swapped["done"]:
            swapped["done"] = True
            report_path.unlink()
            (src / "external.txt").write_text(
                "external source change", encoding="utf-8")
        return real_snapshot(path)

    monkeypatch.setattr(
        e2e, "_snapshot_source_tree", delete_report_and_mutate)
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["ok"] is False
    assert any("e2e_report.ownership" in f for f in report["failures"])
    assert any("sources_unchanged_final" in f for f in report["failures"])


def test_report_fallback_target_existing_never_overwritten(
        tmp_path, monkeypatch):
    """Report bị thay VÀ e2e_report.failed.json đã tồn tại -> 'xb' raise
    FileExistsError -> main exit 2, CẢ HAI file ngoài giữ nguyên bytes —
    không bao giờ ghi đè để báo lỗi (review NAIA-12)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Mẫu tĩnh synthetic"])
    out = tmp_path / "out"
    out.mkdir()
    old_failed = out / "e2e_report.failed.json"
    old_failed.write_bytes(b"OLD FAILED REPORT")
    report_path = out / "e2e_report.json"
    sentinel = b"COMPETING REPORT"
    real_snapshot = e2e._snapshot_source_tree
    swapped = {"done": False}

    def swap_report_and_mutate(path):
        if report_path.exists() and not swapped["done"]:
            swapped["done"] = True
            report_path.unlink()
            report_path.write_bytes(sentinel)
            (src / "external.txt").write_text(
                "external source change", encoding="utf-8")
        return real_snapshot(path)

    monkeypatch.setattr(e2e, "_snapshot_source_tree", swap_report_and_mutate)
    assert e2e.main([str(src), str(out)]) == 2
    assert report_path.read_bytes() == sentinel
    assert old_failed.read_bytes() == b"OLD FAILED REPORT"


# -------------------------------------- repair-2dddc41 --------------------


def test_report_replaced_source_unchanged_still_fails(
        tmp_path, monkeypatch):
    """Report bị thay sau capture nhưng nguồn NGUYÊN -> run vẫn phải
    ok=False + ghi e2e_report.failed.json — file ngoài giữ nguyên bytes
    (repair-2dddc41: chỉ nhánh nguồn đổi từng kiểm ownership nên report
    ngoài được trả ok=true)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Mẫu tĩnh synthetic"])
    out = tmp_path / "out"
    report_path = out / "e2e_report.json"
    sentinel = b"COMPETING REPORT - KHONG DUOC GHI DE"
    real_snapshot = e2e._snapshot_source_tree
    swapped = {"done": False}

    def swap_report_only(path):
        result = real_snapshot(path)
        if report_path.exists() and not swapped["done"]:
            swapped["done"] = True
            report_path.rename(out / "run-report.json")
            report_path.write_bytes(sentinel)
        return result

    monkeypatch.setattr(e2e, "_snapshot_source_tree", swap_report_only)
    assert e2e.main([str(src), str(out)]) == 1
    assert report_path.read_bytes() == sentinel
    assert (out / "run-report.json").exists()
    import json
    report = json.loads(
        (out / "e2e_report.failed.json").read_text(encoding="utf-8"))
    assert report["ok"] is False
    assert any("e2e_report.ownership" in f for f in report["failures"])


def test_mutant_authorized_field_suffix_fails_run(tmp_path, monkeypatch):
    """Engine nối ` SAI` sau tên UQ — oracle so TRỌN đoạn (re.fullmatch)
    phải trượt (repair-2dddc41: substring trước đây chấp nhận suffix)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Đại diện: [Họ tên người được ủy quyền 1]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[Họ tên người được ủy quyền 1]"] += " SAI"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    assert [f for f in report["failures"]
            if f.startswith("batch.content[")]


def test_run_valid_numeric_literal_period_passes(tmp_path, monkeypatch):
    """`[Tài sản - Diện tích].` xuất `120.5.` — dấu chấm câu literal đứng
    sau giá trị số là hợp lệ, không bị rào số đánh trượt (repair-2dddc41:
    `(?![\\d.])` từng gán dấu `.` câu vào giá trị -> exit 1 oan)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [Tài sản - Diện tích]."])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_canonical_row_within_same_asset_fails(
        tmp_path, monkeypatch):
    """`[Loại đất 3.1 - …]` -> `3.2` trong CÙNG tài sản — oracle đối chiếu
    tuple (tài sản, dòng, trường) của occurrence NGUỒN phải trượt
    migration.context_labels (repair-2dddc41: chỉ kiểm asset từng pass
    exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Thửa đất số: [Số thửa 3]",
        "Loại đất 1: [Loại đất 3.1 - Loại đất]; "
        "Diện tích: [Loại đất 3.1 - Diện tích] m2",
    ])
    real_migrate = wtm.migrate_templates

    def mutant_migrate(src_dir, out_dir, **kw):
        manifest = real_migrate(src_dir, out_dir, **kw)
        for item in manifest["files"]:
            path = Path(out_dir) / item["output_name"]
            doc = Document(str(path))
            for p in doc.paragraphs:
                p.text = p.text.replace("Loại đất 3.1 -", "Loại đất 3.2 -")
            doc.save(str(path))
        return manifest

    monkeypatch.setattr(wtm, "migrate_templates", mutant_migrate)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    assert [f for f in report["failures"]
            if f.startswith("migration.context_labels[")]


def test_mutant_canonical_row_only_asset_fails(tmp_path, monkeypatch):
    """Canonical `[Loại đất 3.1 - …]` đổi asset -> `2.1` — KHÔNG cần anchor
    trong đoạn: chính occurrence nguồn đủ tuple (tài sản, dòng, trường)
    để đối chiếu (repair-2dddc41: đoạn thiếu anchor từng bị bỏ kiểm ->
    exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [Loại đất 3.1 - Diện tích] m2"])
    real_migrate = wtm.migrate_templates

    def mutant_migrate(src_dir, out_dir, **kw):
        manifest = real_migrate(src_dir, out_dir, **kw)
        for item in manifest["files"]:
            path = Path(out_dir) / item["output_name"]
            doc = Document(str(path))
            for p in doc.paragraphs:
                p.text = p.text.replace("Loại đất 3.1 -", "Loại đất 2.1 -")
            doc.save(str(path))
        return manifest

    monkeypatch.setattr(wtm, "migrate_templates", mutant_migrate)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    assert [f for f in report["failures"]
            if f.startswith("migration.context_labels[")]


# -------------------------------------- repair-f35c9f6 --------------------


def test_report_swapped_right_before_rplus_open(tmp_path, monkeypatch):
    """Report còn đúng owner lúc ``stat()`` nhưng bị THAY NGAY TRƯỚC
    ``open(r+b)`` trong ``_rewrite_owned_report`` -> ``fstat`` trên fd bắt
    nhầm file ngoài -> từ chối; sentinel giữ nguyên bytes, bản fail ghi
    sang ``e2e_report.failed.json`` (repair-f35c9f6: regression khe
    stat -> r+b mà chỉ ``stat`` không đủ)."""
    import builtins
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Mẫu tĩnh synthetic"])
    out = tmp_path / "out"
    report_path = out / "e2e_report.json"
    sentinel = b"COMPETING REPORT - KHONG DUOC GHI DE"
    real_open = builtins.open
    real_snapshot = e2e._snapshot_source_tree
    swapped = {"done": False}

    def mutate_source(path):
        if report_path.exists():
            (src / "external.txt").write_text(
                "external source change", encoding="utf-8")
        return real_snapshot(path)

    def swap_before_rewrite(path, *args, **kwargs):
        mode = args[0] if args else kwargs.get("mode", "r")
        if (Path(path) == report_path and mode == "r+b"
                and not swapped["done"]):
            swapped["done"] = True
            report_path.rename(out / "run-report.json")
            report_path.write_bytes(sentinel)
        return real_open(path, *args, **kwargs)

    monkeypatch.setattr(e2e, "_snapshot_source_tree", mutate_source)
    monkeypatch.setattr(builtins, "open", swap_before_rewrite)
    assert e2e.main([str(src), str(out)]) == 1
    assert swapped["done"]
    assert report_path.read_bytes() == sentinel          # file ngoài nguyên vẹn
    assert (out / "run-report.json").exists()            # bản owned còn lại
    failed = out / "e2e_report.failed.json"
    assert failed.exists()
    import json
    report = json.loads(failed.read_text(encoding="utf-8"))
    assert report["ok"] is False
    assert any("e2e_report.ownership" in f for f in report["failures"])
    assert any("sources_unchanged_final" in f for f in report["failures"])


def _mutate_migrated_paras(monkeypatch, replacement_fn):
    """Monkeypatch migrate_templates để hậu sửa đoạn output — mô phỏng
    engine đổi/ thêm/ xoá placeholder sau migration."""
    real_migrate = wtm.migrate_templates

    def mutant_migrate(src_dir, out_dir, **kw):
        manifest = real_migrate(src_dir, out_dir, **kw)
        for item in manifest["files"]:
            path = Path(out_dir) / item["output_name"]
            doc = Document(str(path))
            for p in doc.paragraphs:
                new_text = replacement_fn(p.text)
                if new_text != p.text:
                    p.text = new_text
            doc.save(str(path))
        return manifest

    monkeypatch.setattr(wtm, "migrate_templates", mutant_migrate)


def _run_failures(out):
    import json
    report = json.loads(
        (out / "e2e_report.json").read_text(encoding="utf-8"))
    return report["failures"]


def test_alias_same_tuple_passes_context(tmp_path):
    """Canonical `[Loại đất 3.1 - Diện tích]` -> alias
    `[Diện tích loại đất tài sản 3.1]` = CÙNG tuple (tài sản, dòng,
    trường) -> không vi phạm context (repair-f35c9f6: oracle so tuple,
    không so syntax)."""
    src = _docx(tmp_path / "src" / "m.docx",
                ["Diện tích: [Loại đất 3.1 - Diện tích] m2"])
    mig = _docx(tmp_path / "mig" / "m.docx",
                ["Diện tích: [Diện tích loại đất tài sản 3.1] m2"])
    violations, stats = e2e._context_label_violations(mig, src)
    assert violations == []


@pytest.mark.parametrize("replacement", [
    "[Diện tích loại đất tài sản 2.1]",   # sai tài sản
    "[Diện tích loại đất tài sản 3.2]",   # sai dòng trong cùng tài sản
    "[Loại đất tài sản 3.1]",             # sai trường
])
def test_mutant_alias_wrong_tuple_fails(tmp_path, monkeypatch, replacement):
    """Đổi canonical sang alias cùng cụm tài sản nhưng sai tài sản/ dòng/
    trường -> occurrence nguồn (3,1,'dien tich') mất + occurrence lạ ->
    fail (repair-f35c9f6: chỉ so asset từng cho exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [Loại đất 3.1 - Diện tích] m2"])
    _mutate_migrated_paras(monkeypatch, lambda t: t.replace(
        "[Loại đất 3.1 - Diện tích]", replacement))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_ont_to_cln_fails(tmp_path, monkeypatch):
    """`[ONT 3]` -> `[CLN 3]`: cùng dạng vị trí nhưng khác dòng (slot 1 vs
    slot 2) -> tuple nguồn (3,1,'dien tich') bị thay (3,2,...) -> fail
    (repair-f35c9f6: ONT->CLN từng exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Thửa đất số: [Số thửa 3]",
        "Diện tích dòng 1: [ONT 3] m2",
    ])
    _mutate_migrated_paras(
        monkeypatch, lambda t: t.replace("[ONT 3]", "[CLN 3]"))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_extra_row_ref_fails(tmp_path, monkeypatch):
    """Thêm `[Loại đất 3.2 - Diện tích]` mà nguồn chỉ có 3.1 -> multiset
    migrated THỪA occurrence -> fail (repair-f35c9f6)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Diện tích: [Loại đất 3.1 - Diện tích] m2"])
    _mutate_migrated_paras(monkeypatch, lambda t: (
        t + "; thêm: [Loại đất 3.2 - Diện tích] m2"))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_source_ref_deleted_fails(tmp_path, monkeypatch):
    """Xoá `[Loại đất 3.1 - Diện tích]` khỏi đoạn -> occurrence nguồn
    thiếu -> fail (repair-f35c9f6: đoạn sửa không placeholder từng bị bỏ
    kiểm -> exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Diện tích: [Loại đất 3.1 - Diện tích] m2"])
    _mutate_migrated_paras(monkeypatch, lambda t: (
        "DIFFERENT_LITERAL_SENTINEL" if "3.1" in t else t))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_numeric_suffix_swallowed_by_neighbor_fails(
        tmp_path, monkeypatch):
    r"""`[ONT 1]`=80 nhưng engine xuất `80.5`: wildcard của token kế bên
    (`[Ngày]`) không được nuốt `.5` — rào số `(?!\.?\d)` bắt đuôi thập
    phân lạc (repair-f35c9f6: `[^\[\]]*?` từng nuốt `.5` -> exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT 1] [Ngày]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[ONT 1]"] = "80.5"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def test_mutant_numeric_prefix_swallowed_fails(tmp_path, monkeypatch):
    r"""`[ONT 1]`=80 nhưng engine xuất `180`: prefix số trước giá trị phải
    trượt `(?<!\d)` — giá trị 80 không được chấp `180`."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT 1]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[ONT 1]"] = "180"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def test_empty_expected_value_still_patterned(tmp_path):
    """Fixture row trống (tài sản 2 không có dòng 2) -> value "" VẪN phát
    pattern: đoạn chỉ gồm token -> expected "" -> chỉ khớp đoạn rỗng,
    output "999" trượt (repair-f35c9f6: pattern rỗng từng bị bỏ qua)."""
    import re
    fixture = e2e._fixture_case_data(e2e._synthetic_case())
    template = _docx(tmp_path / "m.docx", ["[Loại đất 2.2 - Diện tích]"])
    patterns, missing = e2e._expected_patterns(template, fixture)
    assert not missing
    assert len(patterns) == 1
    assert re.fullmatch(patterns[0][1], "")
    assert not re.fullmatch(patterns[0][1], "999")


def test_mutant_empty_expected_row_filled_fails(tmp_path, monkeypatch):
    """`[Loại đất 2.2 - Diện tích]` expected rỗng nhưng engine xuất 999 ->
    pattern rỗng vẫn kiểm -> fail (repair-f35c9f6)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["[Loại đất 2.2 - Diện tích]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[Loại đất 2.2 - Diện tích]"] = "999"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


# -------------------------------------- repair-83e6d03 --------------------


@pytest.mark.parametrize("asset", [1, 2, 3])
def test_run_valid_dk18_manifest_context_passes(
        tmp_path, monkeypatch, asset):
    """migrate_templates THẬT -> manifest schema v2 -> oracle chấp action
    literal land-label đã phê duyệt — replay đúng ba failure DK18 review
    83e6d03 gây ra vì consumer đọc trường ``kind`` không tồn tại nên bỏ
    qua ngân sách literal/token (repair-83e6d03)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / f"2. DK18, thuế {asset}.docx", [
        f"Thửa đất số: [Số thửa {asset}]",
        f"Loại đất 1: ONT; Diện tích: [ONT {asset}] m2",
    ])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_run_valid_context_token_rename_passes(tmp_path, monkeypatch):
    """`[Loại đất 2]` -> `[Loại đất tài sản 2]` là context rename được
    manifest (tokens.status=context_fix) phê duyệt cho stem DK18-2 ->
    consumer phải đọc budget ở ``tokens``, không phải ``literal_fixes``
    (review 83e6d03: rename hợp lệ bị báo thiếu/thừa ref)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 2.docx", [
        "Thửa đất số: [Số thửa 2]",
        "Loại đất: [Loại đất 2]",
    ])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_rename_budget_not_reused_across_paragraphs(tmp_path, monkeypatch):
    """Manifest ``applied=1`` chỉ chứng nhận ĐÚNG một occurrence — mutant
    chèn thêm `[Loại đất tài sản 2]` ở đoạn khác phải fail anchor-extra,
    không được xài lại budget của đoạn đầu (repair-83e6d03)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 2.docx", [
        "Thửa đất số: [Số thửa 2]",
        "Loại đất: [Loại đất 2]",
        "Ghi chú thêm",
    ])
    _mutate_migrated_paras(monkeypatch, lambda t: (
        t + "; [Loại đất tài sản 2]" if t == "Ghi chú thêm" else t))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_row_order_same_paragraph_fails(tmp_path, monkeypatch):
    """Hoán đổi `[Loại đất 3.1 - Diện tích]` <-> `[Loại đất 3.2 - Diện
    tích]` TRONG CÙNG đoạn: multiset occurrence giữ nguyên nhưng THỨ TỰ
    đổi -> fail context (repair-83e6d03: multiset check từng pass)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Diện tích: [Loại đất 3.1 - Diện tích] - "
        "[Loại đất 3.2 - Diện tích] m2",
    ])

    def swap(text):
        a = "[Loại đất 3.1 - Diện tích]"
        b = "[Loại đất 3.2 - Diện tích]"
        return text.replace(a, "\x00").replace(b, a).replace("\x00", b)

    _mutate_migrated_paras(monkeypatch, swap)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_run_valid_row_order_passes(tmp_path, monkeypatch):
    """Đối chứng của test trên: hai ref đúng thứ tự nguồn -> exit 0."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Diện tích: [Loại đất 3.1 - Diện tích] - "
        "[Loại đất 3.2 - Diện tích] m2",
    ])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_anchor_field_swap_fails(tmp_path, monkeypatch):
    """`[Số thửa 3]` -> `[Số tờ 3]`: anchor đổi TRƯỜNG cùng tài sản —
    anchor nguồn (3,'so thua') mất + anchor (3,'so to') thừa -> fail
    (repair-83e6d03: anchor chỉ so theo set asset nên pass)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Thửa đất số: [Số thửa 3]"])
    _mutate_migrated_paras(
        monkeypatch, lambda t: t.replace("[Số thửa 3]", "[Số tờ 3]"))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_anchor_extra_occurrence_fails(tmp_path, monkeypatch):
    """Thêm một `[Số thửa 3]` nữa vào cùng đoạn: anchor multiset thừa ->
    fail (repair-83e6d03: set asset không thấy occurrence dư)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Thửa đất số: [Số thửa 3]"])
    _mutate_migrated_paras(
        monkeypatch, lambda t: t + " [Số thửa 3]")
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_bare_ont_to_cln_fails(tmp_path, monkeypatch):
    """`[ONT]` trần -> `[CLN]` không được manifest phê duyệt: tuple
    (1,1,'dien tich') -> (1,2,'dien tich') là occurrence khác -> fail
    như `[ONT 3]`->`[CLN 3]` (repair-83e6d03: bare code từng pass)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT] m2"])
    _mutate_migrated_paras(
        monkeypatch, lambda t: t.replace("[ONT]", "[CLN]"))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_run_valid_bare_ont_passes(tmp_path, monkeypatch):
    """`[ONT]` trần hợp lệ = dòng 1 tài sản 1 -> exit 0 (đối chứng của
    mutant trên — token vẫn được context oracle xem là ref dòng)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT] m2"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_authorized_suffix_before_neighbor_fails(
        tmp_path, monkeypatch):
    """` SAI` nối sau `[Họ tên người được ủy quyền 1]` trước token
    neighbor `[Tên 1]`: neighbor được giải từ fixture (Nguyễn Văn Chủ)
    nên wildcard không nuốt được suffix -> batch.content fail
    (repair-83e6d03: wildcard neighbor từng chấp ` SAI`)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Đại diện: [Họ tên người được ủy quyền 1] [Tên 1]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[Họ tên người được ủy quyền 1]"] += " SAI"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def test_run_valid_authorized_neighbor_passes(tmp_path, monkeypatch):
    """Đối chứng: đoạn `[Họ tên người được ủy quyền 1] [Tên 1]` đúng ->
    neighbor slot người giải đúng `Nguyễn Văn Chủ` -> exit 0."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", [
        "Đại diện: [Họ tên người được ủy quyền 1] [Tên 1]"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_empty_expected_neighbor_filled_fails(
        tmp_path, monkeypatch):
    """Row expected RỖNG (`[Loại đất 2.2 - Diện tích]`) đi trước neighbor
    `[Ngày]`: engine ghi 999 -> neighbor ngày thật không chứng nhận thay
    cho ô trống bị điền -> batch.content fail (repair-83e6d03)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["[Loại đất 2.2 - Diện tích] [Ngày]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[Loại đất 2.2 - Diện tích]"] = "999"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def test_mutant_numeric_prefix_decimal_neighbor_fails(
        tmp_path, monkeypatch):
    r"""`[ONT 1]`=80 nhưng engine xuất `1.80` đứng SAU neighbor `[Ngày]`:
    rào số hai phía `(?<![\d.,])` chặn cả prefix có dấu `.`/`,`, neighbor
    ngày thật không bao được số lạc -> batch.content fail
    (repair-83e6d03: `(?<!\d)`+wildcard từng chấp `1.80`)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [Ngày] [ONT 1]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[ONT 1]"] = "1.80"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


# -------------------------------------- repair-1ae81e3 --------------------


def test_run_valid_adjacent_values_literal_comma_passes(
        tmp_path, monkeypatch):
    """`Diện tích: [ONT 3],[CLN 3] m2` xuất `66,12` — dấu ``,`` LITERAL
    nguồn giữa hai giá trị là phân tách, không phải phần số của giá
    trị nào -> oracle phải chấp nhận (repair-1ae81e3: rào số từng coi
    comma literal là phần numeric neighbor -> fail oan output đúng)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT 3],[CLN 3] m2"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


@pytest.mark.parametrize("bad", [
    "Diện tích: 180,12 m2",     # nối số vào giá trị đầu
    "Diện tích: 66,120 m2",     # nối số vào giá trị sau
    "Diện tích: 1.80,12 m2",    # prefix số+dấu vào giá trị đầu
    "Diện tích: 80.5,12 m2",    # đuôi thập phân lạc vào giá trị đầu
    "Diện tích: 805,12 m2",     # nối số liền giá trị đầu
    "Diện tích: 66,1.80 m2",    # suffix số+dấu vào giá trị sau
])
def test_mutant_adjacent_values_numeric_concat_fails(
        tmp_path, monkeypatch, bad):
    """Cùng pattern `[ONT 3],[CLN 3]`: comma literal cho `66,12` qua
    nhưng mọi biến thể số-bị-nối của HAI giá trị vẫn phải trượt —
    nhãn separator không mở lỗ hổng ghép số (repair-1ae81e3)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT 3],[CLN 3] m2"])
    real_export = e2e.wbe.export_batch

    def mutant_export(case, **kw):
        result = real_export(case, **kw)
        for item in result["documents"]:
            if item["status"] != "saved":
                continue
            path = Path(kw["dest_dir"]) / item["actual_filename"]
            doc = Document(str(path))
            for p in doc.paragraphs:
                if p.text == "Diện tích: 66,12 m2":
                    p.text = bad
            doc.save(str(path))
        return result

    monkeypatch.setattr(e2e.wbe, "export_batch", mutant_export)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def _swap_tokens(a, b):
    def swap(text):
        return text.replace(a, "\x00").replace(b, a).replace("\x00", b)
    return swap


def test_run_dk18_action_paragraph_positive(tmp_path, monkeypatch):
    """DK18-3 đoạn có literal label + canonical rows ĐÚNG thứ tự ->
    migrate thật -> manifest v2 -> oracle chấp action literal+token
    tại đúng vị trí -> exit 0 (đối chứng của các mutant bên dưới)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", [
        "Loại đất 1: ONT; Diện tích: [ONT 3] m2; "
        "Dòng 1: [Loại đất 3.1 - Diện tích]; "
        "Dòng 2: [Loại đất 3.2 - Diện tích]",
    ])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_row_swap_in_literal_action_paragraph_fails(
        tmp_path, monkeypatch):
    """Hoán `[Loại đất 3.1 - Diện tích]` <-> `[Loại đất 3.2 - Diện
    tích]` TRONG đoạn chứa literal action (nhãn ONT): multiset giữ
    nguyên nhưng sequence expected-vs-migrated lệch -> context_labels
    fail (review 1ae81e3: đoạn đã 'tiêu budget' từng được miễn order
    check -> mutant pass exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", [
        "Loại đất 1: ONT; Diện tích: [ONT 3] m2; "
        "Dòng 1: [Loại đất 3.1 - Diện tích]; "
        "Dòng 2: [Loại đất 3.2 - Diện tích]",
    ])
    _mutate_migrated_paras(monkeypatch, _swap_tokens(
        "[Loại đất 3.1 - Diện tích]", "[Loại đất 3.2 - Diện tích]"))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_row_swap_in_token_action_paragraph_fails(
        tmp_path, monkeypatch):
    """Cùng hoán 3.1 <-> 3.2 trong đoạn chứa TOKEN context action
    (`[Loại đất 3]` -> `[Loại đất tài sản 3]`): nhánh token budget
    cũng không được miễn order check (review 1ae81e3)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", [
        "Loại đất tổng: [Loại đất 3]; "
        "Dòng 1: [Loại đất 3.1 - Diện tích]; "
        "Dòng 2: [Loại đất 3.2 - Diện tích]",
    ])
    _mutate_migrated_paras(monkeypatch, _swap_tokens(
        "[Loại đất 3.1 - Diện tích]", "[Loại đất 3.2 - Diện tích]"))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_literal_label_relocated_to_other_paragraph_fails(
        tmp_path, monkeypatch):
    """`[Loại đất tài sản 3.1]` được manifest phê duyệt TẠI ĐÚNG span
    literal nguồn — dời sang đoạn khác (đoạn gốc trả lại literal ONT):
    đoạn gốc thiếu ref + đoạn đích thừa ref -> context_labels fail
    (review 1ae81e3: budget literal toàn file từng chứng nhận
    relocation -> mutant pass exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", [
        "Loại đất 1: ONT; Diện tích: [ONT 3] m2",
        "Ghi chú synthetic",
    ])
    label = "[Loại đất tài sản 3.1]"

    def relocate(text):
        if label in text:
            return text.replace(label, "ONT")
        if text == "Ghi chú synthetic":
            return text + "; nhãn: " + label
        return text

    _mutate_migrated_paras(monkeypatch, relocate)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_mutant_anchor_order_swap_fails(tmp_path, monkeypatch):
    """`Thửa: [Số thửa 3]; Tờ: [Số tờ 3]` hoán hai token cùng đoạn:
    multiset anchor giữ nguyên nhưng THỨ TỰ đổi -> fail context
    (review 1ae81e3: anchor chỉ so multiset nên swap pass exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Thửa: [Số thửa 3]; Tờ: [Số tờ 3]"])
    _mutate_migrated_paras(monkeypatch, _swap_tokens(
        "[Số thửa 3]", "[Số tờ 3]"))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_run_purpose_enum_with_kept_row_positive(tmp_path, monkeypatch):
    """PCDS: enum `ONT: [ONT] m2; CLN: [CLN] m2` -> `[Mục đích sử dụng
    1]` VÀ `[ONT]` đứng riêng ở đoạn sau được GIỮ — removal chỉ hạch
    toán occurrence TRONG span -> exit 0 (đối chứng mutant dưới)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "1. PCDS.docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2",
        "Diện tích riêng: [ONT] m2",
    ])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_removal_budget_reused_for_kept_row_fails(
        tmp_path, monkeypatch):
    """Cùng fixture trên nhưng mutant XOÁ luôn `[ONT]` đứng riêng:
    removal của enum chỉ chứng nhận occurrence TRONG span — occurrence
    kept ngoài span bị xoá thêm phải fail 'thiếu ref' (review 1ae81e3:
    removed_tokens + context_fix đếm đôi -> xoá thêm pass exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "1. PCDS.docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; CLN: [CLN] m2",
        "Diện tích riêng: [ONT] m2",
    ])
    _mutate_migrated_paras(
        monkeypatch, lambda t: t.replace("[ONT]", "") if "riêng" in t
        else t)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_ctx_budget_consumption_line_is_live(tmp_path):
    """Regression budget THẬT: cùng ``old_text`` ở hai đoạn, chỉ MỘT
    action producer-applied — occurrence đầu đổi ref, occurrence sau
    phải GIỮ. Mutant xoá dòng tiêu hao ``rn["remaining"] -= 1`` phải
    làm test này fail (mutation-testing của Leader: bản cũ dùng token
    không-ref nên đoạn bị skip, phép tiêu hao không bao giờ chạy ->
    mutant vô hình)."""
    src = _docx(tmp_path / "src" / "m.docx", [
        "Đoạn một: [Loại đất 1.1 - Diện tích]",
        "Ghi chú giữa",
        "Đoạn hai: [Loại đất 1.1 - Diện tích]",
    ])
    mig = _docx(tmp_path / "mig" / "m.docx", [
        "Đoạn một: [Loại đất 1.2 - Diện tích]",
        "Ghi chú giữa",
        "Đoạn hai: [Loại đất 1.1 - Diện tích]",
    ])
    entry = {
        "tokens": [{
            "old": "[Loại đất 1.1 - Diện tích]",
            "new": "[Loại đất 1.2 - Diện tích]",
            "status": "context_fix",
            "applied": 1,   # chỉ MỘT occurrence được đổi
            "count": 2,
        }],
        "literal_fixes": [],
        "source_name": "m.docx",
    }
    violations, _stats = e2e._context_label_violations(
        mig, src, manifest_entry=entry)
    assert violations == []

    # Đối chứng: occurrence thứ hai CŨNG bị đổi (budget applied=1 không
    # đủ chứng nhận) -> phải fail. Nếu dòng tiêu hao bị xoá thì cả hai
    # occurrence đều được coi là đã rename -> nhánh này pass sai.
    mig_bad = _docx(tmp_path / "migbad" / "m.docx", [
        "Đoạn một: [Loại đất 1.2 - Diện tích]",
        "Ghi chú giữa",
        "Đoạn hai: [Loại đất 1.2 - Diện tích]",
    ])
    violations_bad, _ = e2e._context_label_violations(
        mig_bad, src, manifest_entry=entry)
    assert violations_bad


# -------------------------------------- repair-be750f7 --------------------


def test_mutant_label_relocated_same_paragraph_fails(tmp_path, monkeypatch):
    """`[Loại đất tài sản 3.1]` từ literal `Loại đất 1: ONT` bị DỜI sang
    nhãn `Ghi chú:` TRONG CÙNG đoạn — dãy ref giữ nguyên thứ tự nhưng
    mốc literal trật vị trí -> context_labels fail (review be750f7:
    oracle chỉ so dãy ref nên relocation cùng đoạn từng pass exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", [
        "Loại đất 1: ONT; Ghi chú: riêng; Diện tích: [ONT 3] m2"])
    label = "[Loại đất tài sản 3.1]"

    def relocate(text):
        return text.replace(label, "ONT").replace(
            "Ghi chú: riêng", "Ghi chú: " + label)

    _mutate_migrated_paras(monkeypatch, relocate)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_run_label_same_paragraph_positive(tmp_path, monkeypatch):
    """Đối chứng của mutant trên: label ở đúng nhãn `Loại đất 1:` ->
    migrate thật -> oracle replay chấp span tại chỗ -> exit 0."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", [
        "Loại đất 1: ONT; Ghi chú: riêng; Diện tích: [ONT 3] m2"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_area_date_token_swap_fails(tmp_path, monkeypatch):
    """Hoán `Diện tích: [ONT 3]` <-> `Ngày: [Ngày]` TRONG CÙNG đoạn —
    vị trí token đổi qua literal `Ngày:` -> context_labels fail
    (review be750f7: swap cùng đoạn từng pass vì chỉ so multiset ref)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT 3]; Ngày: [Ngày]"])
    _mutate_migrated_paras(monkeypatch, _swap_tokens("[ONT 3]", "[Ngày]"))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_run_area_date_order_positive(tmp_path, monkeypatch):
    """Đối chứng: đoạn `Diện tích: [ONT 3]; Ngày: [Ngày]` đúng thứ tự ->
    exit 0."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT 3]; Ngày: [Ngày]"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_context_fixed_token_deleted_fails(tmp_path, monkeypatch):
    """Nguồn `Địa chỉ: [Địa chỉ đất]` trong DK18-3 được context-fix thành
    `[Địa chỉ đất tài sản 3]`; mutant XOÁ luôn token đó -> đoạn migrated
    không còn ref nào nhưng expected vẫn chứa ref -> context_labels fail
    (review be750f7: early-skip khi cả hai phía không ref từng bỏ qua
    hành động xoá -> pass exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", ["Địa chỉ: [Địa chỉ đất]"])
    _mutate_migrated_paras(
        monkeypatch, lambda t: t.replace("[Địa chỉ đất tài sản 3]", ""))
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out)
            if f.startswith("migration.context_labels[")]


def test_run_context_fixed_token_positive(tmp_path, monkeypatch):
    """Đối chứng: `[Địa chỉ đất]` -> `[Địa chỉ đất tài sản 3]` context-
    fix giữ nguyên -> exit 0."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "2. DK18, thuế 3.docx", ["Địa chỉ: [Địa chỉ đất]"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_run_drop_then_enum_merge_positive(tmp_path, monkeypatch):
    """PCDS `Mục đích sử dụng: ONT: [ONT] m2; [Người ủy quyền2]; CLN:
    [CLN] m2` — token giữa bị Owner GỠ TRƯỚC làm hai mục enum thành cụm
    liên tiếp -> merge `[Mục đích sử dụng 1]` HỢP LỆ: oracle replay
    rename/drop trước rồi mới dự đoán span -> exit 0 (review be750f7:
    span dự đoán trên text gốc từng từ chối oan chuỗi hợp lệ này)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "1. PCDS.docx", [
        "Mục đích sử dụng: ONT: [ONT] m2; [Người ủy quyền2]; "
        "CLN: [CLN] m2"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_mutant_numeric_space_neighbor_suffix_fails(tmp_path, monkeypatch):
    r"""`[ONT 3]`=66 nhưng engine xuất `66.5` trước neighbor `[Ghi chú]`
    cách bằng DẤU CÁCH — literal whitespace-only strip thành `''` và
    `'' in ",."` từng cho sep_after=True -> `(?!\d)` chấp `.5`; giờ sep
    chỉ đúng khi ký tự thật là `,`/`.` -> `(?![\d.,]?\d)` bắt mutant
    (review be750f7)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["Diện tích: [ONT 3] [Ghi chú]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[ONT 3]"] = "66.5"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def test_mutant_numeric_space_neighbor_prefix_fails(tmp_path, monkeypatch):
    r"""`[Ghi chú] [ONT 3]`=66 nhưng engine xuất `1.66` — prefix số lẫn
    dấu sau neighbor chỉ-cách-dấu-trắng phải trượt `(?<![\d.,])` (sep
    whitespace không còn được tính là dấu phân tách — review be750f7)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["[Ghi chú] [ONT 3]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[ONT 3]"] = "1.66"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def test_run_numeric_space_neighbors_positive(tmp_path, monkeypatch):
    """Đối chứng hai layout trên với giá trị ĐÚNG 66 — whitespace chỉ là
    khoảng trắng, không làm rào số chặt hơn/thả lỏng -> exit 0."""
    _fast_com(monkeypatch)
    for i, paras in enumerate([
            ["Diện tích: [ONT 3] [Ghi chú]"],
            ["[Ghi chú] [ONT 3]"]]):
        src = tmp_path / f"src{i}"
        _docx(src / "m.docx", paras)
        out = tmp_path / f"out{i}"
        assert e2e.main([str(src), str(out)]) == 0


def test_whitespace_slice_is_not_numeric_separator():
    """Unit boundary: `_value_pattern` với before/after chỉ-khoảng-trắng
    phải dùng rào số ĐẦY ĐỦ — `'' in ",."` không còn được coi là
    separator (review be750f7)."""
    import re
    pat = e2e._value_pattern("66", "Diện tích: ", " ")
    assert re.search(pat, "Diện tích: 66 Hồ sơ kiểm chứng NAIA-12")
    assert not re.search(pat, "Diện tích: 66.5 Hồ sơ kiểm chứng NAIA-12")
    pat = e2e._value_pattern("66", " ", "")
    assert re.search(pat, "Hồ sơ kiểm chứng NAIA-12 66")
    assert not re.search(pat, "Hồ sơ kiểm chứng NAIA-12 1.66")


# ----------------------------- review NAIA-12 (repair-039a994) -------------


def _migrate_with_alias_map(monkeypatch, alias_map):
    """Bọc ``wtm.migrate_templates`` để truyền ``alias_map`` — service hỗ
    trợ alias nối tiếp còn ``run()`` gọi mặc định, cần inject để phủ
    fixture chuỗi (giống probe review 039a994)."""
    real_migrate = wtm.migrate_templates

    def migrate(src, out, **kw):
        kw["alias_map"] = alias_map
        return real_migrate(src, out, **kw)

    monkeypatch.setattr(wtm, "migrate_templates", migrate)


def test_run_mixed_diacritic_drop_spellings_passes(tmp_path, monkeypatch):
    """`A: [Số công chứng]; B: [So cong chung]` — producer gỡ cả hai qua
    lookup chuẩn hoá; manifest ghi HAI record dropped/applied=1 theo
    spelling nguồn và replay chấp nhận (review 039a994: record không dấu
    `unmapped` -> expected sai -> exit 1 oan)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["A: [Số công chứng]; B: [So cong chung]"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0
    import json
    manifest = json.loads(
        (out / "migrated" / "migration_manifest.json")
        .read_text(encoding="utf-8"))
    by_old = {t["old"]: t for t in manifest["files"][0]["tokens"]}
    assert len(by_old) == 2
    for spelling in ("[Số công chứng]", "[So cong chung]"):
        assert by_old[spelling]["status"] == "dropped"
        assert by_old[spelling]["count"] == 1
        assert by_old[spelling]["applied"] == 1


def test_run_accented_drop_spelling_control(tmp_path, monkeypatch):
    """Đối chứng hai occurrence cùng spelling có dấu — đường drop đã đúng
    từ trước, giữ pass (review 039a994)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["A: [Số công chứng]; B: [Số công chứng]"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_run_chained_alias_niem_yet_passes(tmp_path, monkeypatch):
    """Chuỗi alias `[Legacy niêm yết]` -> `[Niêm Yết]` -> `[Nơi niêm
    yết]`: producer xuất `[Nơi niêm yết]` cho CẢ HAI occurrence — manifest
    ghi đích cuối nên replay tiêu đúng một budget mỗi occurrence nguồn
    (review 039a994: budget `[Niêm Yết]` từng cạn trước khi replay đoạn
    `[Legacy...]` -> exit 1 oan dù output đúng)."""
    _fast_com(monkeypatch)
    _migrate_with_alias_map(
        monkeypatch, {"[Legacy niêm yết]": "[Niêm Yết]"})
    src = tmp_path / "src"
    _docx(src / "m.docx",
          ["Một: [Niêm Yết]", "Hai: [Legacy niêm yết]"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0
    migrated = Document(str(out / "migrated" / "m.docx"))
    assert [p.text for p in migrated.paragraphs] == [
        "Một: [Nơi niêm yết]", "Hai: [Nơi niêm yết]"]


def test_run_direct_alias_niem_yet_control(tmp_path, monkeypatch):
    """Đối chứng alias TRỰC TIẾP cùng fixture — đã pass từ trước, giữ
    exit 0 (review 039a994)."""
    _fast_com(monkeypatch)
    _migrate_with_alias_map(
        monkeypatch, {"[Legacy niêm yết]": "[Nơi niêm yết]"})
    src = tmp_path / "src"
    _docx(src / "m.docx",
          ["Một: [Niêm Yết]", "Hai: [Legacy niêm yết]"])
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 0


def test_context_replay_chained_alias_service_level(tmp_path):
    """`First: [B]` + `Second: [A]` với alias A->B->C: manifest ghi đích
    cuối `[C]` và ``_context_label_violations`` không vi phạm — budget
    của `[B]`->`[C]` không bị replay kéo sang chứng nhận occurrence `[A]`
    (probe service-level review 039a994)."""
    src = _docx(tmp_path / "src" / "m.docx",
                ["First: [B]", "Second: [A]"])
    manifest = wtm.migrate_templates(
        tmp_path / "src", tmp_path / "out", registry=["[C]"],
        no_source=[], alias_map={"[A]": "[B]", "[B]": "[C]"})
    by_old = {t["old"]: t for t in manifest["files"][0]["tokens"]}
    assert by_old["[A]"]["new"] == "[C]"
    assert by_old["[B]"]["new"] == "[C]"
    violations, _stats = e2e._context_label_violations(
        tmp_path / "out" / "m.docx", src,
        manifest_entry=manifest["files"][0])
    assert violations == []


def test_run_numeric_dot_separator_neighbor_positive(
        tmp_path, monkeypatch):
    """`[ONT 3].[Ghi chú]` lẫn `[Ghi chú].[ONT 3]` với giá trị đúng `66`
    — dấu `.` thật là separator và neighbor `[Ghi chú]` giờ giải được từ
    fixture (`case.ghi_chu`) nên positive pass (review 039a994)."""
    _fast_com(monkeypatch)
    for i, paras in enumerate([
            ["[ONT 3].[Ghi chú]"], ["[Ghi chú].[ONT 3]"]]):
        src = tmp_path / f"src{i}"
        _docx(src / "m.docx", paras)
        out = tmp_path / f"out{i}"
        assert e2e.main([str(src), str(out)]) == 0


def test_mutant_numeric_dot_neighbor_suffix_fails(tmp_path, monkeypatch):
    r"""`[ONT 3]`=66 nhưng engine xuất `66.5` trước `.`+`[Ghi chú]` —
    neighbor có expected độc lập nên phần `.5` không bị wildcard nuốt ->
    trượt batch.content (review 039a994: wildcard từng chấp -> exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["[ONT 3].[Ghi chú]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[ONT 3]"] = "66.5"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def test_mutant_numeric_dot_neighbor_prefix_fails(tmp_path, monkeypatch):
    r"""Đảo layout `[Ghi chú].[ONT 3]` xuất `1.66` — prefix số lạc trước
    `.`+neighbor phải trượt (review 039a994: từng exit 0)."""
    _fast_com(monkeypatch)
    src = tmp_path / "src"
    _docx(src / "m.docx", ["[Ghi chú].[ONT 3]"])
    real_mapping = e2e.word_engine.build_template_mapping

    def mutant_mapping(case, **kw):
        mapping = dict(real_mapping(case, **kw))
        mapping["[ONT 3]"] = "1.66"
        return mapping

    monkeypatch.setattr(
        e2e.word_engine, "build_template_mapping", mutant_mapping)
    out = tmp_path / "out"
    assert e2e.main([str(src), str(out)]) == 1
    assert [f for f in _run_failures(out) if f.startswith("batch.content[")]


def test_separator_next_to_wildcard_neighbor_reports_oracle_missing(
        tmp_path):
    """Dấu `.` ĐỨNG MỘT MÌNH giữa giá trị số và token không-giải-được:
    hai slot hai bên separator không đủ bằng chứng độc lập (`66`+`.5`
    hay `66.5` đều khớp) -> báo ``oracle_missing`` thay vì nới rào cho
    wildcard nuốt phần số sai (review 039a994)."""
    fixture = e2e._fixture_case_data(e2e._synthetic_case())
    template = _docx(tmp_path / "m.docx", ["[ONT 3].[Ghi chú khác]"])
    _patterns, missing = e2e._expected_patterns(template, fixture)
    assert missing
