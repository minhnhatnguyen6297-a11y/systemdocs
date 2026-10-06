from pathlib import Path
import json
import sys
ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / 'notary_v2'))
from docx import Document
from docx.oxml import parse_xml
from services.word_engine import collect_placeholder_occurrences, replace_in_doc, find_unresolved_placeholders

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
inner = '<w:r><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>[Nơi niêm yết]</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'
doc = Document()
doc.element.body.append(parse_xml(f'<w:p xmlns:w="{W}" xmlns:v="urn:schemas-microsoft-com:vml">{inner}</w:p>'))
print(json.dumps({'physical_tokens': 1, 'reported_occurrences': collect_placeholder_occurrences(doc)}, ensure_ascii=False))

doc = Document()
doc.element.body.append(parse_xml(f'<w:p xmlns:w="{W}" xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:t>[Họ tên </w:t></w:r>{inner}<w:r><w:t>chủ đất]</w:t></w:r></w:p>'))
print(json.dumps({'separate_paragraph_tokens': ['[Họ tên chủ đất]', '[Nơi niêm yết]'], 'reported_before': collect_placeholder_occurrences(doc)}, ensure_ascii=False))
count = replace_in_doc(doc, {'[Họ tên chủ đất]': 'NGUYỄN VĂN A', '[Nơi niêm yết]': 'XÃ B'})
print(json.dumps({'replacements': count, 'remaining_tokens': find_unresolved_placeholders(doc), 'physical_text': [e.text for e in doc.element.body.iter(f'{{{W}}}t')]}, ensure_ascii=False))
